"""tdb_query: look inside TDB files without writing Python.

* ``query_tdb``: phase models (site ratios, constituents, atoms per formula unit, magnetic / DIS_PART type
  definitions) and parameters grouped by type (G, L0, L1, ..., TC, BMAGN, others) as written or with every FUNCTION
  inlined; every parameter of an element; parameter keys shared by several TDBs and whether they are numerically
  equal (``tdbtables.expression_table``, the comparison compare_system uses); FUNCTION definitions with their users
  and undefined references; type definitions; formula units compared across TDBs (with the proportional factor k);
  end-members without a G parameter.
* ``resolve_tdb_symbol``: where a FUNCTION is defined / referenced across TDB/.dat files (with line numbers), similar
  names (edit distance, same element with another phase tag: GMOFCC <-> GFCCMO, GHSERMO, GMOHCP ...), the SGTE unary
  equivalent (G(FCC_A1,MO) - GHSERMO, or a GFCCMO-style function) and how the referencing database uses it.

Text replies stay short (counts + first rows); ``output`` (.xlsx or .csv) gets the full tables.  Parsing uses
:mod:`calphad_tools.tdbparse`; inlining :mod:`calphad_tools.pwinline` via :class:`calphad_tools.tdbtables.Subsystem`.
"""
from __future__ import annotations

import csv
import difflib
import glob
import itertools
import math
import os
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from .pwinline import MissingRef, inline
from .tdbparse import (TDB, ExpressionError, Parameter, Phase, Piecewise, parse_piecewise, parse_tdb,
                       resolve_keyword)
from .tdbtables import Subsystem, expanded_text, expression_table, is_vv, pw_segments

MODES = ("phase", "element", "collisions", "functions", "typedefs", "formula_unit", "missing_endmembers")
TEXT_BUDGET = 3000
PTYPE_RANK = {"G": 0, "TC": 2, "BMAGN": 3}

PERIODIC = set(("H HE LI BE B C N O F NE NA MG AL SI P S CL AR K CA SC TI V CR MN FE CO NI CU ZN GA GE AS SE BR KR "
                "RB SR Y ZR NB MO TC RU RH PD AG CD IN SN SB TE I XE CS BA LA CE PR ND PM SM EU GD TB DY HO ER TM "
                "YB LU HF TA W RE OS IR PT AU HG TL PB BI PO AT RN FR RA AC TH PA U NP PU AM CM BK CF ES FM MD NO "
                "LR").split())
# phase tags used in function names (GFCCMO, GMOFCC, GLIQFE, GHSERCR) -> phase name in the SGTE unary database
TAG_PHASE = {"FCC": "FCC_A1", "BCC": "BCC_A2", "HCP": "HCP_A3", "LIQ": "LIQUID", "LIQUID": "LIQUID", "L": "LIQUID",
             "DIA": "DIAMOND_A4", "DIAM": "DIAMOND_A4", "CBCC": "CBCC_A12", "CUB": "CUB_A13", "BCT": "BCT_A5",
             "DHCP": "DHCP", "RHO": "RHOMBOHEDRAL_A7", "RHOMB": "RHOMBOHEDRAL_A7", "ORT": "ORTHORHOMBIC_A20",
             "TET": "TETRAGONAL_U", "GAS": "GAS", "SER": None, "HSER": None}


# =============================================================================================== small helpers
def _norm(name: str) -> str:
    return str(name).strip().upper().rstrip("#")


def _empty_view(tdb: TDB) -> Subsystem:
    """A Subsystem without phases: only its FUNCTION inliner (``fn``) and species helpers are used."""
    return Subsystem(tdb, Path(tdb.path).name, tdb.path, [], {})


def inline_parameter(view: Subsystem, p: Parameter):
    """Inlined segments of a parameter (sign-corrected for reordered constituents); None if it cannot be inlined."""
    try:
        return inline(pw_segments(p.expr, p.sign), view.fn, view.inline_notes)
    except (MissingRef, ExpressionError, RecursionError, ValueError, ZeroDivisionError):
        return None


def param_group(p: Parameter) -> str:
    """G (end-member), L0/L1/... (interaction), TC, BMAGN or the parameter type."""
    if p.key.ptype == "G":
        return f"L{p.key.order}" if any(len(sl) > 1 for sl in p.key.constituents) else "G"
    return p.key.ptype


def _group_rank(g: str) -> Tuple:
    if g == "G":
        return (0, 0, g)
    if re.fullmatch(r"L\d+", g):
        return (1, int(g[1:]), g)
    return (PTYPE_RANK.get(g, 4), 0, g)


def t_ranges(pw: Piecewise) -> str:
    if not pw.ranges:
        return ""
    return "-".join(f"{x:g}" for x in [pw.ranges[0][0]] + [hi for _, hi, _ in pw.ranges])


def written(pw: Piecewise) -> str:
    if len(pw.ranges) == 1:
        return pw.ranges[0][2]
    return " | ".join(f"{lo:g}-{hi:g}: {e}" for lo, hi, e in pw.ranges)


def _cut(s: str, n: int) -> str:
    s = " ".join(str(s).split())
    return s if len(s) <= n else s[: n - 1] + "…"


def species_elements(view: Subsystem, sp: str) -> set:
    return {e for e in view.species_elements(sp) if e not in ("VA", "*")}


def param_elements(view: Subsystem, p: Parameter) -> List[str]:
    return sorted({e for sl in p.key.constituents for s in sl for e in species_elements(view, s)})


def load(path: str) -> TDB:
    if Path(path).suffix.lower() == ".dat":
        raise ValueError(f"{Path(path).name} is a ChemSage .dat database; query_tdb reads TDB text - convert it "
                         "first with convert_database")
    return parse_tdb(path)


# =============================================================================================== phase models
def magnetic_of(tdb: TDB, ph: Phase) -> Optional[Tuple[float, float, str]]:
    for td in tdb.phase_typedefs(ph):
        m = re.search(r"MAGNETIC\s+(\S+)\s+(\S+)", td.text.upper())
        if td.kind == "MAGNETIC" and m:
            try:
                return float(m.group(1)), float(m.group(2)), td.code
            except ValueError:
                pass
    return None


def dis_part_of(tdb: TDB, ph: Phase) -> Optional[str]:
    for td in tdb.phase_typedefs(ph):
        m = re.search(r"DIS\w*_P\w*\s+([A-Z0-9_]+)", td.text.upper())
        if m:
            return m.group(1)
    return None


def ordered_partners(tdb: TDB) -> Dict[str, str]:
    """disordered phase -> the ordered phase that names it as DIS_PART."""
    out = {}
    for ph in tdb.phases.values():
        d = dis_part_of(tdb, ph)
        if d:
            out[d] = ph.name
    return out


def atoms_per_fu(ph: Phase) -> Tuple[float, List[int]]:
    """(sum of the site ratios of sublattices that are not vacancy-only, 1-based sublattices where VA mixes with
    other constituents - there the atoms per formula unit depend on the vacancy fraction)."""
    va_only = ph.va_only_sublattices()
    n = sum(r for i, r in enumerate(ph.site_ratios) if i not in va_only)
    mixed = [i + 1 for i, sl in enumerate(ph.constituents) if "VA" in sl and i not in va_only]
    return n, mixed


def model_text(ph: Phase) -> str:
    return "".join(f"({','.join(sl)}){r:g}" for sl, r in zip(ph.constituents, ph.site_ratios))


def phase_summary(tdb: TDB, name: str) -> Dict[str, Any]:
    ph = tdb.phases[name]
    n, mixed = atoms_per_fu(ph)
    mag = magnetic_of(tdb, ph)
    dis = dis_part_of(tdb, ph)
    partners = ordered_partners(tdb)
    tds = []
    for c in ph.type_codes:
        td = tdb.typedefs.get(c)
        if td is not None and td.kind != "SEQ":
            tds.append(f"{c}: {td.text}")
        elif td is None and c != "%":
            tds.append(f"{c}: (no TYPE_DEFINITION)")
    return {"phase": name, "model": model_text(ph), "sublattices": ph.n_sublattices,
            "site_ratios": ":".join(f"{r:g}" for r in ph.site_ratios), "sum_site_ratios": sum(ph.site_ratios),
            "atoms_per_fu": n, "va_mixed_sublattices": ",".join(map(str, mixed)),
            "sum_is_1": abs(n - 1.0) <= 1e-9,
            "magnetic": f"AFM {mag[0]:g}, p {mag[1]:g} (type {mag[2]})" if mag else "",
            "dis_part": dis or "", "disordered_part_of": partners.get(name, ""),
            "type_codes": ph.type_codes, "type_definitions": "; ".join(tds)}


def param_row(tdb: TDB, view: Subsystem, p: Parameter, expand: bool) -> Dict[str, Any]:
    row = {"file": tdb.name, "phase": p.key.phase, "group": param_group(p), "parameter": p.raw_name,
           "constituents": ":".join(",".join(sl) for sl in p.key.constituents), "order": p.key.order,
           "elements": "-".join(e.capitalize() for e in param_elements(view, p)),
           "T_ranges": t_ranges(p.expr), "expression": written(p.expr), "reference": p.reference}
    if expand:
        segs = inline_parameter(view, p)
        row["expanded"] = expanded_text(segs).replace("\n", " | ") if segs else "(cannot inline: undefined function)"
    return row


# =============================================================================================== sublattice matching
def match_sublattices(ra: Sequence[float], rb: Sequence[float], ca: Optional[Sequence[Sequence[str]]] = None,
                      cb: Optional[Sequence[Sequence[str]]] = None, rtol: float = 1e-5,
                      max_perms: int = 50000) -> List[Tuple[List[int], float, int]]:
    """Ways sublattices of model B map onto model A with proportional site ratios: [(perm, k, score)], best first.

    ``perm[i]`` is B's index for A's i-th sublattice, ``k`` = sum(B ratios) / sum(A ratios), ``score`` counts the
    sublattices whose constituent sets agree (2 = equal, 1 = A's set inside B's; when constituents are given, a
    sublattice that shares nothing excludes the permutation).  Vacancy-only sublattices are not passed here."""
    n = len(ra)
    if n != len(rb) or n == 0:
        return []
    k = sum(rb) / sum(ra)
    out = []
    for i, perm in enumerate(itertools.permutations(range(n))):
        if i >= max_perms:
            break
        if any(abs(rb[perm[t]] - k * ra[t]) > rtol * max(abs(k * ra[t]), 1e-12) for t in range(n)):
            continue
        score = 0
        if ca is not None and cb is not None:
            ok = True
            for t in range(n):
                A, B = set(ca[t]), set(cb[perm[t]])
                if A == B:
                    score += 2
                elif A <= B:
                    score += 1
                elif not (A & B):
                    ok = False
                    break
            if not ok:
                continue
        out.append((list(perm), k, score))
    out.sort(key=lambda x: (-x[2], sum(abs(a - b) for a, b in enumerate(x[0]))))
    return out


def model_relation(pa: Phase, pb: Phase) -> Tuple[str, Optional[float], Optional[List[int]]]:
    """('same' | 'permuted' | 'proportional' | 'permuted+proportional' | 'different: ...', k, perm) of B vs A,
    vacancy-only sublattices ignored."""
    ia = [i for i in range(pa.n_sublattices) if i not in pa.va_only_sublattices()]
    ib = [i for i in range(pb.n_sublattices) if i not in pb.va_only_sublattices()]
    if len(ia) != len(ib):
        return f"different: {len(ia)} vs {len(ib)} non-vacancy sublattices", None, None
    ra, rb = [pa.site_ratios[i] for i in ia], [pb.site_ratios[i] for i in ib]
    cands = match_sublattices(ra, rb, [pa.constituents[i] for i in ia], [pb.constituents[i] for i in ib])
    if not cands:
        cands = match_sublattices(ra, rb)            # proportional ratios but disjoint constituents (other system)
    if not cands:
        return "different: site ratios not proportional", None, None
    perm, k, _ = cands[0]
    perm_full = [ib[p] for p in perm]
    moved = perm_full != ia
    scaled = abs(k - 1.0) > 1e-9
    rel = ("permuted+proportional" if moved and scaled else "permuted" if moved else
           "proportional" if scaled else "same")
    return rel, k, perm_full


# =============================================================================================== output tables
def write_tables(path: str, sheets: Dict[str, List[Dict[str, Any]]]) -> List[str]:
    """Write tables to one .xlsx (a sheet each) or .csv file(s) (``<stem>_<sheet>.csv`` when there are several)."""
    sheets = {k: v for k, v in sheets.items() if v}
    if not sheets:
        return []
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    if p.suffix.lower() == ".csv":
        files = []
        for name, rows in sheets.items():
            fp = p if len(sheets) == 1 else p.with_name(f"{p.stem}_{name}.csv")
            cols = list(dict.fromkeys(k for r in rows for k in r))
            with open(fp, "w", newline="") as fh:
                w = csv.DictWriter(fh, fieldnames=cols)
                w.writeheader()
                w.writerows(rows)
            files.append(str(fp))
        return files
    if p.suffix.lower() != ".xlsx":
        p = p.with_suffix(".xlsx")
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", name)[:31])
        cols = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([_cell(r.get(c)) for c in cols])
        for i, c in enumerate(cols, 1):
            width = max([len(str(c))] + [min(len(str(r.get(c, ""))), 80) for r in rows[:200]])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(8, width + 2), 80)
        ws.freeze_panes = "A2"
    wb.save(p)
    return [str(p)]


def _cell(v):
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


class Text:
    """Line collector that stops at a character budget and counts what it left out."""

    def __init__(self, budget: int = TEXT_BUDGET):
        self.lines: List[str] = []
        self.budget = budget
        self.dropped = 0

    def add(self, line: str = "") -> bool:
        if sum(len(x) + 1 for x in self.lines) + len(line) > self.budget:
            self.dropped += 1
            return False
        self.lines.append(line)
        return True

    def text(self) -> str:
        s = "\n".join(self.lines)
        if self.dropped:
            s += f"\n... ({self.dropped} more line(s) not shown; give output='name.xlsx' for the full tables)"
        return s


def _table_lines(T: Text, rows: List[Dict[str, Any]], cols: Sequence[str], max_rows: int, widths: Dict[str, int]):
    for r in rows[:max_rows]:
        T.add("  " + "  ".join(_cut(r.get(c, ""), widths.get(c, 40)) for c in cols if r.get(c, "") != ""))
    if len(rows) > max_rows:
        T.add(f"  ... {len(rows) - max_rows} more row(s)")


# =============================================================================================== modes
def _phase_filter(tdb: TDB, phases: Optional[Sequence[str]]) -> List[str]:
    if not phases:
        return list(tdb.phases)
    want = [_norm(p) for p in phases]
    missing = [p for p in want if p not in tdb.phases]
    if missing and len(missing) == len(want):
        close = {p: difflib.get_close_matches(p, list(tdb.phases), n=3) for p in missing}
        raise ValueError(f"{tdb.name}: no phase {', '.join(missing)}; similar: {close}; phases: {sorted(tdb.phases)}")
    return [p for p in want if p in tdb.phases]


def _element_ok(view: Subsystem, p: Parameter, elements: Optional[Sequence[str]], mode: str = "within") -> bool:
    if not elements:
        return True
    els = set(param_elements(view, p))
    want = {e.upper() for e in elements}
    return want <= els if mode == "all" else els <= want


def mode_phase(tdb: TDB, phases, elements, expand: bool, group_by: str, max_rows: int, T: Text):
    view = _empty_view(tdb)
    names = _phase_filter(tdb, phases)
    models, params = [], []
    for nm in names:
        s = phase_summary(tdb, nm)
        models.append({"file": tdb.name, **s})
        rows = [param_row(tdb, view, p, expand) for p in tdb.parameters
                if p.key.phase == nm and _element_ok(view, p, elements)]
        rows.sort(key=lambda r: (_group_rank(r["group"]), r["constituents"].count(",") + r["constituents"].count(":"),
                                 r["constituents"], r["order"]))
        params += rows
    T.add(f"{tdb.name}: {len(names)} phase(s), {len(params)} parameter(s)"
          + (f" within {'-'.join(e.capitalize() for e in elements)}" if elements else ""))
    show = names if len(names) <= 8 else names[:8]
    for s in [m for m in models if m["phase"] in show]:
        apfu = f"{s['atoms_per_fu']:g} atoms/fu" + (f" (VA mixes on sublattice {s['va_mixed_sublattices']})"
                                                     if s["va_mixed_sublattices"] else "")
        extra = "; ".join(x for x in (s["magnetic"] and f"magnetic {s['magnetic']}",
                                      s["dis_part"] and f"DIS_PART {s['dis_part']}",
                                      s["disordered_part_of"] and f"disordered part of {s['disordered_part_of']}")
                          if x)
        T.add(f"\n{s['phase']} {s['model']}: site ratios {s['site_ratios']}, sum {s['sum_site_ratios']:g}, "
              f"{apfu}" + (f"; {extra}" if extra else ""))
        if s["type_definitions"]:
            T.add(f"  type definitions: {_cut(s['type_definitions'], 160)}")
        rows = [r for r in params if r["phase"] == s["phase"]]
        counts: Dict[str, int] = {}
        for r in rows:
            counts[r["group"]] = counts.get(r["group"], 0) + 1
        T.add("  parameters: " + (", ".join(f"{g} {n}" for g, n in sorted(counts.items(), key=lambda kv:
                                                                           _group_rank(kv[0]))) or "none"))
        if group_by == "type":
            rows = sorted(rows, key=lambda r: (_group_rank(r["group"]), r["constituents"]))
        cols = ["group", "parameter", "T_ranges", "expanded" if expand else "expression", "reference"]
        per = max(3, max_rows // max(1, len(show)))
        _table_lines(T, rows, cols, per, {"parameter": 34, "expression": 60, "expanded": 70, "reference": 12})
    if len(names) > len(show):
        T.add(f"\n... {len(names) - len(show)} more phase(s): {', '.join(names[len(show):][:20])}")
    return {"models": models, "parameters": params}


def mode_element(tdb: TDB, phases, elements, expand: bool, group_by: str, max_rows: int, T: Text):
    if not elements:
        raise ValueError("mode 'element' needs elements, e.g. ['MO'] (every parameter whose constituents include "
                         "all of them)")
    view = _empty_view(tdb)
    names = set(_phase_filter(tdb, phases))
    rows = [param_row(tdb, view, p, expand) for p in tdb.parameters
            if p.key.phase in names and _element_ok(view, p, elements, "all")]
    key = {"type": lambda r: (_group_rank(r["group"]), r["phase"], r["constituents"]),
           "elements": lambda r: (r["elements"].count("-"), r["elements"], r["phase"], _group_rank(r["group"]))}.get(
        group_by, lambda r: (r["phase"] != "LIQUID", r["phase"], _group_rank(r["group"]), r["constituents"]))
    rows.sort(key=key)
    by_phase: Dict[str, int] = {}
    systems: Dict[str, int] = {}
    for r in rows:
        by_phase[r["phase"]] = by_phase.get(r["phase"], 0) + 1
        systems[r["elements"]] = systems.get(r["elements"], 0) + 1
    T.add(f"{tdb.name}: {len(rows)} parameter(s) involving {'+'.join(e.capitalize() for e in elements)} in "
          f"{len(by_phase)} phase(s)")
    T.add("  by phase: " + ", ".join(f"{p} {n}" for p, n in by_phase.items()))
    T.add("  by subsystem: " + ", ".join(f"{s} {n}" for s, n in sorted(systems.items(), key=lambda kv:
                                                                          (kv[0].count("-"), kv[0]))))
    cols = ["phase", "group", "parameter", "T_ranges", "expanded" if expand else "expression", "reference"]
    _table_lines(T, rows, cols, max_rows, {"parameter": 34, "expression": 55, "expanded": 65, "reference": 12})
    return {"parameters": rows}


def mode_functions(tdb: TDB, phases, elements, expand: bool, max_rows: int, T: Text):
    view = _empty_view(tdb)
    names = set(_phase_filter(tdb, phases)) if phases else None
    params = [p for p in tdb.parameters if (names is None or p.key.phase in names) and _element_ok(view, p, elements)]
    users_p: Dict[str, List[str]] = {}
    users_f: Dict[str, List[str]] = {}
    for p in params:
        for s in p.expr.symbols():
            users_p.setdefault(s, []).append(p.raw_name)
    for f in tdb.functions.values():
        for s in f.expr.symbols():
            users_f.setdefault(s, []).append(f.name)
    # functions reachable from the selected parameters (all functions when nothing is filtered)
    reach: List[str] = []
    todo = [s for p in params for s in p.expr.symbols()]
    while todo:
        s = todo.pop(0)
        if s in reach:
            continue
        reach.append(s)
        if s in tdb.functions:
            todo += tdb.functions[s].expr.symbols()
    filtered = names is not None or bool(elements)
    fnames = [f for f in tdb.functions if f in reach] if filtered else list(tdb.functions)
    rows = []
    for nm in fnames:
        f = tdb.functions[nm]
        row = {"file": tdb.name, "function": nm, "T_ranges": t_ranges(f.expr), "expression": written(f.expr),
               "uses": ",".join(f.expr.symbols()), "used_by_parameters": len(users_p.get(nm, [])),
               "used_by_functions": ",".join(users_f.get(nm, [])),
               "parameters": "; ".join(users_p.get(nm, [])[:40]),
               "unused": not users_p.get(nm) and not users_f.get(nm) and nm not in reach}
        if expand:
            try:
                row["expanded"] = expanded_text(view.fn(nm)).replace("\n", " | ")
            except (MissingRef, ExpressionError, RecursionError, ValueError) as e:
                row["expanded"] = f"(cannot inline: {e})"
        rows.append(row)
    referenced = set(users_p) | (set(users_f) if not filtered else {s for s in reach})
    undefined = []
    for s in sorted(referenced):
        if s in tdb.functions or s == "R":
            continue
        undefined.append({"file": tdb.name, "symbol": s, "referenced_by": "; ".join(
            (users_p.get(s, []) + [f"FUNCTION {f}" for f in users_f.get(s, [])])[:30])})
    unused = [r["function"] for r in rows if r["unused"]]
    n_vv = sum(1 for r in rows if is_vv(r["function"]))
    T.add(f"{tdb.name}: {len(rows)} function(s)" + (f" reachable from {len(params)} selected parameter(s)"
                                                     if filtered else "")
          + (f" ({n_vv} VV#### optimisation handles)" if n_vv else "")
          + f"; {len(undefined)} undefined reference(s); {len(unused)} unused")
    for u in undefined[:10]:
        T.add(f"  UNDEFINED {u['symbol']} <- {_cut(u['referenced_by'], 120)}")
    if unused:
        T.add(f"  unused: {_cut(', '.join(unused), 200)}")
    shown = sorted(rows, key=lambda r: -r["used_by_parameters"])
    for r in shown[:max_rows]:
        users = f"{r['used_by_parameters']} param(s)" + (f", fn {r['used_by_functions']}" if r["used_by_functions"]
                                                         else "")
        T.add(f"  {r['function']} [{r['T_ranges']}] used by {_cut(users, 60)}: "
              f"{_cut(r['expanded'] if expand else r['expression'], 80)}")
    if len(rows) > max_rows:
        T.add(f"  ... {len(rows) - max_rows} more function(s)")
    return {"functions": rows, "undefined": undefined}


def mode_typedefs(tdb: TDB, phases, max_rows: int, T: Text):
    names = _phase_filter(tdb, phases)
    users: Dict[str, List[str]] = {}
    for nm in names:
        for c in tdb.phases[nm].type_codes:
            users.setdefault(c, []).append(nm)
    partners = ordered_partners(tdb)
    rows = []
    for code, td in tdb.typedefs.items():
        if td.kind == "SEQ":
            continue
        tp = td.target_phase
        mag = re.search(r"MAGNETIC\s+(\S+)\s+(\S+)", td.text.upper())
        dis = re.search(r"DIS\w*_P\w*\s+([A-Z0-9_]+)", td.text.upper())
        problems = []
        if tp and tp != "@" and tp not in tdb.phases:
            problems.append(f"amends undefined phase {tp}")
        if tp and tp != "@" and users.get(code) and tp not in users[code]:
            problems.append(f"written for {tp} but used by {','.join(users[code])}")
        if dis and dis.group(1) not in tdb.phases:
            problems.append(f"DIS_PART {dis.group(1)} undefined")
        rows.append({"file": tdb.name, "code": code, "kind": td.kind, "target": tp or "",
                     "used_by": ",".join(users.get(code, [])),
                     "afm_factor": float(mag.group(1)) if mag else "", "structure_factor": float(mag.group(2))
                     if mag else "", "dis_part": dis.group(1) if dis else "", "text": td.text,
                     "problems": "; ".join(problems)})
    undefined = sorted({(c, nm) for c, us in users.items() for nm in us if c != "%" and c not in tdb.typedefs})
    phase_rows = []
    for nm in names:
        s = phase_summary(tdb, nm)
        if s["magnetic"] or s["dis_part"] or s["disordered_part_of"] or s["type_definitions"]:
            phase_rows.append({"file": tdb.name, "phase": nm, "magnetic": s["magnetic"], "dis_part": s["dis_part"],
                               "disordered_part_of": s["disordered_part_of"], "type_codes": s["type_codes"]})
    T.add(f"{tdb.name}: {len(rows)} TYPE_DEFINITION(s) (SEQ omitted); {len(phase_rows)} phase(s) with magnetic / "
          f"order-disorder settings; {sum(1 for r in rows if r['problems'])} with problems; "
          f"{len(undefined)} undefined code use(s)")
    for r in phase_rows[:max_rows]:
        T.add(f"  {r['phase']}: " + "; ".join(x for x in (r["magnetic"] and f"magnetic {r['magnetic']}",
                                                          r["dis_part"] and f"DIS_PART {r['dis_part']}",
                                                          r["disordered_part_of"] and
                                                          f"disordered part of {r['disordered_part_of']}") if x)
                  or f"  {r['phase']}: type codes {r['type_codes']}")
    for r in rows:
        if r["problems"] or not r["used_by"]:
            T.add(f"  type {r['code']} ({_cut(r['text'], 60)}): " + (r["problems"] or "not used by any phase"))
    for c, nm in undefined[:10]:
        T.add(f"  phase {nm} uses type code {c!r} with no TYPE_DEFINITION")
    return {"typedefs": rows, "phases": phase_rows,
            "undefined_codes": [{"file": tdb.name, "code": c, "phase": nm} for c, nm in undefined]}


def mode_missing(tdb: TDB, phases, elements, add_constituents, max_rows: int, T: Text):
    if not phases:
        raise ValueError("mode 'missing_endmembers' needs phases")
    view = _empty_view(tdb)
    rows = []
    for nm in _phase_filter(tdb, phases):
        ph = tdb.phases[nm]
        cons = [list(sl) for sl in ph.constituents]
        added = []
        for k, sps in (add_constituents or {}).items():
            i = int(k)
            if not 0 <= i < len(cons):
                raise ValueError(f"add_constituents: {nm} has sublattices 0..{len(cons) - 1} (0-based), got {k}")
            for s in sps:
                if _norm(s) not in cons[i]:
                    cons[i].append(_norm(s))
                    added.append(f"{_norm(s)}@{i}")
        if elements:
            want = {e.upper() for e in elements}
            cons = [[s for s in sl if s == "VA" or species_elements(view, s) <= want] for sl in cons]
        gkeys = [p.key.constituents for p in tdb.parameters if p.key.phase == nm and p.key.ptype == "G"
                 and all(len(sl) == 1 for sl in p.key.constituents)]
        total = math.prod(len(sl) for sl in cons)
        if total > 20000:
            raise ValueError(f"{nm}: {total} end-members; restrict with elements")
        missing = []
        for em in itertools.product(*cons):
            if any(len(k) == len(em) and all(c[0] in ("*", s) for c, s in zip(k, em)) for k in gkeys):
                continue
            missing.append(em)
        for em in missing:
            rows.append({"file": tdb.name, "phase": nm, "end_member": ":".join(em),
                         "parameter": f"G({nm},{':'.join(em)};0)",
                         "note": "all vacancies" if all(s == "VA" for s in em) else
                         ("uses added constituent" if any(f"{s}@{i}" in added for i, s in enumerate(em)) else "")})
        T.add(f"{tdb.name} {nm} {''.join('(' + ','.join(sl) + ')' for sl in cons)}"
              + (f" (+{', '.join(added)})" if added else "")
              + f": {total} end-member(s), {total - len(missing)} with G, {len(missing)} missing")
        for r in [r for r in rows if r["phase"] == nm][:max_rows]:
            T.add(f"  missing {r['parameter']}" + (f"  ({r['note']})" if r["note"] else ""))
        if len(missing) > max_rows:
            T.add(f"  ... {len(missing) - max_rows} more")
    return {"missing_endmembers": rows}


def mode_formula_unit(tdbs: List[TDB], phases, max_rows: int, T: Text):
    names: List[str] = []
    for t in tdbs:
        names += [p for p in t.phases if p not in names]
    if phases:
        want = [_norm(p) for p in phases]
        if not any(p in names for p in want):
            raise ValueError(f"none of {want} is a phase of the given TDB(s)")
        names = [p for p in want if p in names]
    rows, comp = [], []
    for nm in names:
        have = [t for t in tdbs if nm in t.phases]
        for t in have:
            s = phase_summary(t, nm)
            rows.append({"phase": nm, "file": t.name, "model": s["model"], "site_ratios": s["site_ratios"],
                         "atoms_per_fu": s["atoms_per_fu"], "sum_is_1": s["sum_is_1"],
                         "va_mixed_sublattices": s["va_mixed_sublattices"]})
        if len(have) > 1:
            a = have[0]
            for b in have[1:]:
                rel, k, perm = model_relation(a.phases[nm], b.phases[nm])
                comp.append({"phase": nm, "A": a.name, "B": b.name, "relation": rel,
                             "k (B sites / A sites)": k if k is not None else "",
                             "B_sublattices": ",".join(map(str, perm)) if perm else "",
                             "A model": model_text(a.phases[nm]), "B model": model_text(b.phases[nm])})
    n1 = sum(1 for r in rows if r["sum_is_1"])
    T.add(f"{len(names)} phase(s) in {len(tdbs)} TDB(s): {n1}/{len(rows)} phase models with exactly 1 atom per "
          "formula unit (vacancy-only sublattices excluded)")
    for r in rows[:max_rows if len(tdbs) == 1 else max_rows // 2]:
        T.add(f"  {r['phase']} [{r['file']}] {r['model']}: {r['atoms_per_fu']:g} atoms/fu"
              + (" (=1)" if r["sum_is_1"] else "")
              + (f", VA mixes on sublattice {r['va_mixed_sublattices']}" if r["va_mixed_sublattices"] else ""))
    if comp:
        T.add("across TDBs (B vs A):")
        for c in comp[:max_rows]:
            k = c["k (B sites / A sites)"]
            T.add(f"  {c['phase']}: {c['B']} vs {c['A']}: {c['relation']}" + (f", k = {k:.6g}" if k != "" else "")
                  + (f", B's sublattices {c['B_sublattices']} for A's" if c["relation"].startswith("permuted")
                     else ""))
    elif len(rows) > max_rows:
        T.add(f"  ... {len(rows) - max_rows} more")
    return {"formula_units": rows, "comparison": comp}


def mode_collisions(tdbs: List[TDB], phases, elements, max_rows: int, T: Text, rtol: float = 1e-6):
    if len(tdbs) < 2:
        raise ValueError("mode 'collisions' needs two or more TDBs")
    labels = []
    for t in tdbs:
        lab = Path(t.path).stem
        labels.append(lab if lab not in labels else f"{lab}_{len(labels)}")
    subs = []
    for t, lab in zip(tdbs, labels):
        els = [e.upper() for e in elements] if elements else t.real_elements
        subs.append(Subsystem(t, lab, t.path, els, {}))
    table = expression_table(subs, phases=[_norm(p) for p in phases] if phases else None, rtol=rtol)
    rows = []
    for er in table:
        have = [lab for lab, w in er.cells.items() if w is not None]
        if len(have) < 2:
            continue
        same_text = len({er.cells[lab].canon for lab in have}) == 1
        if er.cell_differs:          # numbers compared (all present cells could be inlined)
            status = ("differs" if any(er.cell_differs.get(lab) for lab in have[1:]) else
                      "identical" if same_text and er.status != "differs" else "numerically equal")
        else:
            status = "identical" if same_text and er.status == "same" else "differs (not evaluated)"
        row = {"key": er.key, "phase": er.phase, "parameter": er.label, "in": ",".join(have), "status": status,
               "max_abs_diff": er.max_diff if math.isfinite(er.max_diff) else "", "note": er.note}
        for lab in labels:
            row[lab] = er.shown.get(lab) or ""
        rows.append(row)
    # phases defined in several TDBs: same / proportional / different models
    prow = []
    for nm in sorted({p for t in tdbs for p in t.phases}):
        have = [t for t in tdbs if nm in t.phases]
        if len(have) < 2 or (phases and nm not in {_norm(p) for p in phases}):
            continue
        a = have[0]
        for b in have[1:]:
            rel, k, _ = model_relation(a.phases[nm], b.phases[nm])
            prow.append({"phase": nm, "A": a.name, "B": b.name, "A model": model_text(a.phases[nm]),
                         "B model": model_text(b.phases[nm]), "relation": rel, "k": k if k is not None else ""})
    # functions with the same name
    frow = []
    views = {lab: _empty_view(t) for t, lab in zip(tdbs, labels)}
    import numpy as np
    from .tdbtables import segs_close
    grid = np.linspace(298.15, 3000.0, 121)
    for fn in sorted({f for t in tdbs for f in t.functions if not is_vv(f)}):
        have = [(t, lab) for t, lab in zip(tdbs, labels) if fn in t.functions]
        if len(have) < 2:
            continue
        (t0, l0), worst, status = have[0], 0.0, "identical"
        c0 = t0.functions[fn].expr.canonical()
        for t, lab in have[1:]:
            if t.functions[fn].expr.canonical() == c0:
                continue
            try:
                ok, d = segs_close(views[l0].fn(fn), views[lab].fn(fn), grid, rtol, 1e-6)
            except (MissingRef, ExpressionError, RecursionError, ValueError):
                ok, d = False, math.nan
            worst = max(worst, d) if math.isfinite(d) else worst
            status = "differs" if not ok else ("numerically equal" if status == "identical" else status)
        frow.append({"function": fn, "in": ",".join(lab for _, lab in have), "status": status,
                     "max_abs_diff": worst if status == "differs" else ""})
    nd = sum(1 for r in rows if r["status"].startswith("differs"))
    T.add(f"{len(tdbs)} TDBs ({', '.join(labels)})" + (f", elements {'-'.join(elements)}" if elements else "")
          + f": {len(rows)} parameter key(s) in more than one: "
          f"{sum(1 for r in rows if r['status'] == 'identical')} identical, "
          f"{sum(1 for r in rows if r['status'].startswith('numerically'))} "
          f"numerically equal, {nd} differ")
    if prow:
        T.add(f"phases in several TDBs ({len(prow)}):")
        for r in prow[:max_rows]:
            T.add(f"  {r['phase']}: {r['A model']} [{r['A']}] vs {r['B model']} [{r['B']}]: {r['relation']}"
                  + (f", k = {r['k']:.6g}" if r["k"] not in ("", 1.0) else ""))
    if frow:
        fd = [r for r in frow if r["status"] == "differs"]
        T.add(f"functions with the same name ({len(frow)}): {len(fd)} differ"
              + (": " + ", ".join(f"{r['function']} (max |d| {r['max_abs_diff']:.3g})" for r in fd[:8]) if fd
                 else ""))
    shown = sorted(rows, key=lambda r: (not r["status"].startswith("differs"), r["phase"]))
    T.add("parameter collisions (differences first):")
    for r in shown[:max_rows]:
        d = f", max |d| {r['max_abs_diff']:.4g}" if r["max_abs_diff"] != "" and r["status"].startswith("differs") \
            else ""
        T.add(f"  {r['parameter']} in {r['in']}: {r['status']}{d}")
    if len(rows) > max_rows:
        T.add(f"  ... {len(rows) - max_rows} more")
    return {"collisions": rows, "phases": prow, "functions": frow}


# =============================================================================================== query entry point
def query(tdbs: Sequence[str], mode: str = "phase", phases: Optional[Sequence[str]] = None,
          elements: Optional[Sequence[str]] = None, expand_functions: bool = False, group_by: str = "phase",
          output: Optional[str] = None, max_rows: int = 20, add_constituents: Optional[Dict] = None,
          rel=lambda p: p) -> Tuple[str, Dict[str, List[Dict]], List[str]]:
    """(text, tables, files) for one query_tdb call; ``tdbs`` are resolved paths."""
    mode = (mode or "phase").lower()
    if mode not in MODES:
        raise ValueError(f"mode must be one of {MODES}")
    phases = [phases] if isinstance(phases, str) else phases
    elements = [elements] if isinstance(elements, str) else elements
    elements = [e.upper() for e in elements] if elements else None
    parsed = [load(p) for p in tdbs]
    T = Text()
    tables: Dict[str, List[Dict]] = {}

    def merge(d):
        for k, v in d.items():
            tables.setdefault(k, []).extend(v)

    if mode == "collisions":
        merge(mode_collisions(parsed, phases, elements, max_rows, T))
    elif mode == "formula_unit":
        merge(mode_formula_unit(parsed, phases, max_rows, T))
    else:
        for i, t in enumerate(parsed):
            if i:
                T.add("")
            if mode == "phase":
                merge(mode_phase(t, phases, elements, expand_functions, group_by, max_rows, T))
            elif mode == "element":
                merge(mode_element(t, phases, elements, expand_functions, group_by, max_rows, T))
            elif mode == "functions":
                merge(mode_functions(t, phases, elements, expand_functions, max_rows, T))
            elif mode == "typedefs":
                merge(mode_typedefs(t, phases, max_rows, T))
            elif mode == "missing_endmembers":
                merge(mode_missing(t, phases, elements, add_constituents, max_rows, T))
    files = write_tables(output, tables) if output else []
    text = T.text()
    if files:
        text += "\nFiles:\n" + "\n".join(rel(f) for f in files)
    return text, tables, files


# =============================================================================================== symbol resolution
def scan_commands(text: str) -> List[Tuple[int, str]]:
    """(1-based line where the command starts, command text) for every '!'-terminated command; $-comments are
    dropped the way tdbparse drops them."""
    out: List[Tuple[int, str]] = []
    cur: List[str] = []
    start: Optional[int] = None
    for i, ln in enumerate(text.replace("\r\n", "\n").replace("\r", "\n").split("\n"), 1):
        if ln.strip().startswith("$"):
            continue
        if "$" in ln and ln.count("'") % 2 == 0 and "'" not in ln.split("$", 1)[0]:
            ln = ln.split("$", 1)[0]
        parts = ln.split("!")
        for j, part in enumerate(parts):
            if part.strip() and start is None:
                start = i
            cur.append(part)
            if j < len(parts) - 1:
                cmd = "\n".join(cur).strip()
                if cmd:
                    out.append((start or i, cmd))
                cur, start = [], None
    return out


def _levenshtein(a: str, b: str) -> int:
    if a == b:
        return 0
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def _el(s: str) -> Optional[str]:
    if s in PERIODIC:
        return s
    if len(s) == 2 and s[0] == s[1] and s[0] in PERIODIC:     # GHSERWW, GHSERVV, GHSERCC
        return s[0]
    return None


def parse_symbol(name: str) -> Optional[Tuple[str, str, str]]:
    """(element, phase tag, form) of a lattice-stability style name: GHSERMO -> (MO, SER, GHSER{EL}),
    GMOFCC -> (MO, FCC, G{EL}{TAG}), GFCCMO -> (MO, FCC, G{TAG}{EL}); None if it does not look like one."""
    n = _norm(name)
    m = re.fullmatch(r"GHSER([A-Z]{1,2})", n)
    if m and _el(m.group(1)):
        return _el(m.group(1)), "SER", "GHSER{EL}"
    if not n.startswith("G") or len(n) < 3:
        return None
    rest = n[1:]
    found = []
    for k in (2, 1):
        head, tail = rest[:k], rest[k:]
        if _el(head) and tail:
            found.append((_el(head), tail, "G{EL}{TAG}"))
        head, tail = rest[:-k], rest[-k:]
        if _el(tail) and head:
            found.append((_el(tail), head, "G{TAG}{EL}"))
    known = [f for f in found if f[1] in TAG_PHASE]
    if known:
        return known[0]
    return None


def default_search_paths(workdir: Path) -> List[str]:
    from .compatibility import default_cache_dir
    out = [str(workdir)]
    for extra in (workdir / "tdbs", workdir / "examples" / "tdbs"):
        if extra.is_dir():
            out.append(str(extra))
    cache = default_cache_dir()
    if cache.is_dir():
        out.append(str(cache))
    return out


_SKIP = {".venv", "venv", "__pycache__", "node_modules", "site-packages", ".git"}


def database_paths(search_paths: Sequence[str], workdir: Path, limit: int = 500) -> List[str]:
    """TDB/.dat files (any case of the extension) in directories (recursive), files and glob patterns."""
    out: List[str] = []

    def add(p: str):
        rp = os.path.realpath(p)
        if rp not in out and Path(p).suffix.lower() in (".tdb", ".dat") and os.path.isfile(p):
            out.append(rp)

    for sp in search_paths:
        sp = os.path.expanduser(str(sp))
        if not os.path.isabs(sp):
            sp = str(workdir / sp)
        if os.path.isdir(sp):
            for root, dirs, files in os.walk(sp):
                dirs[:] = sorted(d for d in dirs if not d.startswith(".") and d not in _SKIP)
                for f in sorted(files):
                    add(os.path.join(root, f))
        elif any(ch in sp for ch in "*?["):
            for h in sorted(glob.glob(sp, recursive=True)):
                add(h)
        elif os.path.isfile(sp):
            add(sp)
        if len(out) >= limit:
            break
    return out[:limit]


class FileIndex:
    """Functions defined / referenced in one TDB with line numbers (text scan; nothing evaluated)."""

    def __init__(self, path: str):
        self.path = path
        self.is_dat = Path(path).suffix.lower() == ".dat"
        self.text = Path(path).read_text(encoding="utf-8", errors="replace")
        self.defs: Dict[str, List[Tuple[int, Piecewise, str]]] = {}
        self.refs: Dict[str, List[Tuple[int, str, str, set]]] = {}  # symbol -> (line, user, expression, symbols)
        self.params: List[Tuple[int, str, Piecewise]] = []
        if self.is_dat:
            return
        tlo, thi = 298.15, 6000.0
        cmds = scan_commands(self.text)
        for _, cmd in cmds:
            w = cmd.split()
            if w and resolve_keyword(w[0]) == "TEMPERATURE_LIMITS" and len(w) >= 3:
                try:
                    tlo, thi = float(w[1]), float(w[2])
                except ValueError:
                    pass
        for line, cmd in cmds:
            w = cmd.split()
            kw = resolve_keyword(w[0]) if w else None
            try:
                if kw == "FUNCTION" and len(w) > 2:
                    nm = _norm(w[1])
                    pw = parse_piecewise(cmd.split(None, 2)[2], tlo, thi)
                    self.defs.setdefault(nm, []).append((line, pw, " ".join(cmd.split())))
                    for s in pw.symbols():
                        self.refs.setdefault(s, []).append((line, f"FUNCTION {nm}", written(pw), set(pw.symbols())))
                elif kw == "PARAMETER":
                    body = cmd.split(None, 1)[1]
                    close = body.index(")")
                    raw = re.sub(r"\s+", "", body[: close + 1])
                    pw = parse_piecewise(body[close + 1:], tlo, thi)
                    self.params.append((line, raw, pw))
                    for s in pw.symbols():
                        self.refs.setdefault(s, []).append((line, raw, written(pw), set(pw.symbols())))
            except (ValueError, IndexError, ExpressionError):
                continue

    def mentions(self, name: str) -> List[int]:
        rx = re.compile(rf"(?<![A-Za-z0-9_]){re.escape(name)}(?![A-Za-z0-9_])", re.I)
        return [i for i, ln in enumerate(self.text.splitlines(), 1) if rx.search(ln)]


def dialect_variant(target: str, cand: str) -> Optional[str]:
    """Why ``cand`` may be ``target`` after tdb_dialect renaming (8/16-character FUNCTION limits, '_' for illegal
    characters), or None."""
    t = re.sub(r"[^A-Z0-9_]", "_", target)
    if cand == t and cand != target:
        return "illegal characters replaced by '_'"
    for lim, who in ((8, "Thermo-Calc"), (16, "OpenCalphad")):
        for a, b in ((t, cand), (cand, t)):
            if len(a) > lim and len(b) <= lim:
                m = re.fullmatch(r"(.*?)(\d{1,3})", b)
                if m and a.startswith(m.group(1)) and len(b) == lim:
                    return f"{who} {lim}-character function name ({a} -> {b})"
    return None


def resolve_symbol(name: str, files: Sequence[str], max_similar: int = 8,
                   unary_paths: Sequence[str] = ()) -> Dict[str, Any]:
    target = _norm(name)
    idx = []
    for f in files:
        try:
            idx.append(FileIndex(f))
        except OSError:
            continue
    defined, referenced, mentions, variants = [], [], [], []
    all_defs: Dict[str, List[str]] = {}
    for fi in idx:
        if fi.is_dat:
            lines = fi.mentions(target)
            if lines:
                mentions.append({"file": fi.path, "lines": lines[:20]})
            continue
        for nm in fi.defs:
            all_defs.setdefault(nm, []).append(fi.path)
        for line, pw, raw in fi.defs.get(target, []):
            defined.append({"file": fi.path, "line": line, "T_ranges": t_ranges(pw), "expression": written(pw),
                            "n_ranges": len(pw.ranges)})
        for line, user, expr, syms in fi.refs.get(target, []):
            referenced.append({"file": fi.path, "line": line, "by": user, "expression": expr,
                               "defined_in_same_file": target in fi.defs, "symbols": sorted(syms)})
    for nm, fs in all_defs.items():
        why = dialect_variant(target, nm) if nm != target else None
        if why:
            variants.append({"name": nm, "why": why, "files": fs})
    # similar names
    parsed = parse_symbol(target)
    similar: List[Dict[str, Any]] = []
    seen = {target}
    if parsed:
        el = parsed[0]
        for nm in sorted(all_defs):
            p2 = parse_symbol(nm)
            if p2 and p2[0] == el and nm not in seen:
                similar.append({"name": nm, "why": f"same element {el}, tag {p2[1]} ({p2[2]})",
                                "files": all_defs[nm][:5], "distance": _levenshtein(target, nm)})
                seen.add(nm)
    for nm in sorted(all_defs):
        if nm in seen or abs(len(nm) - len(target)) > 2:
            continue
        d = _levenshtein(target, nm)
        if d <= 2:
            similar.append({"name": nm, "why": f"edit distance {d}", "files": all_defs[nm][:5], "distance": d})
            seen.add(nm)
    similar.sort(key=lambda s: (not s["why"].startswith("same element"), s["distance"], s["name"]))
    similar = similar[:max_similar]
    # how the referencing databases use it, and the SGTE unary equivalent
    usage, unary = [], []
    if parsed and parsed[1] != "SER":
        el, tag, _ = parsed
        ser = f"GHSER{el}"
        for r in referenced:
            has_ser = any(s in r["symbols"] for s in (ser, f"GHSER{el}{el}"))
            usage.append("lattice stability (added to GHSER)" if has_ser else "full Gibbs energy")
        unary = unary_equivalents(el, tag, idx, unary_paths, exclude=target)
    out = {"name": target, "parsed": {"element": parsed[0], "tag": parsed[1], "form": parsed[2]} if parsed else None,
           "defined": defined, "referenced": referenced, "dat_mentions": mentions, "dialect_variants": variants,
           "similar": similar, "unary": unary, "files_searched": len(idx),
           "unary50_searched": any("unary" in Path(u).name.lower() for u in unary_paths),
           "usage": sorted(set(usage), key=usage.count, reverse=True)}
    return out


def unary_equivalents(el: str, tag: str, idx: List[FileIndex], unary_paths: Sequence[str] = (),
                      exclude: str = "") -> List[Dict]:
    """G(<phase of tag>,EL) and G<TAG><EL>-style functions in the searched databases (unary50 first), with the
    lattice stability relative to GHSER<EL> when both are present."""
    phase = TAG_PHASE.get(tag)
    names = {f"G{tag}{el}", f"G{el}{tag}", f"G{tag}{el}{el}" if len(el) == 1 else "", f"GHSER{el}"}
    names.discard("")
    order = sorted(idx, key=lambda fi: (os.path.realpath(fi.path) not in {os.path.realpath(u) for u in unary_paths},
                                        "unary" not in Path(fi.path).name.lower(), fi.path))
    out = []
    for fi in order:
        if fi.is_dat:
            continue
        hits = []
        for nm in sorted(names):
            for line, pw, _ in fi.defs.get(nm, [])[-1:]:
                if nm == f"GHSER{el}" and tag != "SER":
                    continue
                hits.append({"what": f"FUNCTION {nm}", "line": line, "expression": written(pw)})
        if phase:
            for line, raw, pw in fi.params:
                m = re.fullmatch(r"G\(([A-Z0-9_]+),([^;]*?)(?:;0)?\)", raw.upper())
                if not m or m.group(1) != phase:
                    continue
                subls = [s.split(",") for s in m.group(2).split(":")]
                flat = [s for sl in subls for s in sl]
                if exclude and exclude in pw.symbols():
                    continue
                if all(len(sl) == 1 for sl in subls) and el in flat and set(flat) <= {el, "VA"}:
                    hits.append({"what": raw, "line": line, "expression": written(pw), "param": pw})
        if not hits:
            continue
        for h in hits:
            pw = h.pop("param", None)
            if pw is None or tag == "SER":
                continue
            try:
                t = parse_tdb(fi.path)
                view = _empty_view(t)
                ser = next((s for s in (f"GHSER{el}", f"GHSER{el}{el}") if s in t.functions), None)
                if ser:
                    diff = Piecewise([(lo, hi, f"({e})-{ser}") for lo, hi, e in pw.ranges])
                    segs = inline(pw_segments(diff), view.fn, view.inline_notes)
                    h["minus_" + ser] = expanded_text(segs).replace("\n", " | ")
            except (MissingRef, ExpressionError, RecursionError, ValueError, OSError):
                pass
        out.append({"file": fi.path, "hits": hits})
        if len(out) >= 2:
            break
    return out


def resolve_text(r: Dict[str, Any], rel=lambda p: p, max_rows: int = 10) -> str:
    T = Text()
    nm = r["name"]
    T.add(f"{nm}: defined in {len(r['defined'])} place(s), referenced {len(r['referenced'])} time(s) "
          f"({r['files_searched']} files searched)")
    for d in r["defined"][:max_rows]:
        T.add(f"  DEFINED {rel(d['file'])}:{d['line']} [{d['T_ranges']}] {_cut(d['expression'], 110)}")
    for x in r["referenced"][:max_rows]:
        T.add(f"  used by {x['by']} at {rel(x['file'])}:{x['line']}"
              + ("" if x["defined_in_same_file"] else " (NOT defined in that file)")
              + f": {_cut(x['expression'], 70)}")
    if len(r["referenced"]) > max_rows:
        T.add(f"  ... {len(r['referenced']) - max_rows} more reference(s)")
    for m in r["dat_mentions"][:5]:
        T.add(f"  mentioned in {rel(m['file'])} (ChemSage .dat) line(s) {m['lines'][:8]}")
    for v in r["dialect_variants"][:5]:
        T.add(f"  dialect variant {v['name']}: {v['why']} in {', '.join(rel(f) for f in v['files'][:3])}")
    if r["parsed"]:
        p = r["parsed"]
        T.add(f"reads as element {p['element']}, phase tag {p['tag']} ({p['form']}); "
              + (f"the referencing parameters use it as: {', '.join(r['usage'])}" if r["usage"] else ""))
    if r["similar"]:
        T.add("similar names: " + "; ".join(f"{s['name']} ({s['why']}; {rel(s['files'][0])})"
                                            for s in r["similar"]))
    elif not r["defined"]:
        T.add("no similar names found")
    if r["parsed"] and r["parsed"]["tag"] != "SER":
        el, tag = r["parsed"]["element"], r["parsed"]["tag"]
        if r["unary"]:
            T.add(f"unary equivalents (G({TAG_PHASE.get(tag) or tag},{el}); G{tag}{el}-style functions):")
            for u in r["unary"]:
                for h in u["hits"][:4]:
                    extra = "".join(f"; {k.replace('minus_', '- ')} = {_cut(v, 90)}" for k, v in h.items()
                                    if k.startswith("minus_"))
                    T.add(f"  {rel(u['file'])}:{h['line']} {h['what']} = {_cut(h['expression'], 90)}{extra}")
            if not r["defined"]:
                usage = r["usage"][0] if r["usage"] else ""
                T.add(f"  -> to define {nm}: " + (f"FUNCTION {nm} = G({TAG_PHASE.get(tag)},{el}) - GHSER{el} (it is "
                                                  "added to GHSER in the referencing parameters)"
                                                  if usage.startswith("lattice") else
                                                  f"FUNCTION {nm} = the full G({TAG_PHASE.get(tag)},{el}) expression "
                                                  "(the referencing parameters use it without GHSER)"
                                                  if usage else "check how the referencing database uses it"))
        else:
            T.add(f"no G({TAG_PHASE.get(tag) or tag},{el}) / G{tag}{el} found in the searched files"
                  + ("" if r.get("unary50_searched") else " (no unary50 cache at ~/.calphad_agent/sgte/unary50.tdb; "
                     "evaluate_compatibility downloads it, or give search_paths)"))
    return T.text()


# =============================================================================================== registration
def register(tb) -> None:
    S = {"type": "string"}
    STRS = {"type": "array", "items": S}

    @tb.add("query_tdb",
            "Inspect phase models and parameters of TDB(s). Modes: phase (sublattices, site ratios, atoms/formula "
            "unit, magnetic/DIS_PART type defs, parameters grouped G/L0/L1/TC/BMAGN as written or "
            "expand_functions), element (every parameter with the element(s)), collisions (keys in several TDBs; "
            "identical / numerically equal / differs, shared phases and functions), functions (definitions, users, "
            "undefined), typedefs, formula_unit (atoms/fu, sum 1?, proportional k across TDBs), "
            "missing_endmembers (G absent; add_constituents {sublattice 0-based: [EL]}). Short text; full tables "
            "with output (.xlsx/.csv).",
            {"tdb": {"anyOf": [S, STRS], "description": "file, or a list for collisions/formula_unit"},
             "phases": STRS, "elements": STRS,
             "mode": {**S, "enum": list(MODES)},
             "expand_functions": {"type": "boolean", "description": "inline every FUNCTION (piecewise in T)"},
             "group_by": {**S, "enum": ["phase", "type", "elements"]},
             "add_constituents": {"type": "object", "description": "missing_endmembers: {sublattice: [species]}"},
             "output": {**S, "description": ".xlsx or .csv for the full tables"},
             "max_rows": {"type": "integer", "description": "rows shown in the reply (default 20)"}},
            ["tdb"])
    def query_tdb(tdb, mode: str = "phase", phases=None, elements=None, expand_functions: bool = False,
                  group_by: str = "phase", add_constituents=None, output=None, max_rows: int = 20):
        paths = [tb._find_tdb(t) for t in ([tdb] if isinstance(tdb, str) else list(tdb))]
        out = tb.outp(output, output) if output else None
        text, _, _ = query(paths, mode, phases, elements, bool(expand_functions), group_by or "phase", out,
                           int(max_rows or 20), add_constituents, rel=tb.rel)
        return text, []

    @tb.add("resolve_tdb_symbol",
            "Where a FUNCTION/symbol (e.g. GMOFCC) is defined and used across TDB/.dat files. Gives file, line, "
            "expression, T ranges of definitions and references, similar names (edit distance; same element with another phase tag: GFCCMO, "
            "GHSERMO, GMOHCP), dialect-shortened names, and the SGTE unary equivalent (G(FCC_A1,MO) - GHSERMO) with "
            "how the referencing TDB uses it. Default search: workdir TDBs + unary50 cache.",
            {"name": S, "search_paths": {**STRS, "description": "folders, files or globs (default: workdir, "
                                                                "examples/tdbs, ~/.calphad_agent/sgte)"}},
            ["name"])
    def resolve_tdb_symbol(name, search_paths=None):
        from .compatibility import default_cache_dir
        sp = list(search_paths) if search_paths else default_search_paths(tb.workdir)
        files = database_paths(sp, tb.workdir)
        cache = default_cache_dir()
        unary = [str(p) for p in cache.glob("*") if p.suffix.lower() == ".tdb"] if cache.is_dir() else []
        files += [u for u in unary if os.path.realpath(u) not in files]
        r = resolve_symbol(name, files, unary_paths=unary)
        return resolve_text(r, tb.rel), []
