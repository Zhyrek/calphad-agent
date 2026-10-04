"""``evaluate_compatibility``: can two (not yet merged) TDBs be joined?  A staged check with short verdicts.

B (the database to add) is first normalised to A's (the base's) conventions - phase names, sublattice order,
formula-unit size (see :func:`calphad_tools.tdbstage.normalize_to_base`) - then:

1. SGTE unaries: the SER Gibbs energy of each shared element (G of the ELEMENT reference phase, or GHSER<EL>),
   compared numerically over 298.15-6000 K (the shared range).  If they differ, each is checked against the SGTE
   unary database 5.0 (unary50.tdb; downloaded once and cached, or a local copy).
2. Other unary terms: lattice stabilities / magnetic terms of the shared elements in the phases both have.
3. Binary terms of the shared binaries.   4. Ternary and higher terms.
5. Phase models of the phases both have: site ratios, constituents, the merged (union) model, the NEW end-members
   it creates inside each database's own subsystems (and whether anything provides their G), magnetic and
   order-disorder settings, phases that newly appear in the other database's space, name candidates.
6. Replot decision: the A-priority merge (A's value wins every conflict, as staged_merge does) is built with
   merge.plan_merge, and each binary/ternary of A and of B is compared with it (tdbstage.diff_versions): a
   subsystem is provably unchanged when all its parameters are numerically equal and its phase models identical.

Text for the model is short; the details go to an Excel workbook and a short PowerPoint.
"""
from __future__ import annotations

import math
import os
import re
import tempfile
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from . import tdbstage as ts
from .tdbparse import parse_tdb

UNARY50_PAGE = "https://www.sgte.net/en/free-pure-elements-database"
UNARY50_URL = ("https://www.sgte.net/en/free-pure-elements-database?"
               "file=files/sgte/members-area/databases/unary50.tdb&cid=243")
UNARY50_TERMS = ("SGTE Unary Database v5.0 (2009), free download from " + UNARY50_PAGE + ". SGTE's terms: "
                 "'This database can only be used for extracting data for assessment work or to tabulate or plot "
                 "data for the pure elements.' Reference: A.T. Dinsdale, SGTE data for pure elements, Calphad 15 "
                 "(1991) 317-425.")
UNARY50_ENV = "CALPHAD_UNARY50"


def default_cache_dir() -> Path:
    return Path(os.environ.get("CALPHAD_AGENT_CACHE", Path.home() / ".calphad_agent")) / "sgte"


def _looks_like_unary(path: Path) -> bool:
    try:
        head = path.read_text(errors="replace")[:200000].upper()
    except OSError:
        return False
    return "ELEMENT" in head and "GHSERFE" in head and "FUNCTION" in head


def locate_unary50(spec: Optional[str] = "auto", cache_dir: Optional[str] = None, allow_download: bool = True,
                   timeout: float = 30.0) -> Tuple[Optional[str], str]:
    """(path or None, note).  spec: a file path (offline clusters), 'auto' (env CALPHAD_UNARY50, the cache
    ~/.calphad_agent/sgte/unary50.tdb, else download it once from SGTE) or 'off'."""
    if spec is not None and str(spec).lower() == "off":
        return None, "unary50 check off"
    if spec and str(spec).lower() != "auto":
        p = Path(os.path.expanduser(str(spec)))
        return (str(p), f"unary50: {p}") if p.exists() else (None, f"unary50 file {p} not found; check skipped")
    env = os.environ.get(UNARY50_ENV)
    if env and Path(env).exists():
        return env, f"unary50: {env} (${UNARY50_ENV})"
    cdir = Path(cache_dir) if cache_dir else default_cache_dir()
    cached = cdir / "unary50.tdb"
    if cached.exists() and _looks_like_unary(cached):
        return str(cached), f"unary50: cached {cached}"
    if not allow_download:
        return None, "unary50 not available (no local copy, download disabled); check skipped"
    try:
        cdir.mkdir(parents=True, exist_ok=True)
        req = urllib.request.Request(UNARY50_URL, headers={"User-Agent": "Mozilla/5.0 (calphad_tools)"})
        with urllib.request.urlopen(req, timeout=timeout) as r:            # noqa: S310 - fixed https URL
            data = r.read()
        fd, tmp = tempfile.mkstemp(dir=str(cdir), suffix=".tdb")
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        if not _looks_like_unary(Path(tmp)):
            os.unlink(tmp)
            return None, f"unary50 download from SGTE did not return a TDB; check skipped (get it from {UNARY50_PAGE})"
        os.replace(tmp, cached)
        return str(cached), f"unary50 downloaded from SGTE to {cached} (terms: assessment work / pure-element tables)"
    except Exception as e:  # noqa: BLE001 - offline clusters: say so and continue
        return None, (f"unary50 not available (download failed: {type(e).__name__}); check skipped - give a local "
                      f"copy with unary50='path' (from {UNARY50_PAGE})")


# =============================================================================================== data
@dataclass
class Stage:
    n: int
    title: str
    verdict: str                    # OK | INCOMPATIBLE | CONFLICT | WARNING | CHANGES | NONE
    line: str
    rows: List[Dict[str, Any]] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)     # for the overall line (e.g. zero-G end-members)


@dataclass
class Evaluation:
    a_path: str
    b_path: str
    shared: List[str]
    stages: List[Stage]
    norm_notes: List[ts.NormNote]
    replot: Dict[str, List[ts.SubsystemDiff]]
    unary_note: str = ""
    normalised_b: str = ""
    predicted_merge: str = ""
    files: Dict[str, str] = field(default_factory=dict)
    rescaled_a: str = ""

    @property
    def overall(self) -> str:
        v = {s.n: s.verdict for s in self.stages}
        bad = [n for n in (1, 2, 5) if v.get(n) == "INCOMPATIBLE"]
        conf = [n for n in (3, 4) if v.get(n) == "CONFLICT"]
        warn = [w for s in self.stages for w in s.warnings]
        if bad:
            return f"INCOMPATIBLE as is (stage {', '.join(map(str, bad))}: harmonise before merging)" + (
                f"; stage {', '.join(map(str, conf))}: choose one assessment per subsystem" if conf else "") + (
                f"; warnings: {_examples(warn, 3)}" if warn else "")
        warn = [w for s in self.stages for w in s.warnings]
        if conf:
            return f"MERGEABLE WITH DECISIONS (stage {', '.join(map(str, conf))}: differently assessed subsystems)" + (
                f"; warnings: {_examples(warn, 3)}" if warn else "")
        if warn:
            return f"COMPATIBLE with warnings: {_examples(warn, 3)}"
        return "COMPATIBLE"


def _split_top(text: str, sep: str) -> List[str]:
    """Split on ``sep`` outside parentheses / brackets."""
    out, depth, cur, i = [], 0, "", 0
    while i < len(text):
        ch = text[i]
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth = max(0, depth - 1)
        if depth == 0 and text.startswith(sep, i):
            out.append(cur)
            cur, i = "", i + len(sep)
            continue
        cur += ch
        i += 1
    out.append(cur)
    return [x for x in out if x.strip()]


def clip(text: str, n: int, where: str = "") -> str:
    """At most ~n characters, cut only at a clause boundary ('; ', then ', ') and never inside a word:
    'a; b; +3 more (Excel)'."""
    text = " ".join(text.split())
    if len(text) <= n:
        return text
    more = f" ({where})" if where else ""
    for sep in ("; ", ", "):
        parts = _split_top(text, sep)
        if len(parts) < 2:
            continue
        keep: List[str] = []
        for i, p in enumerate(parts):
            suffix = f"{sep}+{len(parts) - i - 1} more{more}" if i < len(parts) - 1 else ""
            if len(sep.join(keep + [p]) + suffix) > n and (keep or len(p) > 1.3 * n):
                break
            keep.append(p)
        if keep:
            return sep.join(keep) + f"{sep}+{len(parts) - len(keep)} more{more}"
    cut = text[: max(10, n - 12 - len(more))].rsplit(" ", 1)[0]
    return cut + " …" + more


def _short(t: str, n: int) -> str:
    return clip(t, n)


STATUS_FILL = {"OK": "same", "NONE": "absent", "CHANGES": "error", "CONFLICT": "differs", "INCOMPATIBLE": "differs",
               "COMPATIBLE": "same", "MERGEABLE": "error"}


STATUS_RGB = {"OK": "C8E6C9", "COMPATIBLE": "C8E6C9", "NONE": "ECEFF1", "CHANGES": "FFE082", "MERGEABLE": "FFE082",
              "WARNING": "FFCC80",
              "CONFLICT": "EF9A9A", "INCOMPATIBLE": "EF9A9A"}


def color_status_cells(prs, col: int = 1) -> None:
    """Strong fills for status words (OK green, CHANGES amber, CONFLICT / INCOMPATIBLE red) in column ``col`` of
    every table."""
    from pptx.dml.color import RGBColor
    for sl in prs.slides:
        for sh in sl.shapes:
            if not getattr(sh, "has_table", False) or not sh.has_table:
                continue
            for r in list(sh.table.rows)[1:]:
                cells = list(r.cells)
                if col < len(cells):
                    rgb = STATUS_RGB.get(cells[col].text.strip().split(" ")[0])
                    if rgb:
                        cells[col].fill.solid()
                        cells[col].fill.fore_color.rgb = RGBColor.from_string(rgb)


def fit_rows(rows: List[List[Any]], widths: Sequence[float], size: float, more, max_h: float = 5.25) -> List:
    """Keep as many GridCell rows as fit on one slide (estimated wrapped height); the rest becomes one row made by
    ``more(k)``."""
    from .pptx_utils import estimate_row_height
    avail = max_h - 0.32
    hs = [estimate_row_height(r, widths, size) for r in rows]
    if sum(hs) <= avail:
        return rows
    out, used = [], 0.0
    extra = estimate_row_height(more(len(rows)), widths, size)
    for r, h in zip(rows, hs):
        if used + h + extra > avail:
            break
        out.append(r)
        used += h
    return out + [more(len(rows) - len(out))]


def _cap(e: str) -> str:
    return e.capitalize()


def _examples(items: Sequence[str], n: int = 3) -> str:
    items = list(items)
    return ", ".join(items[:n]) + (f" (+{len(items) - n})" if len(items) > n else "")


# =============================================================================================== stages
def _stage1(VA: ts.StageView, VB: ts.StageView, shared, unary_spec, allow_download, rtol) -> Tuple[Stage, str, Any]:
    rows, diffs, notes = [], [], []
    for el in shared:
        sa, sb = VA.ser(el), VB.ser(el)
        ea, eb = VA.tdb.elements.get(el), VB.tdb.elements.get(el)
        row = {"element": el, "A reference": (sa.function or sa.what) if sa else "none",
               "B reference": (sb.function or sb.what) if sb else "none",
               "A ELEMENT phase": ea.ref_phase if ea else "", "B ELEMENT phase": eb.ref_phase if eb else ""}
        if sa is None or sb is None:
            row["status"] = "no SER reference found"
            rows.append(row)
            continue
        ok, d, t = ts.compare_segs(sa.segs, sb.segs, rtol, 1e-2, ts.STAGE_GRID)
        row.update({"status": "equal" if ok else "differs", "max |Δ| (J/mol)": d, "T at max (K)": t,
                    "_segs": (sa.segs, sb.segs)})
        if ea and eb and ea.ref_phase.upper() != eb.ref_phase.upper():
            notes.append(f"{_cap(el)} ELEMENT reference phase {ea.ref_phase} (A) vs {eb.ref_phase} (B)")
        rows.append(row)
        if not ok:
            diffs.append(row)
    unary_note, V50 = "", None
    if diffs and unary_spec != "off":
        path, unary_note = locate_unary50(unary_spec, allow_download=allow_download)
        if path:
            V50 = ts.StageView(parse_tdb(path), "unary50", elements=shared)
            for row in diffs:
                s50 = V50.ser(row["element"])
                if s50 is None:
                    row["unary50"] = "element not in unary50"
                    continue
                okA, dA, _ = ts.compare_segs(row["_segs"][0], s50.segs, rtol, 1e-2, ts.STAGE_GRID)
                okB, dB, _ = ts.compare_segs(row["_segs"][1], s50.segs, rtol, 1e-2, ts.STAGE_GRID)
                row.update({"A vs unary50 (J/mol)": dA, "B vs unary50 (J/mol)": dB,
                            "unary50": "A matches" if okA and not okB else "B matches" if okB and not okA else
                            "both match" if okA else "neither matches"})
    for r in rows:
        r.pop("_segs", None)
    equal = [r["element"] for r in rows if r.get("status") == "equal"]
    maxeq = max([r["max |Δ| (J/mol)"] for r in rows if r.get("status") == "equal"] or [0.0])
    if not diffs and all(r.get("status") == "equal" for r in rows):
        line = (f"OK ({', '.join(map(_cap, equal))} equal" +
                (f"; max |Δ| {ts.fmt_energy(maxeq)} within tolerance" if maxeq > 1e-6 else "") + ")")
        verdict = "OK"
    else:
        parts = []
        for r in diffs:
            w = (f"SER of {_cap(r['element'])} ({r['A reference']} / {r['B reference']}) differs by up to "
                 f"{ts.fmt_energy(r['max |Δ| (J/mol)'])} at {r['T at max (K)']:.0f} K")
            u = r.get("unary50")
            if u == "A matches":
                w += " (A matches unary50; B is an older/other unary)"
            elif u == "B matches":
                w += " (B matches unary50; A is an older/other unary)"
            elif u:
                w += f" ({u} unary50)"
            parts.append(w)
        parts += [f"SER of {_cap(r['element'])}: {r['status']}" for r in rows if r.get("status") not in ("equal",
                                                                                                         "differs")]
        line = "INCOMPATIBLE: " + "; ".join(parts)
        if equal:
            line += f"; {', '.join(map(_cap, equal))} equal"
        verdict = "INCOMPATIBLE"
    if notes:
        line += "; " + "; ".join(notes)
    if diffs and unary_note and V50 is None:
        line += f" [{unary_note}]"
    return Stage(1, "SGTE unaries", verdict, line, rows), unary_note, V50


def _compare_rows(VA, VB, keys, rtol, V50=None) -> List[Dict[str, Any]]:
    out = []
    for k in keys:
        ra, rb = VA.rows.get(k), VB.rows.get(k)
        r = ra or rb
        row = {"parameter": r.label, "phase": r.phase, "subsystem": ts.stage_sysname((VA if ra else VB).row_elements(k)),
               "A": ts._written(VA, k) if ra else "", "B": ts._written(VB, k) if rb else "", "_key": k,
               "_ptype": r.ptype}
        if ra is None or rb is None:
            other = VB if rb is None else VA
            row["status"] = "only in A" if rb is None else "only in B"
            if not _can_hold(other, k):
                row["status"] += " (other model lacks it)"
        elif ra.segs is None or rb.segs is None:
            row["status"] = "not evaluated (" + (ra.error or rb.error) + ")"
        else:
            ok, d, t = ts.compare_segs(ra.segs, rb.segs, rtol, ts.stage_atol(r.ptype))
            row.update({"status": "equal" if ok else "differs", "max |Δ|": d, "T at max (K)": t})
            if not ok and V50 is not None and k in V50.rows and V50.rows[k].segs is not None:
                okA, _, _ = ts.compare_segs(ra.segs, V50.rows[k].segs, rtol, ts.stage_atol(r.ptype))
                okB, _, _ = ts.compare_segs(rb.segs, V50.rows[k].segs, rtol, ts.stage_atol(r.ptype))
                row["unary50"] = ("A matches" if okA and not okB else "B matches" if okB and not okA else
                                  "both match" if okA else "neither matches")
        out.append(row)
    return out


def _can_hold(view: ts.StageView, key: str) -> bool:
    """Could ``view``'s phase model hold parameter ``key`` (all its constituents on the right sublattices)?"""
    pv = view.phases.get(key[key.index("(") + 1:].split(",", 1)[0].split("&")[0])
    if pv is None:
        return False
    keep = [i for i in range(pv.n_sub) if i not in pv.va_only]
    cons = ts.key_constituents(key)
    if len(cons) != len(keep):
        return False
    return all(s in ("*", "VA") or s in pv.constituents[i] for i, sl in zip(keep, cons) for s in sl)


def _row_text(r: Dict[str, Any]) -> str:
    s = f"{r['parameter']} Δ{ts.fmt_delta(r.get('max |Δ|', math.nan), r['_ptype'])}"
    if r.get("unary50") in ("A matches", "B matches"):
        s += f" ({r['unary50']} unary50)"
    return s


def _param_stage(n, title, VA, VB, shared, sizes, rtol, V50=None, exclude=()) -> Stage:
    E = set(shared)
    common = set(VA.phases) & set(VB.phases)
    keys = []
    for V in (VA, VB):
        for k, r in V.rows.items():
            els = V.row_elements(k)
            if len(els) in sizes and els <= E and r.phase in common and k not in exclude and k not in keys:
                keys.append(k)
    keys.sort(key=lambda k: ts.row_sort_key(k, (VA.rows.get(k) or VB.rows[k]).ptype))
    rows = _compare_rows(VA, VB, keys, rtol, V50)
    diff = [r for r in rows if r["status"] == "differs"]
    onlyA = [r for r in rows if r["status"].startswith("only in A")]
    onlyB = [r for r in rows if r["status"].startswith("only in B")]
    if not rows:
        return Stage(n, title, "NONE", "none shared", rows)
    if n == 2:
        if diff:
            line = (f"INCOMPATIBLE: {len(diff)} of {len(rows)} differ: "
                    + _examples([_row_text(r) for r in sorted(diff, key=lambda r: -r['max |Δ|'])]))
            verdict = "INCOMPATIBLE"
        else:
            line, verdict = f"OK ({len(rows) - len(onlyA) - len(onlyB)} equal)", "OK"
        extra = []
        if onlyA:
            extra.append(f"{len(onlyA)} only in A ({_examples([r['parameter'] for r in onlyA], 2)})")
        if onlyB:
            extra.append(f"{len(onlyB)} only in B ({_examples([r['parameter'] for r in onlyB], 2)}; the merge adds "
                         f"them to A)")
        return Stage(n, title, verdict, line + ("; " + "; ".join(extra) if extra else ""), rows)
    # binaries / higher: per subsystem
    per: Dict[str, List[Dict]] = {}
    for r in rows:
        per.setdefault(r["subsystem"], []).append(r)
    parts, verdict = [], "OK"
    for sub, rs in per.items():
        d = [r for r in rs if r["status"] == "differs"]
        a = [r for r in rs if r["status"].startswith("only in A")]
        b = [r for r in rs if r["status"].startswith("only in B")]
        lacks = [r for r in rs if "lacks" in r["status"]]
        a = [r for r in a if r not in lacks]
        b = [r for r in b if r not in lacks]
        ex = (f"; {len(lacks)} only where the other model lacks the constituent (stage 5)" if lacks else "")
        if not d and not a and not b:
            parts.append(f"{sub} OK ({len(rs) - len(lacks)} equal{ex})")
            continue
        if len(a) == len(rs) or len(b) == len(rs):
            parts.append(f"{sub}: parameters only in {'A' if a else 'B'} ({len(rs)})")
            continue
        verdict = "CONFLICT"
        s = f"{sub} CONFLICT: {len(d)} of {len(rs)} differ"
        if d:
            s += " (" + _examples([_row_text(r) for r in sorted(d, key=lambda r: -r['max |Δ|'])], 2) + ")"
        if a or b:
            s += f"; {len(a)} only in A, {len(b)} only in B"
        parts.append(s)
    return Stage(n, title, verdict, "; ".join(parts), rows)


def _restrict_pm(pm: ts.PhaseModel, E, views) -> Optional[ts.PhaseModel]:
    E = set(E) | {"VA"}

    def ok(s):
        if s == "VA":
            return True
        for v in views:
            if s in v.tdb.species or s in v.tdb.elements:
                return v.species_elements(s) <= E
        return re.sub(r"[+-]\d*$", "", s) in E
    cons = [tuple(s for s in c if ok(s)) for c in pm.constituents]
    if any(not c for c in cons) or all(s == "VA" for c in cons for s in c):
        return None
    return ts.PhaseModel(pm.ratios, tuple(cons), pm.magnetic, pm.dis_part)


def _model_decision_hint(ph: str) -> str:
    return (f"staged_merge / merge_tdbs decisions {{'phase:{ph}': 'rename'}} keeps both, B's as {ph}_2; 'db1' keeps "
            f"only A's model, 'db2' only B's")


def _assessed_in(view: ts.StageView, ph: str, els) -> bool:
    """Does ``view`` describe phase ``ph`` in subsystem ``els`` (a parameter of the phase with >= 2 of its elements)?"""
    E = set(els)
    return any(r.phase == ph and len(view.row_elements(k)) >= 2 and view.row_elements(k) <= E
               for k, r in view.rows.items())


def _zero_g_warning(ph: str, em, where) -> str:
    return f"{ph} {':'.join(em)} has no G (pycalphad uses 0; {_examples(where, 2)})"


def _stage5(VA: ts.StageView, VB: ts.StageView, norm_notes, rtol) -> Tuple[Stage, List[str]]:
    """Phase models.  A new end-member of the merged model with no G in either database (pycalphad uses G = 0):

    * INCOMPATIBLE when the database whose space it lies in already describes the phase in a binary containing
      the end-member (a parameter of the phase with two elements of that binary): the zero G changes an assessed
      phase there;
    * WARNING otherwise - the phase is new in that binary (a phase of the other database entering it) or has only
      unary end-members there: it was not assessed in that binary; the overall verdict says 'COMPATIBLE with
      warnings', and staged_merge reports whether the phase becomes stable in the mapped diagrams.
    """
    rows, parts, verdict = [], [], "OK"
    warnings: List[str] = []
    EA, EB = VA.own_elements, VB.own_elements
    subsA = ts.stage_subsystems(EA, (2, 3))
    subsB = ts.stage_subsystems(EB, (2, 3))
    for ph in sorted(set(VA.phases) & set(VB.phases), key=lambda p: (p != "LIQUID", p)):
        ma, mb = ts.full_model(VA, ph), ts.full_model(VB, ph)
        if ma is None or mb is None:
            continue
        row = {"phase": ph, "A model": ma.text(), "B model": mb.text(), "A magnetic/ordering": ma.extras(),
               "B magnetic/ordering": mb.extras()}
        unm = next((n for n in norm_notes if n.phase == ph and n.action == "unmapped"), None)
        if unm is not None:
            # the sublattice correspondence is ambiguous: compared as is, but the user has to say how they map
            row["issue"] = f"sublattice mapping undecided: {unm.detail}"
            verdict = "INCOMPATIBLE"
            parts.append(f"{ph}: sublattice mapping undecided ({unm.detail.split(': ', 1)[-1]}) - INCOMPATIBLE as "
                         f"is; give {ts.unmapped_hint(ph)}")
            rows.append(row)
            continue
        merged = ts.union_model(ma, mb)
        if merged is None:
            row["issue"] = "site ratios differ: models cannot be merged"
            row["decision"] = _model_decision_hint(ph)
            verdict = "INCOMPATIBLE"
            parts.append(f"{ph}: site ratios {':'.join(f'{x:g}' for x in ma.ratios)} (A) vs "
                         f"{':'.join(f'{x:g}' for x in mb.ratios)} (B) - INCOMPATIBLE (rename one, or drop one "
                         f"model: {_model_decision_hint(ph)})")
            rows.append(row)
            continue
        row["merged model"] = merged.text()
        issues = []
        if ma.magnetic != mb.magnetic:
            issues.append(f"magnetic {ma.extras() or 'none'} (A) vs {mb.extras() or 'none'} (B)")
        if ma.dis_part != mb.dis_part:
            issues.append(f"DIS_PART {ma.dis_part or '-'} (A) vs {mb.dis_part or '-'} (B)")
        news = []
        for side, own, other, E, subs in (("A", VA, VB, EA, subsA), ("B", VB, VA, EB, subsB)):
            mo = ma if side == "A" else mb
            mr = _restrict_pm(merged, E, (VA, VB))
            new = sorted(ts.model_endmembers(mr) - ts.model_endmembers(mo))
            for em in new:
                key = own.key_for("G", ph, em)
                where = [ts.stage_sysname(s) for s in subs if ts.endmember_elements(own, em) <= set(s)
                         and _restrict_pm(merged, s, (VA, VB)) is not None]
                if key in other.rows:
                    src = "B" if side == "A" else "A"
                    alt = other.segs_in(key, own)
                    okk, d, _ = (ts.compare_segs(other.rows[key].segs, alt, rtol, 1e-2)
                                 if alt is not None and other.rows[key].segs is not None else (False, math.nan, 0))
                    prov = (f"G from {src}" + ("" if okk else f"; with {side}'s functions it shifts by "
                                                               f"{ts.fmt_energy(d)}"))
                else:
                    bins = [s for s in subs if len(s) == 2 and ts.endmember_elements(own, em) <= set(s)
                            and _restrict_pm(merged, s, (VA, VB)) is not None]
                    assessed = [ts.stage_sysname(s) for s in bins if _assessed_in(own, ph, s)]
                    if assessed:
                        prov = (f"no G in either: pycalphad uses 0 - changes {side}'s assessed {ph} in "
                                f"{_examples(assessed, 2)}")
                        verdict = "INCOMPATIBLE"
                    else:
                        prov = "no G in either: pycalphad uses 0"
                        warnings.append(_zero_g_warning(ph, em, where or [ts.stage_sysname(em)]))
                news.append(f"{':'.join(em)} in {side}'s space ({_examples(where, 3) or 'unary'}; {prov})")
                rows.append({"phase": ph, "new end-member": ":".join(em), "in space of": side,
                             "affects subsystems": ", ".join(where), "provided by": prov})
        if issues or news or merged != ma or merged != mb:
            row["issue"] = "; ".join(issues)
            row["new end-members"] = "; ".join(news)
            if news or issues:
                news.sort(key=lambda x: "no G" not in x)
                parts.append(f"{ph}: A {ma.text()}, B {mb.text()} -> merged {merged.text()}"
                             + (f"; new {_examples(news, 2)}" if news else "") + (f"; {'; '.join(issues)}" if issues else ""))
                if verdict == "OK":
                    verdict = "CHANGES"
        rows.append(row)
    # phases of one database that appear in the other's space
    for side, own, other, E in (("A", VA, VB, EA), ("B", VB, VA, EB)):
        newp = []
        for ph in sorted(set(other.phases) - set(own.phases)):
            m = ts.restricted_model(other, ph, E)
            if m is not None:
                miss = [em for em in sorted(ts.model_endmembers(m)) if other.key_for("G", ph, em) not in other.rows]
                zero = []
                for em in miss:
                    where = [ts.stage_sysname(s) for s in ts.stage_subsystems(E, (2, 3))
                             if ts.endmember_elements(other, em) <= set(s) and ts.restricted_model(other, ph, s)]
                    zero.append(":".join(em))
                    warnings.append(_zero_g_warning(ph, em, where or [ts.stage_sysname(
                        ts.endmember_elements(other, em))]))
                    rows.append({"phase": ph, "new end-member": ":".join(em), "in space of": side,
                                 "affects subsystems": ", ".join(where),
                                 "provided by": f"no G in either: pycalphad uses 0 ({ph} is new in {side}'s space)"})
                newp.append(f"{ph} {m.text()}" + (f" [no G for {', '.join(zero[:3])}: pycalphad uses 0]"
                                                  if zero else ""))
                rows.append({"phase": ph, "new phase in space of": side, "model there": m.text()})
        if newp:
            parts.append(f"new phases in {side}'s space: {_examples(newp, 3)}")
            if verdict == "OK":
                verdict = "CHANGES"
    cands = [n for n in norm_notes if n.action == "candidate"]
    try:
        from .compare_system import phase_candidates
        for c in phase_candidates([VA, VB], VA.label):
            if c.verdict.startswith("LIKELY") and not any(c.a in n.detail or c.b in n.detail for n in cands):
                cands.append(ts.NormNote(c.b, "candidate", f"~ {c.a}: {c.verdict} ({'; '.join(c.evidence[:2])}); "
                                                           f"phase_map {c.suggestion}"))
    except Exception:  # noqa: BLE001 - candidates are a hint only
        pass
    for n in cands:
        rows.append({"phase": n.phase, "name candidate": n.detail})
    if cands:
        parts.append(f"name candidates (not merged): {_examples([n.phase + ' ' + n.detail.split(';')[0] for n in cands], 2)}")
    if warnings and verdict in ("OK", "CHANGES"):
        verdict = "WARNING"
    line = (f"{verdict}: " + "; ".join(parts)) if parts else "OK (shared phases have identical models)"
    st = Stage(5, "phase models", verdict, line, rows)
    st.warnings = warnings
    return st, []


def _stage6(a_path, bn_path, VA, VB, rtol, stages, workdir,
            a_rescaled=()) -> Tuple[Stage, Dict[str, List[ts.SubsystemDiff]], str]:
    dec = ts.alias_separate([a_path, bn_path])
    plan, text = ts.merged_text([a_path, bn_path], dec, "first")
    mpath = str(Path(workdir) / f"predicted_merge_{Path(a_path).stem}+{Path(bn_path).stem}.tdb")
    Path(mpath).write_text(text)
    M = parse_tdb(mpath)
    VMa = ts.StageView(M, "merged", elements=sorted(VA.own_elements))
    VMb = ts.StageView(M, "merged", elements=sorted(VB.own_elements))
    da = ts.diff_versions(VA, VMa, ts.stage_subsystems(VA.own_elements, (2, 3)), rtol)
    db = ts.diff_versions(VB, VMb, ts.stage_subsystems(VB.own_elements, (2, 3)), rtol)
    conflicts = [r for s in stages if s.n in (2, 3, 4) for r in s.rows if r.get("status") == "differs"]
    rows, parts = [], []
    for side, d in (("A", da), ("B", db)):
        ch = d.changed
        for s in d.subsystems:
            reason = s.reasons()
            if side == "A" and not s.changed:
                c = [r for r in conflicts if set(r["subsystem"].upper().split("-")) <= set(s.elements)]
                if c:
                    reason += f" - only because A's values win {len(c)} conflict(s)"
            rows.append({"database": side, "subsystem": s.name, "replot": "yes" if s.changed else "no",
                         "reasons": reason})
        if ch:
            parts.append(f"{side}: replot {'; '.join(f'{s.name} ({_short(s.reasons(1), 90)})' for s in ch[:4])}"
                         + (f"; +{len(ch) - 4} more (Excel)" if len(ch) > 4 else ""))
        unchanged = [s.name for s in d.subsystems if not s.changed]
        if unchanged:
            parts.append(f"{side} unchanged: {_examples(unchanged, 4)}")
    verdict = "CHANGES" if da.changed or db.changed else "OK"
    line = "(A's values win conflicts) " + " | ".join(parts) if parts else "no binaries/ternaries"
    for n in a_rescaled:
        line += f" | {n.phase}: A rescaled to B's site ratios (B magnetic, 1 atom/fu): exact, not a change"
    for c in plan.conflicts:
        if c.kind == "phase_model" and c.decided_by.startswith("policy"):
            ph = c.id.split(":", 1)[1]
            line += (f" | {ph}: predicted with A's model only ('{c.decision}'); decision 'rename' keeps B's as "
                     f"{ph}_2 instead (B's subsystems with it then keep it)")
    exact = sorted({s.name for d in (da, db) for s in d.subsystems if s.notes and not s.changed})
    if exact:
        line += (f" | exact-only: {', '.join(exact)} (staged_merge force_replot={exact!r} maps them anyway)")
    return Stage(6, "replot decision", verdict, line, rows), {"A": da.subsystems, "B": db.subsystems}, mpath


def evaluate_pair(a_path: str, b_path: str, *, phase_map: Optional[Dict[str, str]] = None,
             sublattice_map: Optional[Dict] = None, unary50: Optional[str] = "auto", allow_download: bool = True,
             rtol: float = ts.STAGE_RTOL, workdir: Optional[str] = None) -> Evaluation:
    A, B = parse_tdb(a_path), parse_tdb(b_path)
    workdir = Path(workdir or tempfile.mkdtemp(prefix="compat_"))
    workdir.mkdir(parents=True, exist_ok=True)
    An, Bn, notes = ts.normalize_pair(A, B, phase_map, sublattice_map, on_error="keep")
    changed = [n for n in notes if n.action in ("renamed", "reordered", "rescaled")]
    a_rescaled = [n for n in notes if n.action == ts.A_RESCALED]
    an_path = a_path
    if a_rescaled:
        an_path = str(workdir / f"{Path(a_path).stem}_rescaled.tdb")
        ts.write_tdb(An, an_path, [f"{Path(a_path).name} with phases rescaled to "
                                   f"{Path(b_path).name}'s site ratios (magnetic exception):"]
                     + [n.text() for n in a_rescaled])
        An = parse_tdb(an_path)
    bn_path = b_path
    if changed:
        bn_path = str(workdir / f"{Path(b_path).stem}_normalised.tdb")
        ts.write_tdb(Bn, bn_path, [f"{Path(b_path).name} normalised to {Path(a_path).name}'s conventions:"]
                     + [n.text() for n in changed])
        Bn = parse_tdb(bn_path)
    shared = sorted(set(A.real_elements) & set(Bn.real_elements))
    VA = ts.StageView(An, "A")
    VB = ts.StageView(Bn, "B")
    s1, unote, V50 = _stage1(VA, VB, shared, unary50, allow_download, rtol)
    ser_keys = set()
    for el in shared:
        for V in (VA, VB):
            s = V.ser(el)
            if s is not None and s.kind == "param":
                ser_keys.add(s.what)
    s2 = _param_stage(2, "other unaries", VA, VB, shared, (1,), rtol, V50, ser_keys)
    if s2.verdict == "INCOMPATIBLE" and V50 is None and unary50 != "off":
        path, unote2 = locate_unary50(unary50, allow_download=allow_download)
        unote = unote or unote2
        if path:
            V50 = ts.StageView(parse_tdb(path), "unary50", elements=shared)
            s2 = _param_stage(2, "other unaries", VA, VB, shared, (1,), rtol, V50, ser_keys)
    s3 = _param_stage(3, "binaries", VA, VB, shared, (2,), rtol)
    s4 = _param_stage(4, "ternary+", VA, VB, shared, tuple(range(3, 12)), rtol)
    s5, _ = _stage5(VA, VB, notes, rtol)
    stages = [s1, s2, s3, s4, s5]
    VA0 = ts.StageView(A, "A") if a_rescaled else VA
    s6, replot, mpath = _stage6(an_path, bn_path, VA0, VB, rtol, stages, workdir, a_rescaled)
    stages.append(s6)
    ev = Evaluation(str(a_path), str(b_path), shared, stages, notes, replot, unote,
                    bn_path if bn_path != b_path else "", mpath)
    ev.rescaled_a = an_path if an_path != a_path else ""
    return ev


# =============================================================================================== output
def summary_text(ev: Evaluation, rel=lambda p: p) -> str:
    L = [f"evaluate_compatibility: A = {Path(ev.a_path).name} (base), B = {Path(ev.b_path).name}; shared elements "
         f"{', '.join(map(_cap, ev.shared)) or 'none'}"]
    done = [n.text() for n in ev.norm_notes if n.action in ts.NORM_CHANGES]
    kept = [n.text() for n in ev.norm_notes if n.action in ("not_rescaled", "incompatible", "unmapped", "kept")]
    if done or kept:
        L.append("Normalised B to A's conventions: " + (_examples(done, 4) if done else "nothing")
                 + (f"; NOT normalised: {_examples(kept, 3)}" if kept else ""))
    for s in ev.stages:
        L.append(f"Stage {s.n} {s.title}: {s.line}")
    L.append(f"Overall: {ev.overall}.")
    if ev.files:
        L.append("Files: " + ", ".join(f"{k} {rel(v)}" for k, v in ev.files.items()))
    if ev.unary_note and "unary50" in ev.unary_note and "downloaded" in ev.unary_note:
        L.append(f"Note: {ev.unary_note}. {UNARY50_TERMS}")
    return "\n".join(L)


def _clean(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows]


def write_xlsx(ev: Evaluation, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for line in summary_text(ev).split("\n"):
        ws.append([line])
    ws.column_dimensions["A"].width = 160
    red = PatternFill("solid", fgColor="FFEBEE")

    def sheet(name, rows):
        rows = _clean(rows)
        w = wb.create_sheet(name[:31])
        cols: List[str] = []
        for r in rows:
            for k in r:
                if k not in cols:
                    cols.append(k)
        w.append(cols or ["(none)"])
        for c in w[1]:
            c.font = Font(bold=True)
        for r in rows:
            vals = []
            for c in cols:
                x = r.get(c, "")
                vals.append(round(x, 6) if isinstance(x, float) and math.isfinite(x) else
                            ("" if isinstance(x, float) else x))
            w.append(vals)
            st = str(r.get("status", "")) + str(r.get("replot", "")) + str(r.get("issue", ""))
            if "differs" in st or st.startswith("yes") or "INCOMPATIBLE" in st or "cannot" in st:
                for c in w[w.max_row]:
                    c.fill = red
        for i, c in enumerate(cols, 1):
            w.column_dimensions[w.cell(1, i).column_letter].width = min(70, max(12, len(c) + 2))
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    names = {1: "1 SGTE unaries", 2: "2 other unaries", 3: "3 binaries", 4: "4 ternary+", 5: "5 phase models",
             6: "6 replot"}
    for s in ev.stages:
        sheet(names[s.n], s.rows)
    sheet("normalisation", [{"phase": n.phase, "action": n.action, "detail": n.detail} for n in ev.norm_notes])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_pptx(ev: Evaluation, path: str) -> str:
    from . import pptx_utils as pu
    from .pptx_utils import GridCell as C
    prs = pu.new_presentation()
    sheets = {1: "1 SGTE unaries", 2: "2 other unaries", 3: "3 binaries", 4: "4 ternary+", 5: "5 phase models",
              6: "6 replot"}
    sub = (f"A = {Path(ev.a_path).name} (base), B = {Path(ev.b_path).name}; shared "
           f"{', '.join(map(_cap, ev.shared)) or 'none'}")
    rows = []
    for s in ev.stages:
        line = re.sub("^" + s.verdict + r":?\s*", "", s.line)
        if line.startswith("(") and line.endswith(")") and s.verdict == "OK":
            line = line[1:-1]
        rows.append([C(f"{s.n}  {s.title}", bold=True), C(s.verdict, bold=True, align="c"),
                     C(clip(line, 420, f"Excel sheet '{sheets[s.n]}'"))])
    ov = "WARNING" if "with warnings" in ev.overall.split(";")[0] else ev.overall.split(" ")[0]
    rows.append([C("Overall", bold=True), C(ov, bold=True, align="c"), C(ev.overall)])
    pu.add_flow_table_slides(prs, f"Compatibility: {Path(ev.a_path).stem} + {Path(ev.b_path).stem}",
                             ["Stage", "Status", "Verdict"], rows, [2.3, 1.9, 8.6], font_size=13, subtitle=sub,
                             max_height_in=5.6)
    # stages 1-4: differences only
    W = [2.9, 3.7, 3.7, 1.5, 1.4]
    diff = [r for s in ev.stages if s.n in (1, 2, 3, 4) for r in s.rows if r.get("status") not in ("equal", None)]
    drows = []
    for r in diff:
        if "element" in r:
            drows.append([C(f"SER of {_cap(r['element'])}"), C(r.get("A reference", "")), C(r.get("B reference", "")),
                          C(ts.fmt_energy(r.get("max |Δ| (J/mol)", math.nan))), C(r.get("unary50", ""))])
        else:
            st = r["status"] if r["status"] != "differs" else ts.fmt_delta(r.get("max |Δ|"), r["_ptype"])
            drows.append([C(r["parameter"]), C(clip(r["A"], 160) or "—"), C(clip(r["B"], 160) or "—"), C(st),
                          C(r.get("unary50", ""))])
    if not drows:
        drows = [[C("no differences in stages 1-4", span=5, muted=True)]]
    drows = fit_rows(drows, W, 11, lambda k: [C(f"+{k} more rows: Excel sheets 1-4", span=5, muted=True)])
    pu.add_flow_table_slides(prs, "Stages 1-4: differences only", ["Parameter", "A", "B", "max |Δ| / status",
                                                                    "unary50"], drows, W, font_size=11,
                             subtitle="numerically different or missing parameters (298-3000 K; SER 298-6000 K)")
    # stage 5
    W5 = [2.0, 3.6, 3.6, 3.6]
    prow = []
    for r in ev.stages[4].rows:
        if r.get("new end-member"):
            prow.append([C(r["phase"]), C(f"new end-member {r['new end-member']} in {r['in space of']}'s space"),
                         C(r.get("affects subsystems", "")),
                         C(r.get("provided by", ""), fill="differs" if "no G" in r.get("provided by", "") else None)])
        elif r.get("new phase in space of"):
            prow.append([C(r["phase"]), C(f"new phase in {r['new phase in space of']}'s space"),
                         C(r.get("model there", "")), C("")])
        elif r.get("name candidate"):
            prow.append([C(r["phase"]), C("name candidate (not merged)"), C(r["name candidate"], span=2)])
        elif r.get("issue") or r.get("merged model") and r.get("new end-members"):
            prow.append([C(r["phase"]), C(f"A {r.get('A model', '')}\nB {r.get('B model', '')}"),
                         C(r.get("merged model", "") and f"merged {r['merged model']}"),
                         C(r.get("issue", ""), fill="differs" if "cannot" in r.get("issue", "") else None)])
    if not prow:
        prow = [[C("shared phases have identical models; no new phases", span=4, muted=True)]]
    prow = fit_rows(prow, W5, 11, lambda k: [C(f"+{k} more rows: Excel sheet '5 phase models'", span=4,
                                               muted=True)])
    pu.add_flow_table_slides(prs, "Stage 5: phase models", ["Phase", "Change", "Where / merged model", "Note"],
                             prow, W5, font_size=11)
    # stage 6: what to replot
    W6 = [0.9, 2.0, 9.9]
    rp = []
    for side in ("A", "B"):
        rs = [r for r in ev.stages[5].rows if r["database"] == side]
        for r in rs:
            if r["replot"] == "yes":
                rp.append([C(side), C(r["subsystem"], fill="error"), C(clip(r["reasons"], 260, "Excel"))])
        same = [r["subsystem"] for r in rs if r["replot"] == "no"]
        if same:
            rp.append([C(side), C("unchanged", fill="same"), C(", ".join(same))])
    rp = fit_rows(rp, W6, 11, lambda k: [C(f"+{k} more rows: Excel sheet '6 replot'", span=3, muted=True)])
    pu.add_flow_table_slides(prs, "Stage 6: subsystems to replot (A's values win conflicts)",
                             ["DB", "Subsystem", "Reasons"], rp, W6, font_size=11)
    color_status_cells(prs, 1)
    pu.save(prs, path)
    return path


def run(a_path: str, b_path: str, out_prefix: str, **kw) -> Evaluation:
    work = Path(out_prefix).parent / (Path(out_prefix).name + "_work")
    ev = evaluate_pair(a_path, b_path, workdir=str(work), **kw)
    ev.files["xlsx"] = write_xlsx(ev, out_prefix + ".xlsx")
    ev.files["pptx"] = write_pptx(ev, out_prefix + ".pptx")
    if ev.normalised_b:
        ev.files["normalised B"] = ev.normalised_b
    if ev.rescaled_a:
        ev.files["rescaled A"] = ev.rescaled_a
    return ev


# =============================================================================================== tool
def register(tb) -> None:
    S = {"type": "string"}

    @tb.add("evaluate_compatibility",
            "Staged check: can TDB add_tdb (B) be merged into base_tdb (A)? B is first normalised to "
            "A's phase names / sublattice order / formula units (a magnetic 1-atom/fu B phase instead keeps its "
            "ratios and A's non-magnetic one is rescaled), then 1 SGTE unaries (vs SGTE unary50 if they "
            "differ), 2 other unaries, 3 binaries, 4 ternary+, 5 phase models (merged model, new end-members, "
            "magnetic/ordering), 6 which binaries/ternaries change and need replotting. Short verdicts; details "
            "-> Excel + PPTX. Then staged_merge.",
            {"base_tdb": S, "add_tdb": S,
             "phase_map": {"type": "object", "description": "{B phase: A phase} for name candidates the user "
                                                            "confirmed"},
             "sublattice_map": {"type": "object", "description": "{phase: [A position of B's sublattice i, ...]} or "
                                                                 "{phase: 'keep'} when the tool asks for it"},
             "unary50": {**S, "description": "default 'auto' (cached; downloaded from SGTE only if unaries "
                                             "differ) - leave it unless the user asks; 'off' or a local path"},
             "rtol": {"type": "number"}, "output": S}, ["base_tdb", "add_tdb"])
    def evaluate_compatibility(base_tdb, add_tdb, phase_map=None, sublattice_map=None, unary50: str = "auto",
                               rtol: float = ts.STAGE_RTOL, output: Optional[str] = None):
        a, b = tb._find_tdb(base_tdb), tb._find_tdb(add_tdb)
        prefix = tb.outp(output, f"compat_{Path(a).stem}+{Path(b).stem}")
        prefix = str(Path(prefix).with_suffix("")) if Path(prefix).suffix in (".xlsx", ".pptx") else prefix
        ev = run(a, b, prefix, phase_map=phase_map, sublattice_map=sublattice_map, unary50=unary50, rtol=rtol)
        return summary_text(ev, tb.rel), []
