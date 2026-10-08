"""Reports of ``combine_databases``: one deck, one Excel workbook, ``combine_report.md`` and the header block of the
final TDB.  Inputs are the run's state dict (see combine.build_state) plus the in-memory staged results."""
from __future__ import annotations

import datetime as _dt
import json
import math
import os
import re
from pathlib import Path
from typing import Dict, List, Optional

from .compatibility import clip

WHO_ORDER = ("user", "agent", "policy", "rule")


def _f(x, fmt="{:.2f}", none="—") -> str:
    try:
        return none if x is None or not math.isfinite(float(x)) else fmt.format(float(x))
    except (TypeError, ValueError):
        return none


def verdict_line(counts: Dict[str, int]) -> str:
    return (f"{counts.get('reproduced', 0)} reproduced, {counts.get('changed, explained', 0)} changed (explained), "
            f"{counts.get('CHANGED, UNEXPLAINED', 0)} CHANGED (unexplained)")


def who_counts(log: List[Dict]) -> str:
    c = {w: sum(e["who"] == w for e in log) for w in WHO_ORDER}
    return ", ".join(f"{v} {w}" for w, v in c.items() if v)


def flagged(log: List[Dict]) -> List[Dict]:
    return [e for e in log if e["who"] in ("user", "agent") and (not e["reason"] or e["reason"] == "not given")]


def exact_line(exact: List[Dict]) -> str:
    if not exact:
        return "no normalised or rescaled phases"
    bad = [c for c in exact if not c.get("ok")]
    worst = max([c["max_rel"] for c in exact if math.isfinite(c.get("max_rel", math.nan))] or [math.nan])
    return (f"{len(exact)} normalised/rescaled phase(s): GM per mole of atoms identical, max relative error "
            f"{_f(worst, '{:.1e}')}" + (f"; NOT EXACT: {', '.join(c['phase'] for c in bad)}" if bad else ""))


def sanity_lines(s: Dict) -> List[str]:
    L = []
    ld = s.get("load", {})
    L.append("pycalphad loads it and builds every phase model" if ld.get("loaded") and not ld.get("errors") else
             f"pycalphad: {'; '.join(ld.get('errors', [])[:3]) or 'does not load'}")
    li = s.get("lint") or {}
    if li.get("counts"):
        c = li["counts"]
        L.append(f"check_tdb: ERROR {c.get('ERROR', 0)}, WARNING {c.get('WARNING', 0)}, INFO {c.get('INFO', 0)}")
    elif li.get("error"):
        L.append(f"check_tdb failed: {li['error']}")
    z = s.get("zero_g") or []
    L.append("new end-members without G: none" if not z else
             "new end-members without G (pycalphad uses 0): " + "; ".join(x["text"] for x in z[:4])
             + (f" (+{len(z) - 4})" if len(z) > 4 else ""))
    cv = s.get("convergence") or {}
    if cv.get("points"):
        L.append(f"equilibrium grids on F: {cv['not_converged']} of {cv['points']} points did not converge")
    return L


def cannot_prove_lines(ex: Dict, vt) -> List[str]:
    L = []
    for b in ex.get("binaries", []):
        L.append(f"{'-'.join(e.capitalize() for e in sorted(b))}: binary no source describes (only the unaries and "
                 f"whatever phase models meet there) - predicted, not validated")
    for t in ex.get("ternaries", []):
        L.append(f"{'-'.join(e.capitalize() for e in sorted(t))}: ternary no source describes - an extrapolation "
                 f"of assessments that never saw each other's data" + (
                     f"; isotherms of F at {', '.join(f'{x:g}' for x in vt)} °C shown as predicted, not validated"
                     if vt else "; give validate_temperatures_c to map F's isotherms as predictions"))
    if ex.get("higher"):
        L.append(ex["higher"])
    return L


# =============================================================================================== TDB header
def tentative_terms(state: Dict) -> List[Dict]:
    """Tentative parameters of the combined TDB: [{'label', 'ref', 'value', 'id' (decision-log entry)}]."""
    from .tdbparse import is_tentative, parse_tdb
    from .tentative import ref_key
    final = state.get("run_final") or state.get("final")
    if not final or not Path(final).exists():
        return []
    F = parse_tdb(final)
    ids = {e["subject"]: e["id"] for e in state.get("log", []) if e.get("kind") == "tentative"}
    out = []
    for p in F.parameters:
        if is_tentative(p, F):
            out.append({"label": p.key.label(), "ref": ref_key(p), "value": p.expr.short(30),
                        "id": ids.get(p.key.label(), "")})
    return out


def tentative_summary(state: Dict, V: Optional[Dict] = None) -> List[str]:
    """Short lines: the tentative terms in the combined TDB, replacements, how close they bring each source."""
    terms = tentative_terms(state)
    rep_ = [e for e in state.get("log", []) if e.get("kind") == "tentative-replaced"]
    if not terms and not rep_:
        return []
    L = []
    if terms:
        L.append(f"Tentative (not assessed, {len(terms)}): " + "; ".join(f"{t['label']} = {t['value']} ({t['ref']})"
                                                                         for t in terms[:4])
                 + (f"; +{len(terms) - 4} more" if len(terms) > 4 else ""))
    if rep_:
        L.append("Tentative replaced by assessed terms: " + "; ".join(e["subject"] for e in rep_[:4])
                 + (f"; +{len(rep_) - 4} more" if len(rep_) > 4 else ""))
    for r in (V or {}).get("rows", []):
        if r.get("tentative_note"):
            L.append(f"  {r['source']}: {clip(r['tentative_note'], 230)}")
    return L


def tdb_header(state: Dict, V: Optional[Dict]) -> List[str]:
    now = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    L = ["=" * 100, f"Combined by calphad_tools combine_databases on {now}", "Sources (merge order):"]
    for i, s in enumerate(state["sources"]):
        L.append(f"  {s['label']} {'base ' if i == 0 else f'step{i} '}{Path(s['path']).name}  "
                 f"({', '.join(e.capitalize() for e in s['elements'])})  sha256 {s['sha256'][:16]}")
    L.append("Order: " + " | ".join(state["order_reasons"]))
    L.append("Decisions (id -> chosen; who: reason):")
    for e in state["log"]:
        if e["who"] != "rule":
            L.append(f"  {e['id']} -> {e['chosen']}; {e['who']}: {e['reason']}")
    rules = [e for e in state["log"] if e["who"] == "rule"]
    L.append(f"  + {len(rules)} automatic rule entries (unaries, normalisation, constituent unions, ...): see the "
             f"report")
    terms = tentative_terms(state)
    if terms:
        L.append(f"Tentative parameters ({state.get('tentative_settings', '')}): NOT assessed, reference keys "
                 f"TENT_MIX / TENT_VA - replace with assessed values:")
        L += [f"  {t['label']} = {t['value']} ({t['ref']}; {t['id'] or 'carried from a source'})" for t in terms[:30]]
    if V:
        L.append("Validation against every source: " + verdict_line(V["counts"]))
        for r in V["rows"]:
            if r["verdict"] != "reproduced":
                L.append(f"  {r['subsystem']} ({r['source']}): {r['verdict']}"
                         + (f" by {', '.join(r['explained_by'][:4])}" if r["explained_by"] and not r["unexplained"]
                            else ""))
        cp = cannot_prove_lines(V["extrapolations"], state.get("validate_temperatures_c") or [])
        if cp:
            L.append("Not validated (no source describes them): " + "; ".join(x.split(":")[0] for x in cp[:12]))
    L.append("Full report: combine_report.md next to this file (decision log, validation, figures)")
    L.append("=" * 100)
    return [clip(x.replace("!", "."), 196) for x in L]


def write_final_tdb(run_final: str, path: str, header: List[str]) -> str:
    text = Path(run_final).read_text()
    Path(path).write_text("\n".join(f"$ {h}" for h in header) + "\n" + text)
    return path


# =============================================================================================== deck
def write_deck(state: Dict, V: Dict, results: List[Dict], path: str) -> str:
    from . import pptx_utils as pu
    from .pptx_utils import GridCell as C
    from .staged_merge import add_staged_slides
    prs = pu.new_presentation()
    names = [Path(s["path"]).stem for s in state["sources"]]
    c = V["counts"]
    unexpl = [r for r in V["rows"] if r["verdict"] == "CHANGED, UNEXPLAINED"]
    expl = [r for r in V["rows"] if r["verdict"] == "changed, explained"]
    log = state["log"]
    flag = flagged(log)
    rows = [
        [C("Inputs", bold=True), C("\n".join(f"{s['label']} {Path(s['path']).name} "
                                            f"({', '.join(e.capitalize() for e in s['elements'])})"
                                            for s in state["sources"]))],
        [C("Base and order", bold=True), C("\n".join(clip(x, 170) for x in state["order_reasons"]))],
    ] + [[C("Note", bold=True), C(clip(w, 300), fill="error")] for st in state["steps"]
         for w in st.get("warnings", [])] + [
        [C("Verdicts", bold=True), C(verdict_line(c), bold=True,
                                     fill="differs" if c.get("CHANGED, UNEXPLAINED") else "same")],
        [C("CHANGED, UNEXPLAINED", bold=True), C("; ".join(f"{r['subsystem']} ({r['source']}): "
                                                           f"{clip(r['basis'], 120)}" for r in unexpl[:4]) or "none",
                                                 fill="differs" if unexpl else None)],
        [C("Changed, explained", bold=True), C("; ".join(f"{r['subsystem']} ({r['source']}) by "
                                                         f"{', '.join(r['explained_by'][:3])}" for r in expl[:5])
                                               or "none")],
        [C("Decision log", bold=True), C(f"{len(log)} entries ({who_counts(log)})"
                                         + (f"; reason NOT GIVEN for {', '.join(e['id'] for e in flag[:4])}" if flag
                                            else "; every user/agent decision has a reason"),
                                         fill="error" if flag else None)],
        [C("Exact transformations", bold=True), C(exact_line(state.get("exactness", [])))],
        [C("Combined database", bold=True), C("; ".join(sanity_lines(V["sanity"])))],
        [C("Cannot be proven", bold=True), C("; ".join(x.split(" - ")[0] for x in cannot_prove_lines(
            V["extrapolations"], state.get("validate_temperatures_c") or [])) or "nothing: every subsystem of F is "
                                                                                 "described by a source")],
        [C("Final TDB", bold=True), C(Path(state["final"]).name + f"  (hash seed {state.get('hash_seed', '?')})")],
    ]
    pu.add_flow_table_slides(prs, "Combined: " + " + ".join(names), ["Item", "Result"], rows, [2.5, 10.3],
                             font_size=12,
                             subtitle="black = source, red dashed = combined; reproduced = identical description or "
                                      "within ΔT 0.5 K, Δx 0.001, 99.5 % grid, equal H/S/Cp/G",
                             max_height_in=5.8)
    # plan
    W = [2.9, 1.1, 1.1, 1.9, 1.5, 4.4]
    prow = []
    for i, s in enumerate(state["inputs_by_order"]):
        prow.append([C(s["name"], bold=i == 0), C(", ".join(e.capitalize() for e in s["elements"])),
                     C(str(s["n_unary"]), align="c"),
                     C(", ".join("-".join(e.capitalize() for e in b) for b in s["binaries"]) or "—"),
                     C("—" if s["unary50"] is None else f"{s['unary50']}/{len(s['elements'])}", align="c"),
                     C(clip(state["order_reasons"][i], 200))])
    prow.append(("band", "Pairwise compatibility (evaluate_compatibility)"))
    for p in state.get("pairs", []):
        prow.append([C(p["pair"], span=2), C(clip(p["overall"], 300), span=4)])
    prow.append(("band", "Steps and their decisions"))
    for st in state["steps"]:
        prow.append([C(f"step {st['k']}: {st['pair']}", span=2),
                     C(clip("; ".join(f"{n['id']} = {n['answer']} ({n['who']})" for n in st["needed"]) or
                            "no decisions needed (staged rules only)", 300), span=4)])
    pu.add_flow_table_slides(prs, "Plan: base, order and decisions", ["Source", "Elements", "Unary params",
                                                                       "Binaries assessed", "SER = unary50",
                                                                       "Role / reason"], prow, W, font_size=11)
    # staged steps
    for st, r in zip(state["steps"], results):
        from . import tdbstage as ts
        norm = [n for n in r["notes"] if n.action in ts.NORM_CHANGES]
        note = clip("Normalised: " + "; ".join(n.text() for n in norm), 125, "Excel") if norm else ""
        add_staged_slides(prs, r["steps"], r["plots"], st["a_label"], st["b_label"], r["tol_K"],
                          Path(r["files"]["xlsx"]).parent / "figures", note, prefix=f"Step {st['k']} · ")
    # decision log
    WD = [2.6, 1.4, 1.55, 0.7, 2.5, 3.05, 1.1]
    drows = []
    cur = None
    for e in log:
        if e["step"] != cur:
            cur = e["step"]
            drows.append(("band", cur))
        fl = e["who"] in ("user", "agent") and (not e["reason"] or e["reason"] == "not given")
        drows.append([C(e["id"]), C(clip(e["subject"], 60)), C(clip(str(e["chosen"]), 90)),
                      C(e["who"], align="c", fill="differs" if fl else ("ref" if e["who"] != "rule" else None)),
                      C(clip(e["reason"], 170) if not fl else "NOT GIVEN", fill="differs" if fl else None,
                        bold=fl),
                      C(clip("; ".join(e["evidence"]), 230) or "—"),
                      C(clip(", ".join(e["affected"]), 70) or "—")])
    pu.add_flow_table_slides(prs, "Decision log", ["ID", "Subject", "Chosen", "Who", "Why", "Evidence", "Affects"],
                             drows, WD, font_size=10, max_height_in=5.7,
                             subtitle="every decision, including the automatic rules; full text in the Excel sheet "
                                      "'Decision log'")
    # validation per source
    for s in state["sources"]:
        png = V["overlays"].get(s["stem"])
        if not png:
            continue
        rs = [r for r in V["rows"] if r["source"] == s["stem"]]
        cc = {v: sum(r["verdict"] == v for r in rs) for v in ("reproduced", "changed, explained",
                                                               "CHANGED, UNEXPLAINED")}
        title = f"Validation: {s['stem']} reproduced in the combined database?"
        sub = f"{verdict_line(cc)} - black {s['stem']}, red dashed combined"
        if len(rs) > 2:
            pu.add_image_slide(prs, title, png, subtitle=sub)
            continue
        _image_with_details(prs, title, sub, png, rs)
    WV = [1.75, 1.75, 1.05, 1.55, 1.1, 1.0, 1.0, 3.7]
    vrows = []
    for r in V["rows"]:
        fill = {"reproduced": "same", "changed, explained": "error", "CHANGED, UNEXPLAINED": "differs"}[r["verdict"]]
        vrows.append([C(r["verdict"], bold=True, fill=fill), C(r["source"]), C(r["subsystem"]),
                      C(f"ΔT {_f(r['inv_max_dT'])} K, Δx {_f(r['inv_max_dx'], '{:.4f}')}"
                        + (f", {r['inv_unmatched']} unmatched" if r.get("inv_unmatched") else "")),
                      C(_f(r["boundary"], "{:.4f}")), C(_f(100 * r["grid"] if r["grid"] == r["grid"] else math.nan,
                                                          "{:.1f} %")),
                      C("equal" if r.get("props_equal") else "differ"),
                      C(clip("; ".join(r["unexplained"]) if r["unexplained"] else
                             (", ".join(r["explained_by"]) or r["basis"]), 220))])
    pu.add_flow_table_slides(prs, "Validation per subsystem (unexplained first)",
                             ["Verdict", "Source", "Subsystem", "Invariants", "Boundaries", "Grid", "H/S/Cp/G",
                              "Explained by / unexplained"], vrows, WV, font_size=11,
                             subtitle="boundaries: max distance (x; 1 K = 0.002); grid: points with the same stable "
                                      "phases")
    # property check
    WP = [2.0, 1.2, 2.3, 1.45, 1.45, 1.45, 1.45, 1.6]
    agg: Dict = {}
    for p in V["props"]:
        a = agg.setdefault((p["source"], p["subsystem"]), {"n": 0, "eq": 0, **{k: 0.0 for k in
                                                                               ("GM", "HM", "SM", "CPM")}})
        a["n"] += 1
        a["eq"] += bool(p["equal"])
        for k in ("GM", "HM", "SM", "CPM"):
            d = p.get(f"Δ{k}", math.nan)
            a[k] = max(a[k], abs(d)) if d == d else math.inf
    prows = []
    for (src, sub), a in agg.items():
        prows.append([C(src), C(sub), C(f"{a['eq']} of {a['n']} equal", fill="same" if a["eq"] == a["n"] else "error"),
                      C(_f(a["GM"], "{:.3g}")), C(_f(a["HM"], "{:.3g}")), C(_f(a["SM"], "{:.3g}")),
                      C(_f(a["CPM"], "{:.3g}")),
                      C(next((r["verdict"] for r in V["rows"] if r["source"] == src and r["subsystem"] == sub), ""))])
    pu.add_flow_table_slides(prs, "Property spot check: stable assemblage in source vs combined",
                             ["Source", "Subsystem", "Points", "max |ΔG| J/mol", "max |ΔH| J/mol",
                              "max |ΔS| J/mol/K", "max |ΔCp| J/mol/K", "Verdict"],
                             prows or [[C("no property points", span=8, muted=True)]], WP, font_size=11,
                             subtitle="9 points per binary (x 0.25/0.5/0.75 at three temperatures), 3 per ternary "
                                      "isotherm; per point in Excel 'Properties'")
    # exactness + sanity
    srows = [("band", "Exact transformations (GM per mole of atoms, random T and site fractions)")]
    for x in state.get("exactness", []) or []:
        srows.append([C(f"step {x['step']}"), C(x["phase"]), C(x["transform"]),
                      C(_f(x.get("max_rel"), "{:.1e}") if math.isfinite(x.get("max_rel", math.nan)) else
                        x.get("error", "?"), fill="same" if x.get("ok") else "differs")])
    if not state.get("exactness"):
        srows.append([C("no normalised or rescaled phases", span=4, muted=True)])
    srows.append(("band", "Combined database"))
    for line in sanity_lines(V["sanity"]):
        srows.append([C(line, span=4)])
    for t in (V["sanity"].get("lint") or {}).get("top", [])[:5]:
        srows.append([C(t, span=4, muted=True)])
    pu.add_flow_table_slides(prs, "Exactness and sanity of the combined database", ["Step", "Phase", "Transform",
                                                                                     "Max rel. error"],
                             srows, [1.4, 3.6, 4.6, 3.2], font_size=11)
    # what cannot be proven
    cp = cannot_prove_lines(V["extrapolations"], state.get("validate_temperatures_c") or [])
    crow = [[C(x)] for x in cp] or [[C("Every binary and ternary of the combined database is described by a source.")]]
    crow.append([C("Proven here: every source's own subsystems are reproduced or every difference is traced to a "
                   "logged decision; normalisations are exact. Not proven: that the combination is a good "
                   "description where no source had data.", muted=True)])
    pu.add_flow_table_slides(prs, "What this cannot prove", ["Not validated by any source"], crow, [12.6],
                             font_size=13)
    for png in V.get("predicted", []):
        name = Path(png).stem.replace("F_", "", 1)
        m = re.match(r"(.+)_(-?[\d.]+)C$", name)
        name = f"{m.group(1)} at {m.group(2)} °C" if m else name
        pu.add_image_slide(prs, f"Predicted, not validated: {name}", png,
                           subtitle="the combined database where no source has data - an extrapolation")
    pu.save(prs, path)
    return path


def row_details(r: Dict) -> str:
    L = [f"{r['subsystem']}: {r['verdict']}",
         ("Not reproduced: " if r["verdict"] == "changed, explained" else "") + clip(r["basis"], 300)]
    if r.get("n_inv") or r["inv_max_dT"] == r["inv_max_dT"]:
        L.append(f"Invariants: {r.get('n_inv', 0)} compared, max ΔT {_f(r['inv_max_dT'])} K, max Δx "
                 f"{_f(r['inv_max_dx'], '{:.4f}')}" + (f", {r['inv_unmatched']} without partner"
                                                       if r.get("inv_unmatched") else ""))
    L.append(f"Boundaries: max distance {_f(r['boundary'], '{:.4f}')} (x; 1 K = 0.002)"
             + (f", {r['boundary_unmatched']} set(s) in one diagram only" if r.get("boundary_unmatched") else ""))
    L.append(f"Grid: {_f(100 * r['grid'] if r['grid'] == r['grid'] else math.nan, '{:.1f}')} % of "
             f"{r.get('grid_n', 0)} points with the same stable phases"
             + (f" ({r['grid_nan_F']} point(s) did not converge in F)" if r.get("grid_nan_F") else ""))
    pm = r.get("props_max") or {}
    L.append("H/S/Cp/G of the stable assemblage: " + ("equal" if r.get("props_equal") else "differ") +
             (f" (max |ΔG| {_f(pm.get('GM'), '{:.3g}')} J/mol, |ΔH| {_f(pm.get('HM'), '{:.3g}')} J/mol)"
              if pm else ""))
    if r["explained_by"]:
        L.append("Explained by: " + ", ".join(r["explained_by"]))
    if r.get("tentative_note"):
        L.append("Tentative: " + clip(r["tentative_note"], 300))
    if r["unexplained"]:
        L.append("UNEXPLAINED: " + "; ".join(clip(x, 160) for x in r["unexplained"][:4]))
    for c in r.get("changes", [])[:4]:
        L.append("  change: " + clip(c, 150))
    return "\n".join(L)


def _image_with_details(prs, title: str, sub: str, png: str, rows: List[Dict]) -> None:
    from PIL import Image
    from pptx.util import Inches
    from . import pptx_utils as pu
    s = pu._blank(prs)
    pu.add_title(s, title, sub)
    top, box_w, box_h = Inches(1.55), Inches(7.4), Inches(5.75)
    with Image.open(png) as im:
        w, h = im.size
    sc = min(box_w / w, box_h / h)
    s.shapes.add_picture(str(png), Inches(0.35), int(top), int(w * sc), int(h * sc))
    pu._textbox(s, Inches(7.95), top, Inches(5.1), Inches(5.6), "\n\n".join(row_details(r) for r in rows), size=12)


# =============================================================================================== Excel
def write_xlsx(state: Dict, V: Dict, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    red, amber, green = (PatternFill("solid", fgColor=c) for c in ("FFCDD2", "FFE0B2", "C8E6C9"))

    def sheet(name, header, rows, widths=None, first=False):
        w = wb.active if first else wb.create_sheet(name[:31])
        if first:
            w.title = name[:31]
        w.append(header)
        for c in w[1]:
            c.font = Font(bold=True)
        for r in rows:
            w.append([("" if isinstance(x, float) and not math.isfinite(x) else
                       round(x, 9) if isinstance(x, float) else
                       json.dumps(x) if isinstance(x, (list, dict)) else x) for x in r])
        for i, wd in enumerate(widths or [], 1):
            w.column_dimensions[w.cell(1, i).column_letter].width = wd
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        return w
    summ = [["Inputs", "; ".join(f"{s['label']} {Path(s['path']).name}" for s in state["sources"])],
            ["Order", " | ".join(state["order_reasons"])], ["Verdicts", verdict_line(V["counts"])],
            ["Decision log", f"{len(state['log'])} entries ({who_counts(state['log'])})"],
            ["Exactness", exact_line(state.get("exactness", []))]]
    summ += [["Combined database", x] for x in sanity_lines(V["sanity"])]
    summ += [["Cannot be proven", x] for x in cannot_prove_lines(V["extrapolations"],
                                                                 state.get("validate_temperatures_c") or [])]
    summ += [["Final TDB", state["final"]], ["Hash seed", state.get("hash_seed", "")]]
    sheet("Summary", ["Item", "Value"], summ, [22, 160], first=True)
    sheet("Plan", ["Source", "Elements", "Unary parameters", "Binaries assessed", "Ternaries described",
                   "SER = unary50", "Role / reason"],
          [[s["name"], ", ".join(s["elements"]), s["n_unary"], ", ".join("-".join(b) for b in s["binaries"]),
            ", ".join("-".join(t) for t in s["ternaries"]), "" if s["unary50"] is None else s["unary50"], why]
           for s, why in zip(state["inputs_by_order"], state["order_reasons"])], [30, 14, 10, 30, 20, 10, 80])
    sheet("Decisions asked", ["ID", "Step", "Kind", "Question", "Options (consequences)", "Answer", "Who", "Reason"],
          [[n["id"], n["step"], n["kind"], n["question"], "\n".join(f"{k}: {v}" for k, v in n["options"]),
            n["answer"], n["who"], n["reason"]] for st in state["steps"] for n in st["needed"]],
          [26, 6, 12, 50, 60, 10, 8, 50])
    w = sheet("Decision log", ["ID", "Step", "Subject", "Options", "Chosen", "Who", "Reason", "Evidence",
                               "Affected subsystems", "Files / figures", "Kind"],
              [[e["id"], e["step"], e["subject"], "\n".join(e["options"]), e["chosen"], e["who"], e["reason"],
                "\n".join(e["evidence"]), ", ".join(e["affected"]), "\n".join(e["refs"]), e["kind"]]
               for e in state["log"]], [30, 22, 18, 40, 18, 8, 50, 70, 24, 40, 12])
    for row in w.iter_rows(min_row=2):
        if row[5].value in ("user", "agent") and row[6].value in ("", "not given"):
            row[6].fill = red
    w = sheet("Validation", ["Verdict", "Source", "Subsystem", "Basis", "Explained by", "Unexplained", "Changes",
                             "Invariants max |ΔT| (K)", "Invariants max Δx", "Invariants unmatched",
                             "Boundary max distance", "Boundary sets unmatched", "Grid agreement", "Grid points",
                             "Grid differences", "Properties equal", "Overlay", "Source diagram", "Combined diagram"],
              [[r["verdict"], r["source"], r["subsystem"], r["basis"], ", ".join(r["explained_by"]),
                "\n".join(r["unexplained"]), "\n".join(r["changes"]), r["inv_max_dT"], r["inv_max_dx"],
                r.get("inv_unmatched", 0), r["boundary"], r.get("boundary_unmatched", 0), r["grid"],
                r.get("grid_n", 0), "\n".join(r.get("grid_examples", [])), r.get("props_equal"), r.get("overlay", ""),
                r.get("S_png", ""), r.get("F_png", "")] for r in V["rows"]],
              [22, 22, 14, 60, 40, 50, 60, 10, 10, 8, 10, 8, 10, 8, 50, 8, 40, 40, 40])
    for row in w.iter_rows(min_row=2):
        row[0].fill = {"reproduced": green, "changed, explained": amber}.get(row[0].value, red)
    inv_rows = []
    for r in V["rows"]:
        for x in r.get("inv_rows", []):
            inv_rows.append([r["source"], r["subsystem"], x["invariant"], x.get("type", ""), x["T_old_K"],
                             x["T_new_K"], x["dT_K"], x["dx"], x["status"]])
    sheet("Invariants", ["Source", "Subsystem", "Invariant", "Type", "T source (K)", "T combined (K)", "ΔT (K)", "Δx",
                         "Status"], inv_rows, [22, 12, 36, 14, 12, 12, 10, 10, 12])
    keys = ["x", "T_K", "S phases", "F phases"] + [f"{a} {p}" for p in ("GM", "HM", "SM", "CPM") for a in ("S", "F")] \
        + [f"Δ{p}" for p in ("GM", "HM", "SM", "CPM")] + ["equal"]
    sheet("Properties", ["Source", "Subsystem"] + keys,
          [[p["source"], p["subsystem"]] + [("/".join(f"{v:.4g}" for v in p[k]) if k == "x" else p.get(k))
                                            for k in keys] for p in V["props"]], [22, 12, 10, 10, 24, 24] + [12] * 13)
    sheet("GM identity", ["Step", "Files", "Phase", "Transform", "Max relative error", "States", "Exact", "Error"],
          [[x["step"], x["file"], x["phase"], x["transform"], x.get("max_rel"), x.get("n"), x.get("ok"),
            x.get("error", "")] for x in state.get("exactness", [])], [6, 50, 24, 24, 14, 8, 8, 40])
    srows = [[x] for x in sanity_lines(V["sanity"])] + [[x] for x in (V["sanity"].get("lint") or {}).get("top", [])]
    srows += [[f"zero G: {z['text']}"] for z in V["sanity"].get("zero_g", [])]
    srows += [[x] for x in cannot_prove_lines(V["extrapolations"], state.get("validate_temperatures_c") or [])]
    sheet("Sanity", ["Finding"], srows, [160])
    sheet("Attribution", ["Parameter / phase model", "Decision id", "Subsystems"],
          [[k, i, ", ".join(subs)] for part in ("params", "models") for k, v in state["attr"][part].items()
           for i, subs in v.items()], [50, 44, 40])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# =============================================================================================== Markdown
def write_md(state: Dict, V: Dict, files: Dict[str, str], path: str) -> str:
    base = Path(path).parent

    def rel(p) -> str:
        try:
            return os.path.relpath(str(p), str(base))
        except ValueError:
            return str(p)

    def esc(t) -> str:
        return str(t).replace("|", "\\|").replace("\n", " ")
    c = V["counts"]
    L = [f"# Combined database: {' + '.join(Path(s['path']).stem for s in state['sources'])}", "",
         f"**Verdicts:** {verdict_line(c)}." + (" **Some differences are not explained - see below.**"
                                                if c.get("CHANGED, UNEXPLAINED") else ""), "",
         f"- Final TDB: [{Path(state['final']).name}]({rel(state['final'])}) (header block: sources, order, "
         f"decisions, verdicts)", f"- Deck: [{Path(files['deck']).name}]({rel(files['deck'])}); workbook: "
                                    f"[{Path(files['xlsx']).name}]({rel(files['xlsx'])}); reproduction notebook: "
                                    f"[{Path(files['notebook']).name}]({rel(files['notebook'])}) "
                                    f"(hash seed {state.get('hash_seed', '?')})", ""]
    L += ["## Inputs, base and order", ""]
    for s, why in zip(state["inputs_by_order"], state["order_reasons"]):
        L.append(f"- `{s['name']}` ({', '.join(e.capitalize() for e in s['elements'])}; sha256 "
                 f"`{s['sha256'][:16]}`): {why}")
    for p in state.get("pairs", []):
        L.append(f"- compatibility {p['pair']}: {p['overall']}")
    for st in state["steps"]:
        L += [f"- **Note:** {w}" for w in st.get("warnings", [])]
    L.append("")
    L += ["## Decision log", "", f"{len(state['log'])} entries ({who_counts(state['log'])}). Each entry says who "
                                 f"chose (user / agent / policy / rule), why, and the evidence.", ""]
    fl = flagged(state["log"])
    if fl:
        L += [f"**Reason not given** for: {', '.join('`' + e['id'] + '`' for e in fl)} - ask whoever chose.", ""]
    L += ["| ID | Subject | Chosen | Who | Why | Evidence | Affects |", "|---|---|---|---|---|---|---|"]
    for e in state["log"]:
        ev = "; ".join(e["evidence"][:4]) + (f" (+{len(e['evidence']) - 4})" if len(e["evidence"]) > 4 else "")
        refs = " ".join(f"[fig]({rel(x)})" for x in e["refs"][:2])
        L.append(f"| `{e['id']}` | {esc(e['subject'])} | {esc(clip(str(e['chosen']), 80))} | {e['who']} | "
                 f"{esc(e['reason']) if e['reason'] else 'NOT GIVEN'} | {esc(clip(ev, 400))} {refs} | "
                 f"{esc(', '.join(e['affected']))} |")
    L.append("")
    terms = tentative_terms(state)
    trep = [e for e in state["log"] if e.get("kind") == "tentative-replaced"]
    if terms or trep:
        L += ["## Tentative parameters", "",
              f"Settings: {state.get('tentative_settings', '')}. Interaction terms the combination added where the "
              f"combined model creates a mixing no source describes - rule 1 (TENT_MIX): a new pair on a "
              f"substitutional sublattice of an intermetallic, L0 = the rule-1 value per mole of atoms x atoms per "
              f"formula unit; rule 2 (TENT_VA): an element with VA on a substitutional sublattice, L0(X,VA) = the "
              f"rule-2 value as written (per formula unit). They are **not "
              f"assessed**: the reference list of the TDB marks them (TENT_MIX / TENT_VA), and an assessed term "
              f"always replaces them in any later merge.", ""]
        if terms:
            L += ["| Parameter | Value | Reference | Decision-log entry |", "|---|---|---|---|"]
            L += [f"| `{t['label']}` | {esc(t['value'])} | {t['ref']} | `{t['id'] or 'carried from a source'}` |"
                  for t in terms]
            L.append("")
        for e in trep:
            L.append(f"- replaced: `{e['subject']}` {esc(e['chosen'])} (`{e['id']}`)")
        notes = [r for r in V["rows"] if r.get("tentative_note")]
        if notes:
            L += ["", "How close the tentative terms bring each affected source subsystem:", ""]
            L += [f"- {r['source']} {r['subsystem']}: {esc(r['tentative_note'])}" for r in notes]
        L.append("")
    L += ["## Steps (staged merges)", ""]
    for st in state["steps"]:
        L.append(f"- step {st['k']}: {st['pair']} -> [{Path(st['Q']).name}]({rel(st['Q'])}); deck "
                 f"[{Path(st['deck']).name}]({rel(st['deck'])}), workbook [{Path(st['xlsx']).name}]({rel(st['xlsx'])})")
        for img in st.get("images", [])[:6]:
            L.append(f"  - ![{Path(img).stem}]({rel(img)})")
    L.append("")
    L += ["## Validation against every source", "",
          "Each source's binaries (and ternaries at the given temperatures) were mapped from the source and from the "
          "combined database (mapping check on); black = source, red dashed = combined.", "",
          "| Verdict | Source | Subsystem | Invariants max ΔT / Δx | Boundaries | Grid | H/S/Cp/G | Explained by / "
          "unexplained | Overlay |", "|---|---|---|---|---|---|---|---|---|"]
    for r in V["rows"]:
        why = "; ".join(r["unexplained"]) if r["unexplained"] else (", ".join(r["explained_by"]) or r["basis"])
        L.append(f"| **{r['verdict']}** | {r['source']} | {r['subsystem']} | {_f(r['inv_max_dT'])} K / "
                 f"{_f(r['inv_max_dx'], '{:.4f}')} | {_f(r['boundary'], '{:.4f}')} | "
                 f"{_f(100 * r['grid'] if r['grid'] == r['grid'] else math.nan, '{:.1f} %')} | "
                 f"{'equal' if r.get('props_equal') else 'differ'} | {esc(clip(why, 300))} | "
                 f"[png]({rel(r.get('overlay', ''))}) |")
    L.append("")
    for stem, png in V["overlays"].items():
        L += [f"![{stem}]({rel(png)})", ""]
    L += ["## Exactness and sanity", "", f"- {exact_line(state.get('exactness', []))}"]
    for x in state.get("exactness", []):
        L.append(f"  - step {x['step']} {x['phase']} ({x['transform']}): max relative GM error "
                 f"{_f(x.get('max_rel'), '{:.1e}')}" + ("" if x.get("ok") else f" - {x.get('error', 'NOT EXACT')}"))
    L += [f"- {x}" for x in sanity_lines(V["sanity"])]
    L += [f"  - {x}" for x in (V["sanity"].get("lint") or {}).get("top", [])[:8]]
    lr = (V["sanity"].get("lint") or {}).get("report")
    if lr:
        L.append(f"- check_tdb report: [{Path(lr).name}]({rel(lr)})")
    L.append("")
    L += ["## What this proves and what it cannot", "",
          "**Proven** (within the stated tolerances): for every source, every binary (and ternary isotherm) it "
          "describes is either reproduced by the combined database - identical parameters and phase models, or the "
          "same invariants (ΔT ≤ 0.5 K, Δx ≤ 0.001), boundaries, ≥ 99.5 % of an equilibrium grid and equal H, S, "
          "Cp, G of the stable assemblage - or every parameter / phase-model difference is traced to a logged "
          "decision "
          "(the attribution of each merge step). Renames, sublattice reorders and rescales were checked to leave GM "
          "per mole of atoms unchanged (relative 1e-9) on the actual data.", "",
          "**Not proven**: that a changed subsystem is *better* (that is the decision's reason, not a test); that the "
          "diagrams are complete beyond the mapped range and grid resolution; anything about systems no source "
          "describes:", ""]
    L += [f"- {x}" for x in cannot_prove_lines(V["extrapolations"], state.get("validate_temperatures_c") or [])] \
        or ["- nothing: every binary and ternary of the combined database is described by a source"]
    for png in V.get("predicted", []):
        L += ["", f"![predicted, not validated: {Path(png).stem}]({rel(png)})"]
    L += ["", "## Files", ""] + [f"- {k}: [{Path(v).name}]({rel(v)})" for k, v in files.items()]
    L += ["", "## Reproduction", "",
          f"`{Path(files['notebook']).name}` rebuilds the combined TDB from the inputs with the same order and "
          f"decisions (standalone code: the TDB parser and merge planner are copied into the notebook) and checks it "
          f"is parameter-equivalent to the tool's file, recomputes the equilibrium-grid and property tables with "
          f"pycalphad, and regenerates the mapped overlays through calphad_tools (labelled cell). Hash seed: "
          f"{state.get('hash_seed', '?')}.", ""]
    Path(path).write_text("\n".join(L))
    return path
