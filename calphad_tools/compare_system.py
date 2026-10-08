"""N-way comparison of databases restricted to one element subsystem (``compare_system``), plus database discovery
(``find_databases``).

* Selection: for each database keep the phases that can exist with ``elements`` (+VA), the parameters whose
  constituents all lie in the subsystem (unaries, end-members of multi-sublattice phases, interactions) and the
  FUNCTIONs reachable from them.
* Every parameter is shown with all FUNCTION references inlined into one piecewise expression in T (union of the
  breakpoints, terms collected), so databases written differently compare term by term (:mod:`pwinline`).
* Phases are matched by exact name.  Unmatched phases that may be the same phase under another name are reported as
  *candidates* with structural evidence; they are merged only through an explicit ``phase_map``.
* Derived expressions (``"G(SIGMA,FE:CR:CR) - 10*GHSERFE - 20*GHSERCR"``) and formation energies of end-members
  relative to each database's own element reference functions.
* Output: one PowerPoint (summary, phase matrix, differences, rows numerically equal but written differently, all
  parameters, one table per database, expressions), an Excel workbook with the full expanded expressions, and a
  long-form CSV.  Every parameter carries its reference key as written (``6000 N 91Din !``) and, when the TDB's
  reference list has that key, the citation (short form on the slides, full text in Excel/CSV).
"""
from __future__ import annotations

import csv
import difflib
import fnmatch
import math
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np
import symengine as se

from . import pwinline as pwi
from .tdbparse import (TDB, ExpressionError, Piecewise, is_tentative, parse_param_name, parse_tdb, short_citation,
                       tokenize)
from .tdbtables import (ExprRow, ParamRow, PhaseView, Subsystem, database_functions, database_keys,  # noqa: F401
                        database_table, differences, equal_rows, expression_table, key_constituents, make_key,
                        make_label, phase_order, row_sort_key, tdb_to_se)

DEFAULT_TEMPS = [298.15, 1000.0, 1500.0]
DB_EXTS = (".tdb", ".dat")
SKIP_DIRS = {".venv", "venv", "__pycache__", "node_modules", "site-packages", ".git"}


# =============================================================================================== expressions
class DbView(Subsystem):
    """One database (a TDB file) restricted to an element subsystem, with every parameter's functions inlined."""

    def __init__(self, path: str, label: str, rel: str, elements: Sequence[str], rename: Dict[str, str]):
        self.path = path
        if Path(path).suffix.lower() == ".dat":
            raise ValueError(f"{Path(path).name} is a ChemSage .dat database; compare_system reads TDB text - "
                             "convert it first with convert_database")
        super().__init__(parse_tdb(path), label, rel, elements, rename)

    # ------------------------------------------------------------------ element references
    def ser_reference(self, el: str) -> Tuple[str, Any, str]:
        """(kind, what, description) of the lattice-stability reference used for element ``el``.

        kind "param" -> what = row key (the G parameter of the element in its ELEMENT reference phase),
        kind "function" -> what = FUNCTION name (GHSER<EL> style fallback)."""
        e = self.tdb.elements.get(el)
        if e is not None and e.ref_phase:
            cn = self.canon_of(e.ref_phase)
            if cn:
                pv = self.phases[cn]
                red = tuple((el,) for i in range(pv.n_sub) if i not in pv.va_only)
                key = make_key("G", cn, red, 0)
                row = self.rows.get(key)
                if row is not None and row.segs is not None:
                    rng = row.raw_pw.ranges if row.raw_pw is not None else []
                    m = re.fullmatch(r"\+?([A-Z][A-Z0-9_]*)#?", rng[0][2].upper()) if len(rng) == 1 else None
                    desc = f"{m.group(1)} (= {row.label})" if m and m.group(1) in self.tdb.functions else row.label
                    return "param", key, desc
        for cand in (f"GHSER{el}", f"GHSER{el}{el}", f"GHSER{el:_<2}".replace("_", "")):
            if cand in self.tdb.functions:
                return "function", cand, cand
        raise pwi.MissingRef(f"reference function for {el}")

    def reference_states(self) -> Dict[str, str]:
        """element -> "GHSERFE (= G(BCC_A2,FE:VA;0)), SER" or "..., NOT SER ...".

        An SER-type reference is an absolute Gibbs energy (it has a heat-capacity term c*T*ln(T) with c < -10);
        Kaufman-style lattice stabilities (e.g. G(BCC_A2,FE) = GBCCFE relative to fcc Fe) do not."""
        out = {}
        for el in self.elements:
            if el not in self.tdb.elements:
                continue
            try:
                kind, what, desc = self.ser_reference(el)
                segs = self.ref_segs(kind, what)
            except (pwi.MissingRef, ValueError, TypeError) as e:
                out[el] = f"no reference found ({e})"
                continue
            c = pwi.terms(segs[0][2])[0].get((1, 1), 0.0)
            out[el] = f"{desc}, " + ("SER" if c < -10 else
                                    f"NOT SER (T*ln(T) coef. {c:g}: lattice stability vs another state)")
        return out

    def ref_segs(self, kind: str, what: str) -> list:
        if kind == "param":
            return self.rows[what].segs
        return self.fn(what)


# =============================================================================================== discovery
def _walk(base: Path, recursive: bool = True):
    if not recursive:
        for p in sorted(base.iterdir()) if base.is_dir() else []:
            if p.is_file():
                yield p
        return
    for root, dirs, files in os.walk(base):
        dirs[:] = sorted(d for d in dirs if not d.startswith(".") and d not in SKIP_DIRS)
        for f in sorted(files):
            yield Path(root) / f


def database_files(base: Path, pattern: Optional[str] = None, recursive: bool = True) -> List[Path]:
    out = []
    for p in _walk(base, recursive):
        if pattern:
            if fnmatch.fnmatch(p.name.lower(), pattern.lower()) or fnmatch.fnmatch(str(p).lower(),
                                                                                     pattern.lower()):
                out.append(p)
        elif p.suffix.lower() in DB_EXTS:
            out.append(p)
    return out


_PARSE_CACHE: Dict[Tuple[str, float], TDB] = {}


def _parsed(path: Path) -> TDB:
    k = (str(path), path.stat().st_mtime)
    if k not in _PARSE_CACHE:
        _PARSE_CACHE[k] = parse_tdb(str(path))
    return _PARSE_CACHE[k]


def database_summary(path: Path, elements: Optional[Sequence[str]] = None) -> Dict[str, Any]:
    info: Dict[str, Any] = {"size_kb": round(path.stat().st_size / 1024, 1)}
    if path.suffix.lower() == ".dat":
        from .dbformats import dat_info
        d = dat_info(str(path))
        info.update(elements=d["elements"], n_phases=d["n_phases"], format="ChemSage .dat")
        return info
    db = _parsed(path)
    info.update(elements=db.real_elements, n_phases=len(db.phases), n_parameters=len(db.parameters))
    head = [c.strip("$ ").strip() for c in db.header_comments]
    head = [h for h in head if len(h) > 3 and not set(h) <= set("-=*_ ")]
    info["header"] = " | ".join(head[:2])[:140]
    if elements:
        want = {e.upper() for e in elements}
        counts: Dict[str, int] = {}
        for p in db.parameters:
            els = set()
            for sl in p.key.constituents:
                for s in sl:
                    if s in ("VA", "*"):
                        continue
                    els.add(s if s in db.elements else re.sub(r"[+-]\d*$", "", s))
            if els and els <= want:
                k = "-".join(sorted(els))
                counts[k] = counts.get(k, 0) + 1
        info["subsystem_parameters"] = dict(sorted(counts.items(), key=lambda kv: (kv[0].count("-"), kv[0])))
    return info


def find_databases(workdir: Path, folder: Optional[str] = None, pattern: Optional[str] = None,
                   elements: Optional[Sequence[str]] = None, recursive: bool = True) -> Tuple[List[Dict], str]:
    base = (workdir / folder) if folder else workdir
    if not base.exists():
        raise FileNotFoundError(f"folder {folder} not found under {workdir}")
    want = [e.upper() for e in elements] if elements else []
    rows, skipped, failed = [], 0, []
    for p in database_files(base, pattern, recursive):
        try:
            info = database_summary(p, want)
        except Exception as e:  # noqa: BLE001 - one unreadable file must not hide the rest
            failed.append(f"{_rel(p, workdir)}: {type(e).__name__}: {str(e)[:80]}")
            continue
        if want and not set(want) <= set(info.get("elements", [])):
            skipped += 1
            continue
        rows.append({"path": _rel(p, workdir), **info})
    lines = [f"{len(rows)} database(s) under {_rel(base, workdir) or '.'}"
             + (f" containing {'+'.join(want)}" if want else "")
             + (f" ({skipped} others lack some of these elements)" if skipped else "") + ":"]
    for r in rows:
        sub = r.get("subsystem_parameters")
        s = (" | params " + ", ".join(f"{k}:{v}" for k, v in sub.items())) if sub is not None else ""
        lines.append(f"- {r['path']}  [{','.join(r.get('elements', []))}]  {r.get('n_phases', '?')} phases, "
                     f"{r['size_kb']} kB{s}" + (f"\n    {r['header']}" if r.get("header") else ""))
    if failed:
        lines.append("unreadable: " + "; ".join(failed))
    lines.append("Use these exact paths in other tools.")
    return rows, "\n".join(lines)


def _rel(p: Path, workdir: Path) -> str:
    try:
        return str(Path(p).resolve().relative_to(Path(workdir).resolve()))
    except ValueError:
        return str(p)


# =============================================================================================== phase candidates
ALIASES = {"BCC": "bcc", "A2": "bcc", "FCC": "fcc", "A1": "fcc", "HCP": "hcp", "A3": "hcp", "SIGMA": "sigma",
           "D8B": "sigma", "SIG": "sigma", "LIQUID": "liq", "LIQ": "liq", "LAVES": "laves", "C14": "c14",
           "C15": "c15", "C36": "c36", "MU": "mu", "D85": "mu", "CHI": "chi", "A12": "chi", "CBCC": "chi",
           "CUB": "cub", "A13": "cub", "B2": "b2", "L12": "l12", "L1_2": "l12", "GAMMA": "gamma", "PRIME": "prime"}
PROTOTYPES = {"bcc", "fcc", "hcp", "sigma", "liq", "c14", "c15", "c36", "mu", "chi", "cub", "b2", "l12"}


def _fam(name: str) -> set:
    return {ALIASES.get(t, t.lower()) for t in re.split(r"[_#\-]", name.upper()) if t}


@dataclass
class Candidate:
    a: str
    b: str
    dbs_a: List[str]
    dbs_b: List[str]
    name_score: float
    evidence: List[str]
    verdict: str
    suggestion: Dict[str, str]

    def text(self) -> str:
        return (f"{self.a} [{', '.join(self.dbs_a)}] ~ {self.b} [{', '.join(self.dbs_b)}]: {self.verdict}; "
                + "; ".join(self.evidence) + f". If the same phase: phase_map {self.suggestion}")


def phase_candidates(views: List["DbView"], ref_label: str) -> List[Candidate]:
    presence: Dict[str, List[DbView]] = {}
    for v in views:
        for cn in v.phases:
            presence.setdefault(cn, []).append(v)
    names = sorted(presence)
    out = []
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            la, lb = {v.label for v in presence[a]}, {v.label for v in presence[b]}
            if la & lb:
                continue                     # one database has both: different phases
            fa, fb = _fam(a), _fam(b)
            jac = len(fa & fb) / max(1, len(fa | fb))
            ratio = difflib.SequenceMatcher(None, a, b).ratio()
            score = max(jac if fa & fb else 0.0, 0.8 * ratio)
            pa, pb = fa & PROTOTYPES, fb & PROTOTYPES
            proto_conflict = bool(pa and pb and not (pa & pb))
            ref_first = lambda vs: sorted(vs, key=lambda v: v.label != ref_label)   # noqa: E731
            PA = [v.phases[a] for v in ref_first(presence[a])]
            PB = [v.phases[b] for v in ref_first(presence[b])]
            A, B = PA[0], PB[0]
            ev = []
            ra, rb = {p.reduced_ratios() for p in PA}, {p.reduced_ratios() for p in PB}
            same_sites = bool(ra & rb)
            ta = "|".join(sorted({p.ratios_text() for p in PA}))
            tb_ = "|".join(sorted({p.ratios_text() for p in PB}))
            ev.append(f"sites {ta} {'~' if same_sites and len(ra | rb) > 1 else '=' if same_sites else 'vs'} {tb_}")
            same_cons = A.reduced_constituents() == B.reduced_constituents()
            ev.append("constituents " + (" : ".join(",".join(c) for c in A.reduced_constituents()) + " same"
                                         if same_cons else
                                         f"{_cons_text(A)} vs {_cons_text(B)}"))
            if A.magnetic or B.magnetic:
                ev.append(f"magnetic {_mag(A)} {'=' if A.magnetic == B.magnetic else 'vs'} {_mag(B)}")
            if bool(A.dis_part) != bool(B.dis_part) or A.dis_part:
                ev.append(f"DIS_PART {A.dis_part or '-'} vs {B.dis_part or '-'}")
            structure = same_sites and same_cons and A.magnetic == B.magnetic and bool(A.dis_part) == bool(B.dis_part)
            if proto_conflict:
                ev.insert(0, f"different prototype tokens ({'/'.join(sorted(pa))} vs {'/'.join(sorted(pb))})")
                score *= 0.5
            if score >= 0.5 and structure:
                verdict = "LIKELY same phase"
            elif score >= 0.5 or (score >= 0.3 and structure):
                verdict = "possibly the same phase"
            else:
                continue
            ev.insert(0, f"name similarity {score:.2f}")
            in_ref_a, in_ref_b = ref_label in la, ref_label in lb
            if in_ref_a or (not in_ref_b and (len(la) > len(lb) or (len(la) == len(lb) and a < b))):
                target, other = a, b
            else:
                target, other = b, a
            out.append(Candidate(a, b, sorted(la), sorted(lb), score, ev, verdict, {other: target}))
    out.sort(key=lambda c: (c.verdict != "LIKELY same phase", -c.name_score))
    return out


def _cons_text(pv: PhaseView) -> str:
    return " : ".join(",".join(c) for c in pv.reduced_constituents())


def _mag(pv: PhaseView) -> str:
    return f"{pv.magnetic[0]:g}/{pv.magnetic[1]:g}" if pv.magnetic else "none"


# =============================================================================================== user expressions
_PARAM_RE = re.compile(r"\b([A-Z][A-Z0-9_]*)\s*\(\s*([^()]*?[,:][^()]*?)\)")


class NotAvailable(Exception):
    pass


@dataclass
class ExprSpec:
    label: str
    text: str                               # as given (or a description for formation)
    kind: str = "expr"                      # expr | formation
    phase: str = ""
    constituents: str = ""
    per: str = "formula_unit"


@dataclass
class ExprResult:
    spec: ExprSpec
    segs: Dict[str, Optional[list]] = field(default_factory=dict)       # db label -> segments (None: n/a)
    detail: Dict[str, str] = field(default_factory=dict)                # db label -> refs used / reason
    plan: Dict[str, Dict] = field(default_factory=dict)                 # db label -> standalone recipe
    ambiguity: str = ""


def parse_specs(expressions: Optional[Sequence] = None, formation: Optional[Sequence] = None) -> List[ExprSpec]:
    specs = []
    for e in expressions or []:
        if isinstance(e, str):
            specs.append(ExprSpec(e.strip(), e.strip()))
        elif isinstance(e, dict) and e.get("expr"):
            specs.append(ExprSpec((e.get("label") or e["expr"]).strip(), e["expr"].strip()))
        else:
            raise ValueError(f"expression must be a string or {{label, expr}}: {e!r}")
    for f in formation or []:
        if not isinstance(f, dict) or not f.get("phase") or not f.get("constituents"):
            raise ValueError(f"formation entries need phase and constituents: {f!r}")
        per = (f.get("per") or "formula_unit").lower()
        per = "mole_of_atoms" if per.startswith("mol") or "atom" in per else "formula_unit"
        ph, cons = f["phase"].upper(), f["constituents"].upper().replace(" ", "")
        lab = f.get("label") or f"ΔfG {ph}({cons}) per {'mol atoms' if per == 'mole_of_atoms' else 'f.u.'}"
        specs.append(ExprSpec(lab, f"G({ph},{cons}) - Σ site ratio × element reference", "formation", ph, cons, per))
    return specs


class ExprEvaluator:
    def __init__(self, views: List[DbView]):
        self.views = views
        self.all_phase_names = set()
        for v in views:
            self.all_phase_names |= set(v.phases) | set(v.tdb.phases) | set(v.rename)

    # ---------------------------------------------------------------- references
    def _param(self, v: DbView, ref: str) -> Tuple[str, float]:
        try:
            key, sign, _ = parse_param_name(ref)
        except ExpressionError as e:
            raise NotAvailable(str(e))
        cn = v.canon_of(key.phase)
        if cn is None:
            if key.phase not in self.all_phase_names:
                raise _Ambiguous([key.phase] + [s for sl in key.constituents for s in sl])
            raise NotAvailable(f"no phase {key.phase}")
        pv = v.phases[cn]
        cons = key.constituents
        n_red = pv.n_sub - len(pv.va_only)
        stripped = tuple(sl for sl in cons if sl != ("VA",))
        if len(cons) == pv.n_sub:
            red = tuple(sl for i, sl in enumerate(cons) if i not in pv.va_only)
        elif len(cons) == n_red:
            red = cons
        elif len(stripped) == n_red:            # ":VA" written for a database without the vacancy sublattice
            red = stripped
        else:
            raise NotAvailable(f"{ref}: {len(cons)} sublattice(s) given, {cn} has {pv.n_sub}")
        k = make_key(key.ptype, cn, red, key.order, key.diffusing)
        row = v.rows.get(k)
        if row is None:
            raise NotAvailable(f"missing {make_label(key.ptype, cn, cons, key.order, key.diffusing)}")
        if row.segs is None:
            raise NotAvailable(f"{row.label}: {row.error}")
        return k, sign

    def candidates_for(self, species: List[str]) -> str:
        want = [s.upper() for s in species if s.upper() != "VA"]
        exact, loose = {}, {}
        for v in self.views:
            for k, r in v.rows.items():
                if r.ptype != "G" or r.segs is None:
                    continue
                red = [sl for sl in key_constituents(k)]
                if any(len(sl) != 1 for sl in red):
                    continue
                seq = [sl[0] for sl in red if sl[0] != "VA"]
                if seq == want:
                    exact.setdefault(r.label, []).append(v.label)
                elif set(seq) == set(want):
                    loose.setdefault(r.label, []).append(v.label)
        parts = []
        if exact:
            parts.append("end-members " + ":".join(want) + ": " + "; ".join(f"{k} [{', '.join(d)}]"
                                                                             for k, d in sorted(exact.items())))
        if loose:
            parts.append("other end-members of " + "+".join(sorted(set(want))) + ": "
                         + "; ".join(f"{k} [{', '.join(d)}]" for k, d in sorted(loose.items())))
        return " | ".join(parts) or "no end-member with these species in any database"

    # ---------------------------------------------------------------- evaluation
    def evaluate(self, spec: ExprSpec) -> ExprResult:
        res = ExprResult(spec)
        for v in self.views:
            try:
                expr, refs, desc = (self._formation(v, spec) if spec.kind == "formation"
                                    else self._expression(v, spec.text))
                segs = self._combine(v, expr, refs)
                res.segs[v.label] = segs
                res.detail[v.label] = desc
                res.plan[v.label] = {"expr": str(expr), "refs": {n: list(r) for n, r in refs.items()}}
            except _Ambiguous as a:
                res.ambiguity = f"{spec.text}: no phase given/unknown phase - {self.candidates_for(a.species)}"
                res.segs[v.label] = None
                res.detail[v.label] = "n/a: ambiguous (see candidates)"
                res.plan[v.label] = {"na": "ambiguous"}
            except (NotAvailable, pwi.MissingRef) as e:
                why = e.args[0] if e.args else str(e)
                if isinstance(e, pwi.MissingRef) and not str(why).startswith("reference"):
                    why = f"missing function {why}"
                res.segs[v.label] = None
                res.detail[v.label] = f"n/a in {v.label}: {why}"
                res.plan[v.label] = {"na": str(why)}
        return res

    def _expression(self, v: DbView, text: str):
        refs: Dict[str, Tuple] = {}
        names: Dict[str, str] = {}

        def sub(m):
            n = f"ZZREF{len(refs)}"
            k, sign = self._param(v, m.group(0))
            refs[n] = ("param", k, sign)
            names[n] = v.rows[k].label
            return f" {n} "
        s = _PARAM_RE.sub(sub, text.upper())
        expr = tdb_to_se(s)
        used = []
        for sym in sorted(expr.free_symbols, key=lambda x: x.name):
            nm = sym.name
            if nm in ("T", "P") or nm in refs:
                continue
            v.fn(nm)                                   # raises MissingRef
            refs[nm] = ("function", nm)
            used.append(nm)
        desc = ", ".join(list(names.values()) + used) or "(T only)"
        return expr, refs, desc

    def _formation(self, v: DbView, spec: ExprSpec):
        cn = v.canon_of(spec.phase)
        if cn is None:
            if spec.phase not in self.all_phase_names:
                raise _Ambiguous(re.split(r"[:,]", spec.constituents))
            raise NotAvailable(f"no phase {spec.phase}")
        pv = v.phases[cn]
        cons = [c for c in spec.constituents.split(":")]
        if any("," in c for c in cons):
            raise NotAvailable("formation needs one species per sublattice")
        if len(cons) != pv.n_sub and len([c for c in cons if c != "VA"]) == pv.n_sub - len(pv.va_only):
            cons = [c for c in cons if c != "VA"]
        if len(cons) == pv.n_sub - len(pv.va_only):
            full, it = [], iter(cons)
            for i in range(pv.n_sub):
                full.append("VA" if i in pv.va_only else next(it))
        elif len(cons) == pv.n_sub:
            full = cons
        else:
            raise NotAvailable(f"{spec.constituents}: {len(cons)} sublattice(s), {cn} has {pv.n_sub}")
        k, _ = self._param(v, f"G({cn},{':'.join(full)};0)")
        counts: Dict[str, float] = {}
        for r, sp in zip(pv.site_ratios, full):
            if sp == "VA":
                continue
            els = v.species_elements(sp)
            if len(els) != 1:
                raise NotAvailable(f"formation of molecular species {sp} is not supported")
            el = next(iter(els))
            counts[el] = counts.get(el, 0.0) + r
        refs = {"ZZREF0": ("param", k, 1.0)}
        expr = se.Symbol("ZZREF0")
        used = []
        for i, (el, n) in enumerate(sorted(counts.items()), start=1):
            kind, what, desc = v.ser_reference(el)
            nm = f"ZZREF{i}"
            refs[nm] = ("param", what, 1.0) if kind == "param" else ("function", what)
            expr = expr - _num(n) * se.Symbol(nm)
            used.append(f"{el}: {_fmtn(n)}×{desc}")
        atoms = sum(counts.values())
        if spec.per == "mole_of_atoms":
            expr = expr / _num(atoms)
        desc = (f"{v.rows[k].label} − " + " − ".join(used)
                + (f", ÷ {_fmtn(atoms)} atoms" if spec.per == "mole_of_atoms" else f" (per f.u., {_fmtn(atoms)} atoms)"))
        return expr, refs, desc

    def _combine(self, v: DbView, expr, refs: Dict[str, Tuple]):
        seg_of: Dict[str, list] = {}
        for n, r in refs.items():
            if r[0] == "param":
                s = v.rows[r[1]].segs
                seg_of[n] = pwi.scale(s, r[2]) if r[2] != 1.0 else s
            else:
                seg_of[n] = v.fn(r[1])
        los = [s[0][0] for s in seg_of.values() if s]
        his = [s[-1][1] for s in seg_of.values() if s]
        lo = max(los) if los else 298.15
        hi = min(his) if his else 6000.0
        if not lo < hi:
            lo, hi = min(los), max(his)
        return pwi.inline([(lo, hi, expr)], lambda n: seg_of[n] if n in seg_of else v.fn(n), v.inline_notes)


class _Ambiguous(Exception):
    def __init__(self, species):
        super().__init__("ambiguous")
        self.species = [s for s in species if s]


def _num(x: float):
    return se.Integer(int(x)) if float(x) == int(x) else se.RealDouble(float(x))


def _fmtn(x: float) -> str:
    return f"{x:g}"


# =============================================================================================== the comparison
@dataclass
class Cell:
    present: bool
    status: str = ""          # ref | same | differs | absent | no ref | error
    values: Optional[np.ndarray] = None
    extrap: Optional[np.ndarray] = None
    max_diff: float = math.nan
    t_at_max: float = math.nan
    differs_at: Optional[np.ndarray] = None


@dataclass
class TableRow:
    key: str
    label: str
    phase: str
    ptype: str
    rows: Dict[str, Optional[ParamRow]]
    cells: Dict[str, Cell] = field(default_factory=dict)
    is_expression: bool = False

    @property
    def all_same(self) -> bool:
        return all(c.status in ("ref", "same") for c in self.cells.values())


@dataclass
class SystemComparison:
    views: List[DbView]
    elements: List[str]
    temps: List[float]
    ref: str
    phases: List[str]                               # canonical phase names shown (ordered)
    table: Dict[str, List[TableRow]]                # phase -> rows
    expressions: List[ExprResult]
    expr_rows: List[TableRow]
    candidates: List[Candidate]
    phase_map: Dict[str, Dict[str, str]]            # db label -> {native: canonical} actually applied
    rtol: float
    atol: float
    grid: Tuple[float, float]
    warnings: List[str]
    ref_states: Dict[str, Dict[str, str]] = field(default_factory=dict)   # db label -> element -> text
    written: List[ExprRow] = field(default_factory=list)                   # parameters as written, all databases

    def written_differences(self) -> List[ExprRow]:
        return differences(self.written)

    def written_equal(self) -> List[ExprRow]:
        """Rows written differently but numerically equal within the tolerance (not differences)."""
        return equal_rows(self.written)

    @property
    def labels(self) -> List[str]:
        return [v.label for v in self.views]

    def view(self, label: str) -> DbView:
        return next(v for v in self.views if v.label == label)

    def all_rows(self) -> List[TableRow]:
        return [r for ph in self.phases for r in self.table.get(ph, [])]


def _labels_for(paths: List[str], workdir: Path) -> List[str]:
    stems = [Path(p).stem for p in paths]
    out = []
    for p, s in zip(paths, stems):
        out.append(s if stems.count(s) == 1 else _rel(Path(p), workdir))
    return out


def _db_matches(key: str, path: str, label: str, rel: str) -> bool:
    k = key.strip().lower()
    return k in {Path(path).name.lower(), Path(path).stem.lower(), label.lower(), rel.lower(), str(path).lower()}


def split_phase_map(phase_map: Optional[Dict], paths: List[str], labels: List[str], rels: List[str]):
    """-> ({db label: {native: canonical}}, warnings).  Global entries apply to every database."""
    glob = {k.upper(): v.upper() for k, v in (phase_map or {}).items() if isinstance(v, str)}
    per = {lab: dict(glob) for lab in labels}
    warns = []
    for k, v in (phase_map or {}).items():
        if isinstance(v, dict):
            hit = [lab for p, lab, r in zip(paths, labels, rels) if _db_matches(k, p, lab, r)]
            if not hit:
                warns.append(f"phase_map key {k!r} matches no database ({', '.join(labels)})")
            for lab in hit:
                per[lab].update({a.upper(): b.upper() for a, b in v.items()})
        elif not isinstance(v, str):
            warns.append(f"phase_map entry {k!r} ignored (value must be a phase name or a per-database dict)")
    return per, warns


def compare_system(paths: Sequence[str], elements: Sequence[str], temperatures: Optional[Sequence[float]] = None,
                   phases: Optional[Sequence[str]] = None, phase_map: Optional[Dict] = None,
                   expressions: Optional[Sequence] = None, formation: Optional[Sequence] = None,
                   reference: Optional[str] = None, include_identical: bool = True, rtol: float = 1e-4,
                   atol: float = 1e-6, t_grid: Tuple[float, float] = (298.15, 3000.0),
                   workdir: Optional[Path] = None, labels: Optional[List[str]] = None) -> SystemComparison:
    workdir = Path(workdir or Path.cwd())
    paths = [str(p) for p in paths]
    if len(paths) < 2:
        raise ValueError("compare_system needs at least two databases")
    if len(set(os.path.realpath(p) for p in paths)) != len(paths):
        raise ValueError("the same database is listed twice")
    els = sorted({e.upper() for e in elements if e.upper() != "VA"})
    if not els:
        raise ValueError("give the subsystem elements, e.g. ['FE', 'CR']")
    temps = [float(t) for t in (temperatures or DEFAULT_TEMPS)]
    labels = labels or _labels_for(paths, workdir)
    rels = [_rel(Path(p), workdir) for p in paths]
    pmap, warns = split_phase_map(phase_map, paths, labels, rels)
    views = [DbView(p, lab, r, els, pmap[lab]) for p, lab, r in zip(paths, labels, rels)]
    for v in views:
        if v.missing_elements:
            warns.append(f"{v.label} lacks element(s) {', '.join(v.missing_elements)}")
        used = {a: b for a, b in pmap[v.label].items() if a in v.tdb.phases}
        pmap[v.label] = used
    for k, val in (phase_map or {}).items():
        if isinstance(val, str) and not any(k.upper() in v.tdb.phases for v in views):
            warns.append(f"phase_map: no database has a phase {k.upper()}")
    ref = labels[0]
    if reference:
        hit = [lab for p, lab, r in zip(paths, labels, rels) if _db_matches(reference, p, lab, r)]
        if not hit:
            raise ValueError(f"reference {reference!r} is not one of {labels}")
        ref = hit[0]

    # ---- phases
    order = phase_order(cn for v in views for cn in v.phases)
    if phases:
        want = [p.upper() for p in phases]
        unknown = [p for p in want if p not in order]
        if unknown:
            warns.append(f"phases not in any database of the subsystem: {', '.join(unknown)}")
        order = [p for p in order if p in want]

    grid = np.union1d(np.linspace(t_grid[0], t_grid[1], 121), temps)
    table: Dict[str, List[TableRow]] = {}
    for ph in order:
        keys: Dict[str, str] = {}
        for v in sorted(views, key=lambda x: x.label != ref):
            for k, r in v.rows.items():
                if r.phase == ph and k not in keys:
                    keys[k] = r.label
        trs = []
        for k, lab in keys.items():
            tr = TableRow(k, lab, ph, "", {v.label: v.rows.get(k) for v in views})
            anyrow = next(r for r in tr.rows.values() if r is not None)
            tr.ptype = anyrow.ptype
            _fill_cells(tr, ref, temps, grid, rtol, atol)
            trs.append(tr)
        trs.sort(key=lambda t: next(r for r in t.rows.values() if r is not None).sort_key())
        if not include_identical:
            trs = [t for t in trs if not t.all_same]
        table[ph] = trs

    # ---- written expressions (the main tables)
    written = expression_table(views, order, rtol=rtol, t_grid=t_grid)
    if not include_identical:
        written = [r for r in written if r.status == "differs"]

    # ---- expressions
    specs = parse_specs(expressions, formation)
    ev = ExprEvaluator(views)
    exprs = [ev.evaluate(s) for s in specs]
    expr_rows = []
    for er in exprs:
        tr = TableRow(er.spec.label, er.spec.label, "(expressions)", "expr", {}, is_expression=True)
        _fill_cells(tr, ref, temps, grid, rtol, atol, segs_by_db=er.segs)
        expr_rows.append(tr)

    cands = phase_candidates(views, ref)
    for v in views:
        warns += [f"{v.label}: {n}" for n in v.notes]
        if v.inline_notes:
            warns.append(f"{v.label}: " + "; ".join(sorted(v.inline_notes)[:4]))
    return SystemComparison(views, els, temps, ref, order, table, exprs, expr_rows, cands, pmap, rtol, atol,
                            (float(grid[0]), float(grid[-1])), warns, {v.label: v.reference_states() for v in views},
                            written)


def _fill_cells(tr: TableRow, ref: str, temps, grid, rtol, atol, segs_by_db: Optional[Dict] = None):
    def segs_of(lab):
        if segs_by_db is not None:
            return segs_by_db.get(lab)
        r = tr.rows.get(lab)
        return None if r is None else r.segs

    labels = list(segs_by_db) if segs_by_db is not None else list(tr.rows)
    rs = segs_of(ref)
    for lab in labels:
        s = segs_of(lab)
        present = (tr.rows.get(lab) is not None) if segs_by_db is None else s is not None
        if not present:
            tr.cells[lab] = Cell(False, "absent")
            continue
        if s is None:
            tr.cells[lab] = Cell(True, "error")
            continue
        vals, ext = pwi.evaluate(s, temps)
        c = Cell(True, "ref" if lab == ref else "", vals, ext)
        if lab != ref:
            if rs is None:
                c.status = "no ref"
            else:
                lo = max(s[0][0], rs[0][0], grid[0])
                hi = min(s[-1][1], rs[-1][1], grid[-1])
                g = grid[(grid >= lo) & (grid <= hi)] if lo < hi else np.asarray(temps)
                va, _ = pwi.evaluate(s, g)
                vr, _ = pwi.evaluate(rs, g)
                same, d = pwi.close(va, vr, rtol, atol)
                c.status = "same" if same else "differs"
                c.max_diff = d
                if np.isfinite(va).any() and np.isfinite(vr).any():
                    c.t_at_max = float(g[int(np.nanargmax(np.abs(va - vr)))])
                vref, _ = pwi.evaluate(rs, temps)
                tol = np.maximum(atol, rtol * np.maximum(np.abs(vals), np.abs(vref)))
                c.differs_at = np.abs(vals - vref) > tol
        tr.cells[lab] = c


# =============================================================================================== text summary
def summary_text(cmp: SystemComparison, files: Dict[str, str], rel=lambda p: p, max_examples: int = 12) -> str:
    L = []
    sys_ = "-".join(cmp.elements)
    idx = {lab: f"D{i}" for i, lab in enumerate(cmp.labels, 1)}
    L.append(f"compare_system {sys_} (+VA): {len(cmp.views)} databases: "
             + "; ".join(f"{idx[v.label]} {v.label} = {v.rel}" for v in cmp.views))
    diffs = cmp.written_differences()
    n_all = len(cmp.written)
    n_val = sum(1 for r in diffs if r.kind == "value")
    n_abs = sum(1 for r in diffs if r.kind == "absent")
    n_eq = sum(1 for r in cmp.written if r.status == "equal")
    n_exp = sum(1 for r in diffs if r.expanded)
    L.append(f"DIFFERENCES: {n_val} parameters differ numerically; {n_eq} written differently but equal (not "
             f"listed as differences); {n_abs} absent in some database; {n_all - len(diffs) - n_eq} identical as "
             f"written ({n_all} in all; equal = within rtol {cmp.rtol:g} over {cmp.grid[0]:g}-{cmp.grid[1]:g} K). "
             "Cells as written with VV#### replaced by values"
             + (f"; in {n_exp} differing rows the functions are expanded (inlined polynomial in T) because the "
                "databases do not share them" if n_exp else "") + "."
             + (f" The {n_eq} equal-but-written-differently rows are in their own 'Numerically equal' slides/sheet "
                "(not listed here)." if n_eq else ""))
    pick = [r for r in diffs if r.kind == "value"] + [r for r in diffs if r.kind == "absent"]
    for r in pick[:max_examples]:
        groups: Dict[str, List[str]] = {}
        for lab in cmp.labels:
            txt = r.shown.get(lab)
            txt = "—" if txt is None else txt.replace("\n", " / ")
            groups.setdefault(txt, []).append(idx[lab])
        cells = " | ".join(f"{','.join(d)}: {t[:110] + ('…' if len(t) > 110 else '')}" for t, d in groups.items())
        L.append(f"  {r.label}: {cells}" + (f"  [{r.note}]" if r.note else ""))
    if len(pick) > max_examples:
        L.append(f"  ... {len(pick) - max_examples} more in the Differences slides / sheet")
    refs = []
    tent_rows = []
    for lab in cmp.labels:
        v = cmp.view(lab)
        n_ref = sum(1 for r in v.rows.values() if r.reference)
        n_cit = sum(1 for k in v.rows if v.citation(k))
        refs.append(f"{idx[lab]} {n_ref}/{len(v.rows)}"
                    + (f" ({n_cit} resolved by its reference list)" if n_cit else ""))
        tl = [r.label for k, r in v.rows.items() if r.reference and is_tentative(r.reference, citation=v.citation(k))]
        if tl:
            tent_rows.append(f"{idx[lab]}: {', '.join(tl[:4])}" + (f" +{len(tl) - 4}" if len(tl) > 4 else ""))
    L.append("Parameter references (key after the last T range; shown under each cell, full citations in Excel): "
             + ", ".join(refs) + ".")
    if tent_rows:
        L.append("TENTATIVE parameters (not assessed; TENT_MIX / TENT_VA, marked 'tentative' under the cell): "
                 + "; ".join(tent_rows) + ".")
    applied = {lab: m for lab, m in cmp.phase_map.items() if m}
    if applied:
        L.append("phase_map applied: " + "; ".join(f"{lab}: " + ", ".join(f"{a}->{b}" for a, b in m.items())
                                                   for lab, m in applied.items()))
    non_ser = {lab: {el: t for el, t in d.items() if "NOT SER" in t or "no reference" in t}
               for lab, d in cmp.ref_states.items()}
    non_ser = {lab: d for lab, d in non_ser.items() if d}
    if non_ser:
        L.append("Element reference states differ (Gibbs energies of different references are not directly "
                 "comparable): " + "; ".join(f"{lab}: " + ", ".join(f"{el} {t}" for el, t in d.items())
                                             for lab, d in non_ser.items()))
    L.append("Phase presence (" + " ".join(idx[lab] for lab in cmp.labels) + "):")
    for ph in cmp.phases:
        marks = ["x" if v.phases.get(ph) else "." for v in cmp.views]
        hs = sorted({v.phases[ph].ratios_text() for v in cmp.views if v.phases.get(ph)})
        note = f"sites {' vs '.join(hs)}" if hs else ""
        extra = sorted({h for v in cmp.views if v.phases.get(ph) for h in v.phases[ph].hints()})
        if extra:
            note += "; " + ", ".join(extra)
        L.append(f"  {ph:<16s} {'  '.join(marks)}   {note}")
    od = sorted({f"{pv.canon}->{pv.dis_part}" for v in cmp.views for pv in v.phases.values() if pv.dis_part})
    if od:
        L.append("Order/disorder (DIS_PART): " + ", ".join(od) + " - the ordered phase's Gibbs energy also contains "
                 "the parameters listed under its disordered phase; compare those tables together.")
    if cmp.expressions:
        L.append("Derived expressions (values at " + ", ".join(f"{t:g}" for t in cmp.temps) + " K):")
        for er, tr in zip(cmp.expressions, cmp.expr_rows):
            L.append(f"  {er.spec.label}")
            for lab in cmp.labels:
                c = tr.cells.get(lab)
                segs = er.segs.get(lab)
                if c is not None and c.values is not None and segs:
                    vals = ", ".join(pwi.fmt_value(x) + ("*" if e else "") for x, e in zip(c.values, c.extrap))
                    L.append(f"    {lab}: {pwi.fmt_expr(segs[0][2], 7)[:90]}"
                             + (f" [+{len(segs) - 1} T ranges]" if len(segs) > 1 else "")
                             + f" = {vals}  <- {er.detail.get(lab, '')}")
                else:
                    L.append(f"    {lab}: {er.detail.get(lab, 'n/a')}")
            if er.ambiguity:
                L.append(f"    AMBIGUOUS: {er.ambiguity}")
    if cmp.candidates:
        L.append("Phase-correspondence CANDIDATES (NOT merged - judge them, ask the user if unsure, then re-call "
                 "with phase_map):")
        for c in cmp.candidates[:12]:
            L.append("  " + c.text())
    pyc = [f"{v.label}: {', '.join(v.pyc_sign_notes[:4])}" for v in cmp.views if v.pyc_sign_notes]
    if pyc:
        L.append("Note: non-alphabetical constituents with odd RK order (sign-corrected here, Thermo-Calc "
                 "convention; pycalphad sorts them WITHOUT changing the sign): " + "; ".join(pyc))
    if cmp.warnings:
        L.append("Warnings: " + "; ".join(cmp.warnings[:10]))
    L.append("Files: " + ", ".join(f"{k}: {rel(p)}" for k, p in files.items()))
    return "\n".join(L)


def _short_expr(row: Optional[ParamRow], n: int = 70) -> str:
    if row is None:
        return "-"
    if row.segs is None:
        return row.error
    s = pwi.fmt_expr(row.segs[0][2], 7)
    s = s if len(s) <= n else s[: n - 1] + "…"
    return s + (f" [+{len(row.segs) - 1} ranges]" if len(row.segs) > 1 else "")


# =============================================================================================== CSV / Excel
def _tcols(temps) -> List[str]:
    return [f"value_{t:g}K" for t in temps]


def write_csv(cmp: SystemComparison, path: str) -> str:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["kind", "phase", "key", "item", "database", "path", "status_vs_ref", "max_abs_diff_vs_ref",
                    "T_range"] + _tcols(cmp.temps) + ["extrapolated_at", "expanded", "raw_tdb_or_detail",
                                                      "written", "written_status", "written_note",
                                                      "as_written", "reference", "reference_short",
                                                      "reference_citation"])
        wr = {r.key: r for r in cmp.written}
        for tr in cmp.all_rows() + cmp.expr_rows:
            for v in cmp.views:
                lab = v.label
                c = tr.cells.get(lab)
                r = tr.rows.get(lab) if not tr.is_expression else None
                if tr.is_expression:
                    er = next(e for e in cmp.expressions if e.spec.label == tr.key)
                    segs = er.segs.get(lab)
                    raw = er.detail.get(lab, "")
                else:
                    segs = r.segs if r else None
                    raw = (r.raw_name + " " + r.raw_text) if r else ""
                    if r and r.error:
                        raw += f"  [{r.error}]"
                vals = [("" if c is None or c.values is None else f"{x:.10g}") for x in
                        (c.values if c is not None and c.values is not None else [None] * len(cmp.temps))]
                ext = "" if c is None or c.extrap is None else " ".join(f"{t:g}" for t, e in zip(cmp.temps, c.extrap) if e)
                w.writerow(["expression" if tr.is_expression else "parameter", tr.phase, tr.key, tr.label, lab,
                            v.rel, c.status if c else "", "" if c is None or not math.isfinite(c.max_diff)
                            else f"{c.max_diff:.6g}",
                            "" if not segs else f"{segs[0][0]:g}-{segs[-1][1]:g}"] + vals
                           + [ext, pwi.pw_text(segs) if segs else "", raw]
                           + ([((wr[tr.key].shown.get(lab) or "").replace("\n", " | ") if tr.key in wr else ""),
                               wr[tr.key].status if tr.key in wr else "", wr[tr.key].note if tr.key in wr else "",
                               v.written(tr.key).text.replace("\n", " | ") if r else ""]
                              if not tr.is_expression else ["", "", "", ""])
                           + _ref_cols(v, tr.key, r))
    return path


def _ref_cols(v: Subsystem, key: str, r: Optional[ParamRow]) -> List[str]:
    """[reference as written, short citation, full citation] of one database's parameter (blank when none)."""
    if r is None or not r.reference:
        return ["", "", ""]
    cit = v.citation(key)
    if is_tentative(r.reference, citation=cit):
        return [r.reference, "tentative", cit]
    return [r.reference, short_citation(cit) if cit else "", cit]


FILLS = {"differs": "FFEBEE", "same": "E8F5E9", "absent": "F3F4F6", "no ref": "FFF8E1", "error": "FDE68A",
         "ref": "E3F2FD"}


def _sheet_name(name: str, taken: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "sheet"
    s, k = base, 2
    while s.lower() in taken:
        suf = f"~{k}"
        s, k = base[: 31 - len(suf)] + suf, k + 1
    taken.add(s.lower())
    return s


def write_xlsx(cmp: SystemComparison, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="2C5364")
    white = Font(bold=True, color="FFFFFF")
    taken: set = set()

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font, c.fill = white, head_fill
        ws.freeze_panes = f"A{ws.max_row + 1}"

    def fill(cell, status):
        col = FILLS.get(status)
        if col:
            cell.fill = PatternFill("solid", fgColor=col)

    ws = wb.active
    ws.title = _sheet_name("Summary", taken)
    ws.append([f"{'-'.join(cmp.elements)} comparison of {len(cmp.views)} databases; reference {cmp.ref}; "
               f"values in J/mol (K for TC); functions inlined; P = {pwi.P_REF:g} Pa; equality within rtol "
               f"{cmp.rtol:g} over {cmp.grid[0]:g}-{cmp.grid[1]:g} K"])
    ws.append([])
    header(ws, ["#", "label", "path", "phases in subsystem", "parameters", "functions used", "phase_map applied",
                "element references"])
    for i, v in enumerate(cmp.views, 1):
        ws.append([f"D{i}", v.label, v.rel, len(v.phases), len(v.rows), len(v.used_functions()),
                   ", ".join(f"{a}->{b}" for a, b in cmp.phase_map.get(v.label, {}).items()),
                   "; ".join(f"{el}: {t}" for el, t in cmp.ref_states.get(v.label, {}).items())])
    ws.append([])
    ws.append(["Phase"] + cmp.labels + ["notes"])
    for c in ws[ws.max_row]:
        c.font = bold
    for ph in cmp.phases:
        cells = []
        for v in cmp.views:
            pv = v.phases.get(ph)
            cells.append(_presence_text(pv) if pv else "—")
        ws.append([ph] + cells + ["; ".join(sorted({h for v in cmp.views if v.phases.get(ph)
                                                     for h in v.phases[ph].hints()}))])
    if cmp.candidates:
        ws.append([])
        ws.append(["Phase-correspondence candidates (not merged)"])
        ws[ws.max_row][0].font = bold
        for c in cmp.candidates:
            ws.append([c.a, ", ".join(c.dbs_a), c.b, ", ".join(c.dbs_b), c.verdict, "; ".join(c.evidence),
                       str(c.suggestion)])
    if cmp.warnings:
        ws.append([])
        ws.append(["Warnings"])
        for w in cmp.warnings:
            ws.append([w])
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 40

    tcols = [f"{t:g} K" for t in cmp.temps]
    wrap = Alignment(wrap_text=True, vertical="top")

    def written_sheet(name, rows_in):
        ws = wb.create_sheet(_sheet_name(name, taken))
        header(ws, ["phase", "parameter"] + cmp.labels + ["status", "note"]
               + [f"{lab} (as written)" for lab in cmp.labels] + ["max |Δ|"]
               + [f"{lab} reference" for lab in cmp.labels] + [f"{lab} citation" for lab in cmp.labels])
        for er in rows_in:
            ws.append([er.phase, er.label] + [(er.shown[lab] if er.shown.get(lab) is not None else "—")
                                              for lab in cmp.labels] + [er.status, er.note]
                      + [(er.cells[lab].text if er.cells.get(lab) else "—") for lab in cmp.labels]
                      + [er.max_diff if math.isfinite(er.max_diff) else None]
                      + [er.refs.get(lab, "") for lab in cmp.labels]
                      + [er.citations.get(lab, "") for lab in cmp.labels])
            for j, lab in enumerate(cmp.labels):
                c = ws.cell(row=ws.max_row, column=3 + j)
                c.alignment = wrap
                if er.shown.get(lab) is None:
                    fill(c, "absent")
                elif er.cell_differs.get(lab):
                    fill(c, "differs")
                elif er.status == "equal":
                    fill(c, "same")
        _widths(ws, [14, 30] + [48] * len(cmp.labels) + [10, 50] + [40] * len(cmp.labels) + [12]
                + [14] * len(cmp.labels) + [60] * len(cmp.labels))

    written_sheet("Differences", cmp.written_differences())
    written_sheet("Numerically equal", cmp.written_equal())
    written_sheet("All parameters", cmp.written)
    for v in cmp.views:
        ws = wb.create_sheet(_sheet_name(f"DB {v.label}", taken))
        header(ws, ["phase", "parameter", "expression as written", "T range", "written as", "reference",
                    "citation"])
        for k in database_keys(v, cmp.phases):
            r = v.rows[k]
            ws.append([r.phase, r.label, v.written(k).text, f"{r.raw_pw.tmin:g}-{r.raw_pw.tmax:g}", r.raw_name,
                       r.reference, v.citation(k)])
            ws.cell(row=ws.max_row, column=3).alignment = wrap
        fns = database_functions(v, cmp.phases)
        if fns:
            ws.append([])
            ws.append(["FUNCTION", "name", "expression as written", "T range"])
            for c in ws[ws.max_row]:
                c.font = bold
            for n, w in fns:
                f = v.tdb.functions[n].expr
                ws.append(["", n, w.text, f"{f.tmin:g}-{f.tmax:g}"])
                ws.cell(row=ws.max_row, column=3).alignment = wrap
        _widths(ws, [14, 30, 90, 14, 30, 14, 80])

    ws = wb.create_sheet(_sheet_name("Values", taken))
    header(ws, ["phase", "parameter", "database", "status vs ref", "max |Δ| vs ref", "T range"] + tcols
           + ["expanded (functions inlined)", "as written in the TDB"])
    for ph in cmp.phases:
        for tr in cmp.table.get(ph, []):
            for v in cmp.views:
                c, r = tr.cells.get(v.label), tr.rows.get(v.label)
                if r is None:
                    continue
                vals = list(c.values) if c is not None and c.values is not None else [None] * len(cmp.temps)
                ws.append([ph, tr.label, v.label, c.status if c else "", None if c is None or not math.isfinite(
                    c.max_diff) else c.max_diff, f"{r.segs[0][0]:g}-{r.segs[-1][1]:g}" if r.segs else ""]
                          + [None if x is None or not math.isfinite(x) else float(x) for x in vals]
                          + [pwi.pw_text(r.segs) if r.segs else r.error, r.raw_name + " " + r.raw_text])
                fill(ws.cell(row=ws.max_row, column=4), c.status if c else "")
    _widths(ws, [14, 30, 22, 12, 14, 14] + [14] * len(tcols) + [90, 60])

    ws = wb.create_sheet(_sheet_name("Expressions", taken))
    header(ws, ["expression", "database", "status vs ref", "max |Δ| vs ref"] + tcols
           + ["expanded", "references used / reason"])
    for er, tr in zip(cmp.expressions, cmp.expr_rows):
        for v in cmp.views:
            c = tr.cells.get(v.label)
            segs = er.segs.get(v.label)
            vals = list(c.values) if c is not None and c.values is not None else [None] * len(cmp.temps)
            ws.append([er.spec.label, v.label, c.status if c else "", None if c is None or not math.isfinite(
                c.max_diff) else c.max_diff] + [None if x is None or not math.isfinite(x) else float(x) for x in vals]
                      + [pwi.pw_text(segs) if segs else "", er.detail.get(v.label, "")])
            fill(ws.cell(row=ws.max_row, column=3), c.status if c else "")
        if er.ambiguity:
            ws.append([er.spec.label, "AMBIGUOUS", er.ambiguity])
    _widths(ws, [40, 22, 12, 14] + [14] * len(tcols) + [90, 70])

    # term by term
    ws = wb.create_sheet(_sheet_name("Terms", taken))
    recs, keys = [], set()
    for tr in cmp.all_rows():
        for v in cmp.views:
            r = tr.rows.get(v.label)
            if r is None or not r.segs:
                continue
            for a, b, e in r.segs:
                d, other = pwi.terms(e)
                keys |= set(d)
                recs.append((tr.phase, tr.label, v.label, a, b, d, other))
    ks = sorted(keys, key=pwi.term_order)
    header(ws, ["phase", "parameter", "database", "T low", "T high"] + [pwi.term_name(k) for k in ks] + ["other"])
    for ph, lab, db, a, b, d, other in recs:
        ws.append([ph, lab, db, a, b] + [d.get(k) for k in ks] + ["" if other == 0 else str(other)])
    _widths(ws, [14, 30, 22, 9, 9] + [14] * len(ks) + [60])

    ws = wb.create_sheet(_sheet_name("Functions inlined", taken))
    header(ws, ["database", "function", "expanded (inlined)", "as written"])
    for v in cmp.views:
        for fn in v.used_functions():
            try:
                txt = pwi.pw_text(v.fn(fn))
            except Exception as e:  # noqa: BLE001
                txt = f"{type(e).__name__}: {e}"
            ws.append([v.label, fn, txt, v.tdb.functions[fn].expr.to_tdb()])
    _widths(ws, [22, 14, 100, 80])
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and len(c.value) > 60:
                    c.alignment = Alignment(wrap_text=False)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _presence_text(pv: PhaseView) -> str:
    s = pv.ratios_text()
    if pv.magnetic:
        s += " mag"
    if pv.dis_part:
        s += f" ord({pv.dis_part})"
    if pv.native != pv.canon:
        s += f" [{pv.native}]"
    return s


# =============================================================================================== PowerPoint
def write_pptx(cmp: SystemComparison, path: str, diagram_png: Optional[str] = None,
               diagram_note: str = "") -> Tuple[str, Dict[str, Any]]:
    from . import pptx_utils as pu
    from .pptx_utils import GridCell, add_flow_table_slides, add_grid_slides

    prs = pu.new_presentation()
    sys_ = "–".join(e.capitalize() for e in cmp.elements)
    pu.add_title_slide(prs, f"{sys_}: comparison of {len(cmp.views)} databases",
                       "\n".join(f"D{i}  {v.label}" for i, v in enumerate(cmp.views, 1))[:600],
                       note=(f"Parameters inside the {'-'.join(cmp.elements)} subsystem, shown as written in each "
                             "TDB (ESPEI VV#### handles replaced by their values; other functions keep their names). "
                             "'Differences' lists the parameters whose values differ or that are missing "
                             "somewhere; rows written differently but numerically equal follow in their own table."))
    info: Dict[str, Any] = {"slides": []}
    # ---- databases
    rows = []
    for i, v in enumerate(cmp.views, 1):
        pm = ", ".join(f"{a}→{b}" for a, b in cmp.phase_map.get(v.label, {}).items()) or ""
        refs = "\n".join(f"{el}: {t}" for el, t in cmp.ref_states.get(v.label, {}).items())
        rows.append([GridCell(f"D{i}"), GridCell(v.label, bold=True),
                     GridCell(v.rel, mono=True, size=9), GridCell(f"{len(v.phases)} / {len(v.rows)}", align="c"),
                     GridCell(refs, size=8, fill="no ref" if "NOT SER" in refs else None), GridCell(pm, size=9)])
    add_flow_table_slides(prs, "Databases", ["#", "database", "path", "phases / params",
                                             "element reference (lattice stability at SER?)", "phase_map applied"],
                          rows, [0.5, 2.3, 3.0, 1.1, 4.1, 1.6], font_size=10,
                          subtitle=f"Subsystem {'-'.join(cmp.elements)} (+VA)")
    # ---- phase matrix
    rows = []
    ref_v = cmp.views[0]
    for ph in cmp.phases:
        r = [GridCell(ph, bold=True, mono=True)]
        ref_pv = ref_v.phases.get(ph)
        for v in cmp.views:
            pv = v.phases.get(ph)
            if pv is None:
                r.append(GridCell("—", fill="absent", muted=True))
            else:
                diff_model = ref_pv is not None and pv.reduced_ratios() != ref_pv.reduced_ratios()
                r.append(GridCell(_presence_text(pv), fill="no ref" if diff_model else "same"))
        rows.append(r)
    for gi, grp in enumerate(_db_groups(cmp, 8)):
        n = len(grp)
        w = min(2.2, (12.4 - 2.4) / max(n, 1))
        sub_rows = [[r[0]] + [r[1 + cmp.labels.index(lab)] for lab in grp] for r in rows]
        add_grid_slides(prs, "Phases present" + _grp_suffix(cmp, grp), [("phase", 1)] + [(lab, 1) for lab in grp],
                        None, sub_rows, [2.4] + [w] * n, font_size=10 if w >= 1.5 else 8,
                        subtitle="site ratios per database (VA = vacancy-only sublattice); mag = magnetic; "
                                 f"amber = sublattice model differs from {ref_v.label}; [name] = renamed by phase_map",
                        rows_per_slide=14)
    if cmp.candidates:
        rows = [[GridCell(f"{c.a}\n[{', '.join(c.dbs_a)}]", mono=True), GridCell(f"{c.b}\n[{', '.join(c.dbs_b)}]",
                                                                                  mono=True),
                 GridCell(c.verdict, fill="no ref" if c.verdict.startswith("LIKELY") else None),
                 GridCell("; ".join(c.evidence)), GridCell(", ".join(f"{a}→{b}" for a, b in c.suggestion.items()),
                                                           mono=True)] for c in cmp.candidates]
        add_flow_table_slides(prs, "Possible phase correspondences (not merged)",
                              ["phase", "phase", "verdict", "evidence", "phase_map"], rows,
                              [2.3, 2.3, 1.6, 4.6, 1.8], font_size=9,
                              subtitle="Phases are matched by exact name only; re-run with phase_map to merge")
    # ---- the parameter tables
    diffs = cmp.written_differences()
    eq = cmp.written_equal()
    n_all = len(cmp.written)
    foot = ("As written (VV#### replaced by values) unless functions were expanded; — = absent; red = differs from "
            "the first database that has it; green = written differently, numerically equal; grey = reference "
            "(key · author year)")
    _written_slides(prs, cmp, "Differences", diffs, info,
                    f"{len(diffs)} of {n_all} parameters differ in value or are absent somewhere"
                    + (f"; {len(eq)} written differently but numerically equal are listed next" if eq else
                       "; none is written differently but numerically equal"), foot)
    if eq:
        _written_slides(prs, cmp, EQUAL_TITLE, eq, info,
                        f"{len(eq)} parameters: values agree within the tolerance (rtol {cmp.rtol:g}; 1e-3 J/mol, "
                        f"1e-6 for TC/BMAGN) over {cmp.grid[0]:g}–{cmp.grid[1]:g} K; not differences", foot,
                        equal=True)
    _written_slides(prs, cmp, "All parameters", cmp.written, info,
                    f"{n_all} parameters, sorted by phase, then G, L0, L1, …, ternary, TC, BMAGN", foot)
    shown = set(cmp.phases)
    for v in cmp.views:
        rows = []
        band = None
        by_key = {r.key: r for r in cmp.written}
        for key in database_keys(v, cmp.phases):
            pr, w = v.rows[key], v.written(key)
            if pr.phase != band:
                rows.append(("band", pr.phase))
                band = pr.phase
            er = by_key.get(key)
            note = "" if er is None else {"equal": "numerically equal (written differently)"}.get(er.status, "")
            if er is not None and er.kind == "value":
                note = "value differs between databases"
            cit = v.citation(key)
            tsub = "tentative" if is_tentative(pr.reference, citation=cit) else ""
            rows.append([GridCell(pr.label, mono=True, size=8, sub=note), GridCell(w.text, mono=True),
                         GridCell(pr.reference, size=8, sub=tsub or (short_citation(cit, 24) if cit else ""))])
        has_refs = any(r.reference for r in v.rows.values())
        if not has_refs:                                   # e.g. ESPEI output: no reference column
            rows = [r if isinstance(r, tuple) else r[:2] for r in rows]
        k = add_flow_table_slides(prs, f"{v.label}: parameters",
                                  ["parameter", "expression as written"] + (["reference"] if has_refs else []),
                                  rows, [3.0, 7.7, 1.8] if has_refs else [3.0, 9.5], font_size=8,
                                  subtitle=f"{v.rel} · {len(v.rows)} parameters in {'-'.join(cmp.elements)}"
                                           + (" · reference: key as written, first author + year below"
                                              if any(v.citation(k) for k in v.rows) else
                                              "" if has_refs else " · no parameter references"))
        info["slides"].append({"title": f"{v.label}: parameters", "slides": k})
        fns = database_functions(v, cmp.phases if shown != set(v.phases) else None)
        if fns:
            rows = [[GridCell(n, mono=True, bold=True, size=8), GridCell(w.text, mono=True)] for n, w in fns]
            k = add_flow_table_slides(prs, f"{v.label}: functions used", ["function", "expression as written"],
                                      rows, [1.8, 10.7], font_size=8,
                                      subtitle="FUNCTIONs referenced by these parameters (VV#### handles omitted)")
            info["slides"].append({"title": f"{v.label}: functions used", "slides": k})
    # ---- derived expressions
    if cmp.expressions:
        for grp in _db_groups(cmp, 4):
            nw = (12.5 - 2.6) / len(grp)
            rows = []
            for er in cmp.expressions:
                r = [GridCell(er.spec.label, bold=True, size=8)]
                for lab in grp:
                    segs = er.segs.get(lab)
                    if segs:
                        txt = "\n".join((f"{a:g}–{b:g}: " if len(segs) > 1 else "") + pwi.fmt_expr(e, 7, 40)
                                        for a, b, e in segs)
                        r.append(GridCell(txt, mono=True, size=7))
                    else:
                        r.append(GridCell(er.detail.get(lab, "n/a"), fill="absent", muted=True, size=7))
                rows.append(r)
            add_flow_table_slides(prs, "Derived expressions" + _grp_suffix(cmp, grp), ["expression"] + grp, rows,
                                  [2.6] + [nw] * len(grp), font_size=8,
                                  subtitle="functions inlined and terms collected (J/mol); formation relative to each "
                                           "database's own element references")
        rows = []
        for er in cmp.expressions:
            for i, lab in enumerate(cmp.labels):
                rows.append([GridCell(er.spec.label if i == 0 else "", bold=True, size=8), GridCell(lab, size=8),
                             GridCell(er.detail.get(lab, ""), size=8,
                                      fill="error" if str(er.detail.get(lab, "")).startswith("n/a") else None)])
            if er.ambiguity:
                rows.append([GridCell(""), GridCell("ambiguous", fill="no ref", size=8),
                             GridCell(er.ambiguity, size=8)])
        add_flow_table_slides(prs, "Derived expressions: references used", ["expression", "database", "resolved as"],
                              rows, [3.4, 2.4, 6.6], font_size=8)
    if diagram_png and os.path.exists(diagram_png):
        pu.add_image_slide(prs, f"{sys_} phase diagrams overlaid", diagram_png, subtitle=diagram_note or None)
    pu.save(prs, path)
    info["n_slides"] = len(prs.slides)
    info["groups"] = _db_groups(cmp, 4)
    return path, info


def _db_groups(cmp: SystemComparison, per: int) -> List[List[str]]:
    labs = cmp.labels
    n = -(-len(labs) // per)
    size = -(-len(labs) // n)
    return [labs[i:i + size] for i in range(0, len(labs), size)]


def _grp_suffix(cmp: SystemComparison, grp: List[str]) -> str:
    if len(grp) == len(cmp.labels):
        return ""
    i0, i1 = cmp.labels.index(grp[0]) + 1, cmp.labels.index(grp[-1]) + 1
    return f" [D{i0}–D{i1}]"


EQUAL_TITLE = "Numerically equal within tolerance (written differently)"


def equal_note(er: ExprRow) -> str:
    """Note of a row in the numerically-equal table: the largest difference instead of 'numerically equal'."""
    parts = [p for p in er.note.split("; ") if p and not p.startswith("numerically equal")]
    if math.isfinite(er.max_diff):
        parts.append(f"max |Δ| {er.max_diff:.3g}" + {"G": " J/mol", "TC": " K"}.get(er.ptype, ""))
    return "; ".join(parts)


def _written_slides(prs, cmp: SystemComparison, title: str, rows_in: List[ExprRow], info, subtitle, foot,
                    equal: bool = False):
    from .pptx_utils import GridCell, add_flow_table_slides
    for grp in _db_groups(cmp, 4):
        nw = (12.5 - 2.7) / len(grp)
        rows = []
        band = None
        for er in rows_in:
            if er.phase != band:
                rows.append(("band", er.phase))
                band = er.phase
            r = [GridCell(er.label, mono=True, size=8, sub=equal_note(er) if equal else er.note, bold=True)]
            for lab in grp:
                txt = er.shown_short.get(lab)
                if txt is None:
                    r.append(GridCell("—", fill="absent", muted=True))
                    continue
                fill = "differs" if er.cell_differs.get(lab) else ("same" if er.status == "equal" else None)
                r.append(GridCell(txt, mono=True, fill=fill, size=8 if er.expanded else None,
                                  sub=er.ref_text(lab, 24 if len(grp) > 2 else 30)))
            rows.append(r)
        t = title + _grp_suffix(cmp, grp)
        k = add_flow_table_slides(prs, t, ["parameter"] + grp, rows, [2.7] + [nw] * len(grp),
                                  font_size=9 if len(grp) <= 3 else 8, subtitle=subtitle, footnote=foot)
        info["slides"].append({"title": t, "rows": len(rows_in), "slides": k})


# =============================================================================================== diagrams
def overlay_diagrams(views: List[DbView], elements: List[str], out_png: str, tmin: float = 300.0,
                     tmax: float = 2300.0, progress=None) -> Tuple[Optional[str], str]:
    """Map the binary diagram of every database (in parallel) and overlay the phase boundaries."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from .diagrams import run_tasks

    if len(elements) != 2:
        return None, "diagram overlay needs exactly 2 elements"
    outdir = Path(out_png).parent / (Path(out_png).stem + "_maps")
    outdir.mkdir(parents=True, exist_ok=True)
    tasks = []
    for i, v in enumerate(views):
        png = str(outdir / f"D{i + 1}_{re.sub(r'[^A-Za-z0-9_.-]', '_', v.label)}.png")
        tasks.append({"name": v.label, "kind": "map", "kwargs": {"tdb": v.path, "elements": list(elements),
                                                                  "out_png": png, "tmin": tmin, "tmax": tmax,
                                                                  "temperature_unit": "K"}})
    res = run_tasks(tasks, timeout=900, progress=progress)
    fig, ax = plt.subplots(figsize=(8.5, 6.0))
    colors = plt.rcParams["axes.prop_cycle"].by_key()["color"]
    styles = ["-", "--", ":", "-."]
    notes = []
    el = elements[1]
    plotted = 0
    for i, (v, r, t) in enumerate(zip(views, res, tasks)):
        if not r.ok:
            notes.append(f"{v.label}: {r.error[:80]}")
            continue
        csvp = Path(t["kwargs"]["out_png"]).with_name(Path(t["kwargs"]["out_png"]).stem + "_boundaries.csv")
        if not csvp.exists():
            notes.append(f"{v.label}: no boundary data")
            continue
        import pandas as pd
        df = pd.read_csv(csvp)
        first = True
        for (_, _), g in df.groupby(["line", "phase"]):
            g = g.sort_values("point")
            ax.plot(g[f"X({el})"], g["T_K"], styles[i % 4], color=colors[i % len(colors)], lw=1.3,
                    label=v.label if first else None)
            first = False
        plotted += 1
    if not plotted:
        plt.close(fig)
        return None, "; ".join(notes)
    ax.set_xlabel(f"x({el.capitalize()})")
    ax.set_ylabel("T (K)")
    ax.set_xlim(0, 1)
    ax.set_ylim(tmin, tmax)
    ax.grid(alpha=0.25)
    ax.legend(loc="best", fontsize=9)
    ax.set_title("-".join(e.capitalize() for e in elements) + ": phase boundaries of each database")
    fig.savefig(out_png, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out_png, "; ".join(notes)


# =============================================================================================== tool
def run(paths, elements, output_pptx: str, workdir: Path, plot_diagrams: bool = False, progress=None,
        **kw) -> Dict[str, Any]:
    cmp = compare_system(paths, elements, workdir=workdir, **kw)
    base = Path(output_pptx).with_suffix("")
    files = {"pptx": output_pptx, "xlsx": str(base) + ".xlsx", "csv": str(base) + "_long.csv"}
    diagram, dnote = None, ""
    if plot_diagrams:
        diagram, dnote = overlay_diagrams(cmp.views, cmp.elements, str(base) + "_diagrams.png", progress=progress)
        if diagram:
            files["diagram_png"] = diagram
    write_xlsx(cmp, files["xlsx"])
    write_csv(cmp, files["csv"])
    _, info = write_pptx(cmp, files["pptx"], diagram, dnote)
    return {"comparison": cmp, "files": files, "pptx_info": info, "diagram_note": dnote}


def resolve_inputs(tb, tdbs=None, folder=None, pattern=None, elements=None) -> List[str]:
    if tdbs:
        return [tb._find_tdb(t) for t in tdbs]
    if not folder and not pattern:
        raise ValueError("give tdbs (explicit list) or folder/pattern")
    rows, _ = find_databases(tb.workdir, folder, pattern, elements, True)
    rows = [r for r in rows if not r["path"].lower().endswith(".dat")]
    if len(rows) < 2:
        raise ValueError(f"found {len(rows)} TDB(s) with {elements} in {folder or '.'} - need at least 2")
    return [str(tb.workdir / r["path"]) for r in rows]


def register(tb) -> None:
    S = {"type": "string"}
    STRS = {"type": "array", "items": S}

    @tb.add("find_databases", "List TDB/.dat files under the workdir or a folder: exact relative paths, elements, "
            "phases, parameter counts per subsystem of `elements`, header. `elements` keeps only databases with all "
            "of them.",
            {"folder": S, "pattern": {**S, "description": "glob on file name, e.g. 'Fe-Cr*.tdb'"},
             "elements": STRS, "recursive": {"type": "boolean"}}, [])
    def find_databases_tool(folder=None, pattern=None, elements=None, recursive: bool = True):
        _, text = find_databases(tb.workdir, folder, pattern, elements, recursive)
        return text, []

    @tb.add("compare_system", "N-way comparison of databases within one element subsystem: every parameter as "
            "written in each TDB (ESPEI VV#### replaced by values) with its reference, a differences-only table, "
            "a table of rows equal within tolerance but written differently, per-TDB tables, "
            "derived/formation expressions -> PPTX + Excel + CSV. Phases match by exact name; similar ones are "
            "listed as candidates: decide (ask_user if unsure), then re-call with phase_map.",
            {"tdbs": {**STRS, "description": "exact files; or give folder (+pattern)"}, "folder": S, "pattern": S,
             "elements": STRS, "temperatures": {"type": "array", "items": {"type": "number"},
                                                "description": "K, for the Excel values sheet"},
             "phases": STRS,
             "phase_map": {"type": "object", "description": "{native: canonical} or {db file: {native: canonical}}"},
             "expressions": {"type": "array", "items": {"anyOf": [S, {"type": "object"}]},
                             "description": "e.g. 'G(SIGMA,FE:CR:CR;0)-10*GHSERFE-20*GHSERCR' or {label, expr}"},
             "formation": {"type": "array", "items": {"type": "object"},
                           "description": "[{phase, constituents:'FE:CR', per:'formula_unit'|'mole_of_atoms'}]: "
                                          "G(end-member) minus each database's element references"},
             "reference": S, "include_identical": {"type": "boolean"}, "rtol": {"type": "number"},
             "output_pptx": S, "plot_diagrams": {"type": "boolean", "description": "overlay binary diagrams (slow)"}},
            ["elements"])
    def compare_system_tool(elements, tdbs=None, folder=None, pattern=None, temperatures=None, phases=None,
                            phase_map=None, expressions=None, formation=None, reference=None,
                            include_identical: bool = True, rtol: float = 1e-4, output_pptx=None,
                            plot_diagrams: bool = False):
        paths = resolve_inputs(tb, tdbs, folder, pattern, elements)
        sys_ = "".join(e.capitalize() for e in sorted(e.upper() for e in elements))
        out = tb.outp(output_pptx if not output_pptx or output_pptx.endswith(".pptx") else output_pptx + ".pptx",
                      f"compare_system_{sys_}_{len(paths)}db.pptx")
        r = run(paths, elements, out, tb.workdir, plot_diagrams=plot_diagrams, progress=tb.ui.progress,
                temperatures=temperatures, phases=phases, phase_map=phase_map, expressions=expressions,
                formation=formation, reference=reference, include_identical=include_identical, rtol=rtol)
        text = summary_text(r["comparison"], r["files"], tb.rel)
        if r["diagram_note"]:
            text += f"\ndiagrams: {r['diagram_note']}"
        imgs = [r["files"]["diagram_png"]] if r["files"].get("diagram_png") else []
        return text, imgs


# =============================================================================================== show_code support
def standalone_plan(tb, args: Dict[str, Any]) -> Dict[str, Any]:
    """Everything the standalone script needs that pycalphad's Database cannot tell it."""
    paths = resolve_inputs(tb, args.get("tdbs"), args.get("folder"), args.get("pattern"), args.get("elements"))
    cmp = compare_system(paths, args["elements"], temperatures=args.get("temperatures"), phases=args.get("phases"),
                         phase_map=args.get("phase_map"), expressions=args.get("expressions"),
                         formation=args.get("formation"), reference=args.get("reference"),
                         include_identical=args.get("include_identical", True), rtol=args.get("rtol", 1e-4),
                         workdir=tb.workdir)
    fixes: Dict[str, Dict[str, List]] = {}
    for v in cmp.views:
        for k, r in v.rows.items():
            if r.order_changed:
                p = next(x for x in v.tdb.parameters if x.raw_name == r.raw_name)
                pv = v.phases[r.phase]
                # pycalphad sorts constituents but keeps the order index and sign as written
                raw_order = int(re.search(r";\s*(\d+)\s*\)\s*$", r.raw_name).group(1)) if ";" in r.raw_name else 0
                red = tuple(sl for i, sl in enumerate(p.key.constituents) if i not in pv.va_only)
                pk = make_key(p.key.ptype, pv.canon, red, raw_order, p.key.diffusing)
                fixes.setdefault(v.label, {})[pk] = [k, r.sign]
    return {
        "labels": cmp.labels, "paths": [v.path for v in cmp.views], "elements": cmp.elements,
        "temperatures": cmp.temps, "reference": cmp.ref, "phases": cmp.phases if args.get("phases") else None,
        "phase_map": {lab: m for lab, m in cmp.phase_map.items()},
        "include_identical": args.get("include_identical", True), "rtol": cmp.rtol,
        "key_fixes": fixes,
        "expressions": [{"label": er.spec.label, "per_db": er.plan} for er in cmp.expressions],
    }
