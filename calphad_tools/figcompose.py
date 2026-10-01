"""Figure composition tools for the chat agent: multi-panel figures, slide decks, animations, raster edits and
multi-page PDFs built from existing images (tool-made or foreign), offline and deterministic.

Tools (registered by :func:`register`):

* ``compose_images``   - (a)-(d) labelled multi-panel figure from existing images, saved through matplotlib so the
                         composite itself is stored as an editable figure (PNG + optional PDF/SVG).
* ``figures_to_pptx``  - new 16:9 deck or slides appended to an existing deck, 1-6 images per slide (auto grid).
* ``animate``          - GIF/MP4 from image frames or from a parameter sweep of another tool.
* ``image_tools``      - Pillow raster operations (crop/trim, resize, pad, border, transparency, rotate, convert).
* ``images_to_pdf``    - multi-page PDF (vector pages for tool-made figures when the page is fitted).
"""
from __future__ import annotations

import math
import os
import shutil
import subprocess
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np

MAX_FRAMES = 60
MAX_PIXELS = 60_000_000                      # guard against absurd output sizes
RASTER_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp")


# ---------------------------------------------------------------------------
# image helpers
# ---------------------------------------------------------------------------
def check_images(paths: Sequence[str]) -> None:
    """Raise one error listing every missing or unreadable input."""
    missing = [p for p in paths if not os.path.isfile(p)]
    if missing:
        raise FileNotFoundError(f"{len(missing)} image(s) not found: {missing}")
    bad = []
    for p in paths:
        if Path(p).suffix.lower() in (".pdf", ".svg", ".eps"):
            bad.append(f"{p} (vector file; use the PNG made alongside it or view_pdf_page)")
    if bad:
        raise ValueError("unsupported input: " + "; ".join(bad))


def open_image(path: str):
    """First frame of an image file as RGB or RGBA (palette/greyscale/CMYK converted)."""
    from PIL import Image, ImageOps
    with Image.open(path) as im:
        im.seek(0)
        im = ImageOps.exif_transpose(im)
        has_alpha = im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info)
        return im.convert("RGBA" if has_alpha else "RGB")


def flatten(im, color=(255, 255, 255)):
    """RGB copy with any transparency composited onto ``color``."""
    from PIL import Image
    if im.mode != "RGBA":
        return im.convert("RGB")
    bg = Image.new("RGB", im.size, color)
    bg.paste(im, mask=im.getchannel("A"))
    return bg


def trim_box(im, tolerance: float = 12, pad: int = 4) -> Optional[Tuple[int, int, int, int]]:
    """Bounding box (l, t, r, b) of the content that differs from the uniform border colour (or transparency)."""
    arr = np.asarray(im.convert("RGBA")).astype(np.int16)
    h, w = arr.shape[:2]
    corners = [tuple(arr[y, x]) for y, x in ((0, 0), (0, w - 1), (h - 1, 0), (h - 1, w - 1))]
    bg = np.array(max(set(corners), key=corners.count))
    if bg[3] < 16:                                   # transparent background
        mask = arr[:, :, 3] > tolerance
    else:
        mask = (np.abs(arr[:, :, :3] - bg[:3]).max(axis=2) > tolerance) & (arr[:, :, 3] > 16)
    rows, cols = np.where(mask.any(axis=1))[0], np.where(mask.any(axis=0))[0]
    if rows.size == 0:
        return None
    return (max(int(cols[0]) - pad, 0), max(int(rows[0]) - pad, 0),
            min(int(cols[-1]) + 1 + pad, w), min(int(rows[-1]) + 1 + pad, h))


def trim(im, tolerance: float = 12, pad: int = 4):
    box = trim_box(im, tolerance, pad)
    return im.crop(box) if box else im


def parse_color(c, default=(255, 255, 255)):
    """Any matplotlib colour (name, '#hex', [r,g,b] 0-255 or 0-1) -> RGB tuple 0-255."""
    if c is None:
        return default
    if isinstance(c, (list, tuple)):
        v = [float(x) for x in c][:3]
        if len(v) != 3:
            raise ValueError(f"colour {c!r}: need 3 components")
        return tuple(int(round(x * 255)) if max(v) <= 1 else int(round(x)) for x in v)
    from matplotlib.colors import to_rgb
    try:
        return tuple(int(round(x * 255)) for x in to_rgb(c))
    except ValueError:
        raise ValueError(f"{c!r} is not a colour (use a name like 'white', '#rrggbb' or [r, g, b])") from None


def panel_labels(style, n: int) -> List[str]:
    """'a' 'A' '(a)' 'a)' 'A.' 'i' '(i)' '1' ... or an explicit list; None/'none'/False -> no labels."""
    if style in (None, False, "", "none", "None"):
        return [""] * n
    if isinstance(style, (list, tuple)):
        labels = [str(s) for s in style]
        return (labels + [""] * n)[:n]
    s = str(style)
    idx = next((i for i, ch in enumerate(s) if ch in "aAiI1"), None)
    if idx is None:
        raise ValueError(f"labels {style!r}: use 'a', 'A', '(a)', 'a)', 'i', '1', 'none' or a list of strings")
    pre, core, post = s[:idx], s[idx], s[idx + 1:]

    def roman(k):
        vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
        out = ""
        for v, r in vals:
            while k >= v:
                out, k = out + r, k - v
        return out
    out = []
    for k in range(n):
        if core in "aA":
            t = ""
            m = k
            while True:
                t = chr(ord("a") + m % 26) + t
                m = m // 26 - 1
                if m < 0:
                    break
            t = t.upper() if core == "A" else t
        elif core in "iI":
            t = roman(k + 1).upper() if core == "I" else roman(k + 1)
        else:
            t = str(k + 1)
        out.append(pre + t + post)
    return out


def grid_geometry(aspects: Sequence[float], rows: int, cols: int, avail_w: float,
                  gap: float = 0.0) -> Tuple[List[Tuple[float, float]], List[float]]:
    """Justified rows: every image in a row gets the same height and the row fills ``avail_w``.

    Returns ([(x offset, width) per image], [row heights]).  Uniform images give an exact grid; a partly
    filled last row is not taller than the full rows and is centred.
    """
    n = len(aspects)
    boxes: List[Tuple[float, float]] = [(0.0, 0.0)] * n
    heights: List[float] = []
    full_h = []
    for r in range(rows):
        idx = list(range(r * cols, min(r * cols + cols, n)))
        if not idx:
            heights.append(0.0)
            continue
        h = (avail_w - (len(idx) - 1) * gap) / sum(1.0 / aspects[i] for i in idx)
        if len(idx) == cols:
            full_h.append(h)
        heights.append(h)
    cap = max(full_h) if full_h else None
    for r in range(rows):
        idx = list(range(r * cols, min(r * cols + cols, n)))
        if not idx:
            continue
        if cap is not None and len(idx) < cols and heights[r] > cap:
            heights[r] = cap
        widths = [heights[r] / aspects[i] for i in idx]
        x = (avail_w - sum(widths) - (len(idx) - 1) * gap) / 2
        for i, w in zip(idx, widths):
            boxes[i] = (x, w)
            x += w + gap
    return boxes, heights


def auto_layout(aspects: Sequence[float], target: float = 0.7) -> Tuple[int, int]:
    """[rows, cols] whose composite aspect (h/w) is close to ``target``, with few empty cells."""
    n = len(aspects)
    best = None
    for c in range(1, n + 1):
        r = math.ceil(n / c)
        if (r - 1) * c >= n:
            continue
        _, heights = grid_geometry(aspects, r, c, 1.0)
        score = abs(math.log(sum(heights) / target)) + 0.5 * (r * c - n)
        if best is None or score < best[0] - 1e-9:
            best = (score, r, c)
    return best[1], best[2]


def best_grid(sizes: Sequence[Tuple[int, int]], box_w: float, box_h: float, gap: float) -> Tuple[int, int]:
    """[rows, cols] maximising the total displayed image area inside a box."""
    n = len(sizes)
    best = None
    for c in range(1, n + 1):
        r = math.ceil(n / c)
        if (r - 1) * c >= n:
            continue
        cw, ch = (box_w - (c - 1) * gap) / c, (box_h - (r - 1) * gap) / r
        if cw <= 0 or ch <= 0:
            continue
        area = sum(min(cw / w, ch / h) ** 2 * w * h for w, h in sizes)
        if best is None or area > best[0] * 1.0001:
            best = (area, r, c)
    return (best[1], best[2]) if best else (n, 1)


def _font(size: int):
    from PIL import ImageFont
    import matplotlib
    p = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans.ttf"
    try:
        return ImageFont.truetype(str(p), size)
    except OSError:
        return ImageFont.load_default(size=size)


# ---------------------------------------------------------------------------
# compose
# ---------------------------------------------------------------------------
LABEL_POSITIONS = ("outside", "top-left", "top-right", "bottom-left", "bottom-right")


def compose(paths: Sequence[str], out_png: str, layout=None, labels="(a)", titles=None, suptitle=None,
            trim_whitespace: bool = True, spacing: float = 0.15, label_position: str = "outside",
            label_fontsize: float = 12, width_in: Optional[float] = None, dpi: int = 200, formats=None) -> Dict:
    import matplotlib.pyplot as plt
    from PIL import Image

    n = len(paths)
    if n == 0:
        raise ValueError("images is empty")
    if label_position not in LABEL_POSITIONS:
        raise ValueError(f"label_position must be one of {LABEL_POSITIONS}")
    ims = [open_image(p) for p in paths]
    if trim_whitespace:
        ims = [trim(im) for im in ims]
    aspects = [im.size[1] / im.size[0] for im in ims]
    if layout:
        rows, cols = int(layout[0]), int(layout[1])
        if rows * cols < n:
            raise ValueError(f"layout {rows}x{cols} has fewer cells than {n} images")
    else:
        rows, cols = auto_layout(aspects)
    labs = panel_labels(labels, n)
    titles = list(titles or [])
    titles = [str(t) if t else "" for t in (titles + [""] * n)[:n]]
    width_in = float(width_in or min(max(3.5 * cols, 3.5), 16.0))
    margin, gap = 0.06, max(float(spacing), 0.0)
    avail = width_in - 2 * margin - (cols - 1) * gap
    if avail / cols <= 0.2:
        raise ValueError("width_in too small for this layout/spacing")
    boxes, row_h = grid_geometry(aspects, rows, cols, avail, gap)
    fs = float(label_fontsize)
    strip = (fs * 1.6 / 72) if (label_position == "outside" and any(labs)) or any(titles) else 0.0
    sup_h = fs * 2.2 / 72 if suptitle else 0.0
    grid = [list(range(r * cols, min(r * cols + cols, n))) for r in range(rows)]
    height_in = margin * 2 + sup_h + sum(row_h) + rows * strip + (rows - 1) * gap
    if width_in * height_in * dpi * dpi > MAX_PIXELS:
        raise ValueError(f"composite would be {width_in:.1f}x{height_in:.1f} in at {dpi} dpi; lower dpi or width_in")

    fig = plt.figure(figsize=(width_in, height_in), dpi=dpi, facecolor="white")
    y = height_in - margin - sup_h                  # top of the current row, inches from the bottom
    panels = []
    try:
        for r, idx in enumerate(grid):
            y_img_top = y - strip
            for c, i in enumerate(idx):
                w_in = boxes[i][1]
                h_in = row_h[r]
                x0 = margin + boxes[i][0]
                y0 = y_img_top - (row_h[r] - h_in) / 2 - h_in       # vertically centred in the row
                im = ims[i]
                px = (max(int(round(w_in * dpi)), 1), max(int(round(h_in * dpi)), 1))
                if im.size[0] > px[0] * 1.05:                          # keep the stored figure small
                    im = im.resize(px, Image.LANCZOS)
                ax = fig.add_axes([x0 / width_in, y0 / height_in, w_in / width_in, h_in / height_in])
                ax.imshow(np.asarray(im), interpolation="antialiased", aspect="auto")
                ax.set_axis_off()
                ax.set_label(f"panel {labs[i] or i + 1}: {Path(paths[i]).name}")
                ax_l, ax_t = x0 / width_in, (y0 + h_in) / height_in
                if titles[i]:
                    fig.text((x0 + w_in / 2) / width_in, (y_img_top + strip / 2) / height_in, titles[i],
                             ha="center", va="center", fontsize=fs * 0.9)
                if labs[i]:
                    kw = dict(fontsize=fs, fontweight="bold")
                    if label_position == "outside":
                        fig.text(ax_l, (y_img_top + strip / 2) / height_in, labs[i], ha="left", va="center", **kw)
                    else:
                        vert, horiz = label_position.split("-")
                        ax.text(0.02 if horiz == "left" else 0.98, 0.98 if vert == "top" else 0.02, labs[i],
                                transform=ax.transAxes, ha=horiz, va=vert,
                                bbox=dict(boxstyle="round,pad=0.2", fc="white", ec="none", alpha=0.75), **kw)
                panels.append({"panel": labs[i] or str(i + 1), "image": paths[i],
                               "size_in": [round(w_in, 2), round(h_in, 2)]})
            y -= strip + row_h[r] + gap
        if suptitle:
            fig.text(0.5, (height_in - margin - sup_h / 2) / height_in, suptitle, ha="center", va="center",
                     fontsize=fs * 1.15, fontweight="bold")
        files = [out_png]
        fig.savefig(out_png, dpi=dpi, facecolor="white")
        for ext in formats or []:
            ext = ext.lower().lstrip(".")
            if ext not in ("pdf", "svg", "eps", "png", "jpg", "jpeg", "tiff"):
                raise ValueError(f"unsupported format {ext!r} (pdf, svg, eps, png, jpg, tiff)")
            p = str(Path(out_png).with_suffix("." + ext))
            if p != out_png:
                fig.savefig(p, dpi=dpi, facecolor="white")
                files.append(p)
    finally:
        plt.close(fig)
    return {"files": files, "layout": [rows, cols], "size_in": [round(width_in, 2), round(height_in, 2)],
            "size_px": [int(round(width_in * dpi)), int(round(height_in * dpi))], "panels": panels}


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------
def _blank_layout(prs):
    """The layout with the fewest placeholders (layout 6 'Blank' in the default template)."""
    lays = list(prs.slide_layouts)
    if len(lays) > 6 and len(lays[6].placeholders) == 0:
        return lays[6]
    return min(lays, key=lambda l: len(l.placeholders))


def _pptx_picture_path(path: str, tmpdir: Path) -> str:
    """python-pptx reads PNG/JPEG/GIF/BMP/TIFF; convert anything else Pillow can open to PNG."""
    if Path(path).suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"):
        return path
    tmpdir.mkdir(parents=True, exist_ok=True)
    out = tmpdir / (Path(path).stem + "_conv.png")
    open_image(path).save(out)
    return str(out)


def add_images_slide(prs, spec: Dict, tmpdir: Path) -> Dict:
    from PIL import Image
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches
    from . import pptx_utils as pu

    sw, sh = int(prs.slide_width), int(prs.slide_height)
    s = prs.slides.add_slide(_blank_layout(prs))
    for ph in list(s.placeholders):                 # foreign templates: drop empty placeholders
        ph._element.getparent().remove(ph._element)
    title, subtitle = spec.get("title"), spec.get("subtitle")
    m = Inches(0.5)
    top = Inches(0.35)
    if title:
        pu._textbox(s, m, Inches(0.3), sw - 2 * m, Inches(0.7), str(title), size=26, bold=True)
        bar = s.shapes.add_shape(1, m, Inches(1.0), Inches(1.2), Emu(38000))
        bar.fill.solid()
        bar.fill.fore_color.rgb = pu.ACCENT
        bar.line.fill.background()
        top = Inches(1.3)
        if subtitle:
            pu._textbox(s, m, Inches(1.08), sw - 2 * m, Inches(0.4), str(subtitle), size=13, color=pu.MUTED)
            top = Inches(1.55)
    caption = spec.get("caption")
    bottom = sh - Inches(0.25) - (Inches(0.55) if caption else 0)
    left, right = Inches(0.4), sw - Inches(0.4)
    bullets = spec.get("bullets") or []
    if bullets:
        text_w = int((right - left) * (0.36 if spec.get("images") else 1.0))
        pu._textbox(s, left, top, text_w, bottom - top, "\n".join(f"•  {b}" for b in bullets),
                    size=int(spec.get("bullet_size", 16)))
        left += text_w + Inches(0.2)
    paths = spec.get("images") or []
    sizes = []
    for p in paths:
        with Image.open(p) as im:
            sizes.append(im.size)
    placed = 0
    if paths:
        gap = Inches(0.15)
        bw, bh = right - left, bottom - top
        if spec.get("layout"):
            rows, cols = int(spec["layout"][0]), int(spec["layout"][1])
            if rows * cols < len(paths):
                raise ValueError(f"slide {spec.get('title')!r}: layout {rows}x{cols} < {len(paths)} images")
        else:
            rows, cols = best_grid(sizes, bw, bh, gap)
        cw, ch = (bw - (cols - 1) * gap) / cols, (bh - (rows - 1) * gap) / rows
        # centre partially filled last row
        for i, (p, (w, h)) in enumerate(zip(paths, sizes)):
            r, c = divmod(i, cols)
            in_row = min(cols, len(paths) - r * cols)
            x_off = (cols - in_row) * (cw + gap) / 2
            sc = min(cw / w, ch / h)
            pw, ph = int(w * sc), int(h * sc)
            x = int(left + x_off + c * (cw + gap) + (cw - pw) / 2)
            y = int(top + r * (ch + gap) + (ch - ph) / 2)
            s.shapes.add_picture(_pptx_picture_path(p, tmpdir), x, y, pw, ph)
            placed += 1
    if caption:
        pu._textbox(s, Inches(0.5), sh - Inches(0.75), sw - Inches(1.0), Inches(0.5), str(caption), size=13,
                    color=pu.MUTED, align=PP_ALIGN.CENTER)
    if spec.get("notes"):
        s.notes_slide.notes_text_frame.text = str(spec["notes"])
    return {"title": title or "", "images": placed}


# ---------------------------------------------------------------------------
# animation
# ---------------------------------------------------------------------------
def _set_path(args: Dict, path: str, value) -> None:
    """Set args['a']['b'][0] = value for path 'a.b.0' (creates dicts as needed)."""
    keys = path.split(".")
    cur: Any = args
    for k in keys[:-1]:
        if isinstance(cur, list):
            cur = cur[int(k)]
        else:
            if k not in cur or cur[k] is None:
                cur[k] = {}
            cur = cur[k]
    last = keys[-1]
    if isinstance(cur, list):
        cur[int(last)] = value
    else:
        cur[last] = value


def build_frames(images, captions: Optional[List[str]] = None, max_width: int = 1200, font_size: int = 0):
    """Flatten, pad to a common canvas (centred, white) and add a caption strip on top."""
    from PIL import Image, ImageDraw
    ims = [flatten(im) for im in images]
    ims = [im.resize((max_width, max(int(round(im.size[1] * max_width / im.size[0])), 1)), Image.LANCZOS)
           if im.size[0] > max_width else im for im in ims]
    W = max(im.size[0] for im in ims)
    H = max(im.size[1] for im in ims)
    has_cap = bool(captions) and any(captions)
    fsz = font_size or max(14, min(40, W // 32))
    strip = int(fsz * 1.8) if has_cap else 0
    font = _font(fsz) if has_cap else None
    out = []
    for k, im in enumerate(ims):
        canvas = Image.new("RGB", (W, H + strip), "white")
        canvas.paste(im, ((W - im.size[0]) // 2, strip + (H - im.size[1]) // 2))
        if has_cap and captions[k]:
            d = ImageDraw.Draw(canvas)
            d.text((W // 2, strip // 2), captions[k], fill="black", font=font, anchor="mm")
        out.append(canvas)
    return out


def write_gif(frames, out: str, duration_ms: int, loop: int) -> None:
    frames[0].save(out, save_all=True, append_images=frames[1:], duration=int(duration_ms), loop=int(loop),
                   optimize=False, disposal=1)


def write_mp4(frames, out: str, fps: float) -> str:
    """Encode with the local ffmpeg binary (raw RGB frames on stdin, no network)."""
    exe = shutil.which("ffmpeg")
    if not exe:
        raise ValueError("ffmpeg is not installed: use a .gif output")
    W, H = frames[0].size
    W2, H2 = W + W % 2, H + H % 2                     # yuv420p needs even dimensions
    from PIL import Image
    raw = b"".join((f if (W, H) == (W2, H2) else _pad_to(f, W2, H2)).tobytes() for f in frames)
    last = ""
    for codec in ("libx264", "mpeg4"):
        cmd = [exe, "-y", "-loglevel", "error", "-f", "rawvideo", "-pix_fmt", "rgb24", "-s", f"{W2}x{H2}",
               "-r", f"{fps:g}", "-i", "-", "-an", "-c:v", codec, "-pix_fmt", "yuv420p"]
        cmd += ["-crf", "18"] if codec == "libx264" else ["-q:v", "3"]
        r = subprocess.run(cmd + [out], input=raw, capture_output=True, timeout=300)   # noqa: S603
        if r.returncode == 0:
            return codec
        last = r.stderr.decode(errors="replace")[-400:]
    raise RuntimeError(f"ffmpeg failed: {last}")


def _pad_to(im, W, H):
    from PIL import Image
    c = Image.new("RGB", (W, H), "white")
    c.paste(im, (0, 0))
    return c


# ---------------------------------------------------------------------------
# raster operations
# ---------------------------------------------------------------------------
IMAGE_OPS_HELP = ("crop {box:[l,t,r,b] px (or fractions with units:'fraction') | trim:true, tolerance, pad}; "
                  "resize {width|height|scale}; pad {px (or [l,t,r,b]), color}; border {px, color}; "
                  "transparent_background {color:'white', tolerance}; grayscale; rotate {degrees (ccw), expand}; "
                  "convert {format: png|jpg|pdf|tiff|webp, quality}")


def apply_image_ops(im, ops: Sequence[Dict]) -> Tuple[Any, Optional[str], Dict, List[str]]:
    """Returns (image, output format or None, save options, notes)."""
    from PIL import Image, ImageOps
    fmt, save_opts, notes = None, {}, []
    for k, op in enumerate(ops, start=1):
        name = str(op.get("op", "")).lower()
        try:
            if name == "crop":
                if op.get("trim"):
                    box = trim_box(im, float(op.get("tolerance", 12)), int(op.get("pad", 4)))
                    if box is None:
                        notes.append(f"op {k} trim: image is uniform, nothing trimmed")
                        continue
                else:
                    if "box" not in op:
                        raise ValueError("needs box [l,t,r,b] or trim: true")
                    box = [float(v) for v in op["box"]]
                    if op.get("units") == "fraction":
                        w, h = im.size
                        box = [box[0] * w, box[1] * h, box[2] * w, box[3] * h]
                    box = [int(round(v)) for v in box]
                    if not (0 <= box[0] < box[2] <= im.size[0] and 0 <= box[1] < box[3] <= im.size[1]):
                        raise ValueError(f"box {box} outside the {im.size[0]}x{im.size[1]} image (l<r, t<b)")
                im = im.crop(tuple(box))
            elif name == "resize":
                w, h = im.size
                if op.get("scale"):
                    nw, nh = w * float(op["scale"]), h * float(op["scale"])
                elif op.get("width") and op.get("height"):
                    nw, nh = float(op["width"]), float(op["height"])
                elif op.get("width"):
                    nw = float(op["width"])
                    nh = h * nw / w
                elif op.get("height"):
                    nh = float(op["height"])
                    nw = w * nh / h
                else:
                    raise ValueError("needs width, height or scale")
                nw, nh = max(int(round(nw)), 1), max(int(round(nh)), 1)
                if nw * nh > MAX_PIXELS:
                    raise ValueError(f"{nw}x{nh} px is too large")
                im = im.resize((nw, nh), Image.LANCZOS)
            elif name in ("pad", "border"):
                px = op.get("px", 20 if name == "pad" else 2)
                px = [int(px)] * 4 if not isinstance(px, (list, tuple)) else [int(v) for v in px]
                if len(px) != 4 or min(px) < 0:
                    raise ValueError("px must be a non-negative number or [left, top, right, bottom]")
                color = parse_color(op.get("color"), (255, 255, 255) if name == "pad" else (0, 0, 0))
                fill = color + ((255,) if im.mode == "RGBA" else ())
                im = ImageOps.expand(im, border=tuple(px), fill=fill)
            elif name == "transparent_background":
                color = np.array(parse_color(op.get("color", "white")))
                tol = float(op.get("tolerance", 10))
                arr = np.asarray(im.convert("RGBA")).copy()
                m = np.abs(arr[:, :, :3].astype(int) - color).max(axis=2) <= tol
                arr[m, 3] = 0
                im = Image.fromarray(arr, "RGBA")
                notes.append(f"op {k}: {int(m.sum())} px made transparent")
            elif name == "grayscale":
                im = im.convert("LA" if im.mode == "RGBA" else "L")
            elif name == "rotate":
                deg = float(op.get("degrees", 90))
                fill = (255, 255, 255, 0) if im.mode == "RGBA" else "white"
                if im.mode in ("L", "LA"):
                    fill = 255 if im.mode == "L" else (255, 0)
                im = im.rotate(deg, resample=Image.BICUBIC, expand=bool(op.get("expand", True)), fillcolor=fill)
            elif name == "convert":
                fmt = str(op.get("format", "png")).lower().lstrip(".")
                fmt = {"jpeg": "jpg", "tif": "tiff"}.get(fmt, fmt)
                if fmt not in ("png", "jpg", "pdf", "tiff", "webp"):
                    raise ValueError("format must be png, jpg, pdf, tiff or webp")
                if op.get("quality"):
                    save_opts["quality"] = int(op["quality"])
            else:
                raise ValueError(f"unknown op {name!r}")
        except ValueError as e:
            raise ValueError(f"operation {k} ({name}): {e}. Syntax: {IMAGE_OPS_HELP}") from None
    return im, fmt, save_opts, notes


def save_raster(im, out: str, fmt: str, opts: Dict, dpi: Optional[int] = None) -> List[str]:
    notes = []
    if fmt in ("jpg", "pdf") and im.mode in ("RGBA", "LA", "P"):
        if im.mode != "P" and np.asarray(im.getchannel("A")).min() < 255:
            notes.append(f"transparency flattened onto white ({fmt} has no alpha)")
        im = flatten(im.convert("RGBA"))
    pil_fmt = {"jpg": "JPEG", "png": "PNG", "pdf": "PDF", "tiff": "TIFF", "webp": "WEBP"}[fmt]
    kw = dict(opts)
    if fmt == "jpg":
        kw.setdefault("quality", 92)
    if fmt == "pdf":
        kw["resolution"] = float(dpi or 150)
    if dpi and fmt in ("png", "jpg", "tiff"):
        kw["dpi"] = (dpi, dpi)
    im.save(out, pil_fmt, **kw)
    return notes


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
PAGE_SIZES = {"letter": (8.5, 11.0), "a4": (8.27, 11.69)}


def images_pdf(paths: Sequence[str], out: str, page_size: str = "fit", titles=None, vector: bool = True,
               dpi: int = 150) -> Dict:
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from .figures import load_figure

    if page_size not in ("fit", "letter", "a4"):
        raise ValueError("page_size must be fit, letter or a4")
    titles = [str(t) if t else "" for t in (list(titles or []) + [""] * len(paths))[:len(paths)]]
    kinds = []
    with PdfPages(out) as pdf:
        for p, title in zip(paths, titles):
            fig = load_figure(p) if (vector and page_size == "fit" and p.lower().endswith(".png")) else None
            if fig is not None:
                if title:
                    fig.text(0.5, 1.01, title, ha="center", va="bottom", fontsize=13, fontweight="bold")
                pdf.savefig(fig, bbox_inches="tight")
                plt.close(fig)
                kinds.append("vector")
                continue
            im = flatten(open_image(p))
            arr = np.asarray(im)
            h, w = arr.shape[:2]
            th = 0.45 if title else 0.0
            if page_size == "fit":
                pw, ph = w / dpi, h / dpi + th
                box = [0, 0, 1, (h / dpi) / ph]
            else:
                pw, ph = PAGE_SIZES[page_size]
                if w > h:
                    pw, ph = ph, pw                        # landscape page for landscape images
                m = 0.5
                bw, bh = pw - 2 * m, ph - 2 * m - th
                sc = min(bw / w, bh / h)
                iw, ih = w * sc, h * sc
                box = [(pw - iw) / 2 / pw, (m + (bh - ih) / 2) / ph, iw / pw, ih / ph]
            fig = plt.figure(figsize=(pw, ph), dpi=dpi)
            ax = fig.add_axes(box)
            ax.imshow(arr, interpolation="none", aspect="auto")
            ax.set_axis_off()
            if title:
                fig.text(0.5, 1 - (th / 2) / ph if page_size == "fit" else 1 - (0.5 + th / 2) / ph, title,
                         ha="center", va="center", fontsize=13, fontweight="bold")
            pdf.savefig(fig)
            plt.close(fig)
            kinds.append("raster")
    return {"pages": len(paths), "kinds": kinds}


# ---------------------------------------------------------------------------
# tools
# ---------------------------------------------------------------------------
def register(tb) -> None:
    S = {"type": "string"}
    N = {"type": "number"}
    I = {"type": "integer"}
    B = {"type": "boolean"}
    STRS = {"type": "array", "items": S}
    IMGS = {"type": "array", "items": S, "description": "image paths (PNG/JPG/...)"}

    def resolve_images(images) -> List[str]:
        if isinstance(images, str):
            images = [images]
        paths = [tb.p(p) for p in images or []]
        if not paths:
            raise ValueError("no images given")
        check_images(paths)
        return paths

    def _tmpdir() -> Path:
        return Path(tb.out) / ".figcompose_tmp"

    # --------------------------------------------------------- compose_images
    @tb.add("compose_images",
            "Combine existing images (tool plots or foreign PNG/JPG) into one multi-panel figure with (a),(b).. "
            "labels; auto grid, aspect ratios kept, whitespace trimmed. Saved as an editable figure (PNG + formats).",
            {"images": IMGS, "output": S,
             "layout": {"type": "array", "items": I, "description": "[rows, cols] (default auto)"},
             "labels": {"description": "'(a)' (default) | 'a' | 'A' | 'a)' | 'i' | 'none' | list of strings"},
             "titles": {**STRS, "description": "per-panel titles"}, "suptitle": S,
             "formats": {**STRS, "description": "extra formats e.g. pdf, svg"},
             "trim_whitespace": B, "spacing": {**N, "description": "gap between panels, inches (0.15)"},
             "label_position": {"type": "string", "enum": list(LABEL_POSITIONS)},
             "label_fontsize": N, "width_in": {**N, "description": "figure width (default 3.5 in per column)"},
             "dpi": I},
            ["images", "output"])
    def compose_images(images, output: str, layout=None, labels="(a)", titles=None, suptitle=None, formats=None,
                       trim_whitespace: bool = True, spacing: float = 0.15, label_position: str = "outside",
                       label_fontsize: float = 12, width_in: Optional[float] = None, dpi: Optional[int] = None):
        paths = resolve_images(images)
        formats = list(formats or [])
        ext = Path(output).suffix.lower()
        if ext in (".pdf", ".svg", ".eps", ".tiff"):
            formats.append(ext[1:])
            output = str(Path(output).with_suffix(".png"))
        elif ext not in (".png", ".jpg", ".jpeg"):
            output += ".png"
        out = tb.outp(output, "composite.png")
        r = compose(paths, out, layout=layout, labels=labels, titles=titles, suptitle=suptitle,
                    trim_whitespace=trim_whitespace, spacing=spacing, label_position=label_position,
                    label_fontsize=label_fontsize, width_in=width_in, dpi=int(dpi or 200), formats=formats)
        lines = [f"saved {', '.join(tb.rel(f) for f in r['files'])}",
                 f"layout {r['layout'][0]}x{r['layout'][1]}, {r['size_in'][0]}x{r['size_in'][1]} in, "
                 f"{r['size_px'][0]}x{r['size_px'][1]} px"]
        lines += [f"  {p['panel']}: {tb.rel(p['image'])}" for p in r["panels"]]
        lines.append("Editable: edit_figure/describe_figure work on the composite (each panel is an image axes).")
        return "\n".join(lines), [r["files"][0]]

    # --------------------------------------------------------- figures_to_pptx
    slide_schema = {"type": "object", "properties": {
        "title": S, "subtitle": S, "images": {**IMGS, "description": "1-6 images, auto grid"},
        "layout": {"type": "array", "items": I, "description": "[rows, cols]"},
        "caption": S, "notes": {**S, "description": "speaker notes"}, "bullets": STRS}}

    @tb.add("figures_to_pptx",
            "Put images on PowerPoint slides (16:9, same style as the other decks): new deck, or append slides to "
            "an existing .pptx (append_to). Per slide: title, 1-6 images (auto grid, no distortion), caption, "
            "speaker notes, bullets.",
            {"slides": {"type": "array", "items": slide_schema}, "output": {**S, "description": ".pptx (default: "
             "overwrite append_to)"}, "append_to": {**S, "description": "existing .pptx to add slides to"},
             "title_slide": {"type": "object", "properties": {"title": S, "subtitle": S}}},
            ["slides"])
    def figures_to_pptx(slides, output: Optional[str] = None, append_to: Optional[str] = None, title_slide=None):
        from pptx import Presentation
        from . import pptx_utils as pu
        if not slides and not title_slide:
            raise ValueError("no slides given")
        for k, sp in enumerate(slides, start=1):
            imgs = sp.get("images") or []
            if len(imgs) > 6:
                raise ValueError(f"slide {k}: {len(imgs)} images; at most 6 per slide (split into more slides)")
            if not imgs and not sp.get("bullets") and not sp.get("title"):
                raise ValueError(f"slide {k} is empty")
        missing = [f"slide {k}: {p}" for k, sp in enumerate(slides, start=1) for p in (sp.get("images") or [])
                   if not os.path.isfile(tb.p(p))]
        if missing:
            raise FileNotFoundError(f"{len(missing)} image(s) not found: {missing}")
        if append_to:
            src = tb.p(append_to)
            if not os.path.isfile(src):
                raise FileNotFoundError(f"append_to {append_to} not found")
            prs = Presentation(src)
        else:
            prs = pu.new_presentation()
        before = len(prs.slides)
        if not output:
            if not append_to:
                raise ValueError("give output (a .pptx file name)")
            out = tb.p(append_to)
        else:
            out = tb.outp(output if output.lower().endswith(".pptx") else output + ".pptx", "figures.pptx")
        if title_slide:
            pu.add_title_slide(prs, str(title_slide.get("title", "")), str(title_slide.get("subtitle", "") or ""))
        done = []
        for sp in slides:
            spec = dict(sp, images=[tb.p(p) for p in sp.get("images") or []])
            done.append(add_images_slide(prs, spec, _tmpdir()))
        pu.save(prs, out)
        shutil.rmtree(_tmpdir(), ignore_errors=True)
        verb = f"appended {len(prs.slides) - before} slide(s) to {tb.rel(append_to and tb.p(append_to))}" \
            if append_to else f"created deck with {len(prs.slides)} slide(s)"
        lines = [f"{verb}; saved {tb.rel(out)} ({len(prs.slides)} slides total)"]
        lines += [f"  slide {before + (1 if title_slide else 0) + i + 1}: {d['title'] or '(untitled)'} - "
                  f"{d['images']} image(s)" for i, d in enumerate(done)]
        return "\n".join(lines), []

    # --------------------------------------------------------- animate
    @tb.add("animate",
            "Animated GIF (or .mp4 if ffmpeg is installed) from image frames, or from a sweep: call another tool "
            "once per value with sweep.param set (dotted path for nested args) and use its first image as a frame. "
            f"Failed frames are skipped and reported. Max {MAX_FRAMES} frames.",
            {"frames": IMGS, "output": {**S, "description": ".gif or .mp4"},
             "sweep": {"type": "object", "description": "{tool, args, param, values}", "properties": {
                 "tool": S, "args": {"type": "object"}, "param": S, "values": {"type": "array"}}},
             "fps": N, "duration_ms": {**I, "description": "per frame (overrides fps)"},
             "loop": {**I, "description": "GIF repeats, 0 = forever (default)"},
             "caption_template": {**S, "description": "drawn above each frame, e.g. 'T = {value} K' ({value}: sweep "
                                  "value or frame file stem, {i}, {n})"},
             "captions": {**STRS, "description": "explicit caption per frame"},
             "max_width": {**I, "description": "px, default 1200"}},
            ["output"])
    def animate(output: str, frames=None, sweep=None, fps: Optional[float] = None, duration_ms: Optional[int] = None,
                loop: int = 0, caption_template: Optional[str] = None, captions=None, max_width: int = 1200):
        if bool(frames) == bool(sweep):
            raise ValueError("give either frames (image paths) or sweep {tool, args, param, values}")
        ext = Path(output).suffix.lower()
        if ext not in (".gif", ".mp4"):
            output += ".gif"
            ext = ".gif"
        if ext == ".mp4" and not shutil.which("ffmpeg"):
            raise ValueError("ffmpeg is not installed, so .mp4 is unavailable: use a .gif output")
        out = tb.outp(output, "animation.gif")
        notes: List[str] = []
        items: List[Tuple[str, Any]] = []            # (image path, caption value)
        if frames:
            if len(frames) > MAX_FRAMES:
                raise ValueError(f"{len(frames)} frames > {MAX_FRAMES}")
            paths = resolve_images(frames)
            items = [(p, Path(p).stem) for p in paths]
        else:
            tool, param, values = sweep.get("tool"), sweep.get("param"), sweep.get("values")
            if not tool or not param or not values:
                raise ValueError("sweep needs tool, param and values")
            if tool in ("animate", "run_python", "ask_user") or tool not in tb.tools:
                raise ValueError(f"sweep tool {tool!r} is not available" + (
                    "" if tool in tb.tools else f"; tools: {sorted(tb.tools)}"))
            if len(values) > MAX_FRAMES:
                raise ValueError(f"{len(values)} sweep values > {MAX_FRAMES}")
            fdir = Path(out).with_name(Path(out).stem + "_frames")
            fdir.mkdir(parents=True, exist_ok=True)
            import copy
            for k, v in enumerate(values, start=1):
                args = copy.deepcopy(sweep.get("args") or {})
                _set_path(args, param, v)
                progress = getattr(getattr(tb, "ui", None), "progress", None)
                if progress:
                    progress(f"animate: frame {k}/{len(values)} ({param}={v})")
                text, imgs, err = tb.call(tool, args)
                imgs = [i for i in imgs or [] if Path(i).suffix.lower() in RASTER_EXT and os.path.isfile(i)]
                if err or not imgs:
                    why = (text or "no image returned").strip().splitlines()
                    notes.append(f"frame {k} ({param}={v}) skipped: {why[0][:200] if why else 'no image'}")
                    continue
                dst = fdir / f"frame_{k:03d}{Path(imgs[0]).suffix.lower()}"
                shutil.copyfile(imgs[0], dst)       # the tool may overwrite the same file next call
                items.append((str(dst), v))
            if not items:
                raise ValueError("every sweep frame failed:\n" + "\n".join(notes))
        n = len(items)
        if captions:
            caps = [str(c) for c in (list(captions) + [""] * n)[:n]]
        elif caption_template:
            try:
                caps = [caption_template.format(value=v, i=k + 1, n=n) for k, (_, v) in enumerate(items)]
            except (KeyError, IndexError, ValueError) as e:
                raise ValueError(f"caption_template: {e}; placeholders are {{value}}, {{i}}, {{n}}") from None
        else:
            caps = None
        fr = build_frames([open_image(p) for p, _ in items], caps, max_width=int(max_width))
        if duration_ms:
            dur = int(duration_ms)
        else:
            dur = int(round(1000.0 / float(fps or 2)))
        if ext == ".gif":
            write_gif(fr, out, dur, loop)
            how = f"GIF, {dur} ms/frame, loop={'forever' if not loop else loop}"
        else:
            codec = write_mp4(fr, out, 1000.0 / dur)
            how = f"MP4 ({codec}), {1000.0 / dur:g} fps"
        preview = str(Path(out).with_name(Path(out).stem + "_preview.png"))
        fr[0].save(preview)
        lines = [f"saved {tb.rel(out)}: {n} frame(s), {fr[0].size[0]}x{fr[0].size[1]} px, {how}"]
        if sweep:
            lines.append(f"frames kept in {tb.rel(str(Path(items[0][0]).parent))}")
        lines += notes
        return "\n".join(lines), [preview]

    # --------------------------------------------------------- image_tools
    @tb.add("image_tools",
            "Raster edits of any image with Pillow, applied in order: " + IMAGE_OPS_HELP + ".",
            {"image": S, "operations": {"type": "array", "items": {"type": "object"}},
             "output": {**S, "description": "file name (default <name>_edited.<ext>)"}, "dpi": I},
            ["image", "operations"])
    def image_tools(image: str, operations, output: Optional[str] = None, dpi: Optional[int] = None):
        src = resolve_images([image])[0]
        im = open_image(src)
        size0 = im.size
        im, fmt, opts, notes = apply_image_ops(im, operations or [])
        src_ext = Path(src).suffix.lower().lstrip(".")
        src_ext = {"jpeg": "jpg", "tif": "tiff"}.get(src_ext, src_ext)
        if output:
            oext = Path(output).suffix.lower().lstrip(".")
            oext = {"jpeg": "jpg", "tif": "tiff"}.get(oext, oext)
            if fmt is None:
                fmt = oext if oext in ("png", "jpg", "pdf", "tiff", "webp") else (src_ext if src_ext in (
                    "png", "jpg", "tiff", "webp") else "png")
            if oext != fmt:
                output = str(Path(output).with_suffix("." + fmt)) if oext in (
                    "png", "jpg", "pdf", "tiff", "webp") else output + "." + fmt
            out = tb.outp(output, "image.png")
        else:
            fmt = fmt or (src_ext if src_ext in ("png", "jpg", "tiff", "webp") else "png")
            out = tb.outp(None, f"{Path(src).stem}_edited.{fmt}")
        notes += save_raster(im, out, fmt, opts, dpi)
        txt = (f"saved {tb.rel(out)} ({fmt}, {size0[0]}x{size0[1]} -> {im.size[0]}x{im.size[1]} px, "
               f"mode {im.mode})") + ("\n" + "\n".join(notes) if notes else "")
        return txt, ([out] if fmt in ("png", "jpg", "webp") else [])

    # --------------------------------------------------------- images_to_pdf
    @tb.add("images_to_pdf",
            "Multi-page PDF, one image per page (optional page titles). page_size fit (default; tool-made plots "
            "become vector pages), letter or a4 (auto orientation).",
            {"images": IMGS, "output": S, "page_size": {"type": "string", "enum": ["fit", "letter", "a4"]},
             "titles": STRS, "vector": {**B, "description": "use stored editable figures as vector pages (true)"}},
            ["images", "output"])
    def images_to_pdf(images, output: str, page_size: str = "fit", titles=None, vector: bool = True):
        paths = resolve_images(images)
        out = tb.outp(output if output.lower().endswith(".pdf") else output + ".pdf", "images.pdf")
        r = images_pdf(paths, out, page_size=page_size, titles=titles, vector=vector)
        nv = r["kinds"].count("vector")
        return (f"saved {tb.rel(out)}: {r['pages']} page(s), page size {page_size}"
                + (f" ({nv} vector page(s) from editable figures)" if nv else "")), []
