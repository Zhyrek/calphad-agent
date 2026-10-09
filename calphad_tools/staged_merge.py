"""``staged_merge``: merge database B (e.g. Fe-Mo-W) into base A (e.g. Fe-Cr-Ni-Mo) through intermediate databases,
and show exactly what each step does to B's system (and, when B brings new unary/binary terms, to A's).

Versions (B is first normalised to A's phase names / sublattice order / formula units, tdbstage.normalize_pair;
the magnetic basis rule: a magnetic B phase keeps its site ratios and A's phase is rescaled to them, A0 - exactly when
A's is not magnetic, with BMAGN end-members exact and the error between them tolerated and reported when it is):

* B1 = B with A's SGTE unary (SER) Gibbs energies of the shared elements
* B2 = B1 with all of A's unary parameters of the shared elements, in the phases both have
* B3 = B2 + A's binary description of the shared binaries (merge.plan_merge of A restricted to the shared
  elements' unary+binary terms into B2, A's values winning)
* Q  = A + B (merge.plan_merge; conflicts the stages above settle go to A, the rest are asked like merge_tdbs)
* A1 = A + B's unary terms, A2 = A + B's unary and binary terms (only if B has unary/binary terms A lacks)

B3, Q, A1 and A2 get the merge's tentative parameters (calphad_tools.tentative; ``tentative=False`` turns them
off): an assessed parameter always replaces a tentative one, and mixings the union creates that no database
describes get TENT_MIX / TENT_VA terms - listed in the summary and the version logs.

Every version is written to the output folder.  For B, B1, B2, B3, Q all binaries and ternary isotherms of B's
system are mapped (shaded style), but a subsystem that tdbstage.diff_versions proves unchanged from the previous
version is not recomputed - its figure and invariants are reused.  For A, A1, A2 only A's subsystems that change
are mapped.  Consecutive versions are compared (B->B1->B2->B3->Q, A->A1->A2): numerically different parameters,
phase-model changes, invariant reactions that moved (ΔT, Δx).  One deck: a summary slide, then per step one
differences slide and one slide with the changed diagrams of both versions; an Excel workbook has everything.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from . import tdbstage as ts
from .compatibility import clip, fit_rows
from .tdbparse import parse_tdb

KELVIN = 273.15
STEP_DESC = {"B1": "A's SGTE unaries", "B2": "+ A's other unaries", "B3": "+ A's shared binaries",
             "Q": "full merge A+B", "A1": "+ B's unaries", "A2": "+ B's binaries"}


@dataclass
class Version:
    label: str
    path: str
    desc: str
    log: List[str] = field(default_factory=list)
    validation: Dict[str, Any] = field(default_factory=dict)


@dataclass
class Item:
    """One diagram: a binary, or a ternary at one temperature (°C)."""
    elements: Tuple[str, ...]
    T_c: Optional[float] = None

    @property
    def name(self) -> str:
        n = ts.stage_sysname(self.elements)
        return n if self.T_c is None else f"{n} @ {self.T_c:g} °C"

    @property
    def tag(self) -> str:
        n = ts.stage_sysname(self.elements)
        return n if self.T_c is None else f"{n}_{self.T_c:g}C"


def _stem(p: str) -> str:
    return Path(p).stem


# =============================================================================================== versions
def build_versions(a_path: str, b_path: str, out_dir: str, *, phase_map=None, sublattice_map=None,
                   decisions: Optional[Dict[str, str]] = None, policy: str = "ask",
                   extensions: Optional[Dict[str, Dict]] = None, tentative=True,
                   progress: Callable[[str], None] = lambda s: None) -> Dict[str, Any]:
    """``extensions`` {A phase: spec (harmonize.extension_spec)}: 'extend' decisions - A's phase extended with B's
    components (harmonize.staged_extensions; A_eff = the extended A, and B's source phase is left out of the
    merge, B_merge)."""
    from .merge import validate
    from .harmonize import is_extend
    out = Path(out_dir)
    work = out / "_work"
    work.mkdir(parents=True, exist_ok=True)
    # decisions {'phase:X': 'extend' | {'extend': {...}}} are 'extend' decisions (the recommended phase-model answer)
    ext_dec = {k.split(":", 1)[1].upper(): v for k, v in (decisions or {}).items()
               if str(k).startswith("phase:") and is_extend(v)}
    if ext_dec:
        decisions = {k: v for k, v in decisions.items() if not (str(k).startswith("phase:") and is_extend(v))}
        extensions = {**{t: ((v.get("extend") if isinstance(v.get("extend"), dict) else v) if isinstance(v, dict)
                             else ({"phase": v.split(":", 1)[1].strip().upper()} if ":" in str(v) else {}))
                         for t, v in ext_dec.items()}, **(extensions or {})}
    A, B = parse_tdb(a_path), parse_tdb(b_path)
    sa, sb = _stem(a_path), _stem(b_path)
    An, Bn, notes = ts.normalize_pair(A, B, phase_map, sublattice_map, on_error="keep")
    unmapped = [n for n in notes if n.action == "unmapped"]
    if unmapped:
        # B's sublattices cannot be matched to A's unambiguously: a decision, asked like the merge's own
        lines = ["Merge of: db1 (" + Path(a_path).name + "), db2 (" + Path(b_path).name + ")",
                 f"Pending decisions: {len(unmapped)}"]
        for n in unmapped:
            lines += [f"  * sublattice_map:{n.phase}: {n.detail.split(' kept as is: ', 1)[-1]}",
                      f"      - sublattice_map {{'{n.phase}': [...]}}: B's sublattice i -> A's position (as in A's "
                      f"PHASE line; vacancy-only sublattices stay in place)",
                      f"      - sublattice_map {{'{n.phase}': 'keep'}}: leave B's model as is; if the models then "
                      f"differ, decisions {{'phase:{n.phase}': ...}}: 'db1' (A's model), 'db2' (B's) or 'rename' "
                      f"(keep both, B's as {n.phase}_2)"]
        return {"versions": {}, "notes": notes, "unmapped": unmapped, "pending_report": "\n".join(lines),
                "A": a_path, "B": b_path, "B_original": b_path}
    changed = [n for n in notes if n.action in ts.B_CHANGES]
    a_resc = [n for n in notes if n.action in ts.A_CHANGES]
    a_eff = a_path                      # A's description used everywhere below (A0 when the magnetic basis rule applies)
    if a_resc:
        a_eff = ts.write_tdb(An, str(out / f"A0_{sa}_rescaled.tdb"),
                             [f"{Path(a_path).name} with phases rescaled to {Path(b_path).name}'s site ratios "
                              f"(B magnetic: B's basis wins - exact if A is not magnetic, BMAGN end-members exact if "
                              f"it is) or read as exact fractions:"]
                             + [n.text() for n in a_resc])
        An = parse_tdb(a_eff)
    b_eff = b_path
    if changed:
        b_eff = ts.write_tdb(Bn, str(out / f"B0_{sb}_normalised.tdb"),
                             [f"{Path(b_path).name} normalised to {Path(a_path).name}'s conventions:"]
                             + [n.text() for n in changed])
    Bn = parse_tdb(b_eff)
    shared = sorted(set(A.real_elements) & set(Bn.real_elements))
    if not shared:
        raise ValueError(f"{Path(a_path).name} and {Path(b_path).name} share no element: nothing to stage")
    VA, VB = ts.StageView(An, "A"), ts.StageView(Bn, "B")
    ext, b_merge = None, b_eff
    if extensions:
        from .harmonize import extension_spec, jsonable_report, staged_extensions
        progress("staged_merge: phase harmonisation ('extend' decisions)")
        ren = {n.phase: n.detail.split()[1] for n in notes if n.action == "renamed"}
        specs = {}
        for tgt, v in extensions.items():
            if isinstance(v, dict) and "target" in v and "phase" in v and "normalization" in v:
                specs[tgt.upper()] = v
            else:
                sp = extension_spec(tgt, v, An)
                sp["phase"] = ren.get(sp["phase"], sp["phase"])
                specs[tgt.upper()] = extension_spec(tgt, {**{k: x for k, x in sp.items() if k in (
                    "sublattice_map", "add_constituents", "components", "normalization", "collision", "magnetic",
                    "include_functions") and x is not None}, "phase": sp["phase"]}, An, Bn)
        ext = staged_extensions(a_eff, b_eff, specs, str(out), shared, progress=progress, tentative=tentative)
        a_eff, b_merge = ext["A_ext"], ext["B_merge"]
        ext["reports"] = [jsonable_report(x) for x in ext["reports"]]
    # decisions for Q: name aliases stay separate (names only via normalisation / phase_map), what the stages
    # settle goes to A, the rest is the user's
    user = dict(decisions or {})
    dec = ts.alias_separate([a_eff, b_merge], user)
    plan0 = ts.plan_merge([a_eff, b_merge], dec, "ask")
    auto = ts.auto_decisions(plan0, shared)
    dec = {**auto, **dec, **user}
    versions: Dict[str, Version] = {"B": Version("B", b_eff, f"{Path(b_path).name}" + (" (normalised)" if changed
                                                                                       else ""))}
    # B1, B2: text-level substitutions
    progress("staged_merge: B1 (A's SGTE unaries)")
    B1, log1 = ts.substitute_ser(Bn, VA, VB, shared)
    p1 = ts.write_tdb(B1, str(out / f"B1_{sb}+{sa}-SER.tdb"),
                      [f"B1 = {Path(b_eff).name} with the SER (SGTE unary) Gibbs energies of "
                       f"{', '.join(shared)} from {Path(a_path).name}"] + log1)
    versions["B1"] = Version("B1", p1, STEP_DESC["B1"], log1)
    VB1 = ts.StageView(parse_tdb(p1), "B1")
    progress("staged_merge: B2 (A's unaries)")
    B2, log2 = ts.substitute_unaries(parse_tdb(p1), VA, VB1, shared)
    p2 = ts.write_tdb(B2, str(out / f"B2_{sb}+{sa}-unaries.tdb"),
                      [f"B2 = B1 with every unary parameter of {', '.join(shared)} from {Path(a_path).name} "
                       f"(phases both have)"] + log2)
    versions["B2"] = Version("B2", p2, STEP_DESC["B2"], log2)
    plan = ts.plan_merge([a_eff, b_merge], dec, policy, tentative=tentative)
    result: Dict[str, Any] = {"versions": versions, "notes": notes, "shared": shared, "decisions": dec,
                              "auto_decisions": auto, "A": a_path, "A_eff": a_eff, "B": b_eff,
                              "B_original": b_path, "B_merge": b_merge,
                              # 'extend' decisions: reports and {B phase: A phase} of the source phases
                              "extensions": ext, "extension_renames": (ext or {}).get("renames", {}),
                              # B's phases a 'rename' decision kept under a new name in Q: {B name: Q name}
                              "phase_renames": plan.renames_of(1),
                              # tentative parameters of Q (calphad_tools.tentative)
                              "tentative_cfg": tentative, "tentative": tentative_info(plan)}
    if plan.pending:
        result["pending"] = plan
        return result
    # B3: A's unary+binary description of the shared elements merged into B2 (A wins)
    progress("staged_merge: B3 (A's shared binaries)")
    a_src = parse_tdb(a_eff) if ext else An           # the extended A when an 'extend' decision applies
    a_sub = ts.write_tdb(ts.restrict_tdb(a_src, shared, max_elements=2), str(work / f"{sa}_shared_unary_binary.tdb"))
    d3 = ts.alias_separate([a_sub, p2], dec)
    plan3, text3 = ts.merged_text([a_sub, p2], d3, "first:predict", tentative=tentative)
    p3 = str(out / f"B3_{sb}+{sa}-unaries-binaries.tdb")
    Path(p3).write_text(text3)
    versions["B3"] = Version("B3", p3, STEP_DESC["B3"], list(plan3.log))
    progress("staged_merge: Q (full merge)")
    from .merge import render_tdb
    pq = str(out / f"Q_{sa}+{sb}.tdb")
    Path(pq).write_text(render_tdb(plan))
    versions["Q"] = Version("Q", pq, STEP_DESC["Q"], list(plan.log) + [f"decision {c.id} -> {c.decision} "
                                                                       f"({'auto: staged' if c.id in auto or c.group in auto else c.decided_by})"
                                                                       for c in plan.conflicts if c.decision])
    # A1, A2 when B brings unary / binary terms A lacks
    EB = set(Bn.real_elements)
    new_u = [k for k in VB.rows if len(VB.row_elements(k)) <= 1 and k not in VA.rows]
    new_b = [k for k in VB.rows if len(VB.row_elements(k)) == 2 and k not in VA.rows]
    result["new_terms"] = {"unary": new_u, "binary": new_b}
    if new_u or new_b:
        for lab, n_el, desc in (("A1", 1, "unary"), ("A2", 2, "unary and binary")):
            progress(f"staged_merge: {lab} (A + B's {desc} terms)")
            bpart = ts.write_tdb(ts.restrict_tdb(Bn, sorted(EB), max_elements=n_el, constituents_from_params=True),
                                 str(work / f"{sb}_{desc.replace(' and ', '_')}.tdb"))
            d = ts.alias_separate([a_eff, bpart], dec)
            pl, txt = ts.merged_text([a_eff, bpart], d, "first:predict", tentative=tentative)
            p = str(out / f"{lab}_{sa}+{sb}-{desc.replace(' and ', '-').replace('unary', 'unaries').replace('binary', 'binaries')}.tdb")
            Path(p).write_text(txt)
            versions[lab] = Version(lab, p, STEP_DESC[lab], list(pl.log))
    for v in versions.values():
        if v.label != "B":
            v.validation = validate(v.path)
    return result


def extend_possible(A, B, a_phase: str, b_phase: str) -> Optional[List[str]]:
    """The elements an 'extend' decision would bring into A's ``a_phase`` from B's ``b_phase`` (None when extend is not
    possible: no extra element, a different number of sublattices, non-proportional site ratios, order-disorder)."""
    from .harmonize import _real, phase_comps
    if a_phase not in A.phases or b_phase not in B.phases:
        return None
    if ts._order_disorder_related(A, a_phase) or ts._order_disorder_related(B, b_phase):
        return None
    pa, pb = A.phases[a_phase], B.phases[b_phase]
    extra = sorted(set(phase_comps(B, b_phase)) - set(phase_comps(A, a_phase)) - {"VA"})
    if not extra or len(_real(pa)) != len(_real(pb)):
        return None
    try:
        if ts.sublattice_mapping(pa, pb) is None:
            return None
    except ts.SublatticeMapError:
        pass                                # possible with a sublattice_map
    return extra


def recommend_extend(plan, a_path: str, b_path: str) -> None:
    """Pending phase-model conflicts of ``plan`` where 'extend' is possible get it as their FIRST option, marked
    recommended (nothing of either side is dropped; decisions {'phase:X': 'extend'} or extend={'X': {...}})."""
    from .merge import Option
    pend = [c for c in plan.pending if c.kind == "phase_model" and not any(o.key == "extend" for o in c.options)]
    if not pend:
        return
    A, B = parse_tdb(str(a_path)), parse_tdb(str(b_path))
    for c in pend:
        ph = c.id.split(":", 1)[1]
        extra = extend_possible(A, B, ph, ph)
        if extra:
            c.options.insert(0, Option("extend", f"RECOMMENDED - keep A's {ph} model and its parameters, add "
                                                 f"{', '.join(extra)} with only B's parameters that involve "
                                                 f"{'it' if len(extra) == 1 else 'them'} (nothing dropped)",
                                       f"decisions {{'phase:{ph}': 'extend'}} or extend={{'{ph}': {{}}}}"))


def tentative_info(plan) -> Dict[str, Any]:
    """The tentative parameters of a merge plan as plain data (terms added, replacements, notes)."""
    from .tentative import settings_text
    return {"settings": settings_text(plan.tentative_settings),
            "added": [{"label": t.label, "ref": t.ref, "rule": t.rule, "value": t.value, "phase": t.phase,
                       "text": t.text(), "evidence": t.evidence} for t in plan.tentative],
            "replaced": [{"label": r.label, "ref": r.ref, "from": r.tentative_from, "by": r.by, "kind": r.kind,
                          "text": r.text()} for r in plan.tentative_replaced],
            "notes": list(plan.tentative_notes)}


def tentative_lines(info: Optional[Dict[str, Any]], max_items: int = 6) -> List[str]:
    if not info or not (info.get("added") or info.get("replaced") or info.get("notes")):
        return []
    L = []
    for ref, title in (("TENT_MIX", "rule 1, new mixing"), ("TENT_VA", "rule 2, vacancies")):
        xs = [a for a in info["added"] if a["ref"] == ref]
        if xs:
            L.append(f"Tentative {ref} ({title}; not assessed): " + "; ".join(
                f"{a['label']} = {a['value']:+g}" for a in xs[:max_items])
                + (f"; +{len(xs) - max_items} more" if len(xs) > max_items else ""))
    rep = info.get("replaced") or []
    if rep:
        L.append("Tentative replaced: " + "; ".join(r["text"].replace("rule:tentative-replaced: ", "")
                                                   for r in rep[:max_items])
                 + (f"; +{len(rep) - max_items} more" if len(rep) > max_items else ""))
    L += [f"Tentative note: {n}" for n in (info.get("notes") or [])[:2]]
    return L


# =============================================================================================== plots
def _content_key(tdb: str, item: Item, settings: Dict) -> str:
    from .tdbparse import calc_text
    text = calc_text(Path(tdb).read_text(errors="replace"))
    h = hashlib.sha256(text.encode())
    h.update(json.dumps([list(item.elements), item.T_c, settings], sort_keys=True).encode())
    h.update(b"mapping-check-1")        # diagrams made before the mapping check are not reused
    return h.hexdigest()


def _task(label: str, tdb: str, item: Item, png: str, s: Dict) -> Dict:
    els = list(item.elements)
    title = f"{label}: " + "–".join(e.capitalize() for e in els)
    if item.T_c is None:
        kw = dict(tdb=tdb, elements=els, right=els[1], out_png=png, tmin=s["tmin_c"], tmax=s["tmax_c"],
                  temperature_unit="C", composition_step=s["x_step"], temperature_step=s["t_step"], tielines=0,
                  style="shaded", title=title, formats=[], time_limit_s=0.9 * s["timeout"])
    else:
        kw = dict(tdb=tdb, elements=els, corners=dict(zip(("left", "right", "top"), els)), temperature=item.T_c,
                  temperature_unit="C", out_png=png, composition_step=s["tern_step"], tielines=0, style="shaded",
                  formats=[], time_limit_s=0.9 * s["timeout"], title=f"{title}  {item.T_c:g} °C")
    return {"name": f"{label} {item.name}", "kind": "map", "kwargs": kw}


FORCE_TOL_T, FORCE_TOL_X = 0.5, 1e-3        # a forced replot of an "exact" change must not move invariants more


def force_matcher(force) -> Callable[[Sequence[str]], bool]:
    """force_replot -> predicate on a subsystem's elements: True = all, False/None = none, or a list of names
    ('Fe-Mo', 'mo-fe', 'Fe-Mo-W'; any order / case)."""
    if force is True or (isinstance(force, str) and force.lower() in ("true", "all")):
        return lambda els: True
    if not force:
        return lambda els: False
    names = [force] if isinstance(force, str) else list(force)
    want = {tuple(sorted(e.strip().upper() for e in re.split(r"[-–_ ,]+", str(n)) if e.strip())) for n in names}
    return lambda els: tuple(sorted(e.upper() for e in els)) in want


def plan_plots(chain: List[str], versions: Dict[str, Version], diffs: Dict[Tuple[str, str], ts.VersionDiff],
               items: List[Item], forced=lambda els: False) -> Tuple[Dict[Tuple[str, str], Optional[str]], set]:
    """((version, item tag) -> None (compute) or the version whose result is reused (subsystem provably
    unchanged), the (version, tag) pairs recomputed only because force_replot asked for them)."""
    need: Dict[Tuple[str, str], Optional[str]] = {}
    forced_keys = set()
    for i, lab in enumerate(chain):
        for it in items:
            if i == 0:
                need[(lab, it.tag)] = None
                continue
            prev = chain[i - 1]
            sd = diffs[(prev, lab)].get(it.elements)
            if sd is None or sd.changed:
                need[(lab, it.tag)] = None
            elif forced(it.elements):
                need[(lab, it.tag)] = None
                forced_keys.add((lab, it.tag))
            else:
                need[(lab, it.tag)] = need[(prev, it.tag)] or prev
    return need, forced_keys


def run_plots(versions, need, items, fig_dir: Path, settings: Dict, n_workers=None,
              progress=lambda s: None, nocache=()) -> Dict[Tuple[str, str], Dict]:
    from .diagrams import run_tasks
    fig_dir.mkdir(parents=True, exist_ok=True)
    by_tag = {it.tag: it for it in items}
    tasks, results = [], {}
    for (lab, tag), src in need.items():
        if src is not None:
            continue
        it = by_tag[tag]
        png = str(fig_dir / f"{lab}_{tag}.png")
        key = _content_key(versions[lab].path, it, settings)
        kfile = Path(png + ".key")
        jfile = Path(png).with_suffix(".json")
        if ((lab, tag) not in nocache and Path(png).exists() and jfile.exists() and kfile.exists()
                and kfile.read_text() == key):
            results[(lab, tag)] = {"ok": True, "png": png, "cached": True}
            continue
        for f in (kfile, jfile):
            if f.exists():
                f.unlink()
        t = _task(lab, versions[lab].path, it, png, settings)
        t["_key"] = (lab, tag, key)
        tasks.append(t)
    if tasks:
        progress(f"staged_merge: mapping {len(tasks)} diagram(s)")
        res = run_tasks(tasks, n_workers=n_workers, timeout=settings["timeout"], progress=progress)
        for t, r in zip(tasks, res):
            lab, tag, key = t["_key"]
            ok = r.ok and r.png is not None and Path(r.png).exists()
            results[(lab, tag)] = {"ok": ok, "png": r.png if ok else None, "error": r.error, "seconds": r.seconds}
            if ok:
                Path(r.png + ".key").write_text(key)
    for (lab, tag), src in need.items():
        if src is not None:
            results[(lab, tag)] = {**results.get((src, tag), {"ok": False, "error": "source missing"}),
                                   "reused_from": src}
    from .diagrams import _check_note
    for k, r in results.items():
        r["invariants"] = []
        if r.get("ok") and r.get("png"):
            j = Path(r["png"]).with_suffix(".json")
            if j.exists():
                try:
                    data = json.loads(j.read_text())
                    r["invariants"] = data.get("invariants", [])
                    r["phases_found"] = [str(p) for p in data.get("phases_found") or []]
                    r["mapping_check"] = _check_note(data.get("mapping_check"))
                except (ValueError, OSError):
                    pass
    return results


# =============================================================================================== invariants
def _base(p: str) -> str:
    return re.sub(r"#\d+$", "", p)


def _vertices(inv: Dict, els: Sequence[str]) -> List[Tuple[str, Tuple[float, ...]]]:
    comp = inv.get("compositions", {})
    return sorted((_base(p), tuple(float(comp.get(p, {}).get(e, math.nan)) for e in els)) for p in inv["phases"])


def rename_invariants(invs: List[Dict], mapping: Dict[str, str]) -> List[Dict]:
    """Invariants with phases renamed (B's MU_D85 -> Q's MU_D85_2; '#n' suffixes kept)."""
    if not mapping:
        return invs
    out = []
    for inv in invs:
        def ren(p):
            b = _base(p)
            return mapping.get(b, b) + p[len(b):]
        inv = dict(inv)
        inv["phases"] = [ren(p) for p in inv.get("phases", [])]
        if isinstance(inv.get("compositions"), dict):
            inv["compositions"] = {ren(p): v for p, v in inv["compositions"].items()}
        out.append(inv)
    return out


def compare_invariants(old: List[Dict], new: List[Dict], els: Sequence[str], ternary: bool,
                       tol_T: float = 5.0, tol_x: float = 0.01) -> List[Dict[str, Any]]:
    """Match invariants (binary) / tie-triangles (ternary isotherm) by phase set; ΔT, Δx (max over vertices)."""
    def key(inv):
        return " + ".join(sorted(_base(p) for p in inv["phases"]))

    def dx(a, b):
        va, vb = _vertices(a, els), _vertices(b, els)
        return max((abs(x - y) for (_, xa), (_, xb) in zip(va, vb) for x, y in zip(xa, xb)), default=math.nan)
    rows = []
    groups = sorted({key(i) for i in old} | {key(i) for i in new})
    for g in groups:
        o = [i for i in old if key(i) == g]
        n = [i for i in new if key(i) == g]
        pairs = []
        while o and n:
            best = min(((a, b) for a in o for b in n),
                       key=lambda ab: dx(*ab) if ternary else abs(ab[0]["T_K"] - ab[1]["T_K"]))
            pairs.append(best)
            o.remove(best[0])
            n.remove(best[1])
        for a, b in pairs:
            dT = 0.0 if ternary else b["T_K"] - a["T_K"]
            d = dx(a, b)
            moved = (abs(dT) > tol_T) or (math.isfinite(d) and d > tol_x)
            rows.append({"invariant": g, "type": b.get("type", a.get("type", "")), "T_old_K": a["T_K"],
                         "T_new_K": b["T_K"], "dT_K": dT, "dx": d, "status": "moved" if moved else "same"})
        for a in o:
            rows.append({"invariant": g, "type": a.get("type", ""), "T_old_K": a["T_K"], "T_new_K": math.nan,
                         "dT_K": math.nan, "dx": math.nan, "status": "disappeared"})
        for b in n:
            rows.append({"invariant": g, "type": b.get("type", ""), "T_old_K": math.nan, "T_new_K": b["T_K"],
                         "dT_K": math.nan, "dx": math.nan, "status": "appeared"})
    return rows


# =============================================================================================== steps
@dataclass
class Step:
    old: str
    new: str
    diff: ts.VersionDiff
    items: List[Item]                                    # items of this chain
    inv: Dict[str, List[Dict]] = field(default_factory=dict)    # item tag -> invariant comparison rows
    changed_items: List[Item] = field(default_factory=list)
    forced_items: List[Item] = field(default_factory=list)      # exact-only subsystems replotted on request
    forced_inv: Dict[str, List[Dict]] = field(default_factory=dict)
    forced_kind: Dict[str, str] = field(default_factory=dict)
    zero_g: List[str] = field(default_factory=list)            # end-members with no G (pycalphad 0): stable or not

    def forced_shift(self, name: str) -> Tuple[float, float, bool]:
        """(max |ΔT|, max Δx, consistent?) of a forced replot: consistent = nothing appeared/disappeared and the
        shifts are within FORCE_TOL_T / FORCE_TOL_X."""
        rows = self.forced_inv.get(name)
        if rows is None:
            return math.nan, math.nan, False
        dT = max([abs(r["dT_K"]) for r in rows if math.isfinite(r["dT_K"])] or [0.0])
        dx = max([r["dx"] for r in rows if math.isfinite(r["dx"])] or [0.0])
        ok = all(r["status"] in ("moved", "same") for r in rows) and dT <= FORCE_TOL_T and dx <= FORCE_TOL_X
        return dT, dx, ok

    def forced_phrase(self, name: str) -> str:
        return "replotted as a check" if " kept as " in self.forced_kind.get(name, "") else "replotted on request"

    def forced_label(self, it: "Item") -> str:
        dT, dx, ok = self.forced_shift(it.name)
        if not math.isfinite(dT):
            return f"{it.name}: {self.forced_kind.get(it.name, 'exact')}, {self.forced_phrase(it.name)} - not compared"
        s = (f"{it.name}: {self.forced_kind.get(it.name, 'exact')}, {self.forced_phrase(it.name)} - max invariant "
             f"ΔT {dT:.2f} K, Δx {dx:.4f}")
        if not ok:
            s += " - SHIFTED: the exactness claim FAILED"
        return s

    @property
    def title(self) -> str:
        return f"{self.old} → {self.new}"

    def moved(self) -> List[Tuple[str, Dict]]:
        return [(t, r) for t, rs in self.inv.items() for r in rs if r["status"] != "same"]

    def n_compared(self) -> int:
        return sum(len(rs) for rs in self.inv.values())

    def biggest(self) -> Optional[Tuple[str, Dict]]:
        cand = [(t, r) for t, r in self.moved() if r["status"] == "moved"]
        if not cand:
            return None
        return max(cand, key=lambda tr: (abs(tr[1]["dT_K"]) if math.isfinite(tr[1]["dT_K"]) else 0.0,
                                         tr[1]["dx"] if math.isfinite(tr[1]["dx"]) else 0.0))

    def what_changed(self) -> str:
        pc = self.diff.params
        mods = model_groups(self.diff)
        parts = []
        if pc:
            top = sorted(pc, key=ts._pc_rank)[:2]
            parts.append(f"{len(pc)} parameter(s): {', '.join(p.text() for p in top)}"
                         + (f", +{len(pc) - 2} more" if len(pc) > 2 else ""))
        if mods:
            items = [f"{t} in {', '.join(subs[:2])}" + (f" +{len(subs) - 2}" if len(subs) > 2 else "")
                     for (_, _), (t, subs) in list(mods.items())[:2]]
            parts.append(f"{len(mods)} phase-model change(s): {'; '.join(items)}"
                         + (f"; +{len(mods) - 2} more" if len(mods) > 2 else ""))
        exact = [n.split(":", 1)[0] for n in self.diff.notes]
        if exact:
            parts.append(f"{', '.join(exact)}: site ratios rescaled with G/L alike, no magnetic terms - exact, "
                         f"not a change")
        return "; ".join(parts) or "nothing (numerically identical)"


def model_groups(diff: ts.VersionDiff) -> Dict[Tuple[str, str], Tuple[str, List[str]]]:
    """(phase, kind) -> (text in the largest subsystem, subsystems where it shows)."""
    out: Dict[Tuple[str, str], Tuple[str, List[str]]] = {}
    for s in sorted(diff.changed, key=lambda s: (-len(s.elements), s.name)):
        for m in s.models:
            k = (m.phase, m.kind)
            if k not in out:
                out[k] = (m.text(), [])
            out[k][1].append(s.name)
    return out


def _fmt_shift(t: str, r: Dict) -> str:
    s = f"{t}: {r['invariant']}"
    if r["status"] != "moved":
        return f"{s} {r['status']}"
    if math.isfinite(r["dT_K"]) and r["dT_K"] != 0.0:
        s += f" ΔT={r['dT_K']:+.1f} K"
    if math.isfinite(r["dx"]):
        s += f" Δx={r['dx']:.3f}"
    return s


def step_line(st: Step, tol_T: float) -> str:
    ch = [it.name for it in st.changed_items]
    s = f"{st.title} ({STEP_DESC.get(st.new, '')}): {st.what_changed()}; "
    s += f"diagrams changed: {', '.join(ch) if ch else 'none (identical' + ('' if st.forced_items else ', not recomputed') + ')'}"
    if ch:
        mv = [x for x in st.moved()]
        s += f"; invariants moved >{tol_T:g} K / appeared / gone: {len(mv)} of {max(st.n_compared(), len(mv))}"
        b = st.biggest()
        if b:
            s += f", largest {_fmt_shift(*b)}"
        elif mv:
            s += f" ({_fmt_shift(*mv[0])})"
    if st.zero_g:
        s += "\n    no G (pycalphad uses 0): " + "; ".join(st.zero_g)
    if st.forced_items:
        parts = []
        for it in st.forced_items:
            dT, dx, ok = st.forced_shift(it.name)
            parts.append(f"{it.name} ({st.forced_kind.get(it.name, 'exact')}): " + (
                "not compared" if not math.isfinite(dT) else f"max ΔT {dT:.2f} K, Δx {dx:.4f}"
                + ("" if ok else " - SHIFTED, exactness claim FAILED")))
        s += "\n    forced replots: " + "; ".join(parts)
    return s


# =============================================================================================== deck / excel
def _side_by_side(step: Step, plots, out_png: str, max_cols: int = 4) -> Optional[str]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    # changed diagrams first, then the exact-only ones replotted on request (failed exactness checks first)
    forced = sorted(step.forced_items, key=lambda it: step.forced_shift(it.name)[2])
    items = (step.changed_items + forced)[:max_cols + (2 if forced else 0)]
    if not items:
        return None
    n = len(items)
    labs = (step.old, step.new)
    if n <= 2:      # one row per diagram, the two versions side by side
        fig, axes = plt.subplots(n, 2, figsize=(11.0, 4.6 * n), squeeze=False)
        cell = lambda j, i: axes[j][i]          # noqa: E731
    else:           # one row per version
        fig, axes = plt.subplots(2, n, figsize=(4.2 * n, 7.4), squeeze=False)
        cell = lambda j, i: axes[i][j]          # noqa: E731
    for j, it in enumerate(items):
        for i, lab in enumerate(labs):
            ax = cell(j, i)
            ax.axis("off")
            r = plots.get((lab, it.tag), {})
            if r.get("ok") and r.get("png") and Path(r["png"]).exists():
                ax.imshow(plt.imread(r["png"]))
                if r.get("reused_from"):
                    ax.text(0.5, -0.02, f"{lab} = {r['reused_from']} for {it.name} (unchanged, not recomputed)",
                            transform=ax.transAxes, ha="center", va="top", fontsize=8, color="#555555")
                if i == 1 and it in step.forced_items:
                    ok = step.forced_shift(it.name)[2]
                    ax.text(0.5, -0.02 if n <= 2 else -0.04, step.forced_label(it), transform=ax.transAxes,
                            ha="center", va="top", fontsize=10 if n <= 2 else 8, wrap=True,
                            color="#1b5e20" if ok else "#b71c1c", fontweight="normal" if ok else "bold")
            else:
                ax.text(0.5, 0.5, f"{lab} {it.name}\nnot available\n{clip(str(r.get('error', '')), 80)}", ha="center",
                        va="center", fontsize=9, color="#9b1c1c", transform=ax.transAxes)
    for i, lab in enumerate(labs):
        if n <= 2:
            cell(0, i).set_title(lab, fontsize=16, fontweight="bold")
        else:
            ax0 = cell(0, i)
            ax0.text(-0.04, 0.5, lab, transform=ax0.transAxes, fontsize=16, fontweight="bold", rotation=90,
                     va="center", ha="right")
    fig.tight_layout()
    fig.savefig(out_png, dpi=170, bbox_inches="tight")
    plt.close(fig)
    return out_png


def _diff_rows(step: Step) -> List[List[Any]]:
    from .pptx_utils import GridCell as C
    rows = []
    for note in step.diff.notes:
        ph, _, rest = note.partition(": ")
        rows.append([C("exact (no change)"), C(ph), C(rest), C(", ".join(
            x.name for x in step.diff.subsystems if note in x.notes))])
    for it in sorted(step.forced_items, key=lambda it: step.forced_shift(it.name)[2]):
        ok = step.forced_shift(it.name)[2]
        rows.append([C("forced replot"), C(it.name), C(step.forced_label(it).split(": ", 1)[1], bold=not ok),
                     C("exactness check", fill="same" if ok else "differs")])
    for (phase, _), (text, subs) in model_groups(step.diff).items():
        rows.append([C("phase model"), C(phase), C(text), C(", ".join(subs))])
    for t, r in step.moved():
        rows.append([C("invariant"), C(t), C(_fmt_shift(t, r).split(": ", 1)[1]),
                     C(r.get("type", "") or "", fill="differs" if r["status"] != "moved" else None)])
    changed_subs = step.diff.changed
    for p in sorted(step.diff.params, key=ts._pc_rank):
        subs = [x.name for x in changed_subs if any(q.key == p.key for q in x.params)]
        if not subs:
            continue
        det = (f"Δmax {ts.fmt_delta(p.max_diff, p.ptype)} at {p.T_at:.0f} K" if p.kind == "value"
               else p.kind + (f": {clip(p.new, 120)}" if p.kind == "added" and p.new else ""))
        rows.append([C("parameter"), C(p.label), C(det), C(", ".join(subs))])
    return rows


def write_deck(steps: List[Step], plots, path: str, a_name: str, b_name: str, tol_T: float, fig_dir: Path,
               summary_note: str = "") -> Tuple[str, List[str]]:
    from . import pptx_utils as pu
    prs = pu.new_presentation()
    imgs = add_staged_slides(prs, steps, plots, a_name, b_name, tol_T, fig_dir, summary_note)
    pu.save(prs, path)
    return path, imgs


def add_staged_slides(prs, steps: List[Step], plots, a_name: str, b_name: str, tol_T: float, fig_dir: Path,
                      summary_note: str = "", prefix: str = "") -> List[str]:
    """The staged-merge slides (summary, then per step a differences table and the changed diagrams) added to
    ``prs``; ``prefix`` starts every title (combine_databases: "Step 2 · ").  Returns the comparison images."""
    from . import pptx_utils as pu
    from pptx.util import Inches
    from .pptx_utils import GridCell as C
    rows = []
    for st in steps:
        b = st.biggest()
        mv = st.moved()
        shift = _fmt_shift(*b) if b else (_fmt_shift(*mv[0]) if mv else "—")
        failed = [it for it in st.forced_items if not st.forced_shift(it.name)[2]]
        diag = ", ".join(it.name for it in st.changed_items) or "none"
        if st.forced_items:
            for phrase in sorted({st.forced_phrase(it.name) for it in st.forced_items}):
                diag += f"\n{phrase}: " + ", ".join(
                    f"{it.name} (ΔT {st.forced_shift(it.name)[0]:.2f} K)" for it in st.forced_items
                    if st.forced_phrase(it.name) == phrase)
        if failed:
            shift = "FORCED REPLOT SHIFTED - exactness FAILED: " + ", ".join(it.name for it in failed) + \
                ("" if shift == "—" else f"; {shift}")
        wc = st.what_changed() + ("\nno G (pycalphad uses 0): " + "; ".join(st.zero_g) if st.zero_g else "")
        rows.append([C(st.title, bold=True), C(STEP_DESC.get(st.new, "")), C(wc),
                     C(diag, fill="differs" if failed else "error" if st.changed_items else "same"),
                     C(f"{len(mv)} of {max(st.n_compared(), len(mv))}" if st.changed_items else "—", align="c"),
                     C(shift, fill="differs" if failed else None, bold=bool(failed))])
    W = [1.15, 1.6, 4.35, 2.2, 1.6, 2.3]
    pu.add_flow_table_slides(prs, f"{prefix}Staged merge: {a_name} + {b_name}",
                             ["Step", "What", "Changed parameters / models", "Diagrams changed",
                              f"Invariants ΔT>{tol_T:g} K", "Largest shift"], rows, W, font_size=12,
                             subtitle=(summary_note or "A = base, B = added; each version compared with its "
                                                       "predecessor only"), max_height_in=5.5)
    imgs = []
    WD = [1.5, 3.1, 5.7, 2.5]
    for st in steps:
        drows = _diff_rows(st) or [[C("no differences: every parameter numerically equal, same phase models",
                                      span=4, muted=True)]]
        drows = fit_rows(drows, WD, 11, lambda k, st=st: [C(f"+{k} more rows: Excel sheet '{st.old}-{st.new}'",
                                                            span=4, muted=True)])
        pu.add_flow_table_slides(prs, f"{prefix}{st.title}: differences ({STEP_DESC.get(st.new, '')})",
                                 ["Kind", "Item", "Change", "Subsystems / type"], drows, WD, font_size=11,
                                 subtitle="numerically different parameters, phase-model changes and moved "
                                          "invariants only")
        png = _side_by_side(st, plots, str(fig_dir / f"compare_{st.old}-{st.new}.png"))
        if png:
            imgs.append(png)
            n_all = len(st.changed_items) + len(st.forced_items)
            n_shown = min(n_all, 4 + (2 if st.forced_items else 0))
            pu.add_image_slide(prs, f"{prefix}{st.title}: diagrams that changed"
                               + ("" if not st.forced_items else
                                  " + renamed phase replotted as a check" if all(
                                      st.forced_phrase(it.name) == "replotted as a check" for it in st.forced_items)
                                  else " + exact changes replotted on request"), png,
                               subtitle=(f"left {st.old}, right {st.new}" if n_shown <= 2 else
                                         f"top {st.old}, bottom {st.new}")
                               + (f"  (first {n_shown} of {n_all})" if n_all > n_shown else ""))
        else:
            s = pu._blank(prs)
            pu.add_title(s, f"{prefix}{st.title}: diagrams", None)
            pu._textbox(s, Inches(0.8), Inches(2.5), Inches(11.5), Inches(1.5),
                        "No binary or ternary changes: every parameter is numerically equal and every phase model "
                        "identical, so the diagrams are identical (not recomputed).", size=18)
    return imgs


def write_xlsx(steps: List[Step], plots, versions: Dict[str, Version], path: str, notes) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Versions"
    ws.append(["Version", "File", "What", "pycalphad loads", "Errors", "Build log (first 40)"])
    for v in versions.values():
        ws.append([v.label, v.path, v.desc, v.validation.get("loaded", "") if v.validation else "(input)",
                   "; ".join(v.validation.get("errors", [])[:5]) if v.validation else "", "\n".join(v.log[:40])])
    if notes:
        ws.append([])
        ws.append(["Normalisation of B"])
        for n in notes:
            ws.append(["", n.phase, n.action, n.detail])

    def put(w, row):
        w.append([round(x, 6) if isinstance(x, float) and math.isfinite(x) else ("" if isinstance(x, float) else x)
                  for x in row])
    for st in steps:
        w = wb.create_sheet(f"{st.old}-{st.new}"[:31])
        w.append(["Parameter", "Kind", "max |Δ|", "T at max (K)", "Subsystems", f"{st.old}", f"{st.new}"])
        for p in sorted(st.diff.params, key=ts._pc_rank):
            subs = [s.name for s in st.diff.subsystems if any(q.key == p.key for q in s.params)]
            put(w, [p.label, p.kind, p.max_diff, p.T_at, ", ".join(subs), p.old, p.new])
        w.append([])
        w.append(["Subsystem", "Changed", "Reasons"])
        for s in st.diff.subsystems:
            w.append([s.name, "yes" if s.changed else "no", s.reasons(6)])
        w.append([])
        w.append(["Diagram", "Invariant", "Type", "T old (K)", "T new (K)", "ΔT (K)", "Δx", "Status"])
        for t, rs in st.inv.items():
            for r in rs:
                put(w, [t, r["invariant"], r["type"], r["T_old_K"], r["T_new_K"], r["dT_K"], r["dx"], r["status"]])
        for t, rs in st.forced_inv.items():
            for r in rs:
                put(w, [f"{t} (forced replot, {st.forced_kind.get(t, 'exact')})", r["invariant"], r["type"],
                        r["T_old_K"], r["T_new_K"], r["dT_K"], r["dx"], r["status"]])
        for row in w.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for col, wd in zip("ABCDEFG", (34, 10, 14, 12, 22, 50, 50)):
            w.column_dimensions[col].width = wd
    w = wb.create_sheet("Invariants")
    w.append(["Version", "Diagram", "Reused from", "T (K)", "T (°C)", "Phases", "Type", "Compositions"])
    for (lab, tag), r in sorted(plots.items()):
        for inv in r.get("invariants", []):
            put(w, [lab, tag, r.get("reused_from", ""), inv.get("T_K"), inv.get("T_C"), " + ".join(inv["phases"]),
                    inv.get("type", ""), json.dumps(inv.get("compositions", {}))])
    for wsx in wb.worksheets:
        for c in wsx[1]:
            c.font = Font(bold=True)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# =============================================================================================== driver
def staged_merge(a_path: str, b_path: str, out_dir: str, *, phase_map=None, sublattice_map=None, decisions=None,
                 policy: str = "ask", extensions: Optional[Dict[str, Dict]] = None, tentative=True,
                 ternary_temperatures_c: Sequence[float] = (1000.0,), tmin_c: float = 300.0,
                 tmax_c: float = 3000.0, composition_step: float = 0.02, temperature_step: float = 10.0,
                 ternary_step: float = 0.02, plot: bool = True, tol_K: float = 5.0, tol_x: float = 0.01,
                 n_workers: Optional[int] = None, timeout_s: float = 900.0, force_replot=False,
                 progress: Callable[[str], None] = lambda s: None) -> Dict[str, Any]:
    """force_replot: True (every subsystem whose change is exact / nil is mapped again in each step instead of being
    reused), a list of subsystem names ('Fe-Mo', 'Fe-Mo-W'), or False (default).  tentative: True (default) | False
    | {'mixing_per_atom': 50000, 'vacancy': 200000} (calphad_tools.tentative)."""
    forced = force_matcher(force_replot)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    r = build_versions(a_path, b_path, str(out), phase_map=phase_map, sublattice_map=sublattice_map,
                       decisions=decisions, policy=policy, extensions=extensions, tentative=tentative,
                       progress=progress)
    versions: Dict[str, Version] = r["versions"]
    if r.get("unmapped"):
        r["written"] = False
        r["report"] = r["pending_report"]
        return r
    if r.get("pending") is not None:
        r["written"] = False
        recommend_extend(r["pending"], r.get("A_eff") or a_path, r.get("B_merge") or r["B"])
        r["report"] = r["pending"].report()
        return r
    EA = sorted(parse_tdb(a_path).real_elements)
    EB = sorted(parse_tdb(r["B"]).real_elements)
    views = {}
    # a phase of B kept under a new name in Q (decision 'rename', e.g. MU_D85 -> MU_D85_2) is the same phase: the
    # B-side versions that still call it by B's name are viewed (and their invariants matched) under Q's name
    renames = {**(r.get("phase_renames") or {}), **(r.get("extension_renames") or {})}
    inv_rename: Dict[str, Dict[str, str]] = {}
    for lab in ("B", "B1", "B2", "B3", "Q"):
        tdb = parse_tdb(versions[lab].path)
        inv_rename[lab] = {} if lab == "Q" else ts.stage_renames(tdb, renames)
        views[lab] = ts.StageView(tdb, lab, elements=EB, rename=inv_rename[lab])
    renamed_in: Dict[Tuple[str, ...], str] = {}           # B's subsystems where a renamed phase exists -> label
    for old_n, new_n in renames.items():
        for sub in ts.stage_subsystems(EB, (2, 3)):
            if ts.restricted_model(views["B"], new_n, sub) is not None:
                renamed_in.setdefault(tuple(sorted(sub)), f"{old_n} kept as {new_n}")
    forced_user = forced
    chains = [["B", "B1", "B2", "B3", "Q"]]
    if "A1" in versions:
        versions["A"] = Version("A", a_path, f"{Path(a_path).name} (base)")
        for lab in ("A", "A1", "A2"):
            views[lab + "@A"] = ts.StageView(parse_tdb(versions[lab].path), lab, elements=EA)
        chains.append(["A", "A1", "A2"])
    temps = [float(t) for t in (ternary_temperatures_c or [])]
    settings = {"tmin_c": tmin_c, "tmax_c": tmax_c, "x_step": composition_step, "t_step": temperature_step,
                "tern_step": ternary_step, "timeout": timeout_s}
    steps: List[Step] = []
    need_all: Dict[Tuple[str, str], Optional[str]] = {}
    forced_all: set = set()
    items_by_chain = []
    known = set()
    for chain in chains:
        side_A = chain[0] == "A"
        els = EA if side_A else EB
        subs = ts.stage_subsystems(els, (2, 3))
        items = [Item(s) for s in subs if len(s) == 2] + [Item(s, T) for s in subs if len(s) == 3 for T in temps]
        known |= {tuple(sorted(s)) for s in subs}
        diffs = {}
        for old, new in zip(chain[:-1], chain[1:]):
            vo = views[old + "@A"] if side_A else views[old]
            vn = views[new + "@A"] if side_A else views[new]
            diffs[(old, new)] = ts.diff_versions(vo, vn, subs)
            st = Step(old, new, diffs[(old, new)], items)
            st.changed_items = [it for it in items if (st.diff.get(it.elements) or ts.SubsystemDiff(())).changed]
            steps.append(st)
        if side_A:     # only A's subsystems that change somewhere along A -> A1 -> A2 (or are forced)
            items = [it for it in items if forced(it.elements) or any(
                st.diff.get(it.elements) and st.diff.get(it.elements).changed for st in steps if st.old in chain)]
        need, fk = plan_plots(chain, versions, diffs, items, forced)
        if not side_A:      # a renamed phase's subsystems are mapped again in Q as evidence (same diagram expected)
            for it in items:
                if tuple(sorted(it.elements)) in renamed_in and need.get(("Q", it.tag)) is not None:
                    need[("Q", it.tag)] = None
                    fk.add(("Q", it.tag))
        need_all.update(need)
        forced_all |= fk
        for st in steps:
            if st.old in chain and st.new in chain:
                st.forced_items = [it for it in items if (st.new, it.tag) in fk]
                for it in st.forced_items:
                    sd = st.diff.get(it.elements)
                    rn = renamed_in.get(tuple(sorted(it.elements))) if not side_A else None
                    st.forced_kind[it.name] = ("exact rescale" if sd is not None and sd.notes else
                                               f"{rn}, numerically identical" if rn and not forced_user(it.elements)
                                               else "numerically identical")
        items_by_chain.append(items)
    all_items = {it.tag: it for its in items_by_chain for it in its}
    plots: Dict[Tuple[str, str], Dict] = {}
    fig_dir = out / "figures"
    if plot and need_all:
        plots = run_plots(versions, need_all, list(all_items.values()), fig_dir, settings, n_workers, progress,
                          nocache=forced_all)
    for st in steps:
        for its, target in ((st.changed_items, st.inv), (st.forced_items, st.forced_inv)):
            for it in its:
                ro, rn = plots.get((st.old, it.tag)), plots.get((st.new, it.tag))
                if ro and rn and ro.get("ok") and rn.get("ok"):
                    target[it.name] = compare_invariants(
                        rename_invariants(ro["invariants"], inv_rename.get(st.old, {})),
                        rename_invariants(rn["invariants"], inv_rename.get(st.new, {})), list(it.elements),
                        it.T_c is not None, tol_K, tol_x)
    for st in steps:      # end-members without G (pycalphad uses 0): does the phase become stable where mapped?
        done = set()
        for it in sorted(st.items, key=lambda x: (len(x.elements), x.T_c or 0.0)):     # binaries first
            sd = st.diff.get(it.elements)
            for m in (sd.models if sd is not None else []):
                ems = tuple(e for e in m.missing_g if (m.phase, e) not in done)
                if not ems:
                    continue
                done |= {(m.phase, e) for e in ems}
                rn = plots.get((st.new, it.tag)) or {}
                where = it.name if it.T_c is not None else f"{it.name} {tmin_c:g}-{tmax_c:g} °C"
                if not rn.get("ok"):
                    seen = "not mapped"
                elif m.phase in [_base(p) for p in rn.get("phases_found", [])]:
                    seen = f"{m.phase} STABLE in the mapped {where} - check it"
                else:
                    seen = f"{m.phase} not stable in the mapped {where}"
                st.zero_g.append(f"{m.phase} {', '.join(ems)}: {seen}")
    unknown = []
    if force_replot and force_replot is not True and not isinstance(force_replot, str):
        unknown = [n for n in force_replot if not any(force_matcher([n])(k) for k in known)]
    sa, sb = _stem(a_path), _stem(r["B_original"])
    files = {"deck": str(out / f"staged_merge_{sa}+{sb}.pptx"), "xlsx": str(out / f"staged_merge_{sa}+{sb}.xlsx")}
    norm = [n for n in r["notes"] if n.action in ts.NORM_CHANGES]
    note = clip("Normalised: " + "; ".join(n.text() for n in norm), 125, "Excel 'Versions'") if norm else ""
    if renames:
        rn_note = "; ".join(f"{o} of B kept as {n} in Q (decision 'rename'), compared as the same phase"
                            for o, n in renames.items())
        note = (note + " | " if note else "") + clip(rn_note, 125, "Excel 'Versions'")
    imgs: List[str] = []
    if plot:
        fig_dir.mkdir(parents=True, exist_ok=True)
        _, imgs = write_deck(steps, plots, files["deck"], Path(a_path).name, Path(r["B_original"]).name, tol_K,
                             fig_dir, note)
    else:
        files.pop("deck")
    write_xlsx(steps, plots, versions, files["xlsx"], r["notes"])
    r.update({"written": True, "steps": steps, "plots": plots, "files": files, "images": imgs,
              "renamed_in": {ts.stage_sysname(k): v for k, v in renamed_in.items()},
              "tol_K": tol_K, "plotted": plot, "forced": sorted(forced_all), "force_unknown": unknown})
    return r


def summary_text(r: Dict[str, Any], rel=lambda p: p) -> str:
    if not r.get("written") and r.get("unmapped"):
        return ("NOT FINISHED - decisions needed before staging: B's sublattices cannot be matched to A's for "
                + ", ".join(n.phase for n in r["unmapped"]) + ". Nothing written.\n"
                "Ask the user, then call again with sublattice_map (and decisions if the merge asks).\n"
                + r["report"])
    if not r.get("written"):
        v = r["versions"]
        return ("NOT FINISHED - decisions needed for Q (the full merge). B1/B2 written: "
                + ", ".join(rel(v[k].path) for k in ("B1", "B2") if k in v)
                + ".\nAsk the user, then call again with decisions (conflict id or 'system:A-B' -> option).\n"
                + r["report"])
    L = [f"staged_merge: A = {Path(r['A']).name} (base), B = {Path(r['B_original']).name}; shared "
         f"{', '.join(e.capitalize() for e in r['shared'])}"]
    norm = [n.text() for n in r["notes"] if n.action in ts.NORM_CHANGES]
    if norm:
        L.append("Normalised: " + "; ".join(norm[:4]) + (f" (+{len(norm) - 4})" if len(norm) > 4 else ""))
    for old_n, new_n in (r.get("phase_renames") or {}).items():
        where = [k for k, v in (r.get("renamed_in") or {}).items() if v.startswith(old_n + " ")]
        L.append(f"Kept under a new name: {old_n} of {Path(r['B_original']).name} is {new_n} in Q (its model is "
                 f"incompatible with A's {old_n}; decision 'rename'); B..B3's {old_n} is compared with Q's {new_n} as "
                 f"the same phase" + (f", and {', '.join(where)} replotted in Q as a check" if where else ""))
    for x in (r.get("extensions") or {}).get("reports", []):
        L.append("Extended (decision 'extend'): " + clip(" | ".join(x["text"].split("\n")[:3]), 400))
    # candidates only between phases still unpaired: an 'extend' decision pairs its two phases
    ext_pairs = {p for kv in (r.get("extension_renames") or {}).items() for p in kv}
    cands = [n.text() for n in r["notes"] if n.action in ("candidate", "incompatible", "unmapped", "not_rescaled")
             and not (n.action == "candidate" and (n.phase in ext_pairs or any(
                 re.search(rf"~ (?:A's )?{re.escape(p)}\b", n.detail) for p in ext_pairs)))]
    if cands:
        L.append("Not normalised: " + "; ".join(cands[:3]) + (f" (+{len(cands) - 3})" if len(cands) > 3 else ""))
    for st in r["steps"]:
        L.append(step_line(st, r["tol_K"]))
    L += tentative_lines(r.get("tentative"))
    if "A1" not in r["versions"]:
        L.append("A1/A2 not built: B has no unary or binary terms that A lacks.")
    if r.get("force_unknown"):
        L.append(f"force_replot: not a binary/ternary of A or B: {', '.join(map(str, r['force_unknown']))}")
    failed = [f"{st.title} {it.name}" for st in r["steps"] for it in st.forced_items
              if not st.forced_shift(it.name)[2] and math.isfinite(st.forced_shift(it.name)[0])]
    if failed:
        L.append("WARNING: forced replots of 'exact' changes moved invariants (exactness claim FAILED): "
                 + "; ".join(failed))
    bad = [f"{v.label}: {'; '.join(v.validation.get('errors', [])[:2])}" for v in r["versions"].values()
           if v.validation and (v.validation.get("errors") or not v.validation.get("loaded"))]
    if bad:
        L.append("pycalphad problems: " + " | ".join(bad))
    failed = [f"{k[0]} {k[1]}: {clip(v.get('error', ''), 70)}" for k, v in r["plots"].items() if not v.get("ok")]
    if failed:
        L.append(f"diagrams failed: {'; '.join(failed[:4])}")
    checked = [f"{k[0]} {k[1]}: {clip(v['mapping_check'], 160)}" for k, v in r["plots"].items()
               if v.get("mapping_check") and not v.get("reused_from")]
    if checked:
        L.append("mapping check: " + " | ".join(checked[:3]) + (f" (+{len(checked) - 3})" if len(checked) > 3 else ""))
    a0 = [f"A0 {rel(r['A_eff'])}"] if r.get("A_eff") and r["A_eff"] != r["A"] else []
    L.append("TDBs: " + ", ".join(a0 + [f"{k} {rel(v.path)}" for k, v in r["versions"].items() if k != "A"]))
    L.append("Files: " + ", ".join(f"{k} {rel(v)}" for k, v in r["files"].items()))
    return "\n".join(L)


# =============================================================================================== tool
TENTATIVE_SCHEMA = {"anyOf": [{"type": "boolean"}, {"type": "object"}],
                    "description": "tentative L0 terms for new mixings: true (default) | false | "
                                   "{mixing_per_atom: 50000, vacancy: 200000}"}


def tentative_arg(tb, value):
    """A tool call's tentative argument, the [merge] config defaults when not given."""
    if value is not None:
        return value
    f = getattr(tb, "tentative_default", None)
    return f() if callable(f) else True


def register(tb) -> None:
    S = {"type": "string"}
    N = {"type": "number"}

    @tb.add("staged_merge",
            "Merge add_tdb (B) into base_tdb (A) step by step, diagrams per step. B1 = B with A's SGTE unaries, "
            "B2 = + A's other unaries, B3 = + A's shared binaries, Q = full merge (A wins what the stages settle; other "
            "conflicts are asked like merge_tdbs: re-call with decisions), A1/A2 = A + B's unary / binary terms; "
            "A0 = A rescaled to a magnetic B phase's site ratios when needed (BMAGN end-members exact if A is "
            "magnetic too). "
            "Maps B's binaries+ternary for each version (unchanged subsystems reused), compares consecutive "
            "versions (parameters, phase models, invariant shifts) -> one deck (summary + 2 slides per step) + "
            "Excel; force_replot maps exact-only changes anyway as evidence. Run evaluate_compatibility first.",
            {"base_tdb": S, "add_tdb": S, "output_dir": S,
             "decisions": {"type": "object", "additionalProperties": {"type": "string"}},
             "policy": {"type": "string", "enum": ["ask", "first", "last"]},
             "phase_map": {"type": "object"}, "sublattice_map": {"type": "object"},
             "extend": {"type": "object", "description": "{A phase: {phase: B phase, components, ...}} (as "
                                                         "extend_phase)"},
             "ternary_temperatures_c": {"type": "array", "items": N, "description": "isotherms, default [1000]"},
             "tmin_c": N, "tmax_c": N, "composition_step": N, "temperature_step": N,
             "plot": {"type": "boolean", "description": "default true; false = TDBs + parameter/model diffs only"},
             "tol_K": {**N, "description": "invariant shift threshold, default 5 K"}, "n_workers": {"type": "integer"},
             "force_replot": {"anyOf": [{"type": "boolean"}, {"type": "array", "items": S}],
                              "description": "true or ['Fe-Mo', ...]: recompute subsystems whose change is exact "
                                             "(e.g. a rescale) instead of reusing them, and show the invariant shift"},
             "tentative": TENTATIVE_SCHEMA},
            ["base_tdb", "add_tdb"])
    def staged_merge_tool(base_tdb, add_tdb, output_dir=None, decisions=None, policy: str = "ask", phase_map=None,
                          sublattice_map=None, ternary_temperatures_c=None, tmin_c: float = 300.0,
                          tmax_c: float = 3000.0, composition_step: float = 0.02, temperature_step: float = 10.0,
                          plot: bool = True, tol_K: float = 5.0, n_workers: Optional[int] = None,
                          force_replot=False, extend=None, tentative=None):
        a, b = tb._find_tdb(base_tdb), tb._find_tdb(add_tdb)
        key = f"staged|{a}|{b}"
        dec = {**tb.merge_decisions.get(key, {}), **(decisions or {})}
        tb.merge_decisions[key] = dec
        od = tb.outdir(output_dir, f"staged_merge_{Path(a).stem}+{Path(b).stem}")
        r = staged_merge(a, b, str(od), phase_map=phase_map, sublattice_map=sublattice_map, decisions=dec,
                         policy=policy, extensions=extend or None, tentative=tentative_arg(tb, tentative),
                         ternary_temperatures_c=(1000.0,) if ternary_temperatures_c is None
                         else ternary_temperatures_c, tmin_c=tmin_c, tmax_c=tmax_c,
                         composition_step=composition_step, temperature_step=temperature_step, plot=plot,
                         tol_K=tol_K, n_workers=n_workers, force_replot=force_replot, progress=tb.ui.progress)
        return summary_text(r, tb.rel), (r.get("images") or [])[:2]
