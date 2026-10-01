"""system_coverage: which subsystems of an element system are actually assessed in the available databases.

For elements {Fe, Ni, Cr, ...} it counts the binaries C(n,2), ternaries C(n,3) (optionally quaternaries) and reports,
from every TDB in a folder (including databases that contain only some of the elements), which of those subsystems
have parameters of their own (assessed) and which a calculation would only extrapolate from lower-order systems.

A subsystem counts as *assessed in a database* when that database has at least one parameter whose constituents
involve exactly those elements (G of a compound end-member, L, TC/BMAGN, ...).  Examples: a Fe-Cr binary TDB plus a
Fe-Ni binary TDB give, for Cr-Fe-Ni, 2 of 3 binaries assessed and 0 of 1 ternary.  Note that subsystems assessed in
*different* databases are not automatically usable together (different unary data, phase names, models).
"""
from __future__ import annotations

import itertools
from math import comb
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple

from . import compare_system as cs
from .tdbtables import Subsystem, key_constituents

ORDER_NAMES = {1: "unaries", 2: "binaries", 3: "ternaries", 4: "quaternaries", 5: "quinaries"}


def _cap(e: str) -> str:
    return e.capitalize()


def sysname(els) -> str:
    return "-".join(_cap(e) for e in els)


class DbCoverage:
    def __init__(self, path: Path, label: str, rel: str, system: Sequence[str]):
        tdb = cs._parsed(path)
        self.label, self.rel = label, rel
        self.elements = sorted(e for e in tdb.elements if e not in ("VA", "/-"))
        self.in_system = [e for e in system if e in tdb.elements]
        self.counts: Dict[Tuple[str, ...], int] = {}        # subsystem (sorted elements) -> number of parameters
        self.phases: Dict[Tuple[str, ...], Set[str]] = {}
        if not self.in_system:
            return
        sub = Subsystem(tdb, label, rel, self.in_system, {})
        for r in sub.rows.values():
            els = tuple(sorted({e for sl in key_constituents(r.key) for s in sl for e in sub.species_elements(s)}
                               - {"VA", "*"}))
            if els:
                self.counts[els] = self.counts.get(els, 0) + 1
                self.phases.setdefault(els, set()).add(r.phase)


def coverage(paths: Sequence[str], elements: Sequence[str], workdir: Path, max_order: int = 3):
    system = sorted({e.upper() for e in elements})
    labels = cs._labels_for(list(paths), workdir)
    dbs = [DbCoverage(Path(p), lab, cs._rel(Path(p), workdir), system) for p, lab in zip(paths, labels)]
    dbs = [d for d in dbs if d.in_system]
    table: Dict[int, List[Dict]] = {}
    for k in range(1, max_order + 1):
        rows = []
        for combo in itertools.combinations(system, k):
            where = [(d.label, d.counts[combo], sorted(d.phases[combo])) for d in dbs if combo in d.counts]
            row = {"system": combo, "assessed_in": where}
            if k >= 3:     # how many of its binaries are assessed somewhere
                bins = list(itertools.combinations(combo, 2))
                row["binaries_assessed"] = sum(1 for b in bins if any(b in d.counts for d in dbs))
                row["binaries_total"] = len(bins)
                row["binaries_in_one_db"] = max((sum(1 for b in bins if b in d.counts) for d in dbs), default=0)
            rows.append(row)
        table[k] = rows
    return system, dbs, table


def summary(system, dbs, table) -> Dict:
    n = len(system)
    out = {"elements": system, "missing_elements": [e for e in system if not any(e in d.in_system for d in dbs)]}
    for k, rows in table.items():
        out[ORDER_NAMES.get(k, f"order {k}")] = {"total": comb(n, k), "assessed": sum(1 for r in rows if r["assessed_in"])}
    return out


def best_single(dbs, table) -> List[Tuple[str, int, int]]:
    """(database, binaries covered, ternaries covered), best first."""
    b = [r["system"] for r in table.get(2, [])]
    t = [r["system"] for r in table.get(3, [])]
    res = [(d.label, sum(1 for s in b if s in d.counts), sum(1 for s in t if s in d.counts)) for d in dbs]
    return sorted(res, key=lambda x: (-x[1], -x[2], x[0]))


def summary_text(system, dbs, table, files, rel=lambda p: p, max_list: int = 40) -> str:
    s = summary(system, dbs, table)
    lines = [f"{sysname(system)} ({len(system)} elements), {len(dbs)} database(s) containing at least one of them:"]
    for k, rows in table.items():
        name = ORDER_NAMES.get(k, f"order {k}")
        v = s[name]
        if k == 1:
            lines.append(f"  elements: {v['assessed']} of {v['total']} have data in at least one database")
            continue
        lines.append(f"  {name}: {v['total']} in total, {v['assessed']} assessed in at least one database, "
                     f"{v['total'] - v['assessed']} extrapolated only")
    if s["missing_elements"]:
        lines.append("  elements in no database: " + ", ".join(_cap(e) for e in s["missing_elements"]))
    for k in (2, 3):
        rows = table.get(k, [])
        if not rows:
            continue
        name = ORDER_NAMES[k]
        ass = [r for r in rows if r["assessed_in"]]
        miss = [r for r in rows if not r["assessed_in"]]
        lines.append(f"Assessed {name} ({len(ass)}):")
        for r in ass[:max_list]:
            extra = (f"; binaries {r['binaries_assessed']}/{r['binaries_total']}" if k == 3 else "")
            lines.append(f"  {sysname(r['system'])}: " + ", ".join(f"{db} ({n} params)" for db, n, _ in r["assessed_in"])
                         + extra)
        if miss:
            if k == 3:
                txt = ", ".join(f"{sysname(r['system'])} ({r['binaries_assessed']}/{r['binaries_total']} binaries)"
                                for r in miss[:max_list])
            else:
                txt = ", ".join(sysname(r["system"]) for r in miss[:max_list])
            lines.append(f"Not assessed ({len(miss)}): " + txt + (" ..." if len(miss) > max_list else ""))
    bs = best_single(dbs, table)
    if bs:
        nb, nt = len(table.get(2, [])), len(table.get(3, []))
        lines.append("Coverage per database (binaries, ternaries): " + "; ".join(
            f"{db} {b}/{nb}" + (f", {t}/{nt}" if nt else "") for db, b, t in bs[:10]))
    lines.append("Note: subsystems assessed in different databases are not automatically compatible (unary data, "
                 "phase names and models can differ).")
    lines.append("Files: " + ", ".join(f"{k}: {rel(v)}" for k, v in files.items()))
    return "\n".join(lines)


# ------------------------------------------------------------------------------------------------ files
def write_xlsx(system, dbs, table, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    head = PatternFill("solid", fgColor="2C5364")
    yes, no = PatternFill("solid", fgColor="E8F5E9"), PatternFill("solid", fgColor="FFEBEE")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = summary(system, dbs, table)
    ws.append(["system", sysname(system)])
    ws.append(["databases with any of the elements", len(dbs)])
    for k in table:
        name = ORDER_NAMES.get(k, f"order {k}")
        ws.append([name, f"{s[name]['assessed']} assessed of {s[name]['total']}"])
    if s["missing_elements"]:
        ws.append(["elements in no database", ", ".join(_cap(e) for e in s["missing_elements"])])
    ws.append([])
    ws.append(["database", "file", "elements in file", "binaries covered", "ternaries covered"])
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    bs = {db: (b, t) for db, b, t in best_single(dbs, table)}
    for d in dbs:
        ws.append([d.label, d.rel, ",".join(_cap(e) for e in d.elements), bs[d.label][0], bs[d.label][1]])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 30
    # binary matrix
    bm = wb.create_sheet("Binary matrix")
    bm.append([""] + [_cap(e) for e in system])
    by = {r["system"]: r for r in table.get(2, [])}
    for i, a in enumerate(system):
        row = [_cap(a)]
        for j, b in enumerate(system):
            if i == j:
                row.append("")
            else:
                r = by[tuple(sorted((a, b)))]
                row.append(", ".join(db for db, _, _ in r["assessed_in"]) or "—")
        bm.append(row)
    for row in bm.iter_rows(min_row=2, min_col=2):
        for c in row:
            if c.value:
                c.fill = no if c.value == "—" else yes
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(len(system) + 1):
        bm.column_dimensions[bm.cell(1, i + 1).column_letter].width = 8 if i == 0 else 22
    for k, rows in table.items():
        if k == 1:
            continue
        sh = wb.create_sheet(ORDER_NAMES.get(k, f"order {k}").capitalize())
        hdr = ["system", "assessed", "databases (parameters)", "phases with parameters"]
        if k >= 3:
            hdr += ["binaries assessed (any database)", "most binaries in one database"]
        sh.append(hdr)
        for c in sh[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            vals = [sysname(r["system"]), "yes" if r["assessed_in"] else "no",
                    "; ".join(f"{db} ({n})" for db, n, _ in r["assessed_in"]),
                    "; ".join(f"{db}: {','.join(ph)}" for db, _, ph in r["assessed_in"])]
            if k >= 3:
                vals += [f"{r['binaries_assessed']}/{r['binaries_total']}", f"{r['binaries_in_one_db']}/{r['binaries_total']}"]
            sh.append(vals)
            sh.cell(sh.max_row, 2).fill = yes if r["assessed_in"] else no
        for col, w in zip("ABCDEF", (16, 10, 50, 60, 18, 18)):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "B2"
    wb.save(path)
    return path


def write_pptx(system, dbs, table, path: str) -> str:
    from . import pptx_utils as pu
    from .pptx_utils import GridCell as G
    prs = pu.new_presentation()
    s = summary(system, dbs, table)
    pu.add_title_slide(prs, f"{sysname(system)}: subsystem coverage",
                       f"{len(system)} elements, {len(dbs)} database(s) containing at least one of them",
                       note="Assessed = the database has parameters involving exactly those elements. Anything "
                            "else is extrapolated from lower-order systems.")
    rows = []
    for k in table:
        name = ORDER_NAMES.get(k, f"order {k}")
        v = s[name]
        if k == 1:
            rows.append([G("Elements with data", bold=True), G(str(v["total"]), align="r"),
                         G(str(v["assessed"]), align="r", fill="same"),
                         G(str(v["total"] - v["assessed"]) + " (no data)", align="r",
                           fill="differs" if v["total"] - v["assessed"] else None)])
            continue
        rows.append([G(name.capitalize(), bold=True), G(str(v["total"]), align="r"),
                     G(str(v["assessed"]), align="r", fill="same" if v["assessed"] else None),
                     G(str(v["total"] - v["assessed"]), align="r", fill="differs" if v["total"] - v["assessed"] else None)])
    pu.add_flow_table_slides(prs, "Summary", ["", "total", "assessed", "extrapolated only"], rows, [3.0, 2.0, 2.0, 2.4],
                             font_size=12, subtitle=("elements in no database: " + ", ".join(
                                 _cap(e) for e in s["missing_elements"])) if s["missing_elements"] else None)
    # binary matrix (upper triangle)
    by = {r["system"]: r for r in table.get(2, [])}
    if by:
        n = len(system)
        w_first = 1.0
        w = min(1.9, (12.3 - w_first) / max(1, n - 1))
        rows = []
        for i, a in enumerate(system[:-1]):
            cells = [G(_cap(a), bold=True)]
            for b in system[1:]:
                if system.index(b) <= i:
                    cells.append(G(""))
                    continue
                r = by[tuple(sorted((a, b)))]
                cells.append(G("\n".join(db for db, _, _ in r["assessed_in"]) or "—", size=7,
                               fill="same" if r["assessed_in"] else "differs"))
            rows.append(cells)
        pu.add_flow_table_slides(prs, "Binaries", [""] + [_cap(b) for b in system[1:]], rows,
                                 [w_first] + [w] * (n - 1), font_size=8,
                                 subtitle=f"{s['binaries']['assessed']} of {s['binaries']['total']} assessed; "
                                          "cells list the databases")
    for k in (3, 4):
        rs = table.get(k)
        if not rs:
            continue
        rows = []
        for r in sorted(rs, key=lambda r: (not r["assessed_in"], -r.get("binaries_assessed", 0), r["system"])):
            rows.append([G(sysname(r["system"]), bold=True),
                         G(", ".join(f"{db} ({n})" for db, n, _ in r["assessed_in"]) or "—",
                           fill="same" if r["assessed_in"] else "differs"),
                         G(f"{r['binaries_assessed']}/{r['binaries_total']}", align="c"),
                         G(f"{r['binaries_in_one_db']}/{r['binaries_total']}", align="c")])
        name = ORDER_NAMES[k]
        pu.add_flow_table_slides(prs, name.capitalize(), ["system", "assessed in (parameters)",
                                                          "binaries assessed", "most binaries in one DB"],
                                 rows, [2.6, 5.9, 1.8, 2.0], font_size=9,
                                 subtitle=f"{s[name]['assessed']} of {s[name]['total']} assessed")
    bs = best_single(dbs, table)
    nb, nt = len(table.get(2, [])), len(table.get(3, []))
    rows = [[G(db, bold=True), G(next(d.rel for d in dbs if d.label == db)),
             G(",".join(_cap(e) for e in next(d.elements for d in dbs if d.label == db))),
             G(f"{b}/{nb}", align="c"), G(f"{t}/{nt}" if nt else "", align="c")] for db, b, t in bs]
    pu.add_flow_table_slides(prs, "Databases", ["database", "file", "elements", "binaries", "ternaries"], rows,
                             [2.6, 4.0, 3.0, 1.2, 1.2], font_size=9,
                             footnote="Subsystems assessed in different databases are not automatically compatible "
                                      "(unary data, phase names and models can differ).")
    return pu.save(prs, path)


def register(tb) -> None:
    @tb.add("system_coverage",
            "Statistics of an element system across the databases in a folder: number of binaries/ternaries "
            "(quaternaries with max_order 4), which are assessed (have their own parameters) in which database, "
            "which are only extrapolated, and each ternary's assessed binaries. Uses databases containing any of "
            "the elements. Writes xlsx + pptx.",
            {"elements": {"type": "array", "items": {"type": "string"}},
             "folder": {"type": "string"}, "pattern": {"type": "string"},
             "tdbs": {"type": "array", "items": {"type": "string"}},
             "max_order": {"type": "integer", "description": "2-5, default 3"},
             "output": {"type": "string", "description": "base file name"}}, ["elements"])
    def system_coverage(elements, folder=None, pattern=None, tdbs=None, max_order=3, output=None):
        if tdbs:
            paths = [tb._find_tdb(t) for t in tdbs]
        else:
            base = (tb.workdir / folder) if folder else tb.workdir
            if not base.exists():
                raise FileNotFoundError(f"folder {folder} not found under {tb.workdir}")
            paths = [str(p) for p in cs.database_files(base, pattern, True) if p.suffix.lower() != ".dat"]
        max_order = max(2, min(5, int(max_order)))
        system, dbs, table = coverage(paths, elements, tb.workdir, max_order)
        if not dbs:
            raise ValueError(f"no database under {folder or 'the working directory'} contains any of "
                             f"{', '.join(_cap(e) for e in system)}")
        base = tb.outp(output, f"coverage_{sysname(system)}")
        base = str(Path(base).with_suffix("")) if Path(base).suffix.lower() in (".xlsx", ".pptx") else base
        files = {"xlsx": write_xlsx(system, dbs, table, base + ".xlsx"),
                 "pptx": write_pptx(system, dbs, table, base + ".pptx")}
        return summary_text(system, dbs, table, files, tb.rel), []
