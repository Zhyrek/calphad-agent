"""Small python-pptx helpers shared by all report generators (16:9, clean style)."""
from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
DARK = RGBColor(0x1F, 0x2A, 0x37)
ACCENT = RGBColor(0x2C, 0x53, 0x64)
MUTED = RGBColor(0x6B, 0x72, 0x80)
HEADER_FILL = RGBColor(0x2C, 0x53, 0x64)
BAND_FILL = RGBColor(0xF3, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

STATUS_COLORS = {
    "identical": RGBColor(0xE8, 0xF5, 0xE9),
    "equivalent": RGBColor(0xE3, 0xF2, 0xFD),
    "different": RGBColor(0xFF, 0xEB, 0xEE),
    "only A": RGBColor(0xFF, 0xF8, 0xE1),
    "only B": RGBColor(0xF3, 0xE5, 0xF5),
}


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, left, top, width, height, text, size=18, bold=False, color=DARK,
             align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def add_title(slide, title: str, subtitle: Optional[str] = None):
    _textbox(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7), title, size=26, bold=True)
    bar = slide.shapes.add_shape(1, Inches(0.5), Inches(1.0), Inches(1.2), Emu(38000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if subtitle:
        _textbox(slide, Inches(0.5), Inches(1.08), Inches(12.3), Inches(0.4), subtitle, size=13,
                 color=MUTED)


def add_title_slide(prs, title: str, subtitle: str = "", note: str = ""):
    s = _blank(prs)
    _textbox(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.2), title, size=38, bold=True)
    bar = s.shapes.add_shape(1, Inches(0.8), Inches(3.6), Inches(2.0), Emu(50000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if subtitle:
        _textbox(s, Inches(0.8), Inches(3.8), Inches(11.7), Inches(1.0), subtitle, size=18, color=MUTED)
    if note:
        _textbox(s, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.5), note, size=11, color=MUTED)
    return s


def add_bullets_slide(prs, title: str, bullets: Sequence[str], subtitle: Optional[str] = None,
                      size: int = 16):
    s = _blank(prs)
    add_title(s, title, subtitle)
    _textbox(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.6),
             "\n".join(f"•  {b}" for b in bullets), size=size)
    return s


def add_table_slides(prs, title: str, header: Sequence[str], rows: List[Sequence[str]],
                     col_widths_in: Sequence[float], rows_per_slide: int = 14,
                     font_size: int = 10, status_col: Optional[int] = None,
                     subtitle: Optional[str] = None, mono_cols: Sequence[int] = ()) -> int:
    """Add as many slides as needed to show ``rows``.  Returns the slide count."""
    if not rows:
        s = _blank(prs)
        add_title(s, title, subtitle)
        _textbox(s, Inches(0.6), Inches(2.0), Inches(12), Inches(1), "(no rows)", size=16, color=MUTED)
        return 1
    pages = [rows[i:i + rows_per_slide] for i in range(0, len(rows), rows_per_slide)]
    for pi, page in enumerate(pages):
        s = _blank(prs)
        t = title if len(pages) == 1 else f"{title}  ({pi + 1}/{len(pages)})"
        add_title(s, t, subtitle)
        total_w = sum(col_widths_in)
        left = (13.333 - total_w) / 2
        row_h = Inches(0.36)
        shape = s.shapes.add_table(len(page) + 1, len(header), Inches(left), Inches(1.55),
                                   Inches(total_w), row_h * (len(page) + 1))
        tbl = shape.table
        for j, w in enumerate(col_widths_in):
            tbl.columns[j].width = Inches(w)
        for j, h in enumerate(header):
            _cell(tbl.cell(0, j), str(h), font_size + 1, bold=True, color=WHITE, fill=HEADER_FILL)
        for i, row in enumerate(page, start=1):
            fill = BAND_FILL if i % 2 == 0 else WHITE
            for j, val in enumerate(row):
                f = fill
                if status_col is not None and j == status_col:
                    f = STATUS_COLORS.get(str(val), fill)
                _cell(tbl.cell(i, j), str(val), font_size, fill=f,
                      font="Courier New" if j in mono_cols else "Calibri")
    return len(pages)


def _cell(cell, text, size, bold=False, color=DARK, fill=WHITE, font="Calibri"):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.05)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ""
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font


def add_image_slide(prs, title: str, image_path: str, subtitle: Optional[str] = None,
                    notes: Optional[str] = None):
    """Add a slide with one image scaled to fit under the title."""
    from PIL import Image

    s = _blank(prs)
    add_title(s, title, subtitle)
    top = Inches(1.5) if subtitle else Inches(1.3)
    box_w, box_h = SLIDE_W - Inches(0.8), SLIDE_H - top - Inches(0.25)
    with Image.open(image_path) as im:
        w, h = im.size
    scale = min(box_w / w, box_h / h)
    pw, ph = int(w * scale), int(h * scale)
    s.shapes.add_picture(str(image_path), int((SLIDE_W - pw) / 2), int(top), pw, ph)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def save(prs, path: str) -> str:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(p))
    return str(p)


# ---------------------------------------------------------------------------
# Grid tables: two header rows, merged cells, per-cell fills (N-way comparisons)
# ---------------------------------------------------------------------------
GRID_FILLS = {
    "differs": RGBColor(0xFF, 0xEB, 0xEE),
    "same": RGBColor(0xE8, 0xF5, 0xE9),
    "absent": RGBColor(0xF3, 0xF4, 0xF6),
    "no ref": RGBColor(0xFF, 0xF8, 0xE1),
    "error": RGBColor(0xFD, 0xE6, 0x8A),
    "ref": RGBColor(0xE3, 0xF2, 0xFD),
    "band": RGBColor(0xD9, 0xE2, 0xE8),
}


class GridCell:
    """One table cell: text, column span, fill (a GRID_FILLS key), style."""

    def __init__(self, text: str = "", span: int = 1, fill: Optional[str] = None, bold: bool = False,
                 mono: bool = False, muted: bool = False, size: Optional[int] = None, align: str = "l",
                 sub: str = ""):
        self.text, self.span, self.fill, self.bold, self.sub = str(text), int(span), fill, bold, sub
        self.mono, self.muted, self.size, self.align = mono, muted, size, align


def add_grid_slides(prs, title: str, head1: Sequence, head2: Optional[Sequence[str]], rows: List[List[GridCell]],
                    col_widths_in: Sequence[float], font_size: int = 9, subtitle: Optional[str] = None,
                    footnote: Optional[str] = None, rows_per_slide: int = 14, row_h_in: float = 0.32) -> int:
    """Table slides with a grouped header.

    head1: [(text, span), ...] covering all columns; head2: one label per column ("" = merge with the cell above) or
    None.  rows: lists of GridCell whose spans add up to the column count.  Returns the number of slides."""
    ncols = len(col_widths_in)
    if sum(s for _, s in head1) != ncols:
        raise ValueError("header spans do not match the column count")
    n_pages = max(1, -(-len(rows) // rows_per_slide))
    per = max(1, -(-len(rows) // n_pages))                  # balanced pages (no 1-row leftover slide)
    pages = [rows[i:i + per] for i in range(0, len(rows), per)] or [[]]
    nh = 2 if head2 else 1
    total_w = sum(col_widths_in)
    left = max(0.2, (13.333 - total_w) / 2)
    top = 1.55 if subtitle else 1.35
    for pi, page in enumerate(pages):
        s = _blank(prs)
        add_title(s, title if len(pages) == 1 else f"{title}  ({pi + 1}/{len(pages)})", subtitle)
        if not page:
            _textbox(s, Inches(0.6), Inches(2.0), Inches(12), Inches(1), "(no rows)", size=16, color=MUTED)
            continue
        nrows = len(page) + nh
        shape = s.shapes.add_table(nrows, ncols, Inches(left), Inches(top), Inches(total_w),
                                   Inches(row_h_in * nrows))
        tbl = shape.table
        tbl.first_row = False
        tbl.horz_banding = False
        for j, w in enumerate(col_widths_in):
            tbl.columns[j].width = Inches(w)
        for i in range(nrows):
            tbl.rows[i].height = Inches(row_h_in)
        col = 0
        for text, span in head1:
            if nh == 2 and span == 1 and not head2[col]:
                tbl.cell(0, col).merge(tbl.cell(1, col))
            elif span > 1:
                tbl.cell(0, col).merge(tbl.cell(0, col + span - 1))
            _grid_cell(tbl.cell(0, col), GridCell(text, bold=True, align="c"), font_size, header=True)
            col += span
        if nh == 2:
            for j, text in enumerate(head2):
                if text:
                    _grid_cell(tbl.cell(1, j), GridCell(text, bold=True, align="c"), font_size - 1, header=True)
        for i, row in enumerate(page, start=nh):
            band = BAND_FILL if (i - nh) % 2 else WHITE
            col = 0
            for c in row:
                if c.span > 1:
                    tbl.cell(i, col).merge(tbl.cell(i, col + c.span - 1))
                _grid_cell(tbl.cell(i, col), c, font_size, band=band)
                col += c.span
            if col != ncols:
                raise ValueError(f"row {i} spans {col} columns, table has {ncols}")
        if footnote:
            _textbox(s, Inches(0.5), Inches(7.02), Inches(12.3), Inches(0.35), footnote, size=9, color=MUTED)
    return len(pages)


def _grid_cell(cell, c: GridCell, size: int, header: bool = False, band=WHITE):
    fill = HEADER_FILL if header else GRID_FILLS.get(c.fill, band) if c.fill else band
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.04)
    cell.margin_top = cell.margin_bottom = Inches(0.015)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    lines = c.text.split("\n")
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = {"c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}.get(c.align, PP_ALIGN.LEFT)
        if k == 0:
            p.text = ""
        r = p.add_run()
        r.text = line
        r.font.size = Pt(c.size or size)
        r.font.bold = c.bold
        r.font.color.rgb = WHITE if header else (MUTED if c.muted else DARK)
        r.font.name = "Courier New" if c.mono else "Calibri"
    if c.sub and not header:
        for line in c.sub.split("\n"):
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = line
            r.font.size = Pt(max(6, (c.size or size) - 1))
            r.font.italic = True
            r.font.color.rgb = MUTED
            r.font.name = "Calibri"


def _est_lines(text: str, width_in: float, size: float, mono: bool) -> int:
    cpl = max(6, int((width_in - 0.1) / (size * (0.6 if mono else 0.52) / 72.0) * 0.93))
    n = 0
    for line in text.split("\n"):
        n += max(1, -(-len(line) // cpl))
    return n


def estimate_row_height(cells: List[GridCell], widths: Sequence[float], size: float) -> float:
    h, col = 0.0, 0
    for c in cells:
        w = sum(widths[col:col + c.span])
        col += c.span
        sz = c.size or size
        lines_h = _est_lines(c.text, w, sz, c.mono) * sz * 1.2 / 72.0
        if getattr(c, "sub", ""):
            lines_h += _est_lines(c.sub, w, sz - 1, False) * (sz - 1) * 1.2 / 72.0
        h = max(h, lines_h)
    return h + 0.07


def add_flow_table_slides(prs, title: str, header: Sequence[str], rows: List, col_widths_in: Sequence[float],
                          font_size: int = 8, subtitle: Optional[str] = None, footnote: Optional[str] = None,
                          max_height_in: float = 5.3, header_h_in: float = 0.3) -> int:
    """Tables with rows of varying height (wrapped expressions), paginated so nothing runs off the slide.

    rows: lists of GridCell (spans add up to the column count) or ("band", text) group headings, which are repeated
    as "text (cont.)" when a page starts inside the group.  Returns the number of slides."""
    ncols = len(col_widths_in)
    pages: List[List] = []
    cur: List = []
    used = 0.0
    band = None
    band_h = 0.26
    avail = max_height_in - header_h_in
    hs = [band_h if isinstance(r, tuple) else estimate_row_height(r, col_widths_in, font_size) for r in rows]
    total = sum(hs)
    if total > avail:                                   # balance the pages (no near-empty last slide)
        n_pages = -(-total // avail)
        avail = min(avail, total / n_pages * 1.12)
    for idx, (r, h) in enumerate(zip(rows, hs)):
        # a group heading needs its first row on the same page (never a heading alone at the bottom)
        need = h + (hs[idx + 1] if isinstance(r, tuple) and idx + 1 < len(hs) else 0.0)
        if cur and used + need > avail:
            pages.append(cur)
            cur, used = [], 0.0
            if band is not None and not isinstance(r, tuple):
                cur.append(("band", band + " (cont.)"))
                used += band_h
        if isinstance(r, tuple):
            band = r[1]
        cur.append(r)
        used += h
    if cur:
        pages.append(cur)
    if not pages:
        pages = [[]]
    total_w = sum(col_widths_in)
    left = max(0.2, (13.333 - total_w) / 2)
    top = 1.55 if subtitle else 1.35
    for pi, page in enumerate(pages):
        s = _blank(prs)
        add_title(s, title if len(pages) == 1 else f"{title}  ({pi + 1}/{len(pages)})", subtitle)
        if not page:
            _textbox(s, Inches(0.6), Inches(2.0), Inches(12), Inches(1), "(no rows)", size=16, color=MUTED)
            continue
        heights = [header_h_in] + [band_h if isinstance(r, tuple) else estimate_row_height(r, col_widths_in, font_size)
                                   for r in page]
        shape = s.shapes.add_table(len(page) + 1, ncols, Inches(left), Inches(top), Inches(total_w),
                                   Inches(sum(heights)))
        tbl = shape.table
        tbl.first_row = False
        tbl.horz_banding = False
        for j, w in enumerate(col_widths_in):
            tbl.columns[j].width = Inches(w)
        for i, h in enumerate(heights):
            tbl.rows[i].height = Inches(h)
        for j, text in enumerate(header):
            _grid_cell(tbl.cell(0, j), GridCell(text, bold=True, align="c"), font_size + 1, header=True)
        k = 0
        for i, r in enumerate(page, start=1):
            if isinstance(r, tuple):
                tbl.cell(i, 0).merge(tbl.cell(i, ncols - 1))
                _grid_cell(tbl.cell(i, 0), GridCell(r[1], bold=True, fill="band"), font_size + 1)
                k = 0
                continue
            bandfill = BAND_FILL if k % 2 else WHITE
            k += 1
            col = 0
            for c in r:
                if c.span > 1:
                    tbl.cell(i, col).merge(tbl.cell(i, col + c.span - 1))
                _grid_cell(tbl.cell(i, col), c, font_size, band=bandfill)
                col += c.span
            if col != ncols:
                raise ValueError(f"row spans {col} columns, table has {ncols}")
        if footnote:
            _textbox(s, Inches(0.5), Inches(7.02), Inches(12.3), Inches(0.35), footnote, size=9, color=MUTED)
    return len(pages)


# ---------------------------------------------------------------------------
# Line diffs (TDB changes): removed lines in red text, added lines in the slide's normal text colour
# ---------------------------------------------------------------------------
DIFF_RED = RGBColor(0xC0, 0x00, 0x00)              # removed / old, on a light background
DIFF_RED_ON_DARK = RGBColor(0xFF, 0x80, 0x80)      # removed / old, on a dark background
DIFF_BAND = RGBColor(0xE6, 0xE9, 0xED)             # group headers: light grey band
_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
       "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def _luminance(hexrgb: str) -> float:
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _color_hex(el, master) -> Optional[str]:
    """Hex RGB of a colour element (srgbClr / sysClr / schemeClr resolved through the master's colour map and
    theme), None if unknown."""
    for c in el.iter():
        tag = c.tag.split("}")[-1]
        if tag == "srgbClr":
            return c.get("val")
        if tag == "sysClr":
            return c.get("lastClr") or ("FFFFFF" if c.get("val") == "window" else "000000")
        if tag == "schemeClr":
            name = c.get("val")
            cmap = master._element.find("p:clrMap", _NS)
            if cmap is not None and cmap.get(name):
                name = cmap.get(name)
            name = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}.get(name, name)
            try:
                from lxml import etree
                th = master.part.part_related_by(
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
                root = etree.fromstring(th.blob)
                sc = root.find(f".//a:clrScheme/a:{name}", _NS)
                if sc is not None:
                    return _color_hex(sc, master)
            except Exception:                         # noqa: BLE001 - no theme: assume the usual light background
                pass
            return {"lt1": "FFFFFF", "lt2": "EEECE1", "dk1": "000000", "dk2": "1F497D"}.get(name)
    return None


def background_is_dark(slide) -> bool:
    """True if the slide's background (its own, else its layout's, else its master's) is dark.  Read-only: the
    background XML is inspected, never created."""
    master = slide.slide_layout.slide_master
    for owner in (slide, slide.slide_layout, master):
        bg = owner._element.find("p:cSld/p:bg", _NS)
        if bg is None:
            continue
        hx = _color_hex(bg, master)
        if hx:
            return _luminance(hx) < 0.5
    return False


def text_colors(slide) -> tuple:
    """(normal text colour, removed-line red) for the slide's background: black-ish on light, white on dark."""
    return (WHITE, DIFF_RED_ON_DARK) if background_is_dark(slide) else (DARK, DIFF_RED)


def _diff_cell(cell, text: str, size: float, color, bold: bool = False, mono: bool = True, fill=None,
               italic: bool = False, align=PP_ALIGN.LEFT):
    if fill is None:
        cell.fill.background()                    # no fill: the slide's background shows through
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.04)
    cell.margin_top = cell.margin_bottom = Inches(0.015)
    cell.vertical_anchor = MSO_ANCHOR.TOP
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ""
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Courier New" if mono else "Calibri"


def add_diff_slides(prs, title: str, groups: Sequence, *, subtitle: Optional[str] = None, max_slides: int = 8,
                    more=lambda n: f"… {n} more lines", note: str = "", empty_text: str = "no changes",
                    font_size: float = 11, max_height_in: float = 5.7, width_in: float = 12.33,
                    gutter_in: float = 0.3) -> int:
    """Line-diff slides: ``groups`` [(group title, [block, ...])], a block being the lines that stay together
    ([('removed' | 'added', text), ...] - a pair is one block, never split over two slides).  Removed lines in red
    text with a "−" gutter, added lines in the slide's normal text colour with "+", group headers in a light grey
    band (repeated "(cont.)" on the next slide), full lines wrapped inside the row in Courier New.  At most
    ``max_slides`` slides; what does not fit becomes a last row ``more(n lines)``.  Returns the slide count."""
    text_w = width_in - gutter_in
    band_size = font_size

    def h_line(text: str, mono=True, size=font_size) -> float:
        return _est_lines(text, text_w if mono else width_in, size, mono) * size * 1.2 / 72.0 + 0.07

    def h_block(b) -> float:
        return sum(h_line(t) for _, t in b)

    head_h = lambda t: _est_lines(t, width_in, band_size, False) * band_size * 1.2 / 72.0 + 0.08   # noqa: E731
    def paginate(avail: float) -> List[List]:
        pages: List[List] = []
        cur: List = []
        used = 0.0
        if note:
            cur.append(("note", note))
            used += h_line(note, mono=False)
        if not any(blocks for _, blocks in groups):
            cur.append(("note", empty_text))
        for gtitle, blocks in groups:
            if not blocks:
                continue
            hh = head_h(gtitle)
            if cur and used + hh + h_block(blocks[0]) > avail:      # a header never alone at the bottom
                pages.append(cur)
                cur, used = [], 0.0
            cur.append(("head", gtitle))
            used += hh
            for b in blocks:                                         # a block (a pair) never split
                hb = h_block(b)
                if cur and used + hb > avail and not (cur[-1][0] == "head"):
                    pages.append(cur)
                    cont = f"{gtitle} (cont.)"
                    cur, used = [("head", cont)], head_h(cont)
                cur.append(("block", b))
                used += hb
        if cur:
            pages.append(cur)
        return pages

    avail = max_height_in
    pages = paginate(avail)
    if 1 < len(pages) <= max_slides:                  # balance the pages (no near-empty last slide)
        total = (sum(head_h(t) + sum(h_block(b) for b in blocks) for t, blocks in groups if blocks)
                 + (h_line(note, mono=False) if note else 0.0))
        balanced = paginate(min(avail, total / len(pages) * 1.1))
        if len(balanced) == len(pages):
            pages = balanced
    total_lines = sum(len(b) for _, blocks in groups for b in blocks)
    if len(pages) > max_slides:
        pages = pages[:max_slides]
        last = pages[-1]

        def page_h(pg):
            return sum(head_h(x[1]) if x[0] == "head" else h_line(x[1], mono=False) if x[0] == "note"
                       else h_block(x[1]) for x in pg)
        shown = lambda: sum(len(x[1]) for pg in pages for x in pg if x[0] == "block")   # noqa: E731
        while last and (page_h(last) + h_line(more(total_lines - shown()), mono=False) > avail
                        or last[-1][0] == "head"):
            last.pop()
        last.append(("more", more(total_lines - shown())))
    total_w = width_in
    left = max(0.2, (13.333 - total_w) / 2)
    top = 1.55 if subtitle else 1.35
    for pi, page in enumerate(pages):
        s = _blank(prs)
        add_title(s, title if len(pages) == 1 else f"{title}  ({pi + 1}/{len(pages)})", subtitle)
        normal, red = text_colors(s)
        rows = []                                  # (kind, change, text)
        for x in page:
            if x[0] == "block":
                rows += [("line", ch, t) for ch, t in x[1]]
            else:
                rows.append((x[0], "", x[1]))
        heights = [head_h(t) if k == "head" else h_line(t) if k == "line" else h_line(t, mono=False)
                   for k, _, t in rows]
        shape = s.shapes.add_table(len(rows), 2, Inches(left), Inches(top), Inches(total_w), Inches(sum(heights)))
        tbl = shape.table
        tbl.first_row = False
        tbl.horz_banding = False
        tbl.columns[0].width = Inches(gutter_in)
        tbl.columns[1].width = Inches(text_w)
        for i, ((kind, ch, text), h) in enumerate(zip(rows, heights)):
            tbl.rows[i].height = Inches(h)
            if kind == "line":
                col = red if ch == "removed" else normal
                _diff_cell(tbl.cell(i, 0), "−" if ch == "removed" else "+", font_size, col, bold=True,
                           align=PP_ALIGN.CENTER)
                _diff_cell(tbl.cell(i, 1), text, font_size, col)
                continue
            tbl.cell(i, 0).merge(tbl.cell(i, 1))
            if kind == "head":
                _diff_cell(tbl.cell(i, 0), text, band_size, DARK, bold=True, mono=False, fill=DIFF_BAND)
            else:
                _diff_cell(tbl.cell(i, 0), text, font_size, MUTED if not background_is_dark(s) else normal,
                           mono=False, italic=True)
    return len(pages)
