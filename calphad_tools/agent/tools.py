"""Tools the LLM can call.  Every tool is an ordinary Python function executed in
this process (so it inherits exactly this process's permissions / sandbox).

Each tool returns ``(text, images)``; ``images`` are file paths shown to the
model (if vision is on) and to the user.
"""
from __future__ import annotations

import contextlib
import glob
import io
import json
import os
import re
import time
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

from .llm import ToolSpec


# Extra tool modules; each defines register(toolbox) and adds tools with toolbox.add(...).
PLUGIN_MODULES = ["calphad_tools.figures", "calphad_tools.figdata", "calphad_tools.figfit",
                  "calphad_tools.figops_extra", "calphad_tools.figcompose", "calphad_tools.props",
                  "calphad_tools.mapping_tools", "calphad_tools.espei_plot_tools", "calphad_tools.digitize",
                  "calphad_tools.digitize_zpf",
                  "calphad_tools.espei_post", "calphad_tools.espei_data", "calphad_tools.tdb_checks",
                  "calphad_tools.provenance", "calphad_tools.codegen", "calphad_tools.pycalphad_docs",
                  "calphad_tools.jobs", "calphad_tools.espei_run", "calphad_tools.param_tools",
                  "calphad_tools.dbformats", "calphad_tools.pcache", "calphad_tools.compare_system",
                  "calphad_tools.recall", "calphad_tools.phase_inventory", "calphad_tools.system_coverage",
                  "calphad_tools.compatibility", "calphad_tools.staged_merge"]


# tools sent to the model with their full description and argument schema; every other tool is sent as one line
# plus its argument names, and the model reads the details with tool_help (once, when it first needs the tool)
CORE_TOOLS = ("list_files", "find_databases", "tdb_info", "calculate_equilibrium", "map_phase_diagram", "make_plot",
              "edit_figure", "describe_figure", "view_image", "inspect_data", "recent_results", "ask_user",
              "show_code", "tool_help", "list_calphad_capabilities")


def sig_explicit(fn) -> bool:
    """True if the function names its parameters (not just **kwargs)."""
    import inspect
    try:
        return any(p.kind in (p.POSITIONAL_OR_KEYWORD, p.KEYWORD_ONLY) for p in inspect.signature(fn).parameters.values())
    except (TypeError, ValueError):
        return False


def _summary(desc: str, n: int = 170) -> str:
    """First sentence of a tool description (the compact form the model sees)."""
    m = re.match(r"(.+?[.!?])(\s|$)", desc.strip(), re.S)
    first = (m.group(1) if m else desc.strip()).replace("\n", " ")
    return first if len(first) <= n else first[: n - 3].rsplit(" ", 1)[0] + "..."


def _pptx_text(path: str, limit: int = 6000) -> str:
    """Slide titles and table/text content of a PowerPoint (compact; never the raw zip bytes)."""
    from pptx import Presentation
    prs = Presentation(path)
    out = [f"{Path(path).name}: {len(prs.slides)} slides"]
    for i, sl in enumerate(prs.slides, 1):
        texts, tables = [], []
        for sh in sl.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [" | ".join(c.text.strip().replace("\n", " ") for c in r.cells) for r in sh.table.rows]
                tables.append(f"  table {len(rows)} rows x {len(sh.table.columns)} cols: " + " / ".join(rows[:4])
                              + (" / ..." if len(rows) > 4 else ""))
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip().replace("\n", " "))
        out.append(f"slide {i}: " + (texts[0][:120] if texts else "(no title)"))
        out += tables
        if sum(len(x) for x in out) > limit:
            out.append(f"... ({len(prs.slides) - i} more slides)")
            break
    return "\n".join(out)


class UI:
    """Interface the agent uses to talk to whoever is at the keyboard (terminal or web)."""

    def ask(self, question: str, options: Optional[List[str]] = None) -> str:
        raise NotImplementedError

    def confirm(self, title: str, body: str) -> bool:
        raise NotImplementedError

    def progress(self, msg: str) -> None:
        pass

    def notice(self, kind: str, text: str) -> None:
        """Informational message (e.g. a safety-review verdict)."""
        self.progress(text)


@dataclass
class Tool:
    spec: ToolSpec
    fn: Callable[..., Tuple[str, List[str]]]


def _obj(props: Dict, required: List[str]) -> Dict:
    return {"type": "object", "properties": props, "required": required, "additionalProperties": False}


S = {"type": "string"}
N = {"type": "number"}
I = {"type": "integer"}
B = {"type": "boolean"}
STRS = {"type": "array", "items": {"type": "string"}}
NUMS = {"type": "array", "items": {"type": "number"}}
UNIT = {"type": "string", "enum": ["K", "C"], "description": "Temperature unit of the given temperatures"}


class ToolBox:
    def __init__(self, workdir: str, output_dir: str, ui: UI, python_tool: str = "ask",
                 max_chars: int = 20000, llm_cfg=None, interactive: bool = True, diagram_style: str = "shaded"):
        self.workdir = Path(workdir).resolve()
        self.out = (self.workdir / output_dir).resolve()
        self.out.mkdir(parents=True, exist_ok=True)
        self.ui = ui
        self.python_tool = python_tool
        self.max_chars = max_chars
        self.merge_decisions: Dict[str, Dict[str, str]] = {}
        self.py_ns: Dict[str, Any] = {}
        self.py_history: List[str] = []           # code that has actually been executed
        self.llm_cfg = llm_cfg                    # needed for auto_verify reviews
        self.interactive = interactive            # is a human available to confirm?
        from calphad_tools.diagram_style import check_style
        self.diagram_style = check_style(diagram_style)   # phase-diagram style when a call gives none
        self.tools: Dict[str, Tool] = {}
        self.call_log: List[Dict[str, Any]] = []      # every tool call this session (provenance / export)
        self.call_listeners: List[Callable[[Dict[str, Any]], None]] = []
        self._register_all()

    # ------------------------------------------------------------------ utils
    def p(self, path: str) -> str:
        """Resolve a user/LLM supplied path against the working directory."""
        pp = Path(os.path.expanduser(path))
        return str(pp if pp.is_absolute() else (self.workdir / pp))

    def outp(self, name: Optional[str], default: str) -> str:
        if name:
            pp = Path(os.path.expanduser(name))
            if not pp.is_absolute():
                pp = (self.out / pp) if len(pp.parts) == 1 else (self.workdir / pp)
        else:
            pp = self.out / default
        pp.parent.mkdir(parents=True, exist_ok=True)
        return str(pp)

    def outdir(self, output_dir: Optional[str], default: str) -> Path:
        d = Path(self.p(output_dir)) if output_dir else self.out / default
        d.mkdir(parents=True, exist_ok=True)
        return d

    def rel(self, path: str) -> str:
        try:
            return str(Path(path).resolve().relative_to(self.workdir))
        except ValueError:
            return str(path)

    _SKIP_DIRS = {".venv", "venv", "__pycache__", "node_modules", "site-packages", ".git"}
    _DB_EXTS = (".tdb", ".dat")

    def _walk_files(self):
        for root, dirs, files in os.walk(self.workdir):
            dirs[:] = sorted(d for d in dirs if not d.startswith(".") and d not in self._SKIP_DIRS)
            for f in sorted(files):
                yield os.path.join(root, f)

    def _find_tdb(self, name: str) -> str:
        """Resolve a file name EXACTLY: the path itself, else a unique exact basename anywhere under the workdir,
        else a unique case-insensitive basename, else (no extension given) a unique .tdb/.dat with that stem.
        Never substitutes a different file: a miss raises with 'did you mean' candidates."""
        path = self.p(name)
        if os.path.exists(path):
            return path
        base = Path(name).name
        files = list(self._walk_files())
        hits = [h for h in files if Path(h).name == base]
        if not hits:
            hits = [h for h in files if Path(h).name.lower() == base.lower()]
        if not hits and Path(base).suffix.lower() not in self._DB_EXTS:
            dbs = [h for h in files if Path(h).suffix.lower() in self._DB_EXTS]
            hits = [h for h in dbs if Path(h).stem == base] or [h for h in dbs if Path(h).stem.lower() == base.lower()]
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1:
            raise FileNotFoundError(f"{name} is ambiguous: {[self.rel(h) for h in hits]} - give the relative path")
        import difflib
        suf = Path(base).suffix.lower()
        dbs = ([h for h in files if Path(h).suffix.lower() in self._DB_EXTS] if suf in ("",) + self._DB_EXTS
               else [h for h in files if Path(h).suffix.lower() == suf])
        by_name: Dict[str, List[str]] = {}
        for h in dbs:
            by_name.setdefault(Path(h).name.lower(), []).append(h)
            by_name.setdefault(Path(h).stem.lower(), []).append(h)
        close = difflib.get_close_matches(base.lower(), list(by_name), n=6, cutoff=0.6)
        cands: List[str] = []
        for c in close:
            for h in by_name[c]:
                if self.rel(h) not in cands:
                    cands.append(self.rel(h))
        msg = f"{name} not found under {self.workdir} (exact names only; no other file was substituted)."
        if cands:
            msg += f" Did you mean one of: {', '.join(cands[:5])}? Ask the user which one they meant."
        raise FileNotFoundError(msg)

    def add(self, name: str, description: str, params: Dict, required: List[str]):
        def deco(fn):
            self.tools[name] = Tool(ToolSpec(name, description, _obj(params, required)), fn)
            return fn
        return deco

    def specs(self) -> List[ToolSpec]:
        return [t.spec for t in self.tools.values()]

    def model_specs(self, mode: str = "compact", core=CORE_TOOLS) -> List[ToolSpec]:
        """The tool list sent to the model.  "compact": the core tools in full, every other tool as a one-line
        summary + its argument names (details on demand with tool_help) - about 70 % fewer tokens on every call,
        and the list never changes during a conversation (prompt caching).  "full": everything in full."""
        if str(mode).lower() == "full":
            return self.specs()
        out = []
        for t in self.tools.values():
            sp = t.spec
            if sp.name in core:
                out.append(sp)
                continue
            props = (sp.parameters or {}).get("properties", {}) or {}
            req = set((sp.parameters or {}).get("required", []) or [])
            args = ", ".join(n + ("*" if n in req else "") for n in props) or "none"
            out.append(ToolSpec(sp.name, f"{_summary(sp.description, 100)} Args: {args} [tool_help]",
                                {"type": "object"}))
        return out

    def help_text(self, name: str) -> str:
        t = self.tools.get(name)
        if t is None:
            import difflib
            close = difflib.get_close_matches(name, list(self.tools), n=3)
            return f"{name}: no such tool" + (f" (did you mean {', '.join(close)}?)" if close else "")
        sp = t.spec
        params = sp.parameters or {}
        return (f"## {sp.name}\n{sp.description}\nrequired: {params.get('required', [])}\narguments: "
                + json.dumps(params.get("properties", {}), separators=(",", ":"), default=str))

    def _check_args(self, name: str, args: Dict) -> Optional[str]:
        """Unknown or missing arguments, caught before the tool runs (compact tool descriptions make guesses more
        likely): the error carries the tool's full help, so one retry fixes it."""
        t = self.tools.get(name)
        if t is None or not isinstance(args, dict):
            return None
        params = t.spec.parameters or {}
        props = params.get("properties") or {}
        if params.get("additionalProperties", False) is not False and not props:
            return None
        unknown = [k for k in args if props and k not in props]
        # "required" in a schema tells the model what to give; only arguments the tool really cannot do without
        # (no default in the code) are refused here
        import inspect
        try:
            sig = inspect.signature(t.fn).parameters
            has_default = {n for n, p_ in sig.items() if p_.default is not inspect.Parameter.empty}
            takes_kwargs = any(p_.kind == p_.VAR_KEYWORD for p_ in sig.values())
        except (TypeError, ValueError):
            has_default, takes_kwargs = set(), True
        missing = [k for k in params.get("required", []) or [] if k not in args and k not in has_default]
        if takes_kwargs and not sig_explicit(t.fn):
            missing = []                                  # **kwargs-only tools validate their own arguments
        if not unknown and not missing:
            return None
        why = "; ".join(x for x in (f"unknown argument(s) {unknown}" if unknown else "",
                                    f"missing required {missing}" if missing else "") if x)
        return f"bad arguments for {name}: {why}.\n{self.help_text(name)}"

    def call(self, name: str, args: Dict) -> Tuple[str, List[str], bool]:
        t0 = time.time()
        refs: Dict[str, str] = {}
        try:                                   # '@last' / '@call:N' -> real paths (logged resolved)
            from ..recall import resolve_refs
            args, refs = resolve_refs(self, args)
        except ValueError as e:
            text, images, err = str(e), [], True
        else:
            text, images, err = self._call(name, args)
            if refs:
                text = "(" + ", ".join(f"{k} = {v}" for k, v in refs.items()) + ")\n" + text
        rec = {"time": time.strftime("%Y-%m-%dT%H:%M:%S", time.localtime(t0)), "tool": name, "args": args,
               "ok": not err, "seconds": round(time.time() - t0, 3), "images": list(images),
               "result_preview": text[:500],
               # full text kept for failures (tracebacks) and short results - problem reports need it
               "result_full": text if (err or len(text) <= 20000) else None}
        self.call_log.append(rec)
        for fn in self.call_listeners:
            try:
                fn(rec)
            except Exception:  # noqa: BLE001 - listeners must never break a tool call
                pass
        return text, images, err

    def _call(self, name: str, args: Dict) -> Tuple[str, List[str], bool]:
        t = self.tools.get(name)
        if t is None:
            return f"unknown tool {name!r}", [], True
        bad = self._check_args(name, args)
        if bad:
            return bad, [], True
        try:
            text, images = t.fn(**args)
            if len(text) > self.max_chars:
                # keep the conversation small: the full reply goes to a file the model can read in parts
                full = self.out / ".tool_outputs" / f"{len(self.call_log) + 1:04d}_{name}.txt"
                try:
                    full.parent.mkdir(parents=True, exist_ok=True)
                    full.write_text(text)
                    where = (f"full reply ({len(text):,} chars) saved to {self.rel(str(full))} - read_document it "
                             f"with start/length only if you need the rest")
                except OSError:
                    where = f"{len(text) - self.max_chars:,} more chars not shown"
                text = text[: self.max_chars] + f"\n...[truncated: {where}]"
            return text, images, False
        except TypeError as e:
            msg = str(e)
            if "argument" in msg and ("unexpected" in msg or "missing" in msg or "positional" in msg):
                return f"bad arguments for {name}: {e}\n{self.help_text(name)}", [], True
            tb = traceback.format_exc(limit=3)
            return f"TypeError: {e}\n{tb[-1500:]}", [], True
        except Exception as e:  # noqa: BLE001 - report every failure back to the model
            tb = traceback.format_exc(limit=3)
            return f"{type(e).__name__}: {e}\n{tb[-1500:]}", [], True

    @staticmethod
    def j(obj) -> str:
        return json.dumps(obj, indent=1, default=str)

    # ------------------------------------------------------------------ tools
    def _register_all(self):  # noqa: C901 - one long list of small tool definitions
        from .. import calc, compare, diagrams, espei_flux, expdata, merge

        @self.add("tool_help", "Full description and argument schema of one or more tools (give every tool you are "
                  "about to use in ONE call). Tools listed with 'Args: ...' need this before their first use.",
                  {"names": {"type": "array", "items": {"type": "string"}}}, ["names"])
        def tool_help(names):
            names = [names] if isinstance(names, str) else list(names or [])
            return "\n\n".join(self.help_text(str(n)) for n in names[:12]) or "give tool names", []

        @self.add("list_files", "List files under the working directory (glob pattern, recursive). Use to find "
                  "TDB/POP/EXP/PDF/CSV files before using them.",
                  {"pattern": {"type": "string", "description": "glob, e.g. '**/*.tdb' or '*.pdf'"}}, ["pattern"])
        def list_files(pattern: str):
            hits = sorted(glob.glob(str(self.workdir / pattern), recursive=True))
            hits = [h for h in hits if "/.venv/" not in h and "/." not in h[len(str(self.workdir)):]]
            return self.j([self.rel(h) for h in hits[:300]]) + ("" if len(hits) <= 300 else f"\n({len(hits)} total)"), []

        @self.add("tdb_info", "Summarize a TDB: elements, phases with sublattice models, which subsystems "
                  "have parameters, header comments.", {"tdb": S}, ["tdb"])
        def tdb_info(tdb: str):
            return self.j(calc.tdb_info(self._find_tdb(tdb))), []

        @self.add("compare_tdbs", "Compare two TDB files parameter-by-parameter (textual + numerical over "
                  "298-3000 K) and write a PowerPoint with the comparison tables plus a CSV. For 3+ databases or "
                  "one subsystem use compare_system.",
                  {"tdb_a": S, "tdb_b": S, "output_pptx": {"type": "string", "description": "optional file name"},
                   "include_identical": {"type": "boolean", "description": "also list identical items (default false)"}},
                  ["tdb_a", "tdb_b"])
        def compare_tdbs(tdb_a: str, tdb_b: str, output_pptx: Optional[str] = None, include_identical: bool = False):
            a, b = self._find_tdb(tdb_a), self._find_tdb(tdb_b)
            out = self.outp(output_pptx, f"compare_{Path(a).stem}_vs_{Path(b).stem}.pptx")
            r = compare.run(a, b, out, include_identical=include_identical)
            return f"PowerPoint: {self.rel(r['pptx'])}\nCSV: {self.rel(r['csv'])}\n\n{r['summary']}", []

        ORDER = {**STRS, "description": "orientation preference, e.g. [\"FE\", \"CR\", \"NI\"]: earlier elements "
                                        "go left / bottom-left in every subsystem (default alphabetical)"}
        ORIENT = {"type": "object"}
        from calphad_tools.mapping_tools import CHECK_SCHEMA as CHECK, STYLE_SCHEMA as STYLE

        @self.add("plot_all_binaries", "Compute every binary phase diagram in a TDB (T range in °C, 1 atm), "
                  "in parallel, and write a PowerPoint (grid images, 6 per slide).",
                  {"tdb": S, "tmin_c": {"type": "number", "description": "default 100"},
                   "tmax_c": {"type": "number", "description": "default 3000"},
                   "elements": {**STRS, "description": "optional subset of elements"},
                   "element_order": ORDER, "orientation": {**ORIENT, "description": "per-system override, e.g. "
                                                           "{\"CR-FE\": [\"FE\", \"CR\"]} ([left, right])"},
                   "output_pptx": S, "individual_slides": B, "style": STYLE, "check_mapping": CHECK},
                  ["tdb"])
        def plot_all_binaries(tdb: str, tmin_c: float = 100, tmax_c: float = 3000, elements=None,
                              output_pptx: Optional[str] = None, individual_slides: bool = False,
                              element_order=None, orientation=None, style=None, check_mapping="auto"):
            t = self._find_tdb(tdb)
            out = self.outp(output_pptx, f"binaries_{Path(t).stem}.pptx")
            r = diagrams.all_binaries(t, out, tmin_c=tmin_c, tmax_c=tmax_c, elements=elements,
                                      individual_slides=individual_slides, progress=self.ui.progress,
                                      element_order=element_order, orientation=orientation,
                                      style=style or self.diagram_style, check_mapping=check_mapping)
            txt = f"PowerPoint: {self.rel(r['pptx'])}\n" + "\n".join(
                f"  {s['name']}: {'ok' if s['ok'] else 'FAILED ' + s['error']} ({s['seconds']} s)"
                + (f"; mapping check: {s['note']}" if s.get("note") else "") for s in r["systems"])
            return txt, r["grid_images"][:2]

        @self.add("plot_all_ternaries", "Compute isothermal sections of every ternary subsystem of a TDB "
                  "(3+ elements) at the given temperatures (°C), in parallel, and write a PowerPoint.",
                  {"tdb": S, "temperatures_c": NUMS, "elements": STRS, "output_pptx": S,
                   "element_order": ORDER, "orientation": {**ORIENT, "description": "per-system override, e.g. "
                                                           "{\"CR-FE-NI\": {\"left\": \"NI\", \"right\": \"FE\", "
                                                           "\"top\": \"CR\"}}"},
                   "individual_slides": B, "style": STYLE, "check_mapping": CHECK}, ["tdb", "temperatures_c"])
        def plot_all_ternaries(tdb: str, temperatures_c, elements=None, output_pptx: Optional[str] = None,
                               individual_slides: bool = False, element_order=None, orientation=None, style=None,
                               check_mapping="auto"):
            t = self._find_tdb(tdb)
            out = self.outp(output_pptx, f"ternaries_{Path(t).stem}.pptx")
            r = diagrams.all_ternaries(t, temperatures_c, out, elements=elements,
                                       individual_slides=individual_slides, progress=self.ui.progress,
                                       element_order=element_order, orientation=orientation,
                                       style=style or self.diagram_style, check_mapping=check_mapping)
            txt = f"PowerPoint: {self.rel(r['pptx'])}\n" + "\n".join(
                f"  {s['name']}: {'ok' if s['ok'] else 'FAILED ' + s['error']} ({s['seconds']} s)"
                + (f"; mapping check: {s['note']}" if s.get("note") else "") for s in r["systems"])
            return txt, r["grid_images"][:2]

        # (single binary/ternary/isopleth diagrams, step and G-x plots: see mapping_tools.py and props.py)

        @self.add("calculate_equilibrium", "Equilibrium at specific conditions (small grids only, <=60 points). "
                  "Returns stable phases, their fractions and compositions.",
                  {"tdb": S, "elements": STRS,
                   "conditions": {"type": "object", "description": "e.g. {\"T\": [1000, 1200], \"X(CR)\": 0.2, "
                                  "\"P\": 101325}; T in temperature_unit"},
                   "temperature_unit": UNIT, "phases": STRS}, ["tdb", "elements", "conditions", "temperature_unit"])
        def calculate_equilibrium(tdb, elements, conditions, temperature_unit="K", phases=None):
            r = calc.equilibrium_table(self._find_tdb(tdb), elements, conditions, unit=temperature_unit, phases=phases)
            return self.j(r), []

        @self.add("merge_tdbs", "Merge 2+ TDB files (earlier = higher priority). Obvious differences are resolved "
                  "automatically; other conflicts are returned with ids and options. Present pending conflicts to "
                  "the user and call again with their answers in `decisions` (conflict id -> option key; a "
                  "'system:A-B' key answers every parameter conflict of that subsystem). Never invent answers "
                  "for non-obvious conflicts unless the user said to (then use policy 'first' or 'last').",
                  {"tdbs": STRS, "output": S,
                   "decisions": {"type": "object", "additionalProperties": {"type": "string"}},
                   "policy": {"type": "string", "enum": ["ask", "first", "last"]}},
                  ["tdbs", "output"])
        def merge_tdbs(tdbs, output, decisions=None, policy: str = "ask"):
            paths = [self._find_tdb(t) for t in tdbs]
            key = "|".join(paths)
            dec = {**self.merge_decisions.get(key, {}), **(decisions or {})}
            self.merge_decisions[key] = dec
            out = self.outp(output, "merged.tdb")
            r = merge.merge(paths, out, decisions=dec, policy=policy)
            if not r["written"]:
                return ("NOT WRITTEN - decisions needed. Ask the user (group parameter conflicts by subsystem).\n"
                        + r["report"]), []
            return (f"written {self.rel(out)}\nvalidation: {self.j(r['validation'])}\n\n{r['report']}"), []

        # ------------------------------------------------ experimental data
        @self.add("read_document", "Read text from a PDF (optionally page range like '1-3,7'), a text file "
                  "(POP/EXP/TDB/TXT...) or a PowerPoint (slide titles + table text). start/length read a long text "
                  "file in parts. Tables (csv/xlsx) -> inspect_data; images -> view_image. Don't re-read decks the "
                  "tools wrote: their reply already summarises them.",
                  {"path": S, "pages": S, "start": I, "length": I}, ["path"])
        def read_document(path: str, pages: Optional[str] = None, start: int = 0, length: Optional[int] = None):
            fp = self.p(path)
            suf = Path(fp).suffix.lower()
            if suf == ".pdf":
                return expdata.pdf_text(fp, pages), []
            if suf in (".pptx", ".potx"):
                return _pptx_text(fp), []
            if suf in (".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".parquet", ".nc"):
                raise ValueError(f"{Path(fp).name} is a table: use inspect_data (columns, sizes, a preview)")
            if suf in (".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff", ".bmp", ".svg"):
                raise ValueError(f"{Path(fp).name} is an image: use view_image to look at it")
            raw = Path(fp).read_bytes()
            if b"\x00" in raw[:4096] or raw[:4] in (b"PK\x03\x04", b"%PDF"):
                raise ValueError(f"{Path(fp).name} is a binary file ({len(raw):,} bytes), not text - not read "
                                 f"(its bytes would only fill the conversation)")
            text = raw.decode("utf-8", errors="replace")
            start = max(0, int(start or 0))
            part = text[start:start + int(length)] if length else text[start:]
            head = (f"[characters {start:,}-{start + len(part):,} of {len(text):,}]\n"
                    if start or (length and len(part) < len(text) - start) else "")
            return head + part, []

        @self.add("view_pdf_page", "Render one PDF page (1-based) to an image so you can read tables/figures.",
                  {"path": S, "page": I}, ["path", "page"])
        def view_pdf_page(path: str, page: int):
            fp = self.p(path)
            out = self.outp(None, f"pdf_pages/{Path(fp).stem}_p{page}.png")
            png = expdata.render_pdf_page(fp, page, out)
            return f"rendered page {page} -> {self.rel(png)}", [png]

        @self.add("view_image", "Look at an image file (e.g. to check a plot you made).", {"path": S}, ["path"])
        def view_image(path: str):
            return f"image {path}", [self.p(path)]

        @self.add("parse_dataplot_exp", "Parse a Thermo-Calc DATAPLOT .exp file: axis labels, and for each DATASET "
                  "its id, title, comments and points. Use this to decide which phases each dataset represents.",
                  {"path": S}, ["path"])
        def parse_dataplot_exp(path: str):
            info = expdata.parse_exp(self.p(path))
            brief = {"prologue": info["prologue"], "axes_guess": info["axes_guess"], "datasets": [
                {"id": d["id"], "title": d["title"], "comments": d["comments"][:8], "n_points": len(d["points"]),
                 "first_points": d["points"][:5], "strings": d["strings"][:5]} for d in info["datasets"]]}
            return self.j(brief), []

        @self.add("convert_dataplot_exp_to_zpf", "Convert DATAPLOT datasets (binary x vs T) to ESPEI ZPF JSON. "
                  "dataset_map: {dataset_id: {\"phases\": [phase in equilibrium...], \"measured_phase\": phase whose "
                  "composition the x value is, \"x_component\": EL}}. Unmapped datasets are skipped.",
                  {"path": S, "components": STRS, "dataset_map": {"type": "object"}, "output_dir": S,
                   "x_unit": {"type": "string", "enum": ["mole_fraction", "at%", "wt%", "wt_fraction"]},
                   "y_unit": {"type": "string", "enum": ["K", "C"]}, "reference": S},
                  ["path", "components", "dataset_map"])
        def convert_dataplot_exp_to_zpf(path, components, dataset_map, output_dir=None, x_unit=None, y_unit=None,
                                        reference: str = ""):
            parsed = expdata.parse_exp(self.p(path))
            dss = expdata.exp_to_zpf(parsed, components, dataset_map, reference=reference, x_unit=x_unit,
                                     y_unit=y_unit)
            od = self.outdir(output_dir, "datasets")
            files = [expdata.write_dataset(d, str(od / f"{Path(path).stem}_ds{i + 1}_zpf.json"))
                     for i, d in enumerate(dss)]
            return f"wrote {len(files)} ESPEI ZPF datasets (validated):\n" + "\n".join(self.rel(f) for f in files), []

        @self.add("convert_pop_to_zpf", "Convert a Thermo-Calc POP/DOP experiment file to ESPEI ZPF datasets "
                  "(phase-boundary data only; thermochemical points (ACR, HMR, CP...) -> convert_pop_to_thermo). Give a tdb so "
                  "abbreviated phase names (LIQ, BCC...) resolve, or a phase_map.",
                  {"path": S, "components": STRS, "tdb": S,
                   "phase_map": {"type": "object", "additionalProperties": {"type": "string"}}, "output_dir": S},
                  ["path", "components"])
        def convert_pop_to_zpf(path, components, tdb=None, phase_map=None, output_dir=None):
            phases = sorted(calc.get_db(self._find_tdb(tdb)).phases) if tdb else None
            dss, notes = expdata.pop_to_zpf(self.p(path), components, tdb_phases=phases, phase_map=phase_map)
            od = self.outdir(output_dir, "datasets")
            files = [expdata.write_dataset(d, str(od / f"{Path(path).stem}_{i + 1}_zpf.json")) for i, d in enumerate(dss)]
            return (f"wrote {len(files)} datasets:\n" + "\n".join(self.rel(f) for f in files)
                    + ("\nnotes:\n" + "\n".join(notes) if notes else "")), []

        @self.add("convert_csv_to_zpf", "Convert a CSV (columns T_K or T_C, phase_1, X(EL)_1, phase_2, X(EL)_2, ..., "
                  "optional reference) to ESPEI ZPF datasets.",
                  {"path": S, "components": STRS, "output_dir": S,
                   "composition_unit": {"type": "string", "enum": ["mole_fraction", "at%", "wt%", "wt_fraction"]}},
                  ["path", "components"])
        def convert_csv_to_zpf(path, components, output_dir=None, composition_unit="mole_fraction"):
            dss = expdata.csv_to_zpf(self.p(path), components, composition_unit=composition_unit)
            od = self.outdir(output_dir, "datasets")
            files = [expdata.write_dataset(d, str(od / f"{Path(path).stem}_{i + 1}_zpf.json")) for i, d in enumerate(dss)]
            return f"wrote {len(files)} datasets:\n" + "\n".join(self.rel(f) for f in files), []

        @self.add("write_zpf_dataset", "Write one ESPEI ZPF dataset from data you extracted (e.g. from a paper's "
                  "table). values[i] lists the phase regions in equilibrium at temperatures[i]: "
                  "[[phase, [independent components], [mole fractions or null]], ...]. Mole fractions only "
                  "(convert at.%/wt.% first). Validated with ESPEI before writing.",
                  {"components": STRS, "temperatures": NUMS, "temperature_unit": UNIT,
                   "values": {"type": "array", "items": {"type": "array"}}, "reference": S, "comment": S,
                   "output": S}, ["components", "temperatures", "temperature_unit", "values", "reference", "output"])
        def write_zpf_dataset(components, temperatures, temperature_unit, values, reference, output, comment=""):
            ds = expdata.make_zpf_dataset(components, temperatures, values, reference=reference, comment=comment,
                                          temperature_unit=temperature_unit)
            out = self.outp(output if output.endswith(".json") else output + ".json", "zpf.json")
            expdata.write_dataset(ds, out)
            return f"wrote {self.rel(out)} ({len(ds['values'])} points, phases {ds['phases']})", []

        @self.add("check_espei_datasets", "Validate every ESPEI dataset JSON in a directory with ESPEI's checker.",
                  {"directory": S}, ["directory"])
        def check_espei_datasets(directory: str):
            return self.j(expdata.check_datasets_dir(self.p(directory))), []

        # ------------------------------------------------ ESPEI / Flux
        @self.add("make_phase_models", "Create an ESPEI phase_models.json for a set of elements from a TDB's phases.",
                  {"tdb": S, "elements": STRS, "output": S, "phases": STRS}, ["tdb", "elements", "output"])
        def make_phase_models(tdb, elements, output, phases=None):
            out = self.outp(output, "phase_models.json")
            r = espei_flux.phase_models_from_tdb(self._find_tdb(tdb), elements, out, phases=phases)
            return f"wrote {self.rel(out)}\n{self.j(r)}", []

        run_schema = {"type": "object", "properties": {
            "name": S, "phase_models": S, "datasets": S, "input_db": S, "generate_parameters": B,
            "ref_state": S, "iterations": I, "save_interval": I, "chains_per_parameter": I,
            "chain_std_deviation": N, "cores": I, "symbols": STRS, "prior": {"type": "object"},
            "data_weights": {"type": "object"}}, "required": ["name", "phase_models", "datasets"]}
        flux_schema = {"type": "object", "properties": {
            "nodes": I, "cores_per_node": I, "time_limit": S, "queue": S, "bank": S, "mode":
                {"type": "string", "enum": ["per-job", "bundle"]}, "setup": STRS, "gen_time_limit": S}}

        @self.add("generate_espei_flux_batch", "Write ESPEI YAML inputs and Flux batch scripts for one or more ESPEI "
                  "runs (per-job: gen + multi-node dask MCMC chained with afterok; bundle: many runs in one "
                  "allocation). Scripts are syntax-checked and YAML validated against ESPEI's schema.",
                  {"runs": {"type": "array", "items": run_schema}, "flux": flux_schema, "output_dir": S},
                  ["runs", "flux", "output_dir"])
        def generate_espei_flux_batch(runs, flux, output_dir):
            rs = [espei_flux.Run(**{**r, "phase_models": self.p(r["phase_models"]), "datasets": self.p(r["datasets"]),
                                    **({"input_db": self.p(r["input_db"])} if r.get("input_db") else {})}) for r in runs]
            fs = espei_flux.FluxSettings(**flux)
            od = self.p(output_dir) if os.path.isabs(output_dir) or "/" in output_dir else str(self.out / output_dir)
            r = espei_flux.generate(rs, fs, od)
            return self.j({**r, "files": [self.rel(f) for f in r["files"]]}), []

        # ------------------------------------------------ plug-in tool modules
        import importlib
        for mod in PLUGIN_MODULES:
            try:
                m = importlib.import_module(mod)
            except ModuleNotFoundError as e:
                if e.name != mod:
                    raise
                continue
            reg = getattr(m, "register", None)     # op-only extension modules need no tools
            if reg is not None:
                reg(self)

        # ------------------------------------------------ interaction / python
        if isinstance(self.ui, UI):
            @self.add("ask_user", "Ask the user a question and wait for the answer (use for decisions only they "
                      "can make, e.g. merge conflicts or ambiguous data mapping).",
                      {"question": S, "options": STRS}, ["question"])
            def ask_user(question: str, options=None):
                return self.ui.ask(question, options), []

        if self.python_tool != "off":
            @self.add("run_python", "Run Python code in a persistent namespace inside this process for anything the "
                      "other tools don't cover. Preloaded: np, plt (Agg), pycalphad (Database, equilibrium, "
                      "calculate, binplot, ternplot, variables as v), load_db(path) (cached), OUT (output dir Path), "
                      "savefig(name) -> saves current matplotlib figure into OUT and returns the path. print() "
                      "output is returned. Figures saved with savefig are shown to you.",
                      {"code": S}, ["code"])
            def run_python(code: str):
                ok, why = self._python_gate(code)
                if not ok:
                    return why, []
                ns = self._python_ns()
                self.py_history.append(code)
                ns["_saved_figs"] = []
                buf = io.StringIO()
                t0 = time.time()
                with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
                    try:
                        exec(compile(code, "<llm>", "exec"), ns)  # noqa: S102 - this is the tool's purpose
                    except Exception:  # noqa: BLE001
                        traceback.print_exc(limit=4)
                out = buf.getvalue() or "(no output)"
                return f"{out}\n[{time.time() - t0:.1f} s]", list(ns["_saved_figs"])

    def _python_gate(self, code: str) -> Tuple[bool, str]:
        """Decide whether generated code may run, according to self.python_tool."""
        mode = self.python_tool
        if mode == "auto":
            return True, ""
        if mode == "auto_verify":
            from .review import review_code
            if self.llm_cfg is None:
                mode = "ask"
            else:
                self.ui.progress("independent safety review of the generated code…")
                r = review_code(self.llm_cfg, code, str(self.workdir), str(self.out),
                                log_path=str(self.out / "python_reviews.jsonl"),
                                prior_code=self._prior_definitions(code))
                self.ui.notice("review_ok" if r.auto_ok else "review_flag", r.text())
                if r.auto_ok:
                    return True, ""
                reason = r.text()
                if not self.interactive:
                    return False, ("Not run: the independent safety reviewer did not approve this code and no user "
                                   f"is available to override.\n{reason}\nRewrite it to avoid these concerns "
                                   "(or use the dedicated tools).")
                if self.ui.confirm("The safety reviewer did NOT auto-approve this code. Run it anyway?",
                                   f"{reason}\n\n----- code -----\n{code}"):
                    return True, ""
                return False, f"The user declined to run this code after review.\n{reason}"
        if mode == "ask":
            if self.ui.confirm("Run this Python code?", code):
                return True, ""
            return False, "The user declined to run this code."
        return False, f"python_tool mode {mode!r} does not allow running code"

    def _prior_definitions(self, code: str, limit: int = 6000) -> str:
        """Earlier executed snippets that define names the new code uses (the namespace persists)."""
        import ast
        try:
            used = {n.id for n in ast.walk(ast.parse(code)) if isinstance(n, ast.Name)}
        except SyntaxError:
            return ""
        out, total = [], 0
        for snip in reversed(self.py_history):
            try:
                tree = ast.parse(snip)
            except SyntaxError:
                continue
            defined = set()
            for n in ast.walk(tree):
                if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    defined.add(n.name)
                elif isinstance(n, ast.Name) and isinstance(n.ctx, ast.Store):
                    defined.add(n.id)
                elif isinstance(n, (ast.Import, ast.ImportFrom)):
                    defined |= {(a.asname or a.name).split(".")[0] for a in n.names}
            if defined & used and total + len(snip) < limit:
                out.append(snip)
                total += len(snip)
        return "\n# ---- next earlier snippet ----\n".join(reversed(out))

    def _python_ns(self) -> Dict[str, Any]:
        if not self.py_ns:
            import matplotlib.pyplot as plt
            import numpy as np
            import pycalphad
            from pycalphad import Database, binplot, calculate, equilibrium, ternplot, variables as v
            from .. import calc
            ns = self.py_ns
            ns.update(np=np, plt=plt, pycalphad=pycalphad, Database=Database, equilibrium=equilibrium,
                      calculate=calculate, binplot=binplot, ternplot=ternplot, v=v, OUT=self.out,
                      load_db=lambda p: calc.get_db(self._find_tdb(p)))

            def savefig(name: str, fig=None, dpi: int = 150):
                f = fig or plt.gcf()
                path = self.outp(name if name.endswith(".png") else name + ".png", "figure.png")
                f.savefig(path, dpi=dpi, bbox_inches="tight")
                plt.close(f)
                ns["_saved_figs"].append(path)
                return path
            ns["savefig"] = savefig
            os.chdir(self.workdir)
        return self.py_ns
