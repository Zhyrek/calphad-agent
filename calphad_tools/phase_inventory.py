"""phase_inventory: every phase of an element system, across all databases that contain those elements.

For each phase and each database: the sublattice model written in CALPHAD notation ((Cr,Fe)₁(Va)₃), the
constituents of every sublattice (inside the system, and what else the full database allows), the model type
(substitutional / interstitial solution, stoichiometric compound, ordered with a disordered part, ionic liquid,
...), magnetic and order-disorder settings, and how many parameters the phase has inside the system.

Phases are matched across databases by exact name; likely correspondences (BCC vs BCC_A2, SIGMA vs SIGMA_D8B) are
listed as candidates with structural evidence, never merged automatically (same rule as compare_system).
Writes an Excel workbook (one row per phase and database, plus a phase x database matrix), a CSV and optionally a
PowerPoint.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from . import compare_system as cs
from .phase_names import format_phase
from .tdbtables import PhaseView, Subsystem, key_constituents, phase_order

_SUB = str.maketrans("0123456789.", "₀₁₂₃₄₅₆₇₈₉.")
_ION = re.compile(r"[+-]\d*$")


def _sp(s: str, elements: set) -> str:
    """Display form of a constituent: element symbols title case (CR -> Cr, VA -> Va), species as written."""
    u = s.upper()
    if u == "VA":
        return "Va"
    if u in elements or re.fullmatch(r"[A-Z][A-Z]?", u):
        return u.capitalize()
    return s


def _ratio(r: float) -> str:
    return (f"{r:g}").translate(_SUB)


def sublattice_formula(site_ratios: Sequence[float], constituents: Sequence[Sequence[str]], elements: set) -> str:
    """(Cr,Fe)₁(Va)₃ style notation."""
    return "".join("(" + ",".join(_sp(s, elements) for s in sl) + ")" + _ratio(r)
                   for sl, r in zip(constituents, site_ratios))


def model_type(name: str, pv: PhaseView) -> str:
    cons = pv.constituents
    nsub = len(cons)
    ions = any(_ION.search(s) for sl in cons for s in sl if s != "VA")
    liquid = name.upper().startswith(("LIQ", "IONIC_LIQ")) or name.upper() in ("L", "LIQUID")
    if liquid:
        if ions or nsub == 2:
            return "ionic liquid (two-sublattice)"
        return "substitutional liquid" if nsub == 1 else f"liquid, {nsub} sublattices"
    if all(len(sl) == 1 for sl in cons):
        return "stoichiometric compound"
    if pv.dis_part:
        return f"ordered, partitioned model (disordered part {pv.dis_part})"
    if pv.disordered_of:
        return f"disordered part of {pv.disordered_of}"
    if pv.ordered_like:
        return f"ordering on equivalent sublattices ({nsub} sublattices, CEF)"
    if nsub == 1:
        return "substitutional solution"
    if nsub == 2 and "VA" in cons[1] and len(cons[1]) > 1 and "VA" not in cons[0]:
        return "interstitial solution (Va + interstitials on sublattice 2)"
    if nsub == 2 and cons[1] == ["VA"]:
        return "substitutional solution (interstitial sublattice: Va only)"
    return f"compound energy formalism, {nsub} sublattices"


class Entry:
    def __init__(self, sub: Subsystem, pv: PhaseView, elements: set):
        full = sub.tdb.phases[pv.native]
        self.db, self.rel, self.phase, self.native = sub.label, sub.rel, pv.canon, pv.native
        self.display = format_phase(pv.native, sorted(elements))
        # the model type is a property of the phase as defined in the database, not of the restricted system
        fullpv = PhaseView(pv.canon, pv.native, list(full.site_ratios), [list(sl) for sl in full.constituents],
                           [i for i, sl in enumerate(full.constituents) if list(sl) == ["VA"]], pv.magnetic,
                           pv.dis_part, pv.disordered_of)
        red = [(tuple(sorted(c)), round(r, 6)) for i, (c, r) in enumerate(zip(fullpv.constituents, full.site_ratios))
               if i not in fullpv.va_only]
        fullpv.ordered_like = len(red) > 1 and len(set(red)) < len(red) and all(len(c) > 1 for c, _ in red)
        self.model = model_type(pv.native, fullpv)
        if self.model != "stoichiometric compound" and all(len(sl) == 1 for sl in pv.constituents):
            self.model += "; one end-member inside the system"
        self.formula = sublattice_formula(pv.site_ratios, pv.constituents, elements)
        self.full_formula = sublattice_formula(full.site_ratios, full.constituents, elements)
        extra = [[s for s in sl if s not in inside] for sl, inside in zip(full.constituents, pv.constituents)]
        self.other = "; ".join(f"sublattice {i + 1}: " + ",".join(_sp(s, elements) for s in ex)
                               for i, ex in enumerate(extra) if ex)
        self.site_ratios = ":".join(f"{r:g}" for r in pv.site_ratios)
        self.sites = [",".join(_sp(s, elements) for s in sl) for sl in pv.constituents]
        self.magnetic = (f"AFM factor {pv.magnetic[0]:g}, p = {pv.magnetic[1]:g}" if pv.magnetic else "")
        self.ordering = "; ".join(h for h in pv.hints() if not h.startswith("magnetic"))
        rows = [r for r in sub.rows.values() if r.phase == pv.canon]
        self.n_params = len(rows)
        systems = set()
        for r in rows:
            els = sorted({e for sl in key_constituents(r.key) for s in sl for e in sub.species_elements(s)}
                         - {"VA", "*"})
            if els:
                systems.add("-".join(e.capitalize() for e in els))
        self.systems = ", ".join(sorted(systems, key=lambda s: (s.count("-"), s)))

    def as_dict(self) -> Dict[str, str]:
        return {"phase": self.phase, "label": self.display, "database": self.db, "file": self.rel,
                "native_name": self.native, "model": self.model, "sublattices (in system)": self.formula,
                "site_ratios": self.site_ratios,
                **{f"sublattice {i + 1}": s for i, s in enumerate(self.sites)},
                "other constituents in database": self.other, "full model in database": self.full_formula,
                "magnetic": self.magnetic, "ordering": self.ordering, "parameters in system": str(self.n_params),
                "subsystems with parameters": self.systems}


def inventory(paths: Sequence[str], elements: Sequence[str], workdir: Path,
              phases: Optional[Sequence[str]] = None):
    els = [e.upper() for e in elements]
    labels = cs._labels_for(list(paths), workdir)
    subs = []
    for p, lab in zip(paths, labels):
        if Path(p).suffix.lower() == ".dat":
            raise ValueError(f"{Path(p).name} is a ChemSage .dat database; convert it first with convert_database")
        subs.append(Subsystem(cs._parsed(Path(p)), lab, cs._rel(Path(p), workdir), els, {}))
    want = {x.upper() for x in phases} if phases else None
    E = set(els) | {"VA"}
    entries: Dict[str, List[Entry]] = {}
    for s in subs:
        for cn, pv in s.phases.items():
            if want and cn not in want:
                continue
            entries.setdefault(cn, []).append(Entry(s, pv, E))
    order = phase_order(entries)
    try:
        cands = cs.phase_candidates(subs, subs[0].label) if len(subs) > 1 else []
    except Exception:  # noqa: BLE001 - candidates are a hint only
        cands = []
    return subs, {k: entries[k] for k in order}, cands


def write_csv(entries, path: str) -> str:
    rows = [e.as_dict() for es in entries.values() for e in es]
    keys: List[str] = []
    for r in rows:
        keys += [k for k in r if k not in keys]
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=keys)
        w.writeheader()
        w.writerows(rows)
    return path


def write_xlsx(subs, entries, cands, elements, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Phases"
    rows = [e.as_dict() for es in entries.values() for e in es]
    keys: List[str] = []
    for r in rows:
        keys += [k for k in r if k not in keys]
    ws.append(keys)
    for r in rows:
        ws.append([r.get(k, "") for k in keys])
    head = PatternFill("solid", fgColor="2C5364")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for col, k in zip(ws.columns, keys):
        width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[col[0].column_letter].width = width
    ws.freeze_panes = "C2"
    mx = wb.create_sheet("Matrix")
    mx.append(["phase", "label"] + [s.label for s in subs])
    for ph, es in entries.items():
        by = {e.db: e for e in es}
        mx.append([ph, es[0].display] + [(f"{by[s.label].formula}\n{by[s.label].model}" if s.label in by else "—")
                                         for s in subs])
    for c in mx[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for row in mx.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    mx.column_dimensions["A"].width = 18
    mx.column_dimensions["B"].width = 12
    for i in range(len(subs)):
        mx.column_dimensions[mx.cell(1, 3 + i).column_letter].width = 34
    dbs = wb.create_sheet("Databases")
    dbs.append(["label", "file", "elements in file", "phases in system"])
    for s in subs:
        dbs.append([s.label, s.rel, ",".join(sorted(e for e in s.tdb.elements if e not in ("VA", "/-"))),
                    len(s.phases)])
    if cands:
        cw = wb.create_sheet("Phase candidates")
        cw.append(["phase A", "databases A", "phase B", "databases B", "verdict", "evidence"])
        for c in cands:
            cw.append([c.a, ", ".join(c.dbs_a), c.b, ", ".join(c.dbs_b), c.verdict, "; ".join(c.evidence)])
    wb.save(path)
    return path


def write_pptx(subs, entries, cands, elements, path: str) -> str:
    from . import pptx_utils as pu
    from .pptx_utils import GridCell as G
    prs = pu.new_presentation()
    system = "–".join(e.capitalize() for e in elements)
    pu.add_title_slide(prs, f"{system}: phases in {len(subs)} database(s)",
                       "\n".join(f"{s.label}: {s.rel}" for s in subs),
                       note="Sublattice models restricted to the system's elements (+Va). Phases are matched "
                            "across databases by exact name only.")
    # matrix: phase x database -> model formula
    per = 4
    groups = [subs[i:i + per] for i in range(0, len(subs), per)]
    for gi, grp in enumerate(groups):
        rows = []
        for ph, es in entries.items():
            by = {e.db: e for e in es}
            cells = [G(f"{es[0].display}", bold=True, sub=ph if es[0].display != ph else "")]
            for s in grp:
                e = by.get(s.label)
                cells.append(G(e.formula, sub=e.model) if e else G("—", fill="absent", muted=True))
            rows.append(cells)
        w_first = 1.9
        w = (12.3 - w_first) / max(1, len(grp))
        pu.add_flow_table_slides(prs, "Phases present" + (f"  [{grp[0].label}–{grp[-1].label}]"
                                                          if len(groups) > 1 else ""),
                                 ["phase"] + [s.label for s in grp], rows, [w_first] + [w] * len(grp), font_size=9,
                                 subtitle="sublattice model within the system; model type in grey",
                                 footnote="— = phase not in this database (or no constituent of the system on a "
                                          "sublattice)")
    # details per phase
    rows = []
    for ph, es in entries.items():
        rows.append(("band", f"{es[0].display}  ({ph})" if es[0].display != ph else ph))
        for e in es:
            extra = "; ".join(x for x in (e.magnetic and "magnetic: " + e.magnetic, e.ordering,
                                          e.other and "also in DB: " + e.other) if x)
            rows.append([G(e.db, bold=True), G(e.formula), G(e.model),
                         G(extra), G(f"{e.n_params}", align="r", sub=e.systems)])
    pu.add_flow_table_slides(prs, "Phase models", ["database", "sublattices", "model", "magnetic / ordering / "
                             "other constituents", "params"], rows, [2.0, 3.0, 2.9, 3.3, 1.1], font_size=8,
                             subtitle="params = parameters of the phase inside the system (subsystems listed)")
    if cands:
        rows = [[G(c.a), G(", ".join(c.dbs_a)), G(c.b), G(", ".join(c.dbs_b)), G(c.verdict, bold=True),
                 G("; ".join(c.evidence))] for c in cands]
        pu.add_flow_table_slides(prs, "Possible phase correspondences", ["phase", "in", "phase", "in", "verdict",
                                                                         "evidence"],
                                 rows, [1.6, 1.8, 1.6, 1.8, 1.5, 4.0], font_size=8,
                                 subtitle="not merged: names are matched exactly")
    return pu.save(prs, path)


def summary_text(subs, entries, cands, files, rel=lambda p: p) -> str:
    lines = [f"{len(entries)} phase(s) in {len(subs)} database(s): " + "; ".join(f"{s.label} = {s.rel}" for s in subs)]
    for ph, es in entries.items():
        forms = {}
        for e in es:
            forms.setdefault((e.formula, e.model), []).append(e.db)
        missing = [s.label for s in subs if s.label not in {e.db for e in es}]
        lab = f"{ph} ({es[0].display})" if es[0].display != ph else ph
        parts = [f"{f} [{m}] in {', '.join(d)}" for (f, m), d in forms.items()]
        extra = {e.magnetic for e in es if e.magnetic} | {e.ordering for e in es if e.ordering}
        lines.append(f"- {lab}: " + " | ".join(parts) + (f"; {'; '.join(sorted(extra))}" if extra else "")
                     + (f"; absent in {', '.join(missing)}" if missing and len(subs) > 1 else ""))
    if cands:
        lines.append("Possible correspondences (NOT merged; names match exactly):")
        lines += ["  " + c.text().split(". If the same phase")[0] for c in cands[:8]]
    lines.append("Files: " + ", ".join(f"{k}: {rel(v)}" for k, v in files.items()))
    return "\n".join(lines)


def register(tb) -> None:
    @tb.add("phase_inventory",
            "List every phase of an element system across databases: sublattice model ((Cr,Fe)₁(Va)₃), "
            "constituents per site, model type, magnetic/order-disorder settings, parameter count; phase x "
            "database matrix. Give tdbs, or folder (+pattern) to use every database containing the elements. "
            "Writes xlsx + csv (+ pptx).",
            {"elements": {"type": "array", "items": {"type": "string"}},
             "tdbs": {"type": "array", "items": {"type": "string"}},
             "folder": {"type": "string"}, "pattern": {"type": "string"},
             "phases": {"type": "array", "items": {"type": "string"}, "description": "optional subset"},
             "pptx": {"type": "boolean", "description": "also a PowerPoint (default true)"},
             "output": {"type": "string", "description": "base file name"}}, ["elements"])
    def phase_inventory(elements, tdbs=None, folder=None, pattern=None, phases=None, pptx=True, output=None):
        if tdbs:
            paths = [tb._find_tdb(t) for t in tdbs]
        else:
            rows, _ = cs.find_databases(tb.workdir, folder, pattern, elements, True)
            paths = [str(tb.workdir / r["path"]) for r in rows if not r["path"].lower().endswith(".dat")]
            if not paths:
                raise ValueError(f"no TDB containing {'+'.join(e.upper() for e in elements)} under "
                                 f"{folder or 'the working directory'} (see find_databases)")
        subs, entries, cands = inventory(paths, elements, tb.workdir, phases)
        system = "-".join(e.capitalize() for e in elements)
        base = tb.outp(output, f"phases_{system}")
        base = str(Path(base).with_suffix("")) if Path(base).suffix.lower() in (".xlsx", ".pptx", ".csv") else base
        files = {"xlsx": write_xlsx(subs, entries, cands, elements, base + ".xlsx"),
                 "csv": write_csv(entries, base + ".csv")}
        if pptx:
            files["pptx"] = write_pptx(subs, entries, cands, elements, base + ".pptx")
        return summary_text(subs, entries, cands, files, tb.rel), []
