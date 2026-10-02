"""show_code generators for the figure tools (make_plot, edit_figure, compose_images, image_tools, ...).

Figure operations of :mod:`calphad_tools.figures` and its extensions (figdata sources + transform pipelines, figfit
fits and statistics, figops_extra grids/draw/axes ops) are translated into plain matplotlib / pandas / numpy / scipy
code.  Data-dependent *decisions* the toolkit makes on the fly (column name prefixes, delimiter sniffing, fit start
values, number of groups, tick label rotation, ...) are resolved at generation time by running the toolkit's own
helpers on the real data, so the emitted code stays short and reproduces the figure the tool made.

Standalone figures hold the values *as displayed* (the parent cells plot e.g. °C as data), so the toolkit's display
transforms (``ScaledFormatter``) disappear; values an operation gives in raw data units (``units: data``) are
converted where needed.
"""
from __future__ import annotations

import ast
import json
import keyword
import math
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np

from .codegen import Ctx, generator


# =========================================================================================================
# literals / small code utilities
# =========================================================================================================
class Raw(str):
    """A code expression (emitted verbatim by :func:`lit`)."""


def _plain(v):
    if isinstance(v, Raw):
        return v
    if isinstance(v, np.ndarray):
        return [_plain(x) for x in v.tolist()]
    if isinstance(v, (np.bool_,)):
        return bool(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, tuple):
        return tuple(_plain(x) for x in v)
    if isinstance(v, list):
        return [_plain(x) for x in v]
    if isinstance(v, dict):
        return {_plain(k): _plain(x) for k, x in v.items()}
    return v


def _lit1(v) -> str:
    """One-line Python literal (nan/inf as np.nan/np.inf)."""
    if isinstance(v, Raw):
        return str(v)
    if isinstance(v, bool) or v is None:
        return repr(v)
    if isinstance(v, float):
        if math.isnan(v):
            return "np.nan"
        if math.isinf(v):
            return "np.inf" if v > 0 else "-np.inf"
        return repr(v)
    if isinstance(v, (int, str)):
        return repr(v)
    if isinstance(v, tuple):
        inner = ", ".join(_lit1(x) for x in v)
        return f"({inner},)" if len(v) == 1 else f"({inner})"
    if isinstance(v, list):
        return "[" + ", ".join(_lit1(x) for x in v) + "]"
    if isinstance(v, dict):
        return "{" + ", ".join(f"{_lit1(k)}: {_lit1(x)}" for k, x in v.items()) + "}"
    return repr(v)


def lit(v, indent: int = 0, width: int = 110) -> str:
    """Readable Python literal; long lists / dicts are wrapped (continuation lines indented by ``indent``)."""
    v = _plain(v)
    s = _lit1(v)
    if len(s) + indent <= width or not isinstance(v, (list, tuple, dict)):
        return s
    pad = " " * (indent + 4)
    if isinstance(v, dict):
        items = [f"{_lit1(k)}: {lit(x, indent + 4, width)}" for k, x in v.items()]
        return "{\n" + ",\n".join(pad + it for it in items) + "}"
    op, cl = ("[", "]") if isinstance(v, list) else ("(", ")")
    if all(not isinstance(x, (list, tuple, dict)) for x in v):
        lines, cur = [], ""
        for x in v:
            t = _lit1(x)
            if cur and len(pad) + len(cur) + len(t) + 2 > width:
                lines.append(cur)
                cur = ""
            cur = f"{cur}, {t}" if cur else t
        lines.append(cur)
        return op + "\n" + ",\n".join(pad + ln for ln in lines) + cl
    return op + "\n" + ",\n".join(pad + lit(x, indent + 4, width) for x in v) + cl


def kwargs(d: Dict[str, Any], indent: int = 0) -> str:
    """'k=v, ...' for a dict (identifier keys) or '**{...}' for odd keys."""
    parts, odd = [], {}
    for k, v in d.items():
        if isinstance(k, str) and k.isidentifier() and not keyword.iskeyword(k):
            parts.append(f"{k}={lit(v, indent)}")
        else:
            odd[k] = v
    if odd:
        parts.append("**" + lit(odd, indent))
    return ", ".join(parts)


def call(fn: str, *args, indent: int = 0, **kw) -> str:
    a = [str(x) for x in args if x is not None and x != ""]
    k = kwargs(kw, indent)
    return f"{fn}({', '.join(a + ([k] if k else []))})"


def _ident(s: str, default: str = "data") -> str:
    s = re.sub(r"\W+", "_", str(s)).strip("_").lower()
    if not s or s[0].isdigit():
        s = f"{default}_{s}" if s else default
    if keyword.iskeyword(s) or s in _RESERVED:
        s += "_df"
    return s


_RESERVED = {"np", "pd", "plt", "fig", "ax", "axs", "axes", "out", "workdir", "path", "grid", "table", "x", "y", "z",
             "tri", "cs", "im", "mesh", "q", "sc", "cb", "groups", "subsets", "series", "categories", "values", "labels",
             "means", "parts", "lines", "line", "p", "pcov", "resid", "xg", "yg", "t", "w", "ok", "sigma", "e", "lo",
             "hi", "c", "m", "maps", "cax", "wedges", "bp", "hb", "counts", "colors", "keep", "order", "ys", "xs",
             "frames", "img", "bg", "note", "params", "equation", "se", "dof", "r2", "surf", "ang", "pos", "bars",
             "codes", "edges", "i", "g", "xv", "xn", "obj", "masses", "frac", "amount", "total", "data", "json", "re",
             "stats", "mcolors", "mticker", "gs", "rc", "cycler", "k", "d", "v", "r", "s"}


def _first(v):
    return v[0] if isinstance(v, (list, tuple)) else v


# =========================================================================================================
# runtime helpers (emitted once per script, only when needed)
# =========================================================================================================
HELPERS: Dict[str, Tuple[str, Tuple[str, ...]]] = {}


def _helper(name: str, *imports: str):
    def deco(src: str):
        HELPERS[name] = (src.strip("\n") + "\n", imports)
        return src
    return deco


_helper("panel_axes")('''
def panel_axes(fig):
    """The figure's plot axes, numbered like the toolkit: panels/twins first, then insets (no colour bars)."""
    top = [a for a in fig.axes if a.get_label() not in ("<colorbar>", "background image")]
    kids = [c for a in top for c in a.child_axes
            if type(c).__name__ != "SecondaryAxis" and c.get_label() != "<colorbar>" and c not in top]
    return top + kids
''')

_helper("update_legend")('''
def update_legend(ax, only_if_new=True, fallback_fontsize=None, **kw):
    """Legend = existing entries (incl. proxy handles) + newly labelled artists (the toolkit's legend rule)."""
    old = ax.get_legend()
    handles, labels = [], []
    if old is not None:
        handles = list(old.legend_handles)
        labels = [t.get_text() for t in old.get_texts()]
    new_handles, new_labels = ax.get_legend_handles_labels()
    if only_if_new and all(lab in labels for lab in new_labels):
        if fallback_fontsize is not None and old is None and any(
                not a.get_label().startswith("_") for a in list(ax.get_lines()) + list(ax.collections)):
            ax.legend(fontsize=fallback_fontsize)
        return
    if old is not None and not kw:           # keep the previous placement and font size
        kw["loc"] = old._loc
        if old._bbox_to_anchor is not None:
            (x0, y0), (x1, y1) = ax.transAxes.inverted().transform(old.get_bbox_to_anchor().get_points())
            same = abs(x1 - x0) < 1e-9 and abs(y1 - y0) < 1e-9
            kw["bbox_to_anchor"] = (x0, y0) if same else (x0, y0, x1 - x0, y1 - y0)
        if old.get_texts():
            kw["fontsize"] = old.get_texts()[0].get_fontsize()
    for h, lab in zip(new_handles, new_labels):
        if lab not in labels:
            handles.append(h)
            labels.append(lab)
    if handles:
        ax.legend(handles, labels, **kw)
''')

_helper("hide_legend")('''
def hide_legend(ax):
    """Hidden legend (the toolkit keeps one so later edits do not add it back)."""
    leg = ax.get_legend()
    if leg is None and ax.get_legend_handles_labels()[0]:
        leg = ax.legend()
    if leg is not None:
        leg.set_visible(False)
''')

_helper("find_artists", "import re", "import matplotlib.colors as mcolors")('''
def find_artists(ax, match):
    """Artists whose label matches (substring, regex, 'color:#hex'), or drawn in the colour of a legend entry."""
    arts = list(ax.get_lines()) + list(ax.collections) + list(ax.patches) + list(ax.texts)

    def colour_of(a):
        c = a.get_color() if hasattr(a, "get_color") else a.get_facecolor()
        c = c[0] if isinstance(c, np.ndarray) and c.ndim == 2 and len(c) else c
        return mcolors.to_hex(c)

    def by_label(m):
        if m.lower().startswith("color:"):
            want = mcolors.to_hex(m.split(":", 1)[1].strip())
            out = []
            for a in arts:
                try:
                    if colour_of(a) == want:
                        out.append(a)
                except (ValueError, TypeError, IndexError):
                    continue
            return out
        rx = re.compile(m, re.I) if any(ch in m for ch in "^$[]()*+?|\\\\") else None
        out = []
        for a in arts:
            lab = a.get_text() if hasattr(a, "get_text") and not a.get_label() else a.get_label()
            if lab and not lab.startswith("_") and ((rx and rx.search(lab)) or (not rx and m.lower() in lab.lower())):
                out.append(a)
        return out

    hits = by_label(match)
    if hits or match.lower().startswith("color:"):
        return hits
    leg = ax.get_legend()          # phase names that only exist as legend proxies: artists in that colour
    if leg is None:
        return []
    colours = {}
    for h, t in zip(leg.legend_handles, leg.get_texts()):
        try:
            colours[t.get_text()] = colour_of(h)
        except (ValueError, TypeError, IndexError, AttributeError):
            colours[t.get_text()] = ""
    names = [k for k in colours if k.lower() == match.lower()] or [k for k in colours if match.lower() in k.lower()]
    out = []
    for k in names:
        if colours[k]:
            out += [a for a in by_label("color:" + colours[k]) if a not in out]
    if out:                        # keep the legend swatch in sync
        out += [h for h, t in zip(leg.legend_handles, leg.get_texts()) if t.get_text() in names]
    return out
''')

_helper("series_colors", "import matplotlib.colors as mcolors")('''
def series_colors(spec, labels):
    """Colours per label from a list | {label: colour} | colour | colormap name ('viridis' or 'viridis:N')."""
    n = len(labels)
    if isinstance(spec, dict):
        low = {str(k).lower(): v for k, v in spec.items()}
        return [low.get(str(lab).lower(), f"C{i % 10}") for i, lab in enumerate(labels)]
    if isinstance(spec, str):
        if mcolors.is_color_like(spec):
            return [spec] * n
        cmap = plt.get_cmap(spec.split(":")[0])
        return [mcolors.to_hex(cmap(i / max(1, n - 1))) for i in range(n)]
    return [spec[i % len(spec)] for i in range(n)]
''')

_helper("interp1")('''
def interp1(x, y, xn, method="linear", extrapolate=False):
    """1-D interpolation as the toolkit does it: NaNs dropped, duplicate x averaged, NaN outside the data."""
    from scipy import interpolate as si
    ok = ~(np.isnan(x) | np.isnan(y))
    x, y = x[ok], y[ok]
    if len(x) == 0:
        return np.full(len(xn), np.nan)
    ux, inv = np.unique(x, return_inverse=True)
    uy = np.bincount(inv, weights=y) / np.bincount(inv)
    if len(ux) == 1:
        return np.where(np.isclose(xn, ux[0]), uy[0], np.nan)
    if method == "linear" and not extrapolate:
        out = np.interp(xn, ux, uy)
    elif method == "linear":
        out = si.interp1d(ux, uy, bounds_error=False, fill_value="extrapolate")(xn)
    elif method == "cubic" and len(ux) >= 4:
        out = si.CubicSpline(ux, uy, extrapolate=True)(xn)
    elif method in ("pchip", "cubic"):
        out = si.PchipInterpolator(ux, uy, extrapolate=True)(xn)
    elif method == "akima":
        f = si.Akima1DInterpolator(ux, uy, extrapolate=True) if len(ux) >= 3 else si.PchipInterpolator(ux, uy)
        out = f(xn)
    else:
        out = si.interp1d(ux, uy, kind="nearest", bounds_error=False, fill_value="extrapolate")(xn)
    out = np.asarray(out, dtype=float)
    if not extrapolate:
        out = np.where((xn < ux[0]) | (xn > ux[-1]), np.nan, out)
    return out
''')

_helper("savgol_smooth", "from scipy.signal import savgol_filter")('''
def savgol_smooth(y, window, order):
    """Savitzky-Golay with the window shrunk to an odd length <= len(y) (unchanged if too short)."""
    y = np.asarray(y, dtype=float)
    w = min(window, len(y) if len(y) % 2 else len(y) - 1)
    w = w if w % 2 else w - 1
    return y.copy() if w <= order else savgol_filter(y, w, order, mode="interp")
''')

_helper("jacobian")('''
def jacobian(f, x, p):
    """Numerical d f(x, *p) / d p (central differences), for the confidence band."""
    J = np.empty((len(x), len(p)))
    for i in range(len(p)):
        h = 1e-6 * abs(p[i]) if p[i] != 0 else 1e-8
        pp, pm = p.copy(), p.copy()
        pp[i] += h
        pm[i] -= h
        J[:, i] = (f(x, *pp) - f(x, *pm)) / (2 * h)
    return J
''')

_helper("fmt4")('''
def fmt4(v):
    return "nan" if not np.isfinite(v) else f"{v:.4g}"


def poly_str(c, var="x"):
    """'a + b·x − c·x²' with the coefficients in ascending order."""
    out = ""
    for i, v in enumerate(c):
        term = "" if i == 0 else (f"·{var}" if i == 1 else f"·{var}{'²³'[i - 2] if i <= 3 else '^' + str(i)}")
        if i == 0:
            out = ("−" if v < 0 else "") + fmt4(abs(v))
        else:
            out += (" − " if v < 0 else " + ") + fmt4(abs(v)) + term
    return out
''')

_helper("trim_box")('''
def trim_box(im, tolerance=12, pad=4):
    """(left, top, right, bottom) of the content that differs from the uniform border colour."""
    arr = np.asarray(im.convert("RGBA")).astype(np.int16)
    h, w = arr.shape[:2]
    corners = [tuple(arr[y, x]) for y, x in ((0, 0), (0, w - 1), (h - 1, 0), (h - 1, w - 1))]
    bg = np.array(max(set(corners), key=corners.count))
    if bg[3] < 16:
        mask = arr[:, :, 3] > tolerance
    else:
        mask = (np.abs(arr[:, :, :3] - bg[:3]).max(axis=2) > tolerance) & (arr[:, :, 3] > 16)
    rows, cols = np.where(mask.any(axis=1))[0], np.where(mask.any(axis=0))[0]
    if rows.size == 0:
        return None
    return (max(int(cols[0]) - pad, 0), max(int(rows[0]) - pad, 0),
            min(int(cols[-1]) + 1 + pad, w), min(int(rows[-1]) + 1 + pad, h))
''')

_helper("open_image", "from PIL import Image, ImageOps")('''
def open_image(path):
    """First frame of an image as RGB or RGBA."""
    with Image.open(path) as im:
        im.seek(0)
        im = ImageOps.exif_transpose(im)
        alpha = im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info)
        return im.convert("RGBA" if alpha else "RGB")


def flatten(im, color=(255, 255, 255)):
    """RGB copy with transparency composited onto a background colour."""
    if im.mode != "RGBA":
        return im.convert("RGB")
    bg = Image.new("RGB", im.size, color)
    bg.paste(im, mask=im.getchannel("A"))
    return bg
''')

_helper("label_points")('''
def label_points(ax, x, y, labels, offset=(5, 5), avoid=True, transform=None, **style):
    """Label points, trying other offsets when a label would overlap one placed before (the toolkit's rule)."""
    fig = ax.figure
    ax.get_xlim(), ax.get_ylim()                      # apply pending autoscaling before measuring
    dx, dy = offset
    cands = [(dx, dy), (dx, -dy), (-dx, dy), (-dx, -dy), (0, 1.8 * dy), (0, -1.8 * dy), (2 * dx, 0), (-2 * dx, 0),
             (2 * dx, 2 * dy), (-2 * dx, -2 * dy)]
    avoid = avoid and len(labels) <= 400
    renderer = fig.canvas.get_renderer() if avoid else None
    placed = []
    for i, (xi, yi, lab) in enumerate(zip(x, y, labels)):
        if not lab or not (np.isfinite(xi) and np.isfinite(yi)):
            continue
        order = cands if avoid else cands[i % 2:] + cands[:i % 2]
        for j, (ox, oy) in enumerate(order):
            ha = "left" if ox > 0 else ("right" if ox < 0 else "center")
            va = "bottom" if oy > 0 else ("top" if oy < 0 else "center")
            t = ax.annotate(lab, xy=(xi, yi), xycoords=transform or ax.transData, xytext=(ox, oy),
                            textcoords="offset points", ha=ha, va=va, **style)
            if not avoid:
                break
            bb = t.get_window_extent(renderer).expanded(1.05, 1.1)
            if not any(bb.overlaps(p) for p in placed) or j == len(order) - 1:
                placed.append(bb)
                break
            t.remove()
''')

_helper("label_lines")('''
def label_lines(ax, lines, x=None, where="end", offset=None, fontsize=9, bbox=False, avoid=True):
    """Write each line's label next to it in its colour, nudging labels apart vertically."""
    fig = ax.figure
    ax.get_xlim(), ax.get_ylim()
    items = []
    for ln in lines:
        xd, yd = np.asarray(ln.get_xdata(), dtype=float), np.asarray(ln.get_ydata(), dtype=float)
        ok = np.isfinite(xd) & np.isfinite(yd)
        xd, yd = xd[ok], yd[ok]
        if not len(xd):
            continue
        if x is not None:
            o = np.argsort(xd)
            items.append([ln, x, float(np.interp(x, xd[o], yd[o])), "center"])
        elif where in ("max", "min"):
            k = int(np.argmax(yd) if where == "max" else np.argmin(yd))
            items.append([ln, xd[k], yd[k], "center"])
        else:
            k = -1 if where == "end" else 0
            items.append([ln, xd[k], yd[k], "left" if where == "end" else "right"])
    if offset is None:
        offset = [4 if where == "end" else (-4 if where == "start" else 0),
                  0 if where in ("end", "start") and x is None else 4]
    dx, dy = offset
    shifts = [0.0] * len(items)
    if avoid and len(items) > 1:
        disp = [ax.transData.transform((it[1], it[2]))[1] * 72.0 / fig.dpi for it in items]
        last = -np.inf
        for k in np.argsort(disp):
            want = disp[k]
            if want < last + fontsize * 1.15:
                shifts[k] = last + fontsize * 1.15 - want
                want = last + fontsize * 1.15
            last = want
    for (ln, px, py, ha), sh in zip(items, shifts):
        kw = {"bbox": {"boxstyle": "round,pad=0.15", "facecolor": "white", "alpha": 0.8, "edgecolor": "none"}} \\
            if bbox else {}
        ax.annotate(ln.get_label(), xy=(px, py), xycoords="data", xytext=(dx if ha != "center" else 0, dy + sh),
                    textcoords="offset points", ha=ha, va="center" if ha != "center" else "bottom",
                    fontsize=fontsize, color=ln.get_color(), annotation_clip=False, **kw)
''')

_helper("mappables")('''
def mappables(ax):
    """Colour-mapped artists (images, contour sets, meshes, scatter with c) of an axes."""
    from matplotlib.cm import ScalarMappable
    return [a for a in list(ax.images) + list(ax.collections)
            if isinstance(a, ScalarMappable) and a.get_array() is not None]
''')

_helper("shared_colorbar_axes")('''
def shared_colorbar_axes(fig, axes, location="right", fraction=0.05, pad=0.08, wspace=0.3, hspace=0.35):
    """Colour-bar axes next to a rectangular block of grid panels (panels re-gridded to make room)."""
    specs = [a.get_subplotspec() for a in axes]
    if any(s is None for s in specs) or any(s.get_gridspec() is not specs[0].get_gridspec() for s in specs):
        return None
    gs = specs[0].get_gridspec()
    r0, r1 = min(s.rowspan.start for s in specs), max(s.rowspan.stop for s in specs)
    c0, c1 = min(s.colspan.start for s in specs), max(s.colspan.stop for s in specs)
    members = [a for a in fig.axes if a.get_subplotspec() is not None and a.get_subplotspec().get_gridspec() is gs
               and a.get_label() != "<colorbar>" and a.get_subplotspec().rowspan.start < r1
               and a.get_subplotspec().rowspan.stop > r0 and a.get_subplotspec().colspan.start < c1
               and a.get_subplotspec().colspan.stop > c0]
    ratios = [1.0, fraction] if location in ("right", "bottom") else [fraction, 1.0]
    block = gs[r0:r1, c0:c1]
    outer = (block.subgridspec(1, 2, width_ratios=ratios, wspace=pad) if location in ("right", "left")
             else block.subgridspec(2, 1, height_ratios=ratios, hspace=pad))
    main_cell, cb_cell = (outer[0], outer[1]) if location in ("right", "bottom") else (outer[1], outer[0])
    inner = main_cell.subgridspec(r1 - r0, c1 - c0, wspace=wspace, hspace=hspace)
    for a in members:
        s = a.get_subplotspec()
        a.set_subplotspec(inner[s.rowspan.start - r0:s.rowspan.stop - r0, s.colspan.start - c0:s.colspan.stop - c0])
    return fig.add_subplot(cb_cell, label="<colorbar>")
''')

_helper("copy_artists", "from matplotlib import collections as mcoll, patches as mpatches",
        "from matplotlib.lines import Line2D", "from matplotlib.text import Text",
        "from matplotlib.transforms import IdentityTransform")('''
def copy_artists(src, dst):
    """Copy lines, scatter/line/poly collections, rectangles/polygons and texts drawn in data coordinates."""
    td = src.transData
    for ln in src.get_lines():
        if ln.get_label() == "_break" or ln.get_transform() != td:
            continue
        dst.add_line(Line2D(ln.get_xdata(), ln.get_ydata(), color=ln.get_color(), linestyle=ln.get_linestyle(),
                            linewidth=ln.get_linewidth(), marker=ln.get_marker(), markersize=ln.get_markersize(),
                            markerfacecolor=ln.get_markerfacecolor(), markeredgecolor=ln.get_markeredgecolor(),
                            markeredgewidth=ln.get_markeredgewidth(), alpha=ln.get_alpha(), zorder=ln.get_zorder(),
                            drawstyle=ln.get_drawstyle(), label=ln.get_label()))
    for c in src.collections:
        common = dict(alpha=c.get_alpha(), zorder=c.get_zorder(), label=c.get_label())
        try:
            c.update_scalarmappable()
        except Exception:
            pass
        if isinstance(c, mcoll.PathCollection):
            if c.get_offset_transform() != td:
                continue
            nc = mcoll.PathCollection(c.get_paths(), sizes=c.get_sizes(), offsets=c.get_offsets(),
                                      offset_transform=dst.transData, facecolors=c.get_facecolors(),
                                      edgecolors=c.get_edgecolors(), linewidths=c.get_linewidths(), **common)
            nc.set_transform(IdentityTransform())
            dst.add_collection(nc, autolim=False)
        elif isinstance(c, (mcoll.LineCollection, mcoll.PolyCollection)) and c.get_transform() == td:
            segs = [p.vertices for p in c.get_paths()]
            if isinstance(c, mcoll.LineCollection):
                nc = mcoll.LineCollection(segs, colors=c.get_colors(), linewidths=c.get_linewidths(),
                                          linestyles=c.get_linestyles(), **common)
            else:
                nc = mcoll.PolyCollection(segs, facecolors=c.get_facecolors(), edgecolors=c.get_edgecolors(),
                                          linewidths=c.get_linewidths(), **common)
            dst.add_collection(nc, autolim=False)
    for p in src.patches:
        if p.get_transform() != td and not isinstance(p, mpatches.Rectangle):
            continue
        kw = dict(facecolor=p.get_facecolor(), edgecolor=p.get_edgecolor(), linewidth=p.get_linewidth(),
                  alpha=p.get_alpha(), zorder=p.get_zorder(), label=p.get_label(), hatch=p.get_hatch(),
                  fill=p.get_fill())
        if isinstance(p, mpatches.Rectangle):
            dst.add_patch(mpatches.Rectangle(p.get_xy(), p.get_width(), p.get_height(), **kw))
        elif isinstance(p, mpatches.Polygon):
            dst.add_patch(mpatches.Polygon(p.get_xy(), closed=True, **kw))
    for t in src.texts:
        if type(t) is Text and t.get_transform() == td:
            dst.text(*t.get_position(), t.get_text(), color=t.get_color(), fontsize=t.get_fontsize(),
                     ha=t.get_ha(), va=t.get_va(), rotation=t.get_rotation(), zorder=t.get_zorder())
''')

_helper("scatter_matrix", "from scipy.stats import gaussian_kde")('''
def scatter_matrix(fig, ax, parts, cols, diag="hist", bins=15, s=8, alpha=0.7, colors=None, corner=False):
    """Replace ``ax`` by an n x n grid of pairwise scatter plots (histograms / KDEs on the diagonal).

    parts: [(group label or None, DataFrame)]; returns the new axes, row-major."""
    n = len(cols)
    colors = colors or [f"C{i % 10}" for i in range(len(parts))]
    gs = ax.get_subplotspec().subgridspec(n, n, wspace=0.08, hspace=0.08)
    ax.remove()
    axs = [[None] * n for _ in range(n)]
    for i in range(n):
        for j in range(n):
            if corner and j > i:
                continue
            share = {}
            col_top = axs[j][j] if corner else axs[0][j]
            if col_top is not None:
                share["sharex"] = col_top
            if i != j:
                row_first = next((axs[i][k] for k in range(n) if k != i and axs[i][k] is not None), None)
                if row_first is not None:
                    share["sharey"] = row_first
            a = fig.add_subplot(gs[i, j], **share)
            axs[i][j] = a
            for (lab, sub), c in zip(parts, colors):
                xv = sub[cols[j]].to_numpy(dtype=float)
                if i == j:
                    xv = xv[np.isfinite(xv)]
                    if diag == "kde" and len(xv) > 1 and np.ptp(xv) > 0:
                        k = gaussian_kde(xv)
                        g = np.linspace(xv.min(), xv.max(), 128)
                        a.plot(g, k(g), color=c)
                        a.fill_between(g, 0, k(g), color=c, alpha=0.25, linewidth=0)
                    else:
                        a.hist(xv, bins=bins, color=c, alpha=0.6 if len(parts) > 1 else 0.8)
                    a.tick_params(left=False, labelleft=False)
                else:
                    yv = sub[cols[i]].to_numpy(dtype=float)
                    a.scatter(xv, yv, s=s, alpha=alpha, color=c, linewidths=0,
                              label=lab if (lab and i == 1 and j == 0) else None)
            if i < n - 1:
                a.tick_params(labelbottom=False)
            else:
                a.set_xlabel(str(cols[j]))
            if j > 0 and i != j:
                a.tick_params(labelleft=False)
            if j == 0 and i > 0:
                a.set_ylabel(str(cols[i]))
            if i == 0 and j == n - 1 and n > 1 and not corner:
                a.tick_params(labelright=True)
                a.yaxis.set_label_position("right")
                a.set_ylabel(str(cols[0]))
            a.tick_params(labelsize=7)
    if len(parts) > 1:
        axs[1][0].legend(fontsize=7, markerscale=1.5, loc="best")
    return [a for row in axs for a in row if a is not None]
''')


def helper_code(ctx: Ctx, names: Sequence[str]) -> str:
    """Source of the helpers not yet defined earlier in this script (their imports go to the header)."""
    done = ctx.__dict__.setdefault("_plot_helpers", set())
    out = []
    for n in names:
        if n in done:
            continue
        done.add(n)
        src, imps = HELPERS[n]
        ctx.need("import numpy as np", *imps)
        out.append(src)
    return "\n\n".join(out) + ("\n\n" if out else "")


# =========================================================================================================
# static model of a figure's axes (what the toolkit numbers as axes 0, 1, ...)
# =========================================================================================================
class AxModel:
    def __init__(self, expr: str, proj: str = "rectilinear", kind: str = "panel", xscale="linear", yscale="linear",
                 title: str = "", xlabel: str = "", ylabel: str = "", locked: bool = False,
                 xt=(1.0, 0.0), yt=(1.0, 0.0), spec=None, bounds=None):
        self.expr, self.proj, self.kind = expr, proj, kind
        self.xscale, self.yscale = xscale, yscale
        self.title, self.xlabel, self.ylabel = title, xlabel, ylabel
        self.locked = locked
        self.xt, self.yt = tuple(xt), tuple(yt)          # toolkit display transform (scale, offset) data->shown
        self.children: List["AxModel"] = []
        self.spec = spec                                  # (rows, cols) span key: twins share it
        self.bounds = bounds


class FigModel:
    def __init__(self):
        self.top: List[AxModel] = []

    def main(self) -> List[AxModel]:
        return self.top + [c for a in self.top for c in a.children]

    def index(self, a: AxModel) -> int:
        return self.main().index(a)

    @classmethod
    def from_figure(cls, fig, expr_fmt: str = "axes[{i}]") -> "FigModel":
        from .figures import ScaledFormatter, main_axes
        m = cls()
        axs = main_axes(fig)
        models = {}
        for i, a in enumerate(axs):
            def tr(axis):
                f = axis.get_major_formatter()
                return (f.scale, f.offset) if isinstance(f, ScaledFormatter) else (1.0, 0.0)
            ss = a.get_subplotspec() if hasattr(a, "get_subplotspec") else None
            spec = (id(ss.get_gridspec()), ss.rowspan.start, ss.rowspan.stop, ss.colspan.start, ss.colspan.stop) \
                if ss is not None else None
            kind = "broken" if getattr(a, "_calphad_broken_of", None) is not None else "panel"
            models[id(a)] = AxModel(expr_fmt.format(i=i), a.name, kind, a.get_xscale(), a.get_yscale(),
                                    a.get_title(), a.get_xlabel(), a.get_ylabel(),
                                    bool(getattr(a, "_calphad_legend_set", False)), tr(a.xaxis), tr(a.yaxis),
                                    spec, tuple(np.round(a.get_position().bounds, 6)))
        top = [a for a in fig.axes if a.get_label() not in ("<colorbar>", "background image")]
        for a in axs:
            if a in top:
                m.top.append(models[id(a)])
        for a in axs:
            if a not in top:
                parent = next((p for p in top if a in getattr(p, "child_axes", [])), None)
                if parent is not None:
                    models[id(parent)].children.append(models[id(a)])
        # twin axes share a panel with an earlier axes
        seen = []
        for a in m.top:
            if any((a.spec is not None and a.spec == b.spec) or (a.bounds == b.bounds) for b in seen):
                a.kind = "twin" if a.kind == "panel" else a.kind
            else:
                seen.append(a)
        return m


# =========================================================================================================
# data: sources, "where" filters, transform pipelines, safe expressions -> pandas / numpy
# =========================================================================================================
_CONST_CODE = {"pi": "np.pi", "e": "np.e", "R": "8.314462618", "F": "96485.33212", "kB": "1.380649e-23",
               "nan": "np.nan", "inf": "np.inf"}
_NP_FUNCS = {"log": "np.log", "ln": "np.log", "log10": "np.log10", "log2": "np.log2", "exp": "np.exp",
             "sqrt": "np.sqrt", "abs": "np.abs", "sin": "np.sin", "cos": "np.cos", "tan": "np.tan",
             "arcsin": "np.arcsin", "arccos": "np.arccos", "arctan": "np.arctan", "tanh": "np.tanh",
             "floor": "np.floor", "ceil": "np.ceil", "clip": "np.clip", "where": "np.where"}
_PY_OPS = {ast.Add: "+", ast.Sub: "-", ast.Mult: "*", ast.Div: "/", ast.FloorDiv: "//", ast.Mod: "%",
           ast.Pow: "**", ast.Eq: "==", ast.NotEq: "!=", ast.Lt: "<", ast.LtE: "<=", ast.Gt: ">", ast.GtE: ">="}


def _strip_parens(s: str) -> str:
    if not (s.startswith("(") and s.endswith(")")):
        return s
    depth = 0
    for i, ch in enumerate(s):
        depth += ch == "("
        depth -= ch == ")"
        if depth == 0 and i < len(s) - 1:
            return s
    return s[1:-1]


class ExprWriter:
    """Translate a figdata safe expression into pandas / numpy code on DataFrame ``var``.

    The expression is also evaluated (by the toolkit's own interpreter) on the real frame, so every sub-expression's
    type is known: columns (Series) get vectorised operators (& | ~ .isin), scalars plain Python."""

    def __init__(self, frame, var: str):
        from .figdata import SafeExpr
        self.frame, self.var, self.se = frame, var, SafeExpr(frame)

    def __call__(self, expr) -> str:
        if isinstance(expr, (int, float, bool)):
            return lit(expr)
        text = re.sub(r"`([^`]*)`", lambda m: "col(" + repr(m.group(1)) + ")", str(expr).strip())
        return _strip_parens(self.code(ast.parse(text, mode="eval").body))

    def val(self, node):
        return self.se._ev(node)

    def col(self, name: str) -> str:
        from .figures import find_col
        return f"{self.var}[{find_col(self.frame, name)!r}]"

    def as_bool(self, node) -> str:
        import pandas as pd
        v, c = self.val(node), self.code(node)
        if isinstance(v, pd.Series):
            if v.dtype == object:
                return f"{c}.fillna(False).astype(bool)"
            return c if v.dtype == bool else f"{c}.astype(bool)"
        if isinstance(v, np.ndarray):
            return f"np.asarray({c}).astype(bool)"
        return f"bool({c})"

    def _vec(self, *nodes) -> bool:
        import pandas as pd
        return any(isinstance(self.val(n), (pd.Series, np.ndarray)) for n in nodes)

    def code(self, node) -> str:  # noqa: C901 - one branch per node type
        import pandas as pd
        if isinstance(node, ast.Constant):
            v = node.value
            if v is None:
                return "np.nan"
            if isinstance(v, bool) or isinstance(v, str):
                return repr(v)
            return repr(float(v))
        if isinstance(node, ast.Name):
            n = node.id
            if n in self.frame.columns:
                return f"{self.var}[{n!r}]"
            if n in _CONST_CODE:
                return _CONST_CODE[n]
            return self.col(n)
        if isinstance(node, ast.BinOp):
            return f"({self.code(node.left)} {_PY_OPS[type(node.op)]} {self.code(node.right)})"
        if isinstance(node, ast.UnaryOp):
            c = self.code(node.operand)
            if isinstance(node.op, ast.USub):
                return f"(-{c})"
            if isinstance(node.op, ast.UAdd):
                return f"(+{c})"
            return f"(~{c}.astype(bool))" if isinstance(self.val(node.operand), pd.Series) else f"(not {c})"
        if isinstance(node, ast.BoolOp):
            sym = " & " if isinstance(node.op, ast.And) else " | "
            if self._vec(*node.values):
                return "(" + sym.join(self.as_bool(v) for v in node.values) + ")"
            return "(" + sym.join(f"bool({self.code(v)})" for v in node.values) + ")"
        if isinstance(node, ast.Compare):
            parts, left = [], node.left
            for op, comp in zip(node.ops, node.comparators):
                lc = self.code(left)
                if isinstance(op, (ast.In, ast.NotIn)):
                    vals = "[" + ", ".join(self.code(e) for e in comp.elts) + "]"
                    if isinstance(self.val(left), pd.Series):
                        parts.append(f"{'~' if isinstance(op, ast.NotIn) else ''}{lc}.isin({vals})")
                    else:
                        parts.append(f"({lc} {'not in' if isinstance(op, ast.NotIn) else 'in'} {vals})")
                else:
                    parts.append(f"({lc} {_PY_OPS[type(op)]} {self.code(comp)})")
                left = comp
            if len(parts) == 1:
                return parts[0]
            return "(" + " & ".join(parts) + ")"
        if isinstance(node, ast.IfExp):
            return f"np.where({self.as_bool(node.test)}, {self.code(node.body)}, {self.code(node.orelse)})"
        if isinstance(node, ast.Call):
            name = node.func.id
            if name == "col":
                return self.col(node.args[0].value)
            args = [self.code(a) for a in node.args]
            vals = [self.val(a) for a in node.args]
            ser = [isinstance(v, pd.Series) for v in vals]
            if name in _NP_FUNCS:
                return f"{_NP_FUNCS[name]}({', '.join(args)})"
            if name == "round":
                return f"np.round({args[0]}, {int(vals[1]) if len(vals) > 1 else 0})"
            if name in ("min", "max"):
                if len(args) == 1:
                    return f"{args[0]}.{name}()" if ser[0] else f"np.nan{name}({args[0]})"
                out = args[0]
                for a in args[1:]:
                    out = f"np.{'minimum' if name == 'min' else 'maximum'}({out}, {a})"
                return out
            if name in ("mean", "std", "sum", "median"):
                return f"{args[0]}.{name}()" if ser[0] else f"np.nan{name}(np.asarray({args[0]}, dtype=float))"
            if name in ("cumsum", "diff"):
                return f"{args[0]}.{name}()"
            if name in ("isna", "isnan"):
                return f"pd.isna({args[0]})"
            if name == "num":
                return f"pd.to_numeric({args[0]}, errors='coerce')" if ser[0] else f"float({args[0]})"
            if name == "contains":
                if ser[0]:
                    return f"{args[0]}.astype(str).str.contains(str({args[1]}), case=False, regex=False)"
                return f"(str({args[1]}).lower() in str({args[0]}).lower())"
            if name == "startswith":
                return f"{args[0]}.astype(str).str.startswith(str({args[1]}))" if ser[0] else \
                    f"str({args[0]}).startswith(str({args[1]}))"
        raise ValueError(f"cannot translate expression node {type(node).__name__}")


def _cond_code(c: str, spec: Dict) -> List[str]:
    """Masks for one where-column condition dict (see figdata._cond)."""
    out = []
    for k, v in spec.items():
        kk = str(k).lower()
        if kk in (">", "gt"):
            out.append(f"({c} > {lit(v)})")
        elif kk in (">=", "ge", "gte"):
            out.append(f"({c} >= {lit(v)})")
        elif kk in ("<", "lt"):
            out.append(f"({c} < {lit(v)})")
        elif kk in ("<=", "le", "lte"):
            out.append(f"({c} <= {lit(v)})")
        elif kk in ("==", "eq", "="):
            out.append(f"({c} == {lit(v)})")
        elif kk in ("!=", "ne"):
            out.append(f"({c} != {lit(v)})")
        elif kk == "between":
            out.append(f"{c}.between({lit(v[0])}, {lit(v[1])})")
        elif kk == "in":
            out.append(f"{c}.isin({lit(v if isinstance(v, list) else [v])})")
        elif kk in ("not", "not_in"):
            out.append(f"~{c}.isin({lit(v if isinstance(v, list) else [v])})")
        elif kk == "contains":
            out.append(f"{c}.astype(str).str.contains({str(v)!r}, case=False, regex=False)")
        elif kk == "startswith":
            out.append(f"{c}.astype(str).str.startswith({str(v)!r})")
        elif kk == "isna":
            out.append(f"{c}.isna()" if v else f"{c}.notna()")
        else:
            raise ValueError(f"unknown where condition {k!r}")
    return out


def _text_read_opts(path: Optional[str], read: Optional[Dict], src_text: Optional[str] = None) -> Dict:
    """The pandas.read_csv options figdata uses for a text table (delimiter / decimal / comment-line sniffing)."""
    from .figdata import _sniff
    ext = Path(path).suffix.lower() if path else ""
    opts = dict(read or {})
    if "delimiter" in opts:
        opts["sep"] = opts.pop("delimiter")
    if "sep" not in opts:
        if src_text is not None:
            lines = src_text.splitlines()[:40]
        else:
            lines = []
            with open(path, "r", encoding=opts.get("encoding", "utf-8"), errors="replace") as f:
                for ln in f:
                    lines.append(ln.rstrip("\r\n"))
                    if len(lines) >= 60:
                        break
        skip = 0
        while skip < len(lines) and (not lines[skip].strip() or lines[skip].lstrip().startswith("#")):
            skip += 1
        sep, dec = _sniff(lines[skip:skip + 30], ext)
        opts["sep"] = sep
        if dec and "decimal" not in opts:
            opts["decimal"] = dec
        if skip and "skiprows" not in opts:
            opts["skiprows"] = skip
    sep = opts.get("sep")
    if (isinstance(sep, str) and len(sep) > 1 and sep != r"\s+") or opts.get("skipfooter"):
        opts["engine"] = "python"
    return opts


def _read_kw(opts: Dict) -> Dict:
    kw = dict(opts)
    if kw.get("sep") == ",":
        kw.pop("sep")
    if kw.get("sep") == r"\s+":
        kw["sep"] = Raw(r'r"\s+"')
    return kw


class DataWriter:
    """Emits the code that loads an operation's data source (+ where + transform) into a DataFrame variable, while
    tracking the real DataFrame at generation time (to resolve column prefixes and name the outputs)."""

    def __init__(self, ctx: Ctx, emit, names: set, pre: List[str]):
        self.ctx, self.tb, self.emit, self.names, self.pre = ctx, ctx.tb, emit, names, pre
        self.cache: Dict[str, Tuple[str, Any]] = {}
        self._src_name = "data"

    # -- names ----------------------------------------------------------------------------------------
    def fresh(self, base: str) -> str:
        base = _ident(base)
        name, k = base, 2
        while name in self.names:
            name, k = f"{base}_{k}", k + 1
        self.names.add(name)
        return name

    # -- entry point ------------------------------------------------------------------------------------
    def for_op(self, op: Dict) -> Tuple[str, Any]:
        src = op.get("csv") if op.get("csv") is not None else op.get("data")
        key = json.dumps([src, op.get("sheet"), op.get("variables"), op.get("read"), op.get("where"),
                          op.get("transform")], sort_keys=True, default=str)
        if key not in self.cache:
            self.cache[key] = self.load(src, None, sheet=op.get("sheet"), variables=op.get("variables"),
                                        read=op.get("read"), where=op.get("where"), transform=op.get("transform"))
        return self.cache[key]

    def load(self, src, var: Optional[str], sheet=None, variables=None, read=None, where=None, transform=None,
             depth: int = 0) -> Tuple[str, Any]:
        from . import figdata
        self.ctx.need("import numpy as np", "import pandas as pd")
        if figdata._is_spec(src):
            spec = src
            return self.load(spec.get("data", spec.get("csv", spec.get("path"))), var, spec.get("sheet", sheet),
                             spec.get("variables", variables), spec.get("read", read), spec.get("where", where),
                             spec.get("transform", transform), depth + 1)
        if src is None:
            raise ValueError("no data source")
        if isinstance(src, str) and "\n" not in src:
            name = Path(src).stem
            var = var or self.fresh(name)
            df = self._read_path(src, var, sheet, variables, read)
        else:
            name = "data"
            var = var or self.fresh("data")
            df = self._inline(src, var)
        if where:
            df = self.where(df, var, where)
        if transform:
            steps = [transform] if isinstance(transform, dict) else transform
            for st in steps:
                df = self.step(df, var, st, name, depth)
        return var, df

    # -- reading ----------------------------------------------------------------------------------------
    def _read_path(self, src: str, var: str, sheet, variables, read):
        import pandas as pd
        from . import figdata
        ctx = self.ctx
        path = self.tb.p(src)
        pexpr = ctx.path(src)
        ext = Path(path).suffix.lower()
        df = figdata.read_path(path, sheet=sheet, variables=variables, read=read)
        ctx.need("import pandas as pd")
        if ext in figdata.EXCEL_EXT:
            kw = {"sheet_name": 0 if sheet is None else sheet, **(read or {})}
            self.emit(f"{var} = {call('pd.read_excel', pexpr, **kw)}")
            raw_cols = list(pd.read_excel(path, **kw).columns)
        elif ext in figdata.NC_EXT:
            ctx.need("import xarray as xr")
            vs = [variables] if isinstance(variables, str) else variables
            self.emit(f"with xr.open_dataset({pexpr}) as ds:")
            sel = f"ds[{lit(list(vs))}]" if vs else "ds"
            self.emit(f"    ds = {sel}.load()")
            self.emit(f"{var} = (ds.to_dataframe() if ds.data_vars else ds.coords.to_dataset().to_dataframe())"
                      f".reset_index()")
            raw_cols = list(df.columns)
        elif ext in figdata.JSON_EXT:
            if path.lower().endswith((".jsonl", ".ndjson")):
                self.emit(f"{var} = pd.read_json({pexpr}, lines=True)")
            else:
                ctx.need("import json")
                with open(path, "r", encoding="utf-8") as f:
                    txt = f.read()
                try:
                    obj = json.loads(txt)
                except json.JSONDecodeError:
                    ctx.need("import io")
                    self.emit(f"{var} = pd.read_json(io.StringIO({pexpr}.read_text()), lines=True)")
                    obj = None
                if obj is not None:
                    self.emit(f"with open({pexpr}) as f:")
                    self.emit("    obj = json.load(f)")
                    self.emit(f"{var} = {self._inline_code(obj, 'obj')}")
            raw_cols = list(df.columns)
        else:
            opts = _text_read_opts(path, read)
            self.emit(f"{var} = {call('pd.read_csv', pexpr, **_read_kw(opts))}")
            raw_cols = list(pd.read_csv(path, **opts).columns)
        if [str(c).strip() for c in raw_cols] != raw_cols:
            self.emit(f"{var}.columns = [str(c).strip() for c in {var}.columns]")
        return df

    def _inline_code(self, obj, expr: str) -> str:
        """Code turning a parsed JSON object / inline data (named ``expr``) into a DataFrame (figdata rules)."""
        if isinstance(obj, list):
            if obj and all(isinstance(r, dict) for r in obj):
                nested = any(isinstance(v, dict) for r in obj for v in r.values())
                return f"pd.json_normalize({expr})" if nested else f"pd.DataFrame({expr})"
            return f"pd.DataFrame({expr})"
        if isinstance(obj, dict):
            for k in ("data", "records", "rows"):
                v = obj.get(k)
                if isinstance(v, list) and v and all(isinstance(r, dict) for r in v):
                    return self._inline_code(v, f"{expr}[{k!r}]")
            if all(not isinstance(v, (list, tuple, dict)) for v in obj.values()):
                return f"pd.DataFrame([{expr}])"
            return f"pd.DataFrame({expr})"
        raise ValueError("unsupported inline data")

    def _inline(self, src, var: str):
        from . import figdata
        self.ctx.need("import pandas as pd")
        df = figdata._inline_frame(src)
        if isinstance(src, str):
            self.ctx.need("import io")
            opts = _text_read_opts("", {}, src_text=src)
            self.emit(f"{var} = {call('pd.read_csv', f'io.StringIO({src!r})', **_read_kw(opts))}")
            return df
        code = self._inline_code(src, "\x00OBJ\x00")
        self.emit(f"{var} = " + code.replace("\x00OBJ\x00", lit(src, 4)))
        return df

    # -- where ------------------------------------------------------------------------------------------
    def where(self, df, var: str, where: Dict):
        from .figdata import apply_where
        from .figures import find_col
        masks = []
        for col, val in where.items():
            c = f"{var}[{find_col(df, col)!r}]"
            if isinstance(val, dict):
                masks += _cond_code(c, val)
            else:
                masks.append(f"{c}.isin({lit(val if isinstance(val, list) else [val])})")
        if len(masks) == 1:
            self.emit(f"{var} = {var}[{_strip_parens(masks[0])}]")
        else:
            self.emit(f"{var} = {var}[" + " & ".join(masks) + "]")
        return apply_where(df, where)

    # -- transform steps --------------------------------------------------------------------------------
    def step(self, df, var: str, st: Dict, src_name: str, depth: int):
        import pandas as pd
        from . import figdata
        name = figdata.ALIASES.get(str(st["step"]).lower(), str(st["step"]).lower())
        fn = getattr(self, f"_st_{name}", None)
        if fn is None or name not in figdata.STEPS:
            raise ValueError(f"transform step {name!r} has no translation")
        figdata.run_transform(df, [st], self.tb.p, src_name, depth)          # validates the step
        raw = figdata.STEPS[name][0](df, st, {"resolve": self.tb.p, "src": src_name, "depth": depth})
        out = raw if isinstance(raw.index, pd.RangeIndex) else raw.reset_index(drop=True)
        self._src_name = src_name
        fn(df, var, st, out)
        if not isinstance(raw.index, pd.RangeIndex):         # the toolkit resets the index after every step
            self.emit(f"{var} = {var}.reset_index(drop=True)")
        return out

    def c(self, df, name) -> str:
        from .figures import find_col
        return find_col(df, name)

    def cols(self, df, v) -> List[str]:
        from .figdata import _cols
        return _cols(df, v)

    def ocols(self, df, v) -> List[str]:
        return [] if v is None else self.cols(df, v)

    def num(self, df, var, c) -> str:
        import pandas as pd
        s = f"{var}[{c!r}]"
        return f"{s}.to_numpy(float)" if pd.api.types.is_numeric_dtype(df[c]) else f"pd.to_numeric({s}).to_numpy(float)"

    def _st_derive(self, df, var, st, out):
        from .figdata import evaluate
        pairs = list((st.get("exprs") or {}).items())
        if "new" in st:
            pairs.insert(0, (st["new"], st["expr"]))
        cur = df.copy()
        for new, expr in pairs:
            self.emit(f"{var}[{str(new)!r}] = {ExprWriter(cur, var)(expr)}")
            v = evaluate(cur, expr)
            cur[str(new)] = v.to_numpy() if hasattr(v, "to_numpy") else v
        return cur

    def _st_filter(self, df, var, st, out):
        import pandas as pd
        from .figdata import apply_where, evaluate
        cur = df
        if st.get("where"):
            cur = self.where(cur, var, st["where"])
        if st.get("expr") is not None:
            m = evaluate(cur, st["expr"])
            if isinstance(m, pd.Series):
                code = ExprWriter(cur, var)(st["expr"])
                if m.dtype == object:
                    code = f"({code}).fillna(False).astype(bool)"
                self.emit(f"{var} = {var}[{code}]")
                cur = cur[m.fillna(False).astype(bool).to_numpy()]
            elif not bool(m):
                self.emit(f"{var} = {var}.iloc[0:0]   # {st['expr']!r} is always false")
                cur = cur.iloc[0:0]
        return cur if cur is not df else out

    def _st_select(self, df, var, st, out):
        self.emit(f"{var} = {var}[{lit(self.cols(df, st.get('cols')))}]")
        return out

    def _st_drop(self, df, var, st, out):
        self.emit(f"{var} = {var}.drop(columns={lit(self.cols(df, st.get('cols')))})")
        return out

    def _st_rename(self, df, var, st, out):
        m = {self.c(df, k): str(v) for k, v in st["map"].items()}
        self.emit(f"{var} = {var}.rename(columns={lit(m)})")
        return out

    def _st_sort(self, df, var, st, out):
        by = self.cols(df, st.get("by"))
        asc = st.get("ascending", True)
        self.emit(f"{var} = {call(f'{var}.sort_values', lit(by), **({} if asc is True else {'ascending': asc}), kind='stable')}")
        return df.sort_values(by, ascending=asc, kind="stable")

    def _st_dropna(self, df, var, st, out):
        kw = {}
        sub = self.ocols(df, st.get("cols"))
        if sub:
            kw["subset"] = sub
        if st.get("how", "any") != "any":
            kw["how"] = st["how"]
        self.emit(f"{var} = {call(f'{var}.dropna', **kw)}")
        return df.dropna(subset=sub or None, how=st.get("how", "any"))

    def _st_dedupe(self, df, var, st, out):
        kw = {}
        sub = self.ocols(df, st.get("cols"))
        if sub:
            kw["subset"] = sub
        if st.get("keep", "first") != "first":
            kw["keep"] = st["keep"]
        self.emit(f"{var} = {call(f'{var}.drop_duplicates', **kw)}")
        return df.drop_duplicates(subset=sub or None, keep=st.get("keep", "first"))

    def _headtail(self, df, var, st, which):
        by = self.ocols(df, st.get("by"))
        n = int(st.get("n", 5))
        if by:
            self.emit(f"{var} = {var}.groupby({lit(by)}, sort=False).{which}({n})")
            return getattr(df.groupby(by, sort=False), which)(n)
        self.emit(f"{var} = {var}.{which}({n})")
        return getattr(df, which)(n)

    def _st_head(self, df, var, st, out):
        return self._headtail(df, var, st, "head")

    def _st_tail(self, df, var, st, out):
        return self._headtail(df, var, st, "tail")

    def _agg_code(self, spec: Dict) -> str:
        items = []
        for k, (c, fn) in spec.items():
            if isinstance(fn, str):
                items.append((k, f"({c!r}, {fn!r})"))
            else:
                q = float(re.search(r"(\d+(?:\.\d+)?)$", k).group(1)) / 100.0
                items.append((k, f"({c!r}, lambda s: s.quantile({q!r}))"))
        if all(k.isidentifier() and not keyword.iskeyword(k) for k, _ in items):
            return ", ".join(f"{k}={v}" for k, v in items)
        return "**{" + ", ".join(f"{k!r}: {v}" for k, v in items) + "}"

    def _st_groupby(self, df, var, st, out):
        from .figdata import _aggspec
        by = self.cols(df, st.get("by"))
        spec = _aggspec(df, st.get("agg"), by)
        self.emit(f"groups = {var}.groupby({lit(by)}, sort=True, dropna=False)")
        self.emit(f"{var} = groups.agg({self._agg_code(spec)})")
        self.emit(f'{var}["n"] = groups.size()')
        self.emit(f"{var} = {var}.reset_index()")
        return out

    def _st_pivot(self, df, var, st, out):
        idx = self.cols(df, st.get("index"))
        col = self.c(df, st["columns"])
        vals = self.cols(df, st.get("values"))
        agg = st.get("agg", "mean")
        self.emit(f"{var} = {var}.pivot_table(index={lit(idx)}, columns={col!r}, values={lit(vals)}, "
                  f"aggfunc={lit(agg)}, dropna=False, sort=True)")
        new = [c for c in out.columns if c not in idx]
        self.emit(f"{var}.columns = {lit(new, 4)}   # one column per {col!r} value")
        self.emit(f"{var} = {var}.reset_index()")
        return out

    def _st_melt(self, df, var, st, out):
        kw = {"id_vars": self.ocols(df, st.get("id_vars"))}
        vv = self.ocols(df, st.get("value_vars"))
        if vv:
            kw["value_vars"] = vv
        kw["var_name"] = st.get("var_name", "variable")
        kw["value_name"] = st.get("value_name", "value")
        self.emit(f"{var} = {call(f'{var}.melt', **kw)}")
        return out

    def _st_join(self, df, var, st, out):
        from .figures import find_col
        ovar, other = self.load(st["with"], None)
        if st.get("on") is not None:
            lk = self.cols(df, st["on"])
            rk = [find_col(other, c) for c in ([st["on"]] if isinstance(st["on"], str) else st["on"])]
        else:
            lk = self.cols(df, st.get("left_on"))
            rk = [find_col(other, c) for c in ([st["right_on"]] if isinstance(st.get("right_on"), str)
                                                else st.get("right_on") or [])]
        ren = {r: l for l, r in zip(lk, rk) if r != l and l not in other.columns}
        if ren:
            self.emit(f"{ovar} = {ovar}.rename(columns={lit(ren)})")
            other = other.rename(columns=ren)
        rk = [l if l in other.columns else r for l, r in zip(lk, rk)]
        suffixes = tuple(st.get("suffixes", ["", "_2"]))
        how = st.get("how", "inner")
        if how == "nearest":
            by = self.ocols(df, st.get("by")) or None
            self.emit(f"{var} = {var}.assign(**{{{lk[0]!r}: pd.to_numeric({var}[{lk[0]!r}]).astype(float)}})"
                      f".sort_values({lk[0]!r}, kind='stable')")
            self.emit(f"{ovar} = {ovar}.assign(**{{{rk[0]!r}: pd.to_numeric({ovar}[{rk[0]!r}]).astype(float)}})"
                      f".sort_values({rk[0]!r}, kind='stable')")
            keys = {"on": lk[0]} if lk[0] == rk[0] else {"left_on": lk[0], "right_on": rk[0]}
            tol = st.get("tolerance")
            self.emit(f"{var} = " + call("pd.merge_asof", var, ovar, by=by, suffixes=suffixes,
                                          direction=st.get("direction", "nearest"),
                                          tolerance=None if tol is None else float(tol), **keys)
                      + "   # nearest-key join")
            return out
        self.emit(f"{var} = {var}.merge({ovar}, left_on={lit(lk)}, right_on={lit(rk)}, how={how!r}, "
                  f"suffixes={lit(suffixes)})")
        drop = [r for l, r in zip(lk, rk) if r != l]
        if drop:
            self.emit(f"{var} = {var}.drop(columns={lit(drop)})")
        return out

    def _st_concat(self, df, var, st, out):
        import glob as _glob
        from .figdata import _is_spec
        items = st.get("with")
        if not isinstance(items, list) or (items and all(isinstance(i, dict) and not _is_spec(i) for i in items)):
            items = [items]
        srcs, names = [], []
        for it in items:
            if isinstance(it, str) and any(ch in it for ch in "*?["):
                for h in sorted(_glob.glob(self.tb.p(it))):
                    srcs.append(h)
                    names.append(Path(h).stem)
            else:
                srcs.append(it)
                names.append(Path(it).stem if isinstance(it, str) else
                             Path(str(it.get("data", it.get("csv", it.get("path", ""))))).stem
                             if _is_spec(it) else f"data{len(names) + 2}")
        ovars = [self.load(s, None)[0] for s in srcs]
        labels = st.get("labels") or [self._src_name] + names
        frames = [var] + ovars
        if st.get("label_col"):
            frames = [f"{f}.assign(**{{{st['label_col']!r}: {lab!r}}})" for f, lab in zip(frames, labels)]
        self.emit(f"{var} = pd.concat([{', '.join(frames)}], ignore_index=True, sort=False)")
        return out

    def _grid_code(self, onto, xs: str) -> str:
        if isinstance(onto, (list, tuple)):
            return f"np.asarray({lit(list(onto))}, dtype=float)"
        if isinstance(onto, (int, float)):
            onto = {"num": onto}
        lo = lit(float(onto["start"])) if "start" in onto else f"np.nanmin({xs})"
        hi = lit(float(onto["stop"])) if "stop" in onto else f"np.nanmax({xs})"
        if onto.get("step"):
            s = float(onto["step"])
            return f"{lo} + {s!r} * np.arange(int(np.floor(({hi} - {lo}) / {s!r} + 1e-9)) + 1)"
        return f"np.linspace({lo}, {hi}, {int(onto.get('num', 50))})"

    def _st_interpolate(self, df, var, st, out):
        x = self.c(df, st["x"])
        ys = self.cols(df, st.get("y"))
        by = self.ocols(df, st.get("by"))
        method = str(st.get("method", "linear")).lower()
        extra = bool(st.get("extrapolate", False))
        h = helper_code(self.ctx, ["interp1"])
        if h:
            self.pre.append(h)
        grid = self._grid_code(st.get("onto", {"num": 50}), "xv")
        self.emit("parts = []")
        if by:
            self.emit(f"for keys, g in {var}.groupby({lit(by)}, sort=True, dropna=False):")
            ind, g = "    ", "g"
        else:
            ind, g = "", var
        self.emit(f"{ind}xv = pd.to_numeric({g}[{x!r}]).to_numpy(float)")
        self.emit(f"{ind}xn = {grid}")
        cols = (f"**dict(zip({lit(by)}, keys)), " if by else "") + f"{x!r}: xn"
        for y in ys:
            cols += f", {y!r}: interp1(xv, pd.to_numeric({g}[{y!r}]).to_numpy(float), xn, {method!r}, {extra!r})"
        self.emit(f"{ind}parts.append(pd.DataFrame({{{cols}}}))")
        self.emit(f"{var} = pd.concat(parts, ignore_index=True)")
        return out

    def _out_names(self, cols, st, default_suffix=""):
        if st.get("new") is not None:
            return [st["new"]] if isinstance(st["new"], str) else list(st["new"])
        return [c + st.get("suffix", default_suffix) for c in cols]

    def _st_smooth(self, df, var, st, out):
        cols = self.cols(df, st.get("cols", st.get("col")))
        by = self.ocols(df, st.get("by"))
        method = str(st.get("method", "rolling")).lower()
        w = int(st.get("window", 5))
        if st.get("x"):
            self.emit(f"{var} = {var}.sort_values({lit(by + [self.c(df, st['x'])])}, kind='stable')")
        self.emit(f"{var} = {var}.reset_index(drop=True)")
        if method == "rolling":
            f = f"s.rolling({w}, center=True, min_periods=1).mean()"
        elif method == "ewm":
            f = f"s.ewm(span={w}).mean()"
        else:
            h = helper_code(self.ctx, ["savgol_smooth"])
            if h:
                self.pre.append(h)
            f = f"pd.Series(savgol_smooth(s.to_numpy(float), {w}, {int(st.get('order', 2))}), index=s.index)"
        for c, o in zip(cols, self._out_names(cols, st)):
            s = f"pd.to_numeric({var}[{c!r}]).astype(float)"
            if by:
                self.emit(f"{var}[{o!r}] = {s}.groupby([{', '.join(f'{var}[{b!r}]' for b in by)}], dropna=False)"
                          f".transform(lambda s: {f})")
            else:
                self.emit(f"{var}[{o!r}] = (lambda s: {f})({s})")
        return out

    def _st_derivative(self, df, var, st, out):
        x, y = self.c(df, st["x"]), self.c(df, st["y"])
        by = self.ocols(df, st.get("by"))
        new = st.get("new") or f"d{y}_d{x}"
        self.emit(f"{var} = {var}.sort_values({lit(by + [x])}, kind='stable').reset_index(drop=True)")
        self.emit(f"{var}[{new!r}] = np.nan")
        body = [f"if len(g) >= 2:",
                f"    {var}.loc[g.index, {new!r}] = np.gradient(pd.to_numeric(g[{y!r}]).to_numpy(float), "
                f"pd.to_numeric(g[{x!r}]).to_numpy(float), edge_order=2 if len(g) >= 3 else 1)"]
        if by:
            self.emit(f"for _, g in {var}.groupby({lit(by)}, sort=True, dropna=False):")
            for b in body:
                self.emit("    " + b)
        else:
            self.emit(f"g = {var}")
            for b in body:
                self.emit(b)
        return out

    def _st_integrate(self, df, var, st, out):
        self.ctx.need("from scipy.integrate import cumulative_trapezoid")
        x, y = self.c(df, st["x"]), self.c(df, st["y"])
        by = self.ocols(df, st.get("by"))
        new = st.get("new") or f"int_{y}"
        self.emit(f"{var} = {var}.sort_values({lit(by + [x])}, kind='stable').reset_index(drop=True)")
        yv, xv = f"pd.to_numeric(g[{y!r}]).to_numpy(float)", f"pd.to_numeric(g[{x!r}]).to_numpy(float)"
        if not st.get("cumulative", True):
            if by:
                self.emit(f"{var} = pd.DataFrame([{{**dict(zip({lit(by)}, keys)), {new!r}: np.trapezoid({yv}, {xv})}}"
                          f" for keys, g in {var}.groupby({lit(by)}, sort=True, dropna=False)])")
            else:
                self.emit(f"g = {var}")
                self.emit(f"{var} = pd.DataFrame([{{{new!r}: np.trapezoid({yv}, {xv})}}])")
            return out
        self.emit(f"{var}[{new!r}] = np.nan")
        line = f"{var}.loc[g.index, {new!r}] = cumulative_trapezoid({yv}, {xv}, initial=0.0)"
        if by:
            self.emit(f"for _, g in {var}.groupby({lit(by)}, sort=True, dropna=False):")
            self.emit("    " + line)
        else:
            self.emit(f"g = {var}")
            self.emit(line)
        return out

    def _st_normalize(self, df, var, st, out):
        cols = self.cols(df, st.get("cols", st.get("col")))
        by = self.ocols(df, st.get("by"))
        method = str(st.get("method", "max")).lower()
        f = {"max": "s / s.abs().max()", "minmax": "(s - s.min()) / (s.max() - s.min())",
             "zscore": "(s - s.mean()) / s.std()", "first": "s / s.iloc[0]", "sum": "s / s.sum()"}[method]
        self.emit(f"{var} = {var}.reset_index(drop=True)")
        for c, o in zip(cols, self._out_names(cols, st)):
            s = f"pd.to_numeric({var}[{c!r}]).astype(float)"
            if by:
                self.emit(f"{var}[{o!r}] = {s}.groupby([{', '.join(f'{var}[{b!r}]' for b in by)}], dropna=False)"
                          f".transform(lambda s: {f})")
            else:
                self.emit(f"{var}[{o!r}] = (lambda s: {f})({s})")
        return out

    def _st_bin(self, df, var, st, out):
        from .figdata import _aggspec
        c = self.c(df, st["col"])
        by = self.ocols(df, st.get("by"))
        bins = st.get("bins", 10)
        xv = self.num(df, var, c)
        if isinstance(bins, (int, float)):
            self.emit(f"xv = {xv}")
            self.emit(f"edges = np.linspace(np.nanmin(xv), np.nanmax(xv), {int(bins) + 1})")
        else:
            self.emit(f"edges = np.asarray({lit(sorted(float(v) for v in bins))})")
        spec = _aggspec(df, st.get("agg"), by + [c] if st.get("agg") is None else by)
        label = st.get("label", "center")
        self.emit(f"codes = pd.cut({xv}, bins=edges, labels=False, include_lowest=True)")
        self.emit(f"groups = {var}.assign(__bin__=codes).dropna(subset=['__bin__']).groupby({lit(by + ['__bin__'])}, "
                  f"sort=True)")
        self.emit(f"{var} = groups.agg({self._agg_code(spec)})")
        self.emit(f'{var}["n"] = groups.size()')
        self.emit(f"{var} = {var}.reset_index()")
        self.emit(f"i = {var}['__bin__'].astype(int).to_numpy()")
        if label == "interval":
            self.emit(f"{var}['__bin__'] = [f\"{{'[' if j == 0 else '('}}{{a:.10g}}, {{b:.10g}}]\" "
                      f"for a, b, j in zip(edges[i], edges[i + 1], i)]")
        else:
            self.emit(f"{var}['__bin__'] = " + {"center": "(edges[i] + edges[i + 1]) / 2", "left": "edges[i]",
                                                "right": "edges[i + 1]"}[label] + f"   # bin {label}")
        self.emit(f"{var} = {var}.rename(columns={{'__bin__': {c!r}}})[{lit(list(out.columns), 4)}]")
        return out

    def _st_element_fractions(self, df, var, st, out):
        from . import figdata as fd
        given = {str(e).strip().upper(): self.c(df, c) for c, e in st["cols"].items()}
        els = [str(e).strip().upper() for e in (st.get("elements") or list(given))]
        basis, f = fd.COMP[fd._norm_unit(st.get("unit", "mole_fraction"))]
        self.emit(f"frac = {{{', '.join(f'{e!r}: pd.to_numeric({var}[{c!r}]).astype(float) * {f!r}' for e, c in given.items())}}}"
                  f"   # {st.get('unit', 'mole_fraction')} -> fractions")
        missing = [e for e in els if e not in given]
        if missing:
            self.emit(f"frac[{missing[0]!r}] = 1.0 - sum(frac.values())   # balance element")
        if basis == "mass":
            self.emit(f"masses = {lit({e: fd._mass(e) for e in els})}   # g/mol")
            self.emit("amount = {el: w / masses[el] for el, w in frac.items()}   # mass fractions -> moles")
            self.emit("frac = {el: n / sum(amount.values()) for el, n in amount.items()}")
        self.emit(f"for el in {lit(els)}:\n    {var}[f'X({{el}})'] = frac[el]   # mole fractions by element name")
        return out

    def _st_convert(self, df, var, st, out):
        from . import figdata as fd
        cols = self.cols(df, st.get("cols", st.get("col")))
        fu, tu = fd._norm_unit(st["from"]), fd._norm_unit(st["to"])
        outs = self._out_names(cols, st)
        if fu in fd.TEMP or tu in fd.TEMP:
            a, b = fd.TEMP[fu], fd.TEMP[tu]
            to_k = {"K": "{v}", "C": "({v} + 273.15)", "F": "(({v} + 459.67) * 5.0 / 9.0)"}[a]
            from_k = {"K": "{v}", "C": "{v} - 273.15", "F": "{v} * 9.0 / 5.0 - 459.67"}[b]
            for c, o in zip(cols, outs):
                v = f"pd.to_numeric({var}[{c!r}]).astype(float)"
                self.emit(f"{var}[{o!r}] = {from_k.format(v=to_k.format(v=v))}   # {st['from']} -> {st['to']}")
            return out
        if fu in fd.COMP or tu in fd.COMP:
            fb, fs = fd.COMP.get(fu) or (None, fd.UNITS.get(fu, ("", None))[1])
            tb_, ts = fd.COMP.get(tu) or (None, fd.UNITS.get(tu, ("", None))[1])
            fb, tb_ = fb or tb_, tb_ or fb
            vals = [f"pd.to_numeric({var}[{c!r}]).astype(float) * {fs!r}" for c in cols]
            if fb == tb_:
                for o, v in zip(outs, vals):
                    self.emit(f"{var}[{o!r}] = {v} / {ts!r}")
                return out
            if st.get("element"):
                els = [st["element"]] if isinstance(st["element"], str) else list(st["element"])
            else:
                els = st.get("elements") or [fd._element(c) for c in cols]
            bal = st.get("balance") or st.get("other")
            masses = {str(e).upper(): fd._mass(e) for e in els}
            masses[str(bal).upper()] = fd._mass(bal)
            self.emit(f"masses = {lit(masses)}   # g/mol")
            self.emit(f"frac = {{{', '.join(f'{str(e).upper()!r}: {v}' for e, v in zip(els, vals))}}}"
                      f"   # {'mass' if fb == 'mass' else 'mole'} fractions")
            self.emit(f"frac[{str(bal).upper()!r}] = 1.0 - sum(frac.values())   # balance")
            op_ = "/" if fb == "mass" else "*"
            self.emit(f"amount = {{el: f {op_} masses[el] for el, f in frac.items()}}   "
                      f"# {'moles' if fb == 'mass' else 'mass'} per unit")
            self.emit("total = sum(amount.values())")
            for e, o in zip(els, outs):
                self.emit(f"{var}[{o!r}] = amount[{str(e).upper()!r}] / total / {ts!r}")
            if st.get("balance_col"):
                self.emit(f"{var}[{st['balance_col']!r}] = amount[{str(bal).upper()!r}] / total / {ts!r}")
            return out
        f = fd.UNITS[fu][1] / fd.UNITS[tu][1]
        for c, o in zip(cols, outs):
            self.emit(f"{var}[{o!r}] = pd.to_numeric({var}[{c!r}]).astype(float) * {f!r}   "
                      f"# {st['from']} -> {st['to']}")
        return out


# =========================================================================================================
# figure operations -> matplotlib code
# =========================================================================================================
_DRAW_STYLE = {"color", "marker", "markersize", "linestyle", "linewidth", "alpha", "label", "zorder", "markeredgecolor",
               "markerfacecolor", "markeredgewidth", "drawstyle", "fillstyle", "dashes"}
_SCATTER_STYLE = {"color", "marker", "s", "alpha", "label", "zorder", "edgecolor", "edgecolors", "facecolor",
                  "facecolors", "linewidths", "cmap", "vmin", "vmax"}
_TEXT_STYLE = {"color", "fontsize", "fontweight", "ha", "va", "rotation", "alpha", "zorder", "fontstyle", "family"}
_PATCH_STYLE = {"color", "edgecolor", "facecolor", "linewidth", "linestyle", "alpha", "label", "zorder", "fill",
                "hatch"}
_ALIASES = {"plot": "line"}


class OpCodeError(Exception):
    pass


class OpX:
    """Per-operation view: data source, values (as displayed = standalone data units), coords, style keys."""

    def __init__(self, w: "OpWriter", op: Dict, axm: Optional[AxModel], n: int):
        self.w, self.op, self.axm, self.n = w, op, axm, n
        self.coords = op.get("coords", "data")
        self.mode = op.get("units", "displayed") if self.coords == "data" else "data"
        self._src = None
        self._shadow = None

    # -- data -------------------------------------------------------------------------------------------
    @property
    def has_source(self) -> bool:
        return any(k in self.op for k in ("csv", "data"))

    def src(self):
        if self._src is None:
            self._src = self.w.data.for_op(self.op)
        return self._src

    def df(self) -> str:
        return self.src()[0]

    def frame(self):
        return self.src()[1]

    def colname(self, key: str) -> Optional[str]:
        from .figures import find_col
        col = self.op.get(f"{key}_col")
        return find_col(self.frame(), col) if col else None

    def column(self, name: str) -> Raw:
        return Raw(f"{self.df()}[{name!r}]")

    def raw(self, key: str):
        """Code (Raw) for a data column, the inline value, or None."""
        if self.op.get(f"{key}_col") and self.has_source:
            return self.column(self.colname(key))
        return self.op.get(key)

    def shadow(self):
        """The toolkit's own OpContext on a scratch axes like the target (for data-dependent decisions)."""
        if self._shadow is None:
            from matplotlib.figure import Figure
            from .figures import OpContext, ScaledFormatter
            fig = Figure()
            proj = None if self.axm is None or self.axm.proj == "rectilinear" else self.axm.proj
            if proj == "triangular":
                import pycalphad.plot.triangular  # noqa: F401
            ax = fig.add_subplot(projection=proj)
            if self.axm is not None:
                if self.axm.xt != (1.0, 0.0):
                    ax.xaxis.set_major_formatter(ScaledFormatter(*self.axm.xt))
                if self.axm.yt != (1.0, 0.0):
                    ax.yaxis.set_major_formatter(ScaledFormatter(*self.axm.yt))
                for k in ("x", "y"):
                    sc = getattr(self.axm, f"{k}scale")
                    if sc in ("linear", "log", "symlog", "logit") and proj is None:
                        getattr(ax, f"set_{k}scale")(sc)
            self._shadow = OpContext(fig, ax, dict(self.op), self.n, self.w.tb.p, [])
        return self._shadow

    # -- coordinates --------------------------------------------------------------------------------------
    def _t(self, axis: str):
        """(scale, offset) to turn this op's values into standalone (displayed) units, or None if unchanged."""
        if self.coords != "data" or self.mode != "data" or self.axm is None:
            return None
        t = self.axm.xt if axis == "x" else self.axm.yt
        return None if t == (1.0, 0.0) else t

    def conv(self, v, axis: str, delta: bool = False):
        t = self._t(axis)
        if v is None or t is None:
            return v
        s, o = t
        if isinstance(v, Raw):
            return Raw(f"({v} * {s!r})" if delta else f"({v} * {s!r} + {o!r})")
        arr = np.asarray(v, dtype=float) * s + (0.0 if delta else o)
        return _plain(arr)

    def X(self, key="x"):
        return self.conv(self.raw(key), "x")

    def Y(self, key="y"):
        return self.conv(self.raw(key), "y")

    def dX(self, key):
        return self.conv(self.raw(key), "x", delta=True)

    def dY(self, key):
        return self.conv(self.raw(key), "y", delta=True)

    @staticmethod
    def first(v):
        if isinstance(v, Raw):
            return Raw(f"float({v}.iloc[0])")
        v = _first(v)
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v

    def arr(self, v):
        """Code for a float numpy array of values (for arithmetic / masking)."""
        if isinstance(v, Raw):
            return Raw(f"{v}.to_numpy(float)")
        return Raw(f"np.atleast_1d(np.asarray({self.L(v)}, dtype=float))")

    def tr(self) -> Optional[Raw]:
        return {"axes": Raw("ax.transAxes"), "figure": Raw("fig.transFigure")}.get(self.coords)

    def pick(self, allowed) -> Dict:
        return {k: v for k, v in self.op.items() if k in allowed}

    def L(self, v) -> str:
        return lit(v, self.w.indent + 4)

    def args(self, *a, **kw) -> str:
        parts = [self.L(x) for x in a]
        k = kwargs(kw, self.w.indent + 4)
        return ", ".join(parts + ([k] if k else []))


class OpWriter:
    """Translate figure operations into code acting on ``fig`` and ``ax`` (a static model tracks the axes)."""

    def __init__(self, ctx: Ctx, fm: FigModel, mode: str, indent: int = 0, real_fig=None):
        self.ctx, self.tb, self.fm, self.mode = ctx, ctx.tb, fm, mode
        self.lines: List[str] = []
        self.pre: List[str] = []
        self.names: set = set()
        self.indent = indent
        self.data = DataWriter(ctx, self.emit, self.names, self.pre)
        self.cur_ax: Optional[str] = None
        self.touched: List[AxModel] = []
        self.real_fig = real_fig
        self.counters: Dict[str, int] = {}
        ctx.need("import numpy as np", "import matplotlib.pyplot as plt")

    # -- emission ----------------------------------------------------------------------------------------
    def emit(self, line: str = "") -> None:
        for ln in str(line).split("\n"):
            self.lines.append((" " * self.indent + ln) if ln else "")

    def helpers(self, *names: str) -> None:
        h = helper_code(self.ctx, names)
        if h:
            self.pre.append(h)

    def fresh(self, base: str) -> str:
        k = self.counters.get(base, 0) + 1
        self.counters[base] = k
        name = f"{base}{k}"
        self.names.add(name)
        return name

    def set_ax(self, axm: AxModel) -> None:
        if axm.expr != self.cur_ax:
            self.emit(f"ax = {axm.expr}")
            self.cur_ax = axm.expr

    def touch(self, axm: Optional[AxModel]) -> None:
        if axm is not None and axm not in self.touched:
            self.touched.append(axm)

    def add_top(self, axm: AxModel) -> AxModel:
        self.fm.top.append(axm)
        self.touch(axm)
        return axm

    # -- driver ------------------------------------------------------------------------------------------
    def resolve(self, op: Dict, default: Optional[AxModel]) -> Optional[AxModel]:
        main = self.fm.main()
        if "axes" not in op:
            return default
        idx = op["axes"]
        if isinstance(idx, str) and not idx.strip().lstrip("-").isdigit():
            hits = [a for a in main if a.title.lower() == idx.lower()]
            if not hits:
                raise OpCodeError(f"no axes titled {idx!r}")
            return hits[0]
        idx = int(idx)
        return main[idx] if 0 <= idx < len(main) else None

    def apply(self, ops: Sequence[Dict], default: Optional[AxModel]) -> None:
        from .figures import OP_NEEDS_AXES, ensure_extensions
        ensure_extensions()
        for n, op in enumerate(ops, 1):
            kind = _ALIASES.get(str(op.get("op", "")).lower(), str(op.get("op", "")).lower())
            fn = getattr(self, f"op_{kind}", None)
            if fn is None:
                raise OpCodeError(f"operation {n} ({kind}) has no standalone translation")
            try:
                axm = self.resolve(op, default)
                needs = OP_NEEDS_AXES.get(kind, kind != "figure")
                if needs and axm is None:
                    raise OpCodeError("figure has no such axes")
                if axm is not None and needs:
                    self.set_ax(axm)
                    self.touch(axm)
                fn(OpX(self, op, axm, n))
            except OpCodeError:
                raise
            except Exception as e:  # noqa: BLE001 - reported as the reason for the calphad_tools fallback
                raise OpCodeError(f"operation {n} ({kind}): {type(e).__name__}: {e}") from None

    def refresh_legends(self, axes: Sequence[AxModel], fallback: Optional[AxModel] = None,
                        fallback_fontsize=None) -> None:
        """The toolkit's rule after a batch of operations: unlocked axes get newly labelled artists added."""
        todo = [a for a in axes if not a.locked and a in self.fm.main()]
        if not todo:
            return
        self.helpers("update_legend")
        for k, a in enumerate(todo):
            tail = "   # legend: existing entries + newly labelled artists" if k == 0 else ""
            if a is fallback and fallback_fontsize is not None:
                self.emit(f"update_legend({a.expr}, fallback_fontsize={lit(fallback_fontsize)}){tail}")
            else:
                self.emit(f"update_legend({a.expr}){tail}")

    # -- shared pieces -----------------------------------------------------------------------------------
    def norm(self, spec) -> Raw:
        self.ctx.need("import matplotlib.colors as mcolors")
        if isinstance(spec, str):
            spec = {"type": spec}
        t = spec.get("type", "linear").lower()
        vmin, vmax = spec.get("vmin"), spec.get("vmax")
        if t == "linear":
            return Raw(f"mcolors.Normalize({lit(vmin)}, {lit(vmax)})")
        if t == "log":
            return Raw(f"mcolors.LogNorm({lit(vmin)}, {lit(vmax)})")
        if t == "symlog":
            return Raw(f"mcolors.SymLogNorm(linthresh={lit(spec.get('linthresh', 1.0))}, vmin={lit(vmin)}, "
                       f"vmax={lit(vmax)})")
        if t == "power":
            return Raw(f"mcolors.PowerNorm({lit(spec.get('gamma', 0.5))}, {lit(vmin)}, {lit(vmax)})")
        if t in ("diverging", "twoslope", "center"):
            return Raw(f"mcolors.TwoSlopeNorm(vcenter={lit(spec.get('center', 0.0))}, vmin={lit(vmin)}, "
                       f"vmax={lit(vmax)})")
        if t in ("bounds", "discrete", "boundary"):
            return Raw(f"mcolors.BoundaryNorm({lit(spec['boundaries'])}, ncolors={lit(spec.get('ncolors', 256))}, "
                       f"extend={lit(spec.get('extend', 'neither'))})")
        raise OpCodeError(f"unknown norm {t!r}")

    def cmap_kw(self, o: OpX, allow_colors: bool = False) -> Dict:
        op = o.op
        if allow_colors and op.get("colors") is not None:
            return {"colors": op["colors"]}
        kw: Dict[str, Any] = {"cmap": op.get("cmap", "viridis")}
        if op.get("norm") is not None:
            spec = op["norm"]
            if isinstance(spec, dict):
                spec = {"vmin": op.get("vmin"), "vmax": op.get("vmax"), **spec}
            elif isinstance(spec, str) and (op.get("vmin") is not None or op.get("vmax") is not None):
                spec = {"type": spec, "vmin": op.get("vmin"), "vmax": op.get("vmax")}
            kw["norm"] = self.norm(spec)
        else:
            for k in ("vmin", "vmax"):
                if op.get(k) is not None:
                    kw[k] = op[k]
        return kw

    def colorbar(self, o: OpX, mappable: str, ax_expr: str = "ax", spec=None, three_d: bool = False) -> None:
        spec = o.op.get("colorbar") if spec is None else spec
        if not spec:
            return
        spec = spec if isinstance(spec, dict) else {}
        if three_d:
            spec = {"pad": 0.12, "shrink": 0.7, **spec}
        kw = {k: spec[k] for k in ("extend", "orientation", "shrink", "location", "pad", "aspect") if k in spec}
        self.emit(f"cb = fig.colorbar({mappable}, ax={ax_expr}{', ' + kwargs(kw) if kw else ''})")
        label = spec.get("label", o.op.get("colorbar_label"))
        if label:
            self.emit(f"cb.set_label({label!r})")
        if spec.get("ticks") is not None:
            self.emit(f"cb.set_ticks({lit(spec['ticks'])})")

    def formatter(self, spec, op: Dict, scaled: bool = False) -> str:
        """Code for a tick formatter (sets up ``fmt`` when it needs statements)."""
        from .figops_extra import _pyfmt
        self.ctx.need("from matplotlib import ticker as mticker")
        s, low = str(spec), str(spec).lower()
        if scaled:           # the toolkit keeps its display transform; standalone values are already displayed
            if "{" in s:
                return f"mticker.StrMethodFormatter({_pyfmt(s).replace('{', '{x', 1)!r})"
            if low in ("plain", "g"):
                return "mticker.StrMethodFormatter('{x:g}')"
            if low == "sci":
                return "mticker.StrMethodFormatter('{x:.2e}')"
            raise OpCodeError(f"format {spec!r} on a transformed axis")
        if "{" in s:
            return f"mticker.StrMethodFormatter({(s if '{x' in s else _pyfmt(s).replace('{', '{x', 1))!r})"
        if "%" in s and low != "percent":
            return f"mticker.FormatStrFormatter({s!r})"
        if low == "sci":
            self.emit("fmt = mticker.ScalarFormatter(useMathText=True)")
            self.emit("fmt.set_powerlimits((0, 0))")
            return "fmt"
        if low == "plain":
            self.emit("fmt = mticker.ScalarFormatter(useOffset=False)")
            self.emit("fmt.set_scientific(False)")
            return "fmt"
        if low == "percent":
            return f"mticker.PercentFormatter(xmax={lit(op.get('xmax', 1.0))}, decimals={lit(op.get('decimals'))})"
        if low == "eng":
            return f"mticker.EngFormatter(unit={op.get('unit', '')!r}, places={lit(op.get('decimals'))})"
        if low == "log":
            return "mticker.LogFormatterSciNotation()"
        raise OpCodeError(f"unknown format {spec!r}")

    # =====================================================================================================
    # built-in operations (calphad_tools.figures.apply_ops)
    # =====================================================================================================
    def op_scatter(self, o: OpX):
        op = o.op
        kw = o.pick(_SCATTER_STYLE)
        c = o.raw("c") if (op.get("c_col") and o.has_source) else op.get("c")
        if c is not None:
            kw["c"] = c
            kw.pop("color", None)
        if op.get("s_col") and o.has_source:
            kw["s"] = o.raw("s")
        if "norm" in op:
            kw["norm"] = self.norm(op["norm"])
            kw.pop("vmin", None)
            kw.pop("vmax", None)
        if o.tr():
            kw["transform"] = o.tr()
        code = f"ax.scatter({o.args(o.X(), o.Y(), **kw)})"
        if c is not None and op.get("colorbar"):
            self.emit(f"sc = {code}")
            self.colorbar(o, "sc")
        else:
            self.emit(code)

    def op_line(self, o: OpX):
        kw = o.pick(_DRAW_STYLE)
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"ax.plot({o.args(o.X(), o.Y(), **kw)})")

    def op_step(self, o: OpX):
        kw = {"where": o.op.get("where_step", "mid"), **o.pick(_DRAW_STYLE)}
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"ax.step({o.args(o.X(), o.Y(), **kw)})")

    def op_errorbar(self, o: OpX):
        kw = o.pick(_DRAW_STYLE)
        kw.setdefault("linestyle", "none")
        kw.setdefault("marker", "o")
        extra = {}
        if o.raw("yerr") is not None:
            extra["yerr"] = o.dY("yerr")
        if o.raw("xerr") is not None:
            extra["xerr"] = o.dX("xerr")
        extra["capsize"] = o.op.get("capsize", 3)
        if o.tr():
            extra["transform"] = o.tr()
        self.emit(f"ax.errorbar({o.args(o.X(), o.Y(), **extra, **kw)})")

    def op_fill_between(self, o: OpX):
        kw = o.pick(_PATCH_STYLE - {"fill"})
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"ax.fill_between({o.args(o.X(), o.Y('y1'), o.Y('y2'), **kw)})")

    def op_bar(self, o: OpX):
        kw = {"width": o.op.get("width", 0.8), **o.pick(_PATCH_STYLE - {"fill"})}
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"ax.bar({o.args(o.X(), o.dY('height'), **kw)})")

    def op_barh(self, o: OpX):
        kw = {"height": o.op.get("height", 0.8), **o.pick(_PATCH_STYLE - {"fill"})}
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"ax.barh({o.args(o.Y(), o.dX('width'), **kw)})")

    def op_text(self, o: OpX):
        kw = o.pick(_TEXT_STYLE)
        if o.op.get("bbox"):
            kw["bbox"] = {"boxstyle": "round,pad=0.25", "facecolor": "white", "alpha": 0.8, "edgecolor": "0.6"}
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"ax.text({o.args(o.first(o.X()), o.first(o.Y()), str(o.op['text']), **kw)})")

    def _pt(self, o: OpX, xy):
        if o.coords != "data":
            return [float(xy[0]), float(xy[1])]
        return [o.first(o.conv([xy[0]], "x")), o.first(o.conv([xy[1]], "y"))]

    def op_annotate(self, o: OpX):
        op = o.op
        p = self._pt(o, op["xy"])
        pt = self._pt(o, op.get("xytext", op["xy"]))
        kw = o.pick(_TEXT_STYLE)
        if op.get("arrow", True):
            kw["arrowprops"] = {"arrowstyle": op.get("arrowstyle", "->"), "color": op.get("color", "black"),
                                "lw": op.get("linewidth", 1.2)}
        tr = o.tr()
        co = {"xycoords": tr, "textcoords": tr} if tr else {}
        self.emit(f"ax.annotate({o.args(str(op['text']), xy=tuple(p), xytext=tuple(pt), **co, **kw)})")

    def op_arrow(self, o: OpX):
        op = o.op
        p0 = self._pt(o, [op["x"], op["y"]])
        p1 = self._pt(o, [op["x"] + op["dx"], op["y"] + op["dy"]])
        tr = o.tr()
        co = {"xycoords": tr, "textcoords": tr} if tr else {}
        props = {"arrowstyle": op.get("arrowstyle", "->"), "color": op.get("color", "black"),
                 "lw": op.get("linewidth", 1.5)}
        self.emit(f"ax.annotate({o.args('', xy=tuple(p1), xytext=tuple(p0), **co, arrowprops=props)})")

    def op_hline(self, o: OpX):
        kw = {}
        if o.op.get("xmin", 0) != 0:
            kw["xmin"] = o.op["xmin"]
        if o.op.get("xmax", 1) != 1:
            kw["xmax"] = o.op["xmax"]
        self.emit(f"ax.axhline({o.args(o.first(o.Y()), **kw, **o.pick(_DRAW_STYLE))})")

    def op_vline(self, o: OpX):
        kw = {}
        if o.op.get("ymin", 0) != 0:
            kw["ymin"] = o.op["ymin"]
        if o.op.get("ymax", 1) != 1:
            kw["ymax"] = o.op["ymax"]
        self.emit(f"ax.axvline({o.args(o.first(o.X()), **kw, **o.pick(_DRAW_STYLE))})")

    def op_hspan(self, o: OpX):
        self.emit(f"ax.axhspan({o.args(o.first(o.Y('y1')), o.first(o.Y('y2')), **o.pick(_PATCH_STYLE - {'fill'}))})")

    def op_vspan(self, o: OpX):
        self.emit(f"ax.axvspan({o.args(o.first(o.X('x1')), o.first(o.X('x2')), **o.pick(_PATCH_STYLE - {'fill'}))})")

    def _patch(self, o: OpX, cls: str, *args):
        self.ctx.need(f"from matplotlib.patches import {cls}")
        kw = o.pick(_PATCH_STYLE)
        kw.setdefault("fill", "facecolor" in o.op)
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"ax.add_patch({cls}({o.args(*args, **kw)}))")

    def op_rectangle(self, o: OpX):
        self._patch(o, "Rectangle", (o.first(o.X()), o.first(o.Y())), o.first(o.dX("width")), o.first(o.dY("height")))

    def op_circle(self, o: OpX):
        self._patch(o, "Circle", (o.first(o.X()), o.first(o.Y())), float(o.op["radius"]))

    def op_ellipse(self, o: OpX):
        self._patch(o, "Ellipse", (o.first(o.X()), o.first(o.Y())), o.first(o.dX("width")), o.first(o.dY("height")))

    def op_polygon(self, o: OpX):
        pts = np.asarray(o.op["points"], dtype=float)
        if o.coords == "data":
            pts = np.column_stack([np.asarray(o.conv(list(pts[:, 0]), "x")), np.asarray(o.conv(list(pts[:, 1]), "y"))])
        self._patch(o, "Polygon", _plain(pts), )

    def op_set(self, o: OpX):
        op, axm = o.op, o.axm
        if "xlim" in op:
            self.emit(f"ax.set_xlim({o.args(*o.conv(list(op['xlim']), 'x'))})")
        if "ylim" in op:
            self.emit(f"ax.set_ylim({o.args(*o.conv(list(op['ylim']), 'y'))})")
        for k in ("xlabel", "ylabel"):
            if k in op:
                fs = {"fontsize": op["label_fontsize"]} if op.get("label_fontsize") is not None else {}
                self.emit(f"ax.set_{k}({o.args(op[k], **fs)})")
                setattr(axm, k, op[k])
        if "title" in op:
            fs = {"fontsize": op["title_fontsize"]} if op.get("title_fontsize") is not None else {}
            self.emit(f"ax.set_title({o.args(op['title'], **fs)})")
            axm.title = op["title"]
        for k in ("xscale", "yscale"):
            if k in op:
                self.emit(f"ax.set_{k}({op[k]!r})")
                setattr(axm, k, op[k])
        if "grid" in op:
            self.emit(f"ax.grid({bool(op['grid'])}, alpha=0.35)")
        if "aspect" in op:
            self.emit(f"ax.set_aspect({lit(op['aspect'])})")
        if op.get("invert_x"):
            self.emit("ax.invert_xaxis()")
        if op.get("invert_y"):
            self.emit("ax.invert_yaxis()")
        if "xticks" in op:
            self.emit(f"ax.set_xticks({o.L(o.conv(list(op['xticks']), 'x'))})")
        if "yticks" in op:
            self.emit(f"ax.set_yticks({o.L(o.conv(list(op['yticks']), 'y'))})")
        if "xticklabels" in op:
            self.emit(f"ax.set_xticklabels({o.L(op['xticklabels'])})")
        if "yticklabels" in op:
            self.emit(f"ax.set_yticklabels({o.L(op['yticklabels'])})")
        if "tick_fontsize" in op:
            self.emit(f"ax.tick_params(labelsize={lit(op['tick_fontsize'])})")
        if "facecolor" in op:
            self.emit(f"ax.set_facecolor({lit(op['facecolor'])})")

    def op_legend(self, o: OpX):
        op = o.op
        if op.get("show", True) is False:
            self.emit("if ax.get_legend() is not None:")
            self.emit("    ax.get_legend().remove()")
        else:
            kw = {k: op[k] for k in ("loc", "ncol", "fontsize", "title", "frameon", "framealpha") if k in op}
            if op.get("outside"):
                kw.update(loc="center left", bbox_to_anchor=(1.02, 0.5))
            self.helpers("update_legend")
            self.emit(f"update_legend(ax, only_if_new=False{', ' + kwargs(kw) if kw else ''})")
        o.axm.locked = True

    def op_twinx(self, o: OpX):
        name = self.fresh("twin")
        self.emit(f"{name} = ax.twinx()   # twin y-axis")
        if o.op.get("ylabel"):
            self.emit(f"{name}.set_ylabel({o.op['ylabel']!r})")
        if o.op.get("ylim"):
            self.emit(f"{name}.set_ylim({o.args(*o.op['ylim'])})")
        self.add_top(AxModel(name, "rectilinear", "twin", o.axm.xscale, "linear", xt=o.axm.xt,
                             spec=o.axm.spec, bounds=o.axm.bounds))

    def op_inset(self, o: OpX):
        op = o.op
        name = self.fresh("inset")
        self.emit(f"{name} = ax.inset_axes({lit(op.get('bounds', [0.55, 0.55, 0.4, 0.4]))})")
        self.emit("for ln in ax.get_lines():   # zoomed copy of the lines and points")
        self.emit(f"    {name}.plot(ln.get_xdata(), ln.get_ydata(), color=ln.get_color(), linestyle=ln.get_linestyle(),")
        self.emit("                linewidth=ln.get_linewidth(), marker=ln.get_marker(), markersize=ln.get_markersize())")
        self.emit("for c in ax.collections:")
        self.emit("    if hasattr(c, 'get_offsets') and len(c.get_offsets()):")
        self.emit("        off = c.get_offsets()")
        self.emit(f"        {name}.scatter(off[:, 0], off[:, 1], s=12, "
                  f"color=c.get_facecolor()[:1] if len(c.get_facecolor()) else None)")
        if op.get("xlim"):
            self.emit(f"{name}.set_xlim({o.args(*o.conv(list(op['xlim']), 'x'))})")
        if op.get("ylim"):
            self.emit(f"{name}.set_ylim({o.args(*o.conv(list(op['ylim']), 'y'))})")
        self.emit(f"{name}.tick_params(labelsize=7)")
        self.emit(f"ax.indicate_inset_zoom({name}, edgecolor='0.4')")
        kid = AxModel(name, "rectilinear", "inset", o.axm.xscale, o.axm.yscale)
        o.axm.children.append(kid)
        self.touch(kid)

    def op_style(self, o: OpX):
        op = o.op
        props = {k: op[k] for k in ("color", "linewidth", "linestyle", "alpha", "marker", "markersize", "zorder",
                                    "visible", "edgecolor", "facecolor") if k in op}
        self.helpers("find_artists")
        self.emit(f"for artist in find_artists(ax, {op['match']!r}):")
        if props:
            self.emit(f"    for prop, value in {lit(props)}.items():")
            self.emit("        if hasattr(artist, 'set_' + prop):")
            self.emit("            getattr(artist, 'set_' + prop)(value)")
        if "new_label" in op:
            self.emit(f"    artist.set_label({op['new_label']!r})")
        if not props and "new_label" not in op:
            self.emit("    pass")

    def op_remove(self, o: OpX):
        self.helpers("find_artists")
        self.emit(f"for artist in find_artists(ax, {o.op['match']!r}):")
        self.emit("    artist.remove()")

    def op_figure(self, o: OpX):
        op = o.op
        if op.get("size"):
            self.emit(f"fig.set_size_inches({o.args(*op['size'])})")
        if op.get("suptitle") is not None:
            self.emit(f"fig.suptitle({op['suptitle']!r}, fontsize={lit(op.get('fontsize', 14))})")
        if op.get("dpi"):
            self.emit(f"fig.set_dpi({lit(op['dpi'])})")
        if op.get("tight_layout"):
            self.emit("fig.tight_layout()")

    # =====================================================================================================
    # fits (calphad_tools.figfit): scipy.optimize.curve_fit in the units shown on the axes
    # =====================================================================================================
    def _custom_expr(self, expr: str, var: str, pnames: Sequence[str]) -> str:
        funcs = {"exp", "log", "ln", "log10", "sqrt", "abs", "sin", "cos", "tan", "tanh", "sinh", "cosh", "arctan"}
        npname = {"ln": "log"}

        def c(nd):
            if isinstance(nd, ast.Constant):
                return repr(float(nd.value))
            if isinstance(nd, ast.Name):
                if nd.id == "pi" and nd.id not in pnames and nd.id != var:
                    return "np.pi"
                if nd.id == "R" and nd.id not in pnames and nd.id != var:
                    return "8.314462618"
                return nd.id
            if isinstance(nd, ast.BinOp):
                return f"({c(nd.left)} {_PY_OPS[type(nd.op)]} {c(nd.right)})"
            if isinstance(nd, ast.UnaryOp):
                return f"(-{c(nd.operand)})" if isinstance(nd.op, ast.USub) else c(nd.operand)
            if isinstance(nd, ast.Call) and nd.func.id in funcs:
                return f"np.{npname.get(nd.func.id, nd.func.id)}({c(nd.args[0])})"
            raise OpCodeError(f"cannot translate {type(nd).__name__} in the fit formula")
        return _strip_parens(c(ast.parse(expr.strip(), mode="eval").body))

    def _model(self, o: OpX, res, tunit: str) -> Dict[str, Any]:
        """Code pieces for the toolkit's fit model: value function, ln function, equation / report text."""
        op = o.op
        spec = op.get("model", "linear")
        if isinstance(spec, dict):
            op = {**op, **spec, "model": "custom"}
        m = res.model
        name = m.name
        offset = bool(op.get("offset"))
        pn = list(m.pnames)
        var = m.xvar if name == "custom" else "x"
        T = {"K": "x", "C": "(x + 273.15)", "°C": "(x + 273.15)", "1/K": "(1.0 / x)", "1000/K": "(1000.0 / x)",
             "1e3/K": "(1000.0 / x)", "10000/K": "(10000.0 / x)", "1e4/K": "(10000.0 / x)"}.get(tunit.replace(" ", ""))
        plus = " + c" if offset else ""
        pc = ' + ((" − " if p[-1] < 0 else " + ") + fmt4(abs(p[-1])))' if offset else ""
        d: Dict[str, Any] = {"pnames": pn, "var": var, "logf": None, "report": None, "doc": ""}
        if name == "linear":
            d.update(f="intercept + slope * x", eq='"y = " + poly_str(p)', doc="y = intercept + slope·x")
        elif name.startswith("poly"):
            d.update(f="np.polynomial.polynomial.polyval(x, c)", star=True, eq='"y = " + poly_str(p)',
                     doc=f"polynomial of degree {len(pn) - 1}, ascending coefficients")
        elif name == "exp":
            d.update(f=f"A * np.exp(b * x){plus}", eq='f"y = {fmt4(p[0])}·exp({fmt4(p[1])}·x)"' + pc,
                     logf="np.log(np.abs(A)) + b * x" if not offset else None, doc="y = A·exp(b·x)")
        elif name == "power":
            d.update(f=f"A * np.power(x, b){plus}", eq='f"y = {fmt4(p[0])}·x^{fmt4(p[1])}"' + pc,
                     logf="np.log(np.abs(A)) + b * np.log(x)" if not offset else None, doc="y = A·x^b")
        elif name == "log":
            d.update(f="a + b * np.log(x)", doc="y = a + b·ln x",
                     eq='f"y = {fmt4(p[0])} {\'−\' if p[1] < 0 else \'+\'} {fmt4(abs(p[1]))}·ln(x)"')
        elif name == "arrhenius":
            tnote = "" if tunit == "K" else (", T = x + 273.15 K" if "273" in T else f", T = {tunit.replace('/K', '/x')} K")
            d.update(f=f"A * np.exp(-Ea * 1000.0 / (8.314462618 * {T})){plus}",
                     logf=f"np.log(np.abs(A)) - Ea * 1000.0 / (8.314462618 * {T})" if not offset else None,
                     eq='f"y = {fmt4(p[0])}·exp(−{fmt4(p[1])} kJ/mol / (R·T))"' + pc + (f" + {tnote!r}" if tnote else ""),
                     report='[("A", p[0], se[0], ""), ("Ea", p[1], se[1], "kJ/mol"), ("ln A", np.log(abs(p[0])), '
                            'se[0] / abs(p[0]), "")]' + (' + [("c", p[2], se[2], "")]' if offset else ""),
                     doc=f"y = A·exp(−Ea/(R·T)), Ea in kJ/mol, T = {T.strip('()')} K")
        elif name == "logistic":
            d.update(f=f"L / (1.0 + np.exp(-k * (x - x0))){plus}", doc="y = L / (1 + exp(−k·(x − x0)))",
                     eq='f"y = {fmt4(p[0])} / (1 + exp(−{fmt4(p[1])}·(x − {fmt4(p[2])})))"' + pc)
        elif name == "gaussian":
            d.update(f=f"A * np.exp(-0.5 * ((x - mu) / sigma) ** 2){plus}", doc="y = A·exp(−(x − mu)²/(2·sigma²))",
                     eq='f"y = {fmt4(p[0])}·exp(−(x − {fmt4(p[1])})² / (2·{fmt4(abs(p[2]))}²))"' + pc)
        elif name == "custom":
            expr = op.get("expr", "")
            d.update(f=self._custom_expr(expr, var, pn), doc=f"y = {expr}", custom_expr=expr)
        else:
            raise OpCodeError(f"model {name!r}")
        return d

    def _fit(self, o: OpX, need_stats: bool) -> Dict[str, Any]:
        """Emit data + model + curve_fit; returns what the drawing code needs (names of g / back, flags)."""
        from .figfit import _default_t_unit, fit_from_ctx
        op = o.op
        sh = o.shadow()
        res = fit_from_ctx(sh)                     # the toolkit's own fit on the real data (start values, checks)
        tunit = op.get("T_unit") or _default_t_unit(sh)
        md = self._model(o, res, tunit)
        self.ctx.need("from scipy.optimize import curve_fit")
        linearized = res.back is np.exp
        space = res.space
        # data (as displayed)
        self.emit(f"x = {o.arr(o.X())}")
        self.emit(f"y = {o.arr(o.Y())}")
        sig_key = next((k for k in ("sigma", "yerr") if o.raw(k) is not None), None)
        cond = "np.isfinite(x) & np.isfinite(y)"
        if sig_key:
            self.emit(f"sigma = np.abs(np.broadcast_to({o.arr(o.dY(sig_key))}, np.shape(y))).astype(float)")
            cond += " & np.isfinite(sigma) & (sigma > 0)"
        self.emit(f"ok = {cond}")
        self.emit("x, y" + (", sigma" if sig_key else "") + " = x[ok], y[ok]" + (", sigma[ok]" if sig_key else ""))
        # model
        pn = md["pnames"]
        args = "*c" if md.get("star") else ", ".join(pn)
        self.emit(f"def model({md['var']}, {args}):   # {md['doc']}")
        self.emit(f"    return {md['f']}")
        g, t, sig_t, back = "model", "y", ("sigma" if sig_key else None), None
        y_is = str(op.get("y_is", "value")).lower()
        if y_is in ("ln", "log10") or linearized:
            self.emit(f"def ln_model({md['var']}, {args}):   # ln of the model: fitted to ln y")
            self.emit(f"    return {md['logf']}")
            g = "ln_model"
            if y_is == "log10":
                self.emit("def log10_model(x, *p):")
                self.emit("    return ln_model(x, *p) / np.log(10.0)")
                g = "log10_model"
            elif y_is != "ln":
                self.emit("t = np.log(y)   # fit ln y (multiplicative scatter)")
                t, back = "t", "np.exp"
                if sig_key:
                    self.emit("sigma_t = sigma / y")
                    sig_t = "sigma_t"
        # start values exactly as the toolkit chose them
        start = np.asarray(res.model.p0(res.x, res.y), dtype=float)
        for nm, v in (op.get("p0") or {}).items():
            start[pn.index(nm)] = float(v)
        kw = {}
        k = len(pn)
        if op.get("bounds"):
            lo, hi = np.full(k, -np.inf), np.full(k, np.inf)
            for nm, (b0, b1) in op["bounds"].items():
                i = pn.index(nm)
                lo[i] = -np.inf if b0 is None else float(b0)
                hi[i] = np.inf if b1 is None else float(b1)
            start = np.clip(start, np.where(np.isfinite(lo), lo + 1e-12 * np.abs(lo), lo),
                            np.where(np.isfinite(hi), hi - 1e-12 * np.abs(hi), hi))
            kw["bounds"] = (_plain(lo), _plain(hi))
        else:
            kw["maxfev"] = 20000
        if sig_t:
            kw["sigma"] = Raw(sig_t)
            if op.get("absolute_sigma"):
                kw["absolute_sigma"] = True
        self.emit(f"p, pcov = curve_fit({g}, x, {t}, p0={lit(_plain(start))}, {kwargs(kw)})"
                  f"   # start values: the toolkit's automatic guess")
        if need_stats:
            self.emit(f"resid = {t} - {g}(x, *p)")
            self.emit("dof = len(x) - len(p)")
            if sig_t:
                self.emit(f"w = 1.0 / {sig_t} ** 2")
                self.emit(f"r2 = 1.0 - np.sum(w * resid ** 2) / np.sum(w * ({t} - np.sum(w * {t}) / np.sum(w)) ** 2)"
                          "   # weighted R²")
            else:
                self.emit(f"r2 = 1.0 - np.sum(resid ** 2) / np.sum(({t} - np.mean({t})) ** 2)")
        return {"g": g, "t": t, "back": back, "sig_t": sig_t, "space": space, "md": md, "res": res,
                "name": res.model.name, "sig_key": sig_key}

    def op_fit(self, o: OpX):
        op = o.op
        band = op.get("band")
        band = None if band in (None, False, "none") else ("confidence" if band is True else str(band).lower())
        ann = op.get("annotate")
        label = op.get("label")
        info = self._fit(o, need_stats=bool(ann or label is True or band == "prediction"))
        g, back = info["g"], info["back"]
        bk = (lambda s: f"{back}({s})") if back else (lambda s: s)
        # curve over the data range (or the requested range), as displayed
        if op.get("range") is not None:
            r = o.conv(list(op["range"]), "x")
            lo, hi = lit(float(r[0])), lit(float(r[1]))
        else:
            lo, hi = "x.min()", "x.max()"
        space = "np.geomspace" if (o.axm.xscale == "log" and (op.get("range") is None or float(r[0]) > 0)) \
            else "np.linspace"
        self.emit(f"xg = {space}({lo}, {hi}, {int(op.get('npoints', 200))})")
        self.emit(f"yg = {g}(xg, *p)")
        kw = {k: op[k] for k in ("color", "linestyle", "linewidth", "alpha", "zorder") if op.get(k) is not None}
        if label is True:
            kw["label"] = Raw(f'f"{info["name"]} fit (R² = {{r2:.4f}})"')
        elif label:
            kw["label"] = str(label)
        self.emit(f"line, = ax.plot(xg, {bk('yg')}{', ' + kwargs(kw) if kw else ''})")
        if band:
            level = float(op.get("level", 0.95))
            self.ctx.need("from scipy import stats")
            self.helpers("jacobian")
            self.emit(f"if np.all(np.isfinite(pcov)):   # {int(level * 100)}% {band} band (delta method, t distribution)")
            self.emit(f"    J = jacobian({g}, xg, p)")
            self.emit("    var = np.einsum('ij,jk,ik->i', J, pcov, J)")
            if band == "prediction":
                if info["sig_t"] is None:
                    self.emit("    var = var + np.sum(resid ** 2) / dof   # + residual variance")
                else:
                    self.emit("    o_ = np.argsort(x)")
                    self.emit(f"    s_new = np.exp(np.interp(xg, x[o_], np.log({info['sig_t']}[o_])))")
                    f = "1.0" if op.get("absolute_sigma") else f"(np.sum(w * resid ** 2) / dof)"
                    self.emit(f"    var = var + {f} * s_new ** 2")
            self.emit(f"    half = stats.t.ppf({0.5 + level / 2.0!r}, len(x) - len(p)) * np.sqrt(np.clip(var, 0, None))")
            bkw = {"color": Raw("line.get_color()") if op.get("band_color") is None else op["band_color"],
                   "alpha": float(op.get("band_alpha", 0.2)), "linewidth": 0, "label": op.get("band_label"),
                   "zorder": Raw("line.get_zorder() - 0.1")}
            self.emit(f"    ax.fill_between(xg, {bk('yg - half')}, {bk('yg + half')}, {kwargs(bkw)})")
        if op.get("show_data"):
            dkw = {"color": op.get("data_color") if op.get("data_color") is not None else Raw("line.get_color()"),
                   "marker": op.get("marker", "o"), "label": op.get("data_label")}
            if info["sig_key"] and info["space"] == "y":
                self.emit(f"ax.errorbar(x, y, yerr=sigma, linestyle='none', capsize=2, "
                          f"markersize={lit(op.get('markersize', 5))}, {kwargs(dkw)})")
            else:
                self.emit(f"ax.scatter(x, y, s={lit(op.get('s', 25))}, zorder=line.get_zorder() + 0.1, {kwargs(dkw)})")
        if ann:
            from .figfit import _LOCS
            md = info["md"]
            self.helpers("fmt4")
            if md.get("custom_expr"):
                self.ctx.need("import re")
                self.emit(f"equation = {md['custom_expr']!r}")
                self.emit(f"for name, v in sorted(zip({lit(md['pnames'])}, p), key=lambda t: -len(t[0])):")
                self.emit("    equation = re.sub(rf'\\b{re.escape(name)}\\b', f'({fmt4(v)})', equation)")
                self.emit("equation = 'y = ' + equation")
            else:
                self.emit(f"equation = {md['eq']}")
            self.emit("se = np.sqrt(np.clip(np.diag(pcov), 0, None)) if np.all(np.isfinite(pcov)) "
                      "else np.full(len(p), np.nan)")
            rep = md["report"] or f"list(zip({lit(md['pnames'])}, p, se, [''] * len(p)))"
            self.emit(f"params = '\\n'.join(f\"{{n}} = {{fmt4(v)}} ± {{fmt4(s)}}{{(' ' + u) if u else ''}}\" "
                      f"for n, v, s, u in {rep})")
            if ann == "equation":
                self.emit("note = f'{equation}\\nR² = {r2:.4f}'")
            elif ann == "params":
                self.emit("note = params + f'\\nR² = {r2:.4f}'")
            else:
                self.emit("note = f'{equation}\\n' + params + f'\\nR² = {r2:.4f}, n = {len(x)}'")
            x0, y0, ha, va = _LOCS[op.get("annotate_loc", "upper left")]
            self.emit(f"ax.text({x0}, {y0}, note, transform=ax.transAxes, ha={ha!r}, va={va!r}, "
                      f"fontsize={lit(op.get('fontsize', 8))},")
            self.emit("        bbox={'boxstyle': 'round,pad=0.3', 'facecolor': 'white', 'alpha': 0.85, 'edgecolor': '0.6'})")

    def op_residuals(self, o: OpX):
        op = o.op
        info = self._fit(o, need_stats=True)
        if op.get("standardized"):
            self.emit(f"resid = resid / " + (info["sig_t"] if info["sig_t"] else "np.sqrt(np.sum(resid ** 2) / dof)"))
        kw = {"color": op.get("color"), "marker": op.get("marker", "o"), "s": op.get("s", 25)}
        if op.get("label"):
            kw["label"] = op["label"]
        if op.get("alpha") is not None:
            kw["alpha"] = op["alpha"]
        self.emit(f"ax.scatter(x, resid, {kwargs({k: v for k, v in kw.items() if v is not None})})")
        if op.get("zero_line", True):
            self.emit("ax.axhline(0.0, color='0.4', linewidth=0.8, linestyle='--', zorder=0)")
        lab = ("standardized residual" if op.get("standardized") else "residual") + \
            ("" if info["space"] == "y" else f" ({info['space']})")
        self.emit("if not ax.get_ylabel():")
        self.emit(f"    ax.set_ylabel({lab!r})")

    # =====================================================================================================
    # statistics (calphad_tools.figfit)
    # =====================================================================================================
    def groups(self, o: OpX, key: str = "values", extra: Sequence[str] = ()) -> Tuple[List[str], List[str]]:
        """Emit ``groups = {label: values}`` (+ ``<extra> = {label: values}``); returns (labels, extras present)."""
        from .figfit import _groups
        from .figures import find_col
        op = o.op
        glist = _groups(o.shadow(), key, extra)
        labels = [lab for lab, _, _ in glist]
        exs = [e for e in extra if e in glist[0][2]]
        if len(set(labels)) < len(labels):
            raise OpCodeError("duplicate group labels")
        src_cols = op.get(f"{key}_col") or op.get("value_cols") or op.get("group_col")
        if o.has_source and src_cols:
            df, fr = o.df(), o.frame()
            vcols = op.get("value_cols")
            cols = [find_col(fr, c) for c in (vcols if isinstance(vcols, list) else [vcols])] if vcols else \
                [o.colname(key)]
            ecols = {e: find_col(fr, op[f"{e}_col"]) for e in exs}
            if op.get("group_col"):
                gc = find_col(fr, op["group_col"])
                order = op.get("groups")
                key_code = lit(list(order)) if isinstance(order, list) else f"pd.unique({df}[{gc!r}].dropna())"
                as_str = isinstance(order, list) and any(
                    isinstance(g, str) and not (fr[gc] == g).any() and (fr[gc].astype(str) == g).any() for g in order)
                sel = f"{df}[{gc!r}].astype(str)" if as_str else f"{df}[{gc!r}]"
                self.emit(f"subsets = {{str(g): {df}[{sel} == g] for g in {key_code}}}")
                self.emit(f"groups = {{k: sub[{cols[0]!r}].to_numpy(float) for k, sub in subsets.items()}}")
                for e, c in ecols.items():
                    self.emit(f"{e} = {{k: sub[{c!r}].to_numpy(float) for k, sub in subsets.items()}}")
            else:
                self.emit("groups = {" + ", ".join(f"{lab!r}: {df}[{c!r}].to_numpy(float)"
                                                   for lab, c in zip(labels if not isinstance(op.get('labels'), list)
                                                                     else [str(c) if len(cols) > 1 else
                                                                           (op.get('label') or '') for c in cols],
                                                                     cols)) + "}")
                for e, c in ecols.items():
                    self.emit(f"{e} = {{k: {df}[{c!r}].to_numpy(float) for k in groups}}")
        else:
            v = op.get(key)
            grp = op.get("groups")
            if isinstance(grp, dict):
                v = grp
            if isinstance(v, dict):
                self.emit(f"groups = {{str(k): np.asarray(v, dtype=float) for k, v in {o.L(v)}.items()}}")
                raw_labels = [str(k) for k in v]
            elif isinstance(v, (list, tuple)) and v and all(isinstance(t, (list, tuple)) for t in v):
                self.emit(f"groups = {{f'group {{i + 1}}': np.asarray(v, dtype=float) for i, v in enumerate({o.L(v)})}}")
                raw_labels = [f"group {i + 1}" for i in range(len(v))]
            else:
                raw_labels = [op.get("label") or ""]
                self.emit(f"groups = {{{raw_labels[0]!r}: np.atleast_1d(np.asarray({o.L(v)}, dtype=float))}}")
            for e in exs:
                w = op.get(e)
                vals = {lab: (w[k] if isinstance(w, dict) else (w[i] if isinstance(v, (dict, list)) and
                                                                  len(raw_labels) > 1 else w))
                        for i, (lab, k) in enumerate(zip(raw_labels, (list(v) if isinstance(v, dict) else raw_labels)))}
                self.emit(f"{e} = {{k: np.asarray(w, dtype=float) for k, w in {o.L(vals)}.items()}}")
        if isinstance(op.get("labels"), list):
            self.emit(f"groups = dict(zip({lit([str(x) for x in op['labels']])}, groups.values()))")
            for e in exs:
                self.emit(f"{e} = dict(zip(groups, {e}.values()))")
        if exs:
            self.emit("for k in groups:   # drop non-finite values")
            self.emit("    ok = np.isfinite(groups[k])" + "".join(f" & np.isfinite({e}[k])" for e in exs))
            self.emit("    groups[k] = groups[k][ok]")
            for e in exs:
                self.emit(f"    {e}[k] = {e}[k][ok]")
        else:
            self.emit("groups = {k: v[np.isfinite(v)] for k, v in groups.items()}   # drop non-finite values")
        return labels, exs

    def _gconv(self, o: OpX, axis: str):
        t = o._t(axis)
        if t:
            self.emit(f"groups = {{k: v * {t[0]!r} + {t[1]!r} for k, v in groups.items()}}   # data -> shown units")

    def _gdata(self, labels) -> str:
        return "list(groups.values())" if len(labels) > 1 else f"groups[{labels[0]!r}]"

    def op_hist(self, o: OpX):
        from .figfit import _colors, _style
        op = o.op
        labels, exs = self.groups(o, "values", ("weights",))
        horiz = op.get("orientation", "vertical") == "horizontal"
        self._gconv(o, "y" if horiz else "x")
        kw: Dict[str, Any] = {}
        bins = op.get("bins", 20)
        kw["bins"] = o.conv(list(bins), "y" if horiz else "x") if isinstance(bins, list) else bins
        if op.get("range") is not None:
            kw["range"] = tuple(o.conv(list(op["range"]), "y" if horiz else "x"))
        for k, dflt in (("density", False), ("cumulative", False), ("histtype", "bar"), ("stacked", False),
                        ("log", False)):
            val = op.get(k, dflt)
            val = bool(val) if k in ("density", "stacked", "log") else val
            if val != dflt:
                kw[k] = val
        if horiz:
            kw["orientation"] = "horizontal"
        if op.get("rwidth") is not None:
            kw["rwidth"] = op["rwidth"]
        if "weights" in exs:
            kw["weights"] = Raw("list(weights.values())" if len(labels) > 1 else f"weights[{labels[0]!r}]")
        cols = _colors(op, len(labels))
        if cols:
            kw["color"] = cols if len(labels) > 1 else cols[0]
        if any(labels):
            kw["label"] = labels if len(labels) > 1 else labels[0]
        kw.update(_style(op, ("alpha", "edgecolor", "linewidth", "linestyle", "hatch", "zorder")))
        self.emit(f"ax.hist({self._gdata(labels)}, {kwargs(kw, self.indent + 4)})")

    def _points_overlay(self, o: OpX, pos, horiz, colors, dark=False):
        op = o.op
        jit = op.get("jitter")
        if not (jit or op.get("points")):
            return
        w = 0.15 if jit is True else float(jit or 0.0)
        pcols = [op.get("point_color", "k" if (colors is None or dark) else colors[i]) for i in range(len(pos))]
        self.emit(f"rng = np.random.default_rng({int(op.get('seed', 0))})   # jittered points on top")
        self.emit(f"for v, p, c in zip(groups.values(), {lit(pos)}, {lit(pcols)}):")
        self.emit(f"    off = p + rng.uniform(-{w!r}, {w!r}, len(v))" if w else "    off = p")
        xy = "v, off" if horiz else "off, v"
        self.emit(f"    ax.scatter({xy}, s={lit(op.get('s', 10))}, alpha={lit(op.get('point_alpha', 0.6))}, zorder=3, "
                  "color=c, linewidths=0)")

    def op_box(self, o: OpX):
        from .figfit import _colors, _positions, _style
        op = o.op
        labels, _ = self.groups(o)
        horiz = op.get("orientation", "vertical") == "horizontal"
        self._gconv(o, "x" if horiz else "y")
        pos = _positions(op, len(labels))
        kw: Dict[str, Any] = {"positions": pos}
        if horiz:
            kw["orientation"] = "horizontal"
        kw.update(patch_artist=True, showmeans=bool(op.get("showmeans", False)), notch=bool(op.get("notch", False)),
                  showfliers=bool(op.get("showfliers", not (op.get("jitter") or op.get("points")))))
        if op.get("widths") is not None:
            kw["widths"] = op["widths"]
        if op.get("whis") is not None:
            kw["whis"] = op["whis"]
        if any(labels):
            kw["tick_labels"] = labels
        if op.get("median_color"):
            kw["medianprops"] = {"color": _style(op, ("color",), "median_")["color"], "linewidth": 1.5}
        self.emit(f"bp = ax.boxplot(list(groups.values()), {kwargs(kw, self.indent + 4)})")
        cols = _colors(op, len(labels)) or [f"C{i % 10}" for i in range(len(labels))]
        self.emit(f"for patch, c in zip(bp['boxes'], {lit(cols)}):")
        self.emit("    patch.set_facecolor(c)")
        self.emit(f"    patch.set_alpha({float(op.get('alpha', 0.6))!r})")
        if not op.get("median_color"):
            self.emit("for m in bp['medians']:")
            self.emit("    m.set_color('k')")
        self._points_overlay(o, pos, horiz, cols)

    def op_violin(self, o: OpX):
        from .figfit import _colors, _positions
        op = o.op
        labels, _ = self.groups(o)
        horiz = op.get("orientation", "vertical") == "horizontal"
        self._gconv(o, "x" if horiz else "y")
        pos = _positions(op, len(labels))
        kw: Dict[str, Any] = {"positions": pos}
        if horiz:
            kw["orientation"] = "horizontal"
        kw.update(showmeans=bool(op.get("showmeans", False)), showmedians=bool(op.get("showmedians", True)),
                  showextrema=bool(op.get("showextrema", True)), widths=op.get("widths", 0.7))
        if op.get("side", "both") != "both":
            kw["side"] = op["side"]
        if op.get("bw") is not None:
            kw["bw_method"] = op["bw"]
        if op.get("quantiles") is not None:
            q = op["quantiles"]
            kw["quantiles"] = [q] * len(labels) if q and not isinstance(q[0], list) else q
        self.emit(f"parts = ax.violinplot(list(groups.values()), {kwargs(kw, self.indent + 4)})")
        cols = _colors(op, len(labels)) or [f"C{i % 10}" for i in range(len(labels))]
        self.emit(f"for body, c in zip(parts['bodies'], {lit(cols)}):")
        self.emit("    body.set_facecolor(c)")
        self.emit("    body.set_edgecolor(c)")
        self.emit(f"    body.set_alpha({float(op.get('alpha', 0.5))!r})")
        self.emit("for k in ('cmeans', 'cmedians', 'cmins', 'cmaxes', 'cbars', 'cquantiles'):")
        self.emit("    if k in parts:")
        self.emit("        parts[k].set_color('0.2')")
        if any(labels):
            self.emit(f"ax.set_{'y' if horiz else 'x'}ticks({lit(pos)}, {lit(labels)})")
        self._points_overlay(o, pos, horiz, cols, dark=True)

    def op_kde(self, o: OpX):
        from .figfit import _colors, _style
        self.ctx.need("from scipy.stats import gaussian_kde")
        op = o.op
        labels, exs = self.groups(o, "values", ("weights",))
        horiz = op.get("orientation", "vertical") == "horizontal"
        self._gconv(o, "y" if horiz else "x")
        cols = _colors(op, len(labels))
        style = _style(op, ("linestyle", "linewidth", "zorder"))
        for i, lab in enumerate(labels):
            kw = dict(style)
            if cols:
                kw["color"] = cols[i]
            if lab:
                kw["label"] = lab
            self.emit(f"v = groups[{lab!r}]")
            self.emit(f"k = gaussian_kde(v, bw_method={lit(op.get('bw'))}"
                      + (f", weights=weights[{lab!r}])" if "weights" in exs else ")"))
            if op.get("range") is not None:
                r = o.conv(list(op["range"]), "y" if horiz else "x")
                self.emit(f"g = np.linspace({lit(float(r[0]))}, {lit(float(r[1]))}, {int(op.get('npoints', 256))})")
            else:
                cut = float(op.get("cut", 3))
                self.emit("bw = float(np.sqrt(k.covariance[0, 0]))")
                self.emit(f"g = np.linspace(v.min() - {cut!r} * bw, v.max() + {cut!r} * bw, {int(op.get('npoints', 256))})")
            self.emit("d = k(g)" + (" * len(v)   # counts" if op.get("scale") == "count" else ""))
            xy = "d, g" if horiz else "g, d"
            self.emit(f"line, = ax.plot({xy}{', ' + kwargs(kw) if kw else ''})")
            if op.get("fill"):
                fb = "fill_betweenx" if horiz else "fill_between"
                self.emit(f"ax.{fb}(g, 0, d, color=line.get_color(), alpha={float(op.get('alpha', 0.25))!r}, linewidth=0)")
            if op.get("rug"):
                if horiz:
                    self.emit("ax.plot(np.zeros_like(v), v, '_', color=line.get_color(), markersize=10, "
                              "transform=ax.get_yaxis_transform(), alpha=0.7)")
                else:
                    self.emit("ax.plot(v, np.zeros_like(v), '|', color=line.get_color(), markersize=10, "
                              "transform=ax.get_xaxis_transform(), alpha=0.7)")

    def op_ecdf(self, o: OpX):
        from .figfit import _colors, _style
        op = o.op
        labels, exs = self.groups(o, "values", ("weights",))
        horiz = op.get("orientation", "vertical") == "horizontal"
        self._gconv(o, "y" if horiz else "x")
        cols = _colors(op, len(labels))
        for i, lab in enumerate(labels):
            kw = _style(op, ("linestyle", "linewidth", "alpha", "zorder"))
            if cols:
                kw["color"] = cols[i]
            if lab:
                kw["label"] = lab
            extra = {}
            if "weights" in exs:
                extra["weights"] = Raw(f"weights[{lab!r}]")
            if op.get("complementary"):
                extra["complementary"] = True
            if horiz:
                extra["orientation"] = "horizontal"
            self.emit(f"ax.ecdf(groups[{lab!r}]{', ' + kwargs({**extra, **kw}) if (extra or kw) else ''})")

    def op_qq(self, o: OpX):
        from .figfit import _colors
        self.ctx.need("from scipy import stats")
        op = o.op
        labels, _ = self.groups(o)
        dist = op.get("dist", "norm")
        cols = _colors(op, len(labels))
        t = o._t("y")
        for i, lab in enumerate(labels):
            kw = {"marker": op.get("marker", "o"), "s": op.get("s", 18)}
            if cols:
                kw["color"] = cols[i]
            if lab:
                kw["label"] = lab
            self.emit(f"(osm, osr), (slope, intercept, r) = stats.probplot(groups[{lab!r}], dist={dist!r})")
            conv = (lambda s: f"({s}) * {t[0]!r} + {t[1]!r}") if t else (lambda s: s)
            self.emit(f"sc = ax.scatter(osm, {conv('osr')}, {kwargs(kw)})")
            if op.get("line", True):
                self.emit("xx = np.array([osm.min(), osm.max()])")
                self.emit(f"ax.plot(xx, {conv('intercept + slope * xx')}, color=sc.get_facecolor()[0], "
                          "linestyle='--', linewidth=1.2)")
        self.emit("if not ax.get_xlabel():")
        self.emit(f"    ax.set_xlabel({f'theoretical quantiles ({dist})'!r})")
        self.emit("if not ax.get_ylabel():")
        self.emit("    ax.set_ylabel('ordered values')")

    def op_hist2d(self, o: OpX):
        op = o.op
        self.emit(f"x, y = {o.arr(o.X())}, {o.arr(o.Y())}")
        w = o.raw("weights")
        cond = "np.isfinite(x) & np.isfinite(y)"
        if w is not None:
            self.emit(f"w = {o.arr(w)}")
            cond += " & np.isfinite(w)"
        self.emit(f"ok = {cond}")
        kw: Dict[str, Any] = {}
        if w is not None:
            kw["weights"] = Raw("w[ok]")
        kw.update(bins=op.get("bins", 30), density=bool(op.get("density", False)), cmap=op.get("cmap", "viridis"))
        if op.get("range") is not None:
            (x0, x1), (y0, y1) = op["range"]
            kw["range"] = [o.conv([x0, x1], "x"), o.conv([y0, y1], "y")]
        for k in ("cmin", "cmax", "alpha"):
            if op.get(k) is not None:
                kw[k] = op[k]
        if op.get("norm") is not None:
            kw["norm"] = self.norm(op["norm"])
        else:
            kw.update({k: op[k] for k in ("vmin", "vmax") if op.get(k) is not None})
        self.emit(f"counts, xedges, yedges, mesh = ax.hist2d(x[ok], y[ok], {kwargs(kw)})")
        self.colorbar(o, "mesh")

    def op_hexbin(self, o: OpX):
        op = o.op
        red = op.get("reduce", "mean")
        self.emit(f"x, y = {o.arr(o.X())}, {o.arr(o.Y())}")
        C = o.raw("c")
        cond = "np.isfinite(x) & np.isfinite(y)"
        if C is not None:
            self.emit(f"C = {o.arr(C)}")
            cond += " & np.isfinite(C)"
        self.emit(f"ok = {cond}")
        kw: Dict[str, Any] = {}
        if C is not None:
            kw["C"] = Raw("C[ok]")
        kw.update(gridsize=op.get("gridsize", 30), cmap=op.get("cmap", "viridis"), xscale=op.get("xscale", "linear"),
                  yscale=op.get("yscale", "linear"),
                  reduce_C_function=Raw({"count": "np.size"}.get(red, f"np.{red}")),
                  edgecolors=op.get("edgecolor", "face"))
        for k in ("mincnt", "alpha", "vmin", "vmax"):
            if op.get(k) is not None:
                kw[k] = op[k]
        if op.get("linewidth") is not None:
            kw["linewidths"] = op["linewidth"]
        if op.get("extent") is not None:
            e = op["extent"]
            kw["extent"] = list(o.conv(list(e[:2]), "x")) + list(o.conv(list(e[2:4]), "y"))
        if op.get("norm") is not None:
            kw["norm"] = self.norm(op["norm"])
            kw.pop("vmin", None)
            kw.pop("vmax", None)
        elif op.get("bins") is not None:
            kw["bins"] = op["bins"]
        self.emit(f"hb = ax.hexbin(x[ok], y[ok], {kwargs(kw)})")
        self.colorbar(o, "hb")

    def op_band(self, o: OpX):
        from .figfit import _colors, _style
        op = o.op
        from .figfit import _band_agg
        if op.get("agg"):
            series = _band_agg(o.shadow())
            self._band_agg_code(o)
            n = len(series)
        else:
            self.emit(f"x, y = {o.arr(o.X())}, {o.arr(o.Y())}")
            if o.raw("yerr") is not None:
                self.emit(f"e = {o.arr(o.dY('yerr'))}")
                self.emit("lo, hi = y - e, y + e")
            else:
                self.emit(f"lo, hi = {o.arr(o.Y('ylo'))}, {o.arr(o.Y('yhi'))}")
            self.emit("o_ = np.argsort(x)")
            self.emit(f"series = [({op.get('label')!r}, x[o_], y[o_], np.broadcast_to(lo, x.shape)[o_], "
                      f"np.broadcast_to(hi, x.shape)[o_])]")
            n = 1
        cols = _colors(op, n)
        labels = op.get("labels")
        style = _style(op, ("linestyle", "linewidth", "marker", "markersize", "zorder"))
        self.emit(f"for i, (lab, xs, ys, lo, hi) in enumerate(series):")
        kw = dict(style)
        if cols:
            kw["color"] = Raw(f"{lit(cols)}[i]")
        lab_code = f"{lit([str(v) for v in labels])}[i]" if isinstance(labels, list) else "lab"
        self.emit(f"    line, = ax.plot(xs, ys{', ' + kwargs(kw) if kw else ''}, "
                  f"label=str({lab_code}) if {lab_code} else None)")
        self.emit(f"    ax.fill_between(xs, lo, hi, color=line.get_color(), alpha={float(op.get('alpha', 0.25))!r}, "
                  f"linewidth=0, label={lit(op.get('band_label'))} if i == 0 else None)")

    def _band_agg_code(self, o: OpX):
        from .figures import find_col
        op = o.op
        df, fr = o.df(), o.frame()
        xc, yc = o.colname("x"), o.colname("y")
        agg = op["agg"]
        center = op.get("center", "mean")
        if op.get("group_col"):
            gc = find_col(fr, op["group_col"])
            keys = lit(op["groups"]) if op.get("groups") else f"pd.unique({df}[{gc!r}].dropna())"
            self.emit(f"parts = [(str(k), {df}[{df}[{gc!r}] == k]) for k in {keys}]")
        else:
            self.emit(f"parts = [({lit(op.get('label'))}, {df})]")
        self.emit("series = []   # (label, x, centre, low, high) with rows grouped by x")
        self.emit("for lab, sub in parts:")
        self.emit(f"    gb = sub[[{xc!r}, {yc!r}]].astype(float).dropna().groupby({xc!r})[{yc!r}]")
        self.emit("    xs = np.asarray(sorted(gb.groups), dtype=float)")
        self.emit(f"    c = gb.{'median' if center == 'median' else 'mean'}().reindex(xs).to_numpy()")
        if agg in ("std", "sem", "ci"):
            self.emit("    sd = gb.std(ddof=1).reindex(xs).to_numpy()")
            if agg == "std":
                self.emit(f"    h = {float(op.get('nstd', 1.0))!r} * sd")
            else:
                self.emit("    cnt = gb.count().reindex(xs).to_numpy()")
                self.emit("    h = sd / np.sqrt(cnt)")
                if agg == "ci":
                    self.ctx.need("from scipy import stats")
                    self.emit(f"    h = h * stats.t.ppf({0.5 + float(op.get('level', 0.95)) / 2!r}, np.maximum(cnt - 1, 1))")
            self.emit("    lo, hi = c - np.nan_to_num(h), c + np.nan_to_num(h)")
        else:
            q = {"minmax": (0, 100), "iqr": (25, 75)}.get(agg) or tuple(op.get("pct", (2.5, 97.5)))
            self.emit(f"    lo = gb.quantile({q[0] / 100!r}).reindex(xs).to_numpy()")
            self.emit(f"    hi = gb.quantile({q[1] / 100!r}).reindex(xs).to_numpy()")
        self.emit("    series.append((lab, xs, c, lo, hi))")

    def op_corr_heatmap(self, o: OpX):
        from .figfit import _numeric_cols
        op = o.op
        _, cols = _numeric_cols(o.shadow(), 30)
        df = o.df()
        method = op.get("method", "pearson")
        n = len(cols)
        self.emit(f"cols = {lit([str(c) for c in cols])}")
        self.emit(f"C = {df}[cols].astype(float).corr(method={method!r}).to_numpy()")
        self.emit("M = C.copy()")
        tri = op.get("triangle")
        if tri == "lower":
            self.emit("M[np.triu(np.ones_like(M, bool), 1)] = np.nan   # lower triangle only")
        elif tri == "upper":
            self.emit("M[np.tril(np.ones_like(M, bool), -1)] = np.nan   # upper triangle only")
        self.emit(f"im = ax.imshow(M, cmap={op.get('cmap', 'RdBu_r')!r}, vmin=-1, vmax=1, interpolation='nearest')")
        self.emit(f"ax.set_xticks(range({n}), cols, rotation=45, ha='right', rotation_mode='anchor')")
        self.emit(f"ax.set_yticks(range({n}), cols)")
        self.emit("ax.grid(False)")
        if op.get("annotate", True):
            fmt = "{:" + str(op.get("fmt", ".2f")) + "}"
            fs = op.get("fontsize", max(5, 10 - n // 3))
            self.emit(f"for i in range({n}):")
            self.emit(f"    for j in range({n}):")
            self.emit("        if np.isfinite(M[i, j]):")
            self.emit(f"            ax.text(j, i, {fmt!r}.format(M[i, j]), ha='center', va='center', fontsize={lit(fs)},")
            self.emit("                    color='white' if abs(M[i, j]) > 0.6 else 'black')")
        spec = op["colorbar"] if "colorbar" in op else {"label": f"{method} r"}
        self.colorbar(o, "im", spec=spec)

    def op_scatter_matrix(self, o: OpX):
        from .figfit import _numeric_cols
        from .figures import find_col
        op = o.op
        _, cols = _numeric_cols(o.shadow(), 8)
        df, fr = o.df(), o.frame()
        self.helpers("scatter_matrix")
        if op.get("group_col"):
            gc = find_col(fr, op["group_col"])
            keys = lit(op["groups"]) if op.get("groups") else f"pd.unique({df}[{gc!r}].dropna())"
            self.emit(f"parts = [(str(k), {df}[{df}[{gc!r}] == k]) for k in {keys}]")
            import pandas as pd
            nparts = len(op.get("groups") or list(pd.unique(fr[gc].dropna())))
        else:
            self.emit(f"parts = [(None, {df})]")
            nparts = 1
        base = op.get("color")
        colors = [base] if (base and nparts == 1) else None
        name = self.fresh("pairs")
        kw = {"diag": op.get("diag", "hist"), "bins": op.get("bins", 15), "s": op.get("s", 8),
              "alpha": op.get("alpha", 0.7), "colors": colors, "corner": bool(op.get("corner"))}
        self.emit(f"{name} = scatter_matrix(fig, ax, parts, {lit([str(c) for c in cols])}, {kwargs(kw)})")
        self.emit(f"ax = {name}[0]")
        self.cur_ax = f"{name}[0]"
        # the target axes is replaced by the grid (new axes appended, row-major)
        tgt = o.axm
        if tgt in self.fm.top:
            self.fm.top.remove(tgt)
        n = len(cols)
        k = 0
        for i in range(n):
            for j in range(n):
                if op.get("corner") and j > i:
                    continue
                self.add_top(AxModel(f"{name}[{k}]", xlabel=str(cols[j]) if i == n - 1 else "",
                                     ylabel=str(cols[i]) if j == 0 and i > 0 else ""))
                k += 1

    # =====================================================================================================
    # grids & fields (calphad_tools.figops_extra)
    # =====================================================================================================
    def _order_code(self, fr, df: str, c: str) -> str:
        from .figops_extra import _is_num
        import pandas as pd
        u = list(pd.unique(np.asarray(fr[c], dtype=object)))
        return f"sorted(pd.unique({df}[{c!r}]), key=float)" if _is_num(u) else f"list(pd.unique({df}[{c!r}]))"

    def _pivot_code(self, o: OpX, rkey: str, ckey: str, vkey: str, var: str = "grid", fill=None) -> None:
        """Long data -> 2-D table (mean of duplicates), rows/columns in the toolkit's order."""
        df, fr = o.df(), o.frame()
        rc, cc, vc = o.colname(rkey), o.colname(ckey), o.colname(vkey)
        self.emit(f"{var} = {df}.pivot_table(index={rc!r}, columns={cc!r}, values={vc!r}, aggfunc='mean', sort=False)")
        self.emit(f"{var} = {var}.reindex(index={self._order_code(fr, df, rc)}, columns={self._order_code(fr, df, cc)})"
                  + (f".fillna({lit(fill)})" if fill is not None else ""))

    def _grid(self, o: OpX):
        """Emit x, y, Z for grid ops; returns (x, y, Z code, scattered)."""
        from .figops_extra import _grid as tgrid
        op = o.op
        sx, sy, sZ, scattered = tgrid(o.shadow())
        if o.has_source and op.get("z_col"):
            if scattered:
                self.emit(f"# the {op['x_col']}/{op['y_col']} values do not form a grid: scattered points")
                return o.X(), o.Y(), o.raw("z"), True
            self._pivot_code(o, "y", "x", "z")
            self.emit("x, y, Z = grid.columns.to_numpy(float), grid.index.to_numpy(float), grid.to_numpy(float)")
            tx, ty = o._t("x"), o._t("y")
            if tx:
                self.emit(f"x = x * {tx[0]!r} + {tx[1]!r}")
            if ty:
                self.emit(f"y = y * {ty[0]!r} + {ty[1]!r}")
            return Raw("x"), Raw("y"), Raw("Z"), False
        Z = np.asarray(op["z"], dtype=float)
        if Z.ndim == 1:
            return o.X(), o.Y(), op["z"], True
        x, y = o.X(), o.Y()
        xs = "np.arange(%d, dtype=float)" % Z.shape[1] if x is None else o.L(x)
        ys = "np.arange(%d, dtype=float)" % Z.shape[0] if y is None else o.L(y)
        self.emit(f"x = {xs}")
        self.emit(f"y = {ys}")
        transpose = (np.ndim(x) == 1 and np.ndim(y) == 1 and x is not None and y is not None
                     and Z.shape != (len(y), len(x)) and Z.shape == (len(x), len(y)))
        self.emit(f"Z = np.array({o.L(op['z'])}, dtype=float)" + (".T   # given as [len(x)][len(y)]" if transpose else ""))
        return Raw("x"), Raw("y"), Raw("Z"), False

    def op_heatmap(self, o: OpX):
        from .figops_extra import _is_num, _lab, _pivot
        op = o.op
        if o.has_source and op.get("z_col"):
            ylab, xlab, Zs = _pivot(o.shadow(), "y", "x", "z")
            self._pivot_code(o, "y", "x", "z")
            self.emit("Z = grid.to_numpy(float)")
        else:
            Zs = np.asarray(op["z"], dtype=float)
            xlab = op.get("x") or op.get("xlabels") or list(range(Zs.shape[1]))
            ylab = op.get("y") or op.get("ylabels") or list(range(Zs.shape[0]))
            self.emit(f"Z = np.array({o.L(op['z'])}, dtype=float)")
        origin = op.get("origin", "lower" if _is_num(ylab) else "upper")
        kw = {"aspect": op.get("aspect", "auto"), "origin": origin, "interpolation": op.get("interpolation", "nearest"),
              **self.cmap_kw(o), **o.pick({"alpha", "zorder"})}
        self.emit(f"im = ax.imshow(np.ma.masked_invalid(Z), {kwargs(kw)})")
        nx, ny = Zs.shape[1], Zs.shape[0]
        sx, sy = max(1, math.ceil(nx / 25)), max(1, math.ceil(ny / 25))
        self.emit(f"ax.set_xticks(np.arange(0, {nx}, {sx}), {o.L([_lab(v) for v in list(xlab)[::sx]])})")
        self.emit(f"ax.set_yticks(np.arange(0, {ny}, {sy}), {o.L([_lab(v) for v in list(ylab)[::sy]])})")
        rot = op.get("xrotation", 45 if max((len(_lab(v)) for v in xlab), default=0) > 4 else 0)
        if rot:
            self.emit("for t in ax.get_xticklabels():")
            self.emit(f"    t.set_rotation({lit(rot)})")
            self.emit("    t.set_ha('right')")
        if op.get("annotate"):
            fs = op.get("fontsize", 8 if nx * ny <= 100 else 6)
            self.emit("for i in range(Z.shape[0]):   # value labels, dark text on light cells")
            self.emit("    for j in range(Z.shape[1]):")
            self.emit("        if np.isfinite(Z[i, j]):")
            self.emit("            r, g, b, _ = im.cmap(im.norm(Z[i, j]))")
            self.emit(f"            ax.text(j, i, {self._fmt_code(op.get('fmt'), 'Z[i, j]')}, ha='center', va='center', "
                      f"fontsize={lit(fs)},")
            self.emit("                    color='black' if 0.299 * r + 0.587 * g + 0.114 * b > 0.5 else 'white')")
        self.colorbar(o, "im")

    @staticmethod
    def _fmt_code(fmt, v: str) -> str:
        """Code formatting value ``v`` like figops_extra._fmt."""
        if fmt is None:
            fmt = "{:.3g}"
        if "{" in fmt:
            return f"{re.sub(r'{x(?=[:}!])', '{', fmt)!r}.format({v})"
        if "%" in fmt:
            return f"{fmt!r} % {v}"
        return f"format({v}, {fmt!r})"

    def op_pcolormesh(self, o: OpX):
        op = o.op
        x, y, Z, scattered = self._grid(o)
        if scattered:
            raise OpCodeError("pcolormesh needs gridded data")
        kw = self.cmap_kw(o)
        p = o.pick({"edgecolor", "linewidth", "alpha", "zorder"})
        if "edgecolor" in p:
            p["edgecolors"] = p.pop("edgecolor")
        if "linewidth" in p:
            p["linewidths"] = p.pop("linewidth")
        kw.update(p)
        self.emit(f"mesh = ax.pcolormesh({x}, {y}, np.ma.masked_invalid({Z}), shading={op.get('shading', 'auto')!r}, "
                  f"{kwargs(kw)})")
        self.colorbar(o, "mesh")

    def _contour_common(self, o: OpX, filled: bool, x=None, y=None, Z=None, tri: bool = False):
        from .figops_extra import _pyfmt
        op = o.op
        kw = self.cmap_kw(o, allow_colors=True)
        kw.update(o.pick({"alpha", "zorder"}))
        lv = op.get("levels", 10)
        levels = int(lv) if isinstance(lv, (int, float)) else [float(v) for v in lv]
        if filled:
            if op.get("extend"):
                kw["extend"] = op["extend"]
            if op.get("hatches"):
                kw["hatches"] = op["hatches"]
        else:
            for k in ("linestyles", "linewidths"):
                if op.get(k) is not None:
                    kw[k] = op[k]
        fn = ("tricontourf" if filled else "tricontour") if tri else ("contourf" if filled else "contour")
        data = "tri, z" if tri else f"{x}, {y}, np.ma.masked_invalid({Z})"
        self.emit(f"cs = ax.{fn}({data}, levels={lit(levels)}{', ' + kwargs(kw) if kw else ''})")
        lines = None if filled else "cs"
        if filled and op.get("lines"):
            lkw = {"colors": op.get("line_colors", "k"), "linewidths": op.get("linewidths", 0.6)}
            if op.get("linestyles") is not None:
                lkw["linestyles"] = op["linestyles"]
            self.emit(f"lines = ax.{'tricontour' if tri else 'contour'}({data}, levels=cs.levels, {kwargs(lkw)})")
            lines = "lines"
        if lines is not None and op.get("clabel", not filled):
            lf = op.get("clabel_fmt")
            self.emit(f"ax.clabel({lines}, fontsize={lit(op.get('fontsize', 8))}"
                      + (f", fmt={_pyfmt(lf)!r})" if lf else ")"))
        if op.get("label"):
            self.emit(f"{'cs' if filled else (lines or 'cs')}.set_label({op['label']!r})")
        self.colorbar(o, "cs")

    def op_contour(self, o: OpX, filled: bool = False):
        x, y, Z, scattered = self._grid(o)
        if scattered:
            self.emit("# scattered x, y, z: contoured on a triangulation")
            return self._tri(o, "tricontourf" if filled else "tricontour", x, y, Z)
        self._contour_common(o, filled, x, y, Z)

    def op_contourf(self, o: OpX):
        self.op_contour(o, filled=True)

    def _tri(self, o: OpX, kind: str, x, y, z):
        op = o.op
        self.ctx.need("from matplotlib.tri import Triangulation")
        self.emit(f"x, y, z = (np.ravel(np.asarray(v, dtype=float)) for v in ({o.L(x)}, {o.L(y)}, {o.L(z)}))")
        self.emit("ok = np.isfinite(x) & np.isfinite(y) & np.isfinite(z)")
        if op.get("triangle", o.axm.proj == "triangular"):
            self.emit("ok &= (x >= -1e-9) & (y >= -1e-9) & (x + y <= 1 + 1e-9)   # inside the composition triangle")
        self.emit("x, y, z = x[ok], y[ok], z[ok]")
        self.emit("tri = Triangulation(x, y)")
        if op.get("max_edge"):
            self.emit("px, py = x[tri.triangles], y[tri.triangles]")
            self.emit("edges = np.hypot(px - np.roll(px, 1, axis=1), py - np.roll(py, 1, axis=1)).max(axis=1)")
            self.emit(f"tri.set_mask(edges > {float(op['max_edge'])!r})   # hide long, thin triangles")
        if kind == "tripcolor":
            kw = self.cmap_kw(o)
            kw.update(o.pick({"alpha", "zorder", "edgecolors"}))
            self.emit(f"mesh = ax.tripcolor(tri, z, shading={op.get('shading', 'flat')!r}, {kwargs(kw)})")
            self.colorbar(o, "mesh")
            return
        self._contour_common(o, kind == "tricontourf", tri=True)

    def op_tricontour(self, o: OpX):
        self._tri(o, "tricontour", o.X(), o.Y(), o.raw("z"))

    def op_tricontourf(self, o: OpX):
        self._tri(o, "tricontourf", o.X(), o.Y(), o.raw("z"))

    def op_tripcolor(self, o: OpX):
        self._tri(o, "tripcolor", o.X(), o.Y(), o.raw("z"))

    def op_quiver(self, o: OpX):
        op = o.op
        x, y, u, v = o.X(), o.Y(), o.raw("u"), o.raw("v")
        mesh = (not isinstance(u, Raw)) and np.ndim(u) == 2 and np.ndim(x) == 1 and np.ndim(y) == 1
        self.emit(f"x, y = {o.L(x)}, {o.L(y)}")
        if mesh:
            self.emit("x, y = np.meshgrid(x, y)")
        kw = {k: op[k] for k in ("scale", "scale_units", "angles", "width", "pivot", "headwidth", "headlength") if k in op}
        kw.update(o.pick({"color", "alpha", "zorder", "label"}))
        c = o.raw("c")
        if c is not None:
            kw.pop("color", None)
            kw.update(self.cmap_kw(o))
            self.emit(f"q = ax.quiver(x, y, {o.L(u)}, {o.L(v)}, {o.L(c)}, {kwargs(kw)})")
            self.colorbar(o, "q")
        else:
            self.emit(f"q = ax.quiver(x, y, {o.L(u)}, {o.L(v)}{', ' + kwargs(kw) if kw else ''})")
        key = op.get("key")
        if key:
            self.emit(f"ax.quiverkey(q, {lit(key.get('x', 0.9))}, {lit(key.get('y', 1.03))}, {lit(key.get('U', 1.0))}, "
                      f"{str(key.get('label', str(key.get('U', 1))))!r}, labelpos={key.get('labelpos', 'E')!r}, "
                      "coordinates='axes')")

    # =====================================================================================================
    # draw (calphad_tools.figops_extra)
    # =====================================================================================================
    def op_stackplot(self, o: OpX):
        from .figures import find_col
        op = o.op
        if o.has_source and op.get("group_col"):
            vkey = "value" if op.get("value_col") else "y"
            self._pivot_code(o, "x", "group", vkey, var="table", fill=0.0)
            self.emit("x = table.index.to_numpy(float)")
            self.emit("order = np.argsort(x)")
            self.emit("x, ys = x[order], table.to_numpy(float)[order].T   # one layer per group")
            self.emit("labels = [str(v) for v in table.columns]")
        elif o.has_source and op.get("y_cols"):
            df, fr = o.df(), o.frame()
            cols = [find_col(fr, c) for c in op["y_cols"]]
            self.emit(f"x = {o.L(o.raw('x'))}")
            self.emit(f"ys = np.vstack([{', '.join(f'{df}[{c!r}].to_numpy(float)' for c in cols)}])")
            self.emit(f"labels = {lit(op.get('labels') or [str(c) for c in cols])}")
        else:
            ys = op.get("ys", op.get("y"))
            self.emit(f"x = {o.L(o.raw('x'))}")
            self.emit(f"ys = np.atleast_2d(np.asarray({o.L(ys)}, dtype=float))")
            self.emit(f"labels = {lit(op.get('labels') or [f'layer {i + 1}' for i in range(len(np.atleast_2d(ys)))])}")
        self.emit("ys = np.nan_to_num(np.asarray(ys, dtype=float))")
        if op.get("normalize"):
            self.emit("tot = ys.sum(axis=0)")
            self.emit("ys = np.divide(ys, tot, out=np.zeros_like(ys), where=tot != 0)   # layers sum to 1")
        if op.get("drop_zero", True):
            self.emit("keep = np.abs(ys).max(axis=1) > 0   # empty layers are left out")
            self.emit("ys, labels = ys[keep], [lab for lab, k in zip(labels, keep) if k]")
        kw = o.pick({"alpha", "edgecolor", "linewidth", "zorder"})
        if op.get("colors") is not None:
            self.helpers("series_colors")
            kw["colors"] = Raw(f"series_colors({lit(op['colors'])}, labels)")
        tx, ty = o._t("x"), o._t("y")
        xs = f"np.asarray(x, dtype=float) * {tx[0]!r} + {tx[1]!r}" if tx else "x"
        yv = f"ys * {ty[0]!r}" if ty else "ys"
        self.emit(f"ax.stackplot({xs}, {yv}, labels=labels, baseline={op.get('baseline', 'zero')!r}"
                  + (f", {kwargs(kw)}" if kw else "") + ")")

    def _series_code(self, o: OpX, cat_key: str = "category") -> None:
        """Emit ``categories`` and ``series = {name: values}`` like figops_extra._series."""
        op = o.op
        if o.has_source and op.get(f"{cat_key}_col"):
            vkey = next(k for k in ("value", "y", "height") if op.get(f"{k}_col"))
            if op.get("group_col"):
                self._pivot_code(o, cat_key, "group", vkey, var="table")
                self.emit("categories = [str(int(c)) if isinstance(c, float) and c.is_integer() else "
                          "(f'{c:g}' if isinstance(c, float) else str(c)) for c in table.index]")
                self.emit("series = {str(g): table[g].to_numpy(float) for g in table.columns}")
            else:
                df = o.df()
                cc, vc = o.colname(cat_key), o.colname(vkey)
                self.emit(f"means = {df}.groupby({cc!r}, sort=False)[{vc!r}].mean()")
                self.emit("categories = [str(int(c)) if isinstance(c, float) and c.is_integer() else "
                          "(f'{c:g}' if isinstance(c, float) else str(c)) for c in means.index]")
                self.emit(f"series = {{{op.get('label', str(vc))!r}: means.to_numpy(float)}}")
            return
        cats = op.get("categories")
        series = op.get("series")
        if series is None and op.get("values") is not None:
            series = {op.get("label", "values"): op["values"]}
        if not isinstance(series, dict):
            series = {f"series {i + 1}": v for i, v in enumerate(series)}
        self.emit(f"categories = {lit([str(c) for c in cats])}")
        self.emit(f"series = {o.L({str(k): v for k, v in series.items()})}")

    def op_bars(self, o: OpX):
        from .figops_extra import _series
        op = o.op
        cats, series = _series(o.shadow())
        names = list(series)
        self._series_code(o)
        mode = op.get("mode", "grouped")
        horiz = bool(op.get("horizontal"))
        width = float(op.get("width", 0.8))
        t = o._t("x" if horiz else "y")
        kw = o.pick({"edgecolor", "linewidth", "alpha", "zorder", "hatch", "linestyle"})
        errs = op.get("errors") or {}
        vl = op.get("value_labels")
        self.emit("pos = np.arange(len(categories), dtype=float)")
        if op.get("colors") is not None:
            self.helpers("series_colors")
            self.emit(f"colors = series_colors({lit(op['colors'])}, list(series))")
        else:
            self.emit("colors = [None] * len(series)")
        if errs:
            self.emit(f"errors = {lit(errs)}")
        if mode == "stacked":
            self.emit("pos_base, neg_base = np.zeros(len(categories)), np.zeros(len(categories))")
        self.emit("for i, (name, vals) in enumerate(series.items()):")
        self.emit("    vals = np.nan_to_num(np.asarray(vals, dtype=float)" + (f" * {t[0]!r})" if t else ")"))
        if mode == "grouped":
            self.emit(f"    w = {width!r} / len(series)")
            self.emit(f"    p = pos - {width / 2!r} + w * (i + 0.5)   # side by side")
            base = None
        else:
            self.emit(f"    w, p = {width!r}, pos")
            self.emit("    base = np.where(vals >= 0, pos_base, neg_base)   # stacked")
            base = "base"
        ekw = ""
        if errs:
            self.emit("    e = errors.get(name)")
            if t:
                self.emit(f"    e = None if e is None else np.asarray(e, dtype=float) * {t[0]!r}")
            ekw = f", {'xerr' if horiz else 'yerr'}=e, **({{'capsize': {lit(op.get('capsize', 3))}}} if e is not None else {{}})"
        extra = (", " + kwargs(kw)) if kw else ""
        if horiz:
            self.emit(f"    bars = ax.barh(p, vals, height=w{', left=base' if base else ''}{ekw}, color=colors[i], "
                      f"label=name{extra})")
        else:
            self.emit(f"    bars = ax.bar(p, vals, width=w{', bottom=base' if base else ''}{ekw}, color=colors[i], "
                      f"label=name{extra})")
        if mode == "stacked":
            self.emit("    pos_base = pos_base + np.where(vals >= 0, vals, 0)")
            self.emit("    neg_base = neg_base + np.where(vals < 0, vals, 0)")
        if vl:
            from .figops_extra import _pyfmt
            fmt = _pyfmt(vl) if isinstance(vl, str) else "%g"
            self.emit(f"    ax.bar_label(bars, fmt={fmt!r}, label_type={'center' if mode == 'stacked' else 'edge'!r}, "
                      f"fontsize={lit(op.get('label_fontsize', 8))}, padding=2)")
        if horiz:
            self.emit("ax.set_yticks(pos, categories)")
            if op.get("top_down", True):
                self.emit("ax.invert_yaxis()")
        else:
            self.emit("ax.set_xticks(pos, categories)")
            rot = op.get("rotation", 30 if max(len(c) for c in cats) > 6 else 0)
            if rot:
                self.emit("for t in ax.get_xticklabels():")
                self.emit(f"    t.set_rotation({lit(rot)})")
                self.emit("    t.set_ha('right')")
        if len(names) == 1 and not op.get("legend", False):
            self.helpers("hide_legend")
            self.emit("hide_legend(ax)   # a single series needs no legend")
            o.axm.locked = True

    def op_pie(self, o: OpX):
        op = o.op
        self.emit(f"values = {o.L(o.raw('values'))}")
        labels = o.raw("labels")
        if labels is not None:
            self.emit(f"labels = [str(v) for v in {o.L(labels)}]")
        else:
            self.emit("labels = None")
        kw: Dict[str, Any] = {"startangle": op.get("startangle", 90), "counterclock": op.get("counterclock", False)}
        for k in ("explode", "pctdistance", "labeldistance"):
            if k in op:
                kw[k] = op[k]
        if op.get("colors") is not None:
            self.helpers("series_colors")
            kw["colors"] = Raw(f"series_colors({lit(op['colors'])}, labels or list(range(len(values))))")
        ap = op.get("autopct")
        if ap:
            fmt = ap if isinstance(ap, str) else "{:.1f}%"
            kw["autopct"] = Raw(f"lambda pct: {self._fmt_code(fmt, 'pct')}")
        wp = {"edgecolor": op.get("edgecolor", "white"), "linewidth": op.get("linewidth", 1)}
        if op.get("width"):
            wp["width"] = float(op["width"])
            kw.setdefault("pctdistance", 1 - float(op["width"]) / 2)
        as_legend = bool(op.get("legend"))
        self.emit(f"wedges = ax.pie(values, labels={'None' if as_legend else 'labels'}, wedgeprops={lit(wp)}, "
                  f"textprops={{'fontsize': {lit(op.get('fontsize', 9))}}}, {kwargs(kw)})[0]")
        if labels is not None:
            self.emit("for w, lab in zip(wedges, labels):")
            self.emit(f"    w.set_label({'lab' if as_legend else chr(39) + '_' + chr(39) + ' + lab'})")
        if as_legend:
            self.emit(f"ax.legend(wedges, labels, loc='center left', bbox_to_anchor=(1.0, 0.5), "
                      f"fontsize={lit(op.get('fontsize', 9))})")
        else:
            self.helpers("hide_legend")
            self.emit("hide_legend(ax)")
        o.axm.locked = True
        if op.get("center_text"):
            self.emit(f"ax.text(0, 0, {str(op['center_text'])!r}, ha='center', va='center', "
                      f"fontsize={lit(op.get('fontsize', 9) + 2)})")
        self.emit("ax.set_aspect('equal')")

    def op_radar(self, o: OpX):
        op = o.op
        self._series_code(o)
        self.emit("ang = np.linspace(0, 2 * np.pi, len(categories), endpoint=False)")
        self.emit("ax.set_theta_offset(np.pi / 2)")
        self.emit(f"ax.set_theta_direction({-1 if op.get('clockwise', True) else 1})")
        if op.get("colors") is not None:
            self.helpers("series_colors")
            self.emit(f"colors = series_colors({lit(op['colors'])}, list(series))")
        else:
            self.emit("colors = [f'C{i % 10}' for i in range(len(series))]")
        self.emit("for (name, v), col in zip(series.items(), colors):")
        self.emit("    a2, v2 = np.append(ang, ang[0]), np.append(np.asarray(v, dtype=float), v[0])   # closed polygon")
        self.emit(f"    ax.plot(a2, v2, color=col, linewidth={lit(op.get('linewidth', 1.5))}, "
                  f"marker={lit(op.get('marker', 'o'))}, markersize={lit(op.get('markersize', 4))}, label=name)")
        if op.get("fill", True):
            self.emit(f"    ax.fill(a2, v2, color=col, alpha={lit(op.get('alpha', 0.15))}, label='_' + name)")
        self.emit("ax.set_xticks(ang, categories)")
        if op.get("rlim"):
            self.emit(f"ax.set_ylim({o.args(*op['rlim'])})")
        if op.get("rticks") is not None:
            self.emit(f"ax.set_yticks({lit(op['rticks'])})")
        self.emit(f"ax.tick_params(axis='y', labelsize={lit(op.get('fontsize', 8))})")
        if op.get("legend", True):
            self.emit(f"ax.legend(loc='upper left', bbox_to_anchor=(1.08, 1.05), fontsize={lit(op.get('fontsize', 8) + 1)})")
            o.axm.locked = True

    # -- 3-D --------------------------------------------------------------------------------------------
    def op_surface(self, o: OpX):
        op = o.op
        x, y, Z, scattered = self._grid(o)
        kw = o.pick({"alpha", "edgecolor", "linewidth", "zorder"})
        if op.get("color") and not op.get("cmap"):
            kw["color"] = op["color"]
        else:
            kw.update(self.cmap_kw(o))
        if scattered:
            self.emit(f"surf = ax.plot_trisurf(np.ravel({o.L(x)}), np.ravel({o.L(y)}), np.ravel({o.L(Z)}), {kwargs(kw)})")
        else:
            self.emit(f"X, Y = np.meshgrid({x}, {y}) if np.ndim({x}) == 1 else ({x}, {y})")
            self.emit(f"surf = ax.plot_surface(X, Y, np.ma.masked_invalid({Z}), rstride={lit(op.get('rstride', 1))}, "
                      f"cstride={lit(op.get('cstride', 1))}, {kwargs(kw)})")
        if op.get("label"):
            self.emit(f"surf.set_label({op['label']!r})")
        self.colorbar(o, "surf", three_d=True)

    def op_wireframe(self, o: OpX):
        op = o.op
        x, y, Z, scattered = self._grid(o)
        if scattered:
            raise OpCodeError("wireframe needs gridded data")
        self.emit(f"X, Y = np.meshgrid({x}, {y}) if np.ndim({x}) == 1 else ({x}, {y})")
        kw = o.pick({"color", "linewidth", "alpha", "label", "zorder", "linestyle"})
        self.emit(f"ax.plot_wireframe(X, Y, {Z}, rstride={lit(op.get('rstride', 1))}, cstride={lit(op.get('cstride', 1))}"
                  + (f", {kwargs(kw)}" if kw else "") + ")")

    def op_scatter3d(self, o: OpX):
        op = o.op
        kw = o.pick({"color", "marker", "alpha", "label", "zorder", "edgecolor"})
        s = o.raw("s")
        if s is not None:
            kw["s"] = s
        c = o.raw("c")
        if c is not None:
            kw.pop("color", None)
            kw["c"] = c
            kw.update(self.cmap_kw(o))
        kw["depthshade"] = op.get("depthshade", True)
        code = f"ax.scatter({o.args(o.X(), o.Y(), o.raw('z'), **kw)})"
        if c is not None and op.get("colorbar"):
            self.emit(f"sc = {code}")
            self.colorbar(o, "sc", three_d=True)
        else:
            self.emit(code)

    def op_line3d(self, o: OpX):
        kw = o.pick({"color", "linestyle", "linewidth", "alpha", "zorder", "marker", "markersize", "label"})
        self.emit(f"ax.plot({o.args(o.X(), o.Y(), o.raw('z'), **kw)})")

    def op_view(self, o: OpX):
        op = o.op
        if any(k in op for k in ("elev", "azim", "roll")):
            self.emit(f"ax.view_init(elev={lit(op['elev']) if 'elev' in op else 'ax.elev'}, "
                      f"azim={lit(op['azim']) if 'azim' in op else 'ax.azim'}, "
                      f"roll={lit(op['roll']) if 'roll' in op else 'ax.roll'})")
        for k in ("xlabel", "ylabel", "zlabel"):
            if k in op:
                self.emit(f"ax.set_{k}({op[k]!r})")
        if op.get("zlim"):
            self.emit(f"ax.set_zlim({o.args(*op['zlim'])})")
        if op.get("box_aspect"):
            self.emit(f"ax.set_box_aspect({lit(op['box_aspect'])})")
        if op.get("proj_type"):
            self.emit(f"ax.set_proj_type({op['proj_type']!r})")
        if op.get("zscale"):
            self.emit(f"ax.set_zscale({op['zscale']!r})")

    # =====================================================================================================
    # axes & layout (calphad_tools.figops_extra)
    # =====================================================================================================
    def _secondary(self, o: OpX, which: str):
        from .expdata import ATOMIC_MASS
        from .figops_extra import _EL_RX
        op, axm = o.op, o.axm
        conv = op.get("convert", "linear")
        if isinstance(conv, dict):
            op = {**op, **{k: v for k, v in conv.items() if k != "type"}}
            conv = conv.get("type", conv.get("kind", "linear"))
        kind = str(conv).lower()
        t = axm.xt if which == "x" else axm.yt
        pre = None                  # standalone axis values are displayed ones; 'of: data' maps back to data first
        if op.get("of", "displayed") != "displayed" and t != (1.0, 0.0):
            pre = t
        name = self.fresh(f"sec{which}")
        fwd, inv = f"{name}_forward", f"{name}_inverse"
        u = f"(np.asarray(u, dtype=float) - {pre[1]!r}) / {pre[0]!r}" if pre else "np.asarray(u, dtype=float)"
        unpre = (lambda s: f"({s}) * {pre[0]!r} + {pre[1]!r}") if pre else (lambda s: s)
        if kind == "linear":
            a, b = float(op.get("scale", 1.0)), float(op.get("offset", 0.0))
            def lin(v, a, b, inverse=False):
                if inverse:
                    e = v if b == 0 else f"({v} {'-' if b > 0 else '+'} {abs(b)!r})"
                    return e if a == 1 else f"{e} / {a!r}"
                e = v if a == 1 else f"{v} * {a!r}"
                return e if b == 0 else f"{e} {'+' if b > 0 else '-'} {abs(b)!r}"
            self.emit(f"def {fwd}(u):   # secondary {which}-axis value = axis value * {a:g} + {b:g}")
            self.emit(f"    return {lin(u, a, b)}")
            self.emit(f"def {inv}(u):")
            self.emit(f"    return {unpre(lin('np.asarray(u, dtype=float)', a, b, inverse=True))}")
        elif kind == "reciprocal":
            k = float(op.get("k", 1000.0))
            for fn, arg in ((fwd, u), (inv, "np.asarray(u, dtype=float)")):
                self.emit(f"def {fn}(u):   # {k:g} / value (inf at 0)")
                self.emit(f"    u = np.array({arg}, dtype=float, copy=True)")
                self.emit("    near = np.isclose(u, 0.0)")
                self.emit("    u[near] = np.inf")
                self.emit(f"    u[~near] = {k!r} / u[~near]")
                self.emit(f"    return {unpre('u') if fn == inv else 'u'}")
        elif kind in ("log10", "pow10"):
            f1 = "np.log10(np.clip({}, 1e-300, None))"
            f2 = "np.power(10.0, {})"
            a, b = (f1, f2) if kind == "log10" else (f2, f1)
            self.emit(f"def {fwd}(u):")
            self.emit(f"    return {a.format(u)}")
            self.emit(f"def {inv}(u):")
            self.emit(f"    return {unpre(b.format('np.asarray(u, dtype=float)'))}")
        elif kind in ("mole_to_mass", "mass_to_mole"):
            el = op.get("element")
            if not el:
                m = _EL_RX.search((axm.xlabel if which == "x" else axm.ylabel) or "")
                if not m:
                    raise OpCodeError("element of the axis unknown")
                el = m.group(1)
            other = op.get("other")
            if not other:
                for tt in [axm.title, getattr(self, "suptitle", "")]:
                    syms = [s.upper() for s in re.findall(r"\b([A-Z][a-z]?)\b", re.split(r"\s{2,}|\(", tt or "")[0])
                            if s.upper() in ATOMIC_MASS]
                    cands = [s for s in dict.fromkeys(syms) if s != str(el).upper()]
                    if len(syms) == 2 and len(cands) == 1:
                        other = cands[0]
                        break
            if not other:
                raise OpCodeError("second element of the binary unknown")
            m_el, m_other = ATOMIC_MASS[str(el).upper()], ATOMIC_MASS[str(other).upper()]
            self.emit(f"def comp_convert(x, m_el, m_other, to_mass):   # binary mole <-> mass fraction of {str(el).upper()}")
            self.emit("    x = np.asarray(x, dtype=float)")
            self.emit("    a, b = (m_el, m_other) if to_mass else (1.0 / m_el, 1.0 / m_other)")
            self.emit("    xc = np.clip(x, 0.0, 1.0)")
            self.emit("    y = xc * a / (xc * a + (1 - xc) * b)")
            self.emit("    y = np.where(x < 0, x * a / b, y)   # extended linearly outside 0..1")
            self.emit("    return np.where(x > 1, 1 + (x - 1) * b / a, y)")
            if "in_scale" in op:
                self.emit(f"scale_in = {float(op['in_scale'])!r}")
            else:
                self.emit(f"lo, hi = ax.{which}axis.get_view_interval()")
                self.emit("scale_in = 100.0 if max(abs(lo), abs(hi)) > 1.5 else 1.0   # percent or fraction axis")
            pct = op.get("percent")
            if "out_scale" in op:
                self.emit(f"scale_out = {float(op['out_scale'])!r}")
            elif pct is not None:
                self.emit(f"scale_out = {100.0 if pct else 1.0!r}")
            else:
                self.emit("scale_out = scale_in")
            to_mass = kind == "mole_to_mass"
            self.emit(f"def {fwd}(u, s_in=scale_in, s_out=scale_out):")
            self.emit(f"    return comp_convert({u} / s_in, {m_el!r}, {m_other!r}, {to_mass}) * s_out")
            self.emit(f"def {inv}(u, s_in=scale_in, s_out=scale_out):")
            self.emit(f"    return {unpre(f'comp_convert(np.asarray(u, dtype=float) / s_out, {m_el!r}, {m_other!r}, {not to_mass}) * s_in')}")
        else:
            raise OpCodeError(f"convert {kind!r}")
        loc = op.get("location", "top" if which == "x" else "right")
        self.emit(f"{name} = ax.secondary_{which}axis({lit(loc)}, functions=({fwd}, {inv}))")
        if op.get("label"):
            self.emit(f"{name}.set_{which}label({op['label']!r})")
        if op.get("format"):
            self.emit(f"{name}.{which}axis.set_major_formatter({self.formatter(op['format'], op)})")
        if op.get("ticks") is not None:
            self.emit(f"{name}.{which}axis.set_ticks({lit(op['ticks'])})")
        if op.get("color"):
            self.emit(f"{name}.tick_params(axis={which!r}, colors={op['color']!r})")
            self.emit(f"{name}.{which}axis.label.set_color({op['color']!r})")
        if op.get("labelsize"):
            self.emit(f"{name}.tick_params(axis={which!r}, labelsize={lit(op['labelsize'])})")

    def op_secondary_xaxis(self, o: OpX):
        self._secondary(o, "x")

    def op_secondary_yaxis(self, o: OpX):
        self._secondary(o, "y")

    def op_twiny(self, o: OpX):
        name = self.fresh("twin")
        self.emit(f"{name} = ax.twiny()   # twin x-axis (top)")
        if o.op.get("xlabel"):
            self.emit(f"{name}.set_xlabel({o.op['xlabel']!r})")
        if o.op.get("xlim"):
            self.emit(f"{name}.set_xlim({o.args(*o.op['xlim'])})")
        self.add_top(AxModel(name, "rectilinear", "twin", "linear", o.axm.yscale, yt=o.axm.yt, locked=True,
                             spec=o.axm.spec, bounds=o.axm.bounds))

    def op_ticks(self, o: OpX):
        op, axm = o.op, o.axm
        which = {"x": ["x"], "y": ["y"], "both": ["x", "y"], "z": ["z"]}[str(op.get("axis", "both")).lower()]
        self.ctx.need("from matplotlib import ticker as mticker")
        for a in which:
            scaled = (axm.xt if a == "x" else axm.yt) != (1.0, 0.0) if a in ("x", "y") else False
            scale = getattr(axm, f"{a}scale", "linear") if a in ("x", "y") else "linear"
            ax_ = f"ax.{a}axis"
            if op.get("format"):
                self.emit(f"{ax_}.set_major_formatter({self.formatter(op['format'], op, scaled)})")
            if op.get("nbins") or op.get("step"):
                if op.get("step"):
                    self.emit(f"{ax_}.set_major_locator(mticker.MultipleLocator({float(op['step'])!r}))")
                else:
                    self.emit(f"{ax_}.set_major_locator(mticker.MaxNLocator({int(op['nbins'])}))")
            if "minor" in op:
                m = op["minor"]
                if m is False:
                    self.emit(f"{ax_}.set_minor_locator(mticker.NullLocator())")
                elif scale == "linear":
                    self.emit(f"{ax_}.set_minor_locator(mticker.AutoMinorLocator({'' if m is True else int(m)}))")
                else:
                    self.emit(f"{ax_}.set_minor_locator(mticker.LogLocator(subs='auto'))")
            tp = {k: op[k] for k in ("direction", "length", "width", "labelsize", "color", "pad", "labelcolor") if k in op}
            if tp:
                self.emit(f"ax.tick_params(axis={a!r}, which='major', {kwargs(tp)})")
                mtp = dict(tp)
                if "length" in mtp:
                    mtp["length"] = 0.6 * float(mtp["length"])
                mtp.pop("labelsize", None)
                if mtp:
                    self.emit(f"ax.tick_params(axis={a!r}, which='minor', {kwargs(mtp)})")
            if op.get("rotation") is not None:
                self.emit(f"ax.tick_params(axis={a!r}, labelrotation={lit(op['rotation'])})")
                if a == "x" and op.get("rotation") not in (0, 90):
                    self.emit("for t in ax.get_xticklabels():")
                    self.emit(f"    t.set_ha({op.get('ha', 'right')!r})")
            mirror = op.get("mirror")
            if a == "x" and (mirror or "top" in op):
                self.emit(f"ax.tick_params(axis='x', which='both', top={bool(op.get('top', mirror))})")
            if a == "y" and (mirror or "right" in op):
                self.emit(f"ax.tick_params(axis='y', which='both', right={bool(op.get('right', mirror))})")
            g = op.get("grid")
            if g is not None:
                if g in (False, "none", "off"):
                    self.emit(f"ax.grid(False, axis={a!r}, which='both')")
                else:
                    g = "major" if g is True else g
                    if g in ("minor", "both") and "minor" not in op:
                        self.emit(f"if isinstance({ax_}.get_minor_locator(), mticker.NullLocator):")
                        self.emit(f"    {ax_}.set_minor_locator(mticker.AutoMinorLocator())")
                    if g in ("major", "both"):
                        self.emit(f"ax.grid(True, axis={a!r}, which='major', alpha=0.35)")
                    if g in ("minor", "both"):
                        self.emit(f"ax.grid(True, axis={a!r}, which='minor', alpha=0.18, linestyle=':')")

    def op_spines(self, o: OpX):
        op = o.op

        def names(v):
            return "list(ax.spines)" if v == "all" else lit([v] if isinstance(v, str) else list(v))
        if op.get("hide"):
            self.emit(f"for name in {names(op['hide'])}:")
            self.emit("    ax.spines[name].set_visible(False)")
            hide = op["hide"]
            hl = ["top", "right"] if hide == "all" else ([hide] if isinstance(hide, str) else list(hide))
            if "top" in hl:
                self.emit("ax.xaxis.set_ticks_position('bottom')")
            if "right" in hl:
                self.emit("ax.yaxis.set_ticks_position('left')")
        if op.get("show"):
            self.emit(f"for name in {names(op['show'])}:")
            self.emit("    ax.spines[name].set_visible(True)")
        off = op.get("offset")
        if off is not None:
            if isinstance(off, dict):
                for n, v in off.items():
                    self.emit(f"ax.spines[{n!r}].set_position(('outward', {float(v)!r}))")
            else:
                self.emit("for spine in ax.spines.values():")
                self.emit("    if spine.get_visible():")
                self.emit(f"        spine.set_position(('outward', {float(off)!r}))")
        if op.get("zero"):
            self.emit(f"for name in {names(op['zero'])}:")
            self.emit("    ax.spines[name].set_position('zero')")
        if "linewidth" in op or "color" in op:
            self.emit("for spine in ax.spines.values():")
            if "linewidth" in op:
                self.emit(f"    spine.set_linewidth({lit(op['linewidth'])})")
            if "color" in op:
                self.emit(f"    spine.set_edgecolor({op['color']!r})")

    def _axes_list(self, o: OpX, key: str = "axes_list") -> Optional[List[AxModel]]:
        lst = o.op.get(key)
        if lst is None:
            return None
        main = self.fm.main()
        return [main[int(i)] for i in lst]

    def op_colorbar(self, o: OpX):
        op = o.op
        axes = self._axes_list(o) or ([o.axm] if o.axm is not None else [])
        self.helpers("mappables")
        exprs = [a.expr for a in axes]
        self.emit(f"cb_axes = [{', '.join(exprs)}]")
        self.emit("maps = [m for a in cb_axes for m in mappables(a)]")
        if op.get("mappable_axes") is not None:
            src = self.fm.main()[int(op["mappable_axes"])]
            self.emit(f"m = mappables({src.expr})[-1]")
        else:
            self.emit("m = maps[-1]")
        if op.get("sync", True) and len(axes) >= 1:
            self.emit("if len(maps) > 1:   # one colour scale for all panels")
            self.emit("    lims = np.array([mm.get_clim() for mm in maps], dtype=float)")
            self.emit("    for mm in maps:")
            self.emit("        if mm.norm is m.norm or type(mm.norm) is type(m.norm):")
            self.emit("            mm.set_clim(np.nanmin(lims[:, 0]), np.nanmax(lims[:, 1]))")
        location = op.get("location", "right")
        kw: Dict[str, Any] = {k: op[k] for k in ("extend",) if k in op}
        if op.get("format"):
            kw["format"] = Raw(self.formatter(op["format"], op))
        orient = "vertical" if location in ("right", "left") else "horizontal"
        if len(axes) > 1:
            self.helpers("shared_colorbar_axes")
            self.emit(f"cax = shared_colorbar_axes(fig, cb_axes, {location!r}, {float(op.get('fraction', 0.05))!r}, "
                      f"{float(op.get('pad', 0.08))!r}, {lit(op.get('wspace', 0.3))}, {lit(op.get('hspace', 0.35))})")
            self.emit("if cax is not None:")
            self.emit(f"    cb = fig.colorbar(m, cax=cax, orientation={orient!r}, location={location!r}"
                      + (f", {kwargs(kw)})" if kw else ")"))
            self.emit("else:")
            self.emit(f"    cb = fig.colorbar(m, ax=cb_axes, location={location!r}, "
                      f"fraction={float(op.get('fraction', 0.05))!r}, pad={float(op.get('pad', 0.05))!r}"
                      + (f", {kwargs(kw)})" if kw else ")"))
        else:
            self.emit(f"cb = fig.colorbar(m, ax=cb_axes[0], location={location!r}, "
                      f"fraction={float(op.get('fraction', 0.15))!r}, pad={float(op.get('pad', 0.05))!r}"
                      + (f", {kwargs(kw)})" if kw else ")"))
        if op.get("label"):
            self.emit(f"cb.set_label({op['label']!r})")
        if op.get("ticks") is not None:
            self.emit(f"cb.set_ticks({lit(op['ticks'])})")

    def op_panel_labels(self, o: OpX):
        from .figops_extra import _POS
        op = o.op
        axes = self._axes_list(o)
        if axes is None:
            axes = [a for a in self.fm.main() if a.kind not in ("twin", "broken")]
        labels = op.get("labels")
        if labels is None:
            start = ord(str(op.get("start", "a")).lower()[0]) - ord("a")
            labels = []
            for i in range(len(axes)):
                k = start + i
                s = chr(ord("a") + k % 26) * (k // 26 + 1)
                labels.append(s.upper() if op.get("uppercase") else s)
            if op.get("parentheses", True):
                labels = [f"({s})" for s in labels]
        pos = op.get("position", "outside top left")
        if isinstance(pos, str):
            xy, ha, va = _POS[pos]
        else:
            xy, ha, va = tuple(pos), op.get("ha", "left"), op.get("va", "top")
        self.emit(f"for a, label in zip([{', '.join(a.expr for a in axes)}], {lit(list(labels)[:len(axes)])}):")
        self.emit(f"    a.annotate(label, xy={lit(tuple(xy))}, xycoords='axes fraction', ha={ha!r}, va={va!r}, "
                  f"fontsize={lit(op.get('fontsize', 12))},")
        self.emit(f"               fontweight={'bold' if op.get('bold', True) else 'normal'!r}, "
                  f"color={op.get('color', 'k')!r}, annotation_clip=False)")

    def op_broken_yaxis(self, o: OpX):
        op, axm = o.op, o.axm
        (lo1, hi1), (lo2, hi2) = [sorted(o.conv(list(v), "y")) for v in op["ylims"]]
        rb, rt = [float(v) for v in op.get("ratio", [1, 1])]
        gap = float(op.get("gap", 0.08))
        name = self.fresh("upper")
        self.helpers("copy_artists")
        self.emit("# broken y-axis: the panel is split; the upper part is a copy with its own y range")
        self.emit("ss = ax.get_subplotspec()")
        self.emit("if ss is not None:")
        self.emit(f"    sub = ss.subgridspec(2, 1, height_ratios=[{rt!r}, {rb!r}], hspace={gap!r})")
        self.emit("    ax.set_subplotspec(sub[1])")
        self.emit(f"    {name} = fig.add_subplot(sub[0], sharex=ax)")
        self.emit("else:")
        self.emit("    x0, y0, w, h = ax.get_position().bounds")
        tot = rb + rt + gap * (rb + rt) / 2
        self.emit(f"    hb, ht, g = h * {rb / tot!r}, h * {rt / tot!r}, h * {gap * (rb + rt) / 2 / tot!r}")
        self.emit("    ax.set_position([x0, y0, w, hb])")
        self.emit(f"    {name} = fig.add_axes([x0, y0 + hb + g, w, ht], sharex=ax)")
        self.emit(f"copy_artists(ax, {name})")
        self.emit(f"{name}.set_yscale(ax.get_yscale())")
        self.emit(f"ax.set_ylim({lit(lo1)}, {lit(hi1)})")
        self.emit(f"{name}.set_ylim({lit(lo2)}, {lit(hi2)})")
        self.emit("ax.spines['top'].set_visible(False)")
        self.emit(f"{name}.spines['bottom'].set_visible(False)")
        self.emit("ax.xaxis.tick_bottom()")
        self.emit(f"{name}.tick_params(axis='x', which='both', bottom=False, top=False, labelbottom=False)")
        self.emit("if ax.get_title():")
        self.emit(f"    {name}.set_title(ax.get_title())")
        self.emit("    ax.set_title('')")
        self.emit("if ax.yaxis.get_gridlines() and ax.yaxis._major_tick_kw.get('gridOn'):")
        self.emit(f"    {name}.grid(True, alpha=0.35)")
        self.emit("if ax.get_ylabel():   # one label centred on both parts")
        self.emit(f"    ax.yaxis.label.set_y({(rb + rt + gap * (rb + rt) / 2) / (2 * rb)!r})")
        if op.get("marks", True):
            self.emit("marks = dict(marker=[(-1, -0.5), (1, 0.5)], markersize=10, linestyle='none', color='k', mec='k', "
                      "mew=1, clip_on=False, label='_break')")
            self.emit(f"{name}.plot([0, 1], [0, 0], transform={name}.transAxes, **marks)")
            self.emit("ax.plot([0, 1], [1, 1], transform=ax.transAxes, **marks)")
        new = self.add_top(AxModel(name, "rectilinear", "broken", axm.xscale, axm.yscale, locked=True))
        new.title = axm.title
        if op.get("legend", "upper") == "upper":
            self.emit("leg = ax.get_legend()")
            self.emit("if leg is not None:   # the legend moves to the upper part")
            self.emit(f"    {name}.legend(list(leg.legend_handles), [t.get_text() for t in leg.get_texts()], "
                      f"loc={op.get('loc', 'best')!r})")
            self.emit("    leg.set_visible(False)")
            self.emit("elif ax.get_legend_handles_labels()[0]:")
            self.emit(f"    {name}.legend(*{name}.get_legend_handles_labels(), loc={op.get('loc', 'best')!r})")
            self.helpers("hide_legend")
            self.emit("    hide_legend(ax)")
            # locked only if a legend existed or labelled handles exist: decided at run time in the toolkit
            axm.locked = axm.locked or self._has_labels_hint
        if axm.title:
            axm.title = ""

    _has_labels_hint = True

    def op_annotate_points(self, o: OpX):
        op = o.op
        self.helpers("label_points")
        x, y = o.X(), o.Y()
        labels = o.raw("labels")
        self.emit(f"x, y = {o.arr(x)}, {o.arr(y)}")
        self.emit(f"labels = [('' if v is None else (v if isinstance(v, str) else (str(int(v)) if float(v).is_integer() "
                  f"else f'{{v:g}}'))) for v in np.atleast_1d({o.L(labels)})]")
        if op.get("scatter"):
            tr = o.tr()
            self.emit(f"ax.scatter(x, y, {kwargs({**({'transform': tr} if tr else {}), 's': op.get('s', 20), 'marker': op.get('marker', 'o'), 'color': op.get('color', 'k'), 'zorder': 3})})")
        style = {"fontsize": op.get("fontsize", 8), "color": op.get("text_color", op.get("color", "k"))}
        if op.get("bbox"):
            style["bbox"] = {"boxstyle": "round,pad=0.2", "facecolor": "white", "alpha": 0.8, "edgecolor": "0.6"}
        if op.get("arrow"):
            style["arrowprops"] = {"arrowstyle": "-", "color": "0.4", "lw": 0.6}
        kw = {"offset": tuple(float(v) for v in op.get("offset", [5, 5])), "avoid": bool(op.get("avoid", True))}
        if o.tr():
            kw["transform"] = o.tr()
        self.emit(f"label_points(ax, x, y, labels, {kwargs({**kw, **style})})")

    def op_label_lines(self, o: OpX):
        op = o.op
        self.helpers("label_lines")
        m = op.get("match")
        ms = [m] if isinstance(m, str) else (list(m) if m else None)
        cond = "not ln.get_label().startswith('_')"
        if ms:
            cond += f" and any(s.lower() in ln.get_label().lower() for s in {lit(ms)})"
        self.emit(f"lines = [ln for ln in ax.get_lines() if {cond}]")
        kw: Dict[str, Any] = {}
        if op.get("x") is not None:
            kw["x"] = float(o.conv([op["x"]], "x")[0])
        if op.get("where", "end") != "end":
            kw["where"] = op["where"]
        if op.get("offset") is not None:
            kw["offset"] = [float(v) for v in op["offset"]]
        kw["fontsize"] = op.get("fontsize", 9)
        if op.get("bbox"):
            kw["bbox"] = True
        if not op.get("avoid", True):
            kw["avoid"] = False
        self.emit(f"label_lines(ax, lines, {kwargs(kw)})")
        if op.get("legend") is False:
            self.helpers("hide_legend")
            self.emit("hide_legend(ax)")
            o.axm.locked = True

    # -- data placed by composition (orientation-aware) -----------------------------------------------
    def op_composition_points(self, o: OpX):
        from .orientation import (KELVIN as _K, diagram_recipe, masses_for, parse_unit, place_points, point_style,
                                  points_input, recipe_axes_text)
        op = o.op
        real = getattr(o.axm, "real", None)
        try:
            rec = diagram_recipe(real)
        except ValueError as e:
            raise OpCodeError(f"composition_points: {e}") from None
        frame = o.frame() if o.has_source else None
        inp = points_input(op, frame, rec)
        res = place_points(rec, inp, op.get("T_tolerance"))          # validates like the tool
        kind, els = rec["kind"], list(rec["elements"])
        weight, scale = parse_unit(inp["unit"])
        self.emit(f"# composition points -> diagram coordinates ({recipe_axes_text(rec)[13:]})")
        if frame is not None:
            df = o.df()
            self.emit("comp = {" + ", ".join(f"{e!r}: pd.to_numeric({df}[{c!r}]).to_numpy(float)"
                                             for e, c in inp["cols"].items()) + f"}}   # {inp['unit']}")
            T_code = f"pd.to_numeric({df}[{inp['T_col']!r}]).to_numpy(float)" if inp["T_col"] else None
        else:
            self.emit("comp = {" + ", ".join(f"{e!r}: np.array({lit(np.asarray(v).tolist())})"
                                             for e, v in inp["given"].items()) + f"}}   # {inp['unit']}")
            T_code = (f"np.array({lit([None if np.isnan(t) else float(t) for t in inp['T']])}, dtype=float)"
                      if inp["T"] is not None else None)
        if scale != 1.0:
            self.emit(f"comp = {{el: v / {scale!r} for el, v in comp.items()}}   # percent -> fractions")
        if res["balance"]:
            self.emit(f"comp[{res['balance']!r}] = 1.0 - sum(comp.values())   # balance element")
        else:
            self.emit("total = sum(comp.values())\ncomp = {el: v / total for el, v in comp.items()}")
        if weight:
            self.emit(f"masses = {lit(res['masses'])}   # g/mol")
            self.emit("amount = {el: w / masses[el] for el, w in comp.items()}   # mass fractions -> moles")
            self.emit("X = {el: n / sum(amount.values()) for el, n in amount.items()}   # mole fractions")
        else:
            self.emit("X = comp   # mole fractions")
        c = rec.get("comp") or {}
        if c.get("weight"):            # the diagram shows mass fractions
            self.emit(f"m_axis = {lit({e: float(c['masses'][e]) for e in els})}   # g/mol (the diagram's axis)")
            self.emit("shown = {el: X[el] * m_axis[el] / sum(X[e] * m_axis[e] for e in m_axis) for el in m_axis}")
        else:
            self.emit("shown = X")
        s = float(c.get("scale", 1.0))
        if T_code is not None:
            self.emit(f"T_K = {T_code}" + (" + 273.15   # °C -> K" if inp["T_unit"] == "C" else "   # K"))
        if kind == "ternary":
            self.emit(f"x, y = shown[{els[1]!r}]{f' * {s!r}' if s != 1 else ''}, shown[{els[2]!r}]"
                      f"{f' * {s!r}' if s != 1 else ''}   # triangular axes: X(bottom-right), X(top)")
        else:
            x_el = els[1] if kind == "binary" else rec["axis_element"]
            xs, xo = o.axm.xt if o.axm is not None else (1.0, 0.0)
            ys, yo = o.axm.yt if o.axm is not None else (1.0, 0.0)
            a, b = rec.get("tmap", (1.0, 0.0))
            fx = s * xs
            A, B = ys / a, yo - b * ys / a
            xcode = f"shown[{x_el!r}]" + (f" * {fx!r}" if fx != 1 else "") + (f" + {xo!r}" if xo else "")
            if abs(A - 1) < 1e-12 and abs(B) < 1e-9:
                ycode = "T_K   # K as on the y axis"
            elif abs(A - 1) < 1e-12 and abs(B + _K) < 1e-9:
                ycode = "T_K - 273.15   # °C as on the y axis"
            else:
                ycode = f"T_K * {A!r} + {B!r}"
            self.emit(f"x = {xcode}   # x axis: {'mass' if c.get('weight') else 'mole'} "
                      f"{'percent' if s == 100 else 'fraction'} {x_el.capitalize()}")
            self.emit(f"y = {ycode}")
        keep = res["keep"]
        if not keep.all():
            self.emit(f"keep = ~(np.abs(T_K - {float(rec['T_K'])!r}) > {float(op['T_tolerance'])!r})   "
                      "# points near the section temperature (or without T)")
        else:
            self.emit("keep = np.ones(len(x), dtype=bool)")
        kw = point_style(op)
        self.emit(f"ax.scatter(x[keep], y[keep], {kwargs(kw)})")
        if inp["groups"]:
            self.ctx.need("from matplotlib.collections import LineCollection")
            self.emit(f"groups = {lit(inp['groups'])}   # points joined by tie-lines (indices)")
            self.emit("segs = []\nfor g in groups:\n    idx = [i for i in g if keep[i]]\n"
                      "    ring = idx + [idx[0]] if len(idx) >= 3 else idx\n"
                      "    segs += [[(x[i], y[i]), (x[j], y[j])] for i, j in zip(ring, ring[1:])]")
            lc = {"colors": [op.get("tieline_color", kw.get("color", "k"))],
                  "linewidths": float(op.get("tieline_width", 0.9)), "linestyles": op.get("tieline_style", "solid"),
                  "zorder": kw["zorder"] - 0.5}
            self.emit(f"ax.add_collection(LineCollection(segs, {kwargs(lc)}))")

    # -- phase-region labels -------------------------------------------------------------------------
    def _label_helpers(self):
        from .codegen_calc import label_helpers
        h = label_helpers(self.ctx)
        if h:
            self.pre.append(h)

    def op_label_phase_regions(self, o: OpX):
        from .calc import get_db
        from .codegen_calc import label_code
        from .figures import ScaledFormatter
        from .region_labels import STYLE_KEYS
        op = o.op
        real = getattr(o.axm, "real", None)
        d = getattr(real, "_calphad_diagram", None) if real is not None else None
        if not d or "recipe" not in d:
            raise OpCodeError("label_phase_regions: the figure has no phase-diagram recipe")
        rec = dict(d["recipe"])
        if rec["kind"] == "binary":       # standalone figures hold the values as displayed (e.g. °C, not K)
            fy, fx = real.yaxis.get_major_formatter(), real.xaxis.get_major_formatter()
            if isinstance(fy, ScaledFormatter):
                a, b = rec.get("tmap", (1.0, 0.0))
                rec["tmap"] = [a / fy.scale, b - a * fy.offset / fy.scale]
            if isinstance(fx, ScaledFormatter):
                if fx.offset:
                    raise OpCodeError("label_phase_regions: shifted composition axis")
                rec["comp"] = {**rec.get("comp", {}), "scale": float(rec.get("comp", {}).get("scale", 1.0)) * fx.scale}
        self._label_helpers()
        self.ctx.need("from pycalphad import Database")
        style = {k: op[k] for k in STYLE_KEYS if k in op}
        self.emit(f"db_labels = Database(str({self.ctx.tdb(rec['tdb'])}))")
        self.emit(label_code(self.ctx, rec, get_db(rec["tdb"]), None, mapped=False,
                             invariants=op.get("invariants", True), phase_labels=op.get("phase_labels"),
                             style=style, db_expr="db_labels").rstrip("\n"))

    def _find_label(self, o: OpX) -> str:
        op = o.op
        if not op.get("id") and op.get("text") is None:
            raise OpCodeError("label op without id or text")
        self._label_helpers()
        return f"find_labels(ax, id={op.get('id')!r}, text={op.get('text')!r})"

    def op_move_label(self, o: OpX):
        op = o.op
        target = self._find_label(o)
        if op.get("x") is not None and op.get("y") is not None:
            x, y = o.first(o.X()), o.first(o.Y())
            self.emit(f"for t in {target}:\n    move_label(t, xy=({lit(x)}, {lit(y)}))")
        else:
            self.emit(f"for t in {target}:\n    move_label(t, {lit(float(op.get('dx', 0.0)))}, "
                      f"{lit(float(op.get('dy', 0.0)))})   # points")

    def op_delete_label(self, o: OpX):
        self.emit(f"for t in {self._find_label(o)}:\n    remove_label(t)")

    def op_relabel(self, o: OpX):
        op = o.op
        body = []
        if op.get("new_text") is not None:
            body.append(f"t.set_text({str(op['new_text'])!r})")
        for k in ("fontsize", "color", "rotation", "fontweight"):
            if k in op:
                body.append(f"t.set_{k}({lit(op[k])})")
        self.emit(f"for t in {self._find_label(o)}:\n" + "\n".join("    " + b for b in (body or ["pass"])))


# =========================================================================================================
# generators: make_plot / edit_figure
# =========================================================================================================
def _assemble(ctx: Ctx, w: OpWriter, body: List[str]) -> str:
    return "".join(w.pre) + "\n".join(body) + "\n"


def _fallback_on(ctx: Ctx, fn):
    try:
        return fn()
    except OpCodeError as e:
        ctx.fallback(str(e))


def _rc_code(rcp: Dict) -> str:
    items = []
    for k, v in rcp.items():
        if k == "axes.prop_cycle":
            items.append(f"    {k!r}: cycler(color={lit(list(v.by_key()['color']))}),")
        else:
            items.append(f"    {k!r}: {lit(v)},")
    return "{\n" + "\n".join(items) + "\n}"


@generator("make_plot")
def gen_make_plot(args: Dict, ctx: Ctx) -> str:
    from .figures import PRESETS, style_rc
    panels = args["panels"]
    output = args["output"]
    layout, figsize = args.get("layout"), args.get("figsize")
    sharex, sharey = bool(args.get("sharex")), bool(args.get("sharey"))
    preset, palette, font, rc = args.get("preset"), args.get("palette"), args.get("font"), args.get("rc")
    style = args.get("style")
    n = len(panels)
    positioned = any(("row" in p or "col" in p) for p in panels)
    rows, cols = (layout or ([n, 1] if not positioned else [
        max(p.get("row", 0) + p.get("rowspan", 1) for p in panels),
        max(p.get("col", 0) + p.get("colspan", 1) for p in panels)]))[:2]
    rcp = style_rc(preset, palette, font, rc)
    size = figsize or rcp.pop("figure.figsize", None) or (6.4 * cols, 4.4 * rows)
    rcp.pop("figure.figsize", None)
    if preset and not figsize and rows > 1:
        size = (size[0], size[1] * (0.55 * rows + 0.45))
    ctx.need("import numpy as np", "import matplotlib.pyplot as plt", "from matplotlib.gridspec import GridSpec")
    head: List[str] = []
    desc = [f"preset {preset!r}" if preset else "", f"palette {palette!r}" if palette and isinstance(palette, str)
            else ("custom palette" if palette else ""), "font" if font else "", "rc" if rc else ""]
    desc = [d for d in desc if d]
    if rcp:
        if "axes.prop_cycle" in rcp:
            ctx.need("from cycler import cycler")
        head.append(f"# rcParams from the {', '.join(desc)} options")
        head.append(f"rc = {_rc_code(rcp)}")
        ctx_line = f"with plt.style.context({(style or 'default')!r}), plt.rc_context(rc):"
    else:
        ctx_line = f"with plt.style.context({(style or 'default')!r}):   # the toolkit draws on matplotlib's " \
                   f"{'defaults' if not style else 'style ' + repr(style)}"
    fm = FigModel()
    w = OpWriter(ctx, fm, "make_plot", indent=4)
    w.emit(f"fig = plt.figure(figsize={lit(tuple(round(float(v), 6) for v in size))})")
    gkw = {}
    if args.get("width_ratios"):
        gkw["width_ratios"] = args["width_ratios"]
    if args.get("height_ratios"):
        gkw["height_ratios"] = args["height_ratios"]
    w.emit(f"gs = GridSpec({rows}, {cols}, figure=fig{', ' + kwargs(gkw) if gkw else ''})")
    models: List[AxModel] = []
    for i, p in enumerate(panels):
        if positioned:
            r, c = p.get("row", 0), p.get("col", 0)
            rs, cs = p.get("rowspan", 1), p.get("colspan", 1)
            cell = f"gs[{r if rs == 1 else f'{r}:{r + rs}'}, {c if cs == 1 else f'{c}:{c + cs}'}]"
            spec = (r, r + rs, c, c + cs)
        else:
            cell = f"gs[{i // cols}, {i % cols}]"
            spec = (i // cols, i // cols + 1, i % cols, i % cols + 1)
        proj = p.get("projection")
        proj = None if proj in (None, "rectilinear") else proj
        if proj == "triangular":
            ctx.need("import pycalphad.plot.triangular  # registers the 'triangular' projection")
        share = {}
        if models and proj is None:
            if sharex:
                share["sharex"] = Raw("axs[0]")
            if sharey:
                share["sharey"] = Raw("axs[0]")
        kw = {"projection": proj} if proj else {}
        kw.update(share)
        add = f"fig.add_subplot({cell}{', ' + kwargs(kw) if kw else ''})"
        w.emit(f"axs = [{add}]" if i == 0 else f"axs.append({add})")
        am = AxModel(f"axs[{i}]", proj or "rectilinear", spec=spec)
        models.append(am)
        fm.top.append(am)
    fs = rcp.get("legend.fontsize", 9)

    def body():
        for i, panel in enumerate(panels):
            ops = [dict(o, units="data") for o in panel.get("operations", [])]
            w.emit("")
            w.emit(f"# ---- panel {i}" + (f" ({panel['projection']})" if panel.get("projection") not in
                                          (None, "rectilinear") else "") + " ----")
            w.touched = [models[i]] if models[i] in fm.main() else []
            w.cur_ax = None
            w.apply(ops, models[i] if models[i] in fm.main() else None)
            w.refresh_legends(w.touched, fallback=models[i], fallback_fontsize=fs)
    _fallback_on(ctx, body)
    w.emit("")
    if args.get("suptitle"):
        w.emit(f"fig.suptitle({args['suptitle']!r})")
    if any(p.get("projection") not in (None, "rectilinear") for p in panels):
        w.emit("try:")
        w.emit("    fig.tight_layout()")
        w.emit("except Exception:   # 3-D / triangular axes may not support it fully")
        w.emit("    pass")
    else:
        w.emit("fig.tight_layout()")
    out = ctx.tb.outp(output if output.lower().endswith((".png", ".jpg")) else output + ".png", "plot.png")
    _emit_save(w, ctx, out, args.get("formats"), {"dpi": args.get("dpi") or 150, "bbox_inches": "tight",
                                                  **({"transparent": True} if args.get("transparent") else {})})
    main = fm.main()
    tail = [f"ax = {main[0].expr}   # main axes" if main else ""]
    return _assemble(ctx, w, head + [ctx_line] + w.lines + tail)


def _emit_save(w: OpWriter, ctx: Ctx, out: str, formats, opts: Dict) -> None:
    name = ctx.out_name(out, Path(out).name)
    w.emit(f"fig.savefig(OUT / {name!r}, {kwargs(opts)})")
    for ext in formats or []:
        ext = ext.lower().lstrip(".")
        p = str(Path(out).with_suffix("." + ext))
        if p != out:
            line = f"fig.savefig(OUT / {ctx.out_name(p, Path(p).name)!r}, {kwargs(opts)})"
            if ext in ("pdf", "svg", "eps"):     # text stays text (editable in Inkscape / Illustrator)
                line = ("with plt.rc_context({'svg.fonttype': 'none', 'pdf.fonttype': 42, 'ps.fonttype': 42}):\n"
                        "    " + line)
            w.emit(line)


def _would_refresh(ax) -> bool:
    shown = [t.get_text() for t in ax.get_legend().get_texts()] if ax.get_legend() is not None else []
    return any(lab not in shown for lab in ax.get_legend_handles_labels()[1])


@generator("edit_figure")
def gen_edit_figure(args: Dict, ctx: Ctx) -> str:
    from .figures import load_figure, main_axes
    tb = ctx.tb
    image = args["image"]
    ops = args["operations"]
    p = tb.p(image)
    full = str(Path(p).resolve())
    fig_in = load_figure(p)
    ctx.need("import numpy as np", "import matplotlib.pyplot as plt")
    raster = fig_in is None
    if not raster and full not in ctx.produced:
        ctx.fallback("the figure being edited was not rebuilt by an earlier standalone cell "
                     "(its step is not included or not standalone), so there is no `fig` to continue")
    default = Path(p).stem + "_edited.png"
    out = tb.outp(args.get("output"), default) if args.get("output") else str(Path(p).with_name(default))
    if not out.lower().endswith((".png", ".jpg", ".jpeg")):
        out = str(Path(out).with_suffix(".png"))
    head: List[str] = []
    if raster:
        cal = args.get("calibration")
        fm = FigModel()
        am = AxModel("ax", xscale=(cal or {}).get("xscale", "linear"), yscale=(cal or {}).get("yscale", "linear"))
        fm.top.append(am)
        w = OpWriter(ctx, fm, "edit")
        w.emit(f"# {Path(p).name} has no editable figure: draw on top of the image"
               + (" in calibrated data coordinates" if cal else " (coordinates are 0-1 fractions of the image)"))
        w.emit(f"img = plt.imread({ctx.path(image)})")
        w.emit("h, w = img.shape[:2]")
        w.emit("fig = plt.figure(figsize=(w / 100, h / 100), dpi=100)")
        w.emit("bg = fig.add_axes([0, 0, 1, 1])")
        w.emit("bg.imshow(img)")
        w.emit("bg.set_axis_off()")
        w.emit("bg.set_label('background image')")
        if cal:
            l, t, r, b = [float(v) for v in cal["pixel_box"]]
            w.emit(f"ax = fig.add_axes([{l!r} / w, 1 - {b!r} / h, ({r!r} - {l!r}) / w, ({b!r} - {t!r}) / h])"
                   "   # the plot frame (pixel box)")
            w.emit(f"ax.set_xlim({lit(cal['xlim'][0])}, {lit(cal['xlim'][1])})")
            w.emit(f"ax.set_ylim({lit(cal['ylim'][0])}, {lit(cal['ylim'][1])})")
            for k in ("xscale", "yscale"):
                if cal.get(k):
                    w.emit(f"ax.set_{k}({cal[k]!r})")
        else:
            w.emit("ax = fig.add_axes([0, 0, 1, 1])")
            w.emit("ax.set_xlim(0, 1)")
            w.emit("ax.set_ylim(0, 1)")
        w.emit("ax.patch.set_alpha(0)")
        w.emit("ax.set_axis_off()")
        w.emit("ax.set_label('overlay')")
        w.cur_ax = "ax"
    else:
        fm = FigModel.from_figure(fig_in)
        for a, real in zip(fm.main(), main_axes(fig_in)):
            a.real = real
        w = OpWriter(ctx, fm, "edit", real_fig=fig_in)
        simple = not any(str(o.get("op", "")).lower() in ("twinx", "twiny", "inset", "broken_yaxis", "scatter_matrix")
                         or o.get("axes", 0) not in (0, "0") for o in ops)
        if len(fm.main()) == 1 and simple:
            fm.main()[0].expr = "ax"
            w.emit("# continue on the figure of the previous cell; values are given as shown on its axes")
            w.cur_ax = "ax"
        else:
            w.helpers("panel_axes")
            w.emit("# continue on the figure of the previous cell; values are given as shown on its axes")
            w.emit("axes = panel_axes(fig)   # the axes numbered as describe_figure lists them")
    main0 = fm.main()[0] if fm.main() else None

    def body():
        w.apply(ops, main0)
        refresh = []
        for a in fm.main():
            if a in w.touched or getattr(a, "real", None) is None or _would_refresh(a.real):
                refresh.append(a)
        w.refresh_legends(refresh)
    _fallback_on(ctx, body)
    if fig_in is not None:
        import matplotlib.pyplot as _plt
        _plt.close(fig_in)
    opts = {"dpi": Raw("fig.dpi")} if raster else {"dpi": args.get("dpi") or 150, "bbox_inches": "tight"}
    if args.get("transparent"):
        opts["transparent"] = True
    from .figures import LABEL_OPS
    formats = args.get("formats")
    if formats is None and any(str(o.get("op", "")).lower() in LABEL_OPS for o in ops):
        formats = ["svg", "pdf"]
    _emit_save(w, ctx, out, formats, opts)
    main = fm.main()
    if main and w.cur_ax != main[0].expr:
        w.emit(f"ax = {main[0].expr}   # main axes")
    return _assemble(ctx, w, head + w.lines)


# =========================================================================================================
# data tools: transform_data / fit_data
# =========================================================================================================
@generator("transform_data")
def gen_transform_data(args: Dict, ctx: Ctx) -> str:
    from .figdata import write_table
    tb = ctx.tb
    source = args["source"]
    lines: List[str] = []
    pre: List[str] = []
    dw = DataWriter(ctx, lines.append, set(), pre)
    try:
        var, df = dw.load(source, "df", sheet=args.get("sheet"), where=args.get("where"),
                          transform=args.get("transform") or [])
    except (ValueError, KeyError) as e:
        ctx.fallback(f"data pipeline: {e}")
    stem = Path(source).stem if isinstance(source, str) and "\n" not in source else "table"
    out = tb.outp(args.get("output"), f"{stem}_transformed.csv")
    ext = Path(out).suffix.lower()
    if ext not in (".csv", ".tsv", ".txt", ".xlsx", ".json"):
        out = str(Path(out).with_suffix(".csv")) if not ext else out
        ext = Path(out).suffix.lower()
    name = ctx.out_name(out, Path(out).name)
    wr = {".csv": "df.to_csv(OUT / {n!r}, index=False)", ".txt": "df.to_csv(OUT / {n!r}, index=False)",
          ".tsv": "df.to_csv(OUT / {n!r}, index=False, sep='\\t')", ".xlsx": "df.to_excel(OUT / {n!r}, index=False)",
          ".json": "df.to_json(OUT / {n!r}, orient='records', indent=1)"}[ext]
    lines.append(wr.format(n=name))
    lines.append("print(df.head(10))")
    return "".join(pre) + "\n".join(lines) + "\n"


@generator("fit_data")
def gen_fit_data(args: Dict, ctx: Ctx) -> str:
    tb = ctx.tb
    op = {k: v for k, v in args.items() if v is not None and k not in ("output_csv", "npoints", "level", "band", "range")}
    fm = FigModel()
    am = AxModel("ax")
    fm.top.append(am)
    w = OpWriter(ctx, fm, "edit")
    ctx.need("import pandas as pd")

    def body():
        o = OpX(w, op, am, 1)
        info = w._fit(o, need_stats=True)
        g, back = info["g"], info["back"]
        bk = (lambda s: f"{back}({s})") if back else (lambda s: s)
        rng = args.get("range")
        lo, hi = (lit(float(rng[0])), lit(float(rng[1]))) if rng else ("x.min()", "x.max()")
        w.emit(f"xg = np.linspace({lo}, {hi}, {int(args.get('npoints', 200))})")
        w.emit(f"yg = {g}(xg, *p)")
        band = args.get("band", "confidence")
        level = float(args.get("level", 0.95))
        w.emit(f"curve = {{'x': xg, 'y_fit': {bk('yg')}}}")
        ctx.need("from scipy import stats")
        w.helpers("jacobian")
        w.emit(f"if np.all(np.isfinite(pcov)):   # {int(level * 100)}% {band} band")
        w.emit(f"    var = np.einsum('ij,jk,ik->i', jacobian({g}, xg, p), pcov, jacobian({g}, xg, p))")
        if band == "prediction":
            if info["sig_t"] is None:
                w.emit("    var = var + np.sum(resid ** 2) / dof")
            else:
                w.emit("    o_ = np.argsort(x)")
                w.emit(f"    s_new = np.exp(np.interp(xg, x[o_], np.log({info['sig_t']}[o_])))")
                f = "1.0" if op.get("absolute_sigma") else "(np.sum(w * resid ** 2) / dof)"
                w.emit(f"    var = var + {f} * s_new ** 2")
        w.emit(f"    half = stats.t.ppf({0.5 + level / 2.0!r}, dof) * np.sqrt(np.clip(var, 0, None))")
        w.emit(f"    curve[{band + '_lo'!r}], curve[{band + '_hi'!r}] = {bk('yg - half')}, {bk('yg + half')}")
        out = tb.outp(args.get("output_csv"), f"fit_{info['name']}.csv")
        if not out.lower().endswith(".csv"):
            out = str(Path(out).with_suffix(".csv"))
        rout = str(Path(out).with_name(Path(out).stem + "_residuals.csv"))
        w.emit(f"pd.DataFrame(curve).to_csv(OUT / {ctx.out_name(out, Path(out).name)!r}, index=False)")
        rname = "residual" + ("" if info["space"] == "y" else f" ({info['space']})")
        w.emit(f"pd.DataFrame({{'x': x, 'y': y, 'y_fit': {bk(f'{g}(x, *p)')}, {rname!r}: resid}})"
               f".to_csv(OUT / {ctx.out_name(rout, Path(rout).name)!r}, index=False)")
        w.emit("print('parameters:', p, '± SE', np.sqrt(np.diag(pcov)), 'R² =', r2)")
    _fallback_on(ctx, body)
    return "".join(w.pre) + "\n".join(w.lines) + "\n"


# =========================================================================================================
# composition / raster tools (calphad_tools.figcompose)
# =========================================================================================================
@generator("compose_images")
def gen_compose_images(args: Dict, ctx: Ctx) -> str:
    from . import figcompose as fc
    from PIL import Image
    tb = ctx.tb
    images = args["images"] if isinstance(args["images"], list) else [args["images"]]
    paths = [tb.p(x) for x in images]
    output = args["output"]
    formats = list(args.get("formats") or [])
    ext = Path(output).suffix.lower()
    if ext in (".pdf", ".svg", ".eps", ".tiff"):
        formats.append(ext[1:])
        output = str(Path(output).with_suffix(".png"))
    elif ext not in (".png", ".jpg", ".jpeg"):
        output += ".png"
    out = tb.outp(output, "composite.png")
    labels, titles = args.get("labels", "(a)"), args.get("titles")
    trim_ws = args.get("trim_whitespace", True)
    spacing, lpos = float(args.get("spacing", 0.15)), args.get("label_position", "outside")
    fs = float(args.get("label_fontsize", 12))
    dpi = int(args.get("dpi") or 200)
    suptitle = args.get("suptitle")
    # geometry exactly as the toolkit computed it from the (trimmed) images
    ims = [fc.open_image(q) for q in paths]
    if trim_ws:
        ims = [fc.trim(im) for im in ims]
    n = len(ims)
    aspects = [im.size[1] / im.size[0] for im in ims]
    if args.get("layout"):
        rows, cols = int(args["layout"][0]), int(args["layout"][1])
    else:
        rows, cols = fc.auto_layout(aspects)
    labs = fc.panel_labels(labels, n)
    tl = [str(t) if t else "" for t in (list(titles or []) + [""] * n)[:n]]
    width_in = float(args.get("width_in") or min(max(3.5 * cols, 3.5), 16.0))
    margin, gap = 0.06, max(spacing, 0.0)
    avail = width_in - 2 * margin - (cols - 1) * gap
    boxes, row_h = fc.grid_geometry(aspects, rows, cols, avail, gap)
    strip = (fs * 1.6 / 72) if (lpos == "outside" and any(labs)) or any(tl) else 0.0
    sup_h = fs * 2.2 / 72 if suptitle else 0.0
    height_in = margin * 2 + sup_h + sum(row_h) + rows * strip + (rows - 1) * gap
    ctx.need("import numpy as np", "import matplotlib.pyplot as plt", "from PIL import Image")
    pre = helper_code(ctx, ["open_image"] + (["trim_box"] if trim_ws else []))
    L = [f"# {n} images in a {rows}x{cols} grid (aspect ratios kept"
         + (", white margins trimmed" if trim_ws else "") + "); geometry in inches as the toolkit laid it out",
         f"fig = plt.figure(figsize=({width_in!r}, {height_in!r}), dpi={dpi}, facecolor='white')"]
    y = height_in - margin - sup_h
    first = True
    for r in range(rows):
        idx = list(range(r * cols, min(r * cols + cols, n)))
        y_img_top = y - strip
        for i in idx:
            w_in, h_in = boxes[i][1], row_h[r]
            x0 = margin + boxes[i][0]
            y0 = y_img_top - h_in
            px = (max(int(round(w_in * dpi)), 1), max(int(round(h_in * dpi)), 1))
            rect = [x0 / width_in, y0 / height_in, w_in / width_in, h_in / height_in]
            L.append("")
            L.append(f"# panel {labs[i] or i + 1}: {Path(paths[i]).name}")
            L.append(f"im = open_image({ctx.path(images[i])})")
            if trim_ws:
                L.append("box = trim_box(im)")
                L.append("im = im.crop(box) if box else im")
            L.append(f"if im.size[0] > {px[0]} * 1.05:   # keep the figure small")
            L.append(f"    im = im.resize({px}, Image.LANCZOS)")
            L.append(f"{'ax' if first else 'a'} = fig.add_axes({lit(rect)})")
            v = "ax" if first else "a"
            L.append(f"{v}.imshow(np.asarray(im), interpolation='antialiased', aspect='auto')")
            L.append(f"{v}.set_axis_off()")
            L.append(f"{v}.set_label({f'panel {labs[i] or i + 1}: {Path(paths[i]).name}'!r})")
            if tl[i]:
                L.append(f"fig.text({(x0 + w_in / 2) / width_in!r}, {(y_img_top + strip / 2) / height_in!r}, "
                         f"{tl[i]!r}, ha='center', va='center', fontsize={fs * 0.9!r})")
            if labs[i]:
                if lpos == "outside":
                    L.append(f"fig.text({x0 / width_in!r}, {(y_img_top + strip / 2) / height_in!r}, {labs[i]!r}, "
                             f"ha='left', va='center', fontsize={fs!r}, fontweight='bold')")
                else:
                    vert, horiz = lpos.split("-")
                    L.append(f"{v}.text({0.02 if horiz == 'left' else 0.98}, {0.98 if vert == 'top' else 0.02}, "
                             f"{labs[i]!r}, transform={v}.transAxes, ha={horiz!r}, va={vert!r}, fontsize={fs!r}, "
                             "fontweight='bold',")
                    L.append("        bbox=dict(boxstyle='round,pad=0.2', fc='white', ec='none', alpha=0.75))")
            first = False
        y -= strip + row_h[r] + gap
    if suptitle:
        L.append(f"fig.text(0.5, {(height_in - margin - sup_h / 2) / height_in!r}, {suptitle!r}, ha='center', "
                 f"va='center', fontsize={fs * 1.15!r}, fontweight='bold')")
    L.append("")
    name = ctx.out_name(out, Path(out).name)
    L.append(f"fig.savefig(OUT / {name!r}, dpi={dpi}, facecolor='white')")
    for e in formats:
        e = e.lower().lstrip(".")
        q = str(Path(out).with_suffix("." + e))
        if q != out:
            L.append(f"fig.savefig(OUT / {ctx.out_name(q, Path(q).name)!r}, dpi={dpi}, facecolor='white')")
    return pre + "\n".join(L) + "\n"


def _image_ops_code(ops: Sequence[Dict], L: List[str]) -> Tuple[Optional[str], Dict]:
    from .figcompose import parse_color
    fmt, save = None, {}
    for op in ops:
        name = str(op.get("op", "")).lower()
        if name == "crop":
            if op.get("trim"):
                L.append(f"box = trim_box(im, {float(op.get('tolerance', 12))!r}, {int(op.get('pad', 4))})"
                         "   # trim the uniform margin")
                L.append("im = im.crop(box) if box else im")
            elif op.get("units") == "fraction":
                b = [float(v) for v in op["box"]]
                L.append("w, h = im.size")
                L.append(f"im = im.crop(tuple(int(round(v)) for v in ({b[0]!r} * w, {b[1]!r} * h, {b[2]!r} * w, "
                         f"{b[3]!r} * h)))")
            else:
                L.append(f"im = im.crop({tuple(int(round(float(v))) for v in op['box'])})")
        elif name == "resize":
            L.append("w, h = im.size")
            if op.get("scale"):
                L.append(f"nw, nh = w * {float(op['scale'])!r}, h * {float(op['scale'])!r}")
            elif op.get("width") and op.get("height"):
                L.append(f"nw, nh = {float(op['width'])!r}, {float(op['height'])!r}")
            elif op.get("width"):
                L.append(f"nw = {float(op['width'])!r}")
                L.append("nh = h * nw / w")
            else:
                L.append(f"nh = {float(op['height'])!r}")
                L.append("nw = w * nh / h")
            L.append("im = im.resize((max(int(round(nw)), 1), max(int(round(nh)), 1)), Image.LANCZOS)")
        elif name in ("pad", "border"):
            px = op.get("px", 20 if name == "pad" else 2)
            px = [int(px)] * 4 if not isinstance(px, (list, tuple)) else [int(v) for v in px]
            color = parse_color(op.get("color"), (255, 255, 255) if name == "pad" else (0, 0, 0))
            L.append(f"im = ImageOps.expand(im, border={tuple(px)}, fill={color} + ((255,) if im.mode == 'RGBA' else ()))")
        elif name == "transparent_background":
            color = parse_color(op.get("color", "white"))
            L.append("arr = np.asarray(im.convert('RGBA')).copy()")
            L.append(f"arr[np.abs(arr[:, :, :3].astype(int) - np.array({color})).max(axis=2) <= "
                     f"{float(op.get('tolerance', 10))!r}, 3] = 0   # background colour -> transparent")
            L.append("im = Image.fromarray(arr, 'RGBA')")
        elif name == "grayscale":
            L.append("im = im.convert('LA' if im.mode == 'RGBA' else 'L')")
        elif name == "rotate":
            L.append("fill = {'RGBA': (255, 255, 255, 0), 'L': 255, 'LA': (255, 0)}.get(im.mode, 'white')")
            L.append(f"im = im.rotate({float(op.get('degrees', 90))!r}, resample=Image.BICUBIC, "
                     f"expand={bool(op.get('expand', True))}, fillcolor=fill)")
        elif name == "convert":
            fmt = str(op.get("format", "png")).lower().lstrip(".")
            fmt = {"jpeg": "jpg", "tif": "tiff"}.get(fmt, fmt)
            if op.get("quality"):
                save["quality"] = int(op["quality"])
        else:
            raise ValueError(f"image op {name!r}")
    return fmt, save


@generator("image_tools")
def gen_image_tools(args: Dict, ctx: Ctx) -> str:
    tb = ctx.tb
    image, ops, output, dpi = args["image"], args.get("operations") or [], args.get("output"), args.get("dpi")
    src = tb.p(image)
    ctx.need("import numpy as np", "from PIL import Image, ImageOps")
    L: List[str] = []
    fmt, save = _image_ops_code(ops, L)
    src_ext = {"jpeg": "jpg", "tif": "tiff"}.get(Path(src).suffix.lower().lstrip("."), Path(src).suffix.lower().lstrip("."))
    if output:
        oext = {"jpeg": "jpg", "tif": "tiff"}.get(Path(output).suffix.lower().lstrip("."),
                                                   Path(output).suffix.lower().lstrip("."))
        if fmt is None:
            fmt = oext if oext in ("png", "jpg", "pdf", "tiff", "webp") else (
                src_ext if src_ext in ("png", "jpg", "tiff", "webp") else "png")
        if oext != fmt:
            output = str(Path(output).with_suffix("." + fmt)) if oext in ("png", "jpg", "pdf", "tiff", "webp") \
                else output + "." + fmt
        out = tb.outp(output, "image.png")
    else:
        fmt = fmt or (src_ext if src_ext in ("png", "jpg", "tiff", "webp") else "png")
        out = tb.outp(None, f"{Path(src).stem}_edited.{fmt}")
    pre = helper_code(ctx, ["open_image"] + (["trim_box"] if any(o.get("trim") for o in ops) else []))
    head = [f"im = open_image({ctx.path(image)})"]
    if fmt in ("jpg", "pdf"):
        L.append("if im.mode in ('RGBA', 'LA', 'P'):   # no alpha channel in this format")
        L.append("    im = flatten(im.convert('RGBA'))")
    kw = dict(save)
    if fmt == "jpg":
        kw.setdefault("quality", 92)
    if fmt == "pdf":
        kw["resolution"] = float(dpi or 150)
    if dpi and fmt in ("png", "jpg", "tiff"):
        kw["dpi"] = (dpi, dpi)
    pil = {"jpg": "JPEG", "png": "PNG", "pdf": "PDF", "tiff": "TIFF", "webp": "WEBP"}[fmt]
    L.append(f"im.save(OUT / {ctx.out_name(out, Path(out).name)!r}, {pil!r}{', ' + kwargs(kw) if kw else ''})")
    return pre + "\n".join(head + L) + "\n"


@generator("images_to_pdf")
def gen_images_to_pdf(args: Dict, ctx: Ctx) -> str:
    from .figures import load_figure
    tb = ctx.tb
    images = args["images"] if isinstance(args["images"], list) else [args["images"]]
    output, page_size = args["output"], args.get("page_size", "fit")
    titles = [str(t) if t else "" for t in (list(args.get("titles") or []) + [""] * len(images))[:len(images)]]
    out = tb.outp(output if output.lower().endswith(".pdf") else output + ".pdf", "images.pdf")
    ctx.need("import numpy as np", "import matplotlib.pyplot as plt", "from matplotlib.backends.backend_pdf import PdfPages")
    pre = helper_code(ctx, ["open_image"])
    vector = [bool(args.get("vector", True) and page_size == "fit" and x.lower().endswith(".png")
                   and load_figure(tb.p(x)) is not None) for x in images]
    L = []
    if any(vector):
        L.append("# NOTE: the toolkit put tool-made figures on vector pages; here every page shows the PNG image")
    L.append(f"pages = {lit([(Raw(ctx.path(x)), t) for x, t in zip(images, titles)], 0)}")
    L.append(f"with PdfPages(OUT / {ctx.out_name(out, Path(out).name)!r}) as pdf:")
    L.append("    for path, title in pages:")
    L.append("        arr = np.asarray(flatten(open_image(path)))")
    L.append("        h, w = arr.shape[:2]")
    L.append("        th = 0.45 if title else 0.0   # title strip, inches")
    if page_size == "fit":
        L.append("        pw, ph = w / 150, h / 150 + th   # page = image at 150 dpi")
        L.append("        box = [0, 0, 1, (h / 150) / ph]")
        ty = "1 - (th / 2) / ph"
    else:
        from .figcompose import PAGE_SIZES
        pw, ph = PAGE_SIZES[page_size]
        L.append(f"        pw, ph = ({pw!r}, {ph!r}) if w <= h else ({ph!r}, {pw!r})   # {page_size}, landscape for wide images")
        L.append("        bw, bh = pw - 1.0, ph - 1.0 - th")
        L.append("        sc = min(bw / w, bh / h)")
        L.append("        iw, ih = w * sc, h * sc")
        L.append("        box = [(pw - iw) / 2 / pw, (0.5 + (bh - ih) / 2) / ph, iw / pw, ih / ph]")
        ty = "1 - (0.5 + th / 2) / ph"
    L.append("        page = plt.figure(figsize=(pw, ph), dpi=150)")
    L.append("        a = page.add_axes(box)")
    L.append("        a.imshow(arr, interpolation='none', aspect='auto')")
    L.append("        a.set_axis_off()")
    L.append("        if title:")
    L.append(f"            page.text(0.5, {ty}, title, ha='center', va='center', fontsize=13, fontweight='bold')")
    L.append("        pdf.savefig(page)")
    L.append("        plt.close(page)")
    return pre + "\n".join(L) + "\n"


def _slide_placements(spec: Dict, sw: int, sh: int) -> Tuple[List[Tuple[int, int, int, int]], Dict]:
    """(x, y, w, h) EMU per image exactly as figcompose.add_images_slide places them, + text boxes."""
    from pptx.util import Inches
    from PIL import Image
    from .figcompose import best_grid
    title, subtitle = spec.get("title"), spec.get("subtitle")
    top = Inches(0.35)
    if title:
        top = Inches(1.55) if subtitle else Inches(1.3)
    caption = spec.get("caption")
    bottom = sh - Inches(0.25) - (Inches(0.55) if caption else 0)
    left, right = Inches(0.4), sw - Inches(0.4)
    boxes = {}
    if spec.get("bullets"):
        text_w = int((right - left) * (0.36 if spec.get("images") else 1.0))
        boxes["bullets"] = (int(left), int(top), text_w, int(bottom - top))
        left += text_w + Inches(0.2)
    paths = spec.get("images") or []
    sizes = []
    for p in paths:
        with Image.open(p) as im:
            sizes.append(im.size)
    out = []
    if paths:
        gap = Inches(0.15)
        bw, bh = right - left, bottom - top
        if spec.get("layout"):
            rows, cols = int(spec["layout"][0]), int(spec["layout"][1])
        else:
            rows, cols = best_grid(sizes, bw, bh, gap)
        cw, ch = (bw - (cols - 1) * gap) / cols, (bh - (rows - 1) * gap) / rows
        for i, (w, h) in enumerate(sizes):
            r, c = divmod(i, cols)
            in_row = min(cols, len(paths) - r * cols)
            x_off = (cols - in_row) * (cw + gap) / 2
            sc = min(cw / w, ch / h)
            pw, ph = int(w * sc), int(h * sc)
            out.append((int(left + x_off + c * (cw + gap) + (cw - pw) / 2), int(top + r * (ch + gap) + (ch - ph) / 2),
                        pw, ph))
    return out, boxes


@generator("figures_to_pptx")
def gen_figures_to_pptx(args: Dict, ctx: Ctx) -> str:
    from pptx.util import Inches
    tb = ctx.tb
    slides, output, append_to, title_slide = args["slides"], args.get("output"), args.get("append_to"), \
        args.get("title_slide")
    ctx.need("from pptx import Presentation", "from pptx.util import Inches, Pt")
    L = ["# NOTE: slide styling (fonts, accent bar, colours) is simplified; the image placement is the toolkit's",
         "def add_text(slide, left, top, width, height, text, size=18, bold=False):",
         "    box = slide.shapes.add_textbox(left, top, width, height)",
         "    box.text_frame.word_wrap = True",
         "    for i, line in enumerate(str(text).split('\\n')):",
         "        par = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()",
         "        run = par.add_run()",
         "        run.text = line",
         "        run.font.size, run.font.bold = Pt(size), bold",
         ""]
    if append_to:
        L.append(f"prs = Presentation({ctx.path(append_to)})")
        from pptx import Presentation
        prs = Presentation(tb.p(append_to))
        sw, sh = int(prs.slide_width), int(prs.slide_height)
    else:
        L.append("prs = Presentation()")
        L.append("prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # 16:9")
        sw, sh = int(Inches(13.333)), int(Inches(7.5))
    L.append("layouts = list(prs.slide_layouts)")
    L.append("blank = layouts[6] if len(layouts) > 6 and not len(layouts[6].placeholders) else "
             "min(layouts, key=lambda l: len(l.placeholders))")
    if title_slide:
        L.append("s = prs.slides.add_slide(blank)")
        L.append(f"add_text(s, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.2), {str(title_slide.get('title', ''))!r}, "
                 "size=36, bold=True)")
        if title_slide.get("subtitle"):
            L.append(f"add_text(s, Inches(0.8), Inches(3.9), Inches(11.7), Inches(0.8), "
                     f"{str(title_slide['subtitle'])!r}, size=18)")
    for sp in slides:
        spec = dict(sp, images=[tb.p(p) for p in sp.get("images") or []])
        places, boxes = _slide_placements(spec, sw, sh)
        L.append("")
        L.append(f"# slide: {sp.get('title') or '(untitled)'}")
        L.append("s = prs.slides.add_slide(blank)")
        L.append("for ph in list(s.placeholders):")
        L.append("    ph._element.getparent().remove(ph._element)")
        if sp.get("title"):
            L.append(f"add_text(s, Inches(0.5), Inches(0.3), prs.slide_width - Inches(1.0), Inches(0.7), "
                     f"{str(sp['title'])!r}, size=26, bold=True)")
            if sp.get("subtitle"):
                L.append(f"add_text(s, Inches(0.5), Inches(1.08), prs.slide_width - Inches(1.0), Inches(0.4), "
                         f"{str(sp['subtitle'])!r}, size=13)")
        if "bullets" in boxes:
            bx = boxes["bullets"]
            L.append(f"add_text(s, {bx[0]}, {bx[1]}, {bx[2]}, {bx[3]}, "
                     f"{chr(10).join('•  ' + str(b) for b in sp['bullets'])!r}, size={int(sp.get('bullet_size', 16))})")
        for p, (x, y, w, h) in zip(sp.get("images") or [], places):
            L.append(f"s.shapes.add_picture(str({ctx.path(p)}), {x}, {y}, {w}, {h})")
        if sp.get("caption"):
            L.append(f"add_text(s, Inches(0.5), prs.slide_height - Inches(0.75), prs.slide_width - Inches(1.0), "
                     f"Inches(0.5), {str(sp['caption'])!r}, size=13)")
        if sp.get("notes"):
            L.append(f"s.notes_slide.notes_text_frame.text = {str(sp['notes'])!r}")
    if not output:
        out = tb.p(append_to)
    else:
        out = tb.outp(output if output.lower().endswith(".pptx") else output + ".pptx", "figures.pptx")
    L.append("")
    L.append(f"prs.save(OUT / {ctx.out_name(out, Path(out).name)!r})")
    return "\n".join(L) + "\n"


@generator("animate")
def gen_animate(args: Dict, ctx: Ctx) -> str:
    tb = ctx.tb
    if args.get("sweep"):
        ctx.fallback("a sweep calls another toolkit tool once per frame; it has no standalone equivalent")
    output = args["output"]
    ext = Path(output).suffix.lower()
    if ext not in (".gif", ".mp4"):
        output += ".gif"
        ext = ".gif"
    if ext == ".mp4":
        ctx.fallback("MP4 frames are piped to the local ffmpeg binary by the toolkit")
    frames = args["frames"]
    out = tb.outp(output, "animation.gif")
    n = len(frames)
    if args.get("captions"):
        caps = [str(c) for c in (list(args["captions"]) + [""] * n)[:n]]
    elif args.get("caption_template"):
        caps = [args["caption_template"].format(value=Path(tb.p(p)).stem, i=k + 1, n=n) for k, p in enumerate(frames)]
    else:
        caps = None
    dur = int(args["duration_ms"]) if args.get("duration_ms") else int(round(1000.0 / float(args.get("fps") or 2)))
    max_w = int(args.get("max_width", 1200))
    ctx.need("import numpy as np", "import matplotlib", "from PIL import Image, ImageDraw, ImageFont")
    pre = helper_code(ctx, ["open_image"])
    L = [f"frames = [flatten(open_image(p)) for p in {lit([Raw(ctx.path(p)) for p in frames])}]",
         f"frames = [im.resize(({max_w}, max(int(round(im.size[1] * {max_w} / im.size[0])), 1)), Image.LANCZOS)",
         f"          if im.size[0] > {max_w} else im for im in frames]",
         "W, H = max(im.size[0] for im in frames), max(im.size[1] for im in frames)   # common canvas"]
    if caps and any(caps):
        L += [f"captions = {lit(caps)}",
              "fsz = max(14, min(40, W // 32))",
              "strip = int(fsz * 1.8)   # caption strip above each frame",
              "font = ImageFont.truetype(str(Path(matplotlib.get_data_path()) / 'fonts' / 'ttf' / 'DejaVuSans.ttf'), fsz)"]
    else:
        L += ["captions, strip = [''] * len(frames), 0"]
    L += ["canvases = []",
          "for im, cap in zip(frames, captions):",
          "    canvas = Image.new('RGB', (W, H + strip), 'white')",
          "    canvas.paste(im, ((W - im.size[0]) // 2, strip + (H - im.size[1]) // 2))",
          "    if cap:",
          "        ImageDraw.Draw(canvas).text((W // 2, strip // 2), cap, fill='black', font=font, anchor='mm')",
          "    canvases.append(canvas)",
          f"canvases[0].save(OUT / {ctx.out_name(out, Path(out).name)!r}, save_all=True, append_images=canvases[1:], "
          f"duration={dur}, loop={int(args.get('loop', 0))}, optimize=False, disposal=1)"]
    return pre + "\n".join(L) + "\n"


# =========================================================================================================
# digitizing: the image analysis stays in the toolkit, the extracted numbers are recorded
# =========================================================================================================
@generator("digitize_setup")
def gen_digitize_setup(args: Dict, ctx: Ctx) -> str:
    ctx.fallback("frame / tick detection is the toolkit's image-processing pipeline (no data produced)")


def _digitized_csv(ctx: Ctx, tool: str, args: Dict) -> Optional[str]:
    tb = ctx.tb
    rec = ctx.record if ctx.record.get("tool") == tool else {}
    m = re.search(r"->\s*(\S+\.csv)", rec.get("result") or "")
    if m and Path(tb.p(m.group(1))).exists():
        return str(Path(tb.p(m.group(1))))
    key = json.dumps(args, sort_keys=True, default=str)
    for r in reversed(tb.call_log):
        if r.get("tool") == tool and r.get("ok") and json.dumps(r.get("args"), sort_keys=True, default=str) == key:
            m = re.search(r"->\s*(\S+\.csv)", r.get("result_full") or r.get("result_preview") or "")
            if m:
                p = Path(tb.p(m.group(1)))
                if p.exists():
                    return str(p)
    if args.get("output_csv"):
        p = Path(tb.outp(args["output_csv"], "x.csv"))
        if p.exists():
            return str(p)
    return None


def _gen_digitized(tool: str):
    def gen(args: Dict, ctx: Ctx) -> str:
        csv = _digitized_csv(ctx, tool, args)
        if csv is None:
            ctx.fallback(f"{tool} is the toolkit's image analysis and its CSV is no longer available")
        return _embed_digitized(tool, csv, args, ctx)
    return gen


def _gen_digitize_fix(args: Dict, ctx: Ctx) -> str:
    """digitize_fix: the corrected CSV (user-confirmed counts / added / removed points) embedded like digitize_points."""
    p = Path(ctx.tb.p(args["csv"]))
    run = Path(str(p) + ".run.json")
    if not p.exists() or not run.exists():
        ctx.fallback("the corrected CSV (or its digitize_points settings) is no longer available")
    spec = json.loads(run.read_text())
    fix = {k: args[k] for k in ("counts", "add", "remove") if args.get(k)}
    return _embed_digitized("digitize_fix", str(p), {**{k: v for k, v in spec.items() if v is not None}, **fix}, ctx)


def _embed_digitized(tool: str, csv: str, args: Dict, ctx: Ctx) -> str:
        import pandas as pd
        df = pd.read_csv(csv)
        ctx.need("import pandas as pd")
        ctx.note(f"{tool}: the image analysis is not repeated; the digitized values are embedded in the code")
        what = {"digitize_points": "marker detection", "digitize_curve": "curve tracing",
                "digitize_fix": "marker detection corrected with the user's point counts"}[tool]
        shown = {k: v for k, v in args.items() if k not in ("image",)}
        cols = {c: [None if (isinstance(v, float) and math.isnan(v)) else _plain(v) for v in df[c].tolist()]
                for c in df.columns}
        name = ctx.out_name(csv, Path(csv).name)
        code = (f"# {tool}: {what} + axis calibration on {Path(args['image']).name} ran in the toolkit's image\n"
                f"# pipeline (options: {json.dumps(shown, default=str)[:300]}).\n"
                f"# The extracted values are recorded here so this notebook keeps the numbers.\n"
                f"digitized = pd.DataFrame({lit(cols, 0)})\n"
                f"digitized.to_csv(OUT / {name!r}, index=False)\n"
                "print(digitized.head())\n")
        return code + _provenance_code(ctx, args, df, Path(csv).stem)


def _provenance_code(ctx: Ctx, args: Dict, df, stem: str) -> str:
    """Code that redraws the provenance check: the recorded values mapped back onto the figure through the
    calibration (its coefficients are computed here, once, by the toolkit and embedded).  '' if not possible."""
    try:
        from . import digitize as D
        tb = ctx.tb
        work = D.prepare_image(tb.p(args["image"]), tb.outdir(None, "digitize"), args.get("page"), args.get("crop"),
                               args.get("dpi") or 300)
        rgb = D.load_rgb(work)
        tern = (args.get("axes") or {}).get("ternary")
        if tern:
            tcal = D.calibrate_triangle(tern, D.detect_triangle(rgb))
            els = tcal.elements
            if not all(f"X({e})" in df.columns for e in els):
                return ""
            cal = f"_A = np.array({lit(np.round(tcal.A, 12).tolist(), 0)})      # fractions {els} = A @ [px, py, 1]\n"
            xy = (f"_f = _g[{lit([f'X({e})' for e in els], 0)}].to_numpy(float)\n"
                  "    _px, _py = np.linalg.solve(_A[:2, :2], (_f[:, :2] - _A[:2, 2]).T)\n")
        else:
            fr = D.detect_frame(rgb)
            calx, caly = D._cal(args.get("axes") or {}, D.detect_ticks(rgb, fr), fr)
            cal = (f"_cal = {{'x': ({calx.a!r}, {calx.b!r}, {calx.scale!r}), 'y': ({caly.a!r}, {caly.b!r}, "
                   f"{caly.scale!r})}}      # value = a + b * pixel (log10(value) on a log axis)\n"
                   "def _to_px(v, c):\n"
                   "    v = np.asarray(v, float)\n"
                   "    return ((np.log10(v) if c[2] == 'log' else v) - c[0]) / c[1]\n")
            xy = "_px, _py = _to_px(_g['x'], _cal['x']), _to_px(_g['y'], _cal['y'])\n"
    except Exception:  # noqa: BLE001 - the numbers are recorded either way; the check figure is optional
        return ""
    derived = Path(work).stem != Path(str(args["image"])).stem
    img = ctx.path(work if derived else args["image"])
    kind = "'-'" if "x_px" not in df.columns else "'o'"
    ctx.need("import numpy as np", "import matplotlib.pyplot as plt")
    H, W = rgb.shape[:2]
    out = ctx.out_name(None, f"{stem}_provenance.png")
    return ("\n# provenance check: the values above mapped back onto the figure through the calibration\n"
            + (f"# (the toolkit's working copy of the figure: rendered / cropped / deskewed)\n" if derived else "")
            + f"_img = plt.imread({img})\n" + cal
            + f"_pfig, _pax = plt.subplots(figsize=({W / 100:.2f}, {H / 100:.2f}), dpi=100)\n"
            "_pax.imshow(_img)\n"
            "_pax.set_axis_off()\n"
            "for _name, _g in digitized.groupby('series', sort=False):\n"
            f"    {xy}"
            f"    _pax.plot(_px, _py, {kind}, mfc='none', ms=12, mew=1.3, lw=1.2, label=f'{{_name}} ({{len(_g)}})')\n"
            "_pax.legend(fontsize=7, loc='upper right')\n"
            f"_pfig.savefig(OUT / {out!r}, bbox_inches='tight')\n"
            "plt.close(_pfig)\n")


generator("digitize_points")(_gen_digitized("digitize_points"))
generator("digitize_curve")(_gen_digitized("digitize_curve"))
generator("digitize_fix")(_gen_digitize_fix)
