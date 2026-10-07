"""Standalone-code generators (``show_code``) for the calculation tools.

Each generator turns the arguments of one tool call into readable pycalphad / numpy / pandas / scipy / matplotlib
code that repeats the tool's *calculation* (same conditions, phases, grids and derived quantities), writes the same
data files under ``OUT`` and - for figures - leaves ``fig`` / ``ax`` defined with the data plotted in the units shown
on the axes.  Decisions that only depend on the inputs (which phases can form, which phase is the liquid, atomic
masses, phase colours, file names) are made here, at generation time, and emitted as literals so the generated code
stays short.  See the contract in :mod:`calphad_tools.codegen`.
"""
from __future__ import annotations

import os
import re
import textwrap
from pathlib import Path
from typing import Dict, List, Optional

from .codegen import Ctx, generator
from .mapping_check import check_mode

KELVIN = 273.15


# ---------------------------------------------------------------------------------------------------------
# helpers (generation time)
# ---------------------------------------------------------------------------------------------------------
def _T(template: str, **kw) -> str:
    """Fill ``<<name>>`` placeholders of a code template (keeps the template free of brace escaping)."""
    out = textwrap.dedent(template).lstrip("\n")
    for k, val in kw.items():
        out = out.replace(f"<<{k}>>", str(val))
    left = re.findall(r"<<\w+>>", out)
    if left:
        raise KeyError(f"unfilled placeholders {left}")
    return out


def _indent(block: str, n: int) -> str:
    """A code block as a continuation: newline + every line indented by n spaces."""
    return "\n" + textwrap.indent(block.strip("\n"), " " * n)


def _outp(tb, name: Optional[str], default: str) -> str:
    """Path the tool wrote to (same rule as ToolBox.outp, without creating directories)."""
    if name:
        pp = Path(os.path.expanduser(name))
        if not pp.is_absolute():
            pp = (tb.out / pp) if len(pp.parts) == 1 else (tb.workdir / pp)
    else:
        pp = tb.out / default
    return str(pp)


def _db(ctx: Ctx, tdb: str):
    from .calc import get_db
    return get_db(ctx.tb._find_tdb(tdb))


def _el(s: str) -> str:
    s = str(s).strip().upper()
    m = re.fullmatch(r"[XW]\(([A-Z]{1,2})\)", s)
    return m.group(1) if m else s


def _is_celsius(unit) -> bool:
    return bool(unit) and str(unit).strip().upper().startswith("C")


def _num(x) -> str:
    """Short literal for a float."""
    x = float(x)
    return repr(int(x)) + ".0" if x == int(x) and abs(x) < 1e15 else repr(round(x, 10))


def _load_db_code(ctx: Ctx, tdb: str) -> str:
    ctx.need("from pycalphad import Database")
    return f"db = Database(str({ctx.tdb(tdb)}))\n"


def _colors(names, overrides=None) -> Dict[str, str]:
    from .mapping_tools import phase_color
    return {n: phase_color(n, overrides) for n in sorted(set(names))}


def _disordered_names(db, phases, rename: bool = True) -> Dict[str, tuple]:
    """Ordered phase -> (disordered-state name, number of ordered sublattices), as the tools' _Labeler."""
    out = {}
    if not rename:
        return out
    for name in phases:
        ph = db.phases.get(name)
        if ph is None or ph.model_hints.get("ordered_phase") != name:
            continue
        dis = ph.model_hints.get("disordered_phase")
        if not dis or dis not in db.phases:
            continue
        n_sub = len(db.phases[name].sublattices) - (len(db.phases[dis].sublattices) - 1)
        if n_sub < 2:
            continue
        m = re.fullmatch(r"([A-Z]\d+)_([A-Z]+)", dis)
        if m and f"{m.group(2)}_{m.group(1)}" in db.phases:
            dis = f"{m.group(2)}_{m.group(1)}"
        out[name] = (dis, n_sub)
    return out


def _names_code(dis: Dict[str, tuple], models_expr: str = "strategy.models") -> str:
    """Code defining names(compsets) -> phase labels (PHASE, PHASE#2, ...), with disordered-state renaming."""
    if not dis:
        return _T('''
            def names(compsets):
                """Phase name of each composition set; a second set of the same phase is PHASE#2."""
                out, seen = [], {}
                for cs in compsets:
                    n = cs.phase_record.phase_name
                    seen[n] = seen.get(n, 0) + 1
                    out.append(n if seen[n] == 1 else f"{n}#{seen[n]}")
                return out
            ''')
    return _T('''
        # An ordered phase (order/disorder model) whose ordered sublattices have equal site fractions is in its
        # disordered state: it is reported under the disordered phase's name (e.g. B2_BCC -> BCC_A2).
        DISORDERED = <<dis>>   # ordered phase: (name when disordered, number of ordered sublattices)
        site_index = {ph: {(y.sublattice_index, y.species.name): j for j, y in enumerate(<<models>>[ph].site_fractions)}
                      for ph in DISORDERED if ph in <<models>>}


        RENAMED = {}   # ordered phases that were reported under their disordered name


        def phase_name(cs, tol=1e-4):
            ph = cs.phase_record.phase_name
            if ph not in site_index:
                return ph
            dis, n_sub = DISORDERED[ph]
            y = np.asarray(cs.dof)[cs.phase_record.num_statevars:]
            idx = site_index[ph]
            for sp in {sp for (s, sp) in idx if s < n_sub}:
                if np.ptp([y[idx[(s, sp)]] if (s, sp) in idx else 0.0 for s in range(n_sub)]) > tol:
                    return ph
            RENAMED[ph] = dis
            return dis


        def names(compsets):
            """Phase name of each composition set; a second set of the same phase is PHASE#2."""
            out, seen = [], {}
            for cs in compsets:
                n = phase_name(cs)
                seen[n] = seen.get(n, 0) + 1
                out.append(n if seen[n] == 1 else f"{n}#{seen[n]}")
            return out
        ''', dis=Ctx.py(dis, 13), models=models_expr)


def _unit_code(t_celsius: bool, comp_unit: str, masses: Optional[Dict[str, float]]) -> str:
    """Code for T_disp(T_K) and c_disp(X_dict, el): temperature / composition in the displayed units."""
    cu = (comp_unit or "mole_fraction").lower()
    scale = 100 if cu.endswith("percent") else 1
    lines = [("def T_disp(T):\n    return T - KELVIN          # plotted in °C\n" if t_celsius else
              "def T_disp(T):\n    return T                   # plotted in K\n")]
    if cu.startswith("weight"):
        lines.append(f"masses = {Ctx.py(masses)}   # g/mol (TDB ELEMENT lines)\n\n\n"
                     "def to_weight(X):\n"
                     "    total = sum(X[e] * masses[e] for e in masses)\n"
                     "    return {e: X[e] * masses[e] / total for e in masses}\n\n\n"
                     "def c_disp(X, el):\n"
                     f"    return to_weight(X)[el]{' * 100' if scale == 100 else ''}   # mass "
                     f"{'percent' if scale == 100 else 'fraction'}\n")
    else:
        lines.append("def c_disp(X, el):\n"
                     f"    return X[el]{' * 100' if scale == 100 else ''}   # mole "
                     f"{'percent' if scale == 100 else 'fraction'}\n")
    return "\n\n".join(lines)


def _comp_label(el: str, comp_unit: str) -> str:
    cu = (comp_unit or "mole_fraction").lower()
    return f"{'Mass' if cu.startswith('weight') else 'Mole'} {'percent' if cu.endswith('percent') else 'fraction'} " \
           f"{el.capitalize()}"


def _T_label(celsius: bool) -> str:
    return "Temperature (°C)" if celsius else "Temperature (K)"


def _register_siblings(ctx: Ctx, png: str, suffixes: List[str]) -> Dict[str, str]:
    """Register files written next to a PNG (<stem><suffix>) as produced by this cell; returns suffix -> name."""
    p = Path(png)
    return {s: ctx.out_name(str(p.with_name(p.stem + s)), p.stem + s) for s in suffixes}


# ---------------------------------------------------------------------------------------------------------
# map_phase_diagram
# ---------------------------------------------------------------------------------------------------------
_MAP_EXTRACT = '''
    s_els = [e for e in strategy.elements if e != "VA"]   # order of cs.X and of the chemical potentials


    def comp(cs):
        return dict(zip(s_els, np.asarray(cs.X, dtype=float)))


    def temp(point):
        return float(np.squeeze(point.global_conditions[v.T]))


    # ZPF lines -> phase-boundary points (one row per phase and point) and tie-lines (one row per point)
    lines, boundary_rows, tieline_rows, seen = [], [], [], set()
    degenerate = duplicates = 0
    for li, zpf in enumerate(strategy.zpf_lines):
        if not zpf.points:
            continue
        # skip artefacts: zero-length tie-lines along a pure-element/binary edge, and lines traced twice
        if all(max(np.max(np.abs(np.asarray(a.X, float) - np.asarray(b.X, float)))
                   for a, b in itertools.combinations(p.stable_composition_sets, 2)) < 1e-5
               for p in zpf.points if len(p.stable_composition_sets) > 1):
            degenerate += 1
            continue
        ends = tuple(sorted((round(temp(p), 3),) + tuple(np.round(np.concatenate(
            [np.asarray(cs.X, float) for cs in p.stable_composition_sets]), 5)) for p in (zpf.points[0], zpf.points[-1])))
        if ends in seen:
            duplicates += 1
            continue
        seen.add(ends)
        line_phases = names(zpf.points[len(zpf.points) // 2].stable_composition_sets)
        region = " + ".join(sorted(line_phases))
        points = []
        for pi, p in enumerate(zpf.points):
            T = temp(p)
            mu = dict(zip(s_els, map(float, p.chemical_potentials)))
            tie = {"line": li, "point": pi, "region": region, "T_K": T, "T_C": T - KELVIN}
            per_phase = []
            for k, (name, cs) in enumerate(zip(line_phases, p.stable_composition_sets)):
                X = comp(cs)
                per_phase.append((name, X))
                row = {"line": li, "region": region, "point": pi, "phase": name, "T_K": T, "T_C": T - KELVIN}
                row.update({f"X({e})": X[e] for e in elements})<<weight_cols>>
                row.update({f"MU({e})": mu[e] for e in elements})
                boundary_rows.append(row)
                tie[f"phase_{k + 1}"] = name
                tie.update({f"X{k + 1}({e})": X[e] for e in elements})
            tie.update({f"MU({e})": mu[e] for e in elements})
            tieline_rows.append(tie)
            points.append({"T": T, "phases": per_phase})
        lines.append({"phases": line_phases, "points": points})

    # three-phase nodes: invariant reactions (binary) / tie-triangles (ternary), without duplicates
    invariants = []
    for node in strategy.node_queue.nodes:
        if len(node.stable_composition_sets) != 3:
            continue
        node_phases = names(node.stable_composition_sets)
        T = temp(node)
        X = {n: comp(cs) for n, cs in zip(node_phases, node.stable_composition_sets)}
        key = sorted(tuple(round(c[e], 3) for e in elements) for c in X.values())
        dup = next((inv for inv in invariants if abs(inv["T_K"] - T) < 0.05 and inv["key"] == key), None)
        if dup is not None:   # found again (possibly under other names, e.g. a second composition set)
            alias = " + ".join(sorted(node_phases))
            if alias != " + ".join(sorted(dup["phases"])):
                dup.setdefault("also_found_as", []).append(alias)
            continue
        invariants.append({"T_K": round(T, 3), "T_C": round(T - KELVIN, 3), "phases": node_phases, "key": key,
                           "compositions": {n: {e: round(c[e], 6) for e in elements} for n, c in X.items()},
                           "MU": {e: round(float(mu), 3) for e, mu in zip(s_els, node.chemical_potentials)},
                           "P": float(np.squeeze(node.global_conditions.get(v.P, 101325.0))),
                           "type": "three-phase region", "reaction_on_cooling": ""})
    '''

_MAP_CLASSIFY = '''

    LIQUIDS = <<liquids>>   # liquid phases of this system


    def is_liquid(name):
        return name.split("#")[0] in LIQUIDS


    def classify(inv, b=elements[1]):
        """Eutectic-type (middle phase decomposes on cooling) or peritectic-type (end phases form the middle phase)."""
        lo, mid, hi = sorted(inv["phases"], key=lambda n: inv["compositions"][n][b])
        xs = [inv["compositions"][n][b] for n in (lo, mid, hi)]
        # does the two-phase field of the two end phases lie above or below the invariant? (mapped ZPF lines)
        votes = 0
        for ln in lines:
            pts = ln["points"]
            if len(pts) < 2 or len(ln["phases"]) != 2:
                continue
            for end in (pts[0], pts[-1]):
                if abs(end["T"] - inv["T_K"]) > 0.5:
                    continue
                idx = []
                for _, X in end["phases"]:
                    d = [abs(X[b] - x) for x in xs]
                    k = int(np.argmin(d))
                    idx.append(k if d[k] < 5e-3 else -1)
                if sorted(idx) == [0, 2]:
                    votes += 1 if np.mean([p["T"] for p in pts]) > inv["T_K"] else -1
        side, how = int(np.sign(votes)), "zpf-lines"
        if side == 0:   # undecided: equilibrium 1 K below the invariant at the middle-phase composition
            how = "equilibrium check"
            try:
                eq = equilibrium(db, comps, phases, {v.T: inv["T_K"] - 1.0, v.P: <<P>>, v.N: 1,
                                                     v.X(b): min(max(xs[1], 1e-6), 1 - 1e-6)})
                x_eq = np.atleast_1d(np.squeeze(eq.X.sel(component=b).values))
                np_eq = np.atleast_1d(np.squeeze(eq.NP.values))
                ok = np.isfinite(np_eq) & (np_eq > 0)
                side = 1 if np.nansum(np_eq[ok & (np.abs(x_eq - xs[1]) < 5e-3)]) > 0.9 else -1
            except Exception:
                return "unclassified", " + ".join((lo, mid, hi)), None
        liq = {n: is_liquid(n) for n in (lo, mid, hi)}
        if side > 0:    # lo + hi -> mid on cooling
            kind = ("unusual (liquid forms on cooling)" if liq[mid] else "syntectic" if liq[lo] and liq[hi]
                    else "peritectic" if liq[lo] or liq[hi] else "peritectoid")
            return kind, f"{lo} + {hi} -> {mid}", how
        if liq[mid]:    # mid -> lo + hi on cooling
            return ("monotectic" if liq[lo] or liq[hi] else "eutectic"), f"{mid} -> {lo} + {hi}", how
        return ("metatectic" if liq[lo] or liq[hi] else "eutectoid"), f"{mid} -> {lo} + {hi}", how


    for inv in invariants:
        inv["type"], inv["reaction_on_cooling"], how = classify(inv)
        if how:
            inv["classified_by"] = how
    invariants.sort(key=lambda inv: -inv["T_K"])
    '''

_MAP_WRITE = '''

    inv_rows = [{"invariant": k, "T_K": inv["T_K"], "T_C": inv["T_C"], "type": inv["type"],
                 "reaction_on_cooling": inv["reaction_on_cooling"], "phases": " + ".join(inv["phases"]), "phase": n,
                 **{f"X({e})": inv["compositions"][n][e] for e in elements}}
                for k, inv in enumerate(invariants) for n in inv["phases"]]
    pd.DataFrame(boundary_rows).to_csv(OUT / <<bcsv>>, index=False, float_format="%.8g")
    pd.DataFrame(tieline_rows).to_csv(OUT / <<tcsv>>, index=False, float_format="%.8g")
    pd.DataFrame(inv_rows, columns=["invariant", "T_K", "T_C", "type", "reaction_on_cooling", "phases", "phase"]
                 + [f"X({e})" for e in elements]).to_csv(OUT / <<icsv>>, index=False, float_format="%.8g")
    for inv in invariants:
        print(f"{inv['type']}: {inv['reaction_on_cooling'] or ' + '.join(inv['phases'])} at {inv['T_K']:.2f} K "
              f"({inv['T_C']:.2f} °C)")
    '''

_MAP_PLOT = '''

    # plot: phase boundaries coloured by phase, tie-lines in grey, three-phase triangles in red
    colors = <<colors>>


    <<xy>>

    <<figline>>
    all_xy = np.array([xy(X, p["T"]) for ln in lines for p in ln["points"] for _, X in p["phases"]])
    tiny = np.ptp(all_xy, axis=0) * 1e-3
    for ln in lines:
        for k, name in enumerate(ln["phases"]):
            pts = np.array([xy(p["phases"][k][1], p["T"]) for p in ln["points"]])
            c = colors[name.split("#")[0]]
            if np.ptp(pts[:, 0]) < tiny[0] and np.ptp(pts[:, 1]) < tiny[1]:   # line collapsed onto one point
                ax.scatter([pts[:, 0].mean()], [pts[:, 1].mean()], color=c, s=8, zorder=3)
            else:
                ax.plot(pts[:, 0], pts[:, 1], color=c, lw=1.3, solid_capstyle="butt", zorder=2)<<tieline_block>>
    for inv in invariants:
        pts = [xy(inv["compositions"][n], inv["T_K"]) for n in inv["phases"]]
        ax.plot([p[0] for p in pts + pts[:1]], [p[1] for p in pts + pts[:1]], color="#d62728", lw=1.1, zorder=2.5)<<node_block>>
    '''

_MAP_PLOT_SHADED = '''

    # plot (shaded style): black phase boundaries and invariant lines, two-phase fields filled aliceblue between their
    # two boundaries (binary); grey lines, aliceblue two-phase fields (ternary); no tie-lines, no legend
    <<xy>>

    <<figline>>
    boundaries = [[np.array([xy(p["phases"][k][1], p["T"]) for p in ln["points"]]) for k in range(len(ln["phases"]))]
                  for ln in lines]
    invariant_points = [np.array([xy(inv["compositions"][n], inv["T_K"]) for n in inv["phases"]]) for inv in invariants]
    draw_shaded(ax, boundaries, invariant_points, kind=<<kind>>)<<node_block>>
    '''

_NODE_BLOCK_SHADED = '''
    colors = <<colors>>
    for inv in invariants:
        for n in inv["phases"]:
            p = xy(inv["compositions"][n], inv["T_K"])
            ax.scatter([p[0]], [p[1]], color=colors[n.split("#")[0]], s=14, zorder=3, edgecolors="k", linewidths=0.3)'''

_TIELINE_BLOCK = '''
        if len(ln["phases"]) == 2:
            segs = [[xy(p["phases"][0][1], p["T"]), xy(p["phases"][1][1], p["T"])] for p in ln["points"][::<<tl>>]]
            ax.add_collection(LineCollection(segs, colors=[(0.45, 0.45, 0.45, 0.45)], linewidths=0.5, zorder=1))'''

_NODE_BLOCK = '''
        for n, p in zip(inv["phases"], pts):
            ax.scatter([p[0]], [p[1]], color=colors[n.split("#")[0]], s=14, zorder=3, edgecolors="k", linewidths=0.3)'''

_LEGEND = '''
    found = sorted({n for ln in <<lines>> for n in <<ln_phases>>} | {n for inv in <<invs>> for n in inv["phases"]})
    ax.legend(handles=[mpatches.Patch(color=colors[n.split("#")[0]], label=n) for n in found], loc="center left",
              bbox_to_anchor=(1.02, 0.5), fontsize=8, frameon=False<<ltitle>>)
    '''


@generator("map_phase_diagram")
def _gen_map_phase_diagram(args, ctx: Ctx) -> str:
    from . import mapping_tools as mt
    a = {"temperature_unit": "K", "temperature_step": 10.0, "composition_unit": "mole_fraction", "tielines": 3,
         "label_nodes": False, "rename_disordered": True, "pressure": 101325.0, **args}
    shaded = _style(ctx, a) == "shaded"
    if a.get("potential_axis"):
        ctx.fallback("T-MU (potential_axis) diagrams are drawn from chemical potentials with the toolkit's "
                     "dilute-limit axis scaling")
    db = _db(ctx, a["tdb"])
    els = mt._check_elements(db, a["elements"])
    fixed = {_el(k): float(val) for k, val in (a.get("fixed") or {}).items()}
    vary = a.get("vary") or (a.get("right") if fixed else None)
    if fixed or vary:
        kind = "isopleth"
    elif len(els) == 2:
        kind = "binary"
    elif len(els) == 3 and a.get("temperature") is not None:
        kind = "ternary"
    else:
        ctx.fallback("unsupported combination of arguments")
    from .orientation import binary_order, corner_label_code, orientation_text, ternary_order
    if kind == "binary":        # the orientation the tool used: [left, right] / [bottom-left, bottom-right, top]
        els = binary_order(els, a.get("right"))
    elif kind == "ternary":
        els = ternary_order(els, a.get("corners"), a.get("right"))
    ph = mt._select_phases(db, els, a.get("phases"), a.get("exclude_phases"))
    comps = sorted(els) + ["VA"]
    unit = a["temperature_unit"]
    tC = _is_celsius(unit)
    cu = (a["composition_unit"] or "mole_fraction").lower()
    masses = mt._masses(db, sorted(els)) if cu.startswith("weight") else None
    P = float(a["pressure"])
    xr = [float(a["x_range"][0]), float(a["x_range"][1])] if a.get("x_range") else [0.0, 1.0]
    step = float(a.get("composition_step") or 0.02)
    system = "-".join(e.capitalize() for e in els)

    # file names as the tool wrote them
    tag = "-".join(els) if kind != "isopleth" else "-".join(_el(e) for e in a["elements"])
    if fixed or vary:
        tag += f"_isopleth_{_el(vary or '')}_" + "_".join(f"{_el(k)}{val:g}" for k, val in (a.get("fixed") or {}).items())
    elif a.get("temperature") is not None and len(els) == 3:
        tag += f"_{a['temperature']:g}{unit}"
    png = _outp(ctx.tb, (a["output_prefix"] + ".png") if a.get("output_prefix") else None,
                f"map_{tag}_{Path(ctx.tb._find_tdb(a['tdb'])).stem}.png")
    png_name = ctx.out_name(png, "phase_diagram.png")

    ctx.need("import itertools", "import numpy as np", "import pandas as pd", "import matplotlib.pyplot as plt",
             "import matplotlib.patches as mpatches", "from pycalphad import Database", "from pycalphad import variables as v")
    conds = [f"v.P: {_num(P)}", "v.N: 1"]
    if kind in ("binary", "isopleth"):
        if a.get("tmin") is None or a.get("tmax") is None:
            ctx.fallback(f"{kind} needs tmin and tmax")
        tlo, thi = sorted((mt._toK(a["tmin"], unit), mt._toK(a["tmax"], unit)))
        conds.append(f"v.T: ({_num(tlo)}, {_num(thi)}, {_num(a['temperature_step'])})")
    if kind == "binary":
        conds.append(f"v.X({els[1]!r}): ({_num(xr[0])}, {_num(xr[1])}, {_num(step)})")
        strat = "BinaryStrategy"
    elif kind == "ternary":
        Tk = mt._toK(a["temperature"], unit)
        conds.append(f"v.T: {_num(Tk)}")
        conds += [f"v.X({els[1]!r}): (0, 1, {_num(step)})", f"v.X({els[2]!r}): (0, 1, {_num(step)})"]
        strat = "TernaryStrategy"
    else:
        axis_el = _el(vary)
        if not vary or not fixed or axis_el not in els or any(k not in els for k in fixed) or axis_el in fixed \
                or len(fixed) != len(els) - 2:
            ctx.fallback("invalid isopleth arguments")
        xmax = min(xr[1], 1.0 - sum(fixed.values()))
        conds.append(f"v.X({axis_el!r}): ({_num(xr[0])}, {_num(xmax)}, {_num(step)})")
        conds += [f"v.X({k!r}): {_num(val)}" for k, val in fixed.items()]
        strat = "IsoplethStrategy"
    ctx.need(f"from pycalphad.mapping import {strat}")
    dis = _disordered_names(db, ph, a["rename_disordered"])
    all_names = list(ph) + [d for d, _ in dis.values()]
    colors = _colors(all_names, a.get("phase_colors"))
    from .mapping_check import check_mode, grid_size
    check = check_mode(a.get("check_mapping"))
    label_regions = a.get("label_regions")
    if label_regions is None:       # the tool's default: labels on in the shaded style (no legend)
        label_regions = shaded
    if kind == "isopleth" or check == "off":
        ctx.need("import time")
        run_code = ("t_map = time.time()\nstrategy.do_map()\n"
                    "run = {'finished': True, 'iterations': None, 'seconds': round(time.time() - t_map, 1)}\n")
        checked = False
    else:
        unit_ = a["temperature_unit"]
        crec = mt._recipe("", kind, els, ph, P, _conds_values(kind, a, unit_, xr, step),
                          mt._Units(unit_, cu, masses), unit_, True, xr, step, float(a["temperature_step"]))
        run_code = check_code(ctx, crec, check, float(a.get("time_limit_s") or 600.0),
                              grid_size(kind, bool(label_regions), check)).lstrip("\n")
        checked = True

    head = _load_db_code(ctx, a["tdb"]) + _T('''
        elements = <<els>>   # first = balance element<<axes_note>>
        comps = <<comps>>
        phases = <<ph>>   # all phases that can form<<phase_note>>
        conds = {<<conds>>}
        KELVIN = 273.15

        # map the <<what>> with pycalphad's mapping strategy (note: which lines get traced - and how often - can
        # vary slightly with Python's hash seed; the phase boundaries and invariants are the same)
        strategy = <<strat>>(db, comps, phases, conds)
        <<run>>
        ''', run=run_code, els=Ctx.py(els), comps=Ctx.py(comps), ph=Ctx.py(ph), conds=", ".join(conds), strat=strat,
                                                  axes_note={"binary": f"; {orientation_text(kind, els)[13:]}",
                                                             "ternary": f"; {orientation_text(kind, els)[13:]}",
                                                             "isopleth": ""}[kind],
                                                  phase_note=(" (selection as requested)" if a.get("phases") or
                                                              a.get("exclude_phases") else ""),
                                                  what={"binary": "binary phase diagram",
                                                        "ternary": "isothermal section",
                                                        "isopleth": "isopleth"}[kind])
    head = head.replace("KELVIN = 273.15\n", "KELVIN = 273.15\n" + (mapping_check_helpers(ctx) if checked else ""), 1)
    head += _names_code(dis) + "\n\n" + (shaded_helpers(ctx) if shaded else "")
    units = _unit_code(tC, cu, masses) + "\n\n"
    ctx.mapped(str(Path(png).with_suffix(".json")))
    if kind == "isopleth":
        regions = bool(shaded or label_regions)
        return head + units + _isopleth_code(ctx, a, els, axis_el, fixed, png, png_name, colors, tC, cu, tlo, thi,
                                             system, shaded, db=db, ph=ph, P=P, xr=xr, step=step, masses=masses,
                                             check=check, label_regions=bool(label_regions)) + _summary_code(
            ctx, a, "isopleth", els, png, system, cu, shaded, check, checked, bool(dis), bool(label_regions),
            axis_el=axis_el, balance=[e for e in els if e != axis_el and e not in fixed][0], regions=regions)

    # ---- binary / ternary (tie-line strategies)
    weight_cols = ('\n            row.update({f"W({e})": w for e, w in to_weight(X).items()})'
                   if cu.startswith("weight") else "")
    body = _T(_MAP_EXTRACT, weight_cols=weight_cols) + "\n"
    if kind == "binary":
        ctx.need("from pycalphad import equilibrium")
        liquids = sorted(p for p in ph if mt._is_liquid(db, p))
        body += "\n" + _T(_MAP_CLASSIFY, liquids=Ctx.py(set(liquids)) if liquids else "set()", P=_num(P)) + "\n"
    names_ = _register_siblings(ctx, png, ["_boundaries.csv", "_tielines.csv", "_invariants.csv"])
    body += "\n" + _T(_MAP_WRITE, bcsv=repr(names_["_boundaries.csv"]), tcsv=repr(names_["_tielines.csv"]),
               icsv=repr(names_["_invariants.csv"]))
    tl = int(a["tielines"] or 0) if not shaded else 0
    if tl > 0:
        ctx.need("from matplotlib.collections import LineCollection")
    if kind == "binary":
        xy = ("def xy(X, T):\n    return c_disp(X, elements[1]), T_disp(T)   # (composition, temperature) "
              "as displayed\n")
        figline = "fig, ax = plt.subplots(figsize=(7.2, 5.4))"
    else:
        # the triangular projection needs fractions: percent is shown through the tick labels only
        if cu.endswith("percent"):
            xy = ("def xy(X, T):\n    # fractions: the triangular projection needs them (tick labels show percent)\n"
                  "    return c_disp(X, elements[1]) / 100, c_disp(X, elements[2]) / 100")
        else:
            xy = "def xy(X, T):\n    return c_disp(X, elements[1]), c_disp(X, elements[2])"
        ctx.need("from pycalphad.plot import triangular  # registers the 'triangular' projection")
        figline = 'fig = plt.figure(figsize=(7.0, 6.0))\nax = fig.add_subplot(projection="triangular")'
    if shaded:
        body += "\n\n" + _T(_MAP_PLOT_SHADED, xy=xy, figline=figline, kind=repr(kind),
                           node_block=_indent(_T(_NODE_BLOCK_SHADED, colors=Ctx.py(colors, 9)), 0)
                           if a["label_nodes"] else "")
    else:
        body += "\n\n" + _T(_MAP_PLOT, tl=tl, colors=Ctx.py(colors, 9), xy=xy, figline=figline,
                           tieline_block=_indent(_T(_TIELINE_BLOCK, tl=tl), 4) if tl > 0 else "",
                           node_block=_indent(_T(_NODE_BLOCK), 4) if a["label_nodes"] else "")
    body = textwrap.dedent(body)
    # axes
    ax_code = []
    clab = (lambda el: _comp_label(el, cu))
    if kind == "binary":
        ax_code.append(f"ax.set_xlabel({clab(els[1])!r})")
        ax_code.append(f"ax.set_ylabel({_T_label(tC)!r})")
        lo = mt._Units(unit, cu, masses).comp({els[0]: 1 - xr[0], els[1]: xr[0]}, els[1])
        hi = mt._Units(unit, cu, masses).comp({els[0]: 1 - xr[1], els[1]: xr[1]}, els[1])
        ax_code.append(f"ax.set_xlim({_num(lo)}, {_num(hi)})")
        ax_code.append(f"ax.set_ylim({_num(tlo - KELVIN if tC else tlo)}, {_num(thi - KELVIN if tC else thi)})")
        ax_code.append("ticks_in(ax)   # minor ticks on, all ticks pointing in on all four sides" if shaded
                       else "ax.grid(alpha=0.25)")
    else:
        ax_code.append(f"ax.set_xlabel({clab(els[1])!r})")
        ax_code.append(f"ax.set_ylabel({clab(els[2])!r})")
        ax_code.append("ax.set_xlim(0, 1)")
        ax_code.append("ax.set_ylim(0, 1)")
        if cu.endswith("percent"):
            ctx.need("from matplotlib.ticker import FuncFormatter")
            ax_code.append("ax.xaxis.set_major_formatter(FuncFormatter(lambda val, pos: f'{val * 100:g}'))")
            ax_code.append("ax.yaxis.set_major_formatter(FuncFormatter(lambda val, pos: f'{val * 100:g}'))")
        ax_code.append("ax.yaxis.label.set_rotation(60)")
        ax_code.append("ax.yaxis.set_label_coords(x=0.12, y=0.5)")
        ax_code.append(corner_label_code(els).rstrip("\n"))
    if a.get("xlim"):
        ax_code.append(f"ax.set_xlim({_num(a['xlim'][0])}, {_num(a['xlim'][1])})")
    if a.get("ylim"):
        ax_code.append(f"ax.set_ylim({_num(a['ylim'][0])}, {_num(a['ylim'][1])})")
    legend = _T(_LEGEND, lines="lines", ln_phases='ln["phases"]', invs="invariants", ltitle="") if not shaded else ""
    if kind == "ternary":
        title = f"{system}   {Tk - KELVIN:.0f} °C ({Tk:.0f} K)"
    else:
        title = system
    title = a.get("title") or title
    labels = ""
    if label_regions:
        rec = mt._recipe(ctx.tb._find_tdb(a["tdb"]), kind, els, ph, P, _conds_values(kind, a, unit, xr, step),
                         mt._Units(unit, cu, masses), unit, a["rename_disordered"], xr, step,
                         float(a["temperature_step"]))
        labels = label_helpers(ctx) + label_code(
            ctx, rec, db, "strategy.models", mapped=True, invariants=a.get("label_invariants", True),
            phase_labels=a.get("phase_labels"), style=_label_style(a.get("label_style"), shaded),
            check_expr="mapping_check" if checked else None)
        if not shaded:
            labels += f"prettify_legend(ax, {list(els)!r}, {a.get('phase_labels')!r})\n"
    fmts = a.get("formats")
    if fmts is None:
        fmts = ["svg", "pdf"] if label_regions or shaded else []
    tail = "\n".join(ax_code) + "\n" + legend + f"ax.set_title({title!r}, fontsize=11)\n" + labels + (
        f"fig.savefig(OUT / {png_name!r}, dpi=150, bbox_inches='tight')\n") + _vector_saves(ctx, png, fmts)
    return head + units + body + "\n" + tail + _summary_code(ctx, a, kind, els, png, system, cu, shaded, check,
                                                              checked, bool(dis), bool(label_regions))


_REGION_SUMMARY = '''
    def region_summary(lines):
        """Per two-phase region: number of lines, temperature range, composition range of each phase."""
        regions = {}
        for ln in lines:
            rg = regions.setdefault(" + ".join(sorted(ln["phases"])), {"lines": 0, "T_K": [np.inf, -np.inf], "X": {}})
            rg["lines"] += 1
            for p in ln["points"]:
                rg["T_K"] = [min(rg["T_K"][0], p["T"]), max(rg["T_K"][1], p["T"])]
                for n, X in p["phases"]:
                    cr = rg["X"].setdefault(n, {e: [1.0, 0.0] for e in elements[1:]})
                    for e in elements[1:]:
                        cr[e] = [min(cr[e][0], X[e]), max(cr[e][1], X[e])]
        return [{"region": name, "lines": rg["lines"], <<trange>>
                 "X_ranges": {n: {f"X({e})": [round(a, 5), round(b, 5)] for e, (a, b) in cr.items()}
                              for n, cr in rg["X"].items()}} for name, rg in sorted(regions.items())]
    '''

_ISO_REGION_SUMMARY = '''
    def region_summary(lines):
        """Per zero-fraction phase: number of ZPF lines and their temperature range."""
        regions = {}
        for ln in lines:
            rg = regions.setdefault(ln["zero_phase"], {"lines": 0, "T_K": [np.inf, -np.inf]})
            rg["lines"] += 1
            rg["T_K"] = [min(rg["T_K"][0], min(ln["T_K"])), max(rg["T_K"][1], max(ln["T_K"]))]
        return [{"zero_phase_line": k, "lines": r["lines"], "T_range_K": [round(r["T_K"][0], 2), round(r["T_K"][1], 2)]}
                for k, r in sorted(regions.items())]
    '''


def _summary_code(ctx: Ctx, a: Dict, kind: str, els, png: str, system: str, cu: str, shaded: bool, check: str,
                  checked: bool, renames: bool, labelled: bool, axis_el: Optional[str] = None,
                  balance: Optional[str] = None, regions: bool = False) -> str:
    """Code writing the summary JSON map_phase_diagram writes next to the PNG (same keys; times, paths and the
    hash seed are this run's)."""
    from . import mapping_check as mc
    from .orientation import orientation_info
    ctx.need("import json", "import os", "import sys")
    p = Path(png)
    files = {"png": p.name}
    if kind == "isopleth":
        files.update(boundaries_csv=p.stem + "_zpf_lines.csv", invariants_csv=p.stem + "_invariants.csv")
    else:
        files.update(boundaries_csv=p.stem + "_boundaries.csv", tielines_csv=p.stem + "_tielines.csv",
                     invariants_csv=p.stem + "_invariants.csv")
    fmts = a.get("formats")
    if fmts is None:
        fmts = ["svg", "pdf"] if (shaded or labelled) else []
    for ext in fmts:
        ext = str(ext).lower().lstrip(".")
        if ext != "png":
            files[ext] = p.stem + "." + ext
    files["json"] = p.stem + ".json"
    files = {k: ctx.out_name(str(p.with_name(name)), name) for k, name in files.items()}
    files_code = "{" + ", ".join(f"{k!r}: str(OUT / {name!r})" for k, name in files.items()
                                 if k != "json") + "}"     # (the tool lists the JSON only in its reply)
    if kind == "isopleth":
        if not regions:
            check_json = Ctx.py({"mode": check, "note": "not checked (T-MU diagrams, and isopleths without regions, "
                                                        "are not)"})
        elif check == "off":
            check_json = Ctx.py({"mode": check, "note": "off (isopleth regions found, lines not checked)"})
        else:
            check_json = f'{{"mode": {check!r}, "kind": "isopleth", **iso["check"]}}'
        phases_found = 'sorted({ln["zero_phase"] for ln in lines} | {n for inv in invariants for n in inv["phases"]})'
        invs = ('[{**{k: val for k, val in inv.items() if k != "segment"}, f"X({axis_el})_segment": inv["segment"]}\n'
                '                   for inv in invariants]')
        mapping = '{**run, "zpf_lines": len(strategy.zpf_lines), "degenerate_edge_lines_dropped": 0, ' \
                  '"duplicate_lines_dropped": 0}'
        regions = _T(_ISO_REGION_SUMMARY)
        orient = orientation_info(kind, els, axis_el, balance)
    else:
        check_json = ('{k: val for k, val in report_json(mapping_check).items() if k != "run"}' if checked else
                      Ctx.py(mc.report_json(mc.check_and_repair(None, None, {}, mode="off")), 20))
        phases_found = ('sorted({n for ln in lines for n in ln["phases"]} | {n for inv in invariants '
                        'for n in inv["phases"]})')
        drop = '("key",)' if kind == "binary" else '("key", "type", "reaction_on_cooling")'
        invs = f'[{{k: val for k, val in inv.items() if k not in {drop}}} for inv in invariants]'
        mapping = ('{**run, "zpf_lines": len(strategy.zpf_lines), "degenerate_edge_lines_dropped": degenerate,\n'
                   '                "duplicate_lines_dropped": duplicates}')
        regions = _T(_REGION_SUMMARY, trange=('"T_range_K": [round(rg["T_K"][0], 2), round(rg["T_K"][1], 2)],'
                                              if kind == "binary" else "  # (isotherm: no T range)"))
        orient = orientation_info(kind, els)
    code = ("\n\n# ---- the summary JSON the tool writes next to the PNG: conditions, phases, mapping and its check, "
            "phase\n# regions, invariants, labels (run times, paths and the hash seed are this run's)\n"
            + ("run = mapping_check[\"run\"]   # the mapping run (finished, iterations, seconds)\n\n\n"
               if checked else "\n")
            + regions + "\n\n" + _T('''
        summary = {
            "kind": <<kind>>, "system": <<system>>, "tdb": str(<<tdb>>), "elements": elements,
            "phases_considered": phases, "phases_found": <<found>>,
            "conditions": {str(k): (list(val) if isinstance(val, tuple) else val) for k, val in conds.items()},
            "temperature_unit_in_files": "K and C columns", "composition_unit_in_plot": <<cu>>,
            "mapping": <<mapping>>,
            "mapping_check": <<check>>,
            "regions": region_summary(lines),
            "invariants": <<invs>>,
            "files": <<files>>,
            "renamed_when_disordered": <<renamed>>, "style": <<style>>,
            "hash_seed": "0" if sys.flags.hash_randomization == 0 else (os.environ.get("PYTHONHASHSEED")
                                                                         or "random"),<<preg>>
            "orientation": <<orient>>,
        }
        ''', kind=repr(kind), system=repr(system), tdb=ctx.tdb(a["tdb"]), found=phases_found, cu=repr(cu),
                   mapping=mapping, check=check_json, invs=invs, files=files_code,
                   renamed="dict(RENAMED)" if renames else "{}", style=repr("shaded" if shaded else "color"),
                   orient=Ctx.py(orient), preg=("\n    \"phase_regions\": " + ('iso["summary"],' if regions else "[],"))
                   if kind == "isopleth" else ""))
    if labelled:
        code += 'summary["labels"] = {k: report[k] for k in ("labels", "skipped", "seconds", "equilibrium_checks")}\n'
    code += f"(OUT / {files['json']!r}).write_text(json.dumps(summary, indent=1, default=str))\n"
    return code


def _conds_values(kind, a, unit, xr, step):
    """The T condition of map_phase_diagram as the recipe needs it (K)."""
    from pycalphad import variables as v
    from . import mapping_tools as mt
    if kind == "binary":
        tlo, thi = sorted((mt._toK(a["tmin"], unit), mt._toK(a["tmax"], unit)))
        return {v.T: (tlo, thi, float(a.get("temperature_step") or 10.0))}
    return {v.T: mt._toK(a["temperature"], unit)}


def _style(ctx: Ctx, a: Dict) -> str:
    """The diagram style the tool used: the call's style, else the session default (config diagram_style)."""
    from .diagram_style import check_style
    return check_style(a.get("style") or getattr(ctx.tb, "diagram_style", "shaded"))


def _label_style(style: Optional[Dict], shaded: bool) -> Optional[Dict]:
    """Label style as the tool applied it: invariant temperatures in black on a shaded diagram."""
    if not shaded:
        return style
    from .diagram_style import LABEL_COLOR
    return {"invariant_color": LABEL_COLOR, **(style or {})}


def _phase_names_source(ctx: Ctx) -> str:
    """Source of phase_names (format_phase, region_label), once per script (also part of label_helpers)."""
    if ctx.__dict__.get("_label_helpers") or ctx.__dict__.get("_phase_names"):
        return ""
    ctx._phase_names = True
    from . import phase_names
    from .codegen_compare import module_source
    return module_source(ctx, phase_names, "phase names -> labels (BCC_A2 -> BCC, SIGMA -> σ, FE3CR4 -> Fe₃Cr₄)")


def shaded_helpers(ctx: Ctx) -> str:
    """Source of the shaded drawing style (draw_shaded, ticks_in, ...), once per script."""
    if ctx.__dict__.get("_shaded_helpers"):
        return ""
    ctx._shaded_helpers = True
    from . import diagram_style
    from .codegen_compare import module_source
    return module_source(ctx, diagram_style, "shaded phase-diagram style: black boundaries, aliceblue two-phase "
                                             "fields, grey ternaries, ticks pointing in")


VECTOR_RC_CODE = "{'svg.fonttype': 'none', 'pdf.fonttype': 42, 'ps.fonttype': 42}"


def _vector_saves(ctx: Ctx, png: str, fmts) -> str:
    """Extra SVG/PDF/EPS files with the text kept as text (editable in Inkscape / Illustrator)."""
    out = []
    for ext in fmts or []:
        ext = str(ext).lower().lstrip(".")
        if ext == "png":
            continue
        p = str(Path(png).with_suffix("." + ext))
        out.append(f"    fig.savefig(OUT / {ctx.out_name(p, Path(p).name)!r}, dpi=150, bbox_inches='tight')")
    if not out:
        return ""
    if names_in_script(ctx):        # phase names drawn: subscripts typeset, Inkscape subscripts in the SVG
        return ("# text stays text (one editable object per label); SVG subscripts as Inkscape writes them\n"
                + "\n".join("save_vector(fig, " + o.strip()[len("fig.savefig("):] for o in out) + "\n")
    return (f"with plt.rc_context({VECTOR_RC_CODE}):   # text stays text (one editable object per label)\n"
            + "\n".join(out) + "\n")


def names_in_script(ctx: Ctx) -> bool:
    """phase_names (set_name, save_vector ...) is copied into this script."""
    return bool(ctx.__dict__.get("_label_helpers") or ctx.__dict__.get("_phase_names"))


# ---------------------------------------------------------------------------------------------------------
# phase-region labels (shared with the edit_figure generator)
# ---------------------------------------------------------------------------------------------------------
def mapping_check_helpers(ctx: Ctx) -> str:
    """Sources of phase_regions and mapping_check (the check / repair of mapped ZPF lines), once per script."""
    from . import mapping_check, phase_regions
    from .codegen_compare import module_source
    out = ""
    if not ctx.__dict__.get("_phase_regions"):
        ctx._phase_regions = True
        out += module_source(ctx, phase_regions, "phase regions: two-phase polygons from ZPF lines, equilibrium "
                                                 "grid, stable phase set at a point")
    if not ctx.__dict__.get("_mapping_check"):
        ctx._mapping_check = True
        out += module_source(ctx, mapping_check, "mapping check: traced ZPF lines checked against equilibrium "
                                                 "(metastable tie-lines trimmed, duplicate traces dropped, missing "
                                                 "boundaries traced again)")
    return out


def label_helpers(ctx: Ctx) -> str:
    """Source of the label helpers (names, regions, mapping check, placement), once per script."""
    if ctx.__dict__.get("_label_helpers"):
        return ""
    ctx._label_helpers = True
    from . import label_place, phase_names
    from .codegen_compare import module_source
    names = "" if ctx.__dict__.get("_phase_names") else module_source(
        ctx, phase_names, "phase names -> labels (BCC_A2 -> BCC, SIGMA -> σ, FE3CR4 -> Fe₃Cr₄)")
    return (names + mapping_check_helpers(ctx)
            + module_source(ctx, label_place, "label placement: region raster, obstacles, text fitting, names of "
                                              "stoichiometric phases outside the axes"))


def check_code(ctx: Ctx, recipe: Dict, mode: str, time_limit: float, grid_n, strategy: str = "strategy",
               out: str = "mapping_check") -> str:
    """Code that maps with ``strategy`` and checks / repairs its lines as the tools do (mapping_check)."""
    rec = {k: val for k, val in recipe.items() if k != "tdb"}
    return _T('''
        # map, then check the traced lines against equilibrium as the tool does (mapping_check): equilibrium at points
        # of every line (metastable tie-lines trimmed), duplicate traces dropped, two-phase areas of an equilibrium
        # grid that no line bounds traced again
        check_recipe = <<recipe>>
        <<out>> = map_and_check(<<strategy>>, db, check_recipe, <<mode>>, <<tl>>, <<n>>)
        print(<<out>>["text"])
        ''', recipe=Ctx.py(rec, 15), out=out, strategy=strategy, mode=repr(mode), tl=_num(float(time_limit)),
              n=Ctx.py(tuple(grid_n)))


def label_code(ctx: Ctx, recipe: Dict, db, models_expr: Optional[str], *, mapped: bool, invariants=True,
               phase_labels=None, style=None, db_expr: str = "db", check_expr: Optional[str] = None) -> str:
    """Code that labels the phase regions of ``ax`` (a binary T-x diagram or ternary isotherm); check_expr: the
    mapping check report whose equilibrium grid is reused when it covers the axes."""
    from .region_labels import GRID_N, _db_elements, liquids_of
    rec = {k: val for k, val in recipe.items() if k != "tdb"}
    kind = rec["kind"]
    info = (f"ordered_info({db_expr}, sorted(recipe['elements']) + ['VA'], recipe['phases']"
            + (f", {models_expr})" if models_expr else ")")) if rec.get("rename_disordered", True) else "{}"
    if mapped:
        src = (f"regions = regions_from_lines(lines, invariants, recipe, {db_expr})   # exact two-phase fields / "
               "tie-triangles\n")
    else:
        src = ("# map the diagram again (pycalphad mapping) for the exact two-phase fields / tie-triangles\n"
               f"map_zpf, map_nodes = map_lines({db_expr}, recipe, label_info, recipe['composition_step'], "
               "recipe['temperature_step'])\n"
               f"regions = regions_from_lines(map_zpf, map_nodes, recipe, {db_expr})\n")
    kw = {"kind": kind, "elements": _db_elements(db), "liquids": liquids_of(db, rec["phases"])}
    if phase_labels:
        kw["phase_labels"] = dict(phase_labels)
    if not invariants:
        kw["invariants"] = False
    kw["T_unit"] = rec.get("T_unit", "K")
    if style:
        kw["style"] = dict(style)
    ctx.need("import time")
    args = ", ".join(f"{k}={Ctx.py(val)}" for k, val in kw.items())
    return _T('''
        # ---- phase-region labels: two-phase fields and tie-triangles from the mapped lines, the rest of the diagram
        # from an equilibrium grid; each label goes inside its region at the thickest free part (clear of data
        # points), smaller or turned if needed, left out below min_fontsize (8 pt); stoichiometric phases are named
        # outside the axes (covered tick labels hidden); every position is checked by an equilibrium
        recipe = <<recipe>>
        label_info = <<info>>
        <<src>>grid = <<reuse>>equilibrium_grid(<<db>>, recipe, label_info, ax.get_xlim(), ax.get_ylim(), <<n>>)
        t_labels, n_checks = time.time(), [0]


        def probe(x, y):
            """Stable phase set at a point (equilibrium), to check a label position."""
            n_checks[0] += 1
            return key_at(<<db>>, recipe, label_info, x, y)


        report = label_diagram(ax, regions, grid, probe, <<args>>)
        report.update(seconds=round(time.time() - t_labels, 1), equilibrium_checks=n_checks[0])
        for lab in report["labels"]:
            if lab.get("corner") is not None:      # the element at that corner
                lab["corner_element"] = recipe["elements"][lab["corner"]].capitalize()
            print("label:", lab["text"], "at", (round(lab["x"], 4), round(lab["y"], 4)), "gid", lab["gid"])
        for skip in report["skipped"]:
            print("not labelled:", skip["label"], "-", skip["reason"])
        ''', recipe=Ctx.py(rec, 9), info=info, src=src, db=db_expr,
              n=Ctx.py(tuple(GRID_N[kind])), args=args,
              reuse=(f"reuse_grid({check_expr}, ax.get_xlim(), ax.get_ylim(), {Ctx.py(tuple(GRID_N[kind]))}, "
                     f"{bool(rec.get('rename_disordered', True))}) or ") if check_expr else "")


_ISO_BODY = '''
    fixed = <<fixed>>   # fixed mole fractions
    axis_el, balance = <<axis>>, <<bal>>
    s_els = [e for e in strategy.elements if e != "VA"]   # order of cs.X


    def overall(x_axis):
        """Overall mole fractions at a point of the section."""
        d = {**fixed, axis_el: x_axis}
        d[balance] = 1.0 - sum(d.values())
        return d


    def temp(point):
        return float(np.squeeze(point.global_conditions[v.T]))


    def zpf_data():
        """The ZPF lines (along each line one phase - listed first - has zero phase fraction), their CSV rows, and the
        invariant equilibria (zero degrees of freedom) crossed by the section."""
        lines, rows = [], []
        for li, zpf in enumerate(strategy.zpf_lines):
            if not zpf.points:
                continue
            mid_names = names(zpf.points[len(zpf.points) // 2].stable_composition_sets)
            zero, region = mid_names[0], " + ".join(sorted(mid_names))
            xs, ts, tks, point_names = [], [], [], []
            for pi, p in enumerate(zpf.points):
                T = temp(p)
                point_names.append(names(p.stable_composition_sets))   # (the zero-fraction phase first)
                xv = float(np.squeeze(p.global_conditions[v.X(axis_el)]))
                xs.append(c_disp(overall(xv), axis_el))
                ts.append(T_disp(T))
                tks.append(T)
                row = {"line": li, "point": pi, "zero_phase": zero, "phases": region, "T_K": T, "T_C": T - KELVIN,
                       f"X({axis_el})": xv}
                for k, (name, cs) in enumerate(zip(names(p.stable_composition_sets), p.stable_composition_sets)):
                    X = dict(zip(s_els, np.asarray(cs.X, dtype=float)))
                    row[f"phase_{k + 1}"] = name
                    row[f"NP_{k + 1}"] = float(cs.NP)
                    row.update({f"X{k + 1}({e})": X[e] for e in elements})
                rows.append(row)
            lines.append({"zero_phase": zero, "x": xs, "T": ts, "T_K": tks, "names": point_names})
        invariants = []
        nodes = [n for n in strategy.node_queue.nodes
                 if degrees_of_freedom(n, strategy.components, strategy.num_potential_condition) == 0]
        try:
            segments = strategy.get_invariant_data(v.X(axis_el), v.T)
        except Exception:
            segments = [None] * len(nodes)
        for node, seg_data in zip(nodes, segments):
            node_phases = names(node.stable_composition_sets)
            T = temp(node)
            if any(abs(inv["T_K"] - T) < 0.05 and sorted(inv["phases"]) == sorted(node_phases) for inv in invariants):
                continue
            seg = sorted(float(np.squeeze(d.x)) for d in seg_data.data) if seg_data is not None else []
            invariants.append({"T_K": round(T, 3), "T_C": round(T - KELVIN, 3), "phases": node_phases,
                               "compositions": {n: {e: round(float(x), 6) for e, x in zip(s_els, np.asarray(cs.X))}
                                                for n, cs in zip(node_phases, node.stable_composition_sets)},
                               "segment": [round(s, 6) for s in seg]})
        invariants.sort(key=lambda inv: -inv["T_K"])
        return lines, rows, invariants


    lines, rows, invariants = zpf_data()
    fig, ax = plt.subplots(figsize=(7.2, 5.4))
    ax.set_xlabel(<<xlabel>>)
    ax.set_ylabel(<<ylabel>>)
    x_lo, x_hi = (float(np.squeeze(val)) for val in strategy.axis_lims[v.X(axis_el)])   # the mapped range
    ax.set_xlim(c_disp(overall(x_lo), axis_el), c_disp(overall(x_hi), axis_el))
    ax.set_ylim(<<ylim>>)
    <<finish>><<limits>>
    <<regions>>
    inv_rows = [{"invariant": k, "T_K": inv["T_K"], "T_C": inv["T_C"], "phases": " + ".join(inv["phases"]),
                 "phase": n, f"X({axis_el})_segment": " ".join(f"{s:g}" for s in inv["segment"]),
                 **{f"X({e})": inv["compositions"][n][e] for e in elements}}
                for k, inv in enumerate(invariants) for n in inv["phases"]]
    pd.DataFrame(rows).to_csv(OUT / <<zcsv>>, index=False, float_format="%.8g")
    pd.DataFrame(inv_rows).to_csv(OUT / <<icsv>>, index=False, float_format="%.8g")
    for inv in invariants:
        print(f"{' + '.join(inv['phases'])} at {inv['T_K']:.2f} K ({inv['T_C']:.2f} °C)")

    <<plot>>
    '''

_ISO_PLOT = '''
    # plot: each ZPF line coloured by its zero-fraction phase, invariants as red horizontal segments
    colors = <<colors>>
    for ln in lines:
        ax.plot(ln["x"], ln["T"], color=colors[ln["zero_phase"].split("#")[0]], lw=1.3, solid_capstyle="butt")
    for inv in invariants:
        s = [c_disp(overall(x), axis_el) for x in inv["segment"]]
        if len(s) >= 2:
            ax.plot([min(s), max(s)], [T_disp(inv["T_K"])] * 2, color="#d62728", lw=1.1, zorder=2.5)'''

_ISO_PLOT_SHADED = '''
    # plot (shaded style): black ZPF lines and invariant segments, no legend
    segments = [[c_disp(overall(x), axis_el) for x in inv["segment"]] for inv in invariants]
    draw_shaded(ax, [[np.column_stack([ln["x"], ln["T"]])] for ln in lines],
                [np.array([[min(s), T_disp(inv["T_K"])], [max(s), T_disp(inv["T_K"])]])
                 for inv, s in zip(invariants, segments) if len(s) >= 2], kind="isopleth")'''


_ISO_REGIONS = '''
    # ---- phase regions: the areas between the drawn lines, named by the ZPF lines around them (one side of a line has
    # its zero-fraction phase, the other not) and checked against an equilibrium grid (pdens 500; a second phase set
    # counts only if it is lower in Gibbs energy) - isopleth_regions; their outlines are snapped onto the lines
    recipe = <<recipe>>
    region_info = <<info>>
    t_regions = time.time()
    grid = equilibrium_grid(db, recipe, region_info, ax.get_xlim(), ax.get_ylim(), <<n>>, models=strategy.models,
                            phase_records=strategy.phase_records, pdens=<<pdens>>)


    def probe(x, y):
        """Stable phase set at a plotted point (equilibrium)."""
        return key_at(db, recipe, region_info, x, y, <<pdens>>)


    def resolve(x, y, keys):
        """Which of the phase sets is stable at a plotted point, by Gibbs energy (restricted equilibria)."""
        return lowest_key(db, recipe, region_info, x, y, keys)


    def find_regions():
        t0 = time.time()
        segs = [[c_disp(overall(x), axis_el) for x in inv["segment"]] for inv in invariants]
        out = isopleth_regions(ax, [{"xy": np.column_stack([ln["x"], ln["T"]]), "zero_phase": ln["zero_phase"],
                                     "names": ln["names"]} for ln in lines],
                               [{"T_K": inv["T_K"], "y": T_disp(inv["T_K"]), "x0": min(s), "x1": max(s),
                                 "phases": list(inv["phases"])} for inv, s in zip(invariants, segs) if len(s) >= 2],
                               grid, probe, resolve=resolve, info=region_info, check=<<check>>,
                               deadline=t_regions + <<budget>>)
        out["check"]["text"] = check_line(out)
        out["check"]["seconds"] = round(time.time() - t0, 2)
        return out


    def check_line(iso):
        return iso_check_text(iso["check"], lambda val: f"{val:.3g}", lambda val: f"{val:.0f} <<tu>>")


    iso = find_regions()<<retrace>>
    print(iso["check"]["text"])
    '''

_ISO_RETRACE = '''
    seeds = [sd for u in iso["check"]["unbounded"] for sd in u.get("seeds", [])]
    if seeds and run["finished"]:   # boundaries the mapping missed: short step maps through them seed new lines
        found = retrace_isopleth(strategy, recipe, seeds, time.time() + <<budget>>)
        if found["lines_added"]:
            lines, rows, invariants = zpf_data()
            before = iso["check"]["unbounded"]
            iso = find_regions()
            found["regions_before"] = [{k: val for k, val in u.items() if k != "seeds"} for u in before]
        iso["check"]["retraced"] = found
        iso["check"]["text"] = check_line(iso)'''

_ISO_LABELS = '''
    # ---- phase-region labels: inside each region at its thickest free part, at 8 pt or more with a 2.5 pt margin
    # (turned if that lets it be larger); a region too small for that is named right below its lowest point (60°),
    # else in the nearest free space with a leader line, else left out; invariant temperatures next to their line;
    # every position checked by an equilibrium (or, where that misses the stable set, by Gibbs energy)
    t_labels, n_checks = time.time(), [0]


    def check_probe(x, y):
        n_checks[0] += 1
        return probe(x, y)


    def confirm(key, got, x, y):
        n_checks[0] += 1
        return resolve(x, y, [key, got]) == tuple(key)


    report = label_diagram(ax, iso["regions"], None, check_probe, confirm=confirm, <<args>>)
    report.update(seconds=round(time.time() - t_labels, 1), equilibrium_checks=n_checks[0])
    for lab in report["labels"]:
        print("label:", lab["text"], "at", (round(lab["x"], 4), round(lab["y"], 4)), "gid", lab["gid"])
    for skip in report["skipped"]:
        print("not labelled:", skip["label"], "-", skip["reason"])
    '''


def isopleth_helpers(ctx: Ctx) -> str:
    """Sources of phase_regions (equilibrium grid, point equilibria) and isopleth_regions, once per script."""
    from . import isopleth_regions, phase_regions
    from .codegen_compare import module_source
    out = ""
    if not ctx.__dict__.get("_phase_regions"):
        ctx._phase_regions = True
        out += module_source(ctx, phase_regions, "phase regions: equilibrium grid, stable phase set at a point")
    if not ctx.__dict__.get("_mapping_check"):
        ctx._mapping_check = True
        from . import mapping_check
        out += module_source(ctx, mapping_check, "mapping check (here: retrace_isopleth, missing isopleth boundaries "
                                                 "traced from short step maps)")
    if not ctx.__dict__.get("_isopleth_regions"):
        ctx._isopleth_regions = True
        out += module_source(ctx, isopleth_regions, "isopleth regions: areas between the ZPF lines, named by the "
                                                    "lines, checked by equilibrium, outlines snapped to the lines")
    return out


def _isopleth_code(ctx, a, els, axis_el, fixed, png, png_name, colors, tC, cu, tlo, thi, system,
                   shaded=False, *, db=None, ph=None, P=101325.0, xr=(0.0, 1.0), step=0.02, masses=None,
                   check="auto", label_regions=False) -> str:
    from . import mapping_tools as mt
    from .mapping_check import GRID_N, PDENS
    ctx.need("from pycalphad.mapping.utils import degrees_of_freedom")
    balance = [e for e in els if e != axis_el and e not in fixed][0]
    names_ = _register_siblings(ctx, png, ["_zpf_lines.csv", "_invariants.csv"])
    limits = ""
    if a.get("xlim"):
        limits += f"\nax.set_xlim({_num(a['xlim'][0])}, {_num(a['xlim'][1])})"
    if a.get("ylim"):
        limits += f"\nax.set_ylim({_num(a['ylim'][0])}, {_num(a['ylim'][1])})"
    plot = _T(_ISO_PLOT_SHADED) if shaded else _T(_ISO_PLOT, colors=Ctx.py(colors, 9))
    regions = ""
    if shaded or label_regions:     # the regions (fill and labels): as map_phase_diagram finds them
        ctx.need("import time")
        unit = a["temperature_unit"]
        from pycalphad import variables as v
        conds = {v.T: (tlo, thi, float(a["temperature_step"]))}
        rec = mt._recipe(ctx.tb._find_tdb(a["tdb"]), "isopleth", els, ph, P, conds, mt._Units(unit, cu, masses),
                         unit, a["rename_disordered"], xr, step, float(a["temperature_step"]))
        rec.update(axis_element=axis_el, balance=balance, fixed=dict(fixed))
        rec = {k: val for k, val in rec.items() if k != "tdb"}
        info = ("ordered_info(db, sorted(recipe['elements']) + ['VA'], recipe['phases'], strategy.models)"
                if a["rename_disordered"] else "{}")
        tl = float(a.get("time_limit_s") or 600.0)
        regions = _T(_ISO_REGIONS, recipe=Ctx.py(rec, 9), info=info, n=Ctx.py(tuple(GRID_N["isopleth"])), pdens=PDENS,
                     check=repr(check != "off"), budget=_num(max(60.0, 0.25 * tl)), tu="°C" if tC else "K",
                     retrace=("\n" + _T(_ISO_RETRACE, budget=_num(max(30.0, 0.1 * tl))).rstrip("\n"))
                     if check != "off" else "")
        helpers = (label_helpers(ctx) if label_regions else "") + isopleth_helpers(ctx)
    else:
        helpers = ""
    plot = (plot + ("\nfill_regions(ax, iso[\"regions\"][\"polygons\"])   # even numbers of phases (2, 4) aliceblue, "
                    "the rest white" if shaded and regions else ""))
    body = _T(_ISO_BODY, fixed=Ctx.py(fixed), axis=repr(axis_el), bal=repr(balance),
              zcsv=repr(names_["_zpf_lines.csv"]), icsv=repr(names_["_invariants.csv"]),
              plot=plot.strip("\n"),
              xlabel=repr(_comp_label(axis_el, cu)),
              ylabel=repr(_T_label(tC)),
              ylim=f"{_num(tlo - KELVIN if tC else tlo)}, {_num(thi - KELVIN if tC else thi)}",
              finish="ticks_in(ax)" if shaded else "ax.grid(alpha=0.25)", limits=limits,
              regions=regions.strip("\n"))
    body = helpers + textwrap.dedent(body) + "\n"
    weight = cu.startswith("weight")
    title = a.get("title") or (f"{system} isopleth, " + ", ".join(
        f"{'w' if weight else 'x'}({k.capitalize()})={val:g}" for k, val in fixed.items())
        + (" (mole fractions)" if weight else ""))
    legend = _T(_LEGEND, lines="lines", ln_phases='[ln["zero_phase"]]', invs="invariants",
                ltitle=', title="ZPF line of", title_fontsize=8') if not shaded else ""
    fmts = a.get("formats")
    if fmts is None:
        fmts = ["svg", "pdf"] if shaded or label_regions else []
    if shaded and not label_regions:      # no legend and no region labels: each ZPF line is named on the line
        legend = (_phase_names_source(ctx) + f"line_labels(ax, [(format_phase(ln['zero_phase'], {list(els)!r}, "
                  f"{a.get('phase_labels')!r}), np.column_stack([ln['x'], ln['T']])) for ln in lines])\n")
    labels = ""
    if label_regions:
        from .region_labels import _db_elements, liquids_of
        kw = {"kind": "isopleth", "elements": _db_elements(db), "liquids": liquids_of(db, ph)}
        if a.get("phase_labels"):
            kw["phase_labels"] = dict(a["phase_labels"])
        if not a.get("label_invariants", True):
            kw["invariants"] = False
        kw["T_unit"] = "C" if tC else "K"
        st = _label_style(a.get("label_style"), shaded)
        if st:
            kw["style"] = dict(st)
        labels = _T(_ISO_LABELS, args=", ".join(f"{k}={Ctx.py(val)}" for k, val in kw.items()))
        if not shaded:
            labels += f"prettify_legend(ax, {list(els)!r}, {a.get('phase_labels')!r})\n"
    return body + legend + (f"ax.set_title({title!r}, fontsize=11)\n" + labels +
                                      f"fig.savefig(OUT / {png_name!r}, dpi=150, bbox_inches='tight')\n"
                                      + _vector_saves(ctx, png, fmts))


# ---------------------------------------------------------------------------------------------------------
# condition grids (equilibrium_property_plot / export_equilibrium_data)
# ---------------------------------------------------------------------------------------------------------
def _cond_values_code(val, lab: str, celsius: bool):
    """(code expression, numpy values) of a condition value as the tools expand it (props._values)."""
    import numpy as np
    from .props import _values
    arr = _values(val, lab)
    if isinstance(val, dict):
        if {"start", "stop"} <= set(val):
            val = [val["start"], val["stop"], val.get("step", (val["stop"] - val["start"]) / 20.0)]
        else:
            val = list(val["values"])
    if isinstance(val, (list, tuple)):
        vals = [float(x) for x in val]
        if len(vals) == 3:
            a, b, st = vals
            code = f"{_num(a)} + {_num(st)} * np.arange({len(arr)})"
            if celsius:
                code = f"({code}) + KELVIN"
            note = f"{_num(a)} .. {_num(b)} in steps of {_num(st)}"
        else:
            code = f"np.array({Ctx.py(vals)})" + (" + KELVIN" if celsius else "")
            note = ""
    else:
        code = _num(val) + (" + KELVIN" if celsius else "")
        note = ""
    return code, (arr + KELVIN if celsius else arr), note


def _conditions(ctx: Ctx, db, els, conditions: Dict, unit: str):
    """Generation-time parse of a conditions dict -> (list of (label, prop code, value code, values, note),
    note about W->X conversion).  Mirrors props.parse_conditions + props._w_to_x."""
    import numpy as np
    from . import props
    conds, labels, varying, values = props.parse_conditions(conditions, els, unit)
    wnote = props._w_to_x(db, els, conds, labels, values)
    raw = {}
    for k, val in conditions.items():
        ku = str(k).strip().upper().replace(" ", "")
        lab = {"T_K": "T", "T_C": "T", "TEMPERATURE": "T", "PRESSURE": "P", "P_PA": "P"}.get(ku, ku)
        if ku in els:
            lab = f"X({ku})"
        raw[lab] = (val, ku == "T_C")
    out = []
    for lab in labels:
        m = re.fullmatch(r"(X|W|MU)\(([A-Z][A-Z]?)\)", lab)
        if lab in ("T", "P", "N"):
            pcode = f"v.{lab}"
        elif m:
            pcode = f"v.{m.group(1)}({m.group(2)!r})"
        else:
            ctx.fallback(f"condition {lab} has no simple pycalphad variable")
        if lab in raw:
            val, tc = raw[lab]
            code, arr, note = _cond_values_code(val, lab, lab == "T" and (_is_celsius(unit) or tc))
        else:                      # added defaults, or X(EL) converted from a mass fraction
            arr = np.asarray(values[lab], dtype=float)
            code, note = _num(arr[0]), ""
        if not np.allclose(arr, values[lab], rtol=0, atol=0):
            ctx.fallback(f"could not reproduce the values of condition {lab}")
        out.append((lab, pcode, code, np.asarray(values[lab]), note))
    return out, wnote


def _var(lab: str) -> str:
    """Python variable name for a condition label: T, P, N, X_CR, MU_CR."""
    return re.sub(r"\W", "", lab.replace("(", "_"))


def _prop_str(lab: str) -> str:
    """str(pycalphad variable) of a condition label (sort key of pycalphad's condition axes)."""
    m = re.fullmatch(r"(X|W|MU)\(([A-Z][A-Z]?)\)", lab)
    return f"{m.group(1)}_{m.group(2)}" if m else lab


_EQP_HELPERS = '''
def pretty_units(u):
    """pint unit string -> short label ('J / mol' -> 'J/mol'; fractions -> '')."""
    s = str(u or "")
    for a, b in {"kelvin": "K", "pascal": "Pa", "fraction": "", "dimensionless": "", "joule": "J", "mole": "mol",
                 "gram": "g", "kilojoule": "kJ", "degree_Celsius": "°C"}.items():
        s = re.sub(rf"\\b{a}\\b", b, s)
    return s.replace(" / ", "/").replace(" * ", "·").strip()


def label_of(p):
    """Column label of a computed property, e.g. NP(FCC_A1), X(FCC_A1,CR), MU(CR), HM.T."""
    if isinstance(p, v.PhaseFraction):
        return f"NP({p.phase_name})"
    if isinstance(p, v.MoleFraction):
        return f"X({p.phase_name},{p.species.name})" if p.phase_name else f"X({p.species.name})"
    if isinstance(p, v.MassFraction):
        return f"W({p.phase_name},{p.species.name})" if getattr(p, "phase_name", None) else f"W({p.species.name})"
    if isinstance(p, v.ChemicalPotential):
        return f"MU({p.species.name})"
    if isinstance(p, v.SiteFraction):
        return f"Y({p.phase_name},{p.sublattice_index},{p.species.name})"
    if isinstance(p, JanssonDerivative):
        return f"{label_of(p.numerator)}.{label_of(p.denominator)}"
    return str(p)
'''

_GUARDED = '''

class Guarded:
    """Meta-property wrapper: NaN instead of an exception at points where the dormant/isolated solve fails."""

    def __init__(self, inner):
        self.inner = inner
        self.shape = getattr(inner, "shape", tuple())
        self.implementation_units, self.display_units = inner.implementation_units, inner.display_units
        self.display_name = getattr(inner, "display_name", "")

    def compute_property(self, compsets, cur_conds, chemical_potentials):
        if len(compsets) == 0:
            return np.nan
        try:
            return self.inner.compute_property(compsets, cur_conds, chemical_potentials)
        except Exception:
            return np.nan
'''

_PURE_REF = '''

def pure_element_gm(el, phase, T, P):
    """Gibbs energy of pure `el` in `phase` (lowest-energy internal configuration) at each (T, P)."""
    T, P = np.ravel(T), np.broadcast_to(np.ravel(P), np.ravel(T).shape)
    out = np.full(T.shape, np.nan)
    for p in np.unique(P):
        sel = P == p
        tu, inv = np.unique(T[sel], return_inverse=True)
        r = calculate(db, [el, "VA"], phase, T=tu, P=p, N=1, output="GM", pdens=50)
        out[sel] = np.nanmin(r.GM.values.reshape(len(tu), -1), axis=1)[inv]
    return out
'''

_ASSEMBLAGES = '''

def assemblage(phase_names):
    """'A+B' label of the stable phases at one grid point (second composition set of A -> A#1+A#2)."""
    cnt = Counter(str(p) for p in phase_names if str(p) and str(p) != "_FAKE_")
    parts = []
    for n in sorted(cnt):
        parts += [n] if cnt[n] == 1 else [f"{n}#{i + 1}" for i in range(cnt[n])]
    return "+".join(parts) if parts else "(none)"


eq_phase = np.asarray(wks.eq.Phase).reshape(shape + (-1,))
stable = np.empty(shape, dtype=object)
for idx in np.ndindex(*shape):
    stable[idx] = assemblage(eq_phase[idx]) if converged[idx] else "(failed)"
'''

_XSUM = '''
bad = np.sum([<<xgrids>>], axis=0) > 1 + 1e-9      # outside the composition triangle
stable[bad] = "(X sum > 1)"
'''

_MERGE_DISORDERED = '''

# An ordered phase (order/disorder model) that only occurs with degree of order 0 is the same solution as the
# disordered phase: rename it in the phase-region labels and merge its columns into the disordered phase's columns.
ORDERED = <<ordered>>   # ordered phase -> equivalent disordered phase
present = {re.sub(r"#\\d+$", "", n) for a in np.ravel(stable) for n in str(a).split("+")}
renamed = {}
for p, q in ORDERED.items():
    if p in present:
        try:
            doo = wks.get_dict(f"DOO({p})")
        except Exception:
            continue
        if max(float(np.nanmax(np.abs(np.asarray(a, dtype=float)))) for a in doo.values()) < 1e-3:
            renamed[p] = q


def rename(a):
    names = [renamed.get(n, n) for n in (re.sub(r"#\\d+$", "", n) for n in str(a).split("+"))]
    cnt = Counter(names)
    return "+".join(n if cnt[n] == 1 else f"{n}#{i + 1}" for n in sorted(cnt) for i in range(cnt[n]))


if renamed:
    stable = np.vectorize(rename, otypes=[object])(stable) if shape else np.array(rename(stable), dtype=object)
    for p, q in renamed.items():
        if q.endswith("(disordered)"):
            continue
        for lab in list(columns):
            lab2 = re.sub(rf"(?<=[(,]){re.escape(p)}(?=[,)#])", q, lab)
            if lab2 == lab or lab2 not in columns:
                continue
            a, b = np.asarray(columns[lab2], dtype=float), np.asarray(columns[lab], dtype=float)
            if lab.startswith("NP("):
                columns[lab2] = np.where(np.isfinite(a) | np.isfinite(b), np.nan_to_num(a) + np.nan_to_num(b), np.nan)
            else:
                columns[lab2] = np.where(np.isfinite(a), a, b)
            del columns[lab], units[lab]
'''

_ORDER_SETS = '''

# Composition sets of a miscibility gap (PH#1, PH#2) are numbered by the solver in any order: renumber them at every
# grid point by increasing X(<<ref>>) (#1 = leanest) in every column that refers to them.
multi = {}
for lab in columns:
    for ph, n in re.findall(r"([A-Z0-9_]+)#(\\d+)", lab):
        multi.setdefault(ph, set()).add(int(n))
for ph, ns in ((ph, sorted(ns)) for ph, ns in multi.items() if len(ns) > 1):
    xs = {int(re.search(rf"{ph}#(\\d+)", label_of(k)).group(1)): np.asarray(a, dtype=float).reshape(shape)
          for k, a in wks.get_dict(f"X({ph},<<ref>>)").items() if re.search(rf"{ph}#(\\d+)", label_of(k))}
    if sorted(xs) != ns:
        continue
    order = np.argsort(np.stack([np.where(np.isfinite(xs[n]), xs[n], np.inf) for n in ns]), axis=0, kind="stable")
    groups = {}
    for lab in list(columns):
        m = re.search(rf"{ph}#(\\d+)", lab)
        if m and int(m.group(1)) in ns:
            groups.setdefault(lab.replace(m.group(0), f"{ph}#{{n}}"), {})[int(m.group(1))] = lab
    for members in groups.values():
        if sorted(members) == ns:
            new = np.take_along_axis(np.stack([columns[members[n]] for n in ns]), order, axis=0)
            for k, n in enumerate(ns):
                columns[members[n]] = new[k]
'''


def _quantity_code(ctx: Ctx, db, els, q_raw: str):
    """(label or None, property code, kind, extra) for one quantity string (mirrors props.parse_quantities)."""
    from . import props
    body, units = props._split_units(str(q_raw))
    q = props._normalize_q(body)
    q = {"NP": "NP(*)", "PHASE_FRACTIONS": "NP(*)", "CP": "CPM", "H": "HM", "G": "GM", "S": "SM",
         "DOO": "DOO(*)"}.get(q, q)
    ucode = f"[{units!r}]" if units else ""
    m_meta = re.fullmatch(r"(ISOLATED|DORMANT)\(([A-Z0-9_]+)\):(.+)", q)
    m_df = re.fullmatch(r"(?:DF|DRIVINGFORCE|DGM)\(([A-Z0-9_]+)\)", q)
    m_t0 = re.fullmatch(r"T0\(([A-Z0-9_]+),([A-Z0-9_]+)\)", q)
    m_acr = re.fullmatch(r"(ACR|MUR|AC)\(([A-Z][A-Z]?)(?:,([A-Z0-9_]+))?\)", q)
    if m_acr:
        kind = "mur" if m_acr.group(1) == "MUR" else "acr"
        return None, f"as_property('MU({m_acr.group(2)})')", kind, (m_acr.group(2), m_acr.group(3), units)
    if m_df:
        ph = m_df.group(1)
        ctx.need("from pycalphad.property_framework import DormantPhase")
        if units:
            ctx.need("from pycalphad.property_framework.metaproperties import DrivingForce")
            return f"DF({ph})", (f"Guarded(DormantPhase({ph!r}, wks)(driving_force_in({ph!r}, {units!r})))"), \
                "meta", None
        return f"DF({ph})", f"Guarded(DormantPhase({ph!r}, wks).driving_force)", "meta", None
    if m_meta:
        which, ph, inner = m_meta.groups()
        inner_s = props._phase_prop(inner, ph)
        ctx.need(f"from pycalphad.property_framework import {'IsolatedPhase' if which == 'ISOLATED' else 'DormantPhase'}")
        from pycalphad.property_framework import as_property
        lab = f"{which}({ph}):{props._label_of(as_property(inner_s))}"
        cls = "IsolatedPhase" if which == "ISOLATED" else "DormantPhase"
        return lab, f"Guarded({cls}({ph!r}, wks)(as_property({inner_s!r}){ucode}))", "meta", None
    if m_t0:
        p1, p2 = m_t0.groups()
        ctx.need("from pycalphad.property_framework import T0",
                 "from pycalphad.property_framework.metaproperties import find_first_compset")
        return f"T0({p1},{p2})", (f"Guarded(T0(find_first_compset({p1!r}, wks), find_first_compset({p2!r}, wks), "
                                  "None))"), "meta", None
    if q.startswith("Y(") and "*" in q:
        ctx.fallback("wildcard site fractions Y(...*...) are expanded against the toolkit's phase records")
    from pycalphad.property_framework import as_property
    p = as_property(q)
    if units:
        try:
            p[units]
        except Exception:  # noqa: BLE001
            ctx.fallback(f"units {units} for {q}")
    return None, f"as_property({q!r}){ucode}", "pc", ("*" in q, units)


def _qtext(q_raw: str) -> str:
    """Normalized quantity text as the tool keeps it (aliases expanded, units suffix kept)."""
    from . import props
    body, units = props._split_units(str(q_raw))
    q = props._normalize_q(body)
    q = {"NP": "NP(*)", "PHASE_FRACTIONS": "NP(*)", "CP": "CPM", "H": "HM", "G": "GM", "S": "SM",
         "DOO": "DOO(*)"}.get(q, q)
    return q if not units else f"{q}[{units}]"


def _ref_phase(db, el, ref):
    """Reference phase for ACR/MUR as props._pure_ref chooses it."""
    import numpy as np
    from . import props
    if ref:
        return ref
    return props._pure_ref(db, el, None, np.array([298.15]), np.array([101325.0]))[1]


# ---------------------------------------------------------------------------------------------------------
# equilibrium_property_plot
# ---------------------------------------------------------------------------------------------------------
def _csv_header(path: str) -> Optional[List[str]]:
    import csv
    try:
        with open(path, newline="") as fh:
            return next(csv.reader(fh))
    except (OSError, StopIteration):
        return None


def _split_col(col: str):
    m = re.fullmatch(r"(.*) \[([^\]]*)\]", col)
    return (m.group(1), m.group(2)) if m else (col, "")


@generator("equilibrium_property_plot")
def _gen_equilibrium_property_plot(args, ctx: Ctx) -> str:  # noqa: C901 - one linear recipe
    from . import props
    a = {"temperature_unit": "K", "plot": True, "write_netcdf": False, **args}
    unit = a["temperature_unit"] or "K"
    celsius = _is_celsius(unit)
    db = _db(ctx, a["tdb"])
    els = props._check_elements(db, a["elements"])
    ph = props._resolve_phases(db, els, a.get("phases"), a.get("exclude_phases"))
    qs = [a["quantities"]] if isinstance(a["quantities"], str) else list(a["quantities"])
    conds, wnote = _conditions(ctx, db, els, a["conditions"], unit)
    varying = [c[0] for c in conds if len(c[3]) > 1]
    axis_labels = sorted(varying, key=_prop_str)
    shape = tuple(len(c[3]) for lab in axis_labels for c in conds if c[0] == lab)

    # output files as the tool named them
    els_s = sorted(e.upper() for e in a["elements"])
    stem = props._safe_name(f"eqprop_{'-'.join(els_s)}_{'_'.join(props._normalize_q(q) for q in qs)[:40]}")
    base = _outp(ctx.tb, a.get("output_name"), stem)
    if Path(base).suffix.lower() in (".png", ".csv", ".nc"):
        base = str(Path(base).with_suffix(""))
    header = _csv_header(base + ".csv")
    if header is None:
        ctx.fallback("the tool's CSV (used to lay out the same columns and panels) is no longer available")
    csv_name = ctx.out_name(base + ".csv", "eqprop.csv")
    png_name = ctx.out_name(base + ".png", "eqprop.png") if a["plot"] and len(axis_labels) > 0 else None

    ctx.need("import re", "from collections import Counter", "import numpy as np", "import pandas as pd",
             "from pycalphad import Database", "from pycalphad import variables as v",
             "from pycalphad.core.workspace import Workspace",
             "from pycalphad.property_framework import as_property, JanssonDerivative")
    code = [_load_db_code(ctx, a["tdb"]),
            f"elements = {Ctx.py(els)}\n",
            f"comps = {Ctx.py(props._comps(db, els))}\n",
            f"phases = {Ctx.py(ph)}" + ("   # all phases that can form\n" if not (a.get("phases") or
                                                                                   a.get("exclude_phases")) else "\n"),
            "KELVIN = 273.15\n\n"]
    code.append("# equilibrium conditions (T in K" + (", given in °C" if celsius else "") + ")"
                + (f"; {wnote}" if wnote else "") + "\n")
    names = {}
    for lab, pcode, vcode, arr, note in conds:
        var = _var(lab)
        names[lab] = var
        code.append(f"{var} = {vcode}" + (f"   # {note}" if note else "") + "\n")
    code.append("conds = {" + ", ".join(f"{pcode}: {names[lab]}" for lab, pcode, *_ in conds) + "}\n\n")
    code.append("# solve all equilibria once; properties are then evaluated from the stored solution\n"
                "wks = Workspace(database=db, components=comps, phases=phases, conditions=conds)\n")
    if axis_labels:
        code.append(f"axes = {Ctx.py(axis_labels)}   # varying conditions = dimensions of every result array "
                    "(pycalphad sorts them)\n"
                    f"assert [str(k) for k in wks.condition_axis_order] == {Ctx.py([_prop_str(x) for x in axis_labels])}\n"
                    f"grid = dict(zip(axes, np.meshgrid({', '.join('np.atleast_1d(' + names[x] + ')' for x in axis_labels)}, "
                    "indexing='ij')))\n"
                    "shape = grid[axes[0]].shape\n")
    else:
        code.append("axes, grid, shape = [], {}, ()\n")
    code.append("converged = np.isfinite(np.asarray(wks.eq.GM, dtype=float)).reshape(shape)\n")

    # quantities
    qinfo = [(_qtext(q),) + _quantity_code(ctx, db, els, q) for q in qs]
    code.append("\n" + _T(_EQP_HELPERS))
    if any(k == "meta" for _, _, _, k, _ in qinfo):
        code.append(_T(_GUARDED))
    if any("driving_force_in" in pc for _, _, pc, _, _ in qinfo):
        code.append("\n\ndef driving_force_in(phase, units):\n    df = DrivingForce(phase)\n"
                    "    df.display_units = units\n    return df\n")
    refs = {}
    for q, lab, pc, kind, extra in qinfo:
        if kind in ("acr", "mur"):
            el, ref, _ = extra
            refs[q] = _ref_phase(db, el, ref)
    if refs:
        ctx.need("from pycalphad import calculate")
        code.append(_T(_PURE_REF).replace('[el, "VA"]', "[el, 'VA']" if "VA" in {str(e) for e in db.elements}
                                          else "[el]"))
    code.append("\n\n# quantities: (column label, or None = label of each computed property; property; quantity)\n"
                "quantities = [\n")
    for q, lab, pc, kind, extra in qinfo:
        code.append(f"    ({lab!r}, {pc}, {q!r}),\n")
    code.append("]\n")
    wild_units = any(kind == "pc" and extra[0] and extra[1] for q, lab, pc, kind, extra in qinfo)
    if wild_units:
        ctx.need("from pycalphad.property_framework.units import Q_")
    Tgrid = "grid['T']" if "T" in axis_labels else "np.full(shape, float(np.ravel(T)[0]))"
    Pgrid = "grid['P']" if "P" in axis_labels else "np.full(shape, float(np.ravel(P)[0]))"
    loop = [
        "columns, units = {}, {}",
        "for label, prop, q in quantities:",
        "    wildcard = '*' in q",
        "    for key, arr in wks.get_dict(prop).items():",
        "        arr = np.asarray(arr, dtype=float).reshape(shape)",
        "        lab, u = label or label_of(key), pretty_units(getattr(key, 'display_units', ''))",
    ]
    if wild_units:
        loop += ["        want = re.search(r'\\[(.+)\\]$', q)",
                 "        if want and wildcard and u != want.group(1):   # wildcard results come in default units",
                 "            arr, u = Q_(arr, getattr(key, 'display_units', '')).to(want.group(1)).magnitude, "
                 "want.group(1)"]
    if refs:
        loop += ["        if q in REFERENCE:   # ACR / MUR: relative to the pure element in its reference phase",
                 "            el, ref_phase, what = REFERENCE[q]",
                 f"            g_ref = pure_element_gm(el, ref_phase, {Tgrid}, {Pgrid}).reshape(arr.shape)",
                 "            if what == 'mur':",
                 "                lab, u, arr = f'MUR({el},{ref_phase})', 'J/mol', arr - g_ref",
                 "            else:",
                 f"                lab, u = f'ACR({{el}},{{ref_phase}})', ''",
                 f"                arr = np.exp((arr - g_ref) / (8.314462618 * {Tgrid}.reshape(arr.shape)))"]
        code.append("REFERENCE = {" + ", ".join(
            f"{q!r}: ({extra[0]!r}, {refs[q]!r}, {kind!r})" for q, lab, pc, kind, extra in qinfo
            if kind in ("acr", "mur")) + "}   # quantity: (element, reference phase, kind)\n")
    loop += [
        "        if lab.startswith('NP('):",
        "            arr = np.where(converged & ~np.isfinite(arr), 0.0, arr)   # phase absent at a solved point: 0",
        "        if wildcard and (np.all(~np.isfinite(arr)) or (lab.startswith('NP(') and np.nanmax(np.abs(arr)) < 1e-12)):",
        "            continue   # wildcard expansions never present anywhere are dropped",
        "        if lab.startswith('T0(') or lab == 'T':",
        "            u, arr = " + ("'°C', arr - KELVIN" if celsius else "'K', arr"),
        "        if lab not in columns:",
        "            columns[lab], units[lab] = arr, u",
    ]
    code.append("\n".join(loop) + "\n")

    # stable phase assemblages and the tool's post-processing
    code.append(_T(_ASSEMBLAGES))
    xgr = [f"grid[{x!r}]" for x in axis_labels if x.startswith("X(")]
    x_conds = [c for c in conds if c[0].startswith("X(")]
    if xgr and len(x_conds) >= 2:
        allx = [f"grid[{c[0]!r}]" if c[0] in axis_labels else f"np.full(shape, {names[c[0]]})" for c in x_conds]
        code.append(_T(_XSUM, xgrids=", ".join(allx)))
    ordered = {}
    for p in ph:
        if db.phases[p].model_hints.get("ordered_phase") == p:
            tok = set(p.split("_")) & {"FCC", "BCC", "HCP"}
            plain = [q for q in ph if not db.phases[q].model_hints.get("ordered_phase")]
            cands = [q for q in plain if tok and tok <= set(q.split("_"))]
            ordered[p] = cands[0] if len(cands) == 1 else f"{p}(disordered)"
    if ordered:
        code.append(_T(_MERGE_DISORDERED, ordered=Ctx.py(ordered)))
    if any(re.search(r"#\d", c) for c in header) and shape:
        ref_el = next((c[0][2:-1] for c in conds if c[0].startswith("X(")), sorted(els)[0])
        code.append(_T(_ORDER_SETS, ref=ref_el))

    # CSV
    frame = ["", "# CSV: varying conditions, one column per quantity (with units), stable phases",
             "table = {}"]
    for lab in axis_labels:
        if lab == "T":
            frame.append("table['T [K]'] = grid['T'].ravel()")
            if celsius:
                frame.append("table['T [°C]'] = grid['T'].ravel() - KELVIN")
        else:
            frame.append(f"table[{lab + (' [Pa]' if lab == 'P' else '')!r}] = grid[{lab!r}].ravel()")
    frame += ["for lab, arr in columns.items():",
              "    table[f'{lab} [{units[lab]}]' if units[lab] else lab] = np.ravel(arr)",
              "table['stable_phases'] = np.ravel(stable)",
              f"df = pd.DataFrame(table)",
              f"df.to_csv(OUT / {csv_name!r}, index=False, float_format='%.8g')",
              "print(df)"]
    code.append("\n".join(frame) + "\n")
    expect = [c for c in header]
    code.append(f"if list(df.columns) != {Ctx.py(expect, 20)}:\n"
                "    print('note: the columns differ from those of the original tool run')\n")

    # plot
    sysname = "-".join(e.capitalize() for e in els)
    fixed = ", ".join(
        f"{lab}={props._fmt(props._cond_display(lab, arr, unit)[0])}"
        f"{' °C' if lab == 'T' and celsius else (' K' if lab == 'T' else '')}"
        for lab, _, _, arr, _ in conds if len(arr) == 1 and lab not in ("N", "P"))
    title = a.get("title") or (sysname + (f"   {fixed}" if fixed else ""))
    cols = [_split_col(c) for c in header if c not in ("stable_phases",)]
    cond_cols = len(axis_labels) + (1 if "T" in axis_labels and celsius else 0)
    qcols = cols[cond_cols:]
    if not a["plot"] or not axis_labels:
        if not axis_labels:
            code.append("for lab, arr in columns.items():\n    print(lab, float(arr), units[lab])\n"
                        "print('stable phases:', stable)\n")
        return "".join(code)
    ctx.need("import matplotlib.pyplot as plt")
    if len(axis_labels) == 1:
        code.append(_plot_1d_code(a, axis_labels[0], unit, qcols, title, png_name))
    else:
        code.append(_plot_2d_code(a, axis_labels, unit, qcols, title, png_name))
    if a.get("write_netcdf"):
        ctx.note("equilibrium_property_plot: the netCDF copy of the CSV is not written by the standalone code")
    return "".join(code)


def _plot_1d_code(a, clab, unit, qcols, title, png_name) -> str:
    from . import props
    x_lab = clab
    if a.get("x_axis"):
        xa = props._normalize_q(props._split_units(a["x_axis"])[0])
        x_lab = {"TEMPERATURE": "T"}.get(xa, xa)
    groups = {}
    for lab, u in qcols:
        if lab == x_lab:
            continue
        key, ylab = props._family(lab, u)
        groups.setdefault(key, (ylab, []))[1].append((lab, u))
    keys = list(groups)[:8]
    panels = []
    for key in keys:
        ylab, labs = groups[key]
        if key.startswith("u:") and len(labs) == 1:
            ylab = labs[0][0] + (f" ({labs[0][1]})" if labs[0][1] else "")
        elif key.startswith("u:"):
            ylab = f"{', '.join(x for x, _ in labs[:3])}{'…' if len(labs) > 3 else ''}" + \
                   (f" ({labs[0][1]})" if labs[0][1] else "")
        kind = key if key in ("NP", "ACR", "DF") else ""
        panels.append((ylab, [(lab, props._short_label(lab, key)) for lab, _ in labs], kind))
    if x_lab == clab:
        xexpr = "grid['T'] - KELVIN" if (clab == "T" and _is_celsius(unit)) else f"grid[{clab!r}]"
        xlabel = props._cond_axis_label(clab, unit)
    else:
        xu = dict(qcols).get(x_lab, "")
        xexpr = f"columns[{x_lab!r}]"
        xlabel = x_lab + (f" ({xu})" if xu else "")
    return _T('''

        # plot: one panel per kind of quantity (y label, [(column, legend label)], kind)
        panels = <<panels>>
        x = <<xexpr>>
        fig, axs = plt.subplots(len(panels), 1, figsize=(7.2, 1.3 + 2.7 * len(panels)), sharex=True, squeeze=False)
        for ax, (ylabel, curves, kind) in zip(axs[:, 0], panels):
            cmap = plt.get_cmap("tab20" if len(curves) > 10 else "tab10")
            for i, (lab, name) in enumerate(curves):
                ax.plot(x, columns[lab], lw=1.8, color=cmap(i % cmap.N), label=name, marker="o" if len(x) < 12 else None,
                        ms=3)
            if kind == "NP":
                ax.set_ylim(-0.02, 1.02)
            elif kind == "ACR":
                ax.set_ylim(bottom=0)
            elif kind == "DF":
                ax.axhline(0, color="k", lw=0.8, ls="--")
            ax.set_ylabel(ylabel)
            ax.grid(alpha=0.3)
            ax.legend(fontsize=8, ncol=2 if len(curves) > 10 else 1, loc="best", framealpha=0.8)<<vlines>>
        axs[-1, 0].set_xlabel(<<xlabel>>)
        axs[0, 0].set_title(<<title>>, fontsize=10)
        fig.tight_layout()
        fig.savefig(OUT / <<png>>, dpi=150, bbox_inches="tight")
        ax = axs[0, 0]
        ''', panels=Ctx.py(panels, 9), xexpr=xexpr, xlabel=repr(xlabel), title=repr(title), png=repr(png_name),
              vlines=("" if x_lab != clab else _indent(
                  "for i in range(1, len(x)):   # dotted lines where the stable phases change\n"
                  "    if stable[i] != stable[i - 1]:\n"
                  "        ax.axvline(0.5 * (x[i] + x[i - 1]), color='0.6', lw=0.6, ls=':')", 4)))


def _plot_2d_code(a, axis_labels, unit, qcols, title, png_name) -> str:
    from . import props
    a_lab, b_lab = axis_labels
    if a.get("x_axis"):
        xa = props._normalize_q(props._split_units(a["x_axis"])[0])
        x_lab = {"TEMPERATURE": "T"}.get(xa, xa)
    else:
        x_lab = a_lab if not a_lab.startswith(("T", "P")) or b_lab in ("T", "P") else b_lab
        if a_lab in ("T", "P") and b_lab not in ("T", "P"):
            x_lab = b_lab
    y_lab = b_lab if x_lab == a_lab else a_lab
    labs = [lab for lab, _ in qcols][:props.MAX_PANELS - 1]
    units = {lab: u for lab, u in qcols}

    def disp(lab):
        return f"np.atleast_1d({_var(lab)})" + (" - KELVIN" if lab == "T" and _is_celsius(unit) else "")
    return _T('''

        # plot: map of the stable phase regions + one map per quantity, with the region boundaries
        xv, yv = <<xv>>, <<yv>>
        orient = (lambda arr: np.asarray(arr).T) if <<transpose>> else (lambda arr: np.asarray(arr))   # -> (ny, nx)
        regions = orient(stable)
        region_names = sorted({r for r in np.ravel(regions) if not str(r).startswith("(")})
        ids = np.array([region_names.index(r) if r in region_names else np.nan for r in np.ravel(regions)],
                       dtype=float).reshape(regions.shape)


        def boundaries(ax):
            for i in range(len(region_names)):
                m = (ids == i).astype(float)
                if 0 < m.sum() < m.size and len(xv) > 1 and len(yv) > 1:
                    ax.contour(xv, yv, m, levels=[0.5], colors="k", linewidths=0.6)


        labels = <<labs>>
        n = len(labels) + 1
        ncol = min(3, n)
        nrow = -(-n // ncol)
        fig, axs = plt.subplots(nrow, ncol, figsize=(4.6 * ncol, 3.9 * nrow), squeeze=False)
        axs = axs.ravel()
        ax = axs[0]
        ax.pcolormesh(xv, yv, np.ma.masked_invalid(ids), cmap=plt.get_cmap("tab20", max(len(region_names), 2)),
                      shading="auto", vmin=-0.5, vmax=len(region_names) - 0.5, alpha=0.55)
        boundaries(ax)
        for i, name in enumerate(region_names):   # label each region near its centre
            jj, ii = np.nonzero(ids == i)
            if len(ii) >= 2:
                k = np.argsort(np.hypot(ii - ii.mean(), jj - jj.mean()))[0]
                ax.text(xv[ii[k]], yv[jj[k]], name.replace("+", "+\\n") if len(name) > 18 else name, fontsize=6.5,
                        ha="center", va="center", bbox=dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.7))
        ax.set_title("Stable phases", fontsize=9)
        for axq, lab in zip(axs[1:], labels):
            z = orient(columns[lab]).astype(float)
            fin = z[np.isfinite(z)]
            vmin, vmax = (0.0, 1.0) if lab.startswith("NP(") else ((fin.min(), fin.max()) if fin.size else (0.0, 1.0))
            vmax = vmax if vmax - vmin >= 1e-12 else vmin + 1e-9
            cmap = "RdBu_r" if lab.startswith("DF(") and vmin < 0 < vmax else "viridis"
            if cmap == "RdBu_r":
                vmin, vmax = -max(abs(vmin), abs(vmax)), max(abs(vmin), abs(vmax))
            axq.set_facecolor("0.85")
            pm = axq.pcolormesh(xv, yv, np.ma.masked_invalid(z), cmap=cmap, vmin=vmin, vmax=vmax,
                                shading="gouraud" if min(z.shape) > 1 else "auto")
            if min(z.shape) > 2 and np.isfinite(z).sum() > 4:
                axq.contour(xv, yv, np.ma.masked_invalid(z), levels=8, colors="w", linewidths=0.4, alpha=0.7)
            boundaries(axq)
            cb = fig.colorbar(pm, ax=axq, fraction=0.046, pad=0.03)
            if units[lab]:
                cb.set_label(units[lab], fontsize=8)
            axq.set_title(lab, fontsize=9)
        for axq in axs[:n]:
            axq.set_xlabel(<<xlabel>>, fontsize=8)
            axq.set_ylabel(<<ylabel>>, fontsize=8)
            axq.tick_params(labelsize=7)
        for axq in axs[n:]:
            axq.set_visible(False)
        fig.suptitle(<<title>>, fontsize=10)
        fig.tight_layout()
        fig.savefig(OUT / <<png>>, dpi=150, bbox_inches="tight")
        ''', xv=disp(x_lab), yv=disp(y_lab), transpose=str(axis_labels.index(x_lab) == 0), labs=Ctx.py(labs, 9),
              xlabel=repr(props._cond_axis_label(x_lab, unit)), ylabel=repr(props._cond_axis_label(y_lab, unit)),
              title=repr(title), png=repr(png_name))


# ---------------------------------------------------------------------------------------------------------
# step_map
# ---------------------------------------------------------------------------------------------------------
_STEP_BODY = '''
    s_els = [e for e in strategy.elements if e != "VA"]   # order of cs.X and of the chemical potentials
    props = {name: as_property(name) for name in <<props>>}


    def value(point, prop):
        try:
            return float(np.squeeze(point.get_property(prop)))
        except Exception:
            return np.nan


    # one row per calculated point: phase amounts NP(phase), compositions X(phase,EL), MU(EL) and properties
    rows, phases_seen = [], set()
    for zpf in strategy.zpf_lines:
        for p in zpf.points:
            T = float(np.squeeze(p.global_conditions[v.T]))
            row = {<<axis_name>>: float(np.squeeze(p.global_conditions[<<axis>>]))<<T_col>>, "T_C": T - KELVIN}
            point_phases = names(p.stable_composition_sets)
            row["stable_phases"] = " + ".join(point_phases)
            for n, cs in zip(point_phases, p.stable_composition_sets):
                phases_seen.add(n)
                row[f"NP({n})"] = row.get(f"NP({n})", 0.0) + float(cs.NP)
                row.update({f"X({n},{e})": float(x) for e, x in zip(s_els, np.asarray(cs.X))})
            row.update({f"MU({e})": float(mu) for e, mu in zip(s_els, p.chemical_potentials)})
            row.update({name: value(p, prop) for name, prop in props.items()})
            rows.append(row)
    rows.sort(key=lambda r: r[<<axis_name>>])
    phases_seen = sorted(phases_seen)
    for r in rows:
        for n in phases_seen:
            r.setdefault(f"NP({n})", 0.0)
    columns = [<<axis_name>><<T_colname>>, "T_C", "stable_phases"] + [f"NP({n})" for n in phases_seen] + \\
        [f"X({n},{e})" for n in phases_seen for e in s_els] + [f"MU({e})" for e in s_els] + list(props)
    table = pd.DataFrame(rows).reindex(columns=columns)
    table.to_csv(OUT / <<csv>>, index=False, float_format="%.8g")

    # phase regions along the stepped axis and the transitions between them
    segments = sorted([min(a), max(a), tuple(sorted(names(z.points[len(z.points) // 2].stable_composition_sets)))]
                      for z in strategy.zpf_lines if z.points
                      for a in [[float(np.squeeze(p.global_conditions[<<axis>>])) for p in z.points]])
    regions = []
    for s in segments:
        if regions and regions[-1][2] == s[2] and s[0] <= regions[-1][1] + 1e-6 * max(1.0, abs(s[0])):
            regions[-1][1] = max(regions[-1][1], s[1])
        elif s[1] - s[0] > 0 or not regions:
            regions.append(list(s))
    transitions = [round(r1[0], 6) for r1 in regions[1:]]
    for (a0, b0, ph0), (a1, b1, ph1) in zip(regions, regions[1:]):
        print(f"at {<<axis_name>>}={a1:.6g}: {' + '.join(ph0)} -> {' + '.join(ph1)}")
    '''


@generator("step_map")
def _gen_step_map(args, ctx: Ctx) -> str:  # noqa: C901
    from . import mapping_tools as mt
    a = {"temperature_unit": "K", "rename_disordered": True, **args}
    db = _db(ctx, a["tdb"])
    els = mt._check_elements(db, a["elements"])
    ph = mt._select_phases(db, els, a.get("phases"), a.get("exclude_phases"))
    unit = a["temperature_unit"]
    tC = _is_celsius(unit)
    vk = a["vary"].strip().upper()
    conds = ["v.P: 101325.0", "v.N: 1"]
    fx = {}
    P = 101325.0
    for k, val in (a.get("fixed") or {}).items():
        ku = k.strip().upper()
        if ku in ("T", "T_K"):
            fx["T"] = mt._toK(val, unit)
            conds.append(f"v.T: {_num(fx['T'])}")
        elif ku == "P":
            conds[0] = f"v.P: {_num(val)}"
        else:
            e = _el(ku)
            fx[f"X({e})"] = float(val)
            conds.append(f"v.X({e!r}): {_num(val)}")
    start, stop = float(a["start"]), float(a["stop"])
    if vk in ("T", "T_K"):
        lo, hi = mt._toK(start, unit), mt._toK(stop, unit)
        conds = [c for c in conds if not c.startswith("v.T:")]
        conds.append(f"v.T: ({_num(min(lo, hi))}, {_num(max(lo, hi))}, {_num(a.get('step') or 5.0)})")
        axis, axis_name = "v.T", "T_K"
    elif vk == "P":
        conds[0] = (f"v.P: ({_num(min(start, stop))}, {_num(max(start, stop))}, "
                    f"{_num(a.get('step') or abs(stop - start) / 100)})")
        axis, axis_name = "v.P", "P_Pa"
    else:
        e = _el(vk)
        conds = [c for c in conds if not c.startswith(f"v.X({e!r})")]
        conds.append(f"v.X({e!r}): ({_num(min(start, stop))}, {_num(max(start, stop))}, "
                     f"{_num(a.get('step') or 0.005)})")
        axis, axis_name = f"v.X({e!r})", f"X({e})"
    props_ = list(a["properties"]) if a.get("properties") is not None else ["GM", "HM", "SM", "CPM"]
    tag = "-".join(_el(e) for e in a["elements"]) + "_" + re.sub(r"[^A-Za-z0-9.]+", "", str(a["vary"])) + "_" + \
        "_".join(f"{re.sub(r'[^A-Za-z0-9.]+', '', k)}{val:g}" for k, val in (a.get("fixed") or {}).items())
    png = _outp(ctx.tb, (a["output_prefix"] + ".png") if a.get("output_prefix") else None,
                f"step_{tag}_{Path(ctx.tb._find_tdb(a['tdb'])).stem}.png")
    png_name = ctx.out_name(png, "step.png")
    csv_name = _register_siblings(ctx, png, [".csv"])[".csv"]
    dis = _disordered_names(db, ph, a["rename_disordered"])
    colors = _colors(list(ph) + [d for d, _ in dis.values()], a.get("phase_colors"))

    ctx.need("import numpy as np", "import pandas as pd", "import matplotlib.pyplot as plt",
             "import matplotlib.patches as mpatches", "from pycalphad import Database",
             "from pycalphad import variables as v", "from pycalphad.mapping import StepStrategy",
             "from pycalphad.property_framework import as_property")
    head = _load_db_code(ctx, a["tdb"]) + _T('''
        elements = <<els>>
        comps = <<comps>>
        phases = <<ph>>
        conds = {<<conds>>}
        KELVIN = 273.15

        # step calculation along <<axis_name>> (pycalphad StepStrategy)
        strategy = StepStrategy(db, comps, phases, conds)
        strategy.do_map()

        ''', els=Ctx.py(els), comps=Ctx.py(sorted(els) + ["VA"]), ph=Ctx.py(ph), conds=", ".join(conds),
                                             axis_name=axis_name)
    head += _names_code(dis) + "\n\n"
    body = textwrap.dedent(_T(_STEP_BODY, props=Ctx.py(props_), axis_name=repr(axis_name), axis=axis,
                              T_col=', "T_K": T' if axis != "v.T" else "", T_colname=', "T_K"' if axis != "v.T" else "",
                              csv=repr(csv_name)))
    # plot
    panels = list(a["plot"]) if a.get("plot") else ["NP"]
    pcode = []
    for pan in panels:
        pu = pan.strip()
        m = re.fullmatch(r"X\(\*\s*,\s*([A-Za-z]{1,2})\)", pu, re.I)
        if pu.upper() == "NP":
            pcode.append(("NP", None))
        elif m:
            pcode.append(("X", m.group(1).upper()))
        else:
            hdr = _csv_header(str(Path(png).with_suffix(".csv"))) or []
            pcode.append(("col", pu if pu in hdr else next((k for k in hdr if k.upper() == pu.upper()), pu)))
    xexpr = "table['T_K'] - KELVIN" if (axis == "v.T" and tC) else f"table[{axis_name!r}]"
    xlabel = (_T_label(tC) if axis == "v.T" else "Pressure (Pa)" if axis == "v.P"
              else f"Mole fraction {axis_name[2:-1].capitalize()}")
    comp_s = ", ".join(f"{k}={val:g}" for k, val in fx.items() if k != "T")
    if "T" in fx:
        comp_s += (", " if comp_s else "") + (f"T={fx['T'] - KELVIN:g} °C" if tC else f"T={fx['T']:g} K")
    title = a.get("title") or f"{'-'.join(e.capitalize() for e in els)}  {comp_s}"
    lines = ["", "", "# plot: one panel per entry of `plot`",
             f"colors = {Ctx.py(colors, 9)}",
             f"x = {xexpr}",
             f"fig, axs = plt.subplots({len(pcode)}, 1, figsize=(7.2, {3.2 * len(pcode) + 0.6:g}), sharex=True, "
             "squeeze=False)"]
    legend = ("    ax.legend(handles=[mpatches.Patch(color=colors[n.split('#')[0]], label=n) for n in phases_seen],\n"
              "              loc='center left', bbox_to_anchor=(1.02, 0.5), fontsize=8, frameon=False)")
    for i, (kind, what) in enumerate(pcode):
        lines.append(f"ax = axs[{i}, 0]")
        if kind == "NP":
            lines += ["for n in phases_seen:",
                      "    ax.plot(x, table[f'NP({n})'], color=colors[n.split('#')[0]], lw=1.6, label=n)",
                      "ax.set_ylabel('Phase fraction')", "ax.set_ylim(-0.02, 1.02)"]
            lines += ["if True:", legend]
        elif kind == "X":
            lines += ["for n in phases_seen:   # composition of each phase where it is present",
                      f"    ax.plot(x, table[f'X({{n}},{what})'].where(table[f'NP({{n}})'] > 0), "
                      "color=colors[n.split('#')[0]], lw=1.6)",
                      f"ax.set_ylabel('X({what.capitalize()}) in phase')", "if True:", legend]
        else:
            lines += [f"ax.plot(x, table[{what!r}], color='k', lw=1.5)", f"ax.set_ylabel({what!r})"]
        lines += ["ax.grid(alpha=0.25)",
                  "for t in transitions:",
                  "    ax.axvline(" + ("t - KELVIN" if (axis == "v.T" and tC) else "t") + ", color='0.6', lw=0.6, ls=':')"]
    lines += [f"axs[-1, 0].set_xlabel({xlabel!r})", f"axs[0, 0].set_title({title!r}, fontsize=11)",
              "fig.tight_layout()", f"fig.savefig(OUT / {png_name!r}, dpi=150, bbox_inches='tight')",
              "ax = axs[0, 0]"]
    code = "\n".join(lines).replace("if True:\n    ax.legend", "ax.legend").replace(
        "\n              loc='center left'", "\n          loc='center left'")
    # column-name checks: panels that name a column must exist in the table (as in the tool, case-insensitive)
    return head + body + code + "\n"


# ---------------------------------------------------------------------------------------------------------
# phase_property_plot
# ---------------------------------------------------------------------------------------------------------
_PURE_PROP = '''

def pure_element(el, phase, T, prop):
    """<<prop>> of pure `el` in `phase` at each T, in its lowest-G internal configuration."""
    T = np.atleast_1d(np.asarray(T, dtype=float))
    tu, inv = np.unique(T, return_inverse=True)
    r = calculate(db, <<pure_comps>>, phase, T=tu, P=P, N=1, output="GM", pdens=50)
    g = r.GM.values.reshape(len(tu), -1)
    best = np.nanargmin(g, axis=1)
    if prop == "GM":
        vals = g[np.arange(len(tu)), best]
    else:
        y = r.Y.values.reshape(len(tu), g.shape[1], -1)
        vals = np.array([float(np.ravel(calculate(db, <<pure_comps>>, phase, T=t, P=P, N=1, output=prop,
                                                  points=y[i, best[i]][None, :])[prop].values)[0])
                         for i, t in enumerate(tu)])
    return vals[inv]
'''

_PP_COMPOSITION = '''
    # for each phase: sample internal configurations, keep the lowest-G one per composition bin (0.005 wide)
    bins = np.linspace(0, 1, 201)
    fig, ax = plt.subplots(figsize=(7.2, 4.8))
    cmap = plt.get_cmap(<<cmap>>)
    rows, envelope = [], []
    for i, name in enumerate(phases):
        color = cmap(i % cmap.N)
        try:
            r = calculate(db, comps, name, T=T, P=P, N=1, output="GM", pdens=<<points>>)
            x = r.X.sel(component=<<xel>>).values.ravel()
            g = r.GM.values.ravel()<<value>>
        except Exception as e:
            print(f"skipped {name}: {type(e).__name__}")
            continue
        ok = np.isfinite(x) & np.isfinite(g)
        x, g, val = x[ok], g[ok], val[ok]<<refshift>><<scatter>>
        idx = np.digitize(x, bins)
        best = [np.nonzero(idx == k)[0][np.argmin(g[idx == k])] for k in np.unique(idx)]
        order = np.argsort(x[best])
        ex, ev, eg = x[best][order], val[best][order], g[best][order]
        if np.all(~np.isfinite(ev)):
            print(f"skipped {name}: NaN")
            continue
        if np.ptp(ex) < 0.011:   # (nearly) stoichiometric: a point
            ax.plot(ex, ev * <<scale>>, "o", color=color, ms=6, label=name)
        else:
            ax.plot(ex, ev * <<scale>>, "-", color=color, lw=1.8, label=name)
        for xa, va, ga in zip(ex, ev, eg):
            rows.append({"phase": name, <<xcol>>: xa, <<vcol>>: va<<gcol>>})
            envelope.append((xa, va, name))
    '''

_PP_TANGENT = '''

    # common-tangent construction: lower convex hull of all phases' curves; long hull segments are tie-lines
    def lower_hull(x, g):
        o = np.lexsort((g, x))
        pts = np.column_stack([x[o], g[o]])
        hull = []
        for i, p in enumerate(pts):
            while len(hull) >= 2:
                (x1, y1), (x2, y2) = pts[hull[-2]], pts[hull[-1]]
                if (x2 - x1) * (p[1] - y1) - (y2 - y1) * (p[0] - x1) <= 0:
                    hull.pop()
                else:
                    break
            hull.append(i)
        return o[hull]


    xa = np.array([e[0] for e in envelope])
    ga = np.array([e[1] for e in envelope])
    pa = [e[2] for e in envelope]
    hull = lower_hull(xa, ga)
    for i1, i2 in zip(hull[:-1], hull[1:]):
        if xa[i2] - xa[i1] > 0.012:
            ax.plot([xa[i1], xa[i2]], [ga[i1] * <<scale>>, ga[i2] * <<scale>>], "k--", lw=1.1, zorder=5)
            ax.plot([xa[i1], xa[i2]], [ga[i1] * <<scale>>, ga[i2] * <<scale>>], "ko", ms=3.5, zorder=6)
            print(f"two-phase {pa[i1]} + {pa[i2]}: X(<<xel_bare>>) {xa[i1]:.4f} - {xa[i2]:.4f}")
    '''


@generator("phase_property_plot")
def _gen_phase_property_plot(args, ctx: Ctx) -> str:  # noqa: C901
    from . import props
    a = {"temperature_unit": "K", "show_all_points": False, "common_tangent": False, **args}
    db = _db(ctx, a["tdb"])
    els = props._check_elements(db, a["elements"])
    prop = a["prop"].strip().upper()
    prop = {"G_MIX": "GM_MIX", "H_MIX": "HM_MIX", "S_MIX": "SM_MIX", "CP": "CPM", "CP_MIX": "CPM_MIX"}.get(prop, prop)
    ph = props._resolve_phases(db, els, a.get("phases"), a.get("exclude_phases"))
    comps = props._comps(db, els)
    ref = {k.upper(): val.upper() for k, val in (a.get("reference") or {}).items()}
    unit_lab, scale = props._prop_units(prop)
    raw_u = {"kJ/mol": "J/mol", "J/(mol·K)": "J/mol/K"}.get(unit_lab, unit_lab)
    vcol = f"{prop} [{raw_u}]" if raw_u else prop
    celsius = _is_celsius(a["temperature_unit"])
    temperature = a["temperature"]
    tv = temperature if not isinstance(temperature, list) else "T" + "-".join(f"{x:g}" for x in temperature[:2])
    stem = props._safe_name(f"phaseprop_{'-'.join(sorted(e.upper() for e in a['elements']))}_{a['prop'].upper()}_"
                            f"{tv}{a['temperature_unit'] or 'K'}")
    base = _outp(ctx.tb, a.get("output_name"), stem)
    if Path(base).suffix.lower() in (".png", ".csv", ".nc"):
        base = str(Path(base).with_suffix(""))
    png_name = ctx.out_name(base + ".png", "phaseprop.png")
    csv_name = ctx.out_name(base + ".csv", "phaseprop.csv")
    tvals = props._values(temperature, "temperature")
    if celsius:
        tvals = tvals + KELVIN
    sysname = "-".join(e.capitalize() for e in els)
    ref_txt = (" relative to " + ", ".join(f"{e.capitalize()}({p})" for e, p in ref.items())) if ref else ""
    ylabel = f"{prop}{ref_txt and ' (ref.)'}" + (f" ({unit_lab})" if unit_lab else "")
    P = 101325.0 if a.get("pressure") is None else float(a["pressure"])
    ctx.need("import numpy as np", "import pandas as pd", "import matplotlib.pyplot as plt",
             "from pycalphad import Database")
    cmap = repr("tab10" if len(ph) <= 10 else "tab20")
    head = _load_db_code(ctx, a["tdb"]) + (f"comps = {Ctx.py(comps)}\nphases = {Ctx.py(ph)}\n"
                                           f"P = {_num(P)}\nKELVIN = 273.15\n")
    pure = ""
    if ref:
        ctx.need("from pycalphad import calculate")
        pure = _T(_PURE_PROP, prop=prop, pure_comps="[el, 'VA']" if "VA" in comps else "[el]") + "\n\n"
    if len(tvals) == 1:
        if len(els) != 2:
            ctx.fallback("vs-composition curves need a binary")
        ctx.need("from pycalphad import calculate")
        T = float(tvals[0])
        xel = (a.get("x_element") or sorted(els)[1]).upper()
        other = [e for e in els if e != xel][0]
        value = ("\n        val = g.copy()" if prop == "GM" else
                 f"\n        y = r.Y.values.reshape(len(g), -1)   # same configurations, other property\n"
                 f"        val = calculate(db, comps, name, T=T, P=P, N=1, output={prop!r}, points=y)"
                 f"[{prop!r}].values.ravel()")
        refshift = ""
        code = head + ("T = " + (f"{_num(tvals[0] - KELVIN)} + KELVIN" if celsius else _num(T)) + "\n") + pure
        if ref:
            code += (f"reference = {Ctx.py(ref)}   # element -> reference phase (pure element, same T)\n"
                     f"ref_value = {{el: float(pure_element(el, phase, T, {prop!r})[0]) for el, phase in "
                     "reference.items()}\n")
            refshift = (f"\n    val = val - (x * ref_value[{xel!r}] + (1 - x) * ref_value[{other!r}])   "
                        "# relative to the pure-element references")
        scatter = (f"\n    ax.scatter(x, val * {scale!r}, s=2, color=color, alpha=0.25, edgecolors='none')   "
                   "# every sampled configuration") if a["show_all_points"] else ""
        body = _T(_PP_COMPOSITION, cmap=cmap, points=int(a.get("points") or 2000), xel=repr(xel), value=value,
                  refshift=refshift, scatter=scatter, scale=repr(scale), xcol=repr(f"X({xel})"), vcol=repr(vcol),
                  gcol="" if prop == "GM" else ', "GM [J/mol]": ga')
        code += "\n" + body
        if a["common_tangent"]:
            if prop != "GM":
                ctx.fallback("common_tangent needs prop='GM'")
            code += _T(_PP_TANGENT, scale=repr(scale), xel_bare=xel)
        tt = f"{T - KELVIN:.0f} °C" if celsius else f"{T:.0f} K"
        title = a.get("title") or f"{sysname}: {prop} of each phase at {tt}{ref_txt}"
        tail = [f"ax.set_xlim(0, 1)", f"ax.set_xlabel({'Mole fraction ' + xel.capitalize()!r})",
                f"ax.set_ylabel({ylabel!r})", f"ax.set_title({title!r}, fontsize=10)"]
        if ref or prop.endswith("_MIX"):
            tail.append("ax.axhline(0, color='0.5', lw=0.6)")
        tail += ["ax.grid(alpha=0.3)", f"ax.legend(fontsize=8, ncol={2 if len(ph) > 6 else 1})",
                 f"fig.savefig(OUT / {png_name!r}, dpi=150, bbox_inches='tight')",
                 "df = pd.DataFrame(rows)", "df.insert(0, 'T [K]', T)",
                 f"df.to_csv(OUT / {csv_name!r}, index=False, float_format='%.8g')"]
        return code + "\n" + "\n".join(tail) + "\n"

    # ---- vs temperature at fixed composition: each phase alone (IsolatedPhase, internal DOF relaxed)
    ctx.need("from collections import OrderedDict", "from pycalphad import variables as v",
             "from pycalphad.core.workspace import Workspace", "from pycalphad.property_framework import IsolatedPhase")
    comp = {k.upper().replace("X(", "").replace(")", ""): float(x) for k, x in (a.get("composition") or {}).items()}
    if len(els) == 1:
        comp = {}
    bal = [e for e in els if e not in comp][0]
    xfull = {**comp, bal: 1 - sum(comp.values())}
    code = head + pure
    arr_expr, _, note = _cond_values_code(temperature if isinstance(temperature, (list, dict)) else [temperature],
                                          "T", celsius)
    code += f"T = {arr_expr}" + (f"   # {note}" if note else "") + "\n"
    code += f"composition = {Ctx.py(comp)}   # mole fractions (balance {bal})\n"
    code += ("conds = {v.T: T, v.P: P, v.N: 1, **{v.X(el): x for el, x in composition.items()}}\n")
    if ref:
        code += (f"reference = {Ctx.py(ref)}   # element -> reference phase (pure element, same T)\n"
                 f"x_all = {Ctx.py(xfull)}\n"
                 f"shift = sum(x_all[el] * pure_element(el, phase, T, {prop!r}) for el, phase in reference.items())\n")
    else:
        code += "shift = np.zeros_like(T)\n"
    tdisp = "T - KELVIN" if celsius else "T"
    cs = ", ".join(f"X({k})={x:g}" for k, x in comp.items())
    title = a.get("title") or f"{sysname} {cs}: {prop} of each phase (single-phase, internally relaxed){ref_txt}"
    code += _T('''

        # each phase alone at this composition: its own (metastable) internal equilibrium, via IsolatedPhase
        class Guarded:
            """NaN instead of an exception at temperatures where the single-phase solve fails."""

            def __init__(self, inner):
                self.inner, self.shape = inner, getattr(inner, "shape", tuple())
                self.implementation_units, self.display_units = inner.implementation_units, inner.display_units

            def compute_property(self, compsets, cur_conds, chemical_potentials):
                if len(compsets) == 0:
                    return np.nan
                try:
                    return self.inner.compute_property(compsets, cur_conds, chemical_potentials)
                except Exception:
                    return np.nan


        table = OrderedDict([("T [K]", T)<<tc_col>>])
        fig, ax = plt.subplots(figsize=(7.2, 4.8))
        cmap = plt.get_cmap(<<cmap>>)
        for i, name in enumerate(phases):
            try:
                wks = Workspace(database=db, components=comps, phases=[name], conditions=conds)
                isolated = IsolatedPhase(name, wks)
                if isolated._compset is None:
                    raise RuntimeError("phase cannot reach this composition")
                values = np.asarray(wks.get(Guarded(isolated(<<pname>>))), dtype=float).ravel() - shift
            except Exception as e:
                print(f"skipped {name}: {type(e).__name__}: {str(e)[:60]}")
                continue
            if not np.isfinite(values).any():
                print(f"skipped {name}: cannot reach this composition")
                continue
            table[f"<<prop>>({name})"] = values
            ax.plot(<<tdisp>>, values * <<scale>>, lw=1.8, color=cmap(i % cmap.N), label=name)
        ax.set_xlabel(<<xlabel>>)
        ax.set_ylabel(<<ylabel>>)
        ax.set_title(<<title>>, fontsize=9)
        ax.grid(alpha=0.3)
        ax.legend(fontsize=8, ncol=<<ncol>>)
        fig.savefig(OUT / <<png>>, dpi=150, bbox_inches="tight")
        pd.DataFrame(table).to_csv(OUT / <<csv>>, index=False, float_format="%.8g")
        ''', tc_col=', ("T [°C]", T - KELVIN)' if celsius else "", cmap=cmap, pname='f"' + prop + '({name})"',
               prop=prop, tdisp=tdisp, scale=repr(scale), xlabel=repr(_T_label(celsius)), ylabel=repr(ylabel),
               title=repr(title), ncol=2 if len(ph) > 6 else 1, png=repr(png_name), csv=repr(csv_name))
    return code


# ---------------------------------------------------------------------------------------------------------
# export_equilibrium_data
# ---------------------------------------------------------------------------------------------------------
_TO_CELSIUS = '''

def kelvin_to_celsius(ax):
    """Shift everything drawn with temperature (K) on the y axis to °C, so the axis numbers are the data."""
    for ln in ax.get_lines():
        ln.set_ydata(np.asarray(ln.get_ydata(), dtype=float) - KELVIN)
    for c in ax.collections:
        if isinstance(c, LineCollection):
            c.set_segments([np.asarray(s, dtype=float) - [0, KELVIN] for s in c.get_segments()])
        elif hasattr(c, "set_verts"):   # filled areas (fill_between)
            c.set_verts([np.asarray(p.vertices, dtype=float) - [0, KELVIN] for p in c.get_paths()])
        elif len(c.get_offsets()):
            c.set_offsets(np.asarray(c.get_offsets(), dtype=float) - [0, KELVIN])
    lo, hi = ax.get_ylim()
    ax.set_ylim(lo - KELVIN, hi - KELVIN)
    ax.set_ylabel("Temperature (°C)")
'''

_EXPORT_ROWS = '''
# tidy table: one row per grid point and stable phase (conditions, GM, MU, phase, NP, phase composition, site fractions)
cond_dims = list(eq.GM.dims)
components = [str(c) for c in eq.component.values]
phase_arr = np.asarray(eq.Phase.values, dtype=str)
rows = []
for idx in np.ndindex(*eq.GM.shape):
    base = {}
    for d, i in zip(cond_dims, idx):
        val = float(eq.coords[d].values[i])
        if d == "T":
            base["T [K]"] = val<<tc>>
        elif d == "P":
            base["P [Pa]"] = val
        elif d != "N":
            base[d.replace("X_", "X(") + ")" if d.startswith("X_") else d] = val
    base["GM [J/mol]"] = float(eq.GM.values[idx])
    base.update({f"MU({c}) [J/mol]": float(eq.MU.values[idx + (j,)]) for j, c in enumerate(components)})<<extra>>
    for vtx, phase in enumerate(phase_arr[idx]):
        if not phase:
            continue
        row = {**base, "phase": phase, "NP": float(eq.NP.values[idx + (vtx,)])}
        row.update({f"X(phase,{c})": float(eq.X.values[idx + (vtx, j)]) for j, c in enumerate(components)})
        row.update({f"Y{k}": float(y) for k, y in enumerate(eq.Y.values[idx + (vtx,)]) if np.isfinite(y)})
        rows.append(row)
pd.DataFrame(rows).to_csv(OUT / <<csv>>, index=False, float_format="%.8g")
print(pd.DataFrame(rows))
'''


@generator("export_equilibrium_data")
def _gen_export_equilibrium_data(args, ctx: Ctx) -> str:
    from . import props
    a = {"temperature_unit": "K", "eqplot": False, **args}
    unit = a["temperature_unit"] or "K"
    celsius = _is_celsius(unit)
    db = _db(ctx, a["tdb"])
    els = props._check_elements(db, a["elements"])
    ph = props._resolve_phases(db, els, a.get("phases"), a.get("exclude_phases"))
    conds, wnote = _conditions(ctx, db, els, a["conditions"], unit)
    extra = [e.strip().upper() for e in (a.get("extra_outputs") or [])]
    formats = [f.lower().lstrip(".") for f in (a.get("formats") or ["nc", "csv"])]
    base = _outp(ctx.tb, a.get("output_name"), f"eqdata_{'-'.join(sorted(e.upper() for e in a['elements']))}")
    if Path(base).suffix.lower() in (".png", ".csv", ".nc"):
        base = str(Path(base).with_suffix(""))
    ctx.need("import numpy as np", "import pandas as pd", "from pycalphad import Database, equilibrium",
             "from pycalphad import variables as v")
    code = [_load_db_code(ctx, a["tdb"]), f"comps = {Ctx.py(props._comps(db, els))}\n",
            f"phases = {Ctx.py(ph)}\nKELVIN = 273.15\n\n",
            "# conditions (T in K" + (", given in °C" if celsius else "") + ")" + (f"; {wnote}" if wnote else "")
            + "\n"]
    for lab, pcode, vcode, arr, note in conds:
        code.append(f"{_var(lab)} = {vcode}" + (f"   # {note}" if note else "") + "\n")
    code.append("conds = {" + ", ".join(f"{pcode}: {_var(lab)}" for lab, pcode, *_ in conds) + "}\n")
    code.append(f"eq = equilibrium(db, comps, phases, conds{', output=' + Ctx.py(extra) if extra else ''})\n")
    if "nc" in formats:
        nc = ctx.out_name(base + ".nc", "eqdata.nc")
        code.append(_T('''

            # the raw xarray Dataset (netCDF3 via scipy; attributes must be strings)
            ds = eq.copy()
            ds.attrs = {k: str(val) for k, val in ds.attrs.items()}
            ds.attrs["tdb"] = <<tdb>>
            ds.attrs["phases"] = " ".join(phases)
            ds.to_netcdf(OUT / <<nc>>, engine="scipy")
            ''', tdb=repr(Path(ctx.tb._find_tdb(a["tdb"])).name), nc=repr(nc)))
    if "csv" in formats:
        csv = ctx.out_name(base + ".csv", "eqdata.csv")
        code.append("\n" + _T(_EXPORT_ROWS, csv=repr(csv),
                              tc='\n            base["T [°C]"] = val - KELVIN' if celsius else "",
                              extra=("\n    base.update({e: float(eq[e].values[idx]) for e in " + Ctx.py(extra)
                                     + " if e in eq})") if extra else ""))
    if a["eqplot"]:
        varying = [c[0] for c in conds if len(c[3]) > 1]
        indep_x = [x for x in varying if x.startswith("X(")]
        indep_p = [x for x in varying if x in ("T", "P")]
        if (len(indep_x) == 1 and len(indep_p) == 1) or (len(indep_x) == 2 and not indep_p):
            ctx.need("import matplotlib.pyplot as plt", "from pycalphad.plot.eqplot import eqplot")
            png = ctx.out_name(base + ".png", "eqdata.png")
            tern = len(indep_x) == 2
            lines = ["", "", "# pycalphad's legacy eqplot: tie-line end points of the grid"]
            if tern:
                ctx.need("from pycalphad.plot import triangular  # registers the 'triangular' projection")
                lines += ["fig = plt.figure(figsize=(7, 6.5))", "ax = fig.add_subplot(111, projection='triangular')"]
            else:
                lines += ["fig = plt.figure(figsize=(7, 5.5))", "ax = fig.add_subplot(111)"]
            lines.append("eqplot(eq.copy(), ax=ax)")
            if not tern and indep_p == ["T"] and celsius:
                ctx.need("from matplotlib.collections import LineCollection")
                code.append(_T(_TO_CELSIUS))
                lines.append("kelvin_to_celsius(ax)")
            title = f"{'-'.join(e.capitalize() for e in els)} equilibrium grid (eqplot: tie-line end points)"
            lines += [f"ax.set_title({title!r}, fontsize=10)", "ax.xaxis.label.set_size(11)",
                      "ax.yaxis.label.set_size(11)", f"fig.savefig(OUT / {png!r}, dpi=150, bbox_inches='tight')"]
            code.append("\n".join(lines) + "\n")
    return "".join(code)


# ---------------------------------------------------------------------------------------------------------
# scheil_solidification
# ---------------------------------------------------------------------------------------------------------
_LIQUIDUS = '''

def fully_liquid_start(tmax=4000.0):
    """Lowest temperature of a 25 K grid (600 K .. tmax) above which the alloy is fully liquid, + 30 K."""
    T = np.arange(600.0, tmax + 1, 25.0)
    eq = equilibrium(db, comps, phases, {v.T: T, v.P: 101325, v.N: 1, **composition})
    liquid_fraction = np.nansum(np.where(np.squeeze(eq.Phase.values) == liquid, np.squeeze(eq.NP.values), 0.0),
                                axis=-1)
    full = liquid_fraction > 1 - 1e-8
    idx = len(full) - 1
    while idx > 0 and full[idx - 1]:
        idx -= 1
    return float(T[idx] + 30.0)
'''

_SCHEIL_BODY = '''

# scheil names disordered states of order/disorder models A2_BCC / A1_FCC: use BCC_A2 / FCC_A1 if the TDB has them
def rename_disordered(res):
    present = {p for p, amounts in res.cum_phase_amounts.items() if amounts and amounts[-1] > 0}
    renames = {}
    for name in list(res.phase_amounts):
        m = re.fullmatch(r"([A-Z]\\d+)_([A-Z]+)", name)
        alt = f"{m.group(2)}_{m.group(1)}" if m else None
        if alt and alt in db.phases and alt not in present and alt not in renames.values():
            renames[name] = alt
    for old, new in renames.items():
        for d in (res.phase_amounts, res.cum_phase_amounts, res.phase_compositions):
            if old in d:
                d[new] = d.pop(old)


def write_csv(res, name):
    df = res.to_dataframe(include_zero_phases=False)
    df.insert(1, "Temperature (C)", df["Temperature (K)"] - KELVIN)
    df.insert(2, "fraction_solid", res.fraction_solid)
    df.to_csv(OUT / name, index=False, float_format="%.8g")


def summary(res, label):
    T, fs = np.asarray(res.temperatures, dtype=float), np.asarray(res.fraction_solid, dtype=float)
    first = T[np.argmax(fs > 1e-10)] if np.any(fs > 1e-10) else None
    print(f"{label}: first solid at {first} K, ends at {T[-1]:.2f} K, converged={res.converged}")


rename_disordered(scheil)
write_csv(scheil, <<scsv>>)
summary(scheil, "Scheil")<<eqpart>>

# plot: T vs fraction solid (both models) and the cumulative Scheil phase amounts
colors = <<colors>>
fig, (ax, ax2) = plt.subplots(1, 2, figsize=(11, 4.6))
ax.plot(scheil.fraction_solid, <<tconv_s>>, color="#1f77b4", lw=1.8, label="Scheil-Gulliver")<<eqline>>
ax.set_xlabel("Fraction solid")
ax.set_ylabel(<<tlabel>>)
ax.set_xlim(0, 1)
ax.grid(alpha=0.25)
ax.legend(fontsize=9)
for name, cum in sorted(scheil.cum_phase_amounts.items()):
    if cum[-1] > 1e-6:
        ax2.plot(<<tconv_s>>, cum, color=colors.get(name.split("#")[0], "k"), lw=1.6, label=name)
ax2.plot(<<tconv_s>>, scheil.fraction_liquid, color=colors.get(liquid, "k"), lw=1.6, ls=":", label=f"{liquid} (remaining)")
ax2.set_xlabel(<<tlabel>>)
ax2.set_ylabel("Cumulative phase amount (Scheil)")
ax2.set_ylim(-0.02, 1.02)
ax2.invert_xaxis()
ax2.grid(alpha=0.25)
ax2.legend(fontsize=8)
fig.suptitle(<<title>>, fontsize=12)
fig.tight_layout()
fig.savefig(OUT / <<png>>, dpi=150, bbox_inches="tight")
'''


@generator("scheil_solidification")
def _gen_scheil(args, ctx: Ctx) -> str:
    from . import mapping_tools as mt
    a = {"temperature_unit": "K", "step_temperature": 1.0, "include_equilibrium": True, **args}
    db = _db(ctx, a["tdb"])
    comp = {_el(k): float(x) for k, x in a["composition"].items()}
    els = mt._check_elements(db, sorted(set(_el(e) for e in a["elements"]) | set(comp)))
    ph = mt._select_phases(db, els, a.get("phases"), a.get("exclude_phases"))
    liq = (a.get("liquid_phase") or next((p for p in ph if mt._is_liquid(db, p)), "")).upper()
    if liq not in ph:
        ctx.fallback("no liquid phase")
    tC = _is_celsius(a["temperature_unit"])
    tag = "-".join(_el(e) for e in a["elements"]) + "_" + "_".join(f"{_el(k)}{val:g}" for k, val in
                                                                  a["composition"].items())
    png = _outp(ctx.tb, (a["output_prefix"] + ".png") if a.get("output_prefix") else None,
                f"scheil_{tag}_{Path(ctx.tb._find_tdb(a['tdb'])).stem}.png")
    png_name = ctx.out_name(png, "scheil.png")
    csvs = _register_siblings(ctx, png, ["_scheil.csv", "_equilibrium.csv"])
    ctx.need("import re", "import numpy as np", "import matplotlib.pyplot as plt", "from pycalphad import Database",
             "from pycalphad import variables as v",
             "from scheil import simulate_scheil_solidification, simulate_equilibrium_solidification")
    code = _load_db_code(ctx, a["tdb"]) + _T('''
        comps = <<comps>>
        phases = <<ph>>
        liquid = <<liq>>
        composition = {<<comp>>}   # mole fractions (balance <<bal>>)
        KELVIN = 273.15
        ''', comps=Ctx.py(sorted(els) + ["VA"]), ph=Ctx.py(ph), liq=repr(liq),
                                                 comp=", ".join(f"v.X({k!r}): {_num(x)}" for k, x in comp.items()),
                                                 bal=[e for e in els if e not in comp][0])
    if a.get("start_temperature") is not None:
        T0 = mt._toK(a["start_temperature"], a["temperature_unit"])
        code += f"T_start = {_num(T0)}   # K, fully liquid\n"
    else:
        ctx.need("from pycalphad import equilibrium")
        code += "\n\n" + _T(_LIQUIDUS) + "\n\nT_start = fully_liquid_start()   # K\n"
    step = _num(a["step_temperature"])
    stop = _num(a.get("stop", 1e-4) if a.get("stop") is not None else 1e-4)
    code += (f"\n# Scheil-Gulliver (no diffusion in the solid) and equilibrium solidification, {step} K steps\n"
             f"scheil = simulate_scheil_solidification(db, comps, phases, dict(composition), T_start,\n"
             f"                                        step_temperature={step}, liquid_phase_name=liquid, stop={stop})\n")
    if a["include_equilibrium"]:
        code += (f"equil = simulate_equilibrium_solidification(db, comps, phases, dict(composition), T_start,\n"
                 f"                                            step_temperature={step}, liquid_phase_name=liquid)\n")
    colors = _colors(list(ph) + ["BCC_A2", "FCC_A1", "A2_BCC", "A1_FCC", "HCP_A3"], a.get("phase_colors"))
    comp_s = ", ".join(f"X({k})={val:g}" for k, val in comp.items())
    title = a.get("title") or f"{'-'.join(e.capitalize() for e in els)} solidification, {comp_s}"
    tconv = "np.asarray(scheil.temperatures) - KELVIN" if tC else "np.asarray(scheil.temperatures)"
    eqpart = eqline = ""
    if a["include_equilibrium"]:
        eqpart = (f"\nrename_disordered(equil)\nwrite_csv(equil, {csvs['_equilibrium.csv']!r})\n"
                  "summary(equil, 'Equilibrium')")
        eqline = ("\nax.plot(equil.fraction_solid, " + tconv.replace("scheil.", "equil.")
                  + ', color="#ff7f0e", lw=1.8, ls="--", label="Equilibrium")')
    code += "\n\n" + _T(_SCHEIL_BODY, scsv=repr(csvs["_scheil.csv"]), eqpart=eqpart, colors=Ctx.py(colors, 9),
               tconv_s=tconv, eqline=eqline, tlabel=repr(_T_label(tC)), title=repr(title), png=repr(png_name))
    return code


# ---------------------------------------------------------------------------------------------------------
# fit_data
# ---------------------------------------------------------------------------------------------------------
def _fit_model_code(name: str, model, op: Dict) -> tuple:
    """(model function code, log-model code or None, T conversion code) for figfit's built-in models."""
    from .figfit import _t_of_x
    offset = bool(op.get("offset"))
    c = ", c" if offset else ""
    plus_c = " + c" if offset else ""
    pn = ", ".join(model.pnames)
    tconv, logf = "", None
    if name == "linear":
        f = "def model(x, intercept, slope):\n    return intercept + slope * x\n"
    elif name.startswith("poly"):
        f = (f"def model(x, *c):   # polynomial of degree {len(model.pnames) - 1}, c0 + c1·x + ...\n"
             "    return np.polynomial.polynomial.polyval(x, c)\n")
    elif name == "exp":
        f = f"def model(x, A, b{c}):\n    return A * np.exp(b * x){plus_c}\n"
        logf = "def log_model(x, A, b):   # ln of the model\n    return np.log(np.abs(A)) + b * x\n"
    elif name == "power":
        f = f"def model(x, A, b{c}):\n    return A * np.power(x, b){plus_c}\n"
        logf = "def log_model(x, A, b):   # ln of the model\n    return np.log(np.abs(A)) + b * np.log(x)\n"
    elif name == "log":
        f = "def model(x, a, b):\n    return a + b * np.log(x)\n"
    elif name == "arrhenius":
        factor, off, inverse = _t_of_x(op.get("T_unit") or "K")
        if inverse:
            tconv = f"def T_of(x):   # x -> absolute temperature (K)\n    return {_num(factor)} / np.asarray(x, dtype=float)\n"
        else:
            tconv = ("def T_of(x):   # x -> absolute temperature (K)\n    return np.asarray(x, dtype=float)"
                     + (f" * {_num(factor)}" if factor != 1 else "") + (f" + {_num(off)}" if off else "") + "\n")
        f = (f"def model(x, A, Ea{c}):   # Ea in kJ/mol\n"
             f"    return A * np.exp(-Ea * 1000.0 / (R * T_of(x))){plus_c}\n")
        logf = ("def log_model(x, A, Ea):   # ln of the model: ln A - Ea/(R·T)\n"
                "    return np.log(np.abs(A)) - Ea * 1000.0 / (R * T_of(x))\n")
    elif name == "logistic":
        f = f"def model(x, L, k, x0{c}):\n    return L / (1.0 + np.exp(-k * (x - x0))){plus_c}\n"
    elif name == "gaussian":
        f = f"def model(x, A, mu, sigma{c}):\n    return A * np.exp(-0.5 * ((x - mu) / sigma) ** 2){plus_c}\n"
    else:
        return None
    if offset:
        logf = None
    return f, logf, tconv


# fit_data's generator lives in codegen_plot (shared with the 'fit' plot op); this one is kept unregistered
def _gen_fit_data(args, ctx: Ctx) -> str:  # noqa: C901
    import numpy as np
    from matplotlib.figure import Figure
    from . import figfit
    from .figures import OpContext, find_col
    op = {k: val for k, val in args.items() if val is not None and k not in ("output_csv", "npoints", "level",
                                                                              "band", "range")}
    npoints = int(args.get("npoints") or 200)
    level = float(args.get("level") or 0.95)
    band = args.get("band") or "confidence"
    if op.get("transform"):
        ctx.fallback("data transform steps are applied by the toolkit's data loader")
    # replicate the tool's inputs at generation time (data, start values, linearization)
    fig = Figure()
    octx = OpContext(fig, fig.add_subplot(), op, 1, ctx.tb.p, [])
    xd, yd = np.asarray(octx.X(), dtype=float), np.asarray(octx.Y(), dtype=float)
    sig = None
    for key in ("sigma", "yerr"):
        if octx.raw(key) is not None:
            sig = np.abs(octx.dY(key))
            sig_key = key
            break
    model = figfit.build_model(op, "K")
    name = model.name
    mc = _fit_model_code(name, model, op)
    if mc is None:
        ctx.fallback(f"{name} models are evaluated with the toolkit's safe expression parser")
    fcode, logcode, tconv = mc
    lin = figfit.default_linearize(op, yd, sig)
    y_is = (op.get("y_is") or "value").lower()
    # data code
    ctx.need("import numpy as np", "import pandas as pd", "from scipy.optimize import curve_fit",
             "from scipy import stats")
    src = op.get("csv") if op.get("csv") is not None else op.get("data")
    lines = []
    if isinstance(src, str):
        if not src.lower().endswith((".csv", ".txt")) or op.get("sheet") or op.get("read") or op.get("variables"):
            ctx.fallback("this data source is read by the toolkit's data loader")
        import pandas as pd
        df = pd.read_csv(ctx.tb.p(src))
        lines.append(f"df = pd.read_csv({ctx.path(src)})")
        for col, val in (op.get("where") or {}).items():
            cc = find_col(df, col, src)
            vals = val if isinstance(val, list) else [val]
            df = df[df[cc].isin(vals)]
            lines.append(f"df = df[df[{cc!r}].isin({Ctx.py(vals)})]")
        cols = {}
        for key in ("x", "y", sig_key if sig is not None else None):
            if key and op.get(f"{key}_col"):
                cols[key] = find_col(df, op[f"{key}_col"], src)
        for key, arr in (("x", xd), ("y", yd)):
            if key not in cols:
                ctx.fallback("mixing inline values with a data file")
            if not np.array_equal(df[cols[key]].to_numpy(dtype=float), arr, equal_nan=True):
                ctx.fallback("the data could not be re-read identically with pandas.read_csv")
            lines.append(f"{key} = df[{cols[key]!r}].to_numpy(dtype=float)")
        if sig is not None:
            if sig_key in cols:
                lines.append(f"sigma = np.abs(df[{cols[sig_key]!r}].to_numpy(dtype=float))")
            else:
                lines.append(f"sigma = np.broadcast_to(np.abs(np.asarray({Ctx.py(np.atleast_1d(sig).tolist())}, "
                             "dtype=float)), x.shape).copy()")
    else:
        lines.append(f"x = np.array({Ctx.py(xd.tolist())})")
        lines.append(f"y = np.array({Ctx.py(yd.tolist())})")
        if sig is not None:
            lines.append(f"sigma = np.broadcast_to(np.array({Ctx.py(np.atleast_1d(sig).tolist())}), x.shape).copy()")
    if sig is None:
        lines.append("sigma = None")
        lines.append("ok = np.isfinite(x) & np.isfinite(y)")
    else:
        lines.append("ok = np.isfinite(x) & np.isfinite(y) & np.isfinite(sigma) & (sigma > 0)")
    lines.append("x, y = x[ok], y[ok]" + ("\nsigma = sigma[ok]" if sig is not None else ""))
    # start values exactly as the tool computed them
    okm = np.isfinite(xd) & np.isfinite(yd)
    if sig is not None:
        s_b = np.broadcast_to(np.asarray(sig, dtype=float), xd.shape)
        okm &= np.isfinite(s_b) & (s_b > 0)
    start = model.p0(xd[okm], yd[okm])
    for pname, val in (op.get("p0") or {}).items():
        start[model.pnames.index(pname)] = float(val)
    k = len(model.pnames)
    lo, hi = np.full(k, -np.inf), np.full(k, np.inf)
    for pname, (b0, b1) in (op.get("bounds") or {}).items():
        i = model.pnames.index(pname)
        lo[i] = -np.inf if b0 is None else float(b0)
        hi[i] = np.inf if b1 is None else float(b1)
    if op.get("bounds"):
        start = np.clip(start, np.where(np.isfinite(lo), lo + 1e-12 * np.abs(lo), lo),
                        np.where(np.isfinite(hi), hi - 1e-12 * np.abs(hi), hi))
    code = ["\n".join(lines), ""]
    if name == "arrhenius":
        code.append("R = 8.314462618   # J/(mol K)\n")
        code.append(tconv)
    code.append(fcode)
    space = "y"
    if y_is in ("ln", "log10") or lin:
        if logcode is None:
            ctx.fallback("linearized fit of a model without a log form")
        code.append(logcode)
        if y_is == "ln":
            code.append("fit_fn, target, back, sigma_t = log_model, y, (lambda t: t), sigma   # y is given as ln(...)\n")
            space = "ln y"
        elif y_is == "log10":
            code.append("fit_fn, target, back, sigma_t = (lambda x, *p: log_model(x, *p) / np.log(10.0)), y, "
                        "(lambda t: t), sigma   # y is given as log10(...)\n")
            space = "log10 y"
        else:
            code.append("# fitted in ln y (multiplicative scatter" + (", default for arrhenius" if op.get("linearize")
                                                                        is None else "") + ")\n"
                        "fit_fn, target, back = log_model, np.log(y), np.exp\n"
                        "sigma_t = None if sigma is None else sigma / y\n")
            space = "ln y"
    else:
        code.append("fit_fn, target, back, sigma_t = model, y, (lambda t: t), sigma\n")
    kw = (f", bounds=({Ctx.py(lo.tolist())}, {Ctx.py(hi.tolist())})".replace("inf", "np.inf")
          if op.get("bounds") else ", maxfev=20000")
    absolute = bool(op.get("absolute_sigma")) and sig is not None
    code.append(_T('''
        p0 = <<p0>>   # start values (<<pnames>>): the toolkit's automatic guess
        p, pcov = curve_fit(fit_fn, x, target, p0=p0, sigma=sigma_t, absolute_sigma=<<absolute>><<kw>>)
        se = np.sqrt(np.clip(np.diag(pcov), 0, None)) if np.all(np.isfinite(pcov)) else np.full(len(p), np.nan)
        fitted = fit_fn(x, *p)
        resid = target - fitted
        dof = len(x) - len(p)
        rmse = np.sqrt(np.sum(resid ** 2) / len(x))
        if sigma_t is not None:
            w = 1.0 / sigma_t ** 2
            chi2 = np.sum(w * resid ** 2)
            r2 = 1.0 - chi2 / np.sum(w * (target - np.sum(w * target) / np.sum(w)) ** 2)
            print(f"chi2/dof = {chi2 / dof:.4g}")
        else:
            r2 = 1.0 - np.sum(resid ** 2) / np.sum((target - target.mean()) ** 2)
        for name, val, err in zip(<<pnames_list>>, p, se):
            print(f"{name} = {val:.6g} ± {err:.3g}")
        print(f"R² = {r2:.5g}, RMSE = {rmse:.4g} (in <<space>>), n = {len(x)}, dof = {dof}")
        ''', p0=Ctx.py([float(s) for s in start]), pnames=", ".join(model.pnames), absolute=absolute, kw=kw,
                   pnames_list=Ctx.py(model.pnames), space=space))
    # curve + band
    rng = args.get("range")
    xr = f"{_num(rng[0])}, {_num(rng[1])}" if rng else "x.min(), x.max()"
    out = _outp(ctx.tb, args.get("output_csv"), f"fit_{name}.csv")
    if not out.lower().endswith(".csv"):
        out = str(Path(out).with_suffix(".csv"))
    curve_name = ctx.out_name(out, f"fit_{name}.csv")
    rname = ctx.out_name(str(Path(out).with_name(Path(out).stem + "_residuals.csv")), f"fit_{name}_residuals.csv")
    if band in ("none", False):
        band_code = "curve = {'x': xg, 'y_fit': back(fit_fn(xg, *p))}\n"
    else:
        extra = ""
        if band == "prediction":
            if sig is None:
                extra = "var = var + np.sum(resid ** 2) / dof   # new-observation scatter\n"
            else:
                extra = ("o = np.argsort(x)   # new-observation sigma, log-interpolated from the data's sigmas\n"
                         "s_new = np.exp(np.interp(xg, x[o], np.log(sigma_t[o])))\n"
                         + ("var = var + s_new ** 2\n" if absolute else "var = var + chi2 / dof * s_new ** 2\n"))
        band_code = _T('''
            # <<pct>>% <<band>> band from the parameter covariance (numerical Jacobian of the fitted function)
            J = np.empty((len(xg), len(p)))
            for i in range(len(p)):
                h = 1e-6 * abs(p[i]) if p[i] != 0 else 1e-8
                pp, pm = p.copy(), p.copy()
                pp[i] += h
                pm[i] -= h
                J[:, i] = (fit_fn(xg, *pp) - fit_fn(xg, *pm)) / (2 * h)
            var = np.einsum("ij,jk,ik->i", J, pcov, J)
            <<extra>>half = stats.t.ppf(0.5 + <<level>> / 2.0, dof) * np.sqrt(np.clip(var, 0, None))
            yt = fit_fn(xg, *p)
            curve = {"x": xg, "y_fit": back(yt), "<<band>>_lo": back(yt - half), "<<band>>_hi": back(yt + half)}
            ''', pct=int(level * 100), band=band, extra=extra.replace("\n", "\n" + "") if extra else "",
                       level=repr(level))
    code.append("\n" + f"xg = np.linspace({xr}, {npoints})\n" + band_code)
    code.append(f"pd.DataFrame(curve).to_csv(OUT / {curve_name!r}, index=False)\n")
    rcol = "residual" + ("" if space == "y" else f" ({space})")
    code.append(f"pd.DataFrame({{'x': x, 'y': y, 'y_fit': back(fitted), {rcol!r}: resid}}).to_csv(OUT / {rname!r}, "
                "index=False)\n")
    return "\n".join(code)


# ---------------------------------------------------------------------------------------------------------
# plot_all_binaries / plot_all_ternaries
# ---------------------------------------------------------------------------------------------------------
_GRID_AND_DECK = """
# ---- grid images (6 diagrams per image, stitched from the PNGs above as the tool does) and the PowerPoint deck
def make_grid(items, name, dpi=200):
    \"\"\"One image of several diagram PNGs (file names under OUT); a diagram that failed shows why.\"\"\"
    n = len(items)
    ncols = 1 if n == 1 else 2 if n <= 4 else 3 if n <= 9 else 4
    nrows = -(-n // ncols)
    aspect = 0.8
    for _, png, _ in items:
        if (OUT / png).exists():
            h, w = plt.imread(OUT / png).shape[:2]
            aspect = h / w
            break
    cell_w = 16.0 / ncols
    fig, axs = plt.subplots(nrows, ncols, figsize=(16, cell_w * aspect * nrows), squeeze=False)
    for a in axs.flat:
        a.axis("off")
    for a, (label, png, error) in zip(axs.flat, items):
        if (OUT / png).exists():
            a.imshow(plt.imread(OUT / png))
        else:
            a.text(0.5, 0.5, f"{label}\\n\\ncalculation failed:\\n{error[:120]}", ha="center", va="center", fontsize=9,
                   wrap=True, transform=a.transAxes, color="#9b1c1c")
    fig.tight_layout()
    fig.savefig(OUT / name, dpi=dpi, bbox_inches="tight")
    return fig


GRIDS = <<grids>>   # (image, slide title, slide subtitle, [(diagram, PNG, error of a failed calculation)])
grid_figures = [make_grid(items, name) for name, _, _, items in GRIDS]

try:            # the deck: title slide, one slide per grid image<<indiv_note>> (python-pptx)
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print("python-pptx is not installed: <<pptx_name>> is not written")
else:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def slide(title, subtitle=""):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tf = s.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(1.0)).text_frame
        tf.text = title
        tf.paragraphs[0].runs[0].font.size, tf.paragraphs[0].runs[0].font.bold = Pt(26), True
        if subtitle:
            p = tf.add_paragraph()
            p.text = subtitle
            p.runs[0].font.size = Pt(14)
        return s


    def picture(s, png, top=Inches(1.5)):
        h, w = plt.imread(png).shape[:2]
        scale = min((prs.slide_width - Inches(0.8)) / w, (prs.slide_height - top - Inches(0.25)) / h)
        s.shapes.add_picture(str(png), int((prs.slide_width - w * scale) / 2), int(top), int(w * scale),
                             int(h * scale))


    slide(<<title>>, <<subtitle>>)
    for name, title, subtitle, _ in GRIDS:
        picture(slide(title, subtitle), OUT / name)
    for title, png in <<individual>>:
        if (OUT / png).exists():
            picture(slide(title, <<tdb_name>>), OUT / png)
    prs.save(OUT / <<pptx_name>>)
fig = grid_figures[<<first>>]   # the grid image this notebook was made for
ax = fig.axes[0]
"""

_COLOR_BINARY = """
    # pycalphad binplot, as the tool's color style: <<lo>>-<<hi>> °C, x step <<xs>>, T step <<ts>> K
    a, b = <<a>>, <<b>>
    phases = <<ph>>
    fig, ax = plt.subplots(figsize=(6.4, 4.8))
    strategy = BinaryStrategy(db, [a, b, "VA"], phases, {v.X(b): (0, 1, <<xs>>), v.T: (<<tlo>>, <<thi>>, <<ts>>),
                                                         v.P: 101325.0, v.N: 1})
    recipe = {"kind": "binary", "elements": [a, b], "phases": phases, "P": 101325.0, "T_range": [<<tlo>>, <<thi>>],
              "x_range": [0.0, 1.0], "tmap": [1.0, 0.0], "T_unit": "C", "comp": {"scale": 1.0, "weight": False}}
    # mapping check / repair as the tool (mapping_check), then pycalphad's plot (= binplot)
    print(f"{a}-{b}:", map_and_check(strategy, db, recipe, <<mode>>, 600.0, <<grid_n>>)["text"])
    plot_binary(strategy, ax=ax, tielines=4, tieline_color=(0.55, 0.55, 0.55, 0.35))
    ax.grid(False)
    ax.set_xlim(0, 1)
    ax.set_ylim(<<tlo>>, <<thi>>)
    kelvin_to_celsius(ax)
    ax.set_xlabel(f"Mole fraction {b.capitalize()}")
    ax.set_title(<<title>>, fontsize=12)
    fig.savefig(OUT / <<png>>, dpi=150, bbox_inches="tight")
    plt.close(fig)
    """

_COLOR_TERNARY = """
    # pycalphad ternplot, as the tool's color style: <<Tc>> °C, composition step <<step>>
    a, b, c = <<abc>>
    phases = <<ph>>
    T = <<TK>>
    fig = plt.figure(figsize=(6.4, 5.6))
    ax = fig.add_subplot(projection="triangular")
    strategy = TernaryStrategy(db, [a, b, c, "VA"], phases, {v.X(b): (0, 1, <<step>>), v.X(c): (0, 1, <<step>>), v.T: T,
                                                             v.P: 101325.0, v.N: 1})
    recipe = {"kind": "ternary", "elements": [a, b, c], "phases": phases, "P": 101325.0, "T_K": T, "T_unit": "C",
              "comp": {"scale": 1.0, "weight": False}}
    # mapping check / repair as the tool (mapping_check), then pycalphad's plot (= ternplot)
    print(f"{a}-{b}-{c} <<Tc>> °C:", map_and_check(strategy, db, recipe, <<mode>>, 600.0, <<grid_n>>)["text"])
    plot_ternary(strategy, x=v.X(b), y=v.X(c), ax=ax, tielines=4, tieline_color=(0.55, 0.55, 0.55, 0.35))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    <<corners>>
    ax.set_title(<<title>>, fontsize=12)
    fig.savefig(OUT / <<png>>, dpi=150, bbox_inches="tight")
    plt.close(fig)
    """


def _overview_status(rec_text: str) -> Dict[str, str]:
    """name -> error of the subsystems the tool reported as FAILED ("  NAME: FAILED error (12.3 s)")."""
    out = {}
    for m in re.finditer(r"^  (.+?): FAILED (.*) \([\d.]+ s\)", rec_text or "", re.M):
        out[m.group(1)] = m.group(2)
    return out


def _overview_cells(ctx: Ctx, kind: str, a: Dict, tasks, grids, pptx: str, title_slide, individual) -> list:
    """Cells of an overview: one per subsystem (as the tool's worker computed it: map_phase_diagram in the shaded
    style, pycalphad binplot / ternplot in the color style), then the grid images and the deck."""
    from .orientation import corner_label_code
    from .mapping_check import grid_size
    failed = _overview_status(ctx.record.get("result") or "")
    cells = []
    for t in tasks:
        kw = dict(t["kwargs"])
        png = kw.pop("out_png")
        if t["name"] in failed:
            cells.append((t["name"], f"# {t['name']}: the tool's calculation failed ({failed[t['name']]}); "
                                     f"the grid image shows that\n"))
            continue
        if t["kind"] == "map":          # shaded: exactly map_phase_diagram (PNG, SVG, PDF, CSVs, JSON)
            args = {k: val for k, val in kw.items() if val is not None}
            args["output_prefix"] = str(Path(png).with_suffix(""))
            code = _gen_map_phase_diagram(args, ctx)
            cells.append((t["name"], code.rstrip("\n") + "\nplt.close(fig)\n"))
            continue
        name = ctx.out_name(png, Path(png).name)
        db = _db(ctx, kw["tdb"])
        from . import diagrams
        ctx.need("import numpy as np", "import matplotlib.pyplot as plt", "from pycalphad import variables as v")
        ctx.mapped()
        mode = repr(check_mode(kw.get("check_mapping")))
        if t["kind"] == "binary":
            ctx.need("from pycalphad.mapping import BinaryStrategy, plot_binary",
                     "from matplotlib.collections import LineCollection")
            A, B = kw["elements"]
            ph = diagrams.phases_for(db, [A, B], exclude=kw.get("exclude_phases"))
            tlo, thi = float(kw["tmin_c"]) + KELVIN, float(kw["tmax_c"]) + KELVIN
            note = "" if diagrams.has_interaction(db, [A, B]) else "  (no binary parameters in TDB)"
            cells.append((t["name"], _T(_COLOR_BINARY, a=repr(A), b=repr(B), ph=Ctx.py(ph, 9), xs=_num(kw["x_step"]),
                                        ts=_num(kw["t_step"]), tlo=_num(tlo), thi=_num(thi),
                                        lo=f"{kw['tmin_c']:g}", hi=f"{kw['tmax_c']:g}", mode=mode,
                                        grid_n=Ctx.py(tuple(grid_size("binary", labels=False))),
                                        title=repr(f"{A.capitalize()}–{B.capitalize()}{note}"), png=repr(name))))
        else:
            ctx.need("from pycalphad.mapping import TernaryStrategy, plot_ternary",
                     "from pycalphad.plot import triangular  # registers the 'triangular' projection")
            els = list(kw["elements"])
            ph = diagrams.phases_for(db, els, exclude=kw.get("exclude_phases"))
            TK = float(kw["temperature"]) + (KELVIN if str(kw.get("units", "C")).upper() == "C" else 0.0)
            cells.append((t["name"], _T(_COLOR_TERNARY, abc=Ctx.py(tuple(els)), ph=Ctx.py(ph, 9), TK=_num(TK),
                                        Tc=f"{TK - KELVIN:.0f}", step=_num(kw["step"]), mode=mode,
                                        grid_n=Ctx.py(tuple(grid_size("ternary", labels=False))),
                                        corners=corner_label_code(els).strip("\n"),
                                        title=repr("–".join(e.capitalize() for e in els)
                                                   + f"   {TK - KELVIN:.0f} °C ({TK:.0f} K)"), png=repr(name))))
    if any(t["kind"] != "map" for t in tasks):     # color style: helpers once, before the first subsystem
        head = _load_db_code(ctx, a["tdb"]) + "KELVIN = 273.15\n" + mapping_check_helpers(ctx) + _T(_TO_CELSIUS)
        if cells:
            cells[0] = (cells[0][0], head + "\n\n" + cells[0][1])
    target = _target_path(ctx)
    first = next((k for k, g in enumerate(grids) if target and str(Path(g[0]).resolve()) == target), 0)
    grid_lit = [(ctx.out_name(g, Path(g).name), title, sub,
                 [(n, ctx.out_name(png, Path(png).name), failed.get(n, "")) for n, png in items])
                for g, title, sub, items in grids]
    pptx_name = ctx.out_name(pptx, Path(pptx).name)
    ctx.need("import matplotlib.pyplot as plt")
    tdb_name = Path(ctx.tb._find_tdb(a["tdb"])).name
    cells.append(("grid images and deck", _T(
        _GRID_AND_DECK, grids=Ctx.py(grid_lit, 8), pptx_name=repr(pptx_name), title=repr(title_slide[0]),
        subtitle=repr(title_slide[1]), individual=Ctx.py([(t, ctx.out_name(png, Path(png).name))
                                                          for t, png in individual], 17),
        indiv_note=", one per diagram" if individual else "", tdb_name=repr(tdb_name), first=first)))
    return cells


@generator("plot_all_binaries")
def _gen_plot_all_binaries(args, ctx: Ctx) -> list:
    from . import diagrams
    a = {"tmin_c": 100, "tmax_c": 3000, **args}
    tdb_path = ctx.tb._find_tdb(a["tdb"])
    stem = Path(tdb_path).stem
    pptx = _outp(ctx.tb, a.get("output_pptx"), f"binaries_{stem}.pptx")
    tasks, out_dir, els, pairs = diagrams.binary_tasks(
        tdb_path, pptx, a["tmin_c"], a["tmax_c"], elements=a.get("elements"), element_order=a.get("element_order"),
        orientation=a.get("orientation"), style=_style(ctx, a), check_mapping=a.get("check_mapping") or "auto")
    span = f"{a['tmin_c']:g}–{a['tmax_c']:g} °C"
    names = [(t["name"], t["kwargs"]["out_png"]) for t in tasks]
    chunks = [names[i:i + 6] for i in range(0, len(names), 6)]
    grids = [(str(out_dir / f"{stem}_binaries_grid{i + 1}.png"),
              "Binary phase diagrams" + (f" ({i + 1})" if len(names) > 6 else ""),
              f"{Path(tdb_path).name}   •   {span}", ch) for i, ch in enumerate(chunks)]
    title = (f"Binary phase diagrams: {Path(tdb_path).name}",
             f"{len(pairs)} binaries of {'-'.join(e.capitalize() for e in els)}   •   {span}, 1 atm")
    indiv = [(n.replace("-", "–").title(), png) for n, png in names] if a.get("individual_slides") else []
    return _overview_cells(ctx, "binary", a, tasks, grids, pptx, title, indiv)


@generator("plot_all_ternaries")
def _gen_plot_all_ternaries(args, ctx: Ctx) -> list:
    from . import diagrams
    tdb_path = ctx.tb._find_tdb(args["tdb"])
    stem = Path(tdb_path).stem
    pptx = _outp(ctx.tb, args.get("output_pptx"), f"ternaries_{stem}.pptx")
    temps = list(args["temperatures_c"])
    tasks, out_dir, els, triples = diagrams.ternary_tasks(
        tdb_path, temps, pptx, elements=args.get("elements"), element_order=args.get("element_order"),
        orientation=args.get("orientation"), style=_style(ctx, args),
        check_mapping=args.get("check_mapping") or "auto")
    grids = []
    for T in temps:
        names = [(t["name"], t["kwargs"]["out_png"]) for t in tasks if t["T"] == T]
        for i in range(0, len(names), 6):
            grids.append((str(out_dir / f"{stem}_ternaries_{T:g}C_grid{i // 6 + 1}.png"),
                          f"Ternary sections at {T:g} °C" + (f" ({i // 6 + 1})" if len(names) > 6 else ""),
                          Path(tdb_path).name, names[i:i + 6]))
    title = (f"Ternary isothermal sections: {Path(tdb_path).name}",
             f"{len(triples)} ternaries of {'-'.join(e.capitalize() for e in els)}   •   "
             f"T = {', '.join(f'{t:g}' for t in temps)} °C")
    indiv = [(t["name"], t["kwargs"]["out_png"]) for t in tasks] if args.get("individual_slides") else []
    return _overview_cells(ctx, "ternary", args, tasks, grids, pptx, title, indiv)


# ---------------------------------------------------------------------------------------------------------
# ESPEI plots
# ---------------------------------------------------------------------------------------------------------
@generator("plot_diagram_with_data")
def _gen_plot_diagram_with_data(args, ctx: Ctx) -> str:
    from .diagrams import phases_for
    from .espei_plot_tools import _load_datasets
    from .orientation import binary_order, orientation_text, ternary_order
    a = {"tmin": 300, "tmax": 2500, "temperature_unit": "K", **args}
    shaded = _style(ctx, a) == "shaded"
    if a.get("label_regions") is None:      # the tool's default: labels on when shaded (no phase legend)
        a["label_regions"] = shaded and bool(a.get("tdb"))
    els = sorted(e.upper() for e in a["elements"])
    if len(els) == 3:
        return _gen_ternary_with_data(a, ternary_order(els, a.get("corners"), a.get("right")), ctx)
    A, B = binary_order(els, a.get("right"))
    celsius = _is_celsius(a["temperature_unit"])
    off = KELVIN if celsius else 0.0
    lo, hi = float(a["tmin"]) + off, float(a["tmax"]) + off
    datasets, _ = _load_datasets(ctx.tb.p(a["datasets_dir"]))
    data_phases = {p for d in datasets.all() for p in d.get("phases", [])}
    if a.get("phases"):
        ph = [p.upper() for p in a["phases"]]
    elif a.get("tdb"):
        ph = phases_for(_db(ctx, a["tdb"]), [A, B])
    else:
        ph = sorted(data_phases)
    png = _outp(ctx.tb, a.get("output_png"), f"diagram_data_{A}-{B}.png")
    png_name = ctx.out_name(png, "diagram_data.png")
    ctx.need("import numpy as np", "import matplotlib.pyplot as plt", "from pycalphad import variables as v",
             "from pycalphad.plot.utils import phase_legend", "from espei.datasets import load_datasets, recursive_glob",
             "from espei.plot import dataplot")
    code = ""
    if a.get("tdb"):
        ctx.need("from pycalphad import Database")
        code += _load_db_code(ctx, a["tdb"])
    ddir = ctx.path(a["datasets_dir"])
    code += _T('''
        KELVIN = 273.15
        comps = <<comps>>
        phases = <<ph>>
        # <<orient>>: binplot maps X(<<Bc>>), and dataplot converts every ZPF point to X(<<Bc>>) by mass balance
        conds = {v.P: 101325, v.T: (<<lo>>, <<hi>>, 10), v.X(<<B>>): (0, 1, 0.02), v.N: 1}
        datasets = load_datasets(recursive_glob(str(<<ddir>>), "*.json"))   # ESPEI datasets (ZPF data are plotted)

        # one colour per phase, shared by the calculated lines and the data markers
        data_phases = {p for d in datasets.all() for p in d.get("phases", [])}
        handles, phase_colors = phase_legend(sorted(set(phases) | data_phases))


        class PhaseColors(dict):   # a second composition set (BCC_A2#2) gets its phase's colour
            def __missing__(self, key):
                return self[key.split("#")[0]]


        def legend(phase_names):
            want = {str(p) for p in phase_names}
            return [h for h in handles if h.get_label() in want], PhaseColors(phase_colors)


        fig, ax = plt.subplots(figsize=(7.5, 5.2))
        ''', comps=Ctx.py([A, B, "VA"]), ph=Ctx.py(ph), lo=_num(lo), hi=_num(hi), B=repr(B), ddir=ddir,
               orient=orientation_text("binary", [A, B])[13:], Bc=B)
    if a.get("tdb"):            # the tool's mapping + check (mapping_check)
        from .mapping_check import grid_size
        ctx.mapped()
        crec = {"kind": "binary", "elements": [A, B], "phases": list(ph), "P": 101325.0, "rename_disordered": True,
                "composition_step": 0.02, "temperature_step": 10.0, "T_unit": "C" if celsius else "K",
                "tmap": [1.0, 0.0], "T_range": [lo, hi], "x_range": [0.0, 1.0], "comp": {"scale": 1.0, "weight": False}}
        ctx.need("from pycalphad.mapping import BinaryStrategy" + ("" if shaded else ", plot_binary"))
        code = code.replace("KELVIN = 273.15\n", "KELVIN = 273.15\n" + mapping_check_helpers(ctx), 1)
        code += "strategy = BinaryStrategy(db, comps, phases, conds)\n" + check_code(
            ctx, crec, check_mode(a.get("check_mapping")), 600.0, grid_size("binary", bool(a.get("label_regions"))))
    if a.get("tdb") and shaded:
        code = code.replace("KELVIN = 273.15\n", "KELVIN = 273.15\n" + shaded_helpers(ctx), 1)
        code += ("# calculated diagram, shaded style: black boundaries, two-phase fields filled aliceblue, no tie-lines\n"
                 f"draw_shaded(ax, *strategy_boundaries(strategy, v.X({B!r}), v.T), kind='binary')\n"
                 "# data markers keep their phase colours; the legend lists only the references (the phases are "
                 "named by the region labels)\n"
                 "dataplot(comps, phases, conds, datasets, ax=ax, legend_generator=lambda names: ([], legend(names)[1]))\n")
    else:
        if a.get("tdb"):
            code += ("plot_binary(strategy, ax=ax, tielines=4, tieline_color=(0.55, 0.55, 0.55, 0.35), "
                     "legend_generator=legend)   # = pycalphad binplot\nax.grid(False)\n")
        code += "dataplot(comps, phases, conds, datasets, ax=ax, legend_generator=legend)\n"
    code += f"ax.set_xlim(0, 1)\nax.set_ylim({_num(lo)}, {_num(hi)})\n"
    if celsius:
        ctx.need("from matplotlib.collections import LineCollection")
        code = code.replace("KELVIN = 273.15\n", "KELVIN = 273.15\n" + _T(_TO_CELSIUS) + "\n\n", 1)
        code += "kelvin_to_celsius(ax)   # temperatures as data in °C\n"
    elif shaded:
        code += "ax.set_ylabel('Temperature (K)')\n"
    if shaded:
        if not a.get("tdb"):
            code = code.replace("KELVIN = 273.15\n", "KELVIN = 273.15\n" + shaded_helpers(ctx), 1)
        code += _T(_LEGEND_INSIDE, loc='"best"')
    title = f"{A.capitalize()}–{B.capitalize()}" + (f"  ({Path(a['tdb']).name} + data)" if a.get("tdb") else "  (data)")
    code += (f"ax.set_xlabel({'Mole fraction ' + B.capitalize()!r})\nax.set_title({title!r})\n")
    if a.get("label_regions") and a.get("tdb"):
        db = _db(ctx, a["tdb"])
        rec = {"kind": "binary", "tdb": ctx.tb._find_tdb(a["tdb"]), "elements": [A, B], "phases": list(ph),
               "P": 101325.0, "rename_disordered": True, "composition_step": 0.02, "temperature_step": 10.0,
               "T_unit": "C" if celsius else "K", "tmap": [1.0, KELVIN if celsius else 0.0],   # data in °C here
               "T_range": [lo, hi], "x_range": [0.0, 1.0], "comp": {"scale": 1.0, "weight": False}}
        code += (label_helpers(ctx) + "lines, invariants = strategy_lines(strategy, db, check_recipe)   # the checked "
                 "mapping's fields\n"
                 + label_code(ctx, rec, db, None, mapped=True, phase_labels=a.get("phase_labels"),
                              style=_label_style(a.get("label_style"), shaded), check_expr="mapping_check"))
    code += f"fig.savefig(OUT / {png_name!r}, dpi=150, bbox_inches='tight')\n"
    fmts = a.get("formats")
    if fmts is None:
        fmts = ["svg", "pdf"] if a.get("label_regions") or shaded else []
    return code + _vector_saves(ctx, png, fmts)


_LEGEND_INSIDE = '''
    ticks_in(ax)   # shaded style: minor ticks, all ticks pointing in on all four sides
    ax.grid(False)
    leg = ax.get_legend()   # the data legend goes inside the axes, where it hides no data
    if leg is not None:
        leg_handles, leg_texts = list(leg.legend_handles), [t.get_text() for t in leg.get_texts()]
        leg.remove()
        if leg_handles:
            ax.legend(leg_handles, leg_texts, loc=<<loc>>, fontsize=8, framealpha=0.85)
    '''

_TERN_DATA = '''
    KELVIN = 273.15
    elements = <<els>>   # <<orient>>
    phases = <<ph>>
    T_K = <<T>>          # isotherm; ZPF points within <<tol>> K of it are shown
    datasets = load_datasets(recursive_glob(str(<<ddir>>), "*.json"))   # ESPEI datasets


    def zpf_points(rows):
        """ZPF equilibria of this ternary near T_K as complete compositions: the datasets store only the
        independent components, so the missing (balance) element is 1 - sum; unknown compositions stay None."""
        out = []
        for d in rows:
            if d.get("output") != "ZPF" or {c for c in d["components"] if c != "VA"} != set(elements):
                continue
            Ts = np.atleast_1d(np.asarray(d["conditions"]["T"], dtype=float))
            Ts = np.full(len(d["values"]), Ts[0]) if Ts.size == 1 else Ts
            for row, T in zip(d["values"], Ts):
                if abs(T - T_K) > <<tol>>:
                    continue
                eq = []
                for phase, comps_, xs in row:
                    if phase == "__HYPERPLANE__":
                        continue
                    if any(x is None for x in xs):
                        eq.append((phase, None))
                        continue
                    X = dict(zip(comps_, map(float, xs)))
                    missing = [e for e in elements if e not in X]
                    if len(missing) == 1:
                        X[missing[0]] = 1.0 - sum(X.values())
                    eq.append((phase, X if len(missing) <= 1 else None))
                out.append({"reference": d.get("reference", ""), "phases": eq})
        return out


    equilibria = zpf_points(datasets.all())
    data_phases = sorted({p for eq in equilibria for p, _ in eq["phases"]})
    handles, phase_colors = phase_legend(sorted(set(phases) | set(data_phases)))


    def legend(phase_names):
        want = {str(p) for p in phase_names}
        return [h for h in handles if h.get_label() in want], phase_colors


    fig = plt.figure(figsize=(7.4, 6.0))
    ax = fig.add_subplot(projection="triangular")
    shown = set(data_phases)
    '''

_TERN_MAP = '''
    # calculated isotherm: x axis = X(<<B>>) (bottom-right corner), y axis = X(<<C>>) (top corner)
    strategy = TernaryStrategy(db, elements + ["VA"], phases, {v.X(<<Br>>): (0, 1, 0.02), v.X(<<Cr>>): (0, 1, 0.02),
                                                               v.T: T_K, v.P: 101325, v.N: 1})
    '''

_TERN_CALC = '''
    plot_ternary(strategy, x=v.X(<<Br>>), y=v.X(<<Cr>>), ax=ax, tielines=4, tieline_color=(0.55, 0.55, 0.55, 0.35),
                 legend_generator=legend)   # = pycalphad ternplot
    shown |= {cs.phase_record.phase_name for z in strategy.zpf_lines for p in z.points for cs in p.stable_composition_sets}
    '''

_TERN_CALC_SHADED = '''
    # shaded style: grey boundaries and tie-triangle edges, aliceblue two-phase fields
    draw_shaded(ax, *strategy_boundaries(strategy, v.X(<<Br>>), v.X(<<Cr>>)), kind="ternary")
    '''

_TERN_DATA_PLOT = '''
    # data: the triangular axes take (x, y) = (X(bottom-right element), X(top element)); one marker per reference
    refs = sorted({eq["reference"] for eq in equilibria})
    markers = bib_marker_map(refs)
    for eq in equilibria:
        pts = [(p, X[elements[1]], X[elements[2]]) for p, X in eq["phases"] if X is not None]
        m = markers[eq["reference"]]["markers"]
        for p, x, y in pts:
            ax.plot(x, y, marker=m["marker"], fillstyle=m["fillstyle"], linestyle="", color=phase_colors[p],
                    markersize=5.0, markeredgewidth=0.6, zorder=4)
        if len(pts) == 2:     # tie-line
            ax.plot([pts[0][1], pts[1][1]], [pts[0][2], pts[1][2]], color="k", linewidth=0.8, zorder=3.5)
        elif len(pts) == 3:   # tie-triangle
            loop = pts + pts[:1]
            ax.plot([p[1] for p in loop], [p[2] for p in loop], color="r", linewidth=0.8, linestyle="--", zorder=3.5)
        # a single point: broken tie-line (the other phase's composition is unknown)
    # legend: the data references (and, without a region-labelled calculation, the phase colours)
    leg_handles = <<phase_handles>> + [
        mlines.Line2D([], [], linestyle="", color="k", marker=markers[r]["markers"]["marker"],
                      fillstyle=markers[r]["markers"]["fillstyle"], markersize=6, label=markers[r]["formatted"])
        for r in refs]
    if leg_handles:
        ax.legend(handles=leg_handles, <<legend_kw>>)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.set_xlabel(<<xl>>)
    ax.set_ylabel(<<yl>>)
    ax.yaxis.label.set_rotation(60)
    ax.yaxis.set_label_coords(x=0.12, y=0.5)
    '''


def _gen_ternary_with_data(a, els, ctx: Ctx) -> str:
    from .diagrams import phases_for
    from .espei_plot_tools import _load_datasets
    from .orientation import corner_label_code, orientation_text, zpf_equilibria
    A, B, C = els
    celsius = _is_celsius(a["temperature_unit"])
    off = KELVIN if celsius else 0.0
    if a.get("temperature") is None:
        ctx.fallback("ternary without temperature")
    T_K = float(a["temperature"]) + off
    tol = float(a.get("T_tolerance") or 10.0)
    datasets, _ = _load_datasets(ctx.tb.p(a["datasets_dir"]))
    data_phases = {p for eq in zpf_equilibria(datasets, els, T_K, tol) for p, _ in eq["phases"]}
    if a.get("phases"):
        ph = [p.upper() for p in a["phases"]]
    elif a.get("tdb"):
        ph = phases_for(_db(ctx, a["tdb"]), els)
    else:
        ph = sorted(data_phases)
    png = _outp(ctx.tb, a.get("output_png"), f"diagram_data_{A}-{B}-{C}_{T_K - off:g}{'C' if celsius else 'K'}.png")
    png_name = ctx.out_name(png, "diagram_data.png")
    ctx.need("import numpy as np", "import matplotlib.pyplot as plt", "import matplotlib.lines as mlines",
             "from pycalphad import variables as v", "from pycalphad.plot.utils import phase_legend",
             "from pycalphad.plot import triangular  # registers the 'triangular' projection",
             "from espei.datasets import load_datasets, recursive_glob", "from espei.utils import bib_marker_map")
    code = ""
    if a.get("tdb"):
        ctx.need("from pycalphad import Database")
        code += _load_db_code(ctx, a["tdb"])
    code += _T(_TERN_DATA, els=Ctx.py(list(els)), orient=orientation_text("ternary", els)[13:], ph=Ctx.py(ph),
               T=_num(T_K), tol=_num(tol), ddir=ctx.path(a["datasets_dir"]))
    shaded = _style(ctx, a) == "shaded"
    if a.get("tdb"):            # the tool's mapping + check (mapping_check)
        from .mapping_check import grid_size
        ctx.mapped()
        crec = {"kind": "ternary", "elements": list(els), "phases": list(ph), "P": 101325.0, "rename_disordered": True,
                "composition_step": 0.02, "temperature_step": 10.0, "T_unit": "C" if celsius else "K", "T_K": T_K,
                "comp": {"scale": 1.0, "weight": False}}
        ctx.need("from pycalphad.mapping import TernaryStrategy" + ("" if shaded else ", plot_ternary"))
        code = code.replace("KELVIN = 273.15\n", "KELVIN = 273.15\n" + mapping_check_helpers(ctx), 1)
        code += _T(_TERN_MAP, B=B, C=C, Br=repr(B), Cr=repr(C)) + check_code(
            ctx, crec, check_mode(a.get("check_mapping")), 600.0, grid_size("ternary", bool(a.get("label_regions"))))
    if a.get("tdb") and shaded:
        code = code.replace("KELVIN = 273.15\n", "KELVIN = 273.15\n" + shaded_helpers(ctx), 1)
        code += _T(_TERN_CALC_SHADED, Br=repr(B), Cr=repr(C))
    elif a.get("tdb"):
        code += _T(_TERN_CALC, Br=repr(B), Cr=repr(C))
    code += _T(_TERN_DATA_PLOT, xl=repr(f"Mole fraction {B.capitalize()}"),
               yl=repr(f"Mole fraction {C.capitalize()}"),
               phase_handles=("[]" if shaded and a.get("tdb")
                              else "legend(sorted(shown))[0]"),
               legend_kw=('loc="upper right", fontsize=8, framealpha=0.85' if shaded else
                          'loc="center left", bbox_to_anchor=(1.02, 0.5), fontsize=8, frameon=False'))
    code += corner_label_code(els)
    title = (f"{A.capitalize()}–{B.capitalize()}–{C.capitalize()}   {T_K - KELVIN:.0f} °C ({T_K:.0f} K)"
             + (f"  ({Path(a['tdb']).name} + data)" if a.get("tdb") else "  (data)"))
    code += f"ax.set_title({title!r}, fontsize=11)\n"
    if a.get("label_regions") and a.get("tdb"):
        db = _db(ctx, a["tdb"])
        rec = {"kind": "ternary", "tdb": ctx.tb._find_tdb(a["tdb"]), "elements": list(els), "phases": list(ph),
               "P": 101325.0, "rename_disordered": True, "composition_step": 0.02, "temperature_step": 10.0,
               "T_unit": "C" if celsius else "K", "T_K": T_K, "comp": {"scale": 1.0, "weight": False}}
        code += (label_helpers(ctx) + "lines, invariants = strategy_lines(strategy, db, check_recipe)   # the checked "
                 "mapping's fields\n"
                 + label_code(ctx, rec, db, None, mapped=True, phase_labels=a.get("phase_labels"),
                              style=_label_style(a.get("label_style"), shaded), check_expr="mapping_check"))
    code += f"fig.savefig(OUT / {png_name!r}, dpi=150, bbox_inches='tight')\n"
    fmts = a.get("formats")
    if fmts is None:
        fmts = ["svg", "pdf"] if a.get("label_regions") or shaded else []
    return code + _vector_saves(ctx, png, fmts)


@generator("plot_parameter_vs_data")
def _gen_plot_parameter_vs_data(args, ctx: Ctx) -> str:
    phase = args["phase"].upper()
    comps = sorted({e.upper() for e in args["elements"]} | {"VA"})
    cfg = [[s.upper() for s in c] if isinstance(c, list) else c.upper() for c in args["configuration"]]
    cfg_t = tuple(tuple(c) if isinstance(c, list) else c for c in cfg)
    is_interaction = any(isinstance(c, list) and len(c) > 1 for c in cfg)
    output = args["output"].upper()
    label = ":".join(",".join(c) if isinstance(c, list) else c for c in cfg)
    png = _outp(ctx.tb, args.get("output_png"), f"{output}_{phase}_{label.replace(':', '-').replace(',', '')}.png")
    png_name = ctx.out_name(png, "parameter.png")
    fn = "plot_interaction" if is_interaction else "plot_endmember"
    ctx.need("import matplotlib.pyplot as plt", f"from espei.plot import {fn}")
    code = _load_db_code(ctx, args["tdb"])
    if args.get("datasets_dir"):
        ctx.need("from espei.datasets import load_datasets, recursive_glob")
        code += f"datasets = load_datasets(recursive_glob(str({ctx.path(args['datasets_dir'])}), '*.json'))\n"
    else:
        code += "datasets = None\n"
    code += _T('''
        # <<what>> <<phase>>(<<label>>) <<output>> from the TDB vs ESPEI's thermochemical data
        fig, ax = plt.subplots(figsize=(7, 4.8))
        <<fn>>(db, <<comps>>, <<phase>>, <<cfg>>, <<output>>, datasets=datasets, ax=ax)
        fig.savefig(OUT / <<png>>, dpi=150, bbox_inches="tight")
        ''', what="interaction" if is_interaction else "endmember", phase=repr(phase), label=label,
               output=repr(output), fn=fn, comps=Ctx.py(comps), cfg=Ctx.py(cfg_t), png=repr(png_name))
    return code


def _target_path(ctx: Ctx) -> Optional[str]:
    """Resolved path of the output this cell is generated for (the show_code target, or the parent file being
    reproduced); None when the target was a call number.  Used by tools that write several images, so the cell
    leaves ``fig`` / ``ax`` on the requested one."""
    return ctx.target


# ---------------------------------------------------------------------------------------------------------
# export_phase_model / write_database
# ---------------------------------------------------------------------------------------------------------
_MODEL_BODY = '''
CONTRIBUTIONS = {"ref": "reference (surface of reference / endmember) energy", "idmix": "ideal mixing energy",
                 "xsmix": "excess mixing energy (Redlich-Kister / Muggianu)", "mag": "magnetic energy",
                 "2st": "two-state energy", "ein": "Einstein energy", "vol": "volume (pressure) energy",
                 "ord": "atomic ordering energy (partitioned order/disorder model)"}
<<model_line>>
sfs = mod.site_fractions
readable = {y: symengine.Symbol(f"y{y.sublattice_index}_{y.species.name}") for y in sfs}   # Y(PH,s,EL) -> ys_EL
contributions = {k: e for k, e in mod.models.items() if str(e) not in ("0", "0.0")}
physical = {}
for name in ("TC", "BMAG", "NT"):
    try:
        e = getattr(mod, name)
    except Exception:
        continue
    if str(e) not in ("0", "0.0"):
        physical[name] = e
<<functions>>
# parameters of this phase (and of its disordered part) within the components
ph = db.phases[phase]
names = {phase} | ({ph.model_hints["disordered_phase"]} if ph.model_hints.get("ordered_phase") == phase and
                   ph.model_hints.get("disordered_phase") else set())
ok = set(comps) | {"*"}
rows = []
for p in db._parameters.all():
    if p["phase_name"] not in names or not all(s.name in ok or set(s.constituents) <= set(comps)
                                               for subl in p["constituent_array"] for s in subl):
        continue
    cons = ":".join(",".join(sorted(s.name for s in subl)) for subl in p["constituent_array"])
    rows.append({"parameter": f"{p['parameter_type']}({p['phase_name']},{cons};{p['parameter_order']})",
                 "phase": p["phase_name"], "part": "disordered" if p["phase_name"] != phase else "",
                 "type": p["parameter_type"], "constituents": cons, "order": p["parameter_order"],
                 "expression": str(p["parameter"]), "reference": p.get("reference") or ""})
rows.sort(key=lambda r: (r["phase"], r["type"], r["constituents"].count(",") + r["constituents"].count(":"),
                         r["constituents"], r["order"]))
pd.DataFrame(rows, columns=["parameter", "phase", "part", "type", "constituents", "order", "expression",
                            "reference"]).to_csv(OUT / <<pcsv>>, index=False)

# text file: sublattice model and every non-zero Gibbs-energy contribution
lines = [f"Phase {phase} from <<tdbname>>; components {', '.join(comps)}",
         f"Sublattice site ratios: {list(ph.sublattices)}",
         "Constituents (active): " + " : ".join(",".join(sorted(s.name for s in c)) for c in mod.constituents),
         f"Model hints: {ph.model_hints}",
         "Site fractions: " + ", ".join(f"y{y.sublattice_index}_{y.species.name} = {y}" for y in sfs),
         "Energies are per mole of formula unit normalised to one mole of atoms (J/mol); GM = sum of contributions.",
         <<fnote>>, ""]
for k, e in contributions.items():
    lines += [f"=== {k}: {CONTRIBUTIONS.get(k, k)}", str(e.xreplace(readable) if readable else e), ""]
for k, e in physical.items():
    lines += [f"=== physical property {k}", str(e.xreplace(readable) if readable else e), ""]
if used:
    lines += ["=== FUNCTIONS"] + [f"{s} = {db.symbols[s]}" for s in used] + [""]
'''


@generator("export_phase_model")
def _gen_export_phase_model(args, ctx: Ctx) -> str:
    import numpy as np  # noqa: F401
    from . import mapping_tools as mt
    a = {"expand_functions": False, "latex": False, **args}
    db = _db(ctx, a["tdb"])
    phase = a["phase"].upper()
    if phase not in db.phases:
        ctx.fallback("unknown phase")
    ph = db.phases[phase]
    phase_els = sorted({el for subl in ph.constituents for s in subl for el in s.constituents} - {"VA", "/-"})
    els = mt._check_elements(db, a["elements"]) if a.get("elements") else phase_els
    comps = sorted(set(els)) + ["VA"]
    tag = f"{phase}_{'-'.join(_el(e) for e in a['elements']) if a.get('elements') else 'all'}"
    prefix = _outp(ctx.tb, a.get("output_prefix"), f"model_{tag}_{Path(ctx.tb._find_tdb(a['tdb'])).stem}")
    txt = ctx.out_name(prefix + ".txt", "model.txt")
    pcsv = ctx.out_name(prefix + "_parameters.csv", "model_parameters.csv")
    ctx.need("import symengine", "import pandas as pd", "from pycalphad import Database, Model",
             "from pycalphad import variables as v")
    expand = bool(a["expand_functions"])
    code = _load_db_code(ctx, a["tdb"]) + f"phase = {phase!r}\ncomps = {Ctx.py(comps)}\n"
    model_line = ("mod = Model(db, comps, phase, parameters=[])   # FUNCTIONs substituted" if expand else
                  "mod = Model(db, comps, phase, parameters=list(db.symbols))   # FUNCTIONs kept as symbols")
    functions = ("used = []\n" if expand else _T('''

        # FUNCTIONs referenced by the model (recursively)
        used, stack = [], [str(s) for s in mod.ast.free_symbols] + [str(s) for e in physical.values()
                                                                   for s in e.free_symbols]
        while stack:
            s = stack.pop()
            if s in db.symbols and s not in used:
                used.append(s)
                stack += [str(x) for x in db.symbols[s].free_symbols]
        used.sort()
        '''))
    code += _T(_MODEL_BODY, model_line=model_line, functions=functions, pcsv=repr(pcsv),
               tdbname=Path(ctx.tb._find_tdb(a["tdb"])).name,
               fnote=repr("Functions expanded." if expand else "Functions are kept symbolic (see FUNCTIONS below)"))
    if a.get("evaluate_at"):
        # state (T, P, site fractions) exactly as the tool builds it from evaluate_at
        from pycalphad import Model
        with __import__("warnings").catch_warnings():
            __import__("warnings").simplefilter("ignore")
            sfs = Model(db, comps, phase).site_fractions
        ev = {str(k).strip().upper(): float(val) for k, val in a["evaluate_at"].items()}
        T = ev.pop("T")
        P = ev.pop("P", 101325.0)
        sub, xs = {}, {}
        ymap = {str(y).upper(): y for y in sfs}
        ymap.update({f"Y{y.sublattice_index}_{y.species.name}".upper(): y for y in sfs})
        for k, val in ev.items():
            if k in ymap:
                sub[str(ymap[k])] = val
            elif re.fullmatch(r"X\([A-Z]{1,2}\)", k):
                xs[_el(k)] = val
        by_sub = {}
        for y in sfs:
            by_sub.setdefault(y.sublattice_index, []).append(y)
        for s, ys in by_sub.items():
            tot = sum(xs.get(y.species.name, 0.0) for y in ys if y.species.name != "VA")
            for y in ys:
                if str(y) in sub:
                    continue
                sub[str(y)] = (0.0 if tot > 0 else 1.0) if y.species.name == "VA" else \
                    (xs.get(y.species.name, 0.0) / tot if tot > 0 else 0.0)
        code += _T('''

            # evaluation at one state (all functions expanded); site fractions from X(EL): random mixing per sublattice
            full = Model(db, comps, phase)
            site_fractions = <<sub>>
            by_name = {str(y): y for y in full.site_fractions}
            state = {v.P: <<P>>, v.T: <<T>>, **{by_name[k]: val for k, val in site_fractions.items()}}
            lines.append("=== evaluation (all functions expanded)")
            lines.append(f"state = {({str(k): val for k, val in state.items()})}")
            for k, e in list(full.models.items()) + [("GM", full.GM)] + [(n, getattr(full, n)) for n in ("TC", "BMAG")
                                                                         if n in physical]:
                try:
                    lines.append(f"{k} = {float(e.xreplace(state))}")
                except Exception as ex:
                    lines.append(f"{k} = could not evaluate ({type(ex).__name__})")
            ''', T=_num(T), P=_num(P), sub=Ctx.py(sub, 17))
    code += f"(OUT / {txt!r}).write_text('\\n'.join(lines))\nprint('\\n'.join(lines))\n"
    if a["latex"]:
        tex = ctx.out_name(prefix + ".tex", "model.tex")
        code += _T('''

            # LaTeX: site fractions -> y^{(s)}_{EL}, function names upright
            tex = [f"% Gibbs energy contributions of {phase} ({', '.join(comps)}); y<s>_<EL> = site fraction"]
            for k, e in contributions.items():
                rep = {y: symengine.Symbol(f"y^{{({y.sublattice_index})}}_{{\\\\mathrm{{{y.species.name}}}}}") for y in sfs}
                for s in e.free_symbols:
                    if s not in rep and type(s) is symengine.Symbol and len(str(s)) > 1:
                        rep[s] = symengine.Symbol("\\\\mathrm{" + str(s).replace("_", "\\\\_") + "}")
                tex.append(f"% {k}: {CONTRIBUTIONS.get(k, k)}\\n\\\\begin{{dmath}}\\nG^{{\\\\mathrm{{{k}}}}} = "
                           f"{symengine.latex(e.xreplace(rep))}\\n\\\\end{{dmath}}")
            (OUT / <<tex>>).write_text("\\n\\n".join(tex) + "\\n")
            ''', tex=repr(tex))
    return code


@generator("write_database")
def _gen_write_database(args, ctx: Ctx) -> str:
    if args.get("elements") or args.get("phases") or args.get("exclude_phases"):
        ctx.fallback("extracting a subsystem (elements/phases) uses the toolkit's database subsetting")
    if args.get("dialect"):
        from .codegen_formats import write_database_code
        return write_database_code(args, ctx)
    t = ctx.tb._find_tdb(args["tdb"])
    out = _outp(ctx.tb, args["output_tdb"], f"{Path(t).stem}_subset.tdb")
    name = ctx.out_name(out, "database.tdb")
    ctx.need("from pycalphad import Database")
    return _load_db_code(ctx, args["tdb"]) + _T('''
        # rewrite with pycalphad's TDB writer (comments are not kept)
        db.to_file(str(OUT / <<name>>), fmt="tdb", if_exists="overwrite", groupby=<<groupby>>,
                   if_incompatible=<<incompat>>)
        back = Database(str(OUT / <<name>>))   # reload check
        print(len(back.phases), "phases,", len(back._parameters), "parameters,", len(back.symbols), "functions")
        ''', name=repr(name), groupby=repr(args.get("groupby") or "subsystem"),
                                                 incompat=repr("fix" if args.get("fix_incompatible", True) else "warn"))


# ---------------------------------------------------------------------------------------------------------
# ESPEI MCMC post-processing
# ---------------------------------------------------------------------------------------------------------
def _mcmc_inputs(ctx: Ctx, a: Dict):
    """Generation-time: run files, symbols and burn-in exactly as the tool resolved them."""
    from . import espei_post as ep
    t, lpp, d = ep.resolve_run_files(ctx.tb.p(a["trace"]), ctx.tb.p(a["lnprob"]) if a.get("lnprob") else None,
                                     ctx.tb._find_tdb(a["tdb"]) if a.get("tdb") else None)
    ch = ep.load_chains(t, lpp)
    db = ep._load_db(d)
    syms = ep.fitted_symbols(db, ch["trace"].shape[2], a.get("symbols"))
    burn, binfo = ep._burn(a.get("burn_in", "auto"), ch["lnprob"])
    if ch["truncation"] != "espei.analysis.truncate_arrays":
        ctx.fallback(f"the trace was truncated with the toolkit's fallback rule ({ch['truncation']})")
    return t, lpp, d, ch, db, syms, burn, binfo


def _mcmc_load_code(ctx: Ctx, t, lpp, d, syms, given_symbols, burn, binfo) -> str:
    ctx.need("import numpy as np", "from pycalphad import Database", "from espei.analysis import truncate_arrays",
             "from espei.utils import database_symbols_to_fit, unpack_piecewise")
    if binfo.get("method") == "user":
        bnote = "given"
    else:
        bnote = (f"detected when the tool ran: lnprob median within {binfo.get('k', 3):g} noise units of its "
                 f"stable level {binfo.get('reference_lnprob', float('nan')):.6g} +- {binfo.get('tolerance', 0):.3g}")
    return _T('''
        trace = np.load(<<t>>)       # (chains, iterations, parameters)
        lnprob = np.load(<<lp>>)     # (chains, iterations)
        # drop the zero-filled iterations of an unfinished run (keep any iteration with a non-zero lnprob)
        n_alloc = trace.shape[1]
        tr, lp = truncate_arrays(trace, lnprob)
        nz = np.nonzero(np.any(lnprob != 0, axis=0))[0]
        if nz.size and nz[-1] + 1 > tr.shape[1]:
            tr, lp = trace[:, :nz[-1] + 1, :], lnprob[:, :nz[-1] + 1]
        tr, lp = tr.astype(float), np.where(np.isfinite(lp), lp, np.nan).astype(float)
        db = Database(str(<<tdb>>))
        symbols = <<syms>>   # trace columns, in ESPEI's order
        burn_in = <<burn>>   # iterations discarded (<<bnote>>)
        ''', t=ctx.path(t), lp=ctx.path(lpp), tdb=ctx.path(d),
              syms=(f"sorted({Ctx.py(list(given_symbols))})" if given_symbols else
                    f"database_symbols_to_fit(db)   # {Ctx.py(syms)}"), burn=burn, bnote=bnote)


@generator("mcmc_to_tdb")
def _gen_mcmc_to_tdb(args, ctx: Ctx) -> str:
    a = {"statistic": "median", "burn_in": "auto", **args}
    t, lpp, d, ch, db, syms, burn, binfo = _mcmc_inputs(ctx, a)
    st = a["statistic"].lower()
    stem = Path(t).parent.name or Path(t).stem
    output = a.get("output")
    out = _outp(ctx.tb, output if (not output or output.lower().endswith(".tdb")) else output + ".tdb",
                f"{stem}_{st}.tdb")
    name = ctx.out_name(out, f"{stem}_{st}.tdb")
    code = _mcmc_load_code(ctx, t, lpp, d, syms, a.get("symbols"), burn, binfo)
    if st == "map":
        pick = ("# maximum a posteriori: the highest lnprob of the whole run (as ESPEI's own mcmc.tdb)\n"
                "chain, it = divmod(int(np.nanargmax(lp.reshape(-1))), lp.shape[1])\nvalues = tr[chain, it, :]\n")
    else:
        pick = (f"samples = tr[:, burn_in:, :].reshape(-1, tr.shape[2])   # posterior after burn-in\n"
                f"values = np.{'median' if st == 'median' else 'mean'}(samples, axis=0)\n")
    code += "\n" + pick + _T('''
        old = {s: float(unpack_piecewise(db.symbols[s])) for s in symbols}
        for s, val in zip(symbols, values):
            db.symbols[s] = float(val)      # as ESPEI does: dbf.symbols.update(parameters)
            print(f"{s}: {old[s]:.6g} -> {float(val):.6g}")
        db.to_file(str(OUT / <<name>>), if_exists="overwrite")
        ''', name=repr(name))
    return code


_MCMC_STATS = '''

def acf(x):
    """Normalised autocorrelation function (FFT)."""
    x = np.asarray(x, float)
    nfft = 1 << (2 * len(x) - 1).bit_length()
    f = np.fft.rfft(x - x.mean(), n=nfft)
    a = np.fft.irfft(f * np.conjugate(f), n=nfft)[:len(x)]
    return a / a[0]


def integrated_time(x, c=5.0):
    """Integrated autocorrelation time, x (steps, walkers): walker-averaged ACF, Sokal's automatic window."""
    good = [j for j in range(x.shape[1]) if np.ptp(x[:, j]) > 0]
    if not good or x.shape[0] < 2:
        return float("nan")
    taus = 2.0 * np.cumsum(np.mean([acf(x[:, j]) for j in good], axis=0)) - 1.0
    m = np.arange(len(taus)) < c * taus
    return float(max(taus[int(np.argmin(m)) if np.any(~m) else len(taus) - 1], 1.0))


def split_rhat(x):
    """Gelman-Rubin R-hat on split chains, x (walkers, steps) (indicative only for an ensemble sampler)."""
    n = x.shape[1] // 2
    if n < 2:
        return float("nan")
    ch = np.concatenate([x[:, :n], x[:, x.shape[1] - n:]], axis=0)
    W, B = float(np.mean(ch.var(axis=1, ddof=1))), float(n * np.var(ch.mean(axis=1), ddof=1))
    if W <= 0:
        return float("nan") if B <= 0 else float("inf")
    return float(np.sqrt(((n - 1) / n * W + B / n) / W))


post = tr[:, burn_in:, :]
samples = tr[:, burn_in::thin, :].reshape(-1, tr.shape[2])
n_steps = post.shape[1]
chain_map, it_map = divmod(int(np.nanargmax(lp.reshape(-1))), lp.shape[1]) if np.isfinite(lp).any() else (0, 0)
moved = np.any(np.diff(post, axis=1) != 0, axis=2) if n_steps > 1 else np.full((tr.shape[0], 1), np.nan)
acceptance = moved.mean(axis=1)
print(f"acceptance fraction mean {np.nanmean(acceptance):.3f}; MAP lnprob {lp[chain_map, it_map]:.5g}")
rows = []
for k, s in enumerate(symbols):
    x = samples[:, k]
    q = np.percentile(x, [2.5, 16, 50, 84, 97.5])
    chains_k = post[:, :, k]
    tau = integrated_time(chains_k.T) if n_steps >= 4 else float("nan")
    n_good = int(sum(np.ptp(chains_k[j]) > 0 for j in range(chains_k.shape[0])))
    ess = float(min(n_steps * n_good / tau, n_steps * chains_k.shape[0])) if np.isfinite(tau) else float("nan")
    rows.append({"symbol": s, "tdb_parameters": TDB_PARAMETERS[s], "value_in_tdb": float(unpack_piecewise(db.symbols[s])),
                 "mean": float(np.mean(x)), "median": float(q[2]), "std": float(np.std(x, ddof=1)) if len(x) > 1 else 0.0,
                 "ci95_low": float(q[0]), "ci95_high": float(q[4]), "map": float(tr[chain_map, it_map, k]), "tau": tau,
                 "tau_reliable": bool(np.isfinite(tau) and n_steps > 50.0 * tau), "ess": ess,
                 "split_rhat": split_rhat(chains_k)})
summary = pd.DataFrame(rows)
summary.to_csv(OUT / <<csv>>, index=False)
print(summary.to_string())

# Pearson correlation of the posterior samples (constant parameters: NaN)
sd = samples.std(axis=0)
corr = np.full((len(symbols), len(symbols)), np.nan)
okc = sd > 0
if okc.sum() >= 1 and samples.shape[0] > 1:
    z = (samples[:, okc] - samples[:, okc].mean(axis=0)) / sd[okc]
    idx = np.nonzero(okc)[0]
    corr[np.ix_(idx, idx)] = np.clip((z.T @ z) / samples.shape[0], -1, 1)
'''

_MCMC_PLOTS = '''

figures = {}
# 1) lnprob vs iteration
n_ch, n = lp.shape
fig, ax = plt.subplots(figsize=(7.0, 4.2))
for c in range(n_ch):
    ax.plot(np.arange(n), lp[c], color="0.55", lw=0.6, alpha=float(np.clip(8.0 / n_ch, 0.08, 0.5)))
ax.plot(np.arange(n), np.nanmedian(lp, axis=0), color="#d62728", lw=1.8, label="median over chains")
ax.plot([], [], color="0.55", lw=0.8, label=f"{n_ch} chains")
if burn_in > 0:
    ax.axvspan(0, burn_in, color="0.85", alpha=0.6, lw=0, label=f"burn-in ({burn_in})")
fin = lp[:, burn_in:][np.isfinite(lp[:, burn_in:])]
if fin.size:
    lo, hi = np.percentile(fin, [1, 100])
    span = max(hi - lo, 1e-9)
    if lp[np.isfinite(lp)].min() < lo - 3 * span:   # very low early lnprob would squash the plot: clip
        ax.set_ylim(lo - 1.5 * span, hi + 0.15 * span)
        ax.text(0.99, 0.02, "y axis clipped (early iterations lower)", transform=ax.transAxes, ha="right",
                va="bottom", fontsize=7, color="0.4")
ax.set_xlim(0, max(n - 1, 1))
ax.set_xlabel("Iteration")
ax.set_ylabel("ln(probability)")
ax.set_title(<<lptitle>>, fontsize=11)
ax.grid(alpha=0.25)
ax.legend(fontsize=8, loc="lower right" if fin.size else "best")
fig.tight_layout()
fig.savefig(OUT / <<lnprob_png>>, dpi=150, bbox_inches="tight")
figures[<<lnprob_png>>] = (fig, ax)

# 2) parameter traces (grey = burn-in, dashed = posterior median)
k = tr.shape[2]
cols = min(4, max(1, int(np.ceil(np.sqrt(k)))))
rows_ = int(np.ceil(k / cols))
fig, axs = plt.subplots(rows_, cols, figsize=(3.3 * cols, 2.2 * rows_), squeeze=False, sharex=True)
for idx in range(rows_ * cols):
    ax = axs[idx // cols][idx % cols]
    if idx >= k:
        ax.set_visible(False)
        continue
    for c in range(n_ch):
        ax.plot(np.arange(n), tr[c, :, idx], color="#1f77b4", lw=0.5, alpha=float(np.clip(6.0 / n_ch, 0.06, 0.5)))
    if burn_in > 0:
        ax.axvspan(0, burn_in, color="0.85", alpha=0.6, lw=0)
    ax.axhline(summary["median"][idx], color="#d62728", lw=1.0, ls="--")
    ax.set_title(f"{symbols[idx]}\\n{TRACE_LABELS[symbols[idx]]}", fontsize=7.5)
    ax.tick_params(labelsize=7)
    ax.yaxis.set_major_locator(MaxNLocator(4))
    if idx // cols == rows_ - 1 or idx + cols >= k:
        ax.set_xlabel("Iteration", fontsize=8)
        ax.tick_params(labelbottom=True)
fig.suptitle(f"Parameter traces ({n_ch} chains; grey = burn-in, dashed = posterior median)", fontsize=10)
fig.tight_layout()
fig.savefig(OUT / <<traces_png>>, dpi=150, bbox_inches="tight")
figures[<<traces_png>>] = (fig, axs[0][0])

# 3) corner plot of <<corner_what>>: marginals (median, 95 % interval) and pairwise histograms
corner = <<corner>>
cs = samples[:, [symbols.index(s) for s in corner]]
kc = cs.shape[1]
size = max(5.0, min(1.75 * kc, 22.0))
fig, axs = plt.subplots(kc, kc, figsize=(size, size), squeeze=False)
q = np.percentile(cs, [2.5, 50, 97.5], axis=0)
lims = []
for j in range(kc):
    lo, hi = cs[:, j].min(), cs[:, j].max()
    pad = 0.05 * (hi - lo) if hi > lo else max(abs(lo) * 1e-6, 1e-12)
    lims.append((lo - pad, hi + pad))
for i in range(kc):
    for j in range(kc):
        ax = axs[i][j]
        if j > i:
            ax.set_visible(False)
            continue
        if i == j:
            ax.hist(cs[:, j], bins=25, range=lims[j], histtype="stepfilled", color="#9ecae1", edgecolor="#3182bd", lw=0.8)
            for qq, lw, ls in ((q[1, j], 1.0, "-"), (q[0, j], 0.8, "--"), (q[2, j], 0.8, "--")):
                ax.axvline(qq, color="#d62728", lw=lw, ls=ls)
            ax.set_title(f"{corner[j]}\\n{q[1, j]:.4g}\\n[{q[0, j]:.3g}, {q[2, j]:.3g}]", fontsize=6.5)
        else:
            if cs.shape[0] < 400:
                ax.scatter(cs[:, j], cs[:, i], s=3, color="#08519c", alpha=0.35, lw=0)
            else:
                ax.hist2d(cs[:, j], cs[:, i], bins=25, range=[lims[j], lims[i]], cmap="Blues", cmin=1)
            ax.axvline(q[1, j], color="#d62728", lw=0.6)
            ax.axhline(q[1, i], color="#d62728", lw=0.6)
            ax.set_ylim(*lims[i])
        ax.set_xlim(*lims[j])
        ax.xaxis.set_major_locator(MaxNLocator(3, prune="both"))
        ax.yaxis.set_major_locator(NullLocator() if i == j else MaxNLocator(3, prune="both"))
        ax.tick_params(labelsize=6)
        if i == kc - 1:
            ax.set_xlabel(corner[j], fontsize=8)
            for lab in ax.get_xticklabels():
                lab.set_rotation(45)
        else:
            ax.tick_params(labelbottom=False)
        if j == 0 and i > 0:
            ax.set_ylabel(corner[i], fontsize=8)
        elif i != j:
            ax.tick_params(labelleft=False)
fig.text(0.97, 0.97, f"Posterior ({cs.shape[0]} samples)\\ndiagonal: marginal, median (solid)\\n"
         "and 95 % credible interval (dashed)", ha="right", va="top", fontsize=10 if kc > 3 else 8)
fig.subplots_adjust(wspace=0.08, hspace=0.08, top=0.93 if kc < 4 else 0.96)
fig.savefig(OUT / <<corner_png>>, dpi=150, bbox_inches="tight")
figures[<<corner_png>>] = (fig, axs[0][0])

# 4) correlation matrix
k = len(symbols)
size = max(4.5, min(0.55 * k + 2.5, 16))
fig, ax = plt.subplots(figsize=(size + 0.8, size))
im = ax.imshow(np.ma.masked_invalid(corr), cmap="RdBu_r", vmin=-1, vmax=1)
ax.set_xticks(range(k))
ax.set_yticks(range(k))
ax.set_xticklabels(symbols, rotation=90, fontsize=7)
ax.set_yticklabels(symbols, fontsize=7)
if k <= 16:
    for i in range(k):
        for j in range(k):
            if np.isfinite(corr[i, j]):
                ax.text(j, i, f"{corr[i, j]:.2f}", ha="center", va="center", fontsize=6,
                        color="white" if abs(corr[i, j]) > 0.6 else "black")
fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04).set_label("Pearson correlation")
ax.set_title("Posterior parameter correlation", fontsize=11)
fig.tight_layout()
fig.savefig(OUT / <<corr_png>>, dpi=150, bbox_inches="tight")
figures[<<corr_png>>] = (fig, ax)

fig, ax = figures[<<target>>]
'''


@generator("analyze_mcmc")
def _gen_analyze_mcmc(args, ctx: Ctx) -> str:
    from . import espei_post as ep
    a = {"burn_in": "auto", "thin": 1, **args}
    t, lpp, d, ch, db, syms, burn, binfo = _mcmc_inputs(ctx, a)
    thin = max(1, int(a.get("thin") or 1))
    stem = Path(t).parent.name or Path(t).stem
    out = _outp(ctx.tb, a.get("output_prefix"), f"mcmc_analysis/{stem}")
    if Path(out).suffix.lower() in (".png", ".json", ".csv"):
        out = str(Path(out).with_suffix(""))
    names = {s: ctx.out_name(out + s, Path(out + s).name) for s in
             ("_lnprob.png", "_traces.png", "_corner.png", "_correlation.png", "_summary.csv")}
    target = _target_path(ctx)
    tname = next((names[s] for s in names if target and str(Path(out + s).resolve()) == target), names["_lnprob.png"])
    usage = ep.symbol_usage(db, syms)
    tdb_params = {s: "; ".join(f"{p} [{r}]" for p, r in zip(usage[s]["parameters"], usage[s]["roles"]))
                  for s in syms}
    labels = {s: ep._short(usage, s) for s in syms}
    if a.get("corner_parameters"):
        corner = [s.upper() for s in a["corner_parameters"]]
        corner_what = "the chosen parameters"
    else:
        corner = syms[:12]
        corner_what = "the first 12 parameters" if len(syms) > 12 else "all parameters"
    ctx.need("import pandas as pd", "import matplotlib.pyplot as plt", "from matplotlib.ticker import MaxNLocator, "
             "NullLocator")
    code = _mcmc_load_code(ctx, t, lpp, d, syms, a.get("symbols"), burn, binfo)
    code += (f"thin = {thin}\n"
             f"# TDB parameter(s) each symbol belongs to [coefficient of], found by the toolkit in {Path(d).name}\n"
             f"TDB_PARAMETERS = {Ctx.py(tdb_params, 17)}\n"
             f"TRACE_LABELS = {Ctx.py(labels, 15)}   # short form for the trace-plot titles\n")
    lptitle = f"MCMC log-probability ({Path(t).parent.name or Path(t).stem})"
    code += _T(_MCMC_STATS, csv=repr(names["_summary.csv"]))
    code += _T(_MCMC_PLOTS, lptitle=repr(lptitle), lnprob_png=repr(names["_lnprob.png"]),
               traces_png=repr(names["_traces.png"]), corner_png=repr(names["_corner.png"]),
               corr_png=repr(names["_correlation.png"]), corner=Ctx.py(corner), corner_what=corner_what,
               target=repr(tname))
    return code
