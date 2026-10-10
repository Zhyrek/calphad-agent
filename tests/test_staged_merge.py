"""staged_merge: intermediate databases B1/B2/B3/Q/A1/A2, their substitutions, the version diffs, invariant
comparison and the deck - on CrFe-82Her (A) + a Cr-Fe-Ni subset of CoCrFeNiV_Choi2019 (B), and on small synthetic
databases for the decision round trip and the C14 model expansion."""
import math
import json
from pathlib import Path

import pytest

from calphad_tools import tdbstage as ts
from calphad_tools.staged_merge import compare_invariants, staged_merge, summary_text
from calphad_tools.tdbparse import parse_tdb

from test_compatibility import A_FeMo, B_FeMoW, write

EX = Path(__file__).resolve().parent.parent / "examples" / "tdbs"


@pytest.fixture(scope="module")
def her_choi(tmp_path_factory):
    d = tmp_path_factory.mktemp("staged")
    sub = ts.restrict_tdb(parse_tdb(str(EX / "CoCrFeNiV_Choi2019.tdb")), ["CR", "FE", "NI"],
                          phases=["LIQUID", "BCC_A2", "FCC_A1", "SIGMA"])
    b = ts.write_tdb(sub, str(d / "CrFeNi_Choi.tdb"), ["Cr-Fe-Ni subset of CoCrFeNiV_Choi2019"])
    r = staged_merge(str(EX / "CrFe-82Her.tdb"), b, str(d / "out"), policy="ask", plot=True,
                     ternary_temperatures_c=[1000], tmin_c=400, tmax_c=2000, composition_step=0.04,
                     temperature_step=20, ternary_step=0.05, n_workers=3)
    return r


def _row_segs(path, key, elements=None):
    v = ts.StageView(parse_tdb(path), "x", elements=elements)
    return v.rows[key].segs


def _equal(a, b):
    return ts.compare_segs(a, b, 1e-9, 1e-9)[0]


def test_versions_written_and_load(her_choi):
    r = her_choi
    assert r["written"]
    assert set(r["versions"]) == {"B", "B1", "B2", "B3", "Q", "A1", "A2", "A"}
    for lab, v in r["versions"].items():
        if lab in ("A", "B"):
            continue
        assert Path(v.path).exists() and v.validation["loaded"] and not v.validation["errors"], (lab, v.validation)
    names = [Path(r["versions"][k].path).name for k in ("B1", "B2", "B3", "Q")]
    assert names[0].startswith("B1_") and names[3].startswith("Q_")


def test_substitutions_really_happened(her_choi):
    r = her_choi
    A = str(EX / "CrFe-82Her.tdb")
    v = r["versions"]
    k_fcc = "G(FCC_A1,CR;0)"
    # B2: A's lattice stability of fcc Cr (82Her) replaced Choi's (SGTE 91Din)
    assert not _equal(_row_segs(v["B"].path, k_fcc), _row_segs(A, k_fcc))
    assert _equal(_row_segs(v["B2"].path, k_fcc), _row_segs(A, k_fcc))
    # B3: A's Cr-Fe binary description
    k_liq = "G(LIQUID,CR,FE;0)"
    assert not _equal(_row_segs(v["B2"].path, k_liq), _row_segs(A, k_liq))
    assert _equal(_row_segs(v["B3"].path, k_liq), _row_segs(A, k_liq))
    # B's own Cr-Ni / Fe-Ni binaries are untouched
    k_crni = "G(LIQUID,CR,NI;0)"
    assert _equal(_row_segs(v["B3"].path, k_crni), _row_segs(v["B"].path, k_crni))
    # B1: the SER functions were already equal (only a pressure term differs, < tolerance): nothing to change
    st = {s.title: s for s in r["steps"]}
    assert not st["B → B1"].diff.changed and not st["B → B1"].changed_items


def test_unchanged_subsystems_are_reused_not_recomputed(her_choi):
    plots = her_choi["plots"]
    assert all(plots[("B1", t)]["reused_from"] == "B" for t in ("Cr-Fe", "Cr-Ni", "Fe-Ni", "Cr-Fe-Ni_1000C"))
    assert plots[("B2", "Fe-Ni")]["reused_from"] == "B"           # G(FCC,CR) does not touch Fe-Ni
    assert "reused_from" not in plots[("B2", "Cr-Ni")]
    st = {s.title: s for s in her_choi["steps"]}
    assert [it.name for it in st["B1 → B2"].changed_items] == ["Cr-Fe", "Cr-Ni", "Cr-Fe-Ni @ 1000 °C"]


def test_invariant_comparison_consistent(her_choi):
    r = her_choi
    st = {s.title: s for s in r["steps"]}["B1 → B2"]
    rows = st.inv["Cr-Ni"]
    old = json.loads(Path(r["plots"][("B1", "Cr-Ni")]["png"]).with_suffix(".json").read_text())["invariants"]
    new = json.loads(Path(r["plots"][("B2", "Cr-Ni")]["png"]).with_suffix(".json").read_text())["invariants"]
    eut = [x for x in rows if x["invariant"] == "BCC_A2 + FCC_A1 + LIQUID"][0]
    t_old = [i["T_K"] for i in old if "LIQUID" in i["phases"]][0]
    t_new = [i["T_K"] for i in new if "LIQUID" in i["phases"]][0]
    assert eut["T_old_K"] == t_old and eut["T_new_K"] == t_new
    assert math.isclose(eut["dT_K"], t_new - t_old) and eut["status"] == "moved"
    for s in r["steps"]:
        for rs in s.inv.values():
            for x in rs:
                if x["status"] in ("moved", "same"):
                    assert math.isclose(x["dT_K"], x["T_new_K"] - x["T_old_K"], abs_tol=1e-9)
                    assert (x["status"] == "moved") == (abs(x["dT_K"]) > 5 or x["dx"] > 0.01)


def test_deck_one_summary_plus_two_slides_per_step(her_choi):
    from pptx import Presentation
    r = her_choi
    prs = Presentation(r["files"]["deck"])
    assert len(r["steps"]) == 6
    titles = []
    for sl in prs.slides:
        t = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(t[0] if t else "")
    # summary, the normalisation's TDB changes (if B / A were normalised), then per step its TDB changes (one or more
    # slides, "(2/3)") and its diagrams
    n_norm = sum(t.startswith("Normalisation ") for t in titles)
    assert n_norm == sum(1 for d in r["norm_diffs"] for t in titles if t.startswith(f"Normalisation {d.title}"))
    tdb = [t for t in titles if ": TDB changes" in t]
    assert len(prs.slides) == 1 + len(tdb) + len(r["steps"])
    assert titles[0].startswith("Staged merge")
    i = 1 + n_norm
    assert titles[i].startswith("B → B1: TDB changes")
    i = max(j for j, t in enumerate(titles) if t.startswith("B → B1: TDB changes")) + 1
    assert titles[i].startswith("B → B1: diagrams")
    assert titles[-1].startswith("A1 → A2: diagrams")
    for st in r["steps"]:
        assert st.tdb is not None and Path(st.tdb.path).name == f"{st.old}-{st.new}.diff"
        assert sum(t.startswith(f"{st.title}: TDB changes") for t in titles) <= 8
    for sl in prs.slides:            # wrapped, never cut mid-word; tables stay on the slide
        for sh in sl.shapes:
            if sh.has_table:
                assert (sh.top + sh.height) / 914400 < 7.4
                for row in sh.table.rows:
                    for c in row.cells:
                        assert not c.text.rstrip().endswith("…"), c.text
    assert Path(r["files"]["xlsx"]).exists()
    text = summary_text(r)
    assert text.count("\n") < 20 and "B1 → B2" in text and "invariants moved" in text


def test_compare_invariants_matching():
    old = [{"T_K": 1000.0, "phases": ["BCC_A2", "BCC_A2#2", "SIGMA"], "compositions": {
        "BCC_A2": {"CR": 0.1, "FE": 0.9}, "BCC_A2#2": {"CR": 0.9, "FE": 0.1}, "SIGMA": {"CR": 0.5, "FE": 0.5}}}]
    new = [{"T_K": 1012.5, "phases": ["BCC_A2#2", "BCC_A2", "SIGMA"], "compositions": {
        "BCC_A2#2": {"CR": 0.1, "FE": 0.9}, "BCC_A2": {"CR": 0.9, "FE": 0.1}, "SIGMA": {"CR": 0.52, "FE": 0.48}}},
           {"T_K": 1500.0, "phases": ["LIQUID", "BCC_A2", "FCC_A1"], "compositions": {}}]
    rows = compare_invariants(old, new, ["CR", "FE"], ternary=False)
    moved = [r for r in rows if r["status"] == "moved"][0]
    assert moved["dT_K"] == 12.5 and math.isclose(moved["dx"], 0.02)       # '#2' order does not matter
    assert [r["status"] for r in rows if r["invariant"].startswith("BCC_A2 + FCC")] == ["appeared"]


# ------------------------------------------------------------------------------------------ synthetic
def test_decision_round_trip_and_c14_expansion(tmp_path):
    sig_a = [" PHASE SIGMA % 3 10 4 16 !", " CONSTITUENT SIGMA :FE:MO:FE,MO: !"]
    sig_b = [" PHASE SIGMA % 3 8 4 18 !", " CONSTITUENT SIGMA :FE:MO,W:FE,MO,W: !"]
    a_text = A_FeMo().replace(" PHASE C14", "\n".join(sig_a) + "\n PHASE C14") + \
        " PARAMETER G(SIGMA,FE:MO:FE;0) 298.15 26*GHSERFE+4*GHSERMO+100000; 6000 N !\n"
    b_text = B_FeMoW().replace(" PHASE C14", "\n".join(sig_b) + "\n PHASE C14") + \
        " PARAMETER G(SIGMA,FE:W:FE;0) 298.15 26*GHSERFE+4*GHSERWW+90000; 6000 N !\n"
    a, b = write(tmp_path, "A_FeMo.tdb", a_text), write(tmp_path, "B_FeMoW.tdb", b_text)
    r = staged_merge(a, b, str(tmp_path / "out"), policy="ask", plot=False)
    assert not r["written"] and "phase:SIGMA" in summary_text(r)
    assert Path(r["versions"]["B1"].path).exists() and Path(r["versions"]["B2"].path).exists()
    r = staged_merge(a, b, str(tmp_path / "out"), policy="ask", plot=False, decisions={"phase:SIGMA": "db1"})
    assert r["written"]
    st = {s.title: s for s in r["steps"]}
    # A1 = A + B's unary terms: C14 gains FE on the second sublattice -> FE:FE inside A's Fe-Mo
    fe_mo = st["A → A1"].diff.get(("FE", "MO"))
    assert fe_mo.changed
    m = [x for x in fe_mo.models if x.phase == "C14"][0]
    assert m.old == "(FE)2(MO)1" and m.new == "(FE)2(FE,MO)1" and m.new_endmembers == ["FE:FE"]
    assert "G(C14,FE:FE;0)" in {p.key for p in fe_mo.params}
    # B's Fe-W is not touched by any step after B1/B2 (shared elements are Fe, Mo; A has no W)
    assert not st["B3 → Q"].diff.get(("FE", "W")).changed
    q = parse_tdb(r["versions"]["Q"].path)
    assert q.phases["C14"].constituents == [["FE", "W"], ["MO", "FE", "W"]]


def test_tool_and_standalone_code_build_the_same_tdbs(tmp_path):
    import subprocess
    import sys
    from calphad_tools.agent.tools import UI, ToolBox
    write(tmp_path, "A_FeMo.tdb", A_FeMo())
    write(tmp_path, "B_FeMoW.tdb", B_FeMoW())
    tb = ToolBox(str(tmp_path), "out", UI(), python_tool="off")
    text, imgs, err = tb.call("staged_merge", {"base_tdb": "A_FeMo.tdb", "add_tdb": "B_FeMoW.tdb", "plot": False,
                                               "ternary_temperatures_c": [], "composition_step": 0.1,
                                               "temperature_step": 50, "tmin_c": 800, "tmax_c": 2600})
    assert not err, text
    assert "A → A1 (+ B's unaries)" in text and "(+FE:FE)" in text
    text, _, err = tb.call("show_code", {"target": "1", "format": "script"})
    assert not err, text
    script = next(tmp_path.glob("out/code/*_standalone.py"))
    assert "from calphad_tools" not in script.read_text()
    r = subprocess.run([sys.executable, str(script)], cwd=str(tmp_path), capture_output=True, text=True, timeout=900)
    assert r.returncode == 0, r.stderr[-2000:]
    assert "A -> A1" in r.stdout and "(+FE:FE)" in r.stdout
    tool_q = next(tmp_path.glob("out/staged_merge_*/Q_*.tdb"))
    code_q = next(tmp_path.glob("standalone_outputs/Q_*.tdb"))
    v1, v2 = ts.StageView(parse_tdb(str(tool_q)), "t"), ts.StageView(parse_tdb(str(code_q)), "c")
    assert not ts.diff_versions(v1, v2, ts.stage_subsystems(["FE", "MO", "W"])).changed
    assert list(tmp_path.glob("standalone_outputs/A1_Fe-Mo.png"))


def test_magnetic_basis_rule_in_staged_merge(tmp_path):
    from test_compatibility import l12_A, l12_B
    a, b = write(tmp_path, "A_FeMo.tdb", l12_A()), write(tmp_path, "B_FeW.tdb", l12_B())
    r = staged_merge(a, b, str(tmp_path / "out"), policy="ask", plot=False)
    assert r["written"] and r["A_eff"] != r["A"] and Path(r["A_eff"]).name.startswith("A0_")
    text = summary_text(r)
    assert "A rescaled to B's basis - A (non-magnetic) rescaled to B's magnetic basis" in text and "A0 " in text
    q = parse_tdb(r["versions"]["Q"].path)
    assert q.phases["L12"].site_ratios == [0.75, 0.25]                      # B's ratios in A's sublattice order
    assert q.phases["L12"].constituents[0][:2] == ["FE", "MO"]               # A's order kept
    # B's magnetic parameters reach Q unchanged (values; B was only reordered)
    vb = ts.StageView(parse_tdb(r["B"]), "b")
    vq = ts.StageView(q, "q")
    for k, row in vb.rows.items():
        if row.ptype in ("TC", "BMAGN"):
            assert _equal(row.segs, vq.rows[k].segs), k
    # A's own G of the phase in Q is A's value scaled to B's formula unit (x 1/4)
    va = ts.StageView(parse_tdb(a), "a")
    from calphad_tools.pwinline import scale
    assert _equal(scale(va.rows["G(L12,FE:MO;0)"].segs, 0.25), vq.rows["G(L12,FE:MO;0)"].segs)
    # A -> A1 (+ B's unaries) and A -> Q: the exact rescale alone is not a change in A's Fe-Mo
    st = {s.title: s for s in r["steps"]}
    fe_mo = st["A → A1"].diff.get(("FE", "MO"))
    assert not fe_mo.changed and any("exact" in n for n in fe_mo.notes)
    va2 = ts.StageView(parse_tdb(a), "A", elements=["FE", "MO"])
    vq2 = ts.StageView(q, "Q", elements=["FE", "MO"])
    assert not ts.diff_versions(va2, vq2, [("FE", "MO")]).changed


def _fe_w_magnetic_c14():
    from test_compatibility import tdb_text
    return tdb_text(["FE", "W"], [
        " PHASE LIQUID % 1 1 !", " CONSTITUENT LIQUID :FE,W: !", " PHASE BCC_A2 %B 2 1 3 !",
        " CONSTITUENT BCC_A2 :FE,W:VA: !", " PHASE C14 %M 2 0.66667 0.33333 !", " CONSTITUENT C14 :FE,W:W: !"], [
        " PARAMETER G(LIQUID,FE;0) 298.15 12040.17-6.55843*T+GHSERFE; 6000 N !",
        " PARAMETER G(LIQUID,W;0) 298.15 52160.584-14.10999*T+GHSERWW; 6000 N !",
        " PARAMETER G(BCC_A2,FE:VA;0) 298.15 GHSERFE; 6000 N !", " PARAMETER G(BCC_A2,W:VA;0) 298.15 GHSERWW; 6000 N !",
        " PARAMETER TC(BCC_A2,FE:VA;0) 298.15 1043; 6000 N !", " PARAMETER BMAGN(BCC_A2,FE:VA;0) 298.15 2.22; 6000 N !",
        " PARAMETER L(LIQUID,FE,W;0) 298.15 -15000; 6000 N !", " PARAMETER L(BCC_A2,FE,W:VA;0) 298.15 10000; 6000 N !",
        " PARAMETER G(C14,FE:W;0) 298.15 0.66667*GHSERFE+0.33333*GHSERWW-7000; 6000 N !",
        " PARAMETER G(C14,W:W;0) 298.15 GHSERWW+5000; 6000 N !",
        " PARAMETER L(C14,FE,W:W;0) 298.15 -3000; 6000 N !",
        " PARAMETER TC(C14,FE:W;0) 298.15 300; 6000 N !", " PARAMETER BMAGN(C14,FE:W;0) 298.15 0.6; 6000 N !"],
        typedefs=[" TYPE_DEFINITION M GES A_P_D C14 MAGNETIC -3 0.28 !"])


def test_force_replot_exact_rescale_fe_mo(tmp_path):
    from pptx import Presentation
    a, b = write(tmp_path, "A_FeMo.tdb", A_FeMo()), write(tmp_path, "B_FeW.tdb", _fe_w_magnetic_c14())
    kw = dict(policy="ask", ternary_temperatures_c=[], tmin_c=800, tmax_c=2600, composition_step=0.05,
              temperature_step=25, n_workers=3)
    r = staged_merge(a, b, str(tmp_path / "out"), plot=True, force_replot=["mo-FE"], **kw)
    assert r["A_eff"] != r["A"]                                         # the C14 magnetic basis rule applied
    st = {s.title: s for s in r["steps"]}
    a1 = st["A → A1"]
    assert not a1.changed_items and [it.name for it in a1.forced_items] == ["Fe-Mo"]
    assert a1.forced_kind["Fe-Mo"] == "exact rescale"
    dT, dx, ok = a1.forced_shift("Fe-Mo")
    assert ok and dT < 0.5 and dx < 1e-3 and a1.forced_inv["Fe-Mo"]       # really compared, invariants matched
    assert "reused_from" not in r["plots"][("A1", "Fe-Mo")]               # recomputed, not reused
    assert {("A1", "Fe-Mo"), ("A2", "Fe-Mo")} <= set(r["forced"])
    assert not any(k[1] == "Fe-W" for k in r["forced"])                  # only the named subsystem
    label = a1.forced_label(a1.forced_items[0])
    assert label.startswith("Fe-Mo: exact rescale, replotted on request - max invariant ΔT 0.00 K")
    text = summary_text(r)
    assert "forced replots: Fe-Mo (exact rescale): max ΔT 0.00 K" in text and "FAILED" not in text
    prs = Presentation(r["files"]["deck"])
    titles = [[sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()][0]
              for sl in prs.slides]
    # summary, A -> A0 (the magnetic basis rule's rescale), each step's TDB changes, each step's diagrams
    assert len(prs.slides) == 1 + sum(": TDB changes" in t for t in titles) + len(r["steps"])
    assert any(t.startswith("Normalisation A → A0: TDB changes") for t in titles)
    i = titles.index("A → A1: diagrams that changed + exact changes replotted on request")
    assert Path(r["files"]["deck"]).exists() and i > 0
    assert any("compare_A-A1" in p for p in r["images"])
    # the forced replot is in the summary row
    rows = [[c.text for c in row.cells] for sh in prs.slides[0].shapes if sh.has_table for row in sh.table.rows]
    assert any("replotted on request: Fe-Mo" in c for row in rows for c in row)


def test_force_replot_true_covers_all_exact_only_and_default_false(tmp_path):
    a, b = write(tmp_path, "A_FeMo.tdb", A_FeMo()), write(tmp_path, "B_FeW.tdb", _fe_w_magnetic_c14())
    r0 = staged_merge(a, b, str(tmp_path / "o0"), plot=False, ternary_temperatures_c=[])
    assert r0["forced"] == [] and all(not s.forced_items for s in r0["steps"])
    r1 = staged_merge(a, b, str(tmp_path / "o1"), plot=False, ternary_temperatures_c=[], force_replot=True)
    exact_only = {(s.new, x.name) for s in r1["steps"] for x in s.diff.subsystems if not x.changed}
    assert exact_only and set(r1["forced"]) == exact_only               # every unchanged subsystem of every step
    assert ("B1", "Fe-W") in exact_only and ("A1", "Fe-Mo") in exact_only
    r2 = staged_merge(a, b, str(tmp_path / "o2"), plot=False, ternary_temperatures_c=[], force_replot=["Cr-Ni"])
    assert r2["forced"] == [] and "Cr-Ni" in summary_text(r2)


def test_forced_replot_shift_is_flagged():
    from calphad_tools.staged_merge import Item, Step
    st = Step("A", "A1", ts.VersionDiff("A", "A1", [], []), [])
    it = Item(("FE", "MO"))
    st.forced_items, st.forced_kind = [it], {"Fe-Mo": "exact rescale"}
    st.forced_inv["Fe-Mo"] = [{"invariant": "BCC_A2 + C14 + LIQUID", "type": "eutectic", "T_old_K": 1500.0,
                               "T_new_K": 1501.2, "dT_K": 1.2, "dx": 0.0, "status": "same"}]
    assert st.forced_shift("Fe-Mo")[2] is False
    assert st.forced_label(it).endswith("SHIFTED: the exactness claim FAILED")
