"""evaluate_compatibility (staged TDB compatibility check) and the B -> A normalisation (phase names, sublattice
order, formula-unit rescaling) on small synthetic databases."""
import subprocess
import sys
import warnings
from pathlib import Path

import numpy as np
import pytest

from calphad_tools import tdbstage as ts
from calphad_tools.compatibility import evaluate_pair, locate_unary50, run, summary_text
from calphad_tools.tdbparse import parse_tdb

ELS = {"FE": " ELEMENT FE BCC_A2 55.847 4489 27.28 !", "MO": " ELEMENT MO BCC_A2 95.94 4589 28.56 !",
       "W": " ELEMENT W BCC_A2 183.84 4973 32.66 !", "NA": " ELEMENT NA BCC_A2 22.99 6447 51.3 !",
       "CL": " ELEMENT CL GAS 35.45 0 0 !", "K": " ELEMENT K BCC_A2 39.1 7088 64.7 !",
       "BR": " ELEMENT BR LIQUID 79.9 0 0 !"}
GHSER = {
    "FE": "1225.7+124.134*T-23.5143*T*LN(T)-.00439752*T**2-5.8927E-08*T**3+77359*T**(-1); 1811 Y "
          "-25383.581+299.31255*T-46*T*LN(T)+2.29603E+31*T**(-9); 6000 N",
    "MO": "-7746.302+131.9197*T-23.56414*T*LN(T)-.003443396*T**2+5.66283E-07*T**3+65812*T**(-1)"
          "-1.30927E-10*T**4; 2896 Y -30556.41+283.559746*T-42.63829*T*LN(T)-4.849315E+33*T**(-9); 5000 N",
    "W": "-7646.311+130.4*T-24.1*T*LN(T)-.001936*T**2+2.07E-07*T**3-5.33E-11*T**4+44500*T**(-1); 3695 Y "
         "-82868.801+389.362335*T-54*T*LN(T)+1.528621E+33*T**(-9); 6000 N"}
FNAME = {"FE": "GHSERFE", "MO": "GHSERMO", "W": "GHSERWW"}


def tdb_text(els, phases, params, shift=None, typedefs=()):
    shift = shift or {}
    L = [" ELEMENT VA VACUUM 0 0 0 !"] + [ELS[e] for e in els]
    for e in els:
        if e in GHSER:
            L.append(f" FUNCTION {FNAME[e]} 298.15 {shift.get(e, '')}{GHSER[e]} !")
    L += [" TYPE_DEFINITION % SEQ * !", " TYPE_DEFINITION B GES A_P_D BCC_A2 MAGNETIC -1 0.4 !", *typedefs]
    return "\n".join(L + phases + params) + "\n"


def A_FeMo(**kw):
    return tdb_text(["FE", "MO"], [
        " PHASE LIQUID % 1 1 !", " CONSTITUENT LIQUID :FE,MO: !", " PHASE BCC_A2 %B 2 1 3 !",
        " CONSTITUENT BCC_A2 :FE,MO:VA: !", " PHASE C14 % 2 2 1 !", " CONSTITUENT C14 :FE:MO: !"], [
        " PARAMETER G(LIQUID,FE;0) 298.15 12040.17-6.55843*T+GHSERFE; 6000 N !",
        " PARAMETER G(LIQUID,MO;0) 298.15 34085.045-12.47*T+GHSERMO; 5000 N !",
        " PARAMETER G(BCC_A2,FE:VA;0) 298.15 GHSERFE; 6000 N !", " PARAMETER G(BCC_A2,MO:VA;0) 298.15 GHSERMO; 5000 N !",
        " PARAMETER TC(BCC_A2,FE:VA;0) 298.15 1043; 6000 N !", " PARAMETER BMAGN(BCC_A2,FE:VA;0) 298.15 2.22; 6000 N !",
        " PARAMETER L(LIQUID,FE,MO;0) 298.15 -6973-0.37*T; 6000 N !",
        " PARAMETER L(BCC_A2,FE,MO:VA;0) 298.15 36818-9.141*T; 6000 N !",
        " PARAMETER G(C14,FE:MO;0) 298.15 2*GHSERFE+GHSERMO-15000+2*T; 6000 N !"], **kw)


def B_FeMoW(**kw):
    return tdb_text(["FE", "MO", "W"], [
        " PHASE LIQUID % 1 1 !", " CONSTITUENT LIQUID :FE,MO,W: !", " PHASE BCC_A2 %B 2 1 3 !",
        " CONSTITUENT BCC_A2 :FE,MO,W:VA: !", " PHASE C14 % 2 2 1 !", " CONSTITUENT C14 :FE,W:FE,W: !"], [
        " PARAMETER G(LIQUID,FE;0) 298.15 12040.17-6.55843*T+GHSERFE; 6000 N !",
        " PARAMETER G(LIQUID,MO;0) 298.15 34085.045-12.47*T+GHSERMO; 5000 N !",
        " PARAMETER G(LIQUID,W;0) 298.15 52160.584-14.10999*T+GHSERWW; 6000 N !",
        " PARAMETER G(BCC_A2,FE:VA;0) 298.15 GHSERFE; 6000 N !", " PARAMETER G(BCC_A2,MO:VA;0) 298.15 GHSERMO; 5000 N !",
        " PARAMETER G(BCC_A2,W:VA;0) 298.15 GHSERWW; 6000 N !",
        " PARAMETER TC(BCC_A2,FE:VA;0) 298.15 1043; 6000 N !", " PARAMETER BMAGN(BCC_A2,FE:VA;0) 298.15 2.22; 6000 N !",
        " PARAMETER L(LIQUID,FE,MO;0) 298.15 -6973-0.37*T; 6000 N !",
        " PARAMETER L(BCC_A2,FE,MO:VA;0) 298.15 36818-9.141*T; 6000 N !",
        " PARAMETER L(LIQUID,FE,W;0) 298.15 -15000; 6000 N !", " PARAMETER L(BCC_A2,FE,W:VA;0) 298.15 10000; 6000 N !",
        " PARAMETER G(C14,FE:W;0) 298.15 2*GHSERFE+GHSERWW-20000+2*T; 6000 N !",
        " PARAMETER G(C14,FE:FE;0) 298.15 3*GHSERFE+5000; 6000 N !",
        " PARAMETER G(C14,W:W;0) 298.15 3*GHSERWW+5000; 6000 N !",
        " PARAMETER G(C14,W:FE;0) 298.15 2*GHSERWW+GHSERFE+30000; 6000 N !",
        " PARAMETER L(C14,FE:FE,W;0) 298.15 -5000; 6000 N !"], **kw)


def write(tmp_path, name, text):
    p = tmp_path / name
    p.write_text(text)
    return str(p)


@pytest.fixture
def c14_pair(tmp_path):
    return write(tmp_path, "A_FeMo.tdb", A_FeMo()), write(tmp_path, "B_FeMoW.tdb", B_FeMoW())


# ------------------------------------------------------------------------------------------ stages
def test_identical_unaries_stage1_ok_and_c14_expansion(c14_pair, tmp_path):
    a, b = c14_pair
    ev = run(a, b, str(tmp_path / "out" / "compat"), unary50="off")
    st = {s.n: s for s in ev.stages}
    assert st[1].verdict == "OK" and st[1].line.startswith("OK (Fe, Mo equal")
    assert st[2].verdict == "OK"
    # stage 5: merged C14 model and the new FE:FE end-member inside A's Fe-Mo
    assert "merged (FE,W)2(FE,MO,W)1" in st[5].line
    new = [r for r in st[5].rows if r.get("new end-member") == "FE:FE"]
    assert new and new[0]["in space of"] == "A" and "Fe-Mo" in new[0]["affects subsystems"]
    assert new[0]["provided by"] == "G from B"
    # W:MO exists in neither database -> flagged (pycalphad would use G = 0)
    assert any(r.get("new end-member") == "W:MO" and "no G" in r["provided by"] for r in st[5].rows)
    # stage 6: A's Fe-Mo needs a replot because of the model change; B's Fe-W is provably unchanged
    a_fe_mo = next(s for s in ev.replot["A"] if s.name == "Fe-Mo")
    assert a_fe_mo.changed and any(m.phase == "C14" and "FE:FE" in m.new_endmembers for m in a_fe_mo.models)
    b_fe_w = next(s for s in ev.replot["B"] if s.name == "Fe-W")
    assert not b_fe_w.changed and b_fe_w.reasons().startswith("no change")
    row = next(r for r in st[6].rows if r["database"] == "A" and r["subsystem"] == "Fe-Mo")
    assert row["replot"] == "yes"
    text = summary_text(ev)
    assert "Stage 6" in text and len(text) < 3000
    assert Path(ev.files["xlsx"]).exists() and Path(ev.files["pptx"]).exists()
    from pptx import Presentation
    prs = Presentation(ev.files["pptx"])
    first = [sh for sh in prs.slides[0].shapes if sh.has_table][0].table
    assert [r.cells[1].text for r in list(first.rows)[1:7]] == [s.verdict for s in ev.stages]
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_table:
                assert (sh.top + sh.height) / 914400 < 7.4
                assert not any(c.text.rstrip().endswith("…") for r in sh.table.rows for c in r.cells)


def test_clip_cuts_at_clause_boundaries():
    from calphad_tools.compatibility import clip
    t = "a long first clause (with; inner) here; second clause; third clause that is long"
    out = clip(t, 50, "Excel")
    assert out.startswith("a long first clause (with; inner) here") and out.endswith("+2 more (Excel)")
    assert clip("short", 50) == "short"


def test_shifted_ghser_incompatible_with_size_and_unary50(tmp_path):
    a = write(tmp_path, "A.tdb", A_FeMo())
    b = write(tmp_path, "B.tdb", A_FeMo(shift={"MO": "1200+"}))       # GHSERMO + 1.2 kJ/mol
    fake = write(tmp_path, "unary_fake.tdb", A_FeMo())                 # "modern" unary = A's functions
    ev = evaluate_pair(a, b, unary50=fake, workdir=str(tmp_path / "w"))
    s1 = ev.stages[0]
    assert s1.verdict == "INCOMPATIBLE"
    row = next(r for r in s1.rows if r["element"] == "MO")
    assert row["status"] == "differs" and abs(row["max |Δ| (J/mol)"] - 1200) < 1e-6
    assert row["unary50"] == "A matches"
    assert "1.2 kJ/mol" in s1.line and "A matches unary50" in s1.line
    assert next(r for r in s1.rows if r["element"] == "FE")["status"] == "equal"


def test_unary50_locate_offline(tmp_path, monkeypatch):
    monkeypatch.delenv("CALPHAD_UNARY50", raising=False)
    path, note = locate_unary50("auto", cache_dir=str(tmp_path / "cache"), allow_download=False)
    assert path is None and "skipped" in note
    path, note = locate_unary50(str(tmp_path / "missing.tdb"))
    assert path is None and "not found" in note
    assert locate_unary50("off")[0] is None


def test_identical_subsystem_no_change(c14_pair):
    a, _ = c14_pair
    V1, V2 = ts.StageView(parse_tdb(a), "x"), ts.StageView(parse_tdb(a), "y")
    d = ts.diff_versions(V1, V2, [("FE", "MO")])
    assert not d.changed and not d.params
    # one numerically equal rewrite (different text) is still "no change"
    text = Path(a).read_text().replace("36818-9.141*T", "36818.0000-9.1410*T")
    V3 = ts.StageView(parse_tdb(text, is_text=True, name="z"), "z")
    assert not ts.diff_versions(V1, V3, [("FE", "MO")]).changed


# ------------------------------------------------------------------------------------------ normalisation
HDR = ["FE", "MO", "W"]


def _gm(db_path, phase, sitefracs, T, comps):
    from pycalphad import Database, Model, variables as v
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        m = Model(Database(db_path), comps, phase)
    subs = {v.T: T, v.P: 101325.0}
    for (i, sp), y in sitefracs.items():
        subs[v.SiteFraction(phase, i, sp)] = y
    return float(m.GM.subs(subs))


def _check_gm(orig, orig_phase, new, new_phase, perm_new_from_old, species, n=4, seed=0):
    """GM per mole of atoms of B's phase, before and after normalisation, at random T and site fractions."""
    rng = np.random.default_rng(seed)
    comps = sorted({s for sl in species for s in sl}) + ["VA"]
    worst = 0.0
    for _ in range(n):
        T = float(rng.uniform(300, 2500))
        y_old = {}
        for i, sl in enumerate(species):
            f = rng.uniform(0.05, 1.0, len(sl))
            f /= f.sum()
            for sp, y in zip(sl, f):
                y_old[(i, sp)] = float(y)
        y_new = {(perm_new_from_old.index(i), sp): y for (i, sp), y in y_old.items()}
        g0 = _gm(orig, orig_phase, y_old, T, comps)
        g1 = _gm(new, new_phase, y_new, T, comps)
        worst = max(worst, abs(g0 - g1) / abs(g0))
    return worst


def _normalise(tmp_path, a_text, b_text, **kw):
    a, b = write(tmp_path, "A.tdb", a_text), write(tmp_path, "B.tdb", b_text)
    bn, notes = ts.normalize_to_base(parse_tdb(a), parse_tdb(b), **kw)
    return b, ts.write_tdb(bn, str(tmp_path / "Bn.tdb")), notes


def test_laves_c15_rename_parameters_follow(tmp_path):
    a = tdb_text(HDR, [" PHASE LAVES_C15 % 2 2 1 !", " CONSTITUENT LAVES_C15 :FE:MO: !"],
                 [" PARAMETER G(LAVES_C15,FE:MO;0) 298.15 -10000; 6000 N !"])
    b = tdb_text(HDR, [" PHASE C15_LAVES % 2 2 1 !", " CONSTITUENT C15_LAVES :FE,W:W: !"],
                 [" PARAMETER G(C15_LAVES,FE:W;0) 298.15 -20000; 6000 N !",
                  " PARAMETER G(C15_LAVES,W:W;0) 298.15 15000; 6000 N !"])
    _, bn, notes = _normalise(tmp_path, a, b)
    assert any(n.action == "renamed" and n.phase == "C15_LAVES" and "LAVES_C15" in n.detail for n in notes)
    db = parse_tdb(bn)
    assert "LAVES_C15" in db.phases and "C15_LAVES" not in db.phases
    assert sorted(p.raw_name for p in db.parameters) == ["G(LAVES_C15,FE:W;0)", "G(LAVES_C15,W:W;0)"]
    assert ts.phase_tokens("C15") == ts.phase_tokens("LAVES_C15") == ts.phase_tokens("C15_LAVES")


C14_A = tdb_text(HDR, [" PHASE C14 % 2 2 1 !", " CONSTITUENT C14 :FE:MO: !"],
                 [" PARAMETER G(C14,FE:MO;0) 298.15 2*GHSERFE+GHSERMO-15000; 6000 N !"])
C14_B_PARAMS = [" PARAMETER G(C14,W:FE;0) 298.15 GHSERWW+2*GHSERFE-30000+3*T; 6000 N !",
                " PARAMETER G(C14,FE:FE;0) 298.15 3*GHSERFE+15000; 6000 N !",
                " PARAMETER G(C14,W:W;0) 298.15 3*GHSERWW+15000; 6000 N !",
                " PARAMETER G(C14,FE:W;0) 298.15 GHSERFE+2*GHSERWW+40000; 6000 N !",
                " PARAMETER L(C14,FE,W:FE;0) 298.15 -10000+T; 6000 N !",
                " PARAMETER L(C14,FE,W:FE;1) 298.15 3000; 6000 N !",
                " PARAMETER L(C14,W:FE,W;0) 298.15 5000; 6000 N !"]


def test_c14_2to1_vs_1to2_permutation_by_size(tmp_path):
    # A: Fe2Mo (FE)2(MO)1;  B: W:Fe 1:2 written (FE,W)1(FE,W)2 -> B's sublattice 1 (2 sites) goes first
    b_text = tdb_text(HDR, [" PHASE C14 % 2 1 2 !", " CONSTITUENT C14 :FE,W:FE,W: !"], C14_B_PARAMS)
    b, bn, notes = _normalise(tmp_path, C14_A, b_text)
    assert [n.action for n in notes] == ["reordered"]
    ph = parse_tdb(bn).phases["C14"]
    assert ph.site_ratios == [2.0, 1.0]
    raw = {p.raw_name for p in parse_tdb(bn).parameters}
    assert "G(C14,FE:W;0)" in raw and "L(C14,FE:FE,W;0)" in raw and "L(C14,FE,W:W;0)" in raw
    assert _check_gm(b, "C14", bn, "C14", [1, 0], [["FE", "W"], ["FE", "W"]]) < 1e-9


def test_equal_size_sublattices_decided_by_element_overlap(tmp_path):
    a = tdb_text(HDR, [" PHASE PH11 % 2 1 1 !", " CONSTITUENT PH11 :FE:MO: !"],
                 [" PARAMETER G(PH11,FE:MO;0) 298.15 GHSERFE+GHSERMO-5000; 6000 N !"])
    b = tdb_text(HDR, [" PHASE PH11 % 2 1 1 !", " CONSTITUENT PH11 :W:FE,W: !"],
                 [" PARAMETER G(PH11,W:FE;0) 298.15 GHSERWW+GHSERFE-8000; 6000 N !",
                  " PARAMETER G(PH11,W:W;0) 298.15 2*GHSERWW+3000; 6000 N !",
                  " PARAMETER L(PH11,W:FE,W;0) 298.15 -4000+2*T; 6000 N !"])
    b, bn, notes = _normalise(tmp_path, a, b)
    assert [n.action for n in notes] == ["reordered"]
    assert parse_tdb(bn).phases["PH11"].constituents == [["FE", "W"], ["W"]]
    assert _check_gm(b, "PH11", bn, "PH11", [1, 0], [["W"], ["FE", "W"]]) < 1e-9


def test_rock_salt_no_common_elements_raises_clear_error(tmp_path):
    a = tdb_text(["NA", "CL"], [" PHASE ROCKSALT % 2 1 1 !", " CONSTITUENT ROCKSALT :NA:CL: !"],
                 [" PARAMETER G(ROCKSALT,NA:CL;0) 298.15 -400000; 6000 N !"])
    b = tdb_text(["K", "BR"], [" PHASE ROCKSALT % 2 1 1 !", " CONSTITUENT ROCKSALT :K:BR: !"],
                 [" PARAMETER G(ROCKSALT,K:BR;0) 298.15 -390000; 6000 N !"])
    with pytest.raises(ts.SublatticeMapError) as e:
        _normalise(tmp_path, a, b)
    msg = str(e.value)
    assert "ROCKSALT" in msg and "share no element" in msg and "sublattice_map {'ROCKSALT'" in msg
    # the override decides (B's sublattice 0 -> A's position 1)
    _, bn, notes = _normalise(tmp_path, a, b, sublattice_map={"ROCKSALT": [1, 0]})
    assert [n.action for n in notes] == ["reordered"]
    assert [p.raw_name for p in parse_tdb(bn).parameters] == ["G(ROCKSALT,BR:K;0)"]
    # ... or keeps B's model as it is
    _, bn, notes = _normalise(tmp_path, a, b, sublattice_map={"ROCKSALT": "keep"})
    assert [n.action for n in notes] == ["kept"]


def test_overlap_pointing_both_ways_raises(tmp_path):
    a = tdb_text(HDR, [" PHASE PH11 % 2 1 1 !", " CONSTITUENT PH11 :FE:MO: !"], [])
    b = tdb_text(HDR, [" PHASE PH11 % 2 1 1 !", " CONSTITUENT PH11 :FE,MO:W: !"], [])
    with pytest.raises(ts.SublatticeMapError, match="points both ways"):
        _normalise(tmp_path, a, b)


def test_4to2_rescaled_to_2to1_gm_per_atom_identical(tmp_path):
    b_text = tdb_text(HDR, [" PHASE C14 % 2 4 2 !", " CONSTITUENT C14 :FE,W:FE,W: !"],
                      [p.replace("3*GHSER", "6*GHSER").replace("GHSERWW+2*GHSERFE", "2*GHSERWW+4*GHSERFE")
                       .replace("GHSERFE+2*GHSERWW", "2*GHSERFE+4*GHSERWW") for p in C14_B_PARAMS])
    b, bn, notes = _normalise(tmp_path, C14_A, b_text)
    acts = [n.action for n in notes]
    assert "rescaled" in acts and "reordered" not in acts
    assert parse_tdb(bn).phases["C14"].site_ratios == [2.0, 1.0]
    assert _check_gm(b, "C14", bn, "C14", [0, 1], [["FE", "W"], ["FE", "W"]]) < 1e-9


MAG = " TYPE_DEFINITION M GES A_P_D PHM MAGNETIC -1 0.4 !"


def test_magnetic_phase_permuted_exactly_but_never_rescaled(tmp_path):
    a = tdb_text(HDR, [" PHASE PHM %M 2 2 1 !", " CONSTITUENT PHM :FE:MO: !"],
                 [" PARAMETER G(PHM,FE:MO;0) 298.15 2*GHSERFE+GHSERMO; 6000 N !"], typedefs=[MAG])
    mag_params = [" PARAMETER G(PHM,W:FE;0) 298.15 GHSERWW+2*GHSERFE-9000; 6000 N !",
                  " PARAMETER G(PHM,W:W;0) 298.15 3*GHSERWW+9000; 6000 N !",
                  " PARAMETER L(PHM,W:FE,W;0) 298.15 -7000; 6000 N !",
                  " PARAMETER TC(PHM,W:FE;0) 298.15 1043; 6000 N !",
                  " PARAMETER TC(PHM,W:FE,W;0) 298.15 -300; 6000 N !",
                  " PARAMETER BMAGN(PHM,W:FE;0) 298.15 2.22; 6000 N !",
                  " PARAMETER BMAGN(PHM,W:FE,W;1) 298.15 0.5; 6000 N !"]
    # 1:2 written order -> pure permutation (k = 1): exact, TC/BMAGN permuted along
    b = tdb_text(HDR, [" PHASE PHM %M 2 1 2 !", " CONSTITUENT PHM :W:FE,W: !"], mag_params, typedefs=[MAG])
    bp, bn, notes = _normalise(tmp_path, a, b)
    assert [n.action for n in notes] == ["reordered"]
    assert "TC(PHM,FE,W:W;0)" in {p.raw_name for p in parse_tdb(bn).parameters}
    assert _check_gm(bp, "PHM", bn, "PHM", [1, 0], [["W"], ["FE", "W"]]) < 1e-9
    # 2:4 (k = 1/2): pycalphad's IHJ term R*T*ln(beta+1)*f(tau) is per formula unit and not linear in BMAGN ->
    # not rescaled (reported), so the phase's GM is untouched ...
    b2 = tdb_text(HDR, [" PHASE PHM %M 2 2 4 !", " CONSTITUENT PHM :W:FE,W: !"],
                  [p.replace("GHSERWW+2*GHSERFE", "2*GHSERWW+4*GHSERFE").replace("3*GHSERWW", "6*GHSERWW")
                   for p in mag_params], typedefs=[MAG])
    tmp2 = tmp_path / "k"
    tmp2.mkdir()
    bp2, bn2, notes2 = _normalise(tmp2, a, b2)
    assert "not_rescaled" in [n.action for n in notes2]
    assert parse_tdb(bn2).phases["PHM"].site_ratios == [4.0, 2.0]          # reordered only
    assert _check_gm(bp2, "PHM", bn2, "PHM", [1, 0], [["W"], ["FE", "W"]]) < 1e-9
    # ... whereas the naive rescale (G/L x k, TC/BMAGN unchanged) would change GM: the reason it is refused
    naive = parse_tdb(bn2)
    naive.phases["PHM"].site_ratios = [2.0, 1.0]
    for p in naive.parameters:
        if p.key.phase == "PHM" and p.key.ptype == "G":
            p.expr = ts._scale_pw(p.expr, 0.5)
    nv = ts.write_tdb(naive, str(tmp2 / "naive.tdb"))
    assert _check_gm(bp2, "PHM", nv, "PHM", [1, 0], [["W"], ["FE", "W"]]) > 1e-6


def test_not_proportional_is_reported_not_rescaled(tmp_path):
    a = tdb_text(HDR, [" PHASE SIGMA % 3 10 4 16 !", " CONSTITUENT SIGMA :FE:MO:FE,MO: !"], [])
    b = tdb_text(HDR, [" PHASE SIGMA_D8B % 3 8 4 18 !", " CONSTITUENT SIGMA_D8B :FE:MO,W:FE,MO,W: !"], [])
    _, bn, notes = _normalise(tmp_path, a, b)
    assert "SIGMA_D8B" in parse_tdb(bn).phases                  # not renamed: models differ
    assert any(n.action == "candidate" and "site ratios differ" in n.detail for n in notes)


def test_evaluate_reports_normalisation(tmp_path):
    b_text = tdb_text(HDR, [" PHASE LAVES_C14 % 2 1 2 !", " CONSTITUENT LAVES_C14 :FE,W:FE,W: !"],
                      [p.replace("(C14,", "(LAVES_C14,") for p in C14_B_PARAMS])
    a, b = write(tmp_path, "A.tdb", C14_A), write(tmp_path, "B.tdb", b_text)
    ev = evaluate_pair(a, b, unary50="off", workdir=str(tmp_path / "w"))
    text = summary_text(ev)
    assert "LAVES_C14: renamed -> C14" in text and "reordered" in text
    assert ev.normalised_b and Path(ev.normalised_b).exists()
    st5 = ev.stages[4]
    assert any(r.get("new end-member") == "FE:FE" for r in st5.rows)


# ------------------------------------------------------------------------------------------ tool + show_code
def test_tool_registered_and_show_code(c14_pair, tmp_path):
    from calphad_tools.agent.tools import UI, ToolBox
    a, b = c14_pair
    tb = ToolBox(str(tmp_path), "out", UI(), python_tool="off")
    assert "evaluate_compatibility" in tb.tools and "staged_merge" in tb.tools
    compact = {s.name: s.description for s in tb.model_specs()}
    assert len(compact["evaluate_compatibility"]) < 260
    text, imgs, err = tb.call("evaluate_compatibility", {"base_tdb": a, "add_tdb": b, "unary50": "off"})
    assert not err, text
    assert "Stage 5 phase models" in text
    text, _, err = tb.call("show_code", {"target": "1", "format": "script"})
    assert not err, text
    script = next(Path(tmp_path).glob("out/code/*_standalone.py"))
    assert "from calphad_tools" not in script.read_text()          # standalone, not the labelled fallback
    r = subprocess.run([sys.executable, str(script)], cwd=str(tmp_path), capture_output=True, text=True,
                       timeout=600)
    assert r.returncode == 0, r.stderr[-2000:]
    assert "Stage 6 replot decision" in r.stdout and "FE:FE" in r.stdout
