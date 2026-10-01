"""phase_inventory and system_coverage."""
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TDBS = ROOT / "examples" / "tdbs"


def _tb(workdir, out):
    from calphad_tools.agent.tools import UI, ToolBox

    class U(UI):
        def ask(self, q, o=None):
            return ""

        def confirm(self, t, b):
            return False
    return ToolBox(workdir=str(workdir), output_dir=str(out), ui=U(), python_tool="off")


def ok(tb, name, args):
    text, images, err = tb.call(name, args)
    assert not err, text
    return text


def test_coverage_counts_user_example(tmp_path):
    wd = tmp_path / "wd" / "two"
    wd.mkdir(parents=True)
    shutil.copy(TDBS / "Cr-Fe_her.tdb", wd)
    shutil.copy(TDBS / "FeNi-00Ans-2SL.tdb", wd)
    tb = _tb(tmp_path / "wd", tmp_path / "out")
    text = ok(tb, "system_coverage", {"elements": ["Fe", "Cr", "Ni"], "folder": "two"})
    assert "binaries: 3 in total, 2 assessed" in text
    assert "ternaries: 1 in total, 0 assessed" in text
    assert "Not assessed (1): Cr-Ni" in text and "Cr-Fe-Ni (2/3 binaries)" in text
    from calphad_tools import system_coverage as sc
    from calphad_tools.agent.tools import ToolBox  # noqa: F401
    system, dbs, table = sc.coverage([str(wd / "Cr-Fe_her.tdb"), str(wd / "FeNi-00Ans-2SL.tdb")],
                                     ["FE", "CR", "NI"], tmp_path / "wd")
    assessed = {r["system"] for r in table[2] if r["assessed_in"]}
    assert assessed == {("CR", "FE"), ("FE", "NI")}
    out = tmp_path / "out"
    assert (out / "coverage_Cr-Fe-Ni.xlsx").exists() and (out / "coverage_Cr-Fe-Ni.pptx").exists()


def test_coverage_larger_system_and_quaternaries(tmp_path):
    from math import comb
    tb = _tb(ROOT / "examples", tmp_path / "out")
    els = ["FE", "NI", "CR", "MO", "CO", "V", "CU"]
    text = ok(tb, "system_coverage", {"elements": els, "max_order": 4})
    assert f"binaries: {comb(7, 2)} in total" in text and f"ternaries: {comb(7, 3)} in total" in text
    assert f"quaternaries: {comb(7, 4)} in total" in text
    assert "elements in no database: Cu, Mo" in text
    assert "Cr-Fe: " in text and "Cr-Fe_her (" in text and "Co-Fe-Ni (3/3 binaries)" in text
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "out" / "coverage_Co-Cr-Cu-Fe-Mo-Ni-V.xlsx")
    assert {"Summary", "Binary matrix", "Binaries", "Ternaries", "Quaternaries"} <= set(wb.sheetnames)
    assert wb["Ternaries"].max_row == comb(7, 3) + 1


def test_phase_inventory(tmp_path):
    tb = _tb(ROOT / "examples", tmp_path / "out")
    text = ok(tb, "phase_inventory", {"elements": ["CR", "FE"], "folder": "tdbs"})
    assert "3 database(s)" in text
    assert "BCC_A2 (BCC): (Cr,Fe)₁(Va)₃ [substitutional solution" in text
    assert "(Fe)₁₀(Cr)₄(Cr,Fe)₁₆" in text and "AFM factor -1, p = 0.4" in text
    assert "SIGMA [" in text and "~ SIGMA_D8B" in text              # candidate, not merged
    assert "one end-member inside the system" in text               # L12 in Choi2019 is not a line compound
    from openpyxl import load_workbook
    from pptx import Presentation
    wb = load_workbook(tmp_path / "out" / "phases_Cr-Fe.xlsx")
    assert {"Phases", "Matrix", "Databases", "Phase candidates"} <= set(wb.sheetnames)
    prs = Presentation(str(tmp_path / "out" / "phases_Cr-Fe.pptx"))
    W, H = prs.slide_width, prs.slide_height
    for s in prs.slides:
        for sh in s.shapes:
            assert sh.left + sh.width <= W * 1.001 and sh.top + sh.height <= H * 1.001
        # never a group heading alone at the bottom of a table page
        for sh in s.shapes:
            if sh.has_table:
                last = sh.table.rows[len(sh.table.rows) - 1]
                assert last.cells[0].text.strip() and not (last.cells[0].span_width > 1 and
                                                             len(sh.table.columns) > 1 and
                                                             all(not c.text for c in list(last.cells)[1:]))


def test_phase_inventory_single_database_and_subset(tmp_path):
    tb = _tb(ROOT / "examples", tmp_path / "out")
    text = ok(tb, "phase_inventory", {"elements": ["CR", "NI"], "tdbs": ["CrNi-06Hal.tdb"], "phases": ["LIQUID",
                                                                                                          "FCC_A1"],
                                      "pptx": False})
    assert text.startswith("2 phase(s) in 1 database(s)")
    assert not (tmp_path / "out" / "phases_Cr-Ni.pptx").exists()
