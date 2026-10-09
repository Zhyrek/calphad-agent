"""combine_databases: plan -> decisions (asked once, stable ids) -> staged merges -> decision log -> validation against
every source -> one report, on the bundled example TDBs:

* Fe-Nb (90Hua) + Fe-W (15Jac) + Nb-W (97Hua): the μ model conflict (decision 'rename') and C14 gaining Fe:Fe from
  Fe-W (constituent-union rule); the tentative L(C14_LAVES,FE:FE,NB;0) = +150000 (TENT_MIX) keeps Fe-Nb's C14
  stoichiometric, so Fe-Nb is reproduced (without it C14 reaches X(Nb) 0.321 and invariants move ~5 K);
* Cr-Ni (06Hal) + Fe-Ni (00Ans) + Cr-Fe (82Her): no shared binaries; Cr-Ni and Fe-Ni reproduced, Cr-Fe changed by the
  base's fcc-Cr lattice stability (82Her is pre-SGTE), the Cr-Fe-Ni ternary listed as predicted, not validated;
* a hand-edited combined TDB re-validated: CHANGED, UNEXPLAINED.

Binaries only, n_workers 3.  Run:  .venv/bin/python -m pytest -q tests/test_combine.py
"""
import json
import math
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

from calphad_tools import tdbstage as ts
from calphad_tools.combine import (combine_databases, exactness_checks, gm_identity, DecisionLog,
                                   validate_from_state)
from calphad_tools.tdbparse import parse_tdb

ROOT = Path(__file__).resolve().parents[1]
EX = ROOT / "examples" / "tdbs"
TRIO = [str(EX / n) for n in ("FeNb-90Hua-LB.tdb", "FeW-15Jac-static.tdb", "NbW-97Hua-LB.tdb")]
CRFENI = [str(EX / n) for n in ("CrFe-82Her.tdb", "CrNi-06Hal.tdb", "FeNi-00Ans-2SL.tdb")]
REJECT = {"step1:reject:db1:L12_FCC+A1_FCC": "drop", "step1:reject:db2:FCC_A1": "drop"}
REJECT_WHY = {"step1:reject:db1:L12_FCC+A1_FCC": "user: pycalphad ignores REJECT_PHASE; keep all phases",
              "step1:reject:db2:FCC_A1": "user: same"}


def _toolbox(out):
    from calphad_tools.agent.tools import ToolBox, UI

    class U(UI):
        def ask(self, q, o=None):
            return ""

        def confirm(self, t, b):
            return False
    return ToolBox(workdir=str(ROOT / "examples"), output_dir=str(out), ui=U(), python_tool="off")


# ------------------------------------------------------------------------------------------------- plan stage
def test_plan_lists_every_decision_once_with_stable_ids(tmp_path):
    """No decisions given: nothing is merged, and ALL questions of ALL steps come back in one list - here a μ model
    conflict in step 1 (Fe-Nb vs Fe-W) and another in step 2 (a third database with yet another μ)."""
    nbw = (EX / "NbW-97Hua-LB.tdb").read_text() + (
        "\n PHASE MU_D85 % 2 1 1 !\n CONSTITUENT MU_D85 :NB:W: !\n"
        " PARAMETER G(MU_D85,NB:W;0) 298.15 GHSERNB+GHSERWW-5000; 6000 N !\n")
    third = tmp_path / "NbW-mu.tdb"
    third.write_text(nbw)
    srcs = TRIO[:2] + [str(third)]
    kw = dict(order=srcs, unary50="off")
    r1 = combine_databases(srcs, str(tmp_path / "c"), **kw)
    r2 = combine_databases(srcs, str(tmp_path / "c"), **kw)
    assert not r1["written"] and not r2["written"]
    ids = [n.id for n in r1["plan"].pending]
    assert ids == [n.id for n in r2["plan"].pending]                       # stable between calls
    assert ids == ["step1:phase:MU_D85", "step2:phase:MU_D85"]             # both steps, asked at once
    t = r1["text"]
    assert t.startswith("NOT FINISHED") and "* step1:phase:MU_D85" in t and "* step2:phase:MU_D85" in t
    assert "'rename': keep both" in t and "'db1':" in t and "'db2':" in t
    assert "Fe-W" in t                                                     # consequences name the subsystems
    assert not list((tmp_path / "c").glob("step*_*/Q_*.tdb"))             # nothing merged
    # answering step 1 only: step 2 is still asked, step 1 is reported as answered
    r3 = combine_databases(srcs, str(tmp_path / "c"), decisions={"step1:phase:MU_D85": "rename"}, **kw)
    assert [n.id for n in r3["plan"].pending] == ["step2:phase:MU_D85"]
    assert "Already answered: step1:phase:MU_D85 = 'rename'" in r3["text"]
    # a policy fills what is missing - but never a phase-model answer that drops an element from a phase entirely
    # (db1 here drops W from MU_D85 in both steps): those stay questions, and each db option says what it drops
    r4 = combine_databases(srcs, str(tmp_path / "c"), policy="first", **kw)
    assert not r4["written"]
    assert [n.id for n in r4["plan"].pending] == ["step1:phase:MU_D85", "step2:phase:MU_D85"]
    assert all(n.answer is None and n.drops.get("db1") == ["W"] for n in r4["plan"].pending)
    t4 = r4["text"]
    assert "policy='first' not applied: 'db1' would drop W from the phase - asked" in t4
    assert re.search(r"'db1': FeNb-90Hua-LB's MU_D85 model only - drops W from MU_D85 \(FeW-15Jac-static's \d+ "
                     r"MU_D85 parameters lost: ", t4), t4


def test_proposed_base_and_order_have_reasons(tmp_path):
    r = combine_databases(CRFENI, str(tmp_path / "c"), unary50="off")
    plan = r["plan"]
    assert [Path(p).stem for p in plan.order] == ["CrNi-06Hal", "FeNi-00Ans-2SL", "CrFe-82Her"]
    assert "most complete unary set" in plan.order_reasons[0]
    assert "shares Ni" in plan.order_reasons[1] and "shares Cr, Fe" in plan.order_reasons[2]
    assert [n.id for n in plan.pending] == list(REJECT)
    # the unary conflict that base and order decide is said up front
    assert "CrFe-82Her's unary data differ" in r["text"] and "G(FCC_A1,CR:VA;0)" in r["text"]


# ------------------------------------------------------------------------------------------------- the trio
@pytest.fixture(scope="module")
def trio(tmp_path_factory):
    """Through the ToolBox, as the agent calls it (decision recorded with the user's reason)."""
    out = tmp_path_factory.mktemp("trio")
    tb = _toolbox(out)
    args = {"tdbs": ["FeNb-90Hua-LB.tdb", "FeW-15Jac-static.tdb", "NbW-97Hua-LB.tdb"],
            "order": ["FeNb-90Hua-LB.tdb", "FeW-15Jac-static.tdb", "NbW-97Hua-LB.tdb"], "unary50": "off",
            "n_workers": 3, "output_dir": str(out / "combine")}
    text, imgs, err = tb.call("combine_databases", args)
    assert not err and text.startswith("NOT FINISHED") and "step1:phase:MU_D85" in text
    text, imgs, err = tb.call("combine_databases", {
        **args, "decisions": {"step1:phase:MU_D85": "rename"},
        "decision_reasons": {"step1:phase:MU_D85": "user: keep both mu models, each binary as assessed",
                             "plan:order": "user: Fe-Nb first"}})
    assert not err, text
    state = json.loads((out / "combine" / "combine_state.json").read_text())
    return {"tb": tb, "text": text, "images": imgs, "state": state, "out": out / "combine"}


def _row(state, sub):
    return next(r for r in state["validation"]["rows"] if r["subsystem"] == sub)


def test_trio_runs_and_reply_is_short(trio):
    t = trio["text"]
    assert len(t) <= 2200 and t.startswith("combine_databases: FeNb-90Hua-LB -> FeW-15Jac-static -> NbW-97Hua-LB")
    assert "Verdicts: 3 reproduced, 0 changed (explained), 0 CHANGED (unexplained)" in t
    assert "Tentative (not assessed, 2): L(C14_LAVES,FE:FE,NB;0) = +150000 (TENT_MIX)" in t
    assert "Fe-Nb-W" in t and trio["images"]


def test_decision_log_who_why_evidence(trio):
    log = {e["id"]: e for e in trio["state"]["log"]}
    u = log["step1:phase:MU_D85"]
    assert u["who"] == "user" and u["reason"] == "keep both mu models, each binary as assessed"
    assert u["chosen"] == "rename" and any("kept as MU_D85_2" in x for x in u["evidence"])
    assert "Fe-W" in u["affected"]
    assert log["plan:order"]["who"] == "user" and log["plan:order"]["reason"] == "Fe-Nb first"
    c14 = log["step1:rule:constituents:C14_LAVES"]
    assert c14["who"] == "rule" and "union" in c14["reason"] and c14["affected"] == ["Fe-Nb"]
    ev = " ".join(c14["evidence"])
    assert "G(C14_LAVES,FE:FE;0) added" in ev and "FE:FE in A's space (Fe-Nb): G from B" in ev
    # the tentative Fe,Nb term the union needs: one rule entry, with the end-members as evidence
    t1 = log["step1:rule:tentative:L(C14_LAVES,FE:FE,NB;0)"]
    assert t1["who"] == "rule" and t1["kind"] == "tentative" and t1["chosen"] == "+150000 (TENT_MIX, tentative)"
    assert "rule 1 (TENT_MIX)" in t1["reason"] and t1["affected"] == ["Fe-Nb"]
    assert any("G(C14_LAVES,FE:FE) from FeW-15Jac-static.tdb and G(C14_LAVES,FE:NB) from FeNb-90Hua-LB.tdb" in x
               for x in t1["evidence"])
    assert "step1:rule:tentative:L(C14_LAVES,FE:NB,W;0)" in log
    z = [e for e in log.values() if e["kind"] == "zero-g"]
    assert z and "C14_LAVES not stable in the mapped Nb-W" in z[0]["evidence"][0] and z[0]["affected"] == ["Nb-W"]
    assert log["step1:rule:forced:Fe-W"]["evidence"][0].endswith("max invariant ΔT 0.00 K, Δx 0.0000")


def test_trio_validation_verdicts(trio):
    st = trio["state"]
    fenb = _row(st, "Fe-Nb")
    # C14 (Fe)2(Fe,Nb) with the tentative +150000: Fe-Nb as assessed (without it: ΔT 4.97 K, C14 down to 0.321)
    assert fenb["verdict"] == "reproduced" and fenb["explained_by"] == [
        "step1:rule:constituents:C14_LAVES", "step1:rule:tentative:L(C14_LAVES,FE:FE,NB;0)"]
    assert not fenb["unexplained"] and fenb["inv_max_dT"] <= 0.5 and fenb["grid"] >= 0.995
    assert fenb["tentative"] == ["L(C14_LAVES,FE:FE,NB;0)"] and "Fe-Nb is reproduced" in fenb["tentative_note"]
    data = json.loads(Path(fenb["F_png"]).with_suffix(".json").read_text())
    xs = [v for reg in data["regions"] for ph, r in reg["X_ranges"].items() if ph.startswith("C14") for v in r["X(NB)"]]
    assert abs(min(xs) - 1 / 3) < 1e-3 and abs(max(xs) - 1 / 3) < 1e-3
    few = _row(st, "Fe-W")
    assert few["verdict"] == "reproduced" and few["basis"].startswith("identical description")
    assert few["inv_max_dT"] <= 0.5 and few["grid"] >= 0.995
    nbw = _row(st, "Nb-W")
    assert nbw["verdict"] == "reproduced" and "step2:rule:from-base:C14_LAVES" in nbw["explained_by"]
    assert st["validation"]["extrapolations"]["ternaries"] == [["FE", "NB", "W"]]
    s = st["validation"]["sanity"]
    assert s["load"]["loaded"] and not s["load"]["errors"] and s["convergence"]["not_converged"] == 0
    assert [z["phase"] + " " + z["end-member"] for z in s["zero_g"]] == ["C14_LAVES W:NB"]


def test_decisions_in_session_provenance(trio):
    tb = trio["tb"]
    rec = tb.provenance
    ids = {d["id"] for d in rec.decisions}
    assert "step1:phase:MU_D85" in ids and "step1:rule:constituents:C14_LAVES" in ids
    lines = [json.loads(x) for x in Path(rec.log_path).read_text().splitlines()]
    dec = [x for x in lines if x.get("type") == "decision"]
    assert dec and all(x["tool"] == "combine_databases" for x in dec)
    call = [x for x in lines if x.get("type") == "call" and x["tool"] == "combine_databases"][-1]
    assert {x["call_index"] for x in dec} == {call["index"]}            # tied to the call that made them


def test_report_files_and_tdb_header(trio):
    st, out = trio["state"], trio["out"]
    for k in ("final TDB", "deck", "xlsx", "notebook", "report", "state"):
        assert Path(st["files"][k]).exists(), k
    md = Path(st["files"]["report"]).read_text()
    assert "## Decision log" in md and "## What this proves and what it cannot" in md
    assert "validation/overlays/S0_Fe-Nb.png" in md and "step1:rule:constituents:C14_LAVES" in md
    assert "## Tentative parameters" in md and "`L(C14_LAVES,FE:FE,NB;0)` | +150000 | TENT_MIX" in md
    lines = Path(st["final"]).read_text().splitlines()
    end = [i for i, ln in enumerate(lines) if ln.startswith("$ ====")][1]
    h = "\n".join(lines[:end + 1])
    assert "Combined by calphad_tools combine_databases" in h and "step1:phase:MU_D85 -> rename; user:" in h
    assert "Validation against every source: 3 reproduced, 0 changed (explained)" in h
    assert "L(C14_LAVES,FE:FE,NB;0) = +150000 (TENT_MIX; step1:rule:tentative:L(C14_LAVES,FE:FE,NB;0))" in h
    F = parse_tdb(st["final"])
    assert F.reference_list["TENT_MIX"].startswith("Tentative parameter added by calphad_tools when combining")
    assert "!" not in h
    from pycalphad import Database
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        assert {"MU_D85", "MU_D85_2", "C14_LAVES"} <= set(Database(st["final"]).phases)
    from openpyxl import load_workbook
    wb = load_workbook(st["files"]["xlsx"])
    assert {"Decision log", "Validation", "Invariants", "Properties", "GM identity",
            "Attribution"} <= set(wb.sheetnames)


def test_deck_layout(trio):
    from pptx import Presentation
    from pptx.util import Pt
    st = trio["state"]
    prs = Presentation(st["files"]["deck"])
    titles = []
    for sl in prs.slides:
        t = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(t[0] if t else "")
    assert titles[0].startswith("Combined: FeNb-90Hua-LB + FeW-15Jac-static + NbW-97Hua-LB")
    assert titles[1].startswith("Plan: base, order and decisions")
    assert any(t.startswith("Step 1 · Staged merge") for t in titles)
    assert any(t.startswith("Step 2 · B2 → B3: differences") for t in titles)
    assert any(t.startswith("Decision log") for t in titles)
    assert sum(t.startswith("Validation: ") for t in titles) == 3
    assert titles[-1].startswith("What this cannot prove")
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_table:
                assert (sh.top + sh.height) / 914400 < 7.45
                for row in sh.table.rows:
                    for c in row.cells:
                        assert not c.text.rstrip().endswith("…"), c.text
                        for p in c.text_frame.paragraphs:
                            for r in p.runs:
                                assert r.font.size is None or r.font.size >= Pt(10), (c.text, r.font.size)


def test_notebook_rebuilds_the_same_tdb(trio, tmp_path):
    st = trio["state"]
    nb = json.loads(Path(st["files"]["notebook"]).read_text())
    code = "\n\n".join("".join(c["source"]) for c in nb["cells"] if c["cell_type"] == "code")
    assert "HASH_SEED = '0'" in code and "[calphad_tools]" in json.dumps(nb)
    (tmp_path / "nb.py").write_text(code)
    env = {**os.environ, "COMBINE_NOTEBOOK_VALIDATION": "0", "MPLBACKEND": "Agg", "PYTHONHASHSEED": "0"}
    r = subprocess.run([sys.executable, "-W", "ignore", "nb.py"], cwd=tmp_path, capture_output=True, text=True,
                       env=env, timeout=600)
    assert r.returncode == 0, r.stderr[-3000:]
    assert "parameter-equivalent to the tool's combined TDB: True" in r.stdout
    rebuilt = tmp_path / "combine_reproduction" / "combined_rebuilt.tdb"
    X, Y = ts.StageView(parse_tdb(str(rebuilt)), "nb"), ts.StageView(parse_tdb(st["final"]), "tool")
    d = ts.diff_versions(X, Y, ts.stage_subsystems(sorted(X.own_elements), (2, 3)))
    assert not d.params and not d.changed and sorted(X.phases) == sorted(Y.phases)


def test_show_code_finds_the_combination(trio):
    tb = trio["tb"]
    text, _, err = tb.call("show_code", {"target": trio["state"]["files"]["report"]})
    assert not err, text
    code = (tb.out / "code" / "combine_report_standalone.py").read_text()
    assert "merge_step" in code and "STEPS = " in code and "'phase:MU_D85': 'rename'" in code
    assert "[calphad_tools]" in code and "validate_from_state" in code


# ------------------------------------------------------------------------------------------------- Cr-Fe-Ni
@pytest.fixture(scope="module")
def crfeni(tmp_path_factory):
    out = tmp_path_factory.mktemp("crfeni")
    r = combine_databases(CRFENI, str(out / "c"), decisions=REJECT, decision_reasons=REJECT_WHY, unary50="off",
                          n_workers=3, tmax_c=2000.0, lint=False)
    assert r["written"], r["text"]
    return r


def test_crfeni_sources_reproduced(crfeni):
    st, V = crfeni["state"], crfeni["validation"]
    rows = {r["subsystem"]: r for r in V["rows"]}
    assert rows["Cr-Ni"]["verdict"] == "reproduced" and rows["Cr-Ni"]["basis"].startswith("identical description")
    assert rows["Fe-Ni"]["verdict"] == "reproduced", rows["Fe-Ni"]["basis"]
    assert "step1:rule:typedef:HCP_A3" in rows["Fe-Ni"]["explained_by"]          # base's magnetic HCP model
    # Cr-Fe takes the base's (SGTE) fcc-Cr lattice stability instead of 82Her's pre-SGTE one: changed, explained
    cf = rows["Cr-Fe"]
    assert cf["verdict"] == "changed, explained" and cf["explained_by"][0] == "step2:rule:unaries:CR"
    assert cf["changes"][0].startswith("G(FCC_A1,CR:VA;0): Δ")
    assert V["extrapolations"]["ternaries"] == [["CR", "FE", "NI"]] and not V["extrapolations"]["binaries"]
    assert V["counts"]["CHANGED, UNEXPLAINED"] == 0
    log = {e["id"]: e for e in st["log"]}
    assert log["step1:reject:db1:L12_FCC+A1_FCC"]["who"] == "user"
    assert "G(FCC_A1,CR:VA;0)" in " ".join(log["step2:rule:unaries:CR"]["evidence"])
    for r in V["rows"]:
        if r["verdict"] == "reproduced":
            assert r["props_equal"]


def test_hand_edited_tdb_is_unexplained(crfeni, tmp_path):
    st = crfeni["state"]
    text = Path(st["final"]).read_text()
    edited = re.sub(r"(PARAMETER L\(LIQUID,CR,NI;0\)\s+\S+\s+)", r"\g<1>+5000", text, count=1)
    assert edited != text
    f2 = tmp_path / "edited.tdb"
    f2.write_text(edited)
    r = validate_from_state(st["files"]["state"], final_tdb=str(f2), out_dir=str(tmp_path / "v"), only=["CrNi-06Hal"],
                            lint=False, write_reports=False)
    row = r["validation"]["rows"][0]
    assert row["subsystem"] == "Cr-Ni" and row["verdict"] == "CHANGED, UNEXPLAINED"
    assert any(x.startswith("L(LIQUID,CR,NI;0)") and "edited after the run" in x for x in row["unexplained"])
    assert row["inv_max_dT"] > 0.5 and r["validation"]["edited"]
    assert "UNEXPLAINED Cr-Ni (CrNi-06Hal)" in r["text"]


# ------------------------------------------------------------------------------------------------- exactness
def test_gm_identity_of_normalised_phases(tmp_path):
    """A reordered C14 (B's (W)(Fe,W)2 vs A's Fe2Mo order): the run-time check records the error (exact), and it
    catches a real change."""
    from calphad_tools.staged_merge import staged_merge
    from test_compatibility import A_FeMo, B_FeMoW, write
    # B's C14 written W:Fe 1:2, (FE,W)1(FE,W)2: normalised to A's (FE)2(MO)1 order by sublattice size
    b_text = re.sub(r"(\w+)\(C14,([^:;]+):([^:;]+);", lambda m: f"{m.group(1)}(C14,{m.group(3)}:{m.group(2)};",
                    B_FeMoW().replace(" PHASE C14 % 2 2 1 !", " PHASE C14 % 2 1 2 !"))
    a, b = write(tmp_path, "A_FeMo.tdb", A_FeMo()), write(tmp_path, "B_FeMoW.tdb", b_text)
    r = staged_merge(a, b, str(tmp_path / "sm"), plot=False)
    assert r["written"] and any(n.action == "reordered" for n in r["notes"])
    r["k"] = 1
    log = DecisionLog()
    log.rule("step1:norm:C14:reordered", "step 1", "C14", "reorder", "normalisation")
    checks = exactness_checks([r], log)
    c = [x for x in checks if x["phase"] == "C14"][0]
    assert c["ok"] and c["max_rel"] < 1e-9 and c["n"] == 6 and "reordered" in c["transform"]
    assert any("GM identity" in x and "exact" in x for x in log.get("step1:norm:C14:reordered").evidence)
    bad = Path(r["B"]).read_text().replace("3*GHSERFE+5000", "3*GHSERFE+9000")
    assert bad != Path(r["B"]).read_text()
    (tmp_path / "bad.tdb").write_text(bad)
    res = gm_identity(r["B_original"], "C14", str(tmp_path / "bad.tdb"), "C14", [1, 0])
    assert math.isfinite(res["max_rel"]) and res["max_rel"] > 1e-6
