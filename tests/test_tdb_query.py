"""TDB inspection tools: query_tdb (every mode), resolve_tdb_symbol and validate_tdb_edit / validate_edit."""
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
TDBS = ROOT / "examples" / "tdbs"

# synthetic A-B system: a magnetic BCC, a 3-sublattice SIG (10:4:16), a LAVES phase and a liquid
BASE = """$ synthetic A-B system
 ELEMENT /-   ELECTRON_GAS  0.0 0.0 0.0 !
 ELEMENT VA   VACUUM        0.0 0.0 0.0 !
 ELEMENT A    BCC_A2       10.0 0.0 0.0 !
 ELEMENT B    BCC_A2       20.0 0.0 0.0 !
 TYPE_DEFINITION % SEQ * !
 TYPE_DEFINITION M GES A_P_D BCC_A2 MAGNETIC -1.0 0.4 !
 FUNCTION GHSERAA 298.15 -1000+10*T-2*T*LN(T); 3000 N !
 FUNCTION GHSERBB 298.15 -2000+12*T-3*T*LN(T); 3000 N !
 FUNCTION GLAVAB 298.15 -21000+4*T; 3000 N !
 PHASE LIQUID % 1 1 !
 CONSTITUENT LIQUID : A,B : !
 PHASE BCC_A2 %M 2 1 3 !
 CONSTITUENT BCC_A2 : A,B : VA : !
 PHASE SIG % 3 10 4 16 !
 CONSTITUENT SIG : A : B : A,B : !
 PHASE LAVES % 2 2 1 !
 CONSTITUENT LAVES : A,B : A,B : !
 PARAMETER G(LIQUID,A;0) 298.15 GHSERAA+5000-5*T; 3000 N !
 PARAMETER G(LIQUID,B;0) 298.15 GHSERBB+6000-4*T; 3000 N !
 PARAMETER L(LIQUID,A,B;0) 298.15 -10000+2*T; 3000 N !
 PARAMETER G(BCC_A2,A:VA;0) 298.15 GHSERAA; 3000 N !
 PARAMETER G(BCC_A2,B:VA;0) 298.15 GHSERBB; 3000 N !
 PARAMETER TC(BCC_A2,A:VA;0) 298.15 1000; 3000 N !
 PARAMETER BMAGN(BCC_A2,A:VA;0) 298.15 2.2; 3000 N !
 PARAMETER L(BCC_A2,A,B:VA;0) 298.15 -8000+T; 3000 N !
 PARAMETER L(BCC_A2,A,B:VA;1) 298.15 1500; 3000 N !
 PARAMETER G(SIG,A:B:A;0) 298.15 26*GHSERAA+4*GHSERBB-30000; 3000 N !
 PARAMETER G(SIG,A:B:B;0) 298.15 10*GHSERAA+20*GHSERBB-20000; 3000 N !
 PARAMETER L(SIG,A:B:A,B;0) 298.15 -15000+3*T; 3000 N !
 PARAMETER L(SIG,A:B:A,B;1) 298.15 2500; 3000 N !
 PARAMETER G(LAVES,A:A;0) 298.15 3*GHSERAA+5000; 3000 N !
 PARAMETER G(LAVES,B:B;0) 298.15 3*GHSERBB+5000; 3000 N !
 PARAMETER G(LAVES,A:B;0) 298.15 2*GHSERAA+GHSERBB+GLAVAB; 3000 N !
 PARAMETER G(LAVES,B:A;0) 298.15 2*GHSERBB+GHSERAA+20000; 3000 N !
"""


def _edit(text, *pairs):
    for a, b in pairs:
        assert a in text, a
        text = text.replace(a, b)
    return text


# SIG written as (A,B)8 (A)5 (B)2: before's sublattices 0,1,2 -> after's 1,2,0, every site ratio x0.5, G/L x0.5
PERMUTED = _edit(BASE, (" PHASE SIG % 3 10 4 16 !", " PHASE SIG % 3 8 5 2 !"),
                 (" CONSTITUENT SIG : A : B : A,B : !", " CONSTITUENT SIG : A,B : A : B : !"),
                 ("G(SIG,A:B:A;0) 298.15 26*GHSERAA+4*GHSERBB-30000;", "G(SIG,A:A:B;0) 298.15 0.5*(26*GHSERAA+4*GHSERBB-30000);"),
                 ("G(SIG,A:B:B;0) 298.15 10*GHSERAA+20*GHSERBB-20000;", "G(SIG,B:A:B;0) 298.15 0.5*(10*GHSERAA+20*GHSERBB-20000);"),
                 ("L(SIG,A:B:A,B;0) 298.15 -15000+3*T;", "L(SIG,A,B:A:B;0) 298.15 0.5*(-15000+3*T);"),
                 ("L(SIG,A:B:A,B;1) 298.15 2500;", "L(SIG,A,B:A:B;1) 298.15 1250;"))
# magnetic BCC_A2 rescaled to (A,B)2(VA)6: G/L x2, TC/BMAGN unchanged
MAGNETIC = _edit(BASE, (" PHASE BCC_A2 %M 2 1 3 !", " PHASE BCC_A2 %M 2 2 6 !"),
                 ("G(BCC_A2,A:VA;0) 298.15 GHSERAA;", "G(BCC_A2,A:VA;0) 298.15 2*GHSERAA;"),
                 ("G(BCC_A2,B:VA;0) 298.15 GHSERBB;", "G(BCC_A2,B:VA;0) 298.15 2*GHSERBB;"),
                 ("L(BCC_A2,A,B:VA;0) 298.15 -8000+T;", "L(BCC_A2,A,B:VA;0) 298.15 -16000+2*T;"),
                 ("L(BCC_A2,A,B:VA;1) 298.15 1500;", "L(BCC_A2,A,B:VA;1) 298.15 3000;"))

UNARY_MO = """ ELEMENT MO BCC_A2 95.94 4589 28.56 !
 ELEMENT VA VACUUM 0 0 0 !
 FUNCTION GHSERMO 298.15 -7746.302+131.9197*T-23.56414*T*LN(T)-.003443396*T**2+5.66283E-07*T**3+65812*T**(-1);
     2896 Y -30556.41+283.559746*T-42.63829*T*LN(T)-4.849315E+33*T**(-9); 5000 N !
 FUNCTION GFCCMO 298.15 GHSERMO+15200+0.63*T; 5000 N !
 FUNCTION GMOHCP 298.15 GHSERMO+11550; 5000 N !
 PHASE FCC_A1 % 2 1 1 !
 CONSTITUENT FCC_A1 : MO : VA : !
 PARAMETER G(FCC_A1,MO:VA;0) 298.15 GHSERMO+15200+0.63*T; 5000 N 91Din !
"""
USES_GMOFCC = """ ELEMENT MO BCC_A2 95.94 4589 28.56 !
 ELEMENT NI FCC_A1 58.69 4787 29.796 !
 ELEMENT VA VACUUM 0 0 0 !
 FUNCTION GHSERMO 298.15 -7746.302+131.9197*T-23.56414*T*LN(T)-.003443396*T**2+5.66283E-07*T**3+65812*T**(-1);
     2896 Y -30556.41+283.559746*T-42.63829*T*LN(T)-4.849315E+33*T**(-9); 5000 N !
 FUNCTION GHSERNI 298.15 -5179.159+117.854*T-22.096*T*LN(T)-.0048407*T**2; 3000 N !
 PHASE FCC_A1 % 2 1 1 !
 CONSTITUENT FCC_A1 : MO,NI : VA : !
 PARAMETER G(FCC_A1,NI:VA;0) 298.15 GHSERNI; 3000 N !
 PARAMETER G(FCC_A1,MO:VA;0) 298.15 GHSERMO+GMOFCC; 5000 N !
 PARAMETER L(FCC_A1,MO,NI:VA;0) 298.15 -8000; 3000 N !
"""


@pytest.fixture(scope="module")
def tb(tmp_path_factory):
    from calphad_tools.agent.tools import UI, ToolBox

    class NoHumanUI(UI):
        def ask(self, question, options=None):
            return "no"

        def confirm(self, title, body):
            return False

    out = tmp_path_factory.mktemp("tdbq")
    return ToolBox(workdir=str(ROOT / "examples"), output_dir=str(out / "out"), ui=NoHumanUI(), python_tool="off")


@pytest.fixture()
def synth(tmp_path):
    files = {}
    for name, text in (("base", BASE), ("same", BASE), ("permuted", PERMUTED), ("magnetic", MAGNETIC)):
        files[name] = tmp_path / f"{name}.tdb"
        files[name].write_text(text)
    return files


def _call(tb, name, args):
    text, images, err = tb.call(name, args)
    assert not err, text
    assert len(text) <= 3600, len(text)
    return text


# ------------------------------------------------------------------------------------------------ registration
def test_tools_registered_and_in_tool_help(tb):
    for n in ("query_tdb", "resolve_tdb_symbol", "validate_tdb_edit"):
        assert n in tb.tools
    text = _call(tb, "tool_help", {"names": ["query_tdb", "resolve_tdb_symbol", "validate_tdb_edit"]})
    for n in ("## query_tdb", "## resolve_tdb_symbol", "## validate_tdb_edit", "missing_endmembers", "search_paths",
              "affected_phases", "sublattice_map"):
        assert n in text
    compact = {t.name: t for t in tb.model_specs("compact")}
    assert compact["query_tdb"].description.endswith("[tool_help]")


# ------------------------------------------------------------------------------------------------ query_tdb
def test_phase_mode(tb):
    text = _call(tb, "query_tdb", {"tdb": "CoCrFeNiV_Choi2019.tdb", "mode": "phase", "phases": ["FCC_A1", "SIGMA"]})
    assert "FCC_A1 (CO,CR,FE,NI,V)1(VA)1" in text and "1 atoms/fu" in text
    assert "magnetic AFM -3, p 0.28" in text
    assert "parameters: G 5, L0 14, L1 7, L2 2, TC 12, BMAGN 13" in text
    assert "SIGMA (CO,FE,NI)8(CR,V)4(CO,CR,FE,NI,V)18" in text and "30 atoms/fu" in text


def test_phase_mode_ordered_and_expanded(tb, tmp_path):
    text = _call(tb, "query_tdb", {"tdb": "AlCoCrNi_Liu2016.tdb", "phases": ["BCC_B2", "BCC_A2"]})
    assert "DIS_PART BCC_A2" in text and "disordered part of BCC_B2" in text
    assert "VA mixes on sublattice" in text
    out = tmp_path / "febcc.xlsx"
    text = _call(tb, "query_tdb", {"tdb": "FeW-15Jac-static.tdb", "phases": ["BCC_A2"], "expand_functions": True,
                                   "output": str(out)})
    assert "-23.5143*T*ln(T)" in text and "L(BCC_A2,FE,W:VA;1)" in text
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert {"models", "parameters"} <= set(wb.sheetnames)
    assert wb["parameters"].max_row == 8                     # header + 7 parameters


def test_element_mode(tb):
    text = _call(tb, "query_tdb", {"tdb": "AlCoCrNi_Liu2016.tdb", "mode": "element", "elements": ["CO"],
                                   "max_rows": 5})
    assert "parameter(s) involving Co" in text and "by subsystem:" in text and "Al-Co-Cr-Ni" in text
    assert "... " in text                                    # truncated, rest in the table
    text = _call(tb, "query_tdb", {"tdb": "FeW-15Jac-static.tdb", "mode": "element", "elements": ["FE", "W"],
                                   "phases": ["MU_D85"]})
    assert "7 parameter(s) involving Fe+W in 1 phase(s)" in text          # G(MU_D85,W:W:W:W) has no Fe


def test_collisions_feb_few(tb, tmp_path):
    out = tmp_path / "coll.csv"
    text = _call(tb, "query_tdb", {"tdb": ["FeNb-90Hua-LB.tdb", "FeW-15Jac-static.tdb"], "mode": "collisions",
                                   "output": str(out)})
    assert "8 parameter key(s) in more than one: 8 identical, 0 numerically equal, 0 differ" in text
    assert "C14_LAVES: (FE)2(NB)1" in text and "(FE,W)2(FE,W)1" in text
    assert "MU_D85" in text and "different: 3 vs 4 non-vacancy sublattices" in text
    assert "G(BCC_A2,FE:VA;0) in FeNb-90Hua-LB,FeW-15Jac-static: identical" in text
    assert "functions with the same name (6): 0 differ" in text
    assert (tmp_path / "coll_collisions.csv").exists() and (tmp_path / "coll_phases.csv").exists()


def test_collisions_numeric_equality_and_difference(tb, tmp_path):
    a = tmp_path / "a.tdb"
    b = tmp_path / "b.tdb"
    a.write_text(BASE)
    b.write_text(_edit(BASE, ("L(LIQUID,A,B;0) 298.15 -10000+2*T;", "L(LIQUID,B,A;0) 298.15 2*(T-5000);"),
                       ("G(LAVES,B:A;0) 298.15 2*GHSERBB+GHSERAA+20000;", "G(LAVES,B:A;0) 298.15 2*GHSERBB+GHSERAA+20500;"),
                       ("G(LAVES,A:B;0) 298.15 2*GHSERAA+GHSERBB+GLAVAB;", "G(LAVES,A:B;0) 298.15 2*GHSERAA+GHSERBB-21000+4*T;"),
                       ("FUNCTION GHSERBB 298.15 -2000+12*T-3*T*LN(T);", "FUNCTION GHSERBB 298.15 12*T-3*T*LN(T)-2000.0;")))
    text = _call(tb, "query_tdb", {"tdb": [str(a), str(b)], "mode": "collisions", "max_rows": 30})
    assert "17 parameter key(s) in more than one: 14 identical, 2 numerically equal, 1 differ" in text
    assert "G(LAVES,B:A;0) in a,b: differs, max |d| 500" in text
    assert "L(LIQUID,A,B;0) in a,b: numerically equal" in text
    assert "L(LAVES,A:B;0)" not in text and "G(LAVES,A:B;0) in a,b: numerically equal" in text   # GLAVAB inlined
    assert "functions with the same name (3): 0 differ" in text            # GHSERBB: same terms, other order


def test_functions_mode(tb, tmp_path):
    text = _call(tb, "query_tdb", {"tdb": "CoCrFeNiV_Choi2019.tdb", "mode": "functions", "max_rows": 4})
    assert "92 function(s); 0 undefined reference(s); 2 unused" in text and "GCRHCP, GNIHCP" in text
    bad = tmp_path / "uses.tdb"
    bad.write_text(USES_GMOFCC)
    text = _call(tb, "query_tdb", {"tdb": str(bad), "mode": "functions"})
    assert "1 undefined reference(s)" in text and "UNDEFINED GMOFCC <- G(FCC_A1,MO:VA;0)" in text


def test_typedefs_mode(tb):
    text = _call(tb, "query_tdb", {"tdb": "AlCoCrNi_Liu2016.tdb", "mode": "typedefs"})
    assert "BCC_B2: magnetic AFM -1, p 0.4 (type B); DIS_PART BCC_A2" in text
    assert "L12_FCC: magnetic AFM -3, p 0.28 (type A); DIS_PART FCC_A1" in text
    assert "0 with problems" in text


def test_formula_unit_mode(tb):
    text = _call(tb, "query_tdb", {"tdb": ["FeNb-90Hua-LB.tdb", "FeW-15Jac-static.tdb", "CeCo-24Liu.tdb"],
                                   "mode": "formula_unit", "phases": ["C14_LAVES", "MU_D85", "BCC_A2"]})
    assert "C14_LAVES [FeNb-90Hua-LB.tdb] (FE)2(NB)1: 3 atoms/fu" in text
    assert "BCC_A2 [CeCo-24Liu.tdb] (CE,CO)1(VA)3: 1 atoms/fu (=1)" in text
    assert "MU_D85: FeW-15Jac-static.tdb vs FeNb-90Hua-LB.tdb: different" in text


def test_formula_unit_proportional(tb, synth):
    text = _call(tb, "query_tdb", {"tdb": [str(synth["base"]), str(synth["permuted"])], "mode": "formula_unit",
                                   "phases": ["SIG"]})
    assert "permuted+proportional, k = 0.5" in text and "1,2,0" in text


def test_missing_endmembers_mode(tb):
    text = _call(tb, "query_tdb", {"tdb": "FeW-15Jac-static.tdb", "mode": "missing_endmembers",
                                   "phases": ["MU_D85"]})
    assert "8 end-member(s), 8 with G, 0 missing" in text
    text = _call(tb, "query_tdb", {"tdb": "FeNb-90Hua-LB.tdb", "mode": "missing_endmembers", "phases": ["MU_D85"],
                                   "add_constituents": {"0": ["NB"]}})
    assert "(+NB@0)" in text and "4 end-member(s), 2 with G, 2 missing" in text
    assert "missing G(MU_D85,NB:NB:FE;0)  (uses added constituent)" in text


def test_bad_phase_name_suggests(tb):
    text, _, err = tb.call("query_tdb", {"tdb": "FeW-15Jac-static.tdb", "phases": ["MU"]})
    assert err and "MU_D85" in text


# ------------------------------------------------------------------------------------------------ resolve_tdb_symbol
def test_resolve_defined_function(tb):
    text = _call(tb, "resolve_tdb_symbol", {"name": "gfccfe", "search_paths": ["tdbs"]})
    assert "GFCCFE: defined in 5 place(s)" in text
    assert "DEFINED tdbs/FeW-15Jac-static.tdb:84 [298.15-1811-6000]" in text
    assert "used by G(C14_LAVES,FE:NB) at tdbs/FeNb-90Hua-LB.tdb:132" in text


def test_resolve_undefined_suggests_near_misses(tb, tmp_path):
    d = tmp_path / "syms"
    d.mkdir()
    (d / "uses_gmofcc.tdb").write_text(USES_GMOFCC)
    (d / "unary_mo.TDB").write_text(UNARY_MO)                 # upper-case extension is found too
    text = _call(tb, "resolve_tdb_symbol", {"name": "GMOFCC", "search_paths": [str(d), "tdbs/*.tdb"]})
    assert "GMOFCC: defined in 0 place(s), referenced 1 time(s)" in text
    assert "(NOT defined in that file)" in text
    assert "reads as element MO, phase tag FCC" in text and "lattice stability (added to GHSER)" in text
    for near in ("GMOHCP (same element MO", "GFCCMO (same element MO", "GHSERMO (same element MO",
                 "GCOFCC (edit distance 1"):
        assert near in text, near
    assert "FUNCTION GFCCMO = GHSERMO+15200+0.63*T" in text
    assert "- GHSERMO = 15200 +0.63*T" in text
    assert "FUNCTION GMOFCC = G(FCC_A1,MO) - GHSERMO" in text


def test_resolve_dialect_variant(tmp_path):
    from calphad_tools.tdb_query import dialect_variant, parse_symbol
    assert dialect_variant("GHSERMO_LONGNAME", "GHSERMO1") and "8-character" in dialect_variant("GHSERMOLONG",
                                                                                                "GHSERMO1")
    assert dialect_variant("GHSERMO", "GHSERFE") is None
    assert parse_symbol("GHSERWW") == ("W", "SER", "GHSER{EL}")
    assert parse_symbol("GFCCMO")[:2] == ("MO", "FCC") and parse_symbol("GMOFCC")[:2] == ("MO", "FCC")
    assert parse_symbol("GCOFCC")[:2] == ("CO", "FCC") and parse_symbol("GLIQFE")[:2] == ("FE", "LIQ")


# ------------------------------------------------------------------------------------------------ validate_tdb_edit
def test_validate_identical_copy(tb, synth):
    text = _call(tb, "validate_tdb_edit", {"before": str(synth["base"]), "after": str(synth["same"]),
                                           "expect": ["only affected phases changed", "base subsystems unchanged"]})
    assert text.startswith("PASS: base.tdb -> same.tdb; changed phases: none")
    for c in ("pycalphad loads after", "parse round trip", "constituents", "functions", "type definitions",
              "duplicates", "site ratios", "unaffected phases identical"):
        assert f"PASS  {c}" in text, c
    assert "FAIL" not in text


def test_validate_permuted_rescaled_phase(synth, tmp_path):
    from calphad_tools.tdb_validate import validate_edit
    r = validate_edit(str(synth["base"]), str(synth["permuted"]), affected_phases=["SIG"],
                      expect=["only affected phases changed", "energy_per_atom_preserved"], n_samples=12,
                      output=str(tmp_path / "v.xlsx"))
    assert r["verdict"] == "PASS", r["text"]
    assert r["changed_phases"] == ["SIG"]
    g = r["gm"][0]
    assert g["status"] == "pass" and g["perm"] == "1,2,0" and g["k"] == pytest.approx(0.5)
    assert g["max_rel"] <= 1e-9 and g["n"] == 12
    assert "permuted, rescaled x0.5: GM per atom preserved" in r["text"]
    assert Path(r["files"][0]).exists()
    # explicit hints give the same answer; a wrong scale hint finds no mapping
    r2 = validate_edit(str(synth["base"]), str(synth["permuted"]), sublattice_map={"SIG": [1, 2, 0]},
                       scale={"SIG": 0.5}, expect=["energy_per_atom_preserved"], n_samples=4)
    assert r2["verdict"] == "PASS" and r2["gm"][0]["status"] == "pass"
    r3 = validate_edit(str(synth["base"]), str(synth["permuted"]), scale={"SIG": 2.0},
                       expect=["energy_per_atom_preserved"], n_samples=4)
    assert r3["verdict"] == "FAIL" and "no sublattice mapping" in r3["gm"][0]["detail"]


def test_validate_reversed_odd_order_term(synth, tmp_path):
    """L(SIG,B,A:...;1) = -1250 equals L(SIG,A,B:...;1) = 1250 in Thermo-Calc, but pycalphad sorts the constituents
    without flipping the sign: the GM check (pycalphad) sees the change and the order check says why."""
    from calphad_tools.tdb_validate import validate_edit
    p = tmp_path / "reversed.tdb"
    p.write_text(_edit(PERMUTED, ("L(SIG,A,B:A:B;1) 298.15 1250;", "L(SIG,B,A:A:B;1) 298.15 -1250;")))
    r = validate_edit(str(synth["base"]), str(p), affected_phases=["SIG"], expect=["energy_per_atom_preserved"],
                      n_samples=6)
    chk = {c["check"]: c for c in r["checks"]}
    assert chk["constituent order (odd RK terms)"]["status"] == "warn"
    assert "L(SIG,B,A:A:B;1)" in chk["constituent order (odd RK terms)"]["detail"]
    assert r["verdict"] == "FAIL" and r["gm"][0]["status"] == "fail"
    assert "pycalphad sorts them without flipping the sign" in r["gm"][0]["detail"]


def test_validate_flags_change_outside_affected(synth, tmp_path):
    from calphad_tools.tdb_validate import validate_edit
    p = tmp_path / "perm_liq.tdb"
    p.write_text(_edit(PERMUTED, ("L(LIQUID,A,B;0) 298.15 -10000+2*T;", "L(LIQUID,A,B;0) 298.15 -10500+2*T;")))
    r = validate_edit(str(synth["base"]), str(p), affected_phases=["SIG"])
    assert r["verdict"] == "FAIL"
    chk = {c["check"]: c for c in r["checks"]}
    assert chk["unaffected phases identical"]["status"] == "fail"
    assert "LIQUID: 1 changed (G(LIQUID,A,B;0), max |d| 500)" in chk["unaffected phases identical"]["detail"]
    assert chk["GM per atom (affected phases)"]["status"] == "pass"
    # an added parameter outside the affected phases: allowed unless "only affected phases changed"
    q = tmp_path / "added.tdb"
    q.write_text(BASE + " PARAMETER L(LAVES,A,B:A;0) 298.15 -1000; 3000 N !\n")
    r = validate_edit(str(synth["base"]), str(q), affected_phases=["SIG"])
    assert r["verdict"] == "PASS" and "additions: LAVES: 1 added" in r["text"]
    r = validate_edit(str(synth["base"]), str(q), affected_phases=["SIG"], expect=["only affected phases changed"])
    assert r["verdict"] == "FAIL"


def test_validate_flags_undefined_function_and_bad_constituent(tb, synth, tmp_path):
    p = tmp_path / "broken.tdb"
    p.write_text(_edit(BASE, ("G(LAVES,B:A;0) 298.15 2*GHSERBB+GHSERAA+20000;",
                              "G(LAVES,B:A;0) 298.15 2*GHSERBB+GHSERAA+GMOFCC;"),
                       ("G(SIG,A:B:B;0)", "G(SIG,B:B:B;0)")))
    text = _call(tb, "validate_tdb_edit", {"before": str(synth["base"]), "after": str(p)})
    assert text.startswith("FAIL")
    assert "FAIL  functions:" in text and "undefined function GMOFCC referenced by G(LAVES,B:A;0)" in text
    assert "FAIL  constituents:" in text and "G(SIG,B:B:B;0): B is not a constituent of sublattice 1 of SIG" in text


def test_validate_duplicate_and_site_ratio(synth, tmp_path):
    from calphad_tools.tdb_validate import validate_edit
    p = tmp_path / "dup.tdb"
    p.write_text(_edit(BASE, (" PHASE LIQUID % 1 1 !", " PHASE LIQUID % 1 0.9 !"))
                 + " PARAMETER G(LIQUID,A;0) 298.15 GHSERAA+5000-5*T; 3000 N !\n")
    r = validate_edit(str(synth["base"]), str(p))
    chk = {c["check"]: c for c in r["checks"]}
    assert chk["duplicates"]["status"] == "fail" and "defined 2 times" in chk["duplicates"]["detail"]
    assert chk["site ratios"]["status"] == "fail" and "LIQUID (A,B)0.9 sums to 0.9" in chk["site ratios"]["detail"]


def test_validate_magnetic_rescale_reported_separately(synth):
    from calphad_tools.tdb_validate import validate_edit
    r = validate_edit(str(synth["base"]), str(synth["magnetic"]), expect=["energy_per_atom_preserved"],
                      temperatures=[400, 800, 1200], n_samples=6)
    g = {x["phase"]: x for x in r["gm"]}["BCC_A2"]
    assert g["status"] == "warn" and g["k"] == pytest.approx(2.0)
    assert g["nonmag_rel"] <= 1e-9 and g["mag_abs"] > 1.0
    assert "magnetic contribution changed by" in g["detail"] and "non-magnetic part preserved" in g["detail"]
    chk = {c["check"]: c for c in r["checks"]}
    assert chk["GM per atom (affected phases)"]["status"] == "warn"
    assert chk["site ratios"]["status"] == "fail"           # 1 atom/fu before, 2 after
