"""show_code for the figure tools: the generated standalone code must rebuild the same figure DATA.

Every test makes figures with the real tools, asks show_code for the standalone script, runs that script in a
subprocess (no calphad_tools import allowed), dumps the artists of its final figure and compares them with the
same dump of the tool's own stored figure.
"""
import json
import re
import os
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable

# ---------------------------------------------------------------------------------------------------------
# artist dump (run both in this process on the tool's figure and appended to the generated script)
# ---------------------------------------------------------------------------------------------------------
DUMP_SRC = r'''
def _dump_fig(fig):
    import numpy as _np

    def arr(a):
        a = _np.ma.filled(_np.ma.asarray(a, dtype=float), _np.nan)
        return [None if not _np.isfinite(v) else float(v) for v in _np.ravel(a)]

    top = [a for a in fig.axes if a.get_label() not in ("<colorbar>", "background image")]
    kids = [c for a in top for c in a.child_axes
            if type(c).__name__ != "SecondaryAxis" and c.get_label() != "<colorbar>" and c not in top]
    out = {"suptitle": fig._suptitle.get_text() if getattr(fig, "_suptitle", None) else "", "axes": [],
           "n_colorbars": sum(a.get_label() == "<colorbar>" for a in fig.axes)}
    for ax in top + kids:
        d = {"name": ax.name, "title": ax.get_title(), "xlabel": ax.get_xlabel(), "ylabel": ax.get_ylabel(),
             "xscale": ax.get_xscale(), "yscale": ax.get_yscale(),
             "xlim": arr(ax.get_xlim()), "ylim": arr(ax.get_ylim())}
        lines = []
        for ln in ax.get_lines():
            if hasattr(ln, "get_data_3d"):
                xs, ys, zs = ln.get_data_3d()
                lines.append({"x": arr(xs), "y": arr(ys), "z": arr(zs), "label": ln.get_label()})
            else:
                lines.append({"x": arr(ln.get_xdata()), "y": arr(ln.get_ydata()), "label": ln.get_label()})
        d["lines"] = lines
        cols = []
        for c in ax.collections:
            e = {"type": type(c).__name__, "label": c.get_label()}
            try:
                e["offsets"] = arr(c.get_offsets())
            except Exception:
                pass
            a = c.get_array() if hasattr(c, "get_array") else None
            if a is not None:
                e["array"] = arr(a)
            if hasattr(c, "levels"):
                e["levels"] = arr(c.levels)
            try:
                paths = c.get_paths()
                e["npaths"] = len(paths)
                e["verts"] = float(sum(float(_np.nansum(_np.abs(p.vertices))) for p in paths)) if len(paths) < 5000 \
                    else None
            except Exception:
                pass
            cols.append(e)
        d["collections"] = cols
        d["patches"] = [type(p).__name__ for p in ax.patches]
        d["patch_xy"] = [arr(p.get_xy()) if hasattr(p, "get_xy") and not callable(getattr(p, "get_path", None)) is None
                         else arr(p.get_path().vertices.mean(axis=0)) if len(p.get_path().vertices) else []
                         for p in ax.patches]
        d["texts"] = [t.get_text() for t in ax.texts]
        d["images"] = [[list(im.get_array().shape), float(_np.nansum(_np.ma.filled(
            _np.ma.asarray(im.get_array(), dtype=float), 0.0)))] for im in ax.images]
        leg = ax.get_legend()
        d["legend"] = [t.get_text() for t in leg.get_texts()] if leg is not None and leg.get_visible() else None
        d["secondary"] = [(s.get_xlabel(), s.get_ylabel()) for s in ax.child_axes
                          if type(s).__name__ == "SecondaryAxis"]
        out["axes"].append(d)
    return out
'''
ns: dict = {}
exec(DUMP_SRC, ns)
dump_fig = ns["_dump_fig"]


def compare(a, b, path="fig", rtol=1e-6, atol=1e-9):
    """Structural comparison: numbers approximately, everything else exactly."""
    if isinstance(a, dict):
        assert isinstance(b, dict) and set(a) == set(b), f"{path}: keys {sorted(a)} != {sorted(b)}"
        for k in a:
            compare(a[k], b[k], f"{path}.{k}", rtol, atol)
    elif isinstance(a, (list, tuple)):
        assert isinstance(b, (list, tuple)) and len(a) == len(b), \
            f"{path}: length {len(a)} != {len(b) if isinstance(b, (list, tuple)) else b!r}"
        if a and all(isinstance(v, (int, float)) or v is None for v in a) and \
                all(isinstance(v, (int, float)) or v is None for v in b):
            x = np.array([np.nan if v is None else v for v in a], dtype=float)
            y = np.array([np.nan if v is None else v for v in b], dtype=float)
            assert np.allclose(x, y, rtol=rtol, atol=atol, equal_nan=True), \
                f"{path}: max diff {np.nanmax(np.abs(x - y))} ({x[:6]} vs {y[:6]})"
            return
        for i, (u, v) in enumerate(zip(a, b)):
            compare(u, v, f"{path}[{i}]", rtol, atol)
    elif isinstance(a, float) or isinstance(b, float):
        assert a is not None and b is not None and np.isclose(a, b, rtol=rtol, atol=atol, equal_nan=True), \
            f"{path}: {a} != {b}"
    else:
        assert a == b, f"{path}: {a!r} != {b!r}"


# ---------------------------------------------------------------------------------------------------------
# fixtures / helpers
# ---------------------------------------------------------------------------------------------------------
@pytest.fixture
def tb(tmp_path):
    from calphad_tools.agent.tools import ToolBox, UI

    class U(UI):
        def ask(self, q, o=None):
            return ""

        def confirm(self, t, b):
            return False
    work = tmp_path / "work"
    work.mkdir()
    _write_data(work)
    return ToolBox(workdir=str(work), output_dir=str(work / "out"), ui=U(), python_tool="off")


def _write_data(work: Path):
    rng = np.random.default_rng(3)
    T = np.linspace(300, 1200, 40)
    pd.DataFrame({"T [K]": T, "cp": 20 + 0.01 * T + np.sin(T / 50),
                  "phase": np.where(T < 800, "FCC_A1", "BCC_A2"),
                  "x(Cr)": np.linspace(0.05, 0.45, 40)}).to_csv(work / "cp.csv", index=False)
    Tm = np.repeat([600.0, 700, 800, 900, 1000], 4)
    D = 1e-4 * np.exp(-150e3 / (8.314462618 * Tm)) * np.exp(rng.normal(0, 0.1, len(Tm)))
    pd.DataFrame({"T": Tm, "D": D, "alloy": np.tile(["A", "B", "C", "D"], 5),
                  "hv": 200 + 0.1 * Tm + rng.normal(0, 5, len(Tm))}).to_csv(work / "diff.csv", index=False)
    # semicolon separated, decimal comma, comment header (delimiter sniffing)
    (work / "euro.csv").write_text("# measured\n# by someone\nT;val\n300;1,5\n400;2,5\n500;3,0\n")
    # long-format grid (heatmap / contour / stackplot)
    rows = [{"T": t, "x": x, "G": np.sin(t / 300) + x ** 2, "phase": ph, "NP": frac}
            for t in (600, 700, 800, 900) for x in (0.1, 0.2, 0.3)
            for ph, frac in (("FCC", x), ("BCC", 1 - x))]
    pd.DataFrame(rows).to_csv(work / "grid.csv", index=False)
    pd.DataFrame({"t": [0.0, 1, 2, 3, 4], "ref": [10.0, 11, 13, 16, 20]}).to_csv(work / "ref.csv", index=False)
    pd.DataFrame({"x": rng.uniform(0, 1, 60), "y": rng.uniform(0, 1, 60)}).assign(
        z=lambda d: d.x * (1 - d.y), s=lambda d: d.x + d.y).to_json(work / "pts.json", orient="records")


def call(tb, name, args):
    text, imgs, err = tb.call(name, args)
    assert not err, text
    return text, imgs


def standalone(tb, target):
    """show_code script for ``target``, after checking every cell is standalone (no toolkit import)."""
    from calphad_tools import codegen
    cells, ctx = codegen.build_cells(tb, str(target))
    text, _, err = tb.call("show_code", {"target": str(target)})
    assert not err, text
    titles = [c.title for c in cells]
    assert len(set(titles)) == len(titles), titles
    bad = [c.title + ": " + c.note for c in cells if not c.standalone]
    assert not bad, bad
    script = codegen.to_script(tb, cells, ctx)
    assert "calphad_tools" not in script
    assert script in text                 # what show_code returned
    return script, cells


def run_script(script: str, tmp_path: Path, dump: bool = True) -> dict:
    d = tmp_path / "run"
    d.mkdir(parents=True, exist_ok=True)
    code = script + ("\n\n" + DUMP_SRC + "\nimport json as _json\n_json.dump(_dump_fig(fig), open('dump.json', 'w'))\n"
                     if dump else "")
    (d / "script.py").write_text(code)
    env = dict(os.environ, MPLBACKEND="Agg", PYTHONPATH="")
    r = subprocess.run([PY, "script.py"], cwd=d, capture_output=True, text=True, env=env, timeout=240)
    assert r.returncode == 0, f"generated script failed:\n{r.stderr[-4000:]}\n---\n{code[-6000:]}"
    return json.loads((d / "dump.json").read_text()) if dump else {}


def tool_dump(png) -> dict:
    from calphad_tools.figures import load_figure
    fig = load_figure(str(png))
    assert fig is not None
    return dump_fig(fig)


def check(tb, png, tmp_path, n_cells=None):
    script, cells = standalone(tb, png)
    if n_cells is not None:
        assert len(cells) == n_cells, [c.title for c in cells]
    got = run_script(script, tmp_path)
    compare(tool_dump(png), got)
    return script, got


# ---------------------------------------------------------------------------------------------------------
# tests
# ---------------------------------------------------------------------------------------------------------
def test_make_plot_builtin_ops_preset_palette(tb, tmp_path):
    _, imgs = call(tb, "make_plot", {
        "output": "basic.png", "preset": "journal-double", "palette": "okabe-ito", "layout": [2, 2],
        "suptitle": "Overview", "font": {"size": 9}, "rc": {"axes.spines.top": False},
        "panels": [
            {"operations": [
                {"op": "line", "data": "cp.csv", "x_col": "T [", "y_col": "cp", "label": "cp",
                 "where": {"T [K]": {">": 400, "<=": 1100}, "phase": ["FCC_A1", "BCC_A2"]},
                 "transform": [{"step": "derive", "new": "T_C", "expr": "`T [K]` - 273.15"}]},
                {"op": "scatter", "data": "cp.csv", "x_col": "T [", "y_col": "cp", "c_col": "x(Cr)",
                 "cmap": "plasma", "colorbar": {"label": "x(Cr)"}, "s": 12},
                {"op": "errorbar", "x": [500, 700], "y": [25, 27], "yerr": [0.5, 0.8], "color": "k",
                 "label": "exp"},
                {"op": "fill_between", "x": [400, 600, 800], "y1": [20, 21, 22], "y2": [22, 23, 24],
                 "alpha": 0.3, "color": "C2"},
                {"op": "text", "x": 900, "y": 22, "text": "note", "bbox": True},
                {"op": "annotate", "text": "peak", "xy": [700, 27], "xytext": [900, 29]},
                {"op": "arrow", "x": 400, "y": 30, "dx": 100, "dy": -2},
                {"op": "hline", "y": 24, "linestyle": "--", "color": "0.5", "label": "mean"},
                {"op": "vspan", "x1": 800, "x2": 850, "color": "y", "alpha": 0.2},
                {"op": "polygon", "points": [[300, 20], [350, 21], [320, 22]], "facecolor": "b"},
                {"op": "set", "xlabel": "T / K", "ylabel": "cp", "title": "heat capacity", "grid": True,
                 "xlim": [250, 1250]},
                {"op": "legend", "loc": "lower right", "fontsize": 7}]},
            {"operations": [
                {"op": "step", "x": [0, 1, 2, 3], "y": [1, 3, 2, 4], "label": "steps"},
                {"op": "bar", "x": [0.5, 1.5], "height": [2, 1], "width": 0.4, "color": "C3", "label": "bars"},
                {"op": "circle", "x": 2, "y": 2, "radius": 0.3, "facecolor": "none", "edgecolor": "k"},
                {"op": "vline", "x": 2.5, "color": "k"},
                {"op": "hspan", "y1": 0.5, "y2": 1.0, "alpha": 0.1},
                {"op": "set", "yscale": "log", "xticks": [0, 1, 2, 3], "xticklabels": ["a", "b", "c", "d"]}]},
            {"operations": [
                {"op": "barh", "y": [1, 2, 3], "width": [3, 1, 2], "label": "h"},
                {"op": "line", "data": "euro.csv", "x_col": "val", "y_col": "T", "marker": "o"},
                {"op": "twinx", "ylabel": "right", "ylim": [0, 5]},
                {"op": "line", "axes": 4, "x": [0, 1, 2, 3], "y": [4, 3, 2, 1], "color": "m", "label": "on twin"}]},
            {"operations": [
                {"op": "scatter", "data": "pts.json", "x_col": "x", "y_col": "y", "c_col": "z",
                 "norm": {"type": "log", "vmin": 0.01, "vmax": 1}},
                {"op": "inset", "bounds": [0.6, 0.6, 0.35, 0.35], "xlim": [0.2, 0.4], "ylim": [0.2, 0.4]},
                {"op": "figure", "suptitle": "Overview (edited)"}]}]})
    script, got = check(tb, imgs[0], tmp_path)
    assert "sep=';'" in script and "decimal=','" in script and "skiprows=2" in script
    assert "cycler(color=['#E69F00'" in script and "'font.size': 9" in script
    assert got["axes"][0]["legend"] == ["cp", "mean", "exp"]


PIPELINES = {
    "derive_filter": ("cp.csv", None, [
        {"step": "derive", "exprs": {"T_C": "`T [K]` - 273.15", "big": "cp > mean(cp) and phase == 'FCC_A1'",
                                     "cls": "'hot' if `T [K]` > 800 else 'cold'", "r": "round(log10(cp), 3)",
                                     "m": "max(cp, 25)", "lo": "min(cp)", "fcc": "contains(phase, 'fcc')",
                                     "w": "where(cp > 25, 1, 0)", "isin": "phase in ['BCC_A2']"}},
        {"step": "filter", "expr": "T_C >= 100 and not big"},
        {"step": "select", "cols": ["T [", "cp", "cls", "r", "m", "lo", "fcc", "w", "isin"]},
        {"step": "rename", "map": {"cp": "Cp"}}, {"step": "sort", "by": "Cp", "ascending": False},
        {"step": "head", "n": 3, "by": "cls"}]),
    "group_pivot_melt": ("diff.csv", {"T": {"between": [600, 950]}}, [
        {"step": "groupby", "by": ["T"], "agg": {"D": ["mean", "std", "q25"], "hv": "max"}},
        {"step": "melt", "id_vars": ["T"], "value_vars": ["D_mean", "hv_max"]},
        {"step": "pivot", "index": "T", "columns": "variable", "values": "value", "prefix": "v_"}]),
    "join_concat": ("ref.csv", None, [
        {"step": "join", "with": {"data": "ref.csv", "transform": [{"step": "rename", "map": {"ref": "r2"}}]},
         "on": "t", "how": "left"},
        {"step": "join", "with": [{"t": 1.2, "near": 5.0}, {"t": 3.1, "near": 7.0}], "on": "t", "how": "nearest"},
        {"step": "concat", "with": ["ref.csv"], "label_col": "src"},
        {"step": "dedupe", "cols": ["t"]}, {"step": "dropna", "cols": ["ref"]}, {"step": "drop", "cols": ["src"]}]),
    "interp_smooth": ("cp.csv", None, [
        {"step": "interpolate", "x": "T [", "y": ["cp"], "onto": {"start": 400, "stop": 1000, "step": 50},
         "by": ["phase"], "method": "cubic"},
        {"step": "dropna", "cols": ["cp"]},
        {"step": "smooth", "cols": "cp", "method": "savgol", "window": 5, "by": ["phase"], "x": "T [",
         "new": "cp_s"},
        {"step": "smooth", "cols": "cp", "method": "rolling", "window": 3, "suffix": "_r"},
        {"step": "smooth", "cols": "cp", "method": "ewm", "window": 3, "suffix": "_e"},
        {"step": "derivative", "x": "T [", "y": "cp", "by": ["phase"]},
        {"step": "integrate", "x": "T [", "y": "cp", "by": "phase"},
        {"step": "normalize", "cols": ["cp"], "method": "zscore", "by": ["phase"], "suffix": "_z"},
        {"step": "normalize", "cols": ["cp"], "method": "minmax", "suffix": "_mm"}]),
    "bin_convert": ("cp.csv", None, [
        {"step": "convert", "col": "T [", "from": "K", "to": "°F", "new": "T_F"},
        {"step": "convert", "col": "cp", "from": "J/mol/K", "to": "cal/mol/K", "new": "cp_cal"},
        {"step": "convert", "cols": ["x(Cr)"], "from": "mole_fraction", "to": "wt%", "elements": ["CR"],
         "balance": "NI", "balance_col": "Ni wt%", "new": "Cr wt%"},
        {"step": "bin", "col": "T [", "bins": 4, "agg": {"cp": ["mean", "count"]}, "label": "interval"}]),
    "integrate_total": ("diff.csv", None, [
        {"step": "integrate", "x": "T", "y": "hv", "by": ["alloy"], "cumulative": False},
        {"step": "tail", "n": 2}]),
}


def test_transform_data_pipelines(tb, tmp_path):
    outs = {}
    for name, (src, where, steps) in PIPELINES.items():
        args = {"source": src, "transform": steps, "output": f"{name}.csv"}
        if where:
            args["where"] = where
        call(tb, "transform_data", args)
        outs[name] = Path(tb.out) / f"{name}.csv"
    for name, path in outs.items():
        script, _ = standalone(tb, path)
        d = tmp_path / f"run_{name}"
        d.mkdir()
        (d / "script.py").write_text(script)
        r = subprocess.run([PY, "script.py"], cwd=d, capture_output=True, text=True, timeout=120,
                           env=dict(os.environ, MPLBACKEND="Agg"))
        assert r.returncode == 0, f"{name}:\n{r.stderr[-3000:]}\n{script}"
        got = pd.read_csv(d / "standalone_outputs" / f"{name}.csv")
        want = pd.read_csv(path)
        pd.testing.assert_frame_equal(got, want, check_dtype=False, rtol=1e-9, obj=name)


def test_fits(tb, tmp_path):
    _, imgs = call(tb, "make_plot", {"output": "fits.png", "layout": [2, 3], "panels": [
        {"operations": [
            {"op": "scatter", "data": "diff.csv", "x_col": "T", "y_col": "D", "label": "data"},
            {"op": "fit", "data": "diff.csv", "x_col": "T", "y_col": "D", "model": "arrhenius", "band": True,
             "annotate": True, "label": True, "range": [550, 1050]},
            {"op": "set", "yscale": "log"}]},
        {"operations": [
            {"op": "fit", "x": [1, 2, 3, 4, 5, 6], "y": [1.1, 3.9, 9.2, 15.8, 25.3, 35.9], "model": "poly",
             "deg": 2, "band": "prediction", "show_data": True, "annotate": "equation", "color": "C3"},
            {"op": "residuals", "axes": 2, "x": [1, 2, 3, 4, 5, 6], "y": [1.1, 3.9, 9.2, 15.8, 25.3, 35.9],
             "model": "poly", "deg": 2, "standardized": True}]},
        {"operations": []},
        {"operations": [
            {"op": "fit", "x": [1, 2, 3, 4, 5], "y": [2.7, 7.4, 20.1, 54.6, 148.4], "sigma": [0.1, 0.3, 1, 2, 5],
             "model": "exp", "show_data": True, "annotate": "params", "label": "exp"},
            {"op": "fit", "x": [1, 2, 3, 4, 5], "y": [2.0, 5.7, 10.4, 16.0, 22.4], "model": "power",
             "linearize": True, "label": True}]},
        {"operations": [
            {"op": "fit", "x": [-2, -1, 0, 1, 2, 3], "y": [0.2, 1.3, 3.1, 2.4, 0.9, 0.1], "model": "gaussian",
             "offset": True, "npoints": 50, "label": True},
            {"op": "fit", "x": [0.5, 1, 2, 4, 8], "y": [1.2, 2.0, 2.9, 3.6, 4.5], "model": "log", "linestyle": ":"}]},
        {"operations": [
            {"op": "fit", "x": [0, 1, 2, 3, 4, 5], "y": [0.1, 0.4, 1.9, 4.2, 4.8, 5.0], "model": "logistic",
             "label": "logistic"},
            {"op": "fit", "x": [1, 2, 3, 4, 5], "y": [3.1, 4.9, 7.2, 8.8, 11.1],
             "model": {"expr": "a*x + b*sqrt(x)", "params": {"a": 1, "b": 1}}, "annotate": True,
             "bounds": {"a": [0, 10]}, "label": True}]}]})
    script, got = check(tb, imgs[0], tmp_path)
    assert "np.log(y)" in script           # Arrhenius fitted in log space by default
    assert "curve_fit(ln_model" in script


def test_statistics_ops(tb, tmp_path):
    _, imgs = call(tb, "make_plot", {"output": "stats.png", "layout": [3, 3], "panels": [
        {"operations": [{"op": "hist", "data": "diff.csv", "values_col": "hv", "group_col": "alloy", "bins": 5,
                         "histtype": "step", "density": True}]},
        {"operations": [{"op": "box", "data": "diff.csv", "values_col": "hv", "group_col": "alloy", "jitter": True,
                         "showmeans": True}]},
        {"operations": [{"op": "violin", "groups": {"a": [1, 2, 3, 4, 2.5], "b": [2, 3, 5, 4, 4.5]},
                         "jitter": 0.1, "colors": ["C1", "C2"]}]},
        {"operations": [{"op": "kde", "data": "diff.csv", "values_col": "hv", "group_col": "alloy",
                         "groups": ["A", "B"], "fill": True, "rug": True},
                        {"op": "ecdf", "axes": 3, "values": [[1, 2, 3], [2, 2.5, 4, 5]], "labels": ["u", "v"]}]},
        {"operations": [{"op": "qq", "data": "diff.csv", "values_col": "hv", "dist": "norm"}]},
        {"operations": [{"op": "hist2d", "data": "pts.json", "x_col": "x", "y_col": "y", "bins": 6,
                         "colorbar": True}]},
        {"operations": [{"op": "hexbin", "data": "pts.json", "x_col": "x", "y_col": "y", "c_col": "z",
                         "reduce": "max", "gridsize": 5, "colorbar": {"label": "z"}}]},
        {"operations": [{"op": "band", "data": "diff.csv", "x_col": "T", "y_col": "hv", "agg": "ci",
                         "group_col": "alloy", "groups": ["A", "B"]},
                        {"op": "band", "x": [1, 2, 3], "y": [1, 2, 1.5], "yerr": [0.2, 0.3, 0.1],
                         "label": "manual", "band_label": "±σ"}]},
        {"operations": [{"op": "corr_heatmap", "data": "cp.csv", "method": "spearman", "triangle": "lower"}]}]})
    check(tb, imgs[0], tmp_path)


def test_grid_and_draw_ops(tb, tmp_path):
    tri_pts = [(a, b) for a in np.linspace(0, 1, 9) for b in np.linspace(0, 1, 9) if a + b <= 1.0001]
    _, imgs = call(tb, "make_plot", {"output": "extra.png", "layout": [3, 3], "palette": "tol-bright", "panels": [
        {"operations": [{"op": "heatmap", "data": "grid.csv", "x_col": "x", "y_col": "T", "z_col": "G",
                         "where": {"phase": "FCC"}, "annotate": True, "fmt": "{x:.2f}", "colorbar": True}]},
        {"operations": [{"op": "contourf", "data": "grid.csv", "x_col": "x", "y_col": "T", "z_col": "G",
                         "where": {"phase": "FCC"}, "levels": 6, "lines": True, "cmap": "magma"},
                        {"op": "contour", "x": [0, 1, 2], "y": [0, 1], "z": [[0, 1, 2], [1, 2, 3]], "levels": [0.5, 1.5],
                         "colors": "k", "clabel_fmt": "{x:.1f}"}]},
        {"projection": "triangular", "operations": [
            {"op": "tricontourf", "x": [p[0] for p in tri_pts], "y": [p[1] for p in tri_pts],
             "z": [p[0] * p[1] for p in tri_pts], "levels": 5, "colorbar": True},
            {"op": "tricontour", "x": [p[0] for p in tri_pts], "y": [p[1] for p in tri_pts],
             "z": [p[0] - p[1] for p in tri_pts], "levels": 4, "colors": "k", "clabel": False},
            {"op": "scatter", "x": [0.2, 0.3], "y": [0.3, 0.2], "color": "k", "label": "alloys"}]},
        {"operations": [{"op": "stackplot", "data": "grid.csv", "x_col": "T", "group_col": "phase",
                         "value_col": "NP", "where": {"x": 0.2}, "normalize": True,
                         "colors": {"FCC": "tab:blue", "BCC": "tab:red"}},
                        {"op": "secondary_xaxis", "convert": "linear", "offset": -273.15, "label": "T / °C"},
                        {"op": "secondary_yaxis", "convert": {"type": "linear", "scale": 100}, "label": "percent"}]},
        {"operations": [{"op": "bars", "categories": ["alpha", "beta", "gamma-long"],
                         "series": {"A": [1, 2, 3], "B": [2, -1, 1]}, "mode": "stacked",
                         "errors": {"A": [0.1, 0.2, 0.3]}},
                        {"op": "ticks", "axis": "y", "step": 1, "minor": True, "direction": "in", "mirror": True,
                         "grid": "both"},
                        {"op": "spines", "hide": ["top"], "offset": {"left": 4}, "linewidth": 1.5}]},
        {"operations": [{"op": "bars", "data": "diff.csv", "category_col": "alloy", "value_col": "hv",
                         "group_col": "T", "horizontal": True, "colors": "viridis"}]},
        {"operations": [{"op": "pie", "values": [3, 2, 1], "labels": ["a", "b", "c"], "autopct": True,
                         "width": 0.4, "center_text": "total"}]},
        {"projection": "polar", "operations": [
            {"op": "radar", "categories": ["a", "b", "c", "d"], "series": {"s1": [1, 2, 3, 2], "s2": [2, 1, 2, 3]},
             "rlim": [0, 4]}]},
        {"projection": "3d", "operations": [
            {"op": "surface", "x": [0, 1, 2], "y": [0, 1], "z": [[0, 1, 4], [1, 2, 5]], "cmap": "coolwarm"},
            {"op": "scatter3d", "x": [0, 1], "y": [0, 1], "z": [1, 2], "color": "k"},
            {"op": "line3d", "x": [0, 2], "y": [0, 1], "z": [0, 5], "color": "r"},
            {"op": "view", "elev": 20, "azim": 30, "zlabel": "z"}]}]})
    check(tb, imgs[0], tmp_path)


def test_axes_layout_ops(tb, tmp_path):
    _, imgs = call(tb, "make_plot", {
        "output": "layout.png", "figsize": [9, 7], "sharex": True,
        "width_ratios": [2, 1], "height_ratios": [1, 1.5],
        "panels": [
            {"row": 0, "col": 0, "operations": [
                {"op": "pcolormesh", "x": [0, 1, 2, 3], "y": [0, 1, 2], "z": [[1, 2, 3, 4], [2, 3, 4, 5], [3, 4, 5, 9]],
                 "cmap": "cividis"},
                {"op": "quiver", "x": [0.5, 1.5], "y": [0.5, 1.5], "u": [1, 0], "v": [0, 1], "color": "w",
                 "key": {"U": 1, "label": "1 unit"}}]},
            {"row": 0, "col": 1, "operations": [
                {"op": "scatter", "data": "pts.json", "x_col": "x", "y_col": "y", "c_col": "s", "cmap": "viridis"},
                {"op": "annotate_points", "x": [0.1, 0.12, 0.8], "y": [0.1, 0.11, 0.8], "labels": ["p1", "p2", "p3"],
                 "scatter": True, "arrow": True}]},
            {"row": 1, "col": 0, "colspan": 2, "operations": [
                {"op": "line", "x": [0, 1, 2, 3], "y": [1, 50, 2, 60], "label": "spiky"},
                {"op": "line", "x": [0, 1, 2, 3], "y": [2, 3, 4, 3], "label": "calm"},
                {"op": "label_lines", "where": "end", "legend": False},
                {"op": "twiny", "xlabel": "top axis", "xlim": [0, 30]},
                {"op": "colorbar", "axes_list": [0, 1], "label": "shared", "location": "right"},
                {"op": "panel_labels", "position": "top left", "uppercase": True},
                {"op": "broken_yaxis", "axes": 2, "ylims": [[0, 5], [45, 65]]}]}]})
    check(tb, imgs[0], tmp_path)


def test_edit_of_edit_chain(tb, tmp_path):
    _, imgs = call(tb, "make_plot", {"output": "base.png", "panels": [{"operations": [
        {"op": "line", "data": "cp.csv", "x_col": "x(Cr)", "y_col": "cp", "label": "FCC_A1 cp", "color": "C0"},
        {"op": "line", "x": [0.1, 0.4], "y": [22, 30], "label": "guide", "color": "#ff0000"},
        {"op": "scatter", "x": [0.2, 0.3], "y": [25, 26], "label": "points"},
        {"op": "set", "xlabel": "x(Cr)", "title": "Cr-Ni", "ylabel": "cp"}]}]})
    _, e1 = call(tb, "edit_figure", {"image": imgs[0], "output": "edit1.png", "operations": [
        {"op": "style", "match": "guide", "color": "green", "linewidth": 3, "new_label": "guide line"},
        {"op": "remove", "match": "points"},
        {"op": "scatter", "x": [0.25], "y": [27], "marker": "*", "s": 80, "label": "new point"},
        {"op": "twinx", "ylabel": "other", "ylim": [0, 1]},
        {"op": "line", "axes": 1, "x": [0.1, 0.4], "y": [0.2, 0.8], "color": "k", "label": "twin"},
        {"op": "secondary_xaxis", "convert": "mole_to_mass", "label": "w(Cr)"},
        {"op": "legend", "loc": "upper left", "ncol": 2}]})
    _, e2 = call(tb, "edit_figure", {"image": e1[0], "operations": [
        {"op": "inset", "bounds": [0.6, 0.1, 0.3, 0.3], "xlim": [0.2, 0.3], "ylim": [24, 28]},
        {"op": "style", "match": "color:#008000", "linestyle": "--"},
        {"op": "text", "x": 0.05, "y": 0.9, "text": "edited twice", "coords": "axes"},
        {"op": "label_lines", "axes": 0, "match": "FCC", "where": "max"},
        {"op": "annotate", "text": "here", "xy": [0.3, 26], "arrow": False},
        {"op": "set", "axes": "Cr-Ni", "ylim": [18, 34]},
        {"op": "figure", "size": [7, 5], "suptitle": "chain"}]})
    script, got = check(tb, e2[0], tmp_path, n_cells=3)
    assert "comp_convert" in script and "find_artists(ax, 'guide')" in script
    assert got["axes"][0]["legend"] is not None and "guide line" in got["axes"][0]["legend"]


def test_raster_overlay_edit(tb, tmp_path):
    from PIL import Image
    img = Path(tb.workdir) / "foreign.png"
    Image.new("RGB", (400, 300), "white").save(img)
    _, out = call(tb, "edit_figure", {"image": "foreign.png", "calibration": {
        "pixel_box": [50, 30, 370, 260], "xlim": [0, 1], "ylim": [300, 1500]}, "operations": [
        {"op": "scatter", "x": [0.2, 0.5], "y": [800, 1200], "color": "red", "label": "mine"},
        {"op": "hline", "y": 1000, "color": "b"}]})
    check(tb, out[0], tmp_path, n_cells=1)
    _, out2 = call(tb, "edit_figure", {"image": "foreign.png", "output": "frac.png", "operations": [
        {"op": "text", "x": 0.5, "y": 0.5, "text": "centre"}]})
    script, _ = check(tb, out2[0], tmp_path / "b", n_cells=1)
    assert "background image" in script


def _run_files(script, d):
    d.mkdir(parents=True, exist_ok=True)
    (d / "script.py").write_text(script)
    r = subprocess.run([PY, "script.py"], cwd=d, capture_output=True, text=True, timeout=240,
                       env=dict(os.environ, MPLBACKEND="Agg"))
    assert r.returncode == 0, r.stderr[-4000:] + "\n" + script[-4000:]
    return d / "standalone_outputs"


def test_compose_and_raster_tools(tb, tmp_path):
    from PIL import Image
    _, a = call(tb, "make_plot", {"output": "a.png", "figsize": [3, 2.5],
                                  "panels": [{"operations": [{"op": "line", "x": [0, 1], "y": [0, 1]}]}]})
    _, b = call(tb, "make_plot", {"output": "b.png", "figsize": [4, 2.5],
                                  "panels": [{"operations": [{"op": "bar", "x": [0, 1], "height": [1, 2]}]}]})
    _, c = call(tb, "compose_images", {"images": [a[0], b[0], a[0]], "output": "comp.png", "titles": ["one", "", "three"],
                                        "suptitle": "Composite", "labels": "A", "label_position": "top-left"})
    # compose: axes geometry + images; parents (the two make_plot calls) come first
    script, cells = standalone(tb, c[0])
    assert [x.title.split(" ")[0] for x in cells] == ["make_plot", "make_plot", "compose_images"]
    got = run_script(script, tmp_path / "comp")
    want = tool_dump(c[0])
    for g, w in zip(got["axes"], want["axes"]):
        assert g["texts"] == w["texts"] and g["images"][0][0] == w["images"][0][0]
    with Image.open(c[0]) as im1, Image.open(tmp_path / "comp" / "run" / "standalone_outputs" / "comp.png") as im2:
        assert im1.size == im2.size
    # an edit on the composite continues on the composed figure
    _, ce = call(tb, "edit_figure", {"image": c[0], "operations": [{"op": "text", "x": 0.5, "y": 0.5,
                                                                   "text": "X", "coords": "axes", "axes": 1}]})
    script, _ = standalone(tb, ce[0])
    got = run_script(script, tmp_path / "comp_edit")
    assert got["axes"][1]["texts"][-1] == "X"
    # image_tools
    _, t = call(tb, "image_tools", {"image": b[0], "output": "b_small.jpg", "operations": [
        {"op": "crop", "trim": True}, {"op": "pad", "px": [5, 10, 5, 10], "color": "lightgray"},
        {"op": "resize", "scale": 0.5}, {"op": "rotate", "degrees": 90}, {"op": "border", "px": 2}]})
    out = _run_files(standalone(tb, t[0])[0], tmp_path / "imgtools")
    with Image.open(t[0]) as i1, Image.open(out / "b_small.jpg") as i2:
        assert i1.size == i2.size and i1.mode == i2.mode
        assert np.abs(np.asarray(i1, dtype=float) - np.asarray(i2, dtype=float)).mean() < 2.0
    txt, _ = call(tb, "image_tools", {"image": a[0], "operations": [
        {"op": "transparent_background"}, {"op": "grayscale"}, {"op": "crop", "box": [0.1, 0.1, 0.9, 0.9],
                                                                "units": "fraction"}]})
    outp = Path(tb.out) / "a_edited.png"
    out = _run_files(standalone(tb, outp)[0], tmp_path / "imgtools2")
    with Image.open(outp) as i1, Image.open(out / "a_edited.png") as i2:
        assert i1.size == i2.size and i1.mode == i2.mode
        assert np.array_equal(np.asarray(i1), np.asarray(i2))


def test_pdf_pptx_gif(tb, tmp_path):
    from PIL import Image
    _, a = call(tb, "make_plot", {"output": "a.png", "figsize": [3, 2.5],
                                  "panels": [{"operations": [{"op": "line", "x": [0, 1], "y": [0, 1]}]}]})
    _, b = call(tb, "make_plot", {"output": "b.png", "figsize": [4, 2.5],
                                  "panels": [{"operations": [{"op": "line", "x": [0, 1], "y": [1, 0]}]}]})
    call(tb, "images_to_pdf", {"images": [a[0], b[0]], "output": "both.pdf", "titles": ["first", ""],
                               "page_size": "a4"})
    out = _run_files(standalone(tb, Path(tb.out) / "both.pdf")[0], tmp_path / "pdf")
    assert (out / "both.pdf").read_bytes()[:4] == b"%PDF" and (out / "both.pdf").read_bytes().count(b"/Type /Page") \
        >= 2
    call(tb, "figures_to_pptx", {"output": "deck.pptx", "title_slide": {"title": "Deck"}, "slides": [
        {"title": "Two", "images": [a[0], b[0]], "caption": "cap", "notes": "n"},
        {"title": "Bullets", "bullets": ["x", "y"], "images": [a[0]]}]})
    out = _run_files(standalone(tb, Path(tb.out) / "deck.pptx")[0], tmp_path / "pptx")
    from pptx import Presentation
    want, got = Presentation(str(Path(tb.out) / "deck.pptx")), Presentation(str(out / "deck.pptx"))
    assert len(want.slides) == len(got.slides) == 3
    pics = lambda prs: [(s.left, s.top, s.width, s.height) for sl in prs.slides for s in sl.shapes  # noqa: E731
                        if s.shape_type == 13]
    assert pics(want) == pics(got)
    _, g = call(tb, "animate", {"frames": [a[0], b[0]], "output": "anim.gif", "caption_template": "{i}/{n}: {value}",
                                "duration_ms": 300})
    out = _run_files(standalone(tb, Path(tb.out) / "anim.gif")[0], tmp_path / "gif")
    with Image.open(Path(tb.out) / "anim.gif") as g1, Image.open(out / "anim.gif") as g2:
        assert g1.size == g2.size and g1.n_frames == g2.n_frames == 2


def test_value_labels_and_fallbacks(tb, tmp_path):
    from calphad_tools import codegen
    _, imgs = call(tb, "make_plot", {"output": "vl.png", "panels": [{"operations": [
        {"op": "bars", "categories": ["a", "b"], "series": {"s": [1.5, 2.25]}, "value_labels": "{:.1f}"}]}]})
    script, _ = standalone(tb, imgs[0])            # the toolkit keeps no editable copy of this one
    out = _run_files(script, tmp_path / "vl")
    assert (out / "vl.png").exists()
    # an animate sweep has no standalone form: labelled fallback cell
    cells, _ = codegen.build_cells(tb, 1)
    assert cells[0].standalone
    gen = codegen.GENERATORS["animate"]
    ctx = codegen.Ctx(tb)
    with pytest.raises(codegen.NoStandalone):
        gen({"output": "s.gif", "sweep": {"tool": "make_plot", "param": "x", "values": [1]}}, ctx)
    with pytest.raises(codegen.NoStandalone):
        codegen.GENERATORS["digitize_setup"]({"image": "x.png"}, ctx)


def test_digitized_data_is_recorded(tb, tmp_path):
    _, src = call(tb, "make_plot", {"output": "curve.png", "panels": [{"operations": [
        {"op": "line", "x": [0, 0.5, 1], "y": [0.1, 0.8, 0.4], "color": "red", "linewidth": 2},
        {"op": "set", "xlim": [0, 1], "ylim": [0, 1]}]}]})
    t, _ = call(tb, "digitize_curve", {"image": src[0], "color": "red", "series": "red",
                                       "axes": {"x": {"range": [0, 1]}, "y": {"range": [1, 0]}},
                                       "x_values": [0.1, 0.3, 0.5, 0.7, 0.9]})
    csv = re.search(r"-> (\S+\.csv)", t).group(1)
    _, imgs = call(tb, "make_plot", {"output": "redrawn.png", "panels": [{"operations": [
        {"op": "line", "data": csv, "x_col": "x", "y_col": "y", "marker": "o"}]}]})
    script, _ = check(tb, imgs[0], tmp_path, n_cells=3)
    cells = standalone(tb, imgs[0])[1]
    assert [c.title.split(" ")[0] for c in cells] == ["make_plot", "digitize_curve", "make_plot"]
    assert "digitized = pd.DataFrame(" in script and "OUT / " in cells[2].code


def test_more_ops_scatter_matrix_tri_wireframe(tb, tmp_path):
    rng = np.random.default_rng(0)
    xs, ys = rng.uniform(0, 1, 30), rng.uniform(0, 1, 30)
    _, imgs = call(tb, "make_plot", {"output": "more.png", "layout": [1, 3], "panels": [
        {"operations": [{"op": "tripcolor", "x": list(xs), "y": list(ys), "z": list(xs * ys), "max_edge": 0.5,
                         "colorbar": True},
                        {"op": "hist", "values": list(rng.normal(0.5, 0.1, 50)), "bins": [0, 0.25, 0.5, 0.75, 1],
                         "alpha": 0.3, "label": "h"}]},
        {"projection": "3d", "operations": [{"op": "wireframe", "x": [0, 1, 2], "y": [0, 1],
                                             "z": [[0, 1, 0], [1, 2, 1]], "color": "k"}]},
        {"operations": [{"op": "scatter_matrix", "data": "diff.csv", "cols": ["T", "D", "hv"],
                         "group_col": "alloy", "diag": "kde"}]}]})
    check(tb, imgs[0], tmp_path)


def test_values_in_data_units_on_a_transformed_axis(tb):
    """units: data on an axis the toolkit shows transformed (K data shown as °C) -> converted to shown units."""
    from calphad_tools import codegen
    from calphad_tools.codegen_plot import AxModel, FigModel, OpWriter
    fm = FigModel()
    fm.top.append(AxModel("ax", yt=(1.0, -273.15)))
    w = OpWriter(codegen.Ctx(tb), fm, "edit")
    w.cur_ax = "ax"
    w.apply([{"op": "scatter", "x": [0.5], "y": [1573.15], "units": "data"},
             {"op": "hline", "y": 1000},
             {"op": "set", "ylim": [1073.15, 2273.15], "units": "data"}], fm.top[0])
    code = "\n".join(w.lines)
    assert "ax.scatter([0.5], [1300.0" in code and "ax.axhline(1000.0" in code
    assert "ax.set_ylim(800.0" in code


def test_fit_data_and_parentless_edit(tb, tmp_path):
    from calphad_tools import codegen
    call(tb, "fit_data", {"data": "diff.csv", "x_col": "T", "y_col": "D", "model": "arrhenius",
                          "output_csv": "arr.csv", "band": "prediction"})
    out = _run_files(standalone(tb, Path(tb.out) / "arr.csv")[0], tmp_path / "fit")
    for name in ("arr.csv", "arr_residuals.csv"):
        pd.testing.assert_frame_equal(pd.read_csv(out / name), pd.read_csv(Path(tb.out) / name), rtol=1e-6,
                                      check_dtype=False)
    _, imgs = call(tb, "make_plot", {"output": "p.png", "panels": [{"operations": [{"op": "line", "x": [0, 1],
                                                                                    "y": [0, 1]}]}]})
    _, e = call(tb, "edit_figure", {"image": imgs[0], "operations": [{"op": "text", "x": 0, "y": 0, "text": "t"}]})
    cells, _ = codegen.build_cells(tb, e[0], include_parents=False)
    assert len(cells) == 1 and not cells[0].standalone and "earlier standalone cell" in cells[0].note
