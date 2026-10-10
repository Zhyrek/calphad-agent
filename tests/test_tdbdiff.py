"""The "TDB changes" view of staged_merge / combine_databases (calphad_tools.tdbdiff): each step's changed TDB lines,
old (red text) immediately followed by new (normal text), additions and removals alone, unchanged lines left out,
grouped ELEMENT/SPECIES, FUNCTION, TYPE_DEFINITION, then per phase (PHASE, CONSTITUENT, G, L0, L1 ..., TC, BMAGN).

Cases: the 'rename' decision (Fe-W's MU_D85 kept as MU_D85_2 in Q), a sublattice flip with a snap + rescale
(B (FE,W)0.667(W)0.333 vs A (MO)1(FE)2), a rescale without a flip, a tentative TENT_MIX addition, a removal; the
slide colours / gutters, pagination (a pair never split) and the .diff / Excel outputs.

Run:  .venv/bin/python -m pytest -q tests/test_tdbdiff.py
"""
import re
import sys
from pathlib import Path

import pytest

from calphad_tools import pptx_utils as pu
from calphad_tools import tdbstage as ts
from calphad_tools.staged_merge import add_tdb_diff_slides, staged_merge
from calphad_tools.tdbdiff import (KINDS, LEGEND, diff_tdb, perms_from_notes, read_diff, renames_from_notes,
                                   tdb_lines)
from calphad_tools.tdbparse import parse_tdb, resolve_keyword

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_compatibility import write                       # noqa: E402
from test_merge_matching import A_C15, b_c15               # noqa: E402

EX = Path(__file__).resolve().parent.parent / "examples" / "tdbs"
FENB, FEW = str(EX / "FeNb-90Hua-LB.tdb"), str(EX / "FeW-15Jac-static.tdb")
RED = pu.DIFF_RED


@pytest.fixture(scope="module")
def trio(tmp_path_factory):
    d = tmp_path_factory.mktemp("tdbdiff_trio")
    return staged_merge(FENB, FEW, str(d / "out"), decisions={"phase:MU_D85": "rename"}, ternary_temperatures_c=[],
                        plot=False)


def _step(r, title):
    return next(st for st in r["steps"] if st.title == title)


def _flat(d):
    """[(group title, change, text)] in display order."""
    return d.rows()


def _assert_lines_well_formed(d):
    for _, _, t in d.rows():
        kw = resolve_keyword(t.split()[0])
        assert kw in KINDS and t.endswith(" !") and t.count("!") == 1, t
        assert "  " not in t and "\n" not in t, t


# ------------------------------------------------------------------------------------------------ pairing
def test_rename_old_red_then_new(trio):
    """B3 -> Q: B's MU_D85 is MU_D85_2 in Q - each B3 line (removed / old) immediately followed by its Q line; A's
    own MU_D85 comes in as additions only; nothing unchanged is listed."""
    d = _step(trio, "B3 → Q").tdb
    _assert_lines_well_formed(d)
    g = next(g for g in d.groups if g.key == (3, "MU_D85_2"))
    assert g.title == "MU_D85_2 (was MU_D85 in B3)"
    assert all(e.change == "changed" for e in g.entries) and len(g.entries) == 10      # PHASE, CONSTITUENT, 8 G
    rows = [(ch, t) for grp, ch, t in _flat(d) if grp == g.title]
    assert rows[0] == ("removed", "PHASE MU_D85 % 4 1 4 2 6 !") and rows[1] == ("added", "PHASE MU_D85_2 % 4 1 4 2 6 !")
    i = rows.index(("removed", "PARAMETER G(MU_D85,W:W:W:W) 298.15 +13*GHSERWW+469300; 6000.00 N 06Slu !"))
    assert rows[i + 1] == ("added", "PARAMETER G(MU_D85_2,W:W:W:W;0) 298.15 +13*GHSERWW+469300; 6000.00 N 06Slu !")
    assert [ch for ch, _ in rows] == ["removed", "added"] * 10
    a = next(g for g in d.groups if g.key == (3, "MU_D85"))
    assert a.title == "MU_D85 (new in Q)" and all(e.change == "added" for e in a.entries)
    # unchanged lines are absent: B's Fe-W description is in B3 and Q alike
    texts = [t for _, _, t in _flat(d)]
    assert not any("L(LIQUID,FE,W;0)" in t for t in texts) and not any(t.startswith("FUNCTION GHSERFE") for t in texts)
    # B -> B1 / B1 -> B2: B's file vs the rewritten versions differ in formatting only (PAR / CONST, 6000 vs
    # 6000.00, blank-separated constituents): no lines
    assert _step(trio, "B1 → B2").tdb.empty and _step(trio, "B → B1").tdb.empty


def test_tentative_term_is_a_pure_addition(trio):
    d = _step(trio, "B3 → Q").tdb
    e = next(e for e in d.entries if e.new is not None and "L(C14_LAVES,FE:FE,NB;0)" in e.new.text)
    assert e.old is None and e.change == "added"
    assert e.new.text == "PARAMETER L(C14_LAVES,FE:FE,NB;0) 298.15 +150000; 6000.00 N TENT_MIX !"


def test_group_order_matches_the_file_layout(trio):
    d = _step(trio, "A → A1").tdb
    keys = [g.key for g in d.groups]
    assert keys == sorted(keys) and keys[0] == (0, "ELEMENT / SPECIES") and keys[1] == (1, "FUNCTION")
    phases = [k[1] for k in keys if k[0] == 3]
    assert phases == sorted(phases)
    c14 = [t for grp, _, t in _flat(d) if grp.startswith("C14_LAVES")]
    assert c14[0].startswith("CONST ") and c14[1].startswith("CONSTITUENT ")         # old line first
    assert [t.split()[1] for t in c14[2:]] == ["G(C14_LAVES,FE:FE)", "G(C14_LAVES,W:W)", "L(C14_LAVES,FE:FE,NB;0)"]


def test_flip_snap_rescale_pairs(tmp_path):
    """B (FE,W)0.667(W)0.333 normalised to A's (MO)1(FE)2 order and formula unit: every C15 line is a pair, the
    old line (B's sublattice order, per atom) right before the new one (flipped, x3)."""
    a, b = write(tmp_path, "A.tdb", A_C15), write(tmp_path, "B.tdb", b_c15("0.667 0.333", flip=True))
    An, Bn, notes = ts.normalize_pair(parse_tdb(a), parse_tdb(b))
    bn = ts.write_tdb(Bn, str(tmp_path / "Bn.tdb"))
    assert perms_from_notes(notes) == {"C15": [1, 0]}
    d = diff_tdb(b, bn, "B as given", "B", old_renames=renames_from_notes(notes), old_perms=perms_from_notes(notes))
    _assert_lines_well_formed(d)
    assert [g.title for g in d.groups] == ["C15 (sublattices reordered, new from old [1, 0])"]
    rows = [(ch, t) for _, ch, t in _flat(d)]
    assert rows == [
        ("removed", "PHASE C15 % 2 0.667 0.333 !"), ("added", "PHASE C15 % 2 1 2 !"),
        ("removed", "CONSTITUENT C15 :FE,W:W: !"), ("added", "CONSTITUENT C15 :W : FE,W : !"),
        ("removed", "PARAMETER G(C15,FE:W;0) 298.15 (GHSERWW+2*GHSERFE-30000+6*T)/3; 6000 N !"),
        ("added", "PARAMETER G(C15,W:FE;0) 298.15 3*((GHSERWW+2*GHSERFE-30000+6*T)/3); 6000.00 N !"),
        ("removed", "PARAMETER G(C15,W:W;0) 298.15 GHSERWW+5000; 6000 N !"),
        ("added", "PARAMETER G(C15,W:W;0) 298.15 3*(GHSERWW+5000); 6000.00 N !"),
        ("removed", "PARAMETER L(C15,FE,W:W;0) 298.15 -4000+T; 6000 N !"),
        ("added", "PARAMETER L(C15,W:FE,W;0) 298.15 3*(-4000+T); 6000.00 N !")]


def test_rescale_without_flip_values_change(tmp_path):
    a, b = write(tmp_path, "A.tdb", A_C15), write(tmp_path, "B.tdb", b_c15("0.333 0.667"))
    An, Bn, notes = ts.normalize_pair(parse_tdb(a), parse_tdb(b))
    bn = ts.write_tdb(Bn, str(tmp_path / "Bn.tdb"))
    d = diff_tdb(b, bn, "B as given", "B", old_perms=perms_from_notes(notes))
    assert [g.title for g in d.groups] == ["C15"] and all(e.change == "changed" for e in d.entries)
    e = next(e for e in d.entries if e.old.kind == "PARAMETER" and "W:W;0" in e.old.text)
    assert "GHSERWW+5000;" in e.old.text and "3*(GHSERWW+5000);" in e.new.text


def test_removal_unchanged_and_order(tmp_path):
    """A removed parameter is red only; formatting-only differences (keyword abbreviation, number format, blanks,
    type-code letters, a reference moved out of a trailing comment) are unchanged and left out; groups and the
    parameter order G, L0, L1, TC, BMAGN."""
    old = """ ELEMENT VA VACUUM 0 0 0 !
 ELEMENT FE BCC_A2 55.847 4489 27.28 !
 ELEMENT CR BCC_A2 51.996 4050 23.56 !
 FUNCTION GHSERFE 298.15 1225.7+124.134*T; 6000 N !
 TYPE_DEFINITION & GES A_P_D BCC_A2 MAGNETIC -1 0.4 !
 PHASE BCC_A2 %& 2 1 3 !
 CONST BCC_A2 : CR FE : VA : !
 PARAM G(BCC_A2,FE:VA;0) 298.15 GHSERFE; 6000 N ! $ 91Din
 PARAMETER L(BCC_A2,CR,FE:VA;0) 298.15 20500-9.68*T; 6000 N REF1 !
 PARAMETER TC(BCC_A2,FE:VA;0) 298.15 1043; 6000 N REF1 !
 PARAMETER G(BCC_A2,CR:VA;0) 298.15 -8856+157.48*T; 6000 N REF1 !
"""
    new = """ ELEMENT VA VACUUM 0.0 0.0 0.0 !
 ELEMENT CR BCC_A2 5.199600E+01 4.050000E+03 2.356000E+01 !
 ELEMENT FE BCC_A2 5.584700E+01 4.489000E+03 2.728000E+01 !
 FUNCTION GHSERFE 2.98150E+02 +1225.7+124.134*T; 6.00000E+03 N !
 TYPE_DEFINITION A GES AMEND_PHASE_DESCRIPTION BCC_A2 MAGNETIC -1.0 4.00000E-01 !
 PHASE BCC_A2 %A 2 1 3 !
 CONSTITUENT BCC_A2 :FE,CR : VA : !
 PARAMETER G(BCC_A2,FE:VA;0) 298.15 +GHSERFE; 6000.00 N 91Din !
 PARAMETER BMAGN(BCC_A2,FE:VA;0) 298.15 2.22; 6000 N REF1 !
 PARAMETER TC(BCC_A2,FE:VA;0) 298.15 1043; 6000 N REF1 !
 PARAMETER L(BCC_A2,CR,FE:VA;1) 298.15 -500; 6000 N REF2 !
 PARAMETER L(BCC_A2,CR,FE:VA;0) 298.15 20500 - 9.68*T; 6000 N REF2 !
 PARAMETER G(BCC_A2,CR:VA;0) 298.15 -8856+157.48*T; 6000 N REF1 !
 PARAMETER G(LIQUID,CR;0) 298.15 +24339.955-11.420225*T; 6000 N REF1 !
"""
    o, n = write(tmp_path, "old.tdb", old), write(tmp_path, "new.tdb", new)
    d = diff_tdb(o, n, "X", "Y")
    rows = [(g, ch, t) for g, ch, t in _flat(d)]
    assert rows == [
        ("BCC_A2", "removed", "PARAMETER L(BCC_A2,CR,FE:VA;0) 298.15 20500-9.68*T; 6000 N REF1 !"),
        ("BCC_A2", "added", "PARAMETER L(BCC_A2,CR,FE:VA;0) 298.15 20500-9.68*T; 6000 N REF2 !"),
        ("BCC_A2", "added", "PARAMETER L(BCC_A2,CR,FE:VA;1) 298.15 -500; 6000 N REF2 !"),
        ("BCC_A2", "added", "PARAMETER BMAGN(BCC_A2,FE:VA;0) 298.15 2.22; 6000 N REF1 !"),
        ("LIQUID", "added", "PARAMETER G(LIQUID,CR;0) 298.15 +24339.955-11.420225*T; 6000 N REF1 !")]
    # and backwards: the same lines removed (red only), nothing added for them
    back = diff_tdb(n, o, "Y", "X")
    assert [ch for _, ch, t in _flat(back) if "BMAGN" in t or "LIQUID" in t] == ["removed", "removed"]
    assert "L(BCC_A2,CR,FE:VA;1)" in back.text() and "-PARAMETER L(BCC_A2,CR,FE:VA;1)" in back.text()


# ------------------------------------------------------------------------------------------------ slides / files
def _diff_rows_of(prs):
    """Per slide: [(gutter, text, rgb of the text, cell has a solid fill)] of the line rows, and header rows."""
    out = []
    for sl in prs.slides:
        rows = []
        for sh in sl.shapes:
            if not sh.has_table:
                continue
            for r in sh.table.rows:
                c0, c1 = r.cells[0], r.cells[1]
                if c0.is_merge_origin:
                    rows.append(("band", c0.text, None, c0.fill.type))
                    continue
                run = c1.text_frame.paragraphs[0].runs[0]
                g = c0.text_frame.paragraphs[0].runs[0]
                assert g.font.color.rgb == run.font.color.rgb and run.font.name == "Courier New"
                assert run.font.size >= pu.Pt(10)
                rows.append((c0.text, c1.text, run.font.color.rgb, c1.fill.type))
        out.append(rows)
    return out


def test_slide_colours_gutters_and_legend(trio, tmp_path):
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL
    d = _step(trio, "B3 → Q").tdb
    prs = pu.new_presentation()
    n = add_tdb_diff_slides(prs, d, "B3 → Q: TDB changes", "full merge A+B")
    p = pu.save(prs, str(tmp_path / "d.pptx"))
    prs = Presentation(p)
    assert len(prs.slides) == n >= 2
    sub = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text][1]
    assert LEGEND in sub and "diffs/B3-Q.diff" in sub
    pages = _diff_rows_of(prs)
    lines = [r for pg in pages for r in pg if r[0] != "band"]
    assert lines and {r[0] for r in lines} == {"−", "+"}
    for g, text, rgb, fill in lines:
        assert rgb == (RED if g == "−" else pu.DARK), (g, text, rgb)
        assert fill == MSO_FILL.BACKGROUND                      # no tint: the slide background shows through
    bands = [r for pg in pages for r in pg if r[0] == "band"]
    assert bands and all(f == MSO_FILL.SOLID for *_, f in bands)
    assert any(t.startswith("MU_D85_2 (was MU_D85 in B3)") for _, t, _, _ in bands)
    # the rename pair, adjacent
    texts = [(g, t) for g, t, _, _ in lines]
    i = texts.index(("−", "PHASE MU_D85 % 4 1 4 2 6 !"))
    assert texts[i + 1] == ("+", "PHASE MU_D85_2 % 4 1 4 2 6 !")
    for pg in pages:                                         # every page starts with a group header
        assert pg[0][0] == "band"
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_table:
                assert (sh.top + sh.height) / 914400 < 7.45


def test_added_lines_white_on_a_dark_background(trio):
    from pptx.dml.color import RGBColor
    prs = pu.new_presentation()
    bg = prs.slide_masters[0].background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0x10, 0x18, 0x20)
    add_tdb_diff_slides(prs, _step(trio, "A1 → A2").tdb, "A1 → A2: TDB changes")
    lines = [r for pg in _diff_rows_of(prs) for r in pg if r[0] != "band"]
    assert {rgb for g, _, rgb, _ in lines if g == "+"} == {pu.WHITE}
    assert {rgb for g, _, rgb, _ in lines if g == "−"} == {pu.DIFF_RED_ON_DARK}


def test_pagination_never_splits_a_pair_and_caps(tmp_path):
    long = "+".join(f"{i}.123456*T**{i}" for i in range(1, 30))
    groups = [(f"PHASE_{k}", [[("removed", f"PARAMETER G(PHASE_{k},A:B;{j}) 298.15 {long}; 6000 N OLD !"),
                               ("added", f"PARAMETER G(PHASE_{k},A:B;{j}) 298.15 {long}+1; 6000 N NEW !")]
                              for j in range(4)]) for k in range(6)]
    prs = pu.new_presentation()
    n = pu.add_diff_slides(prs, "X → Y: TDB changes", groups, subtitle=LEGEND, max_slides=50)
    assert n > 3
    pages = _diff_rows_of(prs)
    for pg in pages:
        gut = [r[0] for r in pg if r[0] != "band"]
        assert gut == ["−", "+"] * (len(gut) // 2), gut          # (old, new) pairs, both on the same slide
        assert all(not t.rstrip().endswith("…") for _, t, _, _ in pg)
    assert sum(len(pg) for pg in pages) - sum(r[0] == "band" for pg in pages for r in pg) == 6 * 4 * 2
    # capped: the last row says how many lines are left, and where they are
    prs = pu.new_presentation()
    n = pu.add_diff_slides(prs, "X → Y", groups, max_slides=2, more=lambda k: f"… {k} more lines: see X-Y.diff")
    assert n == 2
    last = prs.slides[1].shapes[-1].table
    shown = sum(1 for sl in prs.slides for sh in sl.shapes if sh.has_table for r in sh.table.rows
                if not r.cells[0].is_merge_origin)
    assert last.cell(len(last.rows) - 1, 0).text == f"… {48 - shown} more lines: see X-Y.diff"


def test_diff_files_and_excel(trio):
    from openpyxl import load_workbook
    out = Path(trio["files"]["diffs"])
    names = sorted(p.name for p in out.glob("*.diff"))
    assert names == sorted(f"{st.old}-{st.new}.diff" for st in trio["steps"])
    d = _step(trio, "B3 → Q").tdb
    text = Path(d.path).read_text()
    assert "@@ MU_D85_2 (was MU_D85 in B3)" in text
    assert "-PHASE MU_D85 % 4 1 4 2 6 !\n+PHASE MU_D85_2 % 4 1 4 2 6 !\n" in text
    assert read_diff(d.path) == d.rows()
    assert "# no TDB changes (numerically identical)" in (out / "B1-B2.diff").read_text()
    wb = load_workbook(trio["files"]["xlsx"])
    ws = wb["TDB B3-Q"]
    rows = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=2)]
    assert rows == [tuple(x) for x in d.rows()]
    red = [r for r in ws.iter_rows(min_row=2) if r[1].value == "removed"]
    assert red and all(r[2].font.color.rgb.endswith("C00000") for r in red)


def test_lines_are_the_files_own_text():
    """VV#### and function names stay as written; one logical line from the keyword to '!'."""
    text = (" FUNCTION GHSERFE 298.15 1225.7\n   +124.134*T ; 6000 N !\n"
            " PARAMETER G(LIQUID,FE;0) 298.15 VV0001+GHSERFE+\n    12040.17-6.55843*T; 6000 N REF !\n")
    L = tdb_lines(text, is_text=True)
    assert [x.text for x in L] == ["FUNCTION GHSERFE 298.15 1225.7+124.134*T; 6000 N !",
                                   "PARAMETER G(LIQUID,FE;0) 298.15 VV0001+GHSERFE+12040.17-6.55843*T; 6000 N REF !"]
    assert re.fullmatch(r"PARAMETER .* !", L[1].text)
