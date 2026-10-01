"""Tests for the figure composition tools (compose_images, figures_to_pptx, animate, image_tools, images_to_pdf)."""
import json
import shutil
from pathlib import Path

import numpy as np
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def tb(tmp_path):
    from calphad_tools.agent.tools import ToolBox, UI

    class U(UI):
        def ask(self, q, o=None):
            return ""

        def confirm(self, t, b):
            return False
    return ToolBox(workdir=str(ROOT / "examples"), output_dir=str(tmp_path / "out"), ui=U(), python_tool="off")


def call(tb, name, args):
    text, imgs, err = tb.call(name, args)
    assert not err, text
    return text, imgs


def _plot(tb, name, title, n=1):
    x = np.linspace(0, 1, 20)
    _, imgs = call(tb, "make_plot", {"output": name, "panels": [{"operations": [
        {"op": "line", "x": x.tolist(), "y": (x ** n).tolist(), "label": f"x^{n}"},
        {"op": "set", "title": title, "xlabel": "x", "ylabel": "y"}]}]})
    return imgs[0]


@pytest.fixture
def images(tb, tmp_path):
    from PIL import Image, ImageDraw
    a = _plot(tb, "p1.png", "linear", 1)
    b = _plot(tb, "p2.png", "square", 2)
    c = _plot(tb, "p3.png", "cubic", 3)
    foreign = tmp_path / "foreign.jpg"                   # a tall foreign image with a wide white margin
    im = Image.new("RGB", (400, 700), "white")
    ImageDraw.Draw(im).rectangle([100, 100, 299, 599], fill=(30, 90, 200))
    im.save(foreign, quality=95)
    return [a, b, c, str(foreign)]


def test_compose_2x2_labels_editable_and_formats(tb, images, tmp_path):
    from PIL import Image
    from calphad_tools.figures import load_figure
    text, imgs = call(tb, "compose_images", {"images": images, "layout": [2, 2], "labels": "(a)",
                                             "suptitle": "Demo", "output": "comp.png", "formats": ["pdf", "svg"],
                                             "width_in": 7, "dpi": 150})
    out = Path(imgs[0])
    assert out.name == "comp.png" and out.with_suffix(".pdf").exists() and out.with_suffix(".svg").exists()
    assert "layout 2x2" in text and "(d): " in text
    with Image.open(out) as im:
        assert im.size[0] == 7 * 150
        w, h = im.size
    fig = load_figure(str(out))
    assert fig is not None
    panels = [a for a in fig.axes if a.get_label().startswith("panel")]
    assert len(panels) == 4
    texts = [t.get_text() for t in fig.texts]
    assert {"(a)", "(b)", "(c)", "(d)", "Demo"} <= set(texts)
    # aspect ratios kept: the displayed axes box matches each (trimmed) source image
    for ax in panels:
        arr = ax.images[0].get_array()
        bb = ax.get_position()
        disp = (bb.height * h) / (bb.width * w)
        assert abs(disp - arr.shape[0] / arr.shape[1]) < 0.03
    # the foreign image was trimmed to its blue rectangle (200x500 + pad)
    foreign_ax = [a for a in panels if "foreign" in a.get_label()][0]
    fa = foreign_ax.images[0].get_array()
    assert abs(fa.shape[0] / fa.shape[1] - 508 / 208) < 0.05
    info = json.loads(call(tb, "describe_figure", {"image": str(out)})[0])
    assert len(info["axes"]) == 4
    # svg as the requested output name also works; labels as a list, inside position
    _, imgs2 = call(tb, "compose_images", {"images": images[:3], "labels": ["i", "ii", "iii"], "output": "row.svg",
                                           "label_position": "top-right", "titles": ["A", "B", "C"]})
    assert Path(imgs2[0]).with_suffix(".svg").exists()
    fig2 = load_figure(imgs2[0])
    assert len([a for a in fig2.axes if a.get_label().startswith("panel")]) == 3
    assert "ii" in [t.get_text() for a in fig2.axes for t in a.texts]


def test_compose_errors(tb, images):
    text, _, err = tb.call("compose_images", {"images": images + ["nope1.png", "nope2.png"], "output": "x.png"})
    assert err and "2 image(s) not found" in text and "nope2.png" in text
    text, _, err = tb.call("compose_images", {"images": images, "layout": [1, 2], "output": "x.png"})
    assert err and "fewer cells" in text


def test_figures_to_pptx_new_and_append(tb, images, tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    text, _ = call(tb, "figures_to_pptx", {"output": "figs.pptx", "title_slide": {"title": "Figures", "subtitle": "s"},
                                           "slides": [
        {"title": "One", "images": images[:1], "caption": "cap one", "notes": "speaker note 1"},
        {"title": "Four", "images": images},
        {"title": "Text + fig", "images": images[1:2], "bullets": ["point a", "point b"]}]})
    deck = tb.out / "figs.pptx"
    prs = Presentation(str(deck))
    assert len(prs.slides) == 4 and "4 slide(s)" in text
    pics = [[s for s in sl.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE] for sl in prs.slides]
    assert [len(p) for p in pics] == [0, 1, 4, 1]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "speaker note 1"
    assert any("cap one" in s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame)
    for p in pics[2]:           # inside the slide, no distortion
        assert p.left >= 0 and p.top >= 0 and p.left + p.width <= prs.slide_width
        assert p.top + p.height <= prs.slide_height
        src = p.image.size
        assert abs(p.height / p.width - src[1] / src[0]) < 0.01
    # append to a deck made by pptx_utils (like plot_all_binaries decks)
    from calphad_tools import pptx_utils as pu
    base = pu.new_presentation()
    pu.add_title_slide(base, "Binaries", "existing")
    pu.add_image_slide(base, "CR-NI", images[0])
    old = str(tmp_path / "old.pptx")
    pu.save(base, old)
    text, _ = call(tb, "figures_to_pptx", {"append_to": old, "slides": [{"title": "Extra", "images": images[2:4],
                                                                         "notes": "n2"}]})
    prs2 = Presentation(old)
    assert len(prs2.slides) == 3 and "appended 1 slide" in text
    assert any("Binaries" in s.text_frame.text for s in prs2.slides[0].shapes if s.has_text_frame)
    assert len([s for s in prs2.slides[1].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1
    assert prs2.slides[2].notes_slide.notes_text_frame.text == "n2"
    # append into a new file keeps the source untouched
    call(tb, "figures_to_pptx", {"append_to": old, "output": "copy.pptx", "slides": [{"images": images[:1]}]})
    assert len(Presentation(old).slides) == 3 and len(Presentation(str(tb.out / "copy.pptx")).slides) == 4
    text, _, err = tb.call("figures_to_pptx", {"output": "bad.pptx", "slides": [{"images": images * 2}]})
    assert err and "at most 6" in text


def _gif_frames(path):
    from PIL import Image, ImageSequence
    with Image.open(path) as im:
        return im.n_frames, [np.asarray(f.convert("RGB")).copy() for f in ImageSequence.Iterator(im)]


def test_animate_frames_with_captions(tb, images):
    text, prev = call(tb, "animate", {"frames": images, "output": "anim.gif", "fps": 4,
                                      "caption_template": "frame {i}/{n}"})
    n, fr = _gif_frames(tb.out / "anim.gif")
    assert n == 4 and "4 frame(s)" in text
    assert len({f.shape for f in fr}) == 1                  # padded to a common canvas
    strip = [f[:30].astype(int) for f in fr]                # caption strip differs between frames
    assert np.abs(strip[0] - strip[1]).sum() > 0
    assert Path(prev[0]).exists()
    from PIL import Image
    with Image.open(tb.out / "anim.gif") as im:
        assert im.info.get("duration") == 250 and im.info.get("loop") == 0


@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="ffmpeg not installed")
def test_animate_mp4(tb, images):
    text, _ = call(tb, "animate", {"frames": images[:3], "output": "anim.mp4", "fps": 2})
    assert (tb.out / "anim.mp4").stat().st_size > 1000 and "MP4" in text


def test_animate_sweep_skips_failed_frame(tb):
    text, _ = call(tb, "animate", {"output": "styles.gif", "caption_template": "style = {value}", "sweep": {
        "tool": "make_plot", "param": "style", "values": ["default", "no-such-style", "ggplot"],
        "args": {"output": "sw.png", "panels": [{"operations": [{"op": "line", "x": [0, 1], "y": [0, 1]}]}]}}})
    assert "2 frame(s)" in text and "frame 2 (style=no-such-style) skipped" in text and "unknown style" in text
    n, fr = _gif_frames(tb.out / "styles.gif")
    assert n == 2 and np.abs(fr[0].astype(int) - fr[1].astype(int)).sum() > 0
    # nested parameter path
    text, _ = call(tb, "animate", {"output": "nested.gif", "sweep": {
        "tool": "make_plot", "param": "panels.0.operations.0.y", "values": [[0, 1], [1, 0]],
        "args": {"output": "sw2.png", "panels": [{"operations": [{"op": "line", "x": [0, 1], "y": [0, 0]}]}]}}})
    assert "2 frame(s)" in text


def test_animate_sweep_real_tool(tb):
    text, prev = call(tb, "animate", {"output": "gm.gif", "caption_template": "T = {value} K", "duration_ms": 500,
                                      "sweep": {"tool": "phase_property_plot", "param": "temperature",
                                                "values": [1200, 1600],
                                                "args": {"tdb": "tdbs/CrNi-06Hal.tdb", "elements": ["CR", "NI"],
                                                         "prop": "GM", "temperature_unit": "K"}}})
    n, fr = _gif_frames(tb.out / "gm.gif")
    assert n == 2 and "2 frame(s)" in text
    assert len(list((tb.out / "gm_frames").glob("frame_*.png"))) == 2


def test_animate_guards(tb, images):
    text, _, err = tb.call("animate", {"output": "x.gif", "frames": images * 16})
    assert err and "> 60" in text
    text, _, err = tb.call("animate", {"output": "x.gif", "sweep": {"tool": "run_python", "param": "code",
                                                                     "values": ["1"]}})
    assert err and "not available" in text
    text, _, err = tb.call("animate", {"output": "x.gif", "sweep": {"tool": "make_plot", "param": "style",
                                                                     "values": ["bad1", "bad2"], "args": {}}})
    assert err and "every sweep frame failed" in text


def test_image_tools_ops(tb, images, tmp_path):
    from PIL import Image
    src = images[3]                                           # 400x700 foreign JPG
    _, imgs = call(tb, "image_tools", {"image": src, "operations": [{"op": "crop", "trim": True, "pad": 0, "tolerance": 60}],
                                       "output": "trim.png"})
    assert Image.open(imgs[0]).size == (200, 500)
    _, imgs = call(tb, "image_tools", {"image": src, "operations": [
        {"op": "crop", "box": [0, 0, 0.5, 0.5], "units": "fraction"}, {"op": "resize", "width": 100}],
        "output": "cr.png"})
    assert Image.open(imgs[0]).size == (100, 175)
    _, imgs = call(tb, "image_tools", {"image": src, "operations": [{"op": "resize", "scale": 0.5},
                                                                     {"op": "pad", "px": 10, "color": "red"},
                                                                     {"op": "border", "px": 3}], "output": "pb.png"})
    im = Image.open(imgs[0]).convert("RGB")
    assert im.size == (226, 376) and im.getpixel((0, 0)) == (0, 0, 0) and im.getpixel((5, 5)) == (255, 0, 0)
    text, imgs = call(tb, "image_tools", {"image": images[0], "operations": [
        {"op": "transparent_background", "color": "white", "tolerance": 5}], "output": "tr.png"})
    im = Image.open(imgs[0])
    assert im.mode == "RGBA" and im.getpixel((2, 2))[3] == 0 and "transparent" in text
    _, imgs = call(tb, "image_tools", {"image": src, "operations": [{"op": "rotate", "degrees": 90}],
                                       "output": "rot.png"})
    assert Image.open(imgs[0]).size == (700, 400)
    for fmt in ("jpg", "tiff", "webp", "pdf"):
        text, imgs = call(tb, "image_tools", {"image": images[0], "operations": [
            {"op": "grayscale"}, {"op": "convert", "format": fmt}], "output": "conv.png"})
        assert (tb.out / f"conv.{fmt}").exists(), text
    text, _, err = tb.call("image_tools", {"image": src, "operations": [{"op": "crop", "box": [0, 0, 9999, 10]}]})
    assert err and "outside" in text
    text, _, err = tb.call("image_tools", {"image": src, "operations": [{"op": "sharpen"}]})
    assert err and "unknown op" in text


def test_images_to_pdf(tb, images):
    from pypdf import PdfReader
    text, _ = call(tb, "images_to_pdf", {"images": images, "output": "all.pdf", "titles": ["one", "two"]})
    r = PdfReader(str(tb.out / "all.pdf"))
    assert len(r.pages) == 4 and "3 vector page(s)" in text
    assert "one" in r.pages[0].extract_text()
    call(tb, "images_to_pdf", {"images": images, "output": "a4.pdf", "page_size": "a4", "titles": ["x"] * 4})
    r = PdfReader(str(tb.out / "a4.pdf"))
    assert len(r.pages) == 4
    w, h = float(r.pages[0].mediabox.width), float(r.pages[0].mediabox.height)
    assert abs(w - 11.69 * 72) < 1 and abs(h - 8.27 * 72) < 1          # landscape plot -> landscape page
    w, h = float(r.pages[3].mediabox.width), float(r.pages[3].mediabox.height)
    assert abs(w - 8.27 * 72) < 1                                         # portrait image -> portrait page
