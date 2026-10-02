"""compare_system / find_databases: exact file resolution, subsystem selection, function inlining (checked against
pycalphad), phase-correspondence candidates + phase_map, formation expressions, PowerPoint/Excel output, show_code."""
import csv
import math
import re
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pytest

ROOT = Path(__file__).resolve().parents[1]
TDBS = ROOT / "examples" / "tdbs"
HER, HER82, CHOI = "Cr-Fe_her.tdb", "CrFe-82Her.tdb", "CoCrFeNiV_Choi2019.tdb"


class _UI:
    def ask(self, q, o=None):
        return ""

    def confirm(self, t, b):
        return True

    def progress(self, msg):
        pass

    def notice(self, kind, text):
        pass


def _tb(workdir, out):
    from calphad_tools.agent.tools import ToolBox, UI

    class U(_UI, UI):
        pass
    return ToolBox(workdir=str(workdir), output_dir=str(out), ui=U(), python_tool="off")


@pytest.fixture
def tb(tmp_path):
    return _tb(ROOT / "examples", tmp_path / "out")


def _call(tb, name, args):
    text, imgs, err = tb.call(name, args)
    assert not err, text
    return text


LAVES_TDB = """
 ELEMENT VA   VACUUM   0.0 0.0 0.0 !
 ELEMENT CR   BCC_A2   51.996  4050.0  23.5429 !
 ELEMENT FE   BCC_A2   55.847  4489.0  27.2797 !
 TYPE_DEFINITION % SEQ * !
 FUNCTION GHSERCR   298.15  -8856.94+157.48*T-26.908*T*LN(T)
       +0.00189435*T**2-1.47721E-06*T**3+139250*T**(-1);
      2180.00  Y  -34869.344+344.18*T-50*T*LN(T)-2.88526E+32*T**(-9);   6000.00  N !
 FUNCTION GHSERFE   298.15  +1225.7+124.134*T-23.5143*T*LN(T)
       -0.00439752*T**2-5.8927E-08*T**3+77359*T**(-1);
      1811.00  Y  -25383.581+299.31255*T-46*T*LN(T)+2.29603E+31*T**(-9);   6000.00  N !
 PHASE BCC_A2 % 2 1 3 !
 CONSTITUENT BCC_A2 : CR,FE : VA : !
 PARAMETER G(BCC_A2,CR:VA;0) 298.15 +GHSERCR; 6000 N !
 PARAMETER G(BCC_A2,FE:VA;0) 298.15 +GHSERFE; 6000 N !
 PARAMETER G(BCC_A2,FE,CR:VA;1) 298.15 +1000; 6000 N !
 PHASE LAVES_C15 % 2 1 2 !
 CONSTITUENT LAVES_C15 : FE : CR : !
 PARAMETER G(LAVES_C15,FE:CR;0) 298.15 +GHSERFE+2*GHSERCR+15000-3*T; 6000 N !
"""


# ------------------------------------------------------------------------------------------ file resolution
@pytest.fixture
def fecr_dir(tmp_path):
    d = tmp_path / "work"
    (d / "dbs").mkdir(parents=True)
    shutil.copy(TDBS / HER, d / "dbs" / "Fe-Cr1.tdb")
    shutil.copy(TDBS / CHOI, d / "dbs" / "Fe-Cr2.tdb")
    shutil.copy(TDBS / HER82, d / "dbs" / "Fe-Cr3.tdb")
    return d


def test_exact_files_only(fecr_dir, tmp_path, monkeypatch):
    from calphad_tools import compare_system as cs
    opened = []
    real = cs.parse_tdb
    monkeypatch.setattr(cs, "parse_tdb", lambda p, *a, **k: opened.append(Path(p).name) or real(p, *a, **k))
    tb = _tb(fecr_dir, tmp_path / "out")
    text = _call(tb, "compare_system", {"tdbs": ["Fe-Cr1.tdb", "Fe-Cr3.tdb"], "elements": ["FE", "CR"]})
    assert "Fe-Cr2" not in text and "Fe-Cr2.tdb" not in opened
    assert set(opened) == {"Fe-Cr1.tdb", "Fe-Cr3.tdb"}
    assert "dbs/Fe-Cr1.tdb" in text and "dbs/Fe-Cr3.tdb" in text

    # a typo gives candidates, never a substitution
    text, _, err = tb.call("compare_system", {"tdbs": ["Fe-Cr1.tdb", "Fe-Cr4.tdb"], "elements": ["FE", "CR"]})
    assert err and "Fe-Cr4.tdb not found" in text and "Did you mean" in text
    assert "dbs/Fe-Cr1.tdb" in text and "dbs/Fe-Cr3.tdb" in text
    with pytest.raises(FileNotFoundError, match="no other file was substituted"):
        tb._find_tdb("fe_cr1.tdb")
    # missing extension: unique stem resolves; case-insensitive exact basename resolves
    assert tb._find_tdb("Fe-Cr3").endswith("Fe-Cr3.tdb")
    assert tb._find_tdb("fe-cr3.TDB").endswith("Fe-Cr3.tdb")
    # ambiguous basename -> error listing both
    (fecr_dir / "old").mkdir()
    shutil.copy(TDBS / HER, fecr_dir / "old" / "Fe-Cr1.tdb")
    with pytest.raises(FileNotFoundError, match="ambiguous"):
        tb._find_tdb("Fe-Cr1.tdb")
    assert tb._find_tdb("old/Fe-Cr1.tdb").endswith("old/Fe-Cr1.tdb")


def test_find_databases(tb, fecr_dir, tmp_path):
    text = _call(tb, "find_databases", {"elements": ["FE", "CR"]})
    listed = re.findall(r"^- (\S+)", text, re.M)
    assert sorted(listed) == sorted(f"tdbs/{f}" for f in (HER, HER82, CHOI))
    assert "CR-FE:" in text
    text = _call(tb, "find_databases", {"elements": ["fe", "ni"]})
    assert sorted(re.findall(r"^- (\S+)", text, re.M)) == ["tdbs/CoCrFeNiV_Choi2019.tdb", "tdbs/FeNi-00Ans-2SL.tdb"]
    tb2 = _tb(fecr_dir, tmp_path / "out2")
    text = _call(tb2, "find_databases", {"folder": "dbs", "pattern": "Fe-Cr[13].tdb"})
    assert re.findall(r"^- (\S+)", text, re.M) == ["dbs/Fe-Cr1.tdb", "dbs/Fe-Cr3.tdb"]


# ------------------------------------------------------------------------------------------ selection
def test_choi_subsystem_excludes_other_elements():
    from calphad_tools.compare_system import DbView
    v = DbView(str(TDBS / CHOI), "choi", "choi", ["FE", "CR"], {})
    assert v.rows
    for k in v.rows:
        assert not re.search(r"\b(NI|CO|V)\b", k.replace(",", " ").replace(":", " ").replace("(", " ")), k
    assert {"COV3", "NI2V", "NI2V7"}.isdisjoint(v.phases)
    assert {"LIQUID", "BCC_A2", "FCC_A1", "HCP_A3", "SIGMA", "HIGH_SIGMA", "M3V", "B2_BCC"} <= set(v.phases)
    assert "G(SIGMA,FE:CR:CR;0)" in v.rows and "G(M3V,FE:CR;0)" in v.rows
    assert v.phases["M3V"].constituents == [["FE"], ["CR"]]
    assert v.phases["BCC_A2"].va_only == [1]
    assert v.phases["B2_BCC"].ordered_like
    assert any(p for p in v.tdb.parameters if "NI" in p.raw_name)     # they exist in the file, just not selected
    used = v.used_functions()
    assert "GHSERFE" in used and "GHSERNI" not in used and "GHSERCO" not in used


def test_term_collection_82her():
    from calphad_tools import pwinline as pwi
    from calphad_tools.compare_system import DbView
    v = DbView(str(TDBS / HER82), "82her", "82her", ["FE", "CR"], {})
    segs = v.rows["G(FCC_A1,FE;0)"].segs          # GFCCFE = lattice stability + GHSERFE, inlined
    assert [(a, b) for a, b, _ in segs] == [(298.15, 1811.0), (1811.0, 6000.0)]
    d, other = pwi.terms(segs[0][2])
    assert other == 0
    assert d[(0, 0)] == pytest.approx(1225.7 - 1462.4)
    assert d[(1, 0)] == pytest.approx(124.134 + 8.282)
    assert d[(1, 1)] == pytest.approx(-23.5143 - 1.15)
    assert d[(2, 0)] == pytest.approx(-0.00439752 + 6.4e-4)
    assert d[(-1, 0)] == pytest.approx(77359)
    d2, _ = pwi.terms(segs[1][2])
    assert d2[(-9, 0)] == pytest.approx(2.29603e31 + 4.9251e30)
    assert "T*ln(T)" in pwi.fmt_expr(segs[0][2])


# ------------------------------------------------------------------------------------------ numbers vs pycalphad
@pytest.mark.parametrize("tdb", [HER, HER82, CHOI])
def test_inlined_values_match_pycalphad(tdb):
    from pycalphad import Database, Model, variables as v
    from calphad_tools import pwinline as pwi
    from calphad_tools.compare_system import DbView
    view = DbView(str(TDBS / tdb), "x", "x", ["FE", "CR"], {})
    db = Database(str(TDBS / tdb))
    temps = [400.0, 1000.0, 1700.0, 2500.0]

    def y(phase, occupancy):
        """Site fractions: occupancy = one species per sublattice or {species: fraction} for sublattice 0."""
        m = Model(db, ["CR", "FE", "VA"], phase)
        subs = {}
        for s, subl in enumerate(db.phases[phase].constituents):
            for sp in subl:
                if sp.name not in ("CR", "FE", "VA"):
                    continue
                want = occupancy[s] if s < len(occupancy) else "VA"
                subs[v.Y(phase, s, sp.name)] = (want.get(sp.name, 0.0) if isinstance(want, dict)
                                                else float(sp.name == want))
        return m, subs

    def model_value(m, subs, part, t):
        return float(m.models[part].xreplace({**subs, v.T: t, v.P: 101325.0}))

    checks = [("BCC_A2", ["FE"], "G(BCC_A2,FE;0)"), ("BCC_A2", ["CR"], "G(BCC_A2,CR;0)"),
              ("FCC_A1", ["FE"], "G(FCC_A1,FE;0)"), ("LIQUID", ["CR"], "G(LIQUID,CR;0)")]
    sig = "SIGMA" if "SIGMA" in db.phases else "SIGMA_D8B"
    checks.append((sig, ["FE", "CR", "CR"], f"G({sig},FE:CR:CR;0)"))
    for phase, occ, key in checks:
        m, subs = y(phase, occ)
        mine = pwi.evaluate(view.rows[key].segs, temps)[0]
        atoms = sum(float(r) for r, sp in zip(db.phases[phase].sublattices, occ + ["VA"] * 3) if sp != "VA")
        ref = [atoms * model_value(m, subs, "ref", t) for t in temps]      # Model is per mole of atoms
        assert np.allclose(mine, ref, rtol=1e-9, atol=1e-6), (tdb, key, mine, ref)
    # mixing parameter: xsmix at y=0.5/0.5 is L0/4 (odd orders vanish)
    m, subs = y("BCC_A2", [{"CR": 0.5, "FE": 0.5}])
    mine = pwi.evaluate(view.rows["G(BCC_A2,CR,FE;0)"].segs, temps)[0]
    ref = [4 * model_value(m, subs, "xsmix", t) for t in temps]
    assert np.allclose(mine, ref, rtol=1e-9), (tdb, mine, ref)


# ------------------------------------------------------------------------------------------ phase candidates
def test_phase_candidates_and_phase_map(tmp_path, tb):
    d = tmp_path / "w"
    d.mkdir()
    shutil.copy(TDBS / HER, d / HER)
    (d / "renamed.tdb").write_text((TDBS / HER).read_text().replace("BCC_A2", "BCC"))
    shutil.copy(TDBS / HER82, d / HER82)
    t = _tb(d, tmp_path / "out")
    text = _call(t, "compare_system", {"tdbs": [HER, "renamed.tdb", HER82], "elements": ["FE", "CR"]})
    cand = text.split("CANDIDATES", 1)[1]
    assert re.search(r"BCC \[renamed\] ~ BCC_A2 \[.*Cr-Fe_her.*\]: LIKELY same phase", cand), cand
    assert "{'BCC': 'BCC_A2'}" in cand
    assert re.search(r"SIGMA \[Cr-Fe_her, renamed\] ~ SIGMA_D8B \[CrFe-82Her\]: LIKELY", cand)
    assert re.search(r"^  BCC +\. +x +\.", text, re.M)            # not merged without a phase_map

    from calphad_tools.compare_system import compare_system
    paths = [str(d / HER), str(d / "renamed.tdb"), str(d / HER82)]
    cmp = compare_system(paths, ["FE", "CR"], phase_map={"renamed.tdb": {"BCC": "BCC_A2"}, "SIGMA_D8B": "SIGMA"},
                         workdir=d)
    assert "BCC" not in cmp.phases and "SIGMA_D8B" not in cmp.phases
    bcc = cmp.table["BCC_A2"]
    assert len(bcc) == 10
    for tr in bcc:
        assert tr.cells["renamed"].status == "same", tr.key          # the renamed copy merges row by row
    assert all(c.b != "BCC" for c in cmp.candidates)
    assert cmp.phase_map["renamed"] == {"BCC": "BCC_A2"}


# ------------------------------------------------------------------------------------------ expressions
def _her_sigma_formation(T):
    """Hand calculation: G(SIGMA,FE:CR:CR) - 10 G(bcc Fe) - 20 G(bcc Cr) with Hertzman-Sundman's Kaufman lattice
    stabilities (G(fcc Fe) = G(bcc Cr) = 0): 109990.5 - 142.725 T - 10*GBCCFE."""
    gbccfe = 1462.4 - 8.282 * T + 1.15 * T * math.log(T) - 6.4e-4 * T ** 2
    return 109990.5 - 142.725 * T - 10 * gbccfe


def test_formation_and_expressions():
    from calphad_tools.compare_system import compare_system
    temps = [500.0, 1000.0, 1500.0]
    cmp = compare_system([str(TDBS / HER), str(TDBS / HER82), str(TDBS / CHOI)], ["FE", "CR"], temperatures=temps,
                         phase_map={"SIGMA_D8B": "SIGMA"},
                         expressions=["G(SIGMA,FE:CR:CR;0) - 10*GHSERFE - 20*GHSERCR",
                                      {"label": "informal FeCr2", "expr": "G(FE,CR)-GHSERFE-2*GHSERCR"},
                                      {"label": "L1 written FE,CR", "expr": "L(LIQUID,FE,CR;1)"}],
                         formation=[{"phase": "SIGMA", "constituents": "FE:CR:CR", "per": "formula_unit"},
                                    {"phase": "SIGMA", "constituents": "FE:CR:CR", "per": "mole_of_atoms"}],
                         workdir=ROOT / "examples")
    ex = {e.spec.label: e for e in cmp.expressions}
    rows = {r.key: r for r in cmp.expr_rows}
    hand = np.array([_her_sigma_formation(t) for t in temps])
    fu = "ΔfG SIGMA(FE:CR:CR) per f.u."
    at = "ΔfG SIGMA(FE:CR:CR) per mol atoms"
    for lab in ("Cr-Fe_her", "CrFe-82Her"):
        assert np.allclose(rows[fu].cells[lab].values, hand, rtol=1e-10), lab
        assert np.allclose(rows[at].cells[lab].values, hand / 30, rtol=1e-10), lab
    assert "GBCCFE" in ex[fu].detail["Cr-Fe_her"] and "GHSERFE" in ex[fu].detail["CrFe-82Her"]
    assert "8×" in ex[fu].detail["CoCrFeNiV_Choi2019"] and "22×" in ex[fu].detail["CoCrFeNiV_Choi2019"]
    # explicit expression: 82Her's GHSER functions give the same numbers; her has no GHSER functions
    e = rows["G(SIGMA,FE:CR:CR;0) - 10*GHSERFE - 20*GHSERCR"]
    assert np.allclose(e.cells["CrFe-82Her"].values, hand, rtol=1e-10)
    assert ex[e.key].detail["Cr-Fe_her"].startswith("n/a in Cr-Fe_her: missing function GHSER")
    # informal form without a phase: ambiguity report lists the FE:CR end-members
    amb = ex["informal FeCr2"].ambiguity
    assert "G(M3V,FE:CR;0) [CoCrFeNiV_Choi2019]" in amb and "G(SIGMA,FE:CR:CR;0)" in amb
    assert all(v is None for v in ex["informal FeCr2"].segs.values())
    # odd-order sign: L(LIQUID,FE,CR;1) = -L(LIQUID,CR,FE;1) (Choi: G(LIQUID,CR,FE;1) = -1331)
    l1 = rows["L1 written FE,CR"]
    assert l1.cells["Cr-Fe_her"].status == "absent"
    assert np.allclose(l1.cells["CoCrFeNiV_Choi2019"].values, 1331.0)


def test_laves_formation_and_informal_expression(tmp_path):
    from calphad_tools.compare_system import compare_system
    (tmp_path / "laves.tdb").write_text(LAVES_TDB)
    temps = [400.0, 1200.0, 2000.0]
    cmp = compare_system([str(tmp_path / "laves.tdb"), str(TDBS / HER82)], ["FE", "CR"], temperatures=temps,
                         expressions=["G(LAVES_C15,FE:CR;0) - GHSERFE - 2*GHSERCR", "G(FE,CR)-GHSERFE-2*GHSERCR",
                                      "L(BCC_A2,FE,CR;1)"],
                         formation=[{"phase": "LAVES_C15", "constituents": "FE:CR", "per": "mole_of_atoms"}],
                         workdir=tmp_path)
    rows = {r.key: r for r in cmp.expr_rows}
    hand = np.array([15000 - 3 * t for t in temps])
    assert np.allclose(rows["G(LAVES_C15,FE:CR;0) - GHSERFE - 2*GHSERCR"].cells["laves"].values, hand)
    f = rows["ΔfG LAVES_C15(FE:CR) per mol atoms"]
    assert np.allclose(f.cells["laves"].values, hand / 3)
    assert f.cells["CrFe-82Her"].status == "absent"
    amb = cmp.expressions[1].ambiguity
    assert "end-members FE:CR: G(LAVES_C15,FE:CR;0) [laves]" in amb
    # written FE,CR;1 = +1000 in the file; the user's L(BCC_A2,FE,CR;1) is the same parameter -> +1000
    assert np.allclose(rows["L(BCC_A2,FE,CR;1)"].cells["laves"].values, 1000.0)
    assert np.allclose(cmp.view("laves").rows["G(BCC_A2,CR,FE;1)"].segs[0][2], -1000.0)
    assert cmp.view("laves").pyc_sign_notes == ["G(BCC_A2,FE,CR:VA;1)"]


# ------------------------------------------------------------------------------------------ outputs
def test_pptx_and_excel(tb):
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.util import Emu
    text = _call(tb, "compare_system", {"tdbs": [HER, HER82, CHOI], "elements": ["FE", "CR"],
                                        "phase_map": {"SIGMA_D8B": "SIGMA"},
                                        "formation": [{"phase": "SIGMA", "constituents": "FE:CR:CR"}],
                                        "output_pptx": "fecr.pptx"})
    # the text result leads with the differences
    assert text.splitlines()[1].startswith("DIFFERENCES: 16 parameters differ numerically; 0 written differently "
                                           "but equal (not listed as differences); 15 absent in some database")
    assert "G(LIQUID,CR;0): D1: 16933 -7.95*T | D2: 298.15–2180: 15483.015 +146.059775*T" in text
    assert "[functions expanded for comparison: GCRLIQ, GLIQCR, GPCRLIQ; max |Δ| 2.04e+05]" in text
    assert "L(BCC_A2,CR,FE:VA;0): D1,D2: 22000 - 11.375*T | D3: 20500 - 9.68*T  [max |Δ| 3.58e+03]" in text
    pptx = Path(tb.out) / "fecr.pptx"
    prs = Presentation(str(pptx))
    titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text), "")
              for s in prs.slides]
    base = [re.sub(r"\s+\(\d+/\d+\)$", "", t) for t in titles]
    assert base[0].startswith("Cr–Fe: comparison of 3 databases")
    order = ["Databases", "Phases present", "Possible phase correspondences (not merged)", "Differences",
             "All parameters", "Cr-Fe_her: parameters", "Cr-Fe_her: functions used", "CrFe-82Her: parameters",
             "CoCrFeNiV_Choi2019: parameters", "CoCrFeNiV_Choi2019: functions used", "Derived expressions",
             "Derived expressions: references used"]
    idx = [base.index(t) for t in order if t in base]
    assert len(idx) == len(order) - 1 and idx == sorted(idx), base       # no candidates left after the phase_map
    W, H = prs.slide_width, prs.slide_height
    for s in prs.slides:
        for sh in s.shapes:
            assert sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= W and sh.top + sh.height <= H, sh.name
            if sh.has_table:
                assert sum(c.width for c in sh.table.columns) <= W - Emu(914400 * 0.4)
                assert sum(r.height for r in sh.table.rows) <= Emu(914400 * 5.6)
    assert not any(t.startswith("Numerically equal") for t in base)     # no such rows: no slide
    diff_slides = [s for s, t in zip(prs.slides, base) if t == "Differences"]
    cells = []
    for s in diff_slides:
        tbl = next(sh.table for sh in s.shapes if sh.has_table)
        assert [tbl.cell(0, j).text for j in range(4)] == ["parameter", "Cr-Fe_her", "CrFe-82Her",
                                                            "CoCrFeNiV_Choi2019"]
        cells += [[tbl.cell(i, j).text for j in range(4)] for i in range(1, len(tbl.rows))]
    row = next(c for c in cells if c[0].startswith("L(BCC_A2,CR,FE:VA;0)"))
    # expression, then the reference as a grey sub-line (shorthand · first author + year when the list resolves it)
    assert row[1:] == ["22000 - 11.375*T\nPresent work", "22000 - 11.375*T\n82Her · Hertzman & Sundman 1982",
                       "20500 - 9.68*T\nREF103"]
    row = next(c for c in cells if c[0].startswith("G(LIQUID,CR;0)"))
    assert row[1].endswith("\n1967Kau · Kaufman 1967") and row[2].endswith("\n91Din · Dinsdale 1991")
    assert not any(c[0].startswith("TC(BCC_A2,FE:VA;0)") for c in cells)    # identical everywhere
    assert any(c[0] == "SIGMA" for c in cells)                               # phase band rows

    wb = load_workbook(Path(tb.out) / "fecr.xlsx")
    names = wb.sheetnames
    assert names[:4] == ["Summary", "Differences", "Numerically equal", "All parameters"]
    assert wb["Numerically equal"].max_row == 1
    assert {"DB Cr-Fe_her", "DB CrFe-82Her", "DB CoCrFeNiV_Choi2019", "Values", "Expressions", "Terms",
            "Functions inlined"} <= set(names)
    assert wb["Differences"].max_row == 1 + 31 and wb["All parameters"].max_row == 1 + 37
    head = [c.value for c in wb["All parameters"][1]]
    labs = ["Cr-Fe_her", "CrFe-82Her", "CoCrFeNiV_Choi2019"]
    assert head == (["phase", "parameter"] + labs + ["status", "note"] + [f"{x} (as written)" for x in labs]
                    + ["max |Δ|"] + [f"{x} reference" for x in labs] + [f"{x} citation" for x in labs])
    assert [c.value for c in wb["Differences"][1]] == head == [c.value for c in wb["Numerically equal"][1]]
    liq = next(r for r in wb["All parameters"].iter_rows(values_only=True) if r[1] == "G(LIQUID,CR;0)")
    assert liq[2] == "16933 -7.95*T" and liq[7:9] == ("GLIQCR", "GLIQCR")
    assert liq[11:14] == ("1967Kau", "91Din", "REF279")
    assert liq[14].startswith("Phase stability in metals and alloys, L.Kaufman")      # from the $-comment list
    assert liq[15] == "A.T. Dinsdale, Calphad, 15, 317-425(1991)." and liq[16] is None   # REF279: no list
    db = [[c.value for c in r] for r in wb["DB CrFe-82Her"].iter_rows()]
    assert db[0] == ["phase", "parameter", "expression as written", "T range", "written as", "reference", "citation"]
    assert ["SIGMA", "G(SIGMA,FE:CR:CR;0)", "10*GFCCFE + 20*GHSERCR + 109990.5 - 142.725*T", "298.15-6000",
            "G(SIGMA_D8B,FE:CR:CR)", "82Her", "S. Hertzman, B. Sundman, Calphad, 6, 67-80(1982); Cr-Fe"] in db
    assert ["FUNCTION", "name", "expression as written", "T range", None, None, None] in db
    assert any(r[1] == "GHSERFE" and r[2].startswith("298.15–1811: 1225.7 + 124.134*T") for r in db)
    rows = list(csv.DictReader(open(Path(tb.out) / "fecr_long.csv")))
    assert {"key", "database", "value_298.15K", "value_1000K", "written", "written_status", "reference",
            "reference_short", "reference_citation"} <= set(rows[0])
    r = next(r for r in rows if r["key"] == "G(SIGMA,FE:CR:CR;0)" and r["database"] == "CrFe-82Her")
    assert (r["reference"], r["reference_short"]) == ("82Her", "Hertzman & Sundman 1982")
    assert r["reference_citation"].startswith("S. Hertzman, B. Sundman")
    assert "Parameter references (" in text and "D2 27/27 (27 resolved by its reference list)" in text
    assert "Files:" in text and "fecr.xlsx" in text


def test_many_databases_layout(tmp_path):
    """8 databases: the expression tables are split into column groups (<= 4 databases per slide)."""
    from pptx import Presentation
    from calphad_tools.compare_system import run
    paths = []
    for i in range(8):
        p = tmp_path / f"db{i}.tdb"
        shutil.copy(TDBS / (HER if i % 2 else HER82), p)
        paths.append(str(p))
    r = run(paths, ["FE", "CR"], str(tmp_path / "many.pptx"), tmp_path)
    assert [len(g) for g in r["pptx_info"]["groups"]] == [4, 4]
    prs = Presentation(str(tmp_path / "many.pptx"))
    titles = [s.shapes[0].text_frame.text for s in prs.slides]
    assert any(t.startswith("Differences [D1–D4]") for t in titles)
    assert any(t.startswith("Differences [D5–D8]") for t in titles)
    for s in prs.slides:
        for sh in s.shapes:
            assert sh.left + sh.width <= prs.slide_width and sh.top + sh.height <= prs.slide_height
            if sh.has_table:
                assert len(sh.table.columns) <= 9


# ------------------------------------------------------------------------------------------ written expressions
ESPEI_TDB = """
$ written by ESPEI (pycalphad to_file style)
ELEMENT CR BCC_A2 51.996 4050.0 23.56 !
ELEMENT FE BCC_A2 55.847 4489.0 27.28 !
ELEMENT VA VACUUM 0.0 0.0 0.0 !
FUNCTION GHSERCR 298.15 -8856.94+139250*T**(-1)+0.00189435*T**(2)+157.48*T
  -1.47721E-06*T**(3)-26.908*T*LOG(T); 2180.0 Y -34869.344
  -2.88526E+32*T**(-9)+344.18*T-50*T*LOG(T); 6000.0 N !
FUNCTION GHSERFE 298.15 1225.7+77359*T**(-1)-0.00439752*T**(2)+124.134*T-5.8927E-08*T**(3)-23.5143*T*LOG(T);
  1811.0 Y -25383.581+2.29603E+31*T**(-9)+299.31255*T-46*T*LOG(T); 6000.0 N !
FUNCTION VV0000 1 -1.2139E-05; 10000 N !
FUNCTION VV0001 1 -3.5; 10000 N !
FUNCTION VV0002 1 22780.0; 10000 N !
FUNCTION VV0003 1 -14.6; 10000 N !
FUNCTION VV0004 1 VV0005*2; 10000 N !
FUNCTION VV0005 1 -1000.0; 10000 N !
FUNCTION VV0013 1 2.29997; 10000 N !
FUNCTION VV0014 1 -14190.0; 10000 N !
TYPE_DEFINITION % SEQ * !
PHASE LIQUID % 1 1.0 !
CONSTITUENT LIQUID :CR, FE: !
PHASE BCC_A2 % 2 1.0 3.0 !
CONSTITUENT BCC_A2 :CR, FE:VA: !
PHASE CRFE3 % 2 1.0 3.0 !
CONSTITUENT CRFE3 :CR:FE: !
PARAMETER G(LIQUID,CR;0) 1 1.0*GHSERCR+24339.955-11.420225*T; 10000 N !
PARAMETER G(LIQUID,FE;0) 1 1.0*GHSERFE+12040.17-6.55843*T; 10000 N !
PARAMETER L(LIQUID,CR,FE;0) 1 VV0014+T*VV0013; 10000 N !
PARAMETER L(LIQUID,CR,FE;1) 1 VV0001*T*LOG(T); 10000 N !
PARAMETER L(LIQUID,CR,FE;2) 1 VV0004-T*VV0001; 10000 N !
PARAMETER G(BCC_A2,CR:VA;0) 1 1.0*GHSERCR; 10000 N !
PARAMETER G(BCC_A2,FE:VA;0) 1 1.0*GHSERFE; 10000 N !
PARAMETER TC(BCC_A2,FE:VA;0) 1 1043; 10000 N !
PARAMETER L(BCC_A2,CR,FE:VA;1) 1 T*VV0000; 10000 N !
PARAMETER L(BCC_A2,CR,FE:VA;0) 1 VV0002+T*VV0003; 10000 N !
PARAMETER G(CRFE3,CR:FE;0) 1 3.0*GHSERFE+1.0*GHSERCR+VV0002; 10000 N !
"""


def test_vv_substitution_and_row_order(tmp_path):
    from calphad_tools.tdbparse import parse_tdb
    from calphad_tools.tdbtables import Subsystem, database_functions, database_table, tidy
    (tmp_path / "espei.tdb").write_text(ESPEI_TDB)
    sub = Subsystem(parse_tdb(str(tmp_path / "espei.tdb")), "espei", "espei.tdb", ["CR", "FE"], {})
    rows = [(ph, lab, w.text) for ph, lab, w in database_table(sub)]
    assert rows == [
        ("LIQUID", "G(LIQUID,CR;0)", "GHSERCR + 24339.955 - 11.420225*T"),
        ("LIQUID", "G(LIQUID,FE;0)", "GHSERFE + 12040.17 - 6.55843*T"),
        ("LIQUID", "L(LIQUID,CR,FE;0)", "-14190.0 + 2.29997*T"),
        ("LIQUID", "L(LIQUID,CR,FE;1)", "-3.5*T*LOG(T)"),
        ("LIQUID", "L(LIQUID,CR,FE;2)", "-1000.0*2 + 3.5*T"),               # nested VV0004 = VV0005*2
        ("BCC_A2", "G(BCC_A2,CR:VA;0)", "GHSERCR"),
        ("BCC_A2", "G(BCC_A2,FE:VA;0)", "GHSERFE"),
        ("BCC_A2", "L(BCC_A2,CR,FE:VA;0)", "22780.0 - 14.6*T"),
        ("BCC_A2", "L(BCC_A2,CR,FE:VA;1)", "-1.2139E-05*T"),
        ("BCC_A2", "TC(BCC_A2,FE:VA;0)", "1043"),
        ("CRFE3", "G(CRFE3,CR:FE;0)", "3.0*GHSERFE + GHSERCR + 22780.0"),
    ]
    fns = [n for n, _ in database_functions(sub)]
    assert fns == ["GHSERCR", "GHSERFE"] and not any(n.startswith("VV") for n in sub.used_functions())
    assert tidy("A-VV9*T+VV9", {})[0] == "A - VV9*T + VV9"      # undefined VV: kept as a name
    assert tidy("2*-3+-T")[0] == "2*-3 - T"


GHSERFE_BY_HAND = ("298.15 1225.7+77359*T**(-1)-0.00439752*T**(2)+124.134*T-5.8927E-08*T**(3)-23.5143*T*LOG(T)+5000;"
                   " 1811.0 Y -25383.581+2.29603E+31*T**(-9)+299.31255*T-46*T*LOG(T)+5000; 6000.0 N")


def _fcc_pair(tmp_path, b_param="1 GHSERFE+5000; 10000 N", b_edit=lambda t: t):
    base = ESPEI_TDB.replace("PHASE CRFE3", "PHASE FCC_A1 % 2 1 1 !\nCONSTITUENT FCC_A1 :CR,FE:VA: !\n"
                             "PARAMETER G(FCC_A1,FE:VA;0) 1 GHSERFE+5000; 10000 N !\nPHASE CRFE3")
    (tmp_path / "a.tdb").write_text(base)
    other = base.replace("G(FCC_A1,FE:VA;0) 1 GHSERFE+5000; 10000 N", f"G(FCC_A1,FE:VA;0) {b_param}")
    (tmp_path / "b.tdb").write_text(b_edit(other))
    from calphad_tools.compare_system import compare_system
    cmp = compare_system([str(tmp_path / "a.tdb"), str(tmp_path / "b.tdb")], ["CR", "FE"], workdir=tmp_path)
    return cmp, {r.key: r for r in cmp.written}


def test_written_differently_but_equal(tmp_path):
    # b: L0 typed out instead of VV handles (equal to 1e-7 relative); TC with an extra trailing digit
    cmp, rows = _fcc_pair(tmp_path, b_edit=lambda t: t.replace("VV0002+T*VV0003", "22780.002-14.6*T")
                          .replace("1043; 10000", "1043.00001; 10000"))
    l0 = rows["G(BCC_A2,CR,FE;0)"]
    assert l0.status == "equal" and l0.note == "numerically equal (written differently)"
    assert l0.shown == {"a": "22780.0 - 14.6*T", "b": "22780.002 - 14.6*T"} and not l0.expanded
    assert rows["TC(BCC_A2,FE;0)"].status == "equal"
    diff_keys = {r.key for r in cmp.written_differences()}
    assert "G(BCC_A2,CR,FE;0)" not in diff_keys and "TC(BCC_A2,FE;0)" not in diff_keys
    assert rows["G(BCC_A2,CR;0)"].status == "same"
    assert cmp.written_differences() == []


def test_function_vs_inlined_polynomial(tmp_path):
    # b writes GHSERFE + 5000 out by hand: cells expanded, numerically equal -> not a difference
    cmp, rows = _fcc_pair(tmp_path, GHSERFE_BY_HAND)
    fcc = rows["G(FCC_A1,FE;0)"]
    assert fcc.cells["a"].text == "GHSERFE + 5000" and fcc.cells["b"].text.startswith("298.15–1811: 1225.7 + ")
    assert fcc.expanded == ["GHSERFE"]
    assert fcc.note == "functions expanded for comparison: GHSERFE; numerically equal (written differently)"
    strip = lambda txt: [ln.split(": ", 1)[1] for ln in txt.split("\n")]      # noqa: E731 - drop T ranges
    assert strip(fcc.shown["a"]) == strip(fcc.shown["b"])         # term by term the same
    assert fcc.shown["b"].startswith("298.15–1811: 6225.7 +124.134*T -23.5143*T*ln(T)")
    assert fcc.shown["a"].startswith("1–1811: 6225.7 +124.134*T")   # the parameter's own range (1-10000 K)
    assert fcc.status == "equal" and cmp.written_differences() == []
    # a coefficient changed by hand: kept as a difference, still expanded
    cmp, rows = _fcc_pair(tmp_path, GHSERFE_BY_HAND.replace("+124.134*T", "+124.334*T"))
    fcc = rows["G(FCC_A1,FE;0)"]
    assert fcc.status == "differs" and fcc.kind == "value" and fcc.expanded == ["GHSERFE"]
    assert "+124.334*T" in fcc.shown["b"] and "+124.134*T" in fcc.shown["a"]
    assert fcc.cell_differs == {"b": True}
    assert [r.key for r in cmp.written_differences()] == ["G(FCC_A1,FE;0)"]


def test_shared_function_defined_differently(tmp_path):
    # same "GHSERFE + 5000" in both, GHSERFE differs by 0.1 J/mol: expanded, numerically equal -> not a difference
    cmp, rows = _fcc_pair(tmp_path, b_edit=lambda t: t.replace("1225.7+77359", "1225.8+77359"))
    fcc = rows["G(FCC_A1,FE;0)"]
    assert fcc.expanded == ["GHSERFE"] and fcc.status == "equal"
    assert fcc.shown["a"].startswith("1–1811: 6225.7 ") and fcc.shown["b"].startswith("1–1811: 6225.8 ")
    # ... by 100 J/mol: a real difference
    cmp, rows = _fcc_pair(tmp_path, b_edit=lambda t: t.replace("1225.7+77359", "1325.7+77359"))
    fcc = rows["G(FCC_A1,FE;0)"]
    assert fcc.status == "differs" and fcc.note.startswith("functions expanded for comparison: GHSERFE; max |Δ| 100")
    assert rows["G(BCC_A2,FE;0)"].status == "differs"          # GHSERFE itself


# ------------------------------------------------------------------------------------------ references
REF_TDB = """
$ Test database
$ [Smi24] J. Smith, A. Jones, B. Brown,
$         Calphad, 80 (2024) 102345; Cr-Fe
$
$ REF7: K. Frisk, Metall. Trans. A, 21A, 2477-88(1990)
ELEMENT CR BCC_A2 51.996 4050.0 23.56 !
ELEMENT FE BCC_A2 55.847 4489.0 27.28 !
ELEMENT VA VACUUM 0.0 0.0 0.0 !
FUNCTION GHSERCR 298.15 -8856.94+157.48*T; 6000 N 91Din !
PHASE LIQUID % 1 1.0 !
CONSTITUENT LIQUID :CR,FE: !
PHASE BCC_A2 % 2 1 3 !
CONSTITUENT BCC_A2 :CR,FE:VA: !
PARAMETER G(LIQUID,CR;0) 298.15 +GHSERCR+24339.955-11.420225*T; 6000 N 91Din !
PARAMETER G(LIQUID,FE;0) 298.15 +12040.17-6.55843*T; 6000 N !
PARAMETER L(LIQUID,CR,FE;0) 298.15 -14190+2.29997*T; 6000 N
   REF283 !
PARAMETER L(LIQUID,CR,FE;1) 298.15 -3.5*T;,,N Smi24 !
PARAMETER G(BCC_A2,CR:VA;0) 298.15 +GHSERCR; 6000 N ! $1967Kau
PARAMETER G(BCC_A2,FE:VA;0) 298.15 +1225.7; 6000 N 5 !
PARAMETER L(BCC_A2,CR,FE:VA;0) 298.15 +20500-9.68*T; 6000 N ref7 !
PARAMETER L(BCC_A2,CR,FE:VA;1) 298.15 +1000; 6000 N 90DIN !
LIST_OF_REFERENCES
 NUMBER  SOURCE
  91Din  'A.T. Dinsdale, Calphad, 15, 317-425(1991).'
  REF283 'N. Dupin, I. Ansara, B. Sundman, Calphad, 25,
          279-98(2001); Al-Cr-Ni'
  90DIN  90Din 'Alan Dinsdale, SGTE Data for Pure Elements, NPL Report
          DMA(A)195, Rev. August 1990'
!
ADD_REFERENCES
  5  'SGTE unary database V.4.1, 1999.'
!
"""


def test_reference_parsing(tmp_path):
    from calphad_tools.tdbparse import parse_reference_command, parse_tdb, short_citation
    db = parse_tdb(REF_TDB, is_text=True, name="refs.tdb")
    assert not db.warnings
    refs = {p.raw_name: p.reference for p in db.parameters}
    assert refs == {"G(LIQUID,CR;0)": "91Din", "G(LIQUID,FE;0)": "", "L(LIQUID,CR,FE;0)": "REF283",
                    "L(LIQUID,CR,FE;1)": "Smi24", "G(BCC_A2,CR:VA;0)": "1967Kau", "G(BCC_A2,FE:VA;0)": "5",
                    "L(BCC_A2,CR,FE:VA;0)": "ref7", "L(BCC_A2,CR,FE:VA;1)": "90DIN"}
    cite = {n: db.citation(r) for n, r in refs.items()}
    assert cite["G(LIQUID,CR;0)"] == "A.T. Dinsdale, Calphad, 15, 317-425(1991)."
    assert cite["L(LIQUID,CR,FE;0)"] == "N. Dupin, I. Ansara, B. Sundman, Calphad, 25, 279-98(2001); Al-Cr-Ni"
    assert cite["L(BCC_A2,CR,FE:VA;1)"].startswith("Alan Dinsdale, SGTE Data")          # alias line
    assert db.citation("90din") == cite["L(BCC_A2,CR,FE:VA;1)"]                        # keys case-insensitive
    assert cite["G(BCC_A2,FE:VA;0)"] == "SGTE unary database V.4.1, 1999."            # ADD_REFERENCES, number
    assert cite["L(LIQUID,CR,FE;1)"] == "J. Smith, A. Jones, B. Brown, Calphad, 80 (2024) 102345; Cr-Fe"  # $ [key]
    assert cite["L(BCC_A2,CR,FE:VA;0)"] == "K. Frisk, Metall. Trans. A, 21A, 2477-88(1990)"         # $ KEY: text
    assert cite["G(BCC_A2,CR:VA;0)"] == "" and cite["G(LIQUID,FE;0)"] == ""            # not in any list / none
    assert db.functions["GHSERCR"].expr.reference == "91Din"
    # write-back keeps the reference after the last range (a $-comment reference is not written into the command)
    p = next(p for p in db.parameters if p.raw_name == "L(LIQUID,CR,FE;0)")
    assert p.to_tdb().rstrip(" !").endswith("N REF283")
    assert parse_tdb(REF_TDB.replace("REF283 !", "!"), is_text=True).parameters[2].reference == ""
    # short citations: first author + year, "&" for two authors, "et al." for more
    assert [short_citation(cite[n]) for n in ("G(LIQUID,CR;0)", "L(LIQUID,CR,FE;0)", "L(LIQUID,CR,FE;1)",
                                               "L(BCC_A2,CR,FE:VA;1)", "G(BCC_A2,FE:VA;0)")] == \
        ["Dinsdale 1991", "Dupin et al. 2001", "Smith et al. 2024", "Dinsdale 1990", "SGTE unary database V.4.1 1999"]
    assert short_citation("S. Hertzman, B. Sundman, Calphad, 6, 67-80(1982); Cr-Fe") == "Hertzman & Sundman 1982"
    assert short_citation("C.P.Chin. S.Hertzman, B.Sundman, Internal Report KTH, Sweden.") == "Chin et al."
    assert short_citation("Phase stability in metals and alloys, L.Kaufman, eds. P.S.Rudman, 1967.") == \
        "Kaufman 1967"
    assert short_citation("J.Agren, Met. Trans., 10A (1979) 1847-1852.") == "Agren 1979"     # not the pages
    assert parse_reference_command("LIST_OF_REFERENCES\n R1 first source\n R2 second, 2001") == \
        {"R1": "first source", "R2": "second, 2001"}
    # the example databases: Thermo-Calc LIST-OF-REFERENCE and $-comment lists
    her82 = parse_tdb(str(TDBS / HER82))
    assert her82.citation("82Her") == "S. Hertzman, B. Sundman, Calphad, 6, 67-80(1982); Cr-Fe"
    her = parse_tdb(str(TDBS / HER))
    assert {p.reference for p in her.parameters} == {"1967Kau", "1979Agr", "Present work", "Chin"}
    assert her.citation("1967Kau").startswith("Phase stability in metals and alloys, L.Kaufman")
    assert her.citation("1979Agr") == ""                       # the file's list says [1978Agr]: not guessed


ESPEI_REFS = ESPEI_TDB.replace("$ written by ESPEI (pycalphad to_file style)",
                               "$ ESPEI fit, references added by hand\n$ [Smi24] J. Smith, A. Jones, B. Brown, "
                               "Calphad, 80 (2024) 102345") + """
LIST_OF_REFERENCES
 NUMBER  SOURCE
  REF1  'B. Sundman, ESPEI re-fit of the Cr-Fe liquid, unpublished (2025)'
  REF2  'A.T. Dinsdale, Calphad, 15, 317-425(1991).'
!
"""


def _espei_pair(tmp_path):
    """a: ESPEI output without references; b: the same fit with references, some VV values typed out."""
    (tmp_path / "a.tdb").write_text(ESPEI_TDB)
    b = (ESPEI_REFS
         .replace("L(BCC_A2,CR,FE:VA;0) 1 VV0002+T*VV0003; 10000 N !",
                  "L(BCC_A2,CR,FE:VA;0) 1 22780.002-14.6*T; 10000 N REF1 !")          # equal within tolerance
         .replace("TC(BCC_A2,FE:VA;0) 1 1043; 10000 N !", "TC(BCC_A2,FE:VA;0) 1 1043.00001; 10000 N REF2 !")
         .replace("G(LIQUID,CR;0) 1 1.0*GHSERCR+24339.955-11.420225*T; 10000 N !",
                  "G(LIQUID,CR;0) 1 1.0*GHSERCR+24339.955-11.420225*T; 10000 N REF2 !")  # only the reference
         .replace("L(LIQUID,CR,FE;0) 1 VV0014+T*VV0013; 10000 N !",
                  "L(LIQUID,CR,FE;0) 1 -14000+2.29997*T; 10000 N Smi24 !"))            # a real difference
    (tmp_path / "b.tdb").write_text(b)
    return [str(tmp_path / "a.tdb"), str(tmp_path / "b.tdb")]


def test_reference_columns_and_equal_table(tmp_path):
    from openpyxl import load_workbook
    from pptx import Presentation
    from calphad_tools.compare_system import EQUAL_TITLE, run, summary_text
    r = run(_espei_pair(tmp_path), ["CR", "FE"], str(tmp_path / "espei.pptx"), tmp_path)
    cmp = r["comparison"]
    rows = {x.key: x for x in cmp.written}
    l0, tc, gl, ll = (rows[k] for k in ("G(BCC_A2,CR,FE;0)", "TC(BCC_A2,FE;0)", "G(LIQUID,CR;0)", "G(LIQUID,CR,FE;0)"))
    assert (l0.status, tc.status, gl.status, ll.status) == ("equal", "equal", "same", "differs")
    assert l0.refs == {"a": "", "b": "REF1"} and l0.citations["b"].startswith("B. Sundman, ESPEI re-fit")
    assert l0.ref_text("a") == "" and l0.ref_text("b") == "REF1 · Sundman 2025"
    assert gl.ref_text("b") == "REF2 · Dinsdale 1991" and ll.ref_text("b") == "Smi24 · Smith et al. 2024"
    assert l0.max_diff == pytest.approx(0.002, rel=1e-6) and tc.max_diff == pytest.approx(1e-5, rel=1e-3)
    # the Differences table: never a row equal within tolerance (also not when only the reference differs)
    assert [x.key for x in cmp.written_differences()] == ["G(LIQUID,CR,FE;0)"]
    assert [x.key for x in cmp.written_equal()] == ["G(BCC_A2,CR,FE;0)", "TC(BCC_A2,FE;0)"]
    text = summary_text(cmp, r["files"])
    assert "1 parameters differ numerically; 2 written differently but equal" in text
    assert "TC(BCC_A2,FE:VA;0)" not in text and "L(BCC_A2,CR,FE:VA;0)" not in text      # counted, not listed
    assert "D1 0/11, D2 4/11 (4 resolved by its reference list)" in text
    # Excel: the equal table follows Differences, with max |Δ| and per-database reference / citation columns
    wb = load_workbook(r["files"]["xlsx"])
    assert wb.sheetnames[:4] == ["Summary", "Differences", "Numerically equal", "All parameters"]
    head = [c.value for c in wb["Numerically equal"][1]]
    assert head[-5:] == ["max |Δ|", "a reference", "b reference", "a citation", "b citation"]
    eq = [dict(zip(head, x)) for x in wb["Numerically equal"].iter_rows(min_row=2, values_only=True)]
    assert [(x["parameter"], x["a"], x["b"], x["b reference"]) for x in eq] == [
        ("L(BCC_A2,CR,FE:VA;0)", "22780.0 - 14.6*T", "22780.002 - 14.6*T", "REF1"),
        ("TC(BCC_A2,FE:VA;0)", "1043", "1043.00001", "REF2")]
    assert eq[0]["max |Δ|"] == pytest.approx(0.002, rel=1e-6) and eq[0]["a reference"] is None
    assert eq[1]["b citation"] == "A.T. Dinsdale, Calphad, 15, 317-425(1991)."
    dif = [x for x in wb["Differences"].iter_rows(min_row=2, values_only=True)]
    assert [x[1] for x in dif] == ["L(LIQUID,CR,FE;0)"]
    db = {x[1]: x for x in wb["DB b"].iter_rows(values_only=True)}
    assert db["G(LIQUID,CR;0)"][5:] == ("REF2", "A.T. Dinsdale, Calphad, 15, 317-425(1991).")
    assert db["G(LIQUID,FE;0)"][5:] == (None, None)
    rows_csv = list(csv.DictReader(open(r["files"]["csv"])))
    got = {(x["key"], x["database"]): (x["reference"], x["reference_short"]) for x in rows_csv}
    assert got[("G(LIQUID,CR,FE;0)", "b")] == ("Smi24", "Smith et al. 2024")
    assert got[("G(LIQUID,CR,FE;0)", "a")] == ("", "")
    # slides: Differences, then the equal table (max |Δ| under the parameter, references under the cells)
    prs = Presentation(r["files"]["pptx"])
    titles = [s.shapes[0].text_frame.text for s in prs.slides]
    i_d, i_e, i_a = (titles.index(t) for t in ("Differences", EQUAL_TITLE, "All parameters"))
    assert i_d < i_e < i_a
    tbl = next(sh.table for sh in prs.slides[i_e].shapes if sh.has_table)
    cells = [[tbl.cell(i, j).text for j in range(3)] for i in range(1, len(tbl.rows))]
    assert cells[0][0] == "BCC_A2"                                                           # phase band
    assert cells[1] == ["L(BCC_A2,CR,FE:VA;0)\nmax |Δ| 0.002 J/mol", "22780.0 - 14.6*T",
                        "22780.002 - 14.6*T\nREF1 · Sundman 2025"]
    assert cells[2] == ["TC(BCC_A2,FE:VA;0)\nmax |Δ| 1e-05 K", "1043", "1043.00001\nREF2 · Dinsdale 1991"]
    tbl = next(sh.table for sh in prs.slides[i_d].shapes if sh.has_table)
    assert len(tbl.rows) == 3 and tbl.cell(2, 0).text.startswith("L(LIQUID,CR,FE;0)")
    per_db = next(s for s, t in zip(prs.slides, titles) if t.startswith("b: parameters"))
    tbl = next(sh.table for sh in per_db.shapes if sh.has_table)
    assert tbl.cell(0, 2).text == "reference"
    refs = {tbl.cell(i, 0).text.split("\n")[0]: tbl.cell(i, 2).text for i in range(1, len(tbl.rows))}
    assert refs["G(LIQUID,CR;0)"] == "REF2\nDinsdale 1991" and refs["G(LIQUID,FE;0)"] == ""
    per_db = next(s for s, t in zip(prs.slides, titles) if t.startswith("a: parameters"))
    assert len(next(sh.table for sh in per_db.shapes if sh.has_table).columns) == 2     # no references at all


# ------------------------------------------------------------------------------------------ show_code
def test_show_code_reproduces_numbers(tb, tmp_path):
    args = {"tdbs": [HER, HER82, CHOI], "elements": ["FE", "CR"], "phase_map": {"SIGMA_D8B": "SIGMA"},
            "temperatures": [300, 1000, 1500],
            "formation": [{"phase": "SIGMA", "constituents": "FE:CR:CR", "per": "mole_of_atoms"}],
            "expressions": [{"label": "L0 bcc", "expr": "L(BCC_A2,FE,CR:VA;0)"}], "output_pptx": "sc.pptx"}
    _call(tb, "compare_system", args)
    text = _call(tb, "show_code", {"target": "1"})
    assert "1 standalone" in text
    script = Path(re.search(r"saved (\S+\.py)", text).group(1))
    code = script.read_text()
    assert "calphad_tools" not in code
    (tmp_path / "gen.py").write_text(code)
    r = subprocess.run([sys.executable, "-W", "ignore", "gen.py"], cwd=tmp_path, capture_output=True, text=True,
                       timeout=600)
    assert r.returncode == 0, r.stderr[-3000:]
    tool = {(row["key"], row["database"]): row for row in csv.DictReader(open(Path(tb.out) / "sc_long.csv"))}
    mine = {(row["key"], row["database"]): row
            for row in csv.DictReader(open(tmp_path / "standalone_outputs" / "sc_long.csv"))}
    have = {k for k, v in tool.items() if v["value_1000K"]}
    assert have and have <= set(mine), sorted(have - set(mine))[:5]
    n = 0
    for k in have:
        for col in ("value_300K", "value_1000K", "value_1500K"):
            a, b = float(tool[k][col]), float(mine[k][col])
            assert math.isclose(a, b, rel_tol=1e-8, abs_tol=1e-6), (k, col, a, b)
            n += 1
    assert n > 200
    written = {(row["key"], lab): row[lab] for row in csv.DictReader(open(tmp_path / "standalone_outputs" /
                                                                           "sc_written.csv"))
               for lab in ("Cr-Fe_her", "CrFe-82Her", "CoCrFeNiV_Choi2019")}
    for (key, lab), row in tool.items():
        if row["kind"] == "parameter" and row["written"]:
            assert written[(key, lab)].replace("\n", " | ") == row["written"], (key, lab)
    status = {row["key"]: row["status"] for row in csv.DictReader(open(tmp_path / "standalone_outputs" /
                                                                       "sc_written.csv"))}
    assert all(status[k] == row["written_status"] for (k, _), row in tool.items() if row["written_status"])
    # references: key as written and the full citation, per database
    sc_rows = list(csv.DictReader(open(tmp_path / "standalone_outputs" / "sc_written.csv")))
    refs = {(row["key"], lab): (row[f"{lab} reference"], row[f"{lab} citation"]) for row in sc_rows
            for lab in ("Cr-Fe_her", "CrFe-82Her", "CoCrFeNiV_Choi2019")}
    n = 0
    for (key, lab), row in tool.items():
        if row["kind"] == "parameter" and row["written"] and row["written"] != "—":
            assert refs[(key, lab)] == (row["reference"], row["reference_citation"]), (key, lab)
            n += bool(row["reference_citation"])
    assert n > 20
    assert "16 parameters differ numerically, 0 written differently but equal, 15 absent" in r.stdout
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "standalone_outputs" / "sc.xlsx")
    assert {"Differences", "All parameters", "DB Cr-Fe_her", "Values"} <= set(wb.sheetnames)
    assert (tmp_path / "standalone_outputs" / "sc.pptx").exists()


def test_show_code_equal_table_and_references(tmp_path):
    """The standalone script rebuilds the numerically-equal table and the reference columns (ESPEI-style pair)."""
    from openpyxl import load_workbook
    from pptx import Presentation
    work = tmp_path / "w"
    work.mkdir()
    _espei_pair(work)
    tb = _tb(work, tmp_path / "out")
    _call(tb, "compare_system", {"tdbs": ["a.tdb", "b.tdb"], "elements": ["CR", "FE"], "output_pptx": "ab.pptx"})
    text = _call(tb, "show_code", {"target": "1"})
    script = Path(re.search(r"saved (\S+\.py)", text).group(1))
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "gen.py").write_text(script.read_text())
    r = subprocess.run([sys.executable, "-W", "ignore", "gen.py"], cwd=run_dir, capture_output=True, text=True,
                       timeout=600)
    assert r.returncode == 0, r.stderr[-3000:]
    assert "1 parameters differ numerically, 2 written differently but equal, 0 absent" in r.stdout
    out = run_dir / "standalone_outputs"
    mine = {row["key"]: row for row in csv.DictReader(open(out / "ab_written.csv"))}
    tool = load_workbook(Path(tb.out) / "ab.xlsx")
    wb = load_workbook(out / "ab.xlsx")
    assert wb.sheetnames[:3] == ["Differences", "Numerically equal", "All parameters"]
    for sheet in ("Differences", "Numerically equal", "All parameters"):
        a, b = ([dict(zip([c.value for c in w[sheet][1]], x)) for x in w[sheet].iter_rows(min_row=2, values_only=True)]
                for w in (tool, wb))
        cols = ["parameter", "a", "b", "status", "a reference", "b reference", "a citation", "b citation"]
        assert [[x[c] or "" for c in cols] for x in a] == [[x[c] or "" for c in cols] for x in b], sheet
        for x, y in zip(a, b):
            assert (x["max |Δ|"] is None and (y["max |Δ|"] is None or math.isnan(y["max |Δ|"]))) or \
                math.isclose(x["max |Δ|"], y["max |Δ|"], rel_tol=1e-9)
    assert [k for k, row in mine.items() if row["status"] == "equal"] == ["G(BCC_A2,CR,FE;0)", "TC(BCC_A2,FE;0)"]
    assert mine["G(LIQUID,CR;0)"]["b reference"] == "REF2" and mine["G(LIQUID,CR;0)"]["status"] == "same"
    prs = Presentation(str(out / "ab.pptx"))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles.index("Differences") < titles.index("Numerically equal within tolerance (written differently)") \
        < titles.index("All parameters")
    eq = prs.slides[titles.index("Numerically equal within tolerance (written differently)")]
    tbl = next(sh.table for sh in eq.shapes if sh.has_table)
    assert tbl.cell(1, 0).text == "BCC_A2: L(BCC_A2,CR,FE:VA;0)\nmax |Δ| 0.002"
    assert tbl.cell(1, 2).text == "22780.002 - 14.6*T\n[REF1 · Sundman 2025]"
    per_db = prs.slides[titles.index("b: parameters")]
    tbl = next(sh.table for sh in per_db.shapes if sh.has_table)
    assert [tbl.cell(0, j).text for j in range(4)] == ["phase", "parameter", "expression", "reference"]
    assert any(tbl.cell(i, 3).text == "Smi24 · Smith et al. 2024" for i in range(1, len(tbl.rows)))


def test_folder_input_and_order_disorder(fecr_dir, tmp_path):
    t = _tb(fecr_dir, tmp_path / "out")
    text = _call(t, "compare_system", {"folder": "dbs", "elements": ["FE", "CR"], "include_identical": False})
    assert all(f"dbs/Fe-Cr{i}.tdb" in text for i in (1, 2, 3))
    from calphad_tools.compare_system import compare_system, summary_text
    cmp = compare_system([str(TDBS / "AlCoCrNi_Liu2016.tdb"), str(TDBS / "CrNi-06Hal.tdb")], ["CR", "NI"],
                         include_identical=False, workdir=ROOT)
    liu = cmp.view("AlCoCrNi_Liu2016")
    assert liu.phases["BCC_B2"].dis_part == "BCC_A2" and liu.phases["BCC_A2"].disordered_of == "BCC_B2"
    assert not any(re.search(r"\bAL\b", k.replace(",", " ").replace(":", " ")) for k in liu.rows)
    assert not any(r.all_same for r in cmp.all_rows())
    txt = summary_text(cmp, {})
    assert re.search(r"Order/disorder \(DIS_PART\): .*BCC_B2->BCC_A2", txt)
